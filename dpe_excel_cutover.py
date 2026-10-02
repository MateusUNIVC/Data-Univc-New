from __future__ import annotations

import json
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from io import BytesIO
from typing import Any, Mapping

from openpyxl import load_workbook

from dpe_excel_modern import build_dpe_operational_workbook
from dpe_excel_official import build_dpe_excel_official_artifact_from_payload, build_dpe_snapshot_context
from dpe_excel_parity import DPEParityReport, audit_dpe_payload_parity
from release_info import APP_VERSION, SCHEMA_VERSION

PARITY_REPORT_VERSION = 1


@dataclass(frozen=True, slots=True)
class WorkbookEvidence:
    engine: str
    build_ok: bool
    bytes_size: int = 0
    sheets: int = 0
    formulas: int = 0
    charts: int = 0
    error: str | None = None


@dataclass(frozen=True, slots=True)
class DPECutoverReport:
    source_kind: str
    generated_at: str
    system_version: str
    schema_version: str
    payload_hash: str | None
    semantic_report: DPEParityReport
    legacy_workbook: WorkbookEvidence | None
    new_workbook: WorkbookEvidence | None
    new_release_allowed: bool
    source_quality_ok: bool
    source_quality_notes: tuple[str, ...] = ()
    report_version: int = PARITY_REPORT_VERSION
    directorate: str = "DPE"

    @property
    def failures(self):
        return self.semantic_report.failures

    @property
    def cutover_status(self) -> str:
        builds_ok = bool(
            self.legacy_workbook and self.legacy_workbook.build_ok
            and self.new_workbook and self.new_workbook.build_ok
        )
        if self.failures or not builds_ok or not self.new_release_allowed or not self.source_quality_ok:
            return "BLOCKED"
        return "READY" if self.source_kind == "production" else "CANDIDATE_PASS"

    def to_dict(self) -> dict[str, Any]:
        return {
            "report_version": self.report_version,
            "directorate": self.directorate,
            "source_kind": self.source_kind,
            "generated_at": self.generated_at,
            "system_version": self.system_version,
            "schema_version": self.schema_version,
            "payload_hash": self.payload_hash,
            "semantic": self.semantic_report.as_dict(),
            "legacy_workbook": asdict(self.legacy_workbook) if self.legacy_workbook else None,
            "new_workbook": asdict(self.new_workbook) if self.new_workbook else None,
            "new_release_allowed": self.new_release_allowed,
            "source_quality_ok": self.source_quality_ok,
            "source_quality_notes": list(self.source_quality_notes),
            "summary": {
                "cases": len(self.semantic_report.cases),
                "passed": len(self.semantic_report.cases) - len(self.failures),
                "failures": len(self.failures),
                "semantic_passed": not self.failures,
                "source_quality_ok": self.source_quality_ok,
                "new_release_allowed": self.new_release_allowed,
                "cutover_status": self.cutover_status,
            },
            "cutover_status": self.cutover_status,
        }

    def to_json(self) -> str:
        return json.dumps(self.to_dict(), ensure_ascii=False, indent=2)


def _evidence(engine: str, buffer: BytesIO) -> WorkbookEvidence:
    try:
        raw = buffer.getvalue()
        wb = load_workbook(BytesIO(raw), read_only=False, data_only=False, keep_links=True)
        formulas = sum(1 for ws in wb.worksheets for row in ws.iter_rows() for cell in row if cell.data_type == "f")
        charts = sum(len(getattr(ws, "_charts", ())) for ws in wb.worksheets)
        sheets = len(wb.sheetnames)
        wb.close()
        return WorkbookEvidence(engine, True, len(raw), sheets, formulas, charts)
    except Exception as exc:  # pragma: no cover
        return WorkbookEvidence(engine, False, error=f"{type(exc).__name__}: {exc}")


def audit_dpe_cutover_readiness(
    payload: Mapping[str, Any],
    *,
    source_kind: str = "fixture",
    generated_by: str | None = None,
    build_workbooks: bool = True,
) -> DPECutoverReport:
    normalized_source = str(source_kind or "fixture").strip().lower()
    semantic = audit_dpe_payload_parity(payload, source_kind=normalized_source)
    context = build_dpe_snapshot_context(payload, generated_by=generated_by)

    quality_notes: list[str] = []
    if not payload.get("selected_period"):
        quality_notes.append("Nenhuma competência DPE foi selecionada para a exportação.")
    source_quality_ok = not quality_notes

    legacy_evidence: WorkbookEvidence | None = None
    new_evidence: WorkbookEvidence | None = None
    new_release_allowed = False
    if build_workbooks:
        try:
            legacy_evidence = _evidence("dpe_modern", build_dpe_operational_workbook(payload))
        except Exception as exc:  # pragma: no cover
            legacy_evidence = WorkbookEvidence("dpe_modern", False, error=f"{type(exc).__name__}: {exc}")
        try:
            artifact = build_dpe_excel_official_artifact_from_payload(payload, generated_by=generated_by)
            new_release_allowed = artifact.audit.release_allowed
            new_evidence = _evidence("excel_official", artifact.to_bytes())
        except Exception as exc:  # pragma: no cover
            new_evidence = WorkbookEvidence("excel_official", False, error=f"{type(exc).__name__}: {exc}")

    return DPECutoverReport(
        source_kind=normalized_source,
        generated_at=datetime.now(timezone.utc).isoformat(),
        system_version=APP_VERSION,
        schema_version=str(SCHEMA_VERSION),
        payload_hash=context.payload_hash,
        semantic_report=semantic,
        legacy_workbook=legacy_evidence,
        new_workbook=new_evidence,
        new_release_allowed=new_release_allowed,
        source_quality_ok=source_quality_ok,
        source_quality_notes=tuple(quality_notes),
    )


__all__ = ["DPECutoverReport", "PARITY_REPORT_VERSION", "WorkbookEvidence", "audit_dpe_cutover_readiness"]
