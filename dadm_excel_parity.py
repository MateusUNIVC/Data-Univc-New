from __future__ import annotations

import json
import math
from dataclasses import asdict, dataclass
from datetime import date, datetime, timezone
from io import BytesIO
from typing import Any, Mapping

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from dadm_excel_official import (
    build_dadm_excel_official_artifact_from_payload,
    build_dadm_excel_official_payload_from_db,
    build_dadm_snapshot_context,
)
from dadm_tallos_analytics import build_tallos_dashboard
from dadm_v2_analytics import month_end, month_start
from dadm_v2_report import build_dadm_v2_report
from excel_official.adapters.dadm import DADMAdapter, DADM_COMPARISON_METRICS, DADM_METRICS
from excel_official.audit.quality import evaluate_metric
from excel_official.context import AdapterInput
from release_info import APP_VERSION, SCHEMA_VERSION

PARITY_REPORT_VERSION = 2


@dataclass(frozen=True, slots=True)
class DADMParityCase:
    scope: str
    metric_code: str
    expected: float | int | None
    actual: float | int | None
    tolerance: float
    status: str
    delta: float | None = None


@dataclass(frozen=True, slots=True)
class WorkbookEvidence:
    engine: str
    build_ok: bool
    bytes_size: int = 0
    sheets: int = 0
    formulas: int = 0
    charts: int = 0
    error: str | None = None


@dataclass(frozen=True, slots=True)
class DADMParityReport:
    source_kind: str
    generated_at: str
    system_version: str
    schema_version: str
    payload_hash: str | None
    cases: tuple[DADMParityCase, ...]
    legacy_workbook: WorkbookEvidence | None
    new_workbook: WorkbookEvidence | None
    new_release_allowed: bool
    source_quality_ok: bool
    source_quality_notes: tuple[str, ...] = ()
    report_version: int = PARITY_REPORT_VERSION
    directorate: str = "DADM"

    @property
    def failures(self) -> tuple[DADMParityCase, ...]:
        return tuple(case for case in self.cases if case.status == "FAIL")

    @property
    def cutover_status(self) -> str:
        builds_ok = bool(
            self.legacy_workbook and self.legacy_workbook.build_ok
            and self.new_workbook and self.new_workbook.build_ok
        )
        if self.failures or not builds_ok or not self.new_release_allowed or not self.source_quality_ok:
            return "BLOCKED"
        return "READY" if self.source_kind == "production" else "CANDIDATE_PASS"

    def to_dict(self) -> dict[str, Any]:
        data = asdict(self)
        data["summary"] = {
            "cases": len(self.cases),
            "passed": sum(1 for case in self.cases if case.status == "PASS"),
            "failures": len(self.failures),
            "semantic_passed": not self.failures,
            "source_quality_ok": self.source_quality_ok,
            "new_release_allowed": self.new_release_allowed,
            "cutover_status": self.cutover_status,
        }
        return data

    def to_json(self) -> str:
        return json.dumps(self.to_dict(), ensure_ascii=False, indent=2)

    def to_markdown(self) -> str:
        lines = [
            "# DADM Excel Parity · 04B",
            "",
            f"- Source kind: **{self.source_kind}**",
            f"- Cases: **{len(self.cases)}**",
            f"- Failures: **{len(self.failures)}**",
            f"- Source quality: **{'PASS' if self.source_quality_ok else 'FAIL'}**",
            f"- Excel Official release audit: **{'PASS' if self.new_release_allowed else 'FAIL'}**",
            f"- Cutover status: **{self.cutover_status}**",
            "",
            "## Contract",
            "",
            "- Interactive/offline metrics are limited to additive or exactly recomposable TALLOS components.",
            "- TME/TMA use sum/count of valid values; the workbook does not average averages.",
            "- Rating uses valid TALLOS scores 1–10 only; no satisfaction threshold is inferred.",
            "- Distinct protocols/people/operators stay in the backend summary because they are not safely additive across offline dimensions.",
            "- The previous-period comparison preserves the DADM V2 `previous_period` backend semantics for the export context.",
            "- Median/P90 remain backend snapshot statistics because percentiles are not additive.",
            "",
        ]
        if self.failures:
            lines.extend(["## Failures", ""])
            for case in self.failures:
                lines.append(
                    f"- `{case.scope}` · `{case.metric_code}`: expected={case.expected!r}, actual={case.actual!r}, delta={case.delta!r}"
                )
        if self.source_quality_notes:
            lines.extend(["", "## Source quality", ""])
            lines.extend(f"- {note}" for note in self.source_quality_notes)
        lines.append("")
        return "\n".join(lines)


def _evidence(engine: str, buffer: BytesIO) -> WorkbookEvidence:
    try:
        raw = buffer.getvalue()
        wb = load_workbook(BytesIO(raw), read_only=False, data_only=False, keep_links=True)
        formulas = sum(1 for ws in wb.worksheets for row in ws.iter_rows() for cell in row if cell.data_type == "f")
        charts = sum(len(getattr(ws, "_charts", ())) for ws in wb.worksheets)
        sheets = len(wb.sheetnames)
        wb.close()
        return WorkbookEvidence(engine, True, len(raw), sheets, formulas, charts)
    except Exception as exc:  # pragma: no cover
        return WorkbookEvidence(engine, False, error=f"{type(exc).__name__}: {exc}")


def _same(expected: float | int | None, actual: float | int | None, tolerance: float) -> tuple[bool, float | None]:
    if expected is None or actual is None:
        return expected is None and actual is None, None
    delta = float(actual) - float(expected)
    return math.isclose(float(expected), float(actual), abs_tol=tolerance, rel_tol=0.0), delta


def _expected(summary: Mapping[str, Any]) -> dict[str, float | int | None]:
    def pct(key: str) -> float | None:
        value = summary.get(key)
        return None if value is None else float(value) / 100.0

    def minutes(key: str) -> float | None:
        value = summary.get(key)
        return None if value is None else float(value) / 60.0

    return {
        DADM_METRICS["attendances"]: int(summary.get("attendances") or 0),
        DADM_METRICS["finalization"]: pct("finalization_rate_pct"),
        DADM_METRICS["open"]: int(summary.get("open") or 0),
        DADM_METRICS["tme"]: minutes("tme_avg_seconds"),
        DADM_METRICS["tma"]: minutes("tma_avg_seconds"),
        DADM_METRICS["rating"]: summary.get("rating_avg"),
        DADM_METRICS["rating_coverage"]: pct("rating_coverage_pct"),
        DADM_METRICS["transfer_rate"]: pct("transfer_rate_pct"),
    }


def audit_dadm_parity(
    db: Session,
    directorate_id: int,
    from_month: str | None = None,
    to_month: str | None = None,
    *,
    allowed_departments: tuple[str, ...] | None = None,
    targets: list[dict[str, Any]] | tuple[dict[str, Any], ...] | None = None,
    actions: list[dict[str, Any]] | tuple[dict[str, Any], ...] | None = None,
    source_kind: str = "fixture",
    generated_by: str | None = None,
    build_workbooks: bool = True,
) -> DADMParityReport:
    payload = build_dadm_excel_official_payload_from_db(
        db,
        directorate_id,
        from_month,
        to_month,
        allowed_departments=allowed_departments,
        targets=targets,
        actions=actions,
    )
    context = build_dadm_snapshot_context(payload, generated_by=generated_by)
    period = payload["period"]
    start_date = date.fromisoformat(period["start"])
    end_date = date.fromisoformat(period["end"])

    cube = [dict(row) for row in payload.get("cube", ())]
    department_labels = {
        str(row.get("department") or "").strip(): str(row.get("department_name") or row.get("department") or "").strip()
        for row in cube
        if str(row.get("department") or "").strip()
    }
    departments = sorted(department_labels)
    months = sorted({str(row.get("period") or "").strip() for row in cube if str(row.get("period") or "").strip()})

    scopes: list[tuple[str, date, date, str | None, str]] = [
        ("DADM · período exportado", start_date, end_date, None, "(todos)"),
    ]
    for department in departments:
        label = department_labels.get(department) or department
        scopes.append((f"Departamento {label}", start_date, end_date, department, "(todos)"))
    for month in months:
        scopes.append((f"Mês {month}", month_start(month), month_end(month), None, month))

    cases: list[DADMParityCase] = []
    tolerance_by_metric = {
        DADM_METRICS["attendances"]: 0.0,
        DADM_METRICS["finalization"]: 0.000001,
        DADM_METRICS["open"]: 0.0,
        DADM_METRICS["tme"]: 0.000001,
        DADM_METRICS["tma"]: 0.000001,
        DADM_METRICS["rating"]: 0.000001,
        DADM_METRICS["rating_coverage"]: 0.000001,
        DADM_METRICS["transfer_rate"]: 0.000001,
    }

    for label, scope_start, scope_end, department, period_value in scopes:
        backend = build_tallos_dashboard(
            db,
            directorate_id,
            scope_start,
            scope_end,
            department=department,
            comparison_mode="none",
            grain="month",
            allowed_departments=allowed_departments,
        )
        expected = _expected(backend["summary"])
        initial = {
            "period": period_value,
            "department": department or "(todos)",
            "employee": "(todos)",
            "channel": "(todos)",
            "status": "(todos)",
            "tabulation": "(todos)",
        }
        spec = DADMAdapter().build_spec(AdapterInput(context, payload, initial))
        for metric_code, expected_value in expected.items():
            actual = evaluate_metric(spec, metric_code)
            tolerance = tolerance_by_metric[metric_code]
            passed, delta = _same(expected_value, actual, tolerance)
            cases.append(DADMParityCase(label, metric_code, expected_value, actual, tolerance, "PASS" if passed else "FAIL", delta))

    comparison_expected = _expected(dict(payload.get("comparison", {}) or {}))
    comparison_spec = DADMAdapter().build_spec(AdapterInput(context, payload, dict(payload.get("initial", {}) or {})))
    comparison_pairs = (
        (DADM_COMPARISON_METRICS["attendances"], comparison_expected[DADM_METRICS["attendances"]], 0.0),
        (DADM_COMPARISON_METRICS["tme"], comparison_expected[DADM_METRICS["tme"]], 0.000001),
        (DADM_COMPARISON_METRICS["tma"], comparison_expected[DADM_METRICS["tma"]], 0.000001),
        (DADM_COMPARISON_METRICS["rating"], comparison_expected[DADM_METRICS["rating"]], 0.000001),
        (DADM_COMPARISON_METRICS["rating_coverage"], comparison_expected[DADM_METRICS["rating_coverage"]], 0.000001),
    )
    for metric_code, expected_value, tolerance in comparison_pairs:
        actual = evaluate_metric(comparison_spec, metric_code)
        passed, delta = _same(expected_value, actual, tolerance)
        cases.append(DADMParityCase("Período anterior · recorte exportado", metric_code, expected_value, actual, tolerance, "PASS" if passed else "FAIL", delta))

    quality_notes: list[str] = []
    if not cube:
        quality_notes.append("Cubo TALLOS vazio.")
    if not payload.get("summary"):
        quality_notes.append("Resumo backend ausente.")
    source_quality_ok = not quality_notes

    legacy_evidence: WorkbookEvidence | None = None
    new_evidence: WorkbookEvidence | None = None
    new_release_allowed = False
    if build_workbooks:
        try:
            legacy, _ = build_dadm_v2_report(
                db,
                directorate_id,
                period["from_month"],
                period["to_month"],
                allowed_departments=allowed_departments,
            )
            legacy_evidence = _evidence("dadm_v2", legacy)
        except Exception as exc:  # pragma: no cover
            legacy_evidence = WorkbookEvidence("dadm_v2", False, error=f"{type(exc).__name__}: {exc}")
        try:
            artifact = build_dadm_excel_official_artifact_from_payload(payload, generated_by=generated_by)
            new_release_allowed = artifact.audit.release_allowed
            new_evidence = _evidence("excel_official", artifact.to_bytes())
        except Exception as exc:  # pragma: no cover
            new_evidence = WorkbookEvidence("excel_official", False, error=f"{type(exc).__name__}: {exc}")

    return DADMParityReport(
        source_kind=str(source_kind or "fixture").strip().lower(),
        generated_at=datetime.now(timezone.utc).isoformat(),
        system_version=APP_VERSION,
        schema_version=str(SCHEMA_VERSION),
        payload_hash=context.payload_hash,
        cases=tuple(cases),
        legacy_workbook=legacy_evidence,
        new_workbook=new_evidence,
        new_release_allowed=new_release_allowed,
        source_quality_ok=source_quality_ok,
        source_quality_notes=tuple(quality_notes),
    )


__all__ = ["DADMParityCase", "DADMParityReport", "PARITY_REPORT_VERSION", "WorkbookEvidence", "audit_dadm_parity"]
