from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Mapping

from openpyxl import Workbook
from openpyxl.styles import Protection

from .audit import AuditResult, WorkbookReleaseBlockedError, audit_workbook
from .components import (
    write_action_plan_sheet,
    write_dashboard_charts,
    write_dashboard_sheet,
    write_dataset_sheet,
    write_matrix_sheet,
    write_parameter_system,
    write_quality_sheet,
)
from .components.readme import ReadmeRef, write_readme_sheet
from .components.technical import TechnicalRegistryRef, ensure_technical_layers
from .constants import OPTIONAL_TECHNICAL_SHEETS, REQUIRED_INSTITUTIONAL_SHEETS, STANDARD_TECHNICAL_SHEETS
from .contract import (
    ActionPlanRef,
    ChartSystemRef,
    DashboardRef,
    DatasetRef,
    MatrixRef,
    ParameterSystemRef,
    QualityRef,
    WorkbookSpec,
)
from .metadata import apply_workbook_metadata
from .protection import protect_sheet
from .styles import apply_cell_role, apply_sheet_defaults, register_named_styles, set_subtitle_row_height, set_title_row_height
from .theme import CellRole, DEFAULT_THEME, ExcelOfficialTheme
from .validation import assert_valid_workbook_spec


@dataclass(frozen=True, slots=True)
class WorkbookBuildRefs:
    readme: ReadmeRef
    datasets: Mapping[str, DatasetRef]
    parameters: ParameterSystemRef
    dashboard: DashboardRef
    charts: ChartSystemRef
    matrix: MatrixRef | None
    quality: QualityRef
    action_plan: ActionPlanRef | None
    technical: TechnicalRegistryRef


@dataclass(slots=True)
class WorkbookArtifact:
    workbook: Workbook
    spec: WorkbookSpec
    refs: WorkbookBuildRefs
    audit: AuditResult

    def _audit_before_release(self) -> None:
        self.audit = audit_workbook(
            self.workbook,
            self.spec,
            parameter_system=self.refs.parameters,
            action_plan_ref=self.refs.action_plan,
        )
        if not self.audit.release_allowed:
            raise WorkbookReleaseBlockedError(self.audit)

    def save(self, path: str | Path) -> Path:
        self._audit_before_release()
        output = Path(path)
        output.parent.mkdir(parents=True, exist_ok=True)
        self.workbook.save(output)
        return output

    def to_bytes(self) -> BytesIO:
        self._audit_before_release()
        output = BytesIO()
        self.workbook.save(output)
        output.seek(0)
        return output


def _create_required_placeholder(
    workbook: Workbook,
    sheet_name: str,
    *,
    title: str,
    message: str,
    theme: ExcelOfficialTheme,
) -> None:
    if sheet_name in workbook.sheetnames:
        return
    ws = workbook.create_sheet(sheet_name)
    apply_sheet_defaults(ws, theme=theme, zoom=90, freeze_panes="A4", tab_color=theme.palette.green)
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:F1")
    ws["A1"] = title
    apply_cell_role(ws["A1"], CellRole.TITLE, theme=theme)
    set_title_row_height(ws, 1, theme)
    ws.merge_cells("A2:F2")
    ws["A2"] = message
    apply_cell_role(ws["A2"], CellRole.NOTE, theme=theme)
    set_subtitle_row_height(ws, 2, theme)
    ws.column_dimensions["A"].width = 28
    for cells in ws.iter_rows():
        for cell in cells:
            cell.protection = Protection(locked=True)
    protect_sheet(ws)


def _reorder_sheets(workbook: Workbook, spec: WorkbookSpec) -> None:
    required = list(REQUIRED_INSTITUTIONAL_SHEETS)
    domain_codes = [item.dataset_code for item in spec.domain_sheets]
    dataset_by_code = {item.code: item for item in spec.datasets}
    domain_names = [dataset_by_code[code].sheet_name for code in domain_codes if code in dataset_by_code]
    other_domain = [item.sheet_name for item in spec.datasets if not item.technical and item.sheet_name not in domain_names]
    technical_order = list(STANDARD_TECHNICAL_SHEETS) + list(OPTIONAL_TECHNICAL_SHEETS)
    extra_technical = [
        item.sheet_name
        for item in spec.datasets
        if item.technical and item.sheet_name not in technical_order
    ]
    preferred = required + domain_names + other_domain + technical_order + extra_technical
    seen: set[str] = set()
    ordered_names: list[str] = []
    for name in preferred + list(workbook.sheetnames):
        if name in workbook.sheetnames and name not in seen:
            seen.add(name)
            ordered_names.append(name)
    workbook._sheets = [workbook[name] for name in ordered_names]


class ExcelOfficialCore:
    def __init__(self, *, theme: ExcelOfficialTheme = DEFAULT_THEME):
        self.theme = theme

    def build(self, spec: WorkbookSpec, *, enforce_release: bool = True) -> WorkbookArtifact:
        assert_valid_workbook_spec(spec)
        workbook = Workbook()
        workbook.remove(workbook.active)
        register_named_styles(workbook, self.theme)

        dataset_refs: dict[str, DatasetRef] = {}
        for dataset in spec.datasets:
            dataset_refs[dataset.code] = write_dataset_sheet(workbook, dataset, theme=self.theme)

        parameter_system = write_parameter_system(workbook, spec, theme=self.theme)
        dashboard = write_dashboard_sheet(workbook, spec, dataset_refs, parameter_system, theme=self.theme)
        charts = write_dashboard_charts(
            workbook,
            spec,
            dataset_refs,
            parameter_system,
            dashboard,
            theme=self.theme,
        )
        matrix = write_matrix_sheet(workbook, spec, dataset_refs, parameter_system, theme=self.theme)
        if matrix is None:
            _create_required_placeholder(
                workbook,
                "MATRIZ",
                title="MATRIZ DE INDICADORES",
                message="Nenhuma matriz foi configurada para este snapshot.",
                theme=self.theme,
            )

        quality = write_quality_sheet(workbook, spec, dataset_refs, parameter_system, theme=self.theme)
        action_plan = write_action_plan_sheet(workbook, spec, theme=self.theme)
        if action_plan is None:
            _create_required_placeholder(
                workbook,
                "PLANO_DE_ACAO",
                title="PLANO DE ACAO",
                message="Plano de acao nao habilitado para este snapshot.",
                theme=self.theme,
            )

        readme = write_readme_sheet(workbook, spec, theme=self.theme)
        technical = ensure_technical_layers(workbook, spec, theme=self.theme)
        apply_workbook_metadata(workbook, spec)
        _reorder_sheets(workbook, spec)

        workbook.calculation.calcMode = "auto"
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True

        audit = audit_workbook(
            workbook,
            spec,
            parameter_system=parameter_system,
            action_plan_ref=action_plan,
        )
        refs = WorkbookBuildRefs(
            readme=readme,
            datasets=dataset_refs,
            parameters=parameter_system,
            dashboard=dashboard,
            charts=charts,
            matrix=matrix,
            quality=quality,
            action_plan=action_plan,
            technical=technical,
        )
        artifact = WorkbookArtifact(workbook=workbook, spec=spec, refs=refs, audit=audit)
        if enforce_release and not audit.release_allowed:
            raise WorkbookReleaseBlockedError(audit)
        return artifact


__all__ = ["ExcelOfficialCore", "WorkbookArtifact", "WorkbookBuildRefs"]
