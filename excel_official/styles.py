from __future__ import annotations

from copy import copy
from dataclasses import dataclass
from typing import Iterable

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, NamedStyle, PatternFill, Protection, Side
from openpyxl.worksheet.worksheet import Worksheet

from .theme import CellRole, DEFAULT_THEME, ExcelOfficialTheme

STYLE_PREFIX = "du_"


@dataclass(frozen=True, slots=True)
class StyleNames:
    title: str = f"{STYLE_PREFIX}title"
    subtitle: str = f"{STYLE_PREFIX}subtitle"
    section: str = f"{STYLE_PREFIX}section"
    table_header: str = f"{STYLE_PREFIX}table_header"
    body: str = f"{STYLE_PREFIX}body"
    body_center: str = f"{STYLE_PREFIX}body_center"
    input: str = f"{STYLE_PREFIX}input"
    imported: str = f"{STYLE_PREFIX}imported"
    derived: str = f"{STYLE_PREFIX}derived"
    static: str = f"{STYLE_PREFIX}static"
    control: str = f"{STYLE_PREFIX}control"
    visualization: str = f"{STYLE_PREFIX}visualization"
    technical: str = f"{STYLE_PREFIX}technical"
    warning: str = f"{STYLE_PREFIX}warning"
    error: str = f"{STYLE_PREFIX}error"
    success: str = f"{STYLE_PREFIX}success"
    kpi_label: str = f"{STYLE_PREFIX}kpi_label"
    kpi_value: str = f"{STYLE_PREFIX}kpi_value"
    note: str = f"{STYLE_PREFIX}note"


STYLE_NAMES = StyleNames()

ROLE_TO_STYLE: dict[CellRole, str] = {
    CellRole.TITLE: STYLE_NAMES.title,
    CellRole.SUBTITLE: STYLE_NAMES.subtitle,
    CellRole.SECTION: STYLE_NAMES.section,
    CellRole.TABLE_HEADER: STYLE_NAMES.table_header,
    CellRole.BODY: STYLE_NAMES.body,
    CellRole.BODY_CENTER: STYLE_NAMES.body_center,
    CellRole.INPUT: STYLE_NAMES.input,
    CellRole.IMPORTED: STYLE_NAMES.imported,
    CellRole.DERIVED: STYLE_NAMES.derived,
    CellRole.STATIC: STYLE_NAMES.static,
    CellRole.CONTROL: STYLE_NAMES.control,
    CellRole.VISUALIZATION: STYLE_NAMES.visualization,
    CellRole.TECHNICAL: STYLE_NAMES.technical,
    CellRole.WARNING: STYLE_NAMES.warning,
    CellRole.ERROR: STYLE_NAMES.error,
    CellRole.SUCCESS: STYLE_NAMES.success,
    CellRole.KPI_LABEL: STYLE_NAMES.kpi_label,
    CellRole.KPI_VALUE: STYLE_NAMES.kpi_value,
    CellRole.NOTE: STYLE_NAMES.note,
}


def _solid(color: str) -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=color)


def _border_bottom(color: str, style: str = "thin") -> Border:
    return Border(bottom=Side(style=style, color=color))


def _outline_border(color: str) -> Border:
    side = Side(style="thin", color=color)
    return Border(left=side, right=side, top=side, bottom=side)


def _style_objects(theme: ExcelOfficialTheme) -> dict[str, NamedStyle]:
    p = theme.palette
    t = theme.typography

    styles: dict[str, NamedStyle] = {}

    styles[STYLE_NAMES.title] = NamedStyle(
        name=STYLE_NAMES.title,
        font=Font(name=t.family, size=t.title_size, bold=True, color=p.green_dark),
        alignment=Alignment(horizontal="left", vertical="center"),
        border=_border_bottom(p.green, "medium"),
        protection=Protection(locked=True),
    )
    styles[STYLE_NAMES.subtitle] = NamedStyle(
        name=STYLE_NAMES.subtitle,
        font=Font(name=t.family, size=t.body_size, color=p.gray_600),
        alignment=Alignment(horizontal="left", vertical="top", wrap_text=True),
        protection=Protection(locked=True),
    )
    styles[STYLE_NAMES.section] = NamedStyle(
        name=STYLE_NAMES.section,
        font=Font(name=t.family, size=t.section_size, bold=True, color=p.white),
        fill=_solid(p.green_dark),
        alignment=Alignment(horizontal="left", vertical="center"),
        protection=Protection(locked=True),
    )
    styles[STYLE_NAMES.table_header] = NamedStyle(
        name=STYLE_NAMES.table_header,
        font=Font(name=t.family, size=t.body_size, bold=True, color=p.white),
        fill=_solid(p.green),
        alignment=Alignment(horizontal="center", vertical="center", wrap_text=True),
        border=_border_bottom(p.green_dark, "thin"),
        protection=Protection(locked=True),
    )
    styles[STYLE_NAMES.body] = NamedStyle(
        name=STYLE_NAMES.body,
        font=Font(name=t.family, size=t.body_size, color=p.black),
        alignment=Alignment(vertical="center", wrap_text=True),
        border=_border_bottom(p.gray_300, "hair"),
        protection=Protection(locked=True),
    )
    styles[STYLE_NAMES.body_center] = NamedStyle(
        name=STYLE_NAMES.body_center,
        font=Font(name=t.family, size=t.body_size, color=p.black),
        alignment=Alignment(horizontal="center", vertical="center", wrap_text=True),
        border=_border_bottom(p.gray_300, "hair"),
        protection=Protection(locked=True),
    )
    styles[STYLE_NAMES.input] = NamedStyle(
        name=STYLE_NAMES.input,
        font=Font(name=t.family, size=t.parameter_size, bold=True, color=p.input_text),
        fill=_solid(p.input_fill),
        alignment=Alignment(horizontal="center", vertical="center", wrap_text=True),
        border=_outline_border(p.gray_300),
        protection=Protection(locked=False),
    )
    styles[STYLE_NAMES.imported] = NamedStyle(
        name=STYLE_NAMES.imported,
        font=Font(name=t.family, size=t.body_size, color=p.imported_text),
        alignment=Alignment(vertical="center", wrap_text=True),
        border=_border_bottom(p.gray_300, "hair"),
        protection=Protection(locked=True),
    )
    styles[STYLE_NAMES.derived] = NamedStyle(
        name=STYLE_NAMES.derived,
        font=Font(name=t.family, size=t.body_size, color=p.derived_text),
        alignment=Alignment(vertical="center", wrap_text=True),
        border=_border_bottom(p.gray_300, "hair"),
        protection=Protection(locked=True),
    )
    styles[STYLE_NAMES.static] = NamedStyle(
        name=STYLE_NAMES.static,
        font=Font(name=t.family, size=t.body_size, color=p.static_text),
        alignment=Alignment(vertical="center", wrap_text=True),
        protection=Protection(locked=True),
    )
    styles[STYLE_NAMES.control] = NamedStyle(
        name=STYLE_NAMES.control,
        font=Font(name=t.family, size=t.body_size, bold=True, color=p.control_text),
        fill=_solid(p.gray_50),
        alignment=Alignment(horizontal="center", vertical="center", wrap_text=True),
        border=_outline_border(p.gray_300),
        protection=Protection(locked=True),
    )
    styles[STYLE_NAMES.visualization] = NamedStyle(
        name=STYLE_NAMES.visualization,
        font=Font(name=t.family, size=t.body_size, bold=True, color=p.visualization_text),
        fill=_solid(p.green_light),
        alignment=Alignment(horizontal="center", vertical="center", wrap_text=True),
        protection=Protection(locked=True),
    )
    styles[STYLE_NAMES.technical] = NamedStyle(
        name=STYLE_NAMES.technical,
        font=Font(name=t.family, size=t.body_size, color=p.static_text),
        fill=_solid(p.gray_100),
        alignment=Alignment(vertical="center", wrap_text=True),
        border=_border_bottom(p.gray_300, "hair"),
        protection=Protection(locked=True),
    )
    styles[STYLE_NAMES.warning] = NamedStyle(
        name=STYLE_NAMES.warning,
        font=Font(name=t.family, size=t.body_size, bold=True, color=p.ink),
        fill=_solid(p.warning_fill),
        alignment=Alignment(horizontal="center", vertical="center", wrap_text=True),
        border=_outline_border(p.orange),
        protection=Protection(locked=True),
    )
    styles[STYLE_NAMES.error] = NamedStyle(
        name=STYLE_NAMES.error,
        font=Font(name=t.family, size=t.body_size, bold=True, color=p.red),
        fill=_solid(p.error_fill),
        alignment=Alignment(horizontal="center", vertical="center", wrap_text=True),
        border=_outline_border(p.red),
        protection=Protection(locked=True),
    )
    styles[STYLE_NAMES.success] = NamedStyle(
        name=STYLE_NAMES.success,
        font=Font(name=t.family, size=t.body_size, bold=True, color=p.green_dark),
        fill=_solid(p.green_pale),
        alignment=Alignment(horizontal="center", vertical="center", wrap_text=True),
        border=_outline_border(p.green_mid),
        protection=Protection(locked=True),
    )
    styles[STYLE_NAMES.kpi_label] = NamedStyle(
        name=STYLE_NAMES.kpi_label,
        font=Font(name=t.family, size=t.small_size, bold=True, color=p.gray_600),
        fill=_solid(p.green_light),
        alignment=Alignment(horizontal="center", vertical="center", wrap_text=True),
        protection=Protection(locked=True),
    )
    styles[STYLE_NAMES.kpi_value] = NamedStyle(
        name=STYLE_NAMES.kpi_value,
        font=Font(name=t.family, size=t.kpi_value_size, bold=True, color=p.green_dark),
        fill=_solid(p.green_light),
        alignment=Alignment(horizontal="center", vertical="center"),
        protection=Protection(locked=True),
    )
    styles[STYLE_NAMES.note] = NamedStyle(
        name=STYLE_NAMES.note,
        font=Font(name=t.family, size=t.small_size, color=p.gray_600),
        alignment=Alignment(horizontal="left", vertical="top", wrap_text=True),
        protection=Protection(locked=True),
    )
    return styles


def register_named_styles(workbook: Workbook, theme: ExcelOfficialTheme = DEFAULT_THEME) -> tuple[str, ...]:
    """Register the institutional styles once and return their names.

    Safe to call repeatedly on the same workbook.
    """
    existing = set(workbook.named_styles)
    registered: list[str] = []
    for name, style in _style_objects(theme).items():
        if name not in existing:
            workbook.add_named_style(style)
            existing.add(name)
        registered.append(name)
    return tuple(registered)


def style_name_for_role(role: CellRole) -> str:
    return ROLE_TO_STYLE[CellRole(role)]


def apply_cell_role(
    cell: Cell,
    role: CellRole,
    *,
    number_format: str | None = None,
    theme: ExcelOfficialTheme = DEFAULT_THEME,
) -> Cell:
    register_named_styles(cell.parent.parent, theme)
    cell.style = style_name_for_role(role)
    if number_format is not None:
        cell.number_format = number_format
    return cell


def apply_cells_role(
    cells: Iterable[Cell],
    role: CellRole,
    *,
    number_format: str | None = None,
    theme: ExcelOfficialTheme = DEFAULT_THEME,
) -> None:
    for cell in cells:
        apply_cell_role(cell, role, number_format=number_format, theme=theme)


def apply_range_role(
    worksheet: Worksheet,
    cell_range: str,
    role: CellRole,
    *,
    number_format: str | None = None,
    theme: ExcelOfficialTheme = DEFAULT_THEME,
) -> None:
    for row in worksheet[cell_range]:
        apply_cells_role(row, role, number_format=number_format, theme=theme)


def apply_sheet_defaults(
    worksheet: Worksheet,
    *,
    theme: ExcelOfficialTheme = DEFAULT_THEME,
    zoom: int | None = None,
    freeze_panes: str | None = None,
    show_gridlines: bool = False,
    tab_color: str | None = None,
) -> None:
    layout = theme.layout
    worksheet.sheet_view.showGridLines = show_gridlines
    worksheet.sheet_view.zoomScale = zoom if zoom is not None else layout.default_zoom
    worksheet.freeze_panes = freeze_panes
    if tab_color:
        worksheet.sheet_properties.tabColor = tab_color
    worksheet.sheet_format.defaultRowHeight = layout.default_row_height
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.page_margins.left = layout.margin_left
    worksheet.page_margins.right = layout.margin_right
    worksheet.page_margins.top = layout.margin_top
    worksheet.page_margins.bottom = layout.margin_bottom
    worksheet.page_margins.header = layout.margin_header
    worksheet.page_margins.footer = layout.margin_footer


def set_title_row_height(worksheet: Worksheet, row: int, theme: ExcelOfficialTheme = DEFAULT_THEME) -> None:
    worksheet.row_dimensions[row].height = theme.layout.title_row_height


def set_subtitle_row_height(worksheet: Worksheet, row: int, theme: ExcelOfficialTheme = DEFAULT_THEME) -> None:
    worksheet.row_dimensions[row].height = theme.layout.subtitle_row_height


def set_section_row_height(worksheet: Worksheet, row: int, theme: ExcelOfficialTheme = DEFAULT_THEME) -> None:
    worksheet.row_dimensions[row].height = theme.layout.section_row_height


def set_table_header_row_height(worksheet: Worksheet, row: int, theme: ExcelOfficialTheme = DEFAULT_THEME) -> None:
    worksheet.row_dimensions[row].height = theme.layout.table_header_row_height


def copy_cell_style(source: Cell, target: Cell) -> None:
    """Copy a style safely without sharing mutable style components."""
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.protection = copy(source.protection)
    target.number_format = source.number_format


__all__ = [
    "ROLE_TO_STYLE",
    "STYLE_NAMES",
    "STYLE_PREFIX",
    "StyleNames",
    "apply_cell_role",
    "apply_cells_role",
    "apply_range_role",
    "apply_sheet_defaults",
    "copy_cell_style",
    "register_named_styles",
    "set_section_row_height",
    "set_subtitle_row_height",
    "set_table_header_row_height",
    "set_title_row_height",
    "style_name_for_role",
]
