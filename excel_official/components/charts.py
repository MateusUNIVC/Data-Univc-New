from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from typing import Any, Mapping

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from ..constants import TECHNICAL_LAYER_NOTICE
from ..contract import (
    ChartRef,
    ChartSpec,
    ChartSystemRef,
    ChartType,
    DatasetRef,
    MetricSpec,
    MetricUnit,
    ParameterSystemRef,
    WorkbookSpec,
)
from ..protection import protect_sheet
from ..styles import apply_cell_role, apply_sheet_defaults, set_section_row_height
from ..theme import CellRole, DEFAULT_THEME, ExcelOfficialTheme
from ..validation import assert_valid_workbook_spec
from .dashboard import DASHBOARD_SHEET, _metric_expression, _metric_number_format
from .parameters import _dimension_members

CALC_SHEET = "CALC"


@dataclass(frozen=True, slots=True)
class ChartWriteError(ValueError):
    code: str
    message: str

    def __str__(self) -> str:
        return f"{self.code}: {self.message}"


def _metric_by_code(spec: WorkbookSpec):
    return {item.code: item for item in spec.metrics}


def _binding_by_metric(spec: WorkbookSpec):
    return {item.metric_code: item for item in spec.metric_bindings}


def _dataset_by_code(spec: WorkbookSpec):
    return {item.code: item for item in spec.datasets}


def _dimension_by_code(spec: WorkbookSpec):
    return {item.code: item for item in spec.dimensions}


def _parameter_by_code(spec: WorkbookSpec):
    return {item.code: item for item in spec.parameters}


def _excel_sheet(sheet_name: str) -> str:
    return "'" + sheet_name.replace("'", "''") + "'"


def _excel_string(value: str) -> str:
    return '"' + value.replace('"', '""') + '"'


def _formula_literal(value: Any) -> str:
    if value is None:
        return '""'
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        base = f"DATE({value.year},{value.month},{value.day})"
        if value.hour or value.minute or value.second:
            base += f"+TIME({value.hour},{value.minute},{value.second})"
        return base
    if isinstance(value, date):
        return f"DATE({value.year},{value.month},{value.day})"
    if isinstance(value, (int, float, Decimal)):
        return str(value)
    return _excel_string(str(value))


def _absolute_cell(sheet_name: str, column: int, row: int) -> str:
    return f"{_excel_sheet(sheet_name)}!${get_column_letter(column)}${row}"


def _sorted_members(spec: WorkbookSpec, dimension_code: str, sort: str | None, top_n: int | None):
    dimension = _dimension_by_code(spec)[dimension_code]
    members = list(_dimension_members(spec, dimension))
    if sort in {None, "dimension_asc"}:
        pass
    elif sort == "dimension_desc":
        members.reverse()
    elif sort == "label_asc":
        members.sort(key=lambda item: str(item.label).casefold())
    elif sort == "label_desc":
        members.sort(key=lambda item: str(item.label).casefold(), reverse=True)
    else:
        raise ChartWriteError("chart.unsupported_sort", f"Ordenacao de grafico nao suportada no Contract V1: {sort!r}.")
    if top_n is not None:
        members = members[:top_n]
    return tuple(members)


def _dimension_fact_column(spec: WorkbookSpec, chart: ChartSpec, dimension_code: str) -> str:
    dataset = _dataset_by_code(spec)[chart.dataset_code]
    columns = {column.code for column in dataset.columns}
    column_code = dataset.dimension_columns.get(dimension_code, dimension_code)
    if column_code not in columns:
        raise ChartWriteError(
            "chart.dimension_not_available_in_dataset",
            f"Dataset {dataset.code!r} nao possui coluna para a dimensao {dimension_code!r}.",
        )
    return column_code


def _preflight_chart_system(
    workbook: Workbook,
    spec: WorkbookSpec,
    dataset_refs: Mapping[str, DatasetRef],
    parameter_system: ParameterSystemRef,
    dashboard_ref,
) -> None:
    assert_valid_workbook_spec(spec)
    if not spec.dashboard.charts:
        return
    if DASHBOARD_SHEET not in workbook.sheetnames:
        raise ChartWriteError("chart.dashboard_missing", "PAINEL precisa existir antes da geracao dos graficos.")

    metrics = _metric_by_code(spec)
    bindings = _binding_by_metric(spec)
    parameters = _parameter_by_code(spec)
    dimensions = _dimension_by_code(spec)

    for chart in spec.dashboard.charts:
        if chart.code not in dashboard_ref.chart_anchors:
            raise ChartWriteError("chart.anchor_missing", f"DashboardRef nao possui ancora para {chart.code!r}.")
        metric = metrics.get(chart.metric_code)
        if metric is None:
            raise ChartWriteError("chart.metric_spec_required", f"Grafico {chart.code!r} precisa de MetricSpec.")
        binding = bindings.get(chart.metric_code)
        if binding is None:
            raise ChartWriteError("chart.metric_binding_required", f"Grafico {chart.code!r} precisa de MetricBinding.")
        if binding.dataset_code != chart.dataset_code:
            raise ChartWriteError(
                "chart.dataset_binding_mismatch",
                f"Grafico {chart.code!r} usa {chart.dataset_code!r}, mas MetricBinding usa {binding.dataset_code!r}.",
            )
        if chart.dataset_code not in dataset_refs:
            raise ChartWriteError("chart.missing_dataset_ref", f"Dataset {chart.dataset_code!r} nao foi materializado.")
        if chart.dimension_code not in dimensions:
            raise ChartWriteError("chart.unknown_dimension", f"Dimensao {chart.dimension_code!r} nao existe.")
        _dimension_fact_column(spec, chart, chart.dimension_code)
        if metric.allowed_dimensions and chart.dimension_code not in metric.allowed_dimensions:
            raise ChartWriteError(
                "chart.dimension_not_allowed",
                f"Metrica {metric.code!r} nao permite grafico por {chart.dimension_code!r}.",
            )
        if not _sorted_members(spec, chart.dimension_code, chart.sort, chart.top_n):
            raise ChartWriteError("chart.empty_dimension", f"Dimensao {chart.dimension_code!r} nao possui membros.")

        if chart.series_dimension_code:
            if chart.comparison:
                raise ChartWriteError(
                    "chart.series_and_comparison_not_supported",
                    "Contract V1 nao combina series_dimension_code e comparison no mesmo grafico.",
                )
            _dimension_fact_column(spec, chart, chart.series_dimension_code)
            if metric.allowed_dimensions and chart.series_dimension_code not in metric.allowed_dimensions:
                raise ChartWriteError(
                    "chart.series_dimension_not_allowed",
                    f"Metrica {metric.code!r} nao permite serie por {chart.series_dimension_code!r}.",
                )
            if not _sorted_members(spec, chart.series_dimension_code, None, None):
                raise ChartWriteError(
                    "chart.empty_series_dimension",
                    f"Dimensao de serie {chart.series_dimension_code!r} nao possui membros.",
                )

        if chart.comparison:
            parameter = parameters.get(chart.comparison)
            if parameter is None:
                raise ChartWriteError(
                    "chart.unknown_comparison_parameter",
                    f"Parametro de comparacao {chart.comparison!r} nao existe.",
                )
            comparison_dimension = parameter.values_source
            if comparison_dimension == chart.dimension_code:
                raise ChartWriteError(
                    "chart.comparison_matches_category_dimension",
                    "Parametro de comparacao nao pode controlar a mesma dimensao usada como categoria.",
                )
            if comparison_dimension not in binding.filter_parameters:
                raise ChartWriteError(
                    "chart.comparison_dimension_not_bound",
                    f"Dimensao {comparison_dimension!r} do parametro de comparacao nao participa do MetricBinding.",
                )

        if chart.window_parameter:
            window_parameter = parameters.get(chart.window_parameter)
            reference_parameter = parameters.get(chart.window_reference_parameter or "")
            if window_parameter is None:
                raise ChartWriteError("chart.unknown_window_parameter", f"Parametro de janela {chart.window_parameter!r} nao existe.")
            if reference_parameter is None:
                raise ChartWriteError("chart.unknown_window_reference_parameter", f"Parametro de referencia da janela {chart.window_reference_parameter!r} nao existe.")
            if reference_parameter.values_source != chart.dimension_code:
                raise ChartWriteError("chart.window_reference_dimension_mismatch", "Parametro de referencia da janela precisa usar a mesma dimensao das categorias.")
            if chart.window_max_categories is not None and chart.window_max_categories <= 0:
                raise ChartWriteError("chart.invalid_window_max_categories", "window_max_categories deve ser maior que zero.")
            if chart.series_dimension_code or chart.comparison:
                raise ChartWriteError("chart.window_multi_series_not_supported", "Janela dinamica V1 nao combina series_dimension_code/comparison.")

        if chart.chart_type in {ChartType.PIE, ChartType.DOUGHNUT} and (chart.series_dimension_code or chart.comparison):
            raise ChartWriteError(
                "chart.circular_multi_series_not_supported",
                "Pizza/rosca do Contract V1 aceita somente uma serie.",
            )

        for parameter_code in binding.filter_parameters.values():
            if parameter_code not in parameter_system.parameters:
                raise ChartWriteError(
                    "chart.missing_parameter_ref",
                    f"Parametro {parameter_code!r} nao foi materializado pelo sistema de PARAMETROS.",
                )


def _ensure_calc_sheet(workbook: Workbook, theme: ExcelOfficialTheme):
    if CALC_SHEET in workbook.sheetnames:
        return workbook[CALC_SHEET]
    ws = workbook.create_sheet(CALC_SHEET)
    apply_sheet_defaults(ws, theme=theme, zoom=85, freeze_panes="A4", tab_color=theme.palette.gray_600)
    ws.sheet_view.showGridLines = False
    ws["A1"] = TECHNICAL_LAYER_NOTICE
    apply_cell_role(ws["A1"], CellRole.SECTION, theme=theme)
    set_section_row_height(ws, 1, theme)
    ws["A2"] = "Séries auxiliares dos gráficos do Excel Oficial. Valores derivados; não editar."
    apply_cell_role(ws["A2"], CellRole.NOTE, theme=theme)
    return ws


def _next_calc_row(ws) -> int:
    if ws.max_row <= 2:
        return 4
    return ws.max_row + 2


def _write_helper_block(
    calc_ws,
    spec: WorkbookSpec,
    chart: ChartSpec,
    metric: MetricSpec,
    dataset_ref: DatasetRef,
    parameter_system: ParameterSystemRef,
    *,
    start_row: int,
    theme: ExcelOfficialTheme,
) -> tuple[str, int, int, int, int]:
    binding = _binding_by_metric(spec)[chart.metric_code]
    all_categories = _sorted_members(spec, chart.dimension_code, chart.sort, None if chart.window_parameter else chart.top_n)

    title_row = start_row
    header_row = start_row + 1
    data_start = start_row + 2

    calc_ws.cell(title_row, 1, f"GRÁFICO · {chart.code} · {chart.title}")
    apply_cell_role(calc_ws.cell(title_row, 1), CellRole.TECHNICAL, theme=theme)

    calc_ws.cell(header_row, 1, _dimension_by_code(spec)[chart.dimension_code].label)
    apply_cell_role(calc_ws.cell(header_row, 1), CellRole.TABLE_HEADER, theme=theme)

    comparison_mode = chart.comparison is not None
    series_members = ()
    if chart.series_dimension_code:
        series_members = _sorted_members(spec, chart.series_dimension_code, None, None)
        series_labels = [member.label for member in series_members]
    elif comparison_mode:
        series_labels = ["Referência", "Comparação"]
    else:
        series_labels = [metric.label]

    for idx, label in enumerate(series_labels, start=2):
        calc_ws.cell(header_row, idx, label)
        apply_cell_role(calc_ws.cell(header_row, idx), CellRole.TABLE_HEADER, theme=theme)

    key_col = 2 + len(series_labels)
    calc_ws.cell(header_row, key_col, f"CHAVE · {chart.dimension_code}")
    apply_cell_role(calc_ws.cell(header_row, key_col), CellRole.TABLE_HEADER, theme=theme)

    categories = all_categories
    catalog_end = data_start - 1
    if chart.window_parameter:
        slot_count = min(chart.window_max_categories or len(all_categories), len(all_categories))
        categories = all_categories[:slot_count]
        catalog_label_col = key_col + 1
        catalog_key_col = key_col + 2
        calc_ws.cell(header_row, catalog_label_col, f"CATÁLOGO · {_dimension_by_code(spec)[chart.dimension_code].label}")
        calc_ws.cell(header_row, catalog_key_col, f"CATÁLOGO CHAVE · {chart.dimension_code}")
        apply_cell_role(calc_ws.cell(header_row, catalog_label_col), CellRole.TABLE_HEADER, theme=theme)
        apply_cell_role(calc_ws.cell(header_row, catalog_key_col), CellRole.TABLE_HEADER, theme=theme)
        for index, member in enumerate(all_categories):
            crow = data_start + index
            calc_ws.cell(crow, catalog_label_col, member.label)
            calc_ws.cell(crow, catalog_key_col, member.key)
            apply_cell_role(calc_ws.cell(crow, catalog_label_col), CellRole.TECHNICAL, theme=theme)
            apply_cell_role(calc_ws.cell(crow, catalog_key_col), CellRole.TECHNICAL, theme=theme)
        catalog_end = data_start + len(all_categories) - 1
        qsheet = _excel_sheet(calc_ws.title)
        label_range = f"{qsheet}!${get_column_letter(catalog_label_col)}${data_start}:${get_column_letter(catalog_label_col)}${catalog_end}"
        key_range = f"{qsheet}!${get_column_letter(catalog_key_col)}${data_start}:${get_column_letter(catalog_key_col)}${catalog_end}"
        ref_name = parameter_system.parameters[chart.window_reference_parameter or ""].defined_name
        window_name = parameter_system.parameters[chart.window_parameter].defined_name
        ref_idx = f"MATCH({ref_name},{label_range},0)"
        all_literal = _formula_literal(chart.window_all_value)
        win_size = f"IF({window_name}={all_literal},{slot_count},MIN({slot_count},VALUE({window_name})))"
        start_idx = f"MAX(1,{ref_idx}-{win_size}+1)"

    value_format = _metric_number_format(metric, theme)
    for offset, member in enumerate(categories):
        row = data_start + offset
        if chart.window_parameter:
            slot = offset + 1
            visible_count = f"({ref_idx}-{start_idx}+1)"
            calc_ws.cell(row, 1, f'=IF({slot}<={visible_count},INDEX({label_range},{start_idx}+{slot}-1),"")')
            calc_ws.cell(row, key_col, f'=IF({slot}<={visible_count},INDEX({key_range},{start_idx}+{slot}-1),"")')
        else:
            calc_ws.cell(row, 1, member.label)
            calc_ws.cell(row, key_col, member.key)
        apply_cell_role(calc_ws.cell(row, 1), CellRole.TECHNICAL, theme=theme)
        apply_cell_role(calc_ws.cell(row, key_col), CellRole.TECHNICAL, theme=theme)
        category_key_expr = _absolute_cell(calc_ws.title, key_col, row)

        base_overrides = {chart.dimension_code: category_key_expr}
        if series_members:
            for series_index, series_member in enumerate(series_members, start=2):
                overrides = dict(base_overrides)
                overrides[chart.series_dimension_code or ""] = _formula_literal(series_member.key)
                formula = _metric_expression(
                    spec,
                    metric,
                    binding,
                    dataset_ref,
                    parameter_system,
                    dimension_overrides=overrides,
                )
                cell = calc_ws.cell(row, series_index, formula)
                apply_cell_role(cell, CellRole.DERIVED, number_format=value_format, theme=theme)
        elif comparison_mode:
            current = calc_ws.cell(
                row,
                2,
                _metric_expression(
                    spec,
                    metric,
                    binding,
                    dataset_ref,
                    parameter_system,
                    dimension_overrides=base_overrides,
                ),
            )
            apply_cell_role(current, CellRole.DERIVED, number_format=value_format, theme=theme)
            comparison = calc_ws.cell(
                row,
                3,
                _metric_expression(
                    spec,
                    metric,
                    binding,
                    dataset_ref,
                    parameter_system,
                    comparison_parameter=chart.comparison,
                    dimension_overrides=base_overrides,
                ),
            )
            apply_cell_role(comparison, CellRole.DERIVED, number_format=value_format, theme=theme)
        else:
            value = calc_ws.cell(
                row,
                2,
                _metric_expression(
                    spec,
                    metric,
                    binding,
                    dataset_ref,
                    parameter_system,
                    dimension_overrides=base_overrides,
                ),
            )
            apply_cell_role(value, CellRole.DERIVED, number_format=value_format, theme=theme)

    data_end = data_start + len(categories) - 1
    series_end_col = 1 + len(series_labels)
    calc_last_row = max(data_end, catalog_end)
    calc_last_col = key_col + (2 if chart.window_parameter else 0)
    calc_range = f"A{title_row}:{get_column_letter(calc_last_col)}{calc_last_row}"
    return calc_range, header_row, data_start, data_end, series_end_col


def _new_chart(chart_spec: ChartSpec, theme: ExcelOfficialTheme):
    if chart_spec.chart_type == ChartType.LINE:
        chart = LineChart()
    elif chart_spec.chart_type == ChartType.COLUMN:
        chart = BarChart()
        chart.type = "col"
        chart.grouping = "clustered"
    elif chart_spec.chart_type == ChartType.BAR:
        chart = BarChart()
        chart.type = "bar"
        chart.grouping = "clustered"
    elif chart_spec.chart_type == ChartType.PIE:
        chart = PieChart()
    elif chart_spec.chart_type == ChartType.DOUGHNUT:
        chart = DoughnutChart()
        chart.holeSize = theme.charts.doughnut_hole_size
    else:  # pragma: no cover - enum protects this branch
        raise ChartWriteError("chart.unsupported_type", f"Tipo de grafico nao suportado: {chart_spec.chart_type}.")

    chart.title = chart_spec.title
    chart.style = 10
    chart.width = theme.charts.width
    chart.height = theme.charts.height
    if hasattr(chart, "roundedCorners"):
        chart.roundedCorners = True
    return chart


def _apply_axis_semantics(chart, metric: MetricSpec, theme: ExcelOfficialTheme) -> None:
    if not hasattr(chart, "y_axis"):
        return
    axis = chart.y_axis
    if metric.valid_min is not None:
        axis.scaling.min = metric.valid_min
    if metric.valid_max is not None:
        axis.scaling.max = metric.valid_max

    if metric.unit == MetricUnit.NPS:
        axis.scaling.min = -100 if metric.valid_min is None else metric.valid_min
        axis.scaling.max = 100 if metric.valid_max is None else metric.valid_max
        axis.majorUnit = 20
        axis.numFmt = "0"
    elif metric.unit == MetricUnit.PERCENT:
        axis.scaling.min = 0 if metric.valid_min is None else metric.valid_min
        axis.scaling.max = 1 if metric.valid_max is None else metric.valid_max
        axis.majorUnit = 0.2
        axis.numFmt = "0%"
    elif metric.unit == MetricUnit.SCORE_0_10:
        axis.scaling.min = 0 if metric.valid_min is None else metric.valid_min
        axis.scaling.max = 10 if metric.valid_max is None else metric.valid_max
        axis.majorUnit = 2
        axis.numFmt = "0.0"
    elif metric.unit == MetricUnit.COUNT:
        axis.scaling.min = 0 if metric.valid_min is None else metric.valid_min
        axis.numFmt = "#,##0"
    elif metric.unit == MetricUnit.CURRENCY_BRL:
        axis.numFmt = 'R$ #,##0'

    # Keep the default gridline object for broad Excel compatibility.
    # Styling it requires creating drawing primitives that differ across openpyxl versions.
    if hasattr(chart, "x_axis"):
        chart.x_axis.delete = False
    axis.delete = False


def _style_series(chart, chart_type: ChartType, theme: ExcelOfficialTheme, category_count: int) -> None:
    colors = theme.charts.series_colors
    for index, series in enumerate(chart.series):
        color = colors[index % len(colors)]
        if chart_type == ChartType.LINE:
            series.graphicalProperties.line.solidFill = color
            series.graphicalProperties.line.width = theme.charts.line_width
            series.marker.symbol = "circle"
            series.marker.size = theme.charts.marker_size
            series.marker.graphicalProperties.solidFill = color
            series.marker.graphicalProperties.line.solidFill = color
        elif chart_type in {ChartType.COLUMN, ChartType.BAR}:
            series.graphicalProperties.solidFill = color
            series.graphicalProperties.line.solidFill = color

    if chart_type in {ChartType.PIE, ChartType.DOUGHNUT} and chart.series:
        points: list[DataPoint] = []
        for index in range(category_count):
            point = DataPoint(idx=index)
            color = colors[index % len(colors)]
            point.graphicalProperties.solidFill = color
            point.graphicalProperties.line.solidFill = "FFFFFF"
            points.append(point)
        chart.series[0].data_points = points


def _configure_labels_and_legend(chart, chart_spec: ChartSpec, series_count: int) -> None:
    if chart_spec.chart_type in {ChartType.PIE, ChartType.DOUGHNUT}:
        chart.legend.position = "r"
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showPercent = True
        chart.dataLabels.showLeaderLines = True
    elif chart_spec.chart_type in {ChartType.COLUMN, ChartType.BAR}:
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showVal = True
        if series_count <= 1:
            chart.legend = None
        else:
            chart.legend.position = "r"
    else:
        if series_count <= 1:
            chart.legend = None
        else:
            chart.legend.position = "r"


def write_dashboard_charts(
    workbook: Workbook,
    spec: WorkbookSpec,
    dataset_refs: Mapping[str, DatasetRef],
    parameter_system: ParameterSystemRef,
    dashboard_ref,
    *,
    theme: ExcelOfficialTheme = DEFAULT_THEME,
) -> ChartSystemRef:
    """Materialize dashboard charts and their auditable CALC helper ranges."""
    _preflight_chart_system(workbook, spec, dataset_refs, parameter_system, dashboard_ref)
    if not spec.dashboard.charts:
        return ChartSystemRef(calc_sheet_name=CALC_SHEET, charts=(), next_row=1)

    dashboard_ws = workbook[DASHBOARD_SHEET]
    calc_ws = _ensure_calc_sheet(workbook, theme)
    start_row = _next_calc_row(calc_ws)
    chart_refs: list[ChartRef] = []

    metrics = _metric_by_code(spec)
    for chart_spec in spec.dashboard.charts:
        metric = metrics[chart_spec.metric_code]
        dataset_ref = dataset_refs[chart_spec.dataset_code]
        calc_range, header_row, data_start, data_end, series_end_col = _write_helper_block(
            calc_ws,
            spec,
            chart_spec,
            metric,
            dataset_ref,
            parameter_system,
            start_row=start_row,
            theme=theme,
        )

        chart = _new_chart(chart_spec, theme)
        data = Reference(
            calc_ws,
            min_col=2,
            max_col=series_end_col,
            min_row=header_row,
            max_row=data_end,
        )
        categories = Reference(calc_ws, min_col=1, min_row=data_start, max_row=data_end)
        chart.add_data(data, titles_from_data=True, from_rows=False)
        chart.set_categories(categories)
        _apply_axis_semantics(chart, metric, theme)
        _style_series(chart, chart_spec.chart_type, theme, data_end - data_start + 1)
        _configure_labels_and_legend(chart, chart_spec, series_end_col - 1)

        anchor = dashboard_ref.chart_anchors[chart_spec.code]
        dashboard_ws.add_chart(chart, anchor)
        chart_refs.append(
            ChartRef(
                code=chart_spec.code,
                sheet_name=DASHBOARD_SHEET,
                anchor=anchor,
                chart_type=chart_spec.chart_type,
                calc_sheet_name=CALC_SHEET,
                calc_range=calc_range,
                category_count=data_end - data_start + 1,
                series_count=series_end_col - 1,
            )
        )
        start_row = data_end + 3

    # Keep the technical source readable/auditable without turning it into a presentation layer.
    max_col = max((calc_ws.max_column or 1), 1)
    for col in range(1, max_col + 1):
        letter = get_column_letter(col)
        calc_ws.column_dimensions[letter].width = 20 if col == 1 else 16
    for row in calc_ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
    protect_sheet(calc_ws)

    # openpyxl does not calculate formulas; force the desktop Excel client to
    # refresh helper formulas/chart caches when the snapshot is opened.
    try:
        workbook.calculation.calcMode = "auto"
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
    except AttributeError:  # pragma: no cover - compatibility with older openpyxl
        pass

    return ChartSystemRef(calc_sheet_name=CALC_SHEET, charts=tuple(chart_refs), next_row=start_row)


__all__ = [
    "CALC_SHEET",
    "ChartWriteError",
    "write_dashboard_charts",
]
