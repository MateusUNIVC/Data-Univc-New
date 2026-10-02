from __future__ import annotations

from dataclasses import dataclass
from typing import Any, Mapping

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter

from ..constants import REQUIRED_INSTITUTIONAL_SHEETS
from ..contract import (
    DatasetRef,
    MetricSpec,
    ParameterSystemRef,
    QualityCheckRef,
    QualityCheckSpec,
    QualityRef,
    QualitySeverity,
    WorkbookSpec,
)
from ..protection import protect_sheet
from ..styles import (
    apply_cell_role,
    apply_sheet_defaults,
    set_section_row_height,
    set_subtitle_row_height,
    set_table_header_row_height,
    set_title_row_height,
)
from ..theme import CellRole, DEFAULT_THEME, ExcelOfficialTheme
from ..validation import assert_valid_workbook_spec
from .dashboard import _metric_expression, _metric_number_format
from .parameters import _dimension_members

QUALITY_SHEET = REQUIRED_INSTITUTIONAL_SHEETS[3]


@dataclass(frozen=True, slots=True)
class QualityWriteError(ValueError):
    code: str
    message: str

    def __str__(self) -> str:
        return f"{self.code}: {self.message}"


def _metric_by_code(spec: WorkbookSpec) -> dict[str, MetricSpec]:
    return {item.code: item for item in spec.metrics}


def _binding_by_metric(spec: WorkbookSpec):
    return {item.metric_code: item for item in spec.metric_bindings}


def _dataset_by_code(spec: WorkbookSpec):
    return {item.code: item for item in spec.datasets}


def _dimension_by_code(spec: WorkbookSpec):
    return {item.code: item for item in spec.dimensions}


def _snapshot_field(spec: WorkbookSpec, field_name: str) -> Any:
    if not hasattr(spec.snapshot, field_name):
        raise QualityWriteError(
            "quality.unknown_snapshot_field",
            f"Campo de snapshot inexistente: {field_name!r}.",
        )
    return getattr(spec.snapshot, field_name)


def _is_present(value: Any) -> bool:
    if value is None:
        return False
    if isinstance(value, str):
        return bool(value.strip())
    if isinstance(value, (tuple, list, dict, set, frozenset)):
        return bool(value)
    return True


def _display_value(value: Any) -> Any:
    if value is None:
        return "Nao informado"
    if isinstance(value, (tuple, list, set, frozenset)):
        return ", ".join(str(item) for item in value) if value else "Nao informado"
    if isinstance(value, dict):
        return "; ".join(f"{key}={val}" for key, val in value.items()) if value else "Nao informado"
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value


def _dataset_required_missing_count(spec: WorkbookSpec, dataset_code: str) -> int:
    dataset = _dataset_by_code(spec)[dataset_code]
    required = [column for column in dataset.columns if not column.nullable]
    if not required:
        return 0
    missing = 0
    for row in dataset.rows:
        for column in required:
            value = row.get(column.code)
            if value is None or (isinstance(value, str) and not value.strip()):
                missing += 1
    return missing


def _check_category_rows(spec: WorkbookSpec):
    return (
        ("Integridade das bases", spec.quality.dataset_checks),
        ("Integridade dos indicadores", spec.quality.metric_checks),
        ("Cobertura", spec.quality.coverage_checks),
        ("Snapshot", spec.quality.snapshot_checks),
    )


def _preflight_quality(
    workbook: Workbook,
    spec: WorkbookSpec,
    dataset_refs: Mapping[str, DatasetRef],
    parameter_system: ParameterSystemRef,
) -> None:
    assert_valid_workbook_spec(spec)
    if any(name.casefold() == QUALITY_SHEET.casefold() for name in workbook.sheetnames):
        raise QualityWriteError("quality.sheet_exists", f"Aba {QUALITY_SHEET!r} ja existe.")

    datasets = _dataset_by_code(spec)
    metrics = _metric_by_code(spec)
    bindings = _binding_by_metric(spec)
    dimensions = _dimension_by_code(spec)

    for dataset in spec.datasets:
        if dataset.code not in dataset_refs:
            raise QualityWriteError(
                "quality.missing_dataset_ref",
                f"Dataset {dataset.code!r} nao foi materializado.",
            )
    for parameter in spec.parameters:
        if parameter.code not in parameter_system.parameters:
            raise QualityWriteError(
                "quality.missing_parameter_ref",
                f"Parametro {parameter.code!r} nao foi materializado.",
            )

    for _, checks in _check_category_rows(spec):
        for check in checks:
            if check.check_type in {"dataset_non_empty", "dataset_required_fields"}:
                if check.source_code not in datasets:
                    raise QualityWriteError("quality.unknown_dataset", f"Dataset inexistente: {check.source_code!r}.")
                if check.source_code not in dataset_refs:
                    raise QualityWriteError("quality.missing_dataset_ref", f"Dataset {check.source_code!r} nao foi materializado.")
            elif check.check_type in {"metric_has_data", "metric_valid_range"}:
                metric = metrics.get(check.source_code)
                binding = bindings.get(check.source_code)
                if metric is None:
                    raise QualityWriteError("quality.unknown_metric", f"Metrica inexistente: {check.source_code!r}.")
                if binding is None:
                    raise QualityWriteError("quality.metric_binding_required", f"Metrica {check.source_code!r} precisa de MetricBinding.")
                dataset_ref = dataset_refs.get(binding.dataset_code)
                if dataset_ref is None:
                    raise QualityWriteError("quality.missing_dataset_ref", f"Dataset {binding.dataset_code!r} da metrica nao foi materializado.")
                _metric_expression(spec, metric, binding, dataset_ref, parameter_system)
                if check.check_type == "metric_valid_range" and metric.valid_min is None and metric.valid_max is None:
                    raise QualityWriteError(
                        "quality.metric_range_required",
                        f"Metrica {metric.code!r} precisa de valid_min e/ou valid_max para metric_valid_range.",
                    )
            elif check.check_type == "dimension_non_empty":
                dimension = dimensions.get(check.source_code)
                if dimension is None:
                    raise QualityWriteError("quality.unknown_dimension", f"Dimensao inexistente: {check.source_code!r}.")
                _dimension_members(spec, dimension)
            elif check.check_type == "snapshot_field_present":
                _snapshot_field(spec, check.source_code)
            else:
                raise QualityWriteError(
                    "quality.unsupported_check_type",
                    f"Check type nao suportado no Contract V1: {check.check_type!r}.",
                )


def _status_conditional_formatting(ws, cell_range: str, theme: ExcelOfficialTheme) -> None:
    p = theme.palette
    rules = (
        ("OK", p.green_pale, p.green_dark),
        ("INFO", p.info_fill, p.blue),
        ("WARNING", p.warning_fill, p.orange),
        ("ERROR", p.error_fill, p.red),
        ("BLOCKING", p.error_fill, p.red),
        ("APTO", p.green_pale, p.green_dark),
        ("APTO COM OBSERVACOES", p.info_fill, p.blue),
        ("APTO COM ALERTAS", p.warning_fill, p.orange),
        ("REVISAR", p.error_fill, p.red),
        ("NAO UTILIZAR", p.error_fill, p.red),
        ("SEM CHECKS DECLARADOS", p.warning_fill, p.orange),
    )
    first = cell_range.split(":", 1)[0]
    for text, fill, font in rules:
        ws.conditional_formatting.add(
            cell_range,
            FormulaRule(
                formula=[f'{first}="{text}"'],
                fill=PatternFill(fill_type="solid", fgColor=fill),
                font=Font(name=theme.typography.family, bold=True, color=font),
            ),
        )


def _write_check(
    ws,
    *,
    row: int,
    category: str,
    check: QualityCheckSpec,
    spec: WorkbookSpec,
    dataset_refs: Mapping[str, DatasetRef],
    parameter_system: ParameterSystemRef,
    theme: ExcelOfficialTheme,
) -> QualityCheckRef:
    metrics = _metric_by_code(spec)
    bindings = _binding_by_metric(spec)
    datasets = _dataset_by_code(spec)
    dimensions = _dimension_by_code(spec)
    severity = check.severity.value

    ws.cell(row, 1, category)
    ws.cell(row, 2, check.label)
    ws.cell(row, 3, check.source_code)
    ws.cell(row, 4, severity)

    evidence = ws.cell(row, 5)
    result = ws.cell(row, 6)
    message = ws.cell(row, 7)

    if check.check_type == "dataset_non_empty":
        count = dataset_refs[check.source_code].row_count
        evidence.value = count
        evidence.number_format = theme.number_formats.integer
        result.value = "OK" if count > 0 else severity
    elif check.check_type == "dataset_required_fields":
        missing = _dataset_required_missing_count(spec, check.source_code)
        evidence.value = missing
        evidence.number_format = theme.number_formats.integer
        result.value = "OK" if missing == 0 else severity
    elif check.check_type in {"metric_has_data", "metric_valid_range"}:
        metric = metrics[check.source_code]
        binding = bindings[check.source_code]
        dataset_ref = dataset_refs[binding.dataset_code]
        evidence.value = _metric_expression(spec, metric, binding, dataset_ref, parameter_system)
        evidence.number_format = _metric_number_format(metric, theme)
        if check.check_type == "metric_has_data":
            result.value = f'=IF(E{row}="","{severity}","OK")'
        else:
            tests: list[str] = []
            if metric.valid_min is not None:
                tests.append(f"E{row}>={metric.valid_min}")
            if metric.valid_max is not None:
                tests.append(f"E{row}<={metric.valid_max}")
            condition = tests[0] if len(tests) == 1 else "AND(" + ",".join(tests) + ")"
            result.value = f'=IF(E{row}="","{severity}",IF({condition},"OK","{severity}"))'
    elif check.check_type == "dimension_non_empty":
        count = len(_dimension_members(spec, dimensions[check.source_code]))
        evidence.value = count
        evidence.number_format = theme.number_formats.integer
        result.value = "OK" if count > 0 else severity
    elif check.check_type == "snapshot_field_present":
        value = _snapshot_field(spec, check.source_code)
        evidence.value = _display_value(value)
        result.value = "OK" if _is_present(value) else severity
    else:  # pragma: no cover - preflight guards this branch
        raise QualityWriteError("quality.unsupported_check_type", check.check_type)

    message.value = f'=IF(F{row}="OK","{check.message_ok.replace(chr(34), chr(34) * 2)}","{check.message_error.replace(chr(34), chr(34) * 2)}")'

    for col in range(1, 8):
        apply_cell_role(ws.cell(row, col), CellRole.BODY, theme=theme)
        ws.cell(row, col).protection = Protection(locked=True)
    apply_cell_role(ws.cell(row, 4), CellRole.STATIC, theme=theme)
    apply_cell_role(evidence, CellRole.DERIVED if evidence.data_type == "f" else CellRole.IMPORTED, number_format=evidence.number_format, theme=theme)
    apply_cell_role(result, CellRole.DERIVED if result.data_type == "f" else CellRole.STATIC, theme=theme)
    apply_cell_role(message, CellRole.DERIVED, theme=theme)
    _status_conditional_formatting(ws, result.coordinate, theme)

    return QualityCheckRef(
        code=check.code,
        category=category,
        severity=check.severity,
        evidence_cell=evidence.coordinate,
        result_cell=result.coordinate,
        message_cell=message.coordinate,
    )


def _write_section_header(ws, row: int, title: str, *, end_col: int, theme: ExcelOfficialTheme) -> int:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=end_col)
    ws.cell(row, 1, title)
    apply_cell_role(ws.cell(row, 1), CellRole.SECTION, theme=theme)
    set_section_row_height(ws, row, theme)
    return row + 1


def _write_table_header(ws, row: int, headers: tuple[str, ...], *, theme: ExcelOfficialTheme) -> int:
    for col, label in enumerate(headers, 1):
        ws.cell(row, col, label)
        apply_cell_role(ws.cell(row, col), CellRole.TABLE_HEADER, theme=theme)
    set_table_header_row_height(ws, row, theme)
    return row + 1


def _countif_sum(ranges: list[str], value: str) -> str:
    if not ranges:
        return "0"
    return "+".join(f'COUNTIF({cell_range},"{value}")' for cell_range in ranges)


def write_quality_sheet(
    workbook: Workbook,
    spec: WorkbookSpec,
    dataset_refs: Mapping[str, DatasetRef],
    parameter_system: ParameterSystemRef,
    *,
    theme: ExcelOfficialTheme = DEFAULT_THEME,
) -> QualityRef:
    """Materialize the institutional Quality & Governance sheet.

    The sheet is a snapshot audit surface. Explicit quality checks may be static
    (dataset/snapshot/dimension) or parameter-sensitive (metric checks).
    """

    _preflight_quality(workbook, spec, dataset_refs, parameter_system)

    ws = workbook.create_sheet(QUALITY_SHEET)
    apply_sheet_defaults(ws, theme=theme, freeze_panes="A8", tab_color=theme.palette.orange)
    ws.sheet_state = "visible"

    widths = {"A": 23, "B": 34, "C": 24, "D": 16, "E": 24, "F": 18, "G": 48}
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width

    ws.merge_cells("A1:G1")
    ws["A1"] = "QUALIDADE E GOVERNAN\u00c7A"
    apply_cell_role(ws["A1"], CellRole.TITLE, theme=theme)
    set_title_row_height(ws, 1, theme)

    ws.merge_cells("A2:G2")
    ws["A2"] = f"Auditoria do snapshot | {spec.identity.directorate_label} | {spec.snapshot.export_id}"
    apply_cell_role(ws["A2"], CellRole.SUBTITLE, theme=theme)
    set_subtitle_row_height(ws, 2, theme)

    ws["A4"] = "Status geral"
    apply_cell_role(ws["A4"], CellRole.KPI_LABEL, theme=theme)
    ws.merge_cells("B4:G5")
    overall = ws["B4"]
    overall.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    overall.protection = Protection(locked=True)
    apply_cell_role(overall, CellRole.KPI_VALUE, theme=theme)

    row = 7
    check_refs: list[QualityCheckRef] = []
    check_result_ranges: list[str] = []

    row = _write_section_header(ws, row, "VERIFICA\u00c7\u00d5ES DE QUALIDADE", end_col=7, theme=theme)
    row = _write_table_header(
        ws,
        row,
        ("Categoria", "Verifica\u00e7\u00e3o", "Fonte", "Severidade", "Evid\u00eancia", "Resultado", "Mensagem"),
        theme=theme,
    )
    checks_start = row
    for category, checks in _check_category_rows(spec):
        for check in checks:
            check_refs.append(
                _write_check(
                    ws,
                    row=row,
                    category=category,
                    check=check,
                    spec=spec,
                    dataset_refs=dataset_refs,
                    parameter_system=parameter_system,
                    theme=theme,
                )
            )
            row += 1
    if row == checks_start:
        ws.cell(row, 1, "Nenhuma verifica\u00e7\u00e3o declarada.")
        apply_cell_role(ws.cell(row, 1), CellRole.NOTE, theme=theme)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        row += 1
    else:
        check_result_ranges.append(f"F{checks_start}:F{row - 1}")

    row += 1
    row = _write_section_header(ws, row, "COBERTURA", end_col=7, theme=theme)
    row = _write_table_header(ws, row, ("Item", "Valor", "Detalhe", "", "", "", ""), theme=theme)
    coverage_start = row
    coverage_rows: list[tuple[str, Any, str]] = [
        ("Per\u00edodo m\u00ednimo", spec.snapshot.minimum_period or "Nao informado", "Snapshot"),
        ("Per\u00edodo m\u00e1ximo", spec.snapshot.maximum_period or "Nao informado", "Snapshot"),
        ("Dimens\u00f5es interativas", len(spec.capabilities.interactive_dimensions), ", ".join(spec.capabilities.interactive_dimensions) or "Nenhuma"),
    ]
    for dimension in spec.dimensions:
        coverage_rows.append((dimension.label, len(_dimension_members(spec, dimension)), f"Dimensao {dimension.code}"))
    for label, value, detail in coverage_rows:
        ws.cell(row, 1, label)
        ws.cell(row, 2, value)
        ws.cell(row, 3, detail)
        for col in range(1, 8):
            apply_cell_role(ws.cell(row, col), CellRole.IMPORTED if col == 2 else CellRole.BODY, theme=theme)
        row += 1

    row += 1
    row = _write_section_header(ws, row, "FONTES E BASES", end_col=7, theme=theme)
    row = _write_table_header(
        ws,
        row,
        ("Base", "Origem", "Registros", "Sensibilidade", "Escopo autorizado", "Gr\u00e3o", "C\u00f3digo"),
        theme=theme,
    )
    sources_start = row
    for dataset in spec.datasets:
        ref = dataset_refs[dataset.code]
        values = (
            dataset.label,
            dataset.source or "Nao informado",
            ref.row_count,
            dataset.sensitivity,
            ", ".join(dataset.authorization_scope) if dataset.authorization_scope else ", ".join(spec.snapshot.authorization_scope),
            " x ".join(dataset.grain) if dataset.grain else "Nao informado",
            dataset.code,
        )
        for col, value in enumerate(values, 1):
            ws.cell(row, col, value)
            apply_cell_role(ws.cell(row, col), CellRole.IMPORTED, theme=theme)
        ws.cell(row, 3).number_format = theme.number_formats.integer
        row += 1
    sources_end = max(sources_start, row - 1)

    row += 1
    row = _write_section_header(ws, row, "SNAPSHOT", end_col=7, theme=theme)
    row = _write_table_header(ws, row, ("Campo", "Valor", "", "", "", "", ""), theme=theme)
    snapshot_rows = (
        ("Export ID", spec.snapshot.export_id),
        ("Gerado em", spec.snapshot.generated_at.isoformat()),
        ("Gerado por", spec.snapshot.generated_by or "Nao informado"),
        ("Sistema", spec.identity.system_name),
        ("Vers\u00e3o do sistema", spec.snapshot.system_version),
        ("Schema", str(spec.snapshot.schema_version)),
        ("Excel Official Contract", spec.snapshot.contract_version),
        ("Adapter", f"{spec.identity.adapter_code} v{spec.snapshot.adapter_version}"),
        ("Diretoria", f"{spec.identity.directorate_code} - {spec.identity.directorate_label}"),
        ("Escopo autorizado", ", ".join(spec.snapshot.authorization_scope) or "Nao informado"),
        ("Payload hash", spec.snapshot.payload_hash or "Nao informado"),
    )
    snapshot_start = row
    for label, value in snapshot_rows:
        ws.cell(row, 1, label)
        ws.cell(row, 2, value)
        apply_cell_role(ws.cell(row, 1), CellRole.STATIC, theme=theme)
        apply_cell_role(ws.cell(row, 2), CellRole.IMPORTED, theme=theme)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=7)
        row += 1

    row += 1
    row = _write_section_header(ws, row, "LIMITA\u00c7\u00d5ES CONHECIDAS", end_col=7, theme=theme)
    row = _write_table_header(
        ws,
        row,
        ("Severidade", "T\u00edtulo", "Descri\u00e7\u00e3o", "M\u00e9trica", "Dimens\u00e3o", "C\u00f3digo", ""),
        theme=theme,
    )
    limitation_start = row
    limitations = spec.quality.limitations + spec.limitations
    limitation_severity_range: str | None = None
    if limitations:
        for limitation in limitations:
            values = (
                limitation.severity.value,
                limitation.title,
                limitation.description,
                limitation.affected_metric or "",
                limitation.affected_dimension or "",
                limitation.code,
            )
            for col, value in enumerate(values, 1):
                ws.cell(row, col, value)
                apply_cell_role(ws.cell(row, col), CellRole.BODY, theme=theme)
            _status_conditional_formatting(ws, f"A{row}", theme)
            row += 1
        limitation_severity_range = f"A{limitation_start}:A{row - 1}"
    else:
        ws.cell(row, 1, "Nenhuma limita\u00e7\u00e3o declarada.")
        apply_cell_role(ws.cell(row, 1), CellRole.NOTE, theme=theme)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        row += 1

    status_ranges = list(check_result_ranges)
    limitation_ranges = [limitation_severity_range] if limitation_severity_range else []
    if not status_ranges and not limitation_ranges:
        overall.value = "SEM CHECKS DECLARADOS"
    else:
        blocking = _countif_sum(status_ranges + limitation_ranges, "BLOCKING")
        errors = _countif_sum(status_ranges + limitation_ranges, "ERROR")
        warnings = _countif_sum(status_ranges + limitation_ranges, "WARNING")
        infos = _countif_sum(status_ranges + limitation_ranges, "INFO")
        overall.value = (
            f'=IF(({blocking})>0,"NAO UTILIZAR",'
            f'IF(({errors})>0,"REVISAR",'
            f'IF(({warnings})>0,"APTO COM ALERTAS",'
            f'IF(({infos})>0,"APTO COM OBSERVACOES","APTO"))))'
        )
    _status_conditional_formatting(ws, overall.coordinate, theme)

    # Keep every non-input cell locked; this sheet is an audit surface.
    for row_cells in ws.iter_rows():
        for cell in row_cells:
            cell.protection = Protection(locked=True)
            if cell.alignment is None:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    protect_sheet(ws)

    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True

    return QualityRef(
        sheet_name=QUALITY_SHEET,
        overall_status_cell=overall.coordinate,
        checks=tuple(check_refs),
        coverage_range=f"A{coverage_start}:C{coverage_start + len(coverage_rows) - 1}",
        sources_range=f"A{sources_start}:G{sources_end}",
        snapshot_range=f"A{snapshot_start}:G{snapshot_start + len(snapshot_rows) - 1}",
        limitations_range=(f"A{limitation_start}:F{limitation_start + len(limitations) - 1}" if limitations else None),
    )


__all__ = [
    "QUALITY_SHEET",
    "QualityWriteError",
    "write_quality_sheet",
]
