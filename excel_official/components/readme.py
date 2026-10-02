from __future__ import annotations

from dataclasses import dataclass

from openpyxl import Workbook
from openpyxl.styles import Alignment, Protection

from ..constants import REQUIRED_INSTITUTIONAL_SHEETS, TECHNICAL_LAYER_NOTICE
from ..contract import WorkbookSpec
from ..protection import protect_sheet
from ..styles import (
    apply_cell_role,
    apply_sheet_defaults,
    set_section_row_height,
    set_subtitle_row_height,
    set_title_row_height,
)
from ..theme import CellRole, DEFAULT_THEME, ExcelOfficialTheme

README_SHEET = REQUIRED_INSTITUTIONAL_SHEETS[0]


@dataclass(frozen=True, slots=True)
class ReadmeRef:
    sheet_name: str
    snapshot_cell: str
    contract_cell: str


def write_readme_sheet(
    workbook: Workbook,
    spec: WorkbookSpec,
    *,
    theme: ExcelOfficialTheme = DEFAULT_THEME,
) -> ReadmeRef:
    if README_SHEET in workbook.sheetnames:
        raise ValueError(f"Aba {README_SHEET!r} ja existe.")

    ws = workbook.create_sheet(README_SHEET)
    apply_sheet_defaults(ws, theme=theme, zoom=90, freeze_panes="A4", tab_color=theme.palette.green_dark)
    ws.sheet_view.showGridLines = False
    widths = {"A": 24, "B": 26, "C": 26, "D": 26, "E": 26, "F": 26}
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width

    ws.merge_cells("A1:F1")
    ws["A1"] = "DATA UNIVC - EXCEL OFICIAL"
    apply_cell_role(ws["A1"], CellRole.TITLE, theme=theme)
    set_title_row_height(ws, 1, theme)

    ws.merge_cells("A2:F2")
    ws["A2"] = f"{spec.identity.workbook_title} | {spec.identity.directorate_label}"
    apply_cell_role(ws["A2"], CellRole.SUBTITLE, theme=theme)
    set_subtitle_row_height(ws, 2, theme)

    row = 4
    sections = (
        (
            "COMO USAR",
            (
                ("1. PARAMETROS", "Escolha os filtros amarelos permitidos para explorar o snapshot offline."),
                ("2. PAINEL", "Consulte KPIs e graficos calculados a partir do mesmo contrato de metricas."),
                ("3. MATRIZ", "Compare a metrica declarada entre duas dimensoes quando a matriz estiver configurada."),
                ("4. QUALIDADE E GOVERNANCA", "Verifique cobertura, fontes, checks, snapshot e limitacoes antes de interpretar os resultados."),
                ("5. PLANO_DE_ACAO", "Acoes oficiais sao somente leitura. Linhas LOCAL permanecem apenas neste arquivo."),
            ),
        ),
        (
            "REGRAS IMPORTANTES",
            (
                ("Snapshot offline", "Depois do download, este arquivo nao consulta banco, API ou outro arquivo Excel."),
                ("Fonte de verdade", "O Data UNIVC continua sendo a fonte oficial. O XLSX representa o snapshot autorizado no momento da exportacao."),
                ("Protecao", "Dados importados, formulas e camadas tecnicas ficam bloqueados contra edicao acidental."),
                ("Acompanhamento local", "Edicoes em linhas LOCAL do plano de acao nao sincronizam automaticamente com o Data UNIVC."),
                ("Camada tecnica", TECHNICAL_LAYER_NOTICE),
            ),
        ),
    )
    for title, items in sections:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        ws.cell(row, 1, title)
        apply_cell_role(ws.cell(row, 1), CellRole.SECTION, theme=theme)
        set_section_row_height(ws, row, theme)
        row += 1
        for label, description in items:
            ws.cell(row, 1, label)
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
            ws.cell(row, 2, description)
            apply_cell_role(ws.cell(row, 1), CellRole.STATIC, theme=theme)
            apply_cell_role(ws.cell(row, 2), CellRole.BODY, theme=theme)
            ws.cell(row, 2).alignment = Alignment(vertical="top", wrap_text=True)
            row += 1
        row += 1

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    ws.cell(row, 1, "IDENTIFICACAO DO SNAPSHOT")
    apply_cell_role(ws.cell(row, 1), CellRole.SECTION, theme=theme)
    set_section_row_height(ws, row, theme)
    row += 1
    snapshot_start = row
    metadata = (
        ("Export ID", spec.snapshot.export_id),
        ("Gerado em", spec.snapshot.generated_at.isoformat()),
        ("Gerado por", spec.snapshot.generated_by or "Nao informado"),
        ("Escopo autorizado", ", ".join(spec.snapshot.authorization_scope) or "Nao informado"),
        ("Periodo minimo", spec.snapshot.minimum_period or "Nao informado"),
        ("Periodo maximo", spec.snapshot.maximum_period or "Nao informado"),
        ("Versao do sistema", spec.snapshot.system_version),
        ("Schema", str(spec.snapshot.schema_version)),
        ("Excel Contract", str(spec.snapshot.contract_version)),
        ("Adapter", f"{spec.identity.adapter_code} v{spec.snapshot.adapter_version}"),
        ("Payload hash", spec.snapshot.payload_hash or "Nao informado"),
    )
    for label, value in metadata:
        ws.cell(row, 1, label)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
        ws.cell(row, 2, value)
        apply_cell_role(ws.cell(row, 1), CellRole.STATIC, theme=theme)
        apply_cell_role(ws.cell(row, 2), CellRole.IMPORTED, theme=theme)
        row += 1

    limitations = spec.quality.limitations + spec.limitations
    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    ws.cell(row, 1, "LIMITACOES DECLARADAS")
    apply_cell_role(ws.cell(row, 1), CellRole.SECTION, theme=theme)
    set_section_row_height(ws, row, theme)
    row += 1
    if limitations:
        for limitation in limitations:
            ws.cell(row, 1, limitation.severity.value)
            ws.cell(row, 2, limitation.title)
            ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=6)
            ws.cell(row, 3, limitation.description)
            apply_cell_role(ws.cell(row, 1), CellRole.WARNING if limitation.severity.value in {"WARNING", "ERROR", "BLOCKING"} else CellRole.NOTE, theme=theme)
            apply_cell_role(ws.cell(row, 2), CellRole.BODY, theme=theme)
            apply_cell_role(ws.cell(row, 3), CellRole.BODY, theme=theme)
            row += 1
    else:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        ws.cell(row, 1, "Nenhuma limitacao adicional foi declarada para este snapshot.")
        apply_cell_role(ws.cell(row, 1), CellRole.NOTE, theme=theme)

    for row_cells in ws.iter_rows():
        for cell in row_cells:
            cell.protection = Protection(locked=True)
    protect_sheet(ws)
    return ReadmeRef(sheet_name=README_SHEET, snapshot_cell=f"B{snapshot_start}", contract_cell=f"B{snapshot_start + 8}")


__all__ = ["README_SHEET", "ReadmeRef", "write_readme_sheet"]
