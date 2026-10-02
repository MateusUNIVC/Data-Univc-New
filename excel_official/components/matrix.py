from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from typing import Any, Mapping

from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from ..contract import DatasetRef, MatrixMetricOptionSpec, MatrixRef, MetricBinding, MetricSpec, ParameterSystemRef, WorkbookSpec
from ..protection import protect_sheet
from ..styles import (
    apply_cell_role,
    apply_sheet_defaults,
    set_section_row_height,
    set_subtitle_row_height,
    set_table_header_row_height,
    set_title_row_height,
)
from ..theme import CellRole, DEFAULT_THEME, ExcelOfficialTheme
from ..validation import assert_valid_workbook_spec
from .dashboard import _delta_number_format, _metric_expression, _metric_number_format
from .parameters import _dimension_members
from .targets import target_binding_by_metric, target_expression

MATRIX_SHEET = "MATRIZ"


@dataclass(frozen=True, slots=True)
class MatrixWriteError(ValueError):
    code: str
    message: str

    def __str__(self) -> str:
        return f"{self.code}: {self.message}"


def _metric_by_code(spec: WorkbookSpec) -> dict[str, MetricSpec]:
    return {item.code: item for item in spec.metrics}


def _binding_by_metric(spec: WorkbookSpec) -> dict[str, MetricBinding]:
    return {item.metric_code: item for item in spec.metric_bindings}


def _matrix_options(spec: WorkbookSpec):
    matrix = spec.matrix
    if matrix is None:
        return ()
    if matrix.metric_options:
        return matrix.metric_options
    metric = _metric_by_code(spec).get(matrix.metric_code)
    return (MatrixMetricOptionSpec(selector_value=metric.label if metric else matrix.metric_code, metric_code=matrix.metric_code),)


def _dataset_by_code(spec: WorkbookSpec):
    return {item.code: item for item in spec.datasets}


def _dimension_by_code(spec: WorkbookSpec):
    return {item.code: item for item in spec.dimensions}


def _formula_literal(value: Any) -> str:
    if value is None:
        return '""'
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        base = f"DATE({value.year},{value.month},{value.day})"
        if value.hour or value.minute or value.second:
            base += f"+TIME({value.hour},{value.minute},{value.second})"
        return base
    if isinstance(value, date):
        return f"DATE({value.year},{value.month},{value.day})"
    if isinstance(value, (int, float, Decimal)):
        return str(value)
    return '"' + str(value).replace('"', '""') + '"'


def _sorted_members(spec: WorkbookSpec, dimension_code: str, sorting: str | None):
    dimension = _dimension_by_code(spec)[dimension_code]
    members = list(_dimension_members(spec, dimension))
    if sorting in {None, "dimension_asc"}:
        return tuple(members)
    if sorting == "dimension_desc":
        members.reverse()
        return tuple(members)
    if sorting == "label_asc":
        members.sort(key=lambda item: str(item.label).casefold())
        return tuple(members)
    if sorting == "label_desc":
        members.sort(key=lambda item: str(item.label).casefold(), reverse=True)
        return tuple(members)
    raise MatrixWriteError("matrix.unsupported_sort", f"Ordenacao da matriz nao suportada no Contract V1: {sorting!r}.")


def _fact_column_for_dimension(spec: WorkbookSpec, binding: MetricBinding, dimension_code: str) -> str:
    dataset = _dataset_by_code(spec)[binding.dataset_code]
    columns = {column.code for column in dataset.columns}
    column_code = dataset.dimension_columns.get(dimension_code, dimension_code)
    if column_code not in columns:
        raise MatrixWriteError(
            "matrix.dimension_not_available_in_dataset",
            f"Dataset {dataset.code!r} nao possui coluna para a dimensao {dimension_code!r}.",
        )
    return column_code


def _preflight_matrix(
    workbook: Workbook,
    spec: WorkbookSpec,
    dataset_refs: Mapping[str, DatasetRef],
    parameter_system: ParameterSystemRef,
) -> None:
    assert_valid_workbook_spec(spec)
    matrix = spec.matrix
    if matrix is None:
        return
    if any(name.casefold() == MATRIX_SHEET.casefold() for name in workbook.sheetnames):
        raise MatrixWriteError("matrix.sheet_exists", f"Aba {MATRIX_SHEET!r} ja existe.")

    metrics = _metric_by_code(spec)
    bindings = _binding_by_metric(spec)
    dimensions = _dimension_by_code(spec)

    options = _matrix_options(spec)
    if not options:
        raise MatrixWriteError("matrix.metric_spec_required", "Matriz precisa declarar ao menos uma metrica.")
    if matrix.metric_selector_parameter:
        if matrix.metric_selector_parameter not in parameter_system.parameters:
            raise MatrixWriteError("matrix.selector_parameter_missing", f"Parametro seletor {matrix.metric_selector_parameter!r} nao foi materializado.")
    elif len(options) > 1:
        raise MatrixWriteError("matrix.selector_parameter_required", "Matriz com multiplas metricas exige metric_selector_parameter.")
    if matrix.row_dimension == matrix.column_dimension:
        raise MatrixWriteError("matrix.same_dimensions", "Dimensoes de linha e coluna precisam ser diferentes.")

    for option in options:
        metric = metrics.get(option.metric_code)
        if metric is None:
            raise MatrixWriteError("matrix.metric_spec_required", f"Matriz exige MetricSpec para {option.metric_code!r}.")
        binding = bindings.get(option.metric_code)
        if binding is None:
            raise MatrixWriteError("matrix.metric_binding_required", f"Matriz exige MetricBinding para {option.metric_code!r}.")
        if binding.dataset_code not in dataset_refs:
            raise MatrixWriteError("matrix.missing_dataset_ref", f"Dataset {binding.dataset_code!r} nao foi materializado.")
        for dimension_code, role in ((matrix.row_dimension, "row"), (matrix.column_dimension, "column")):
            if dimension_code not in dimensions:
                raise MatrixWriteError("matrix.unknown_dimension", f"Dimensao {dimension_code!r} ({role}) nao existe.")
            if dimension_code not in option.skip_dimensions and metric.allowed_dimensions and dimension_code not in metric.allowed_dimensions:
                raise MatrixWriteError("matrix.dimension_not_allowed", f"Metrica {metric.code!r} nao permite recorte por {dimension_code!r}.")
            if dimension_code not in option.skip_dimensions:
                _fact_column_for_dimension(spec, binding, dimension_code)
        for parameter_code in binding.filter_parameters.values():
            parameter = next((item for item in spec.parameters if item.code == parameter_code), None)
            if parameter and parameter.values_source in {matrix.row_dimension, matrix.column_dimension}:
                continue
            if parameter and parameter.values_source in option.skip_dimensions:
                continue
            if parameter_code not in parameter_system.parameters:
                raise MatrixWriteError("matrix.missing_parameter_ref", f"Parametro {parameter_code!r} nao foi materializado pelo sistema de PARAMETROS.")

    for dimension_code in (matrix.row_dimension, matrix.column_dimension):
        if dimension_code in dimensions and not _dimension_members(spec, dimensions[dimension_code]):
            raise MatrixWriteError("matrix.empty_dimension", f"Dimensao {dimension_code!r} nao possui membros.")

    if matrix.delta and len(_dimension_members(spec, dimensions[matrix.column_dimension])) < 2:
        raise MatrixWriteError("matrix.delta_requires_two_columns", "Delta exige ao menos dois membros na dimensao de colunas.")

    if matrix.empty_behavior not in {"blank", "zero"}:
        raise MatrixWriteError(
            "matrix.invalid_empty_behavior",
            f"empty_behavior nao suportado: {matrix.empty_behavior!r}.",
        )

def _matrix_formula(
    spec: WorkbookSpec,
    metric: MetricSpec,
    binding: MetricBinding,
    dataset_ref: DatasetRef,
    parameter_system: ParameterSystemRef,
    *,
    row_dimension: str,
    row_key: Any,
    column_dimension: str,
    column_key: Any,
    empty_behavior: str,
    skip_dimensions: tuple[str, ...] = (),
    display_scale: float = 1.0,
) -> str:
    formula = _metric_expression(
        spec,
        metric,
        binding,
        dataset_ref,
        parameter_system,
        dimension_overrides={
            row_dimension: _formula_literal(row_key),
            column_dimension: _formula_literal(column_key),
        },
        skip_dimensions=skip_dimensions,
    )
    expression = formula[1:]
    if display_scale != 1.0:
        expression = f'IF(({expression})="","",({expression})*{display_scale})'
    if empty_behavior == "blank":
        return f'={expression}'
    return f'=IF(({expression})="",0,({expression}))'


def _selected_formula(selector_name: str | None, options, formulas: list[str]) -> str:
    if len(options) == 1 or not selector_name:
        return formulas[0]
    expression = '""'
    for option, formula in reversed(list(zip(options, formulas))):
        inner = formula[1:] if formula.startswith("=") else formula
        expression = f'IF({selector_name}={_formula_literal(option.selector_value)},{inner},{expression})'
    return f'={expression}'


def _matrix_target_formula(
    spec: WorkbookSpec,
    options,
    dataset_refs: Mapping[str, DatasetRef],
    parameter_system: ParameterSystemRef,
    *,
    selector_name: str | None,
    row_key: Any,
    row_dimension: str,
) -> str | None:
    target_bindings = target_binding_by_metric(spec)
    formulas: list[str] = []
    any_target = False
    for option in options:
        binding = target_bindings.get(option.metric_code)
        if binding is None:
            formulas.append('=""')
            continue
        any_target = True
        # Convention: target datasets use the same semantic dimension code as
        # the matrix row when they expose a criterion with that name.
        overrides = {}
        if row_dimension in binding.criteria_parameters or row_dimension in binding.criteria_constants:
            overrides[row_dimension] = _formula_literal(row_key)
        formula = target_expression(
            spec, option.metric_code, dataset_refs, parameter_system,
            criteria_overrides=overrides,
        ) or '=""'
        if option.display_scale != 1.0:
            inner = formula[1:]
            formula = f'=IF(({inner})="","",({inner})*{option.display_scale})'
        formulas.append(formula)
    if not any_target:
        return None
    return _selected_formula(selector_name, options, formulas)


def write_matrix_sheet(
    workbook: Workbook,
    spec: WorkbookSpec,
    dataset_refs: Mapping[str, DatasetRef],
    parameter_system: ParameterSystemRef,
    *,
    theme: ExcelOfficialTheme = DEFAULT_THEME,
) -> MatrixRef | None:
    """Materialize the institutional row-dimension x column-dimension matrix."""
    _preflight_matrix(workbook, spec, dataset_refs, parameter_system)
    matrix = spec.matrix
    if matrix is None:
        return None

    metrics = _metric_by_code(spec)
    bindings = _binding_by_metric(spec)
    options = _matrix_options(spec)
    metric = metrics[options[0].metric_code]
    dimensions = _dimension_by_code(spec)
    selector_name = parameter_system.parameters[matrix.metric_selector_parameter].defined_name if matrix.metric_selector_parameter else None
    row_dimension = dimensions[matrix.row_dimension]
    column_dimension = dimensions[matrix.column_dimension]

    row_members = _sorted_members(spec, matrix.row_dimension, matrix.sorting)
    column_members = tuple(_dimension_members(spec, column_dimension))

    ws = workbook.create_sheet(MATRIX_SHEET)
    apply_sheet_defaults(ws, theme=theme, zoom=90, freeze_panes="B7", tab_color=theme.palette.green)
    ws.sheet_view.showGridLines = False

    has_dynamic_target = any(option.metric_code in target_binding_by_metric(spec) for option in options)
    total_columns = 1 + len(column_members) + (1 if (matrix.target is not None or has_dynamic_target) else 0) + (1 if matrix.delta else 0)
    total_columns = max(total_columns, 4)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_columns)
    ws.cell(1, 1, "MATRIZ DE INDICADORES")
    apply_cell_role(ws.cell(1, 1), CellRole.TITLE, theme=theme)
    set_title_row_height(ws, 1, theme)

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=total_columns)
    ws.cell(2, 1, f"{spec.identity.directorate_label} · {row_dimension.label} × {column_dimension.label}")
    apply_cell_role(ws.cell(2, 1), CellRole.SUBTITLE, theme=theme)
    set_subtitle_row_height(ws, 2, theme)

    ws.cell(4, 1, "Indicador")
    apply_cell_role(ws.cell(4, 1), CellRole.SECTION, theme=theme)
    ws.merge_cells(start_row=4, start_column=2, end_row=4, end_column=min(total_columns, 4))
    ws.cell(4, 2, f"={selector_name}" if selector_name else metric.label)
    apply_cell_role(ws.cell(4, 2), CellRole.CONTROL, theme=theme)
    set_section_row_height(ws, 4, theme)

    header_row = 6
    first_data_row = 7
    first_value_column = 2
    last_value_column = first_value_column + len(column_members) - 1

    ws.cell(header_row, 1, row_dimension.label)
    apply_cell_role(ws.cell(header_row, 1), CellRole.TABLE_HEADER, theme=theme)

    for offset, member in enumerate(column_members, start=first_value_column):
        ws.cell(header_row, offset, member.label)
        apply_cell_role(ws.cell(header_row, offset), CellRole.TABLE_HEADER, theme=theme)

    target_column = None
    delta_column = None
    next_column = last_value_column + 1
    if matrix.target is not None or has_dynamic_target:
        target_column = next_column
        ws.cell(header_row, target_column, "Meta")
        apply_cell_role(ws.cell(header_row, target_column), CellRole.TABLE_HEADER, theme=theme)
        next_column += 1
    if matrix.delta:
        delta_column = next_column
        ws.cell(header_row, delta_column, "Δ")
        apply_cell_role(ws.cell(header_row, delta_column), CellRole.TABLE_HEADER, theme=theme)

    set_table_header_row_height(ws, header_row, theme)
    if matrix.selector_number_format:
        value_format = matrix.selector_number_format
        delta_format = matrix.selector_number_format
    else:
        value_format = _metric_number_format(metric, theme)
        delta_format = _delta_number_format(metric, theme)

    for row_offset, row_member in enumerate(row_members):
        row = first_data_row + row_offset
        ws.cell(row, 1, row_member.label)
        apply_cell_role(ws.cell(row, 1), CellRole.IMPORTED, theme=theme)

        for col_offset, column_member in enumerate(column_members, start=first_value_column):
            formulas = []
            for option in options:
                option_metric = metrics[option.metric_code]
                option_binding = bindings[option.metric_code]
                formulas.append(_matrix_formula(
                    spec,
                    option_metric,
                    option_binding,
                    dataset_refs[option_binding.dataset_code],
                    parameter_system,
                    row_dimension=matrix.row_dimension,
                    row_key=row_member.key,
                    column_dimension=matrix.column_dimension,
                    column_key=column_member.key,
                    empty_behavior=matrix.empty_behavior,
                    skip_dimensions=option.skip_dimensions,
                    display_scale=option.display_scale,
                ))
            formula = _selected_formula(selector_name, options, formulas)
            cell = ws.cell(row, col_offset, formula)
            apply_cell_role(cell, CellRole.DERIVED, number_format=value_format, theme=theme)
            cell.alignment = Alignment(horizontal="center", vertical="center")

        if target_column is not None:
            dynamic_target = _matrix_target_formula(
                spec, options, dataset_refs, parameter_system,
                selector_name=selector_name, row_key=row_member.key, row_dimension=matrix.row_dimension,
            )
            target_value = matrix.target if matrix.target is not None else dynamic_target
            target_cell = ws.cell(row, target_column, target_value)
            apply_cell_role(target_cell, CellRole.DERIVED if dynamic_target is not None and matrix.target is None else CellRole.STATIC, number_format=value_format, theme=theme)
            target_cell.alignment = Alignment(horizontal="center", vertical="center")

        if delta_column is not None:
            latest = f"{get_column_letter(last_value_column)}{row}"
            previous = f"{get_column_letter(last_value_column - 1)}{row}"
            delta_formula = f'=IF(OR({latest}="",{previous}=""),"",{latest}-{previous})'
            delta_cell = ws.cell(row, delta_column, delta_formula)
            apply_cell_role(delta_cell, CellRole.DERIVED, number_format=delta_format, theme=theme)
            delta_cell.alignment = Alignment(horizontal="center", vertical="center")

    last_data_row = first_data_row + len(row_members) - 1
    ws.column_dimensions["A"].width = max(22, min(42, max(len(str(member.label)) for member in row_members) + 3))
    for column in range(first_value_column, last_value_column + 1):
        ws.column_dimensions[get_column_letter(column)].width = 14
    if target_column:
        ws.column_dimensions[get_column_letter(target_column)].width = 12
    if delta_column:
        ws.column_dimensions[get_column_letter(delta_column)].width = 12

    protect_sheet(ws)
    try:
        workbook.calculation.calcMode = "auto"
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
    except AttributeError:  # pragma: no cover
        pass

    value_range = f"{get_column_letter(first_value_column)}{first_data_row}:{get_column_letter(last_value_column)}{last_data_row}"
    return MatrixRef(
        sheet_name=MATRIX_SHEET,
        metric_code=matrix.metric_code,
        row_dimension=matrix.row_dimension,
        column_dimension=matrix.column_dimension,
        header_row=header_row,
        first_data_row=first_data_row,
        last_data_row=last_data_row,
        first_value_column=first_value_column,
        last_value_column=last_value_column,
        value_range=value_range,
        target_column=target_column,
        delta_column=delta_column,
    )


__all__ = [
    "MATRIX_SHEET",
    "MatrixWriteError",
    "write_matrix_sheet",
]
