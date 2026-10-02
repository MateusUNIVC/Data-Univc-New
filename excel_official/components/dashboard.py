from __future__ import annotations

from dataclasses import dataclass
from typing import Mapping

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter

from ..constants import REQUIRED_INSTITUTIONAL_SHEETS
from ..contract import (
    DashboardRef,
    DatasetRef,
    KpiRef,
    KpiSpec,
    MetricAggregation,
    MetricBinding,
    MetricSpec,
    MetricUnit,
    ParameterSystemRef,
    WorkbookSpec,
)
from ..protection import protect_sheet
from ..styles import (
    apply_cell_role,
    apply_range_role,
    apply_sheet_defaults,
    set_section_row_height,
    set_subtitle_row_height,
    set_title_row_height,
)
from ..theme import CellRole, DEFAULT_THEME, ExcelOfficialTheme
from ..validation import assert_valid_workbook_spec
from .targets import target_expression

DASHBOARD_SHEET = "PAINEL"


@dataclass(frozen=True, slots=True)
class DashboardWriteError(ValueError):
    code: str
    message: str

    def __str__(self) -> str:
        return f"{self.code}: {self.message}"


def _excel_sheet(sheet_name: str) -> str:
    return "'" + sheet_name.replace("'", "''") + "'"


def _excel_string(value: str) -> str:
    return '"' + value.replace('"', '""') + '"'


def _parameter_by_code(spec: WorkbookSpec):
    return {item.code: item for item in spec.parameters}


def _dataset_by_code(spec: WorkbookSpec):
    return {item.code: item for item in spec.datasets}


def _metric_by_code(spec: WorkbookSpec):
    return {item.code: item for item in spec.metrics}


def _binding_by_metric(spec: WorkbookSpec):
    return {item.metric_code: item for item in spec.metric_bindings}


def _numeric_column_range(dataset_ref: DatasetRef, column_code: str) -> str:
    if dataset_ref.row_count <= 0:
        raise DashboardWriteError("dashboard.empty_dataset", f"Dataset {dataset_ref.dataset_code!r} nao possui registros.")
    try:
        letter = dataset_ref.column_letter(column_code)
    except KeyError as exc:
        raise DashboardWriteError(
            "dashboard.unknown_dataset_column",
            f"Coluna {column_code!r} nao existe no DatasetRef {dataset_ref.dataset_code!r}.",
        ) from exc
    sheet = _excel_sheet(dataset_ref.sheet_name)
    return f"{sheet}!${letter}${dataset_ref.first_data_row}:${letter}${dataset_ref.last_data_row}"


def _dimension_column_code(spec: WorkbookSpec, binding: MetricBinding, dimension_code: str) -> str:
    dataset = _dataset_by_code(spec)[binding.dataset_code]
    return dataset.dimension_columns.get(dimension_code, dimension_code)


def _filter_masks(
    spec: WorkbookSpec,
    binding: MetricBinding,
    dataset_ref: DatasetRef,
    parameter_system: ParameterSystemRef,
    *,
    comparison_parameter: str | None = None,
    dimension_overrides: Mapping[str, str] | None = None,
    skip_dimensions: tuple[str, ...] = (),
) -> tuple[str, ...]:
    parameters = _parameter_by_code(spec)
    parameter_overrides: dict[str, str] = {}
    if comparison_parameter:
        comparison = parameters[comparison_parameter]
        parameter_overrides[comparison.values_source] = comparison_parameter

    exact_overrides = dict(dimension_overrides or {})
    skipped = set(skip_dimensions)
    masks: list[str] = []
    handled_dimensions: set[str] = set()
    for dimension_code, default_parameter_code in binding.filter_parameters.items():
        if dimension_code in skipped:
            continue
        column_code = _dimension_column_code(spec, binding, dimension_code)
        value_range = _numeric_column_range(dataset_ref, column_code)
        if dimension_code in exact_overrides:
            masks.append(f"--({value_range}={exact_overrides[dimension_code]})")
            handled_dimensions.add(dimension_code)
            continue

        parameter_code = parameter_overrides.get(dimension_code, default_parameter_code)
        parameter = parameters[parameter_code]
        try:
            parameter_ref = parameter_system.parameters[parameter_code]
        except KeyError as exc:
            raise DashboardWriteError(
                "dashboard.missing_parameter_ref",
                f"Parametro {parameter_code!r} nao foi materializado pelo sistema de PARAMETROS.",
            ) from exc
        p_name = parameter_ref.defined_name
        all_value = parameter.empty_option
        if all_value is None and not parameter.required:
            all_value = ""
        if all_value is None:
            masks.append(f"--({value_range}={p_name})")
        else:
            masks.append(f"--((({p_name}={_excel_string(str(all_value))})+({value_range}={p_name}))>0)")
        handled_dimensions.add(dimension_code)

    for dimension_code, expression in exact_overrides.items():
        if dimension_code in handled_dimensions or dimension_code in skipped:
            continue
        column_code = _dimension_column_code(spec, binding, dimension_code)
        value_range = _numeric_column_range(dataset_ref, column_code)
        masks.append(f"--({value_range}={expression})")
    return tuple(masks)


def _sum_expr(value_range: str, masks: tuple[str, ...]) -> str:
    if masks:
        return f"SUMPRODUCT({value_range},{','.join(masks)})"
    return f"SUM({value_range})"


def _match_count_expr(dataset_ref: DatasetRef, masks: tuple[str, ...]) -> str:
    if dataset_ref.row_count <= 0:
        return "0"
    if masks:
        return f"SUMPRODUCT({','.join(masks)})"
    return str(dataset_ref.row_count)


def _metric_expression(
    spec: WorkbookSpec,
    metric: MetricSpec,
    binding: MetricBinding,
    dataset_ref: DatasetRef,
    parameter_system: ParameterSystemRef,
    *,
    comparison_parameter: str | None = None,
    dimension_overrides: Mapping[str, str] | None = None,
    skip_dimensions: tuple[str, ...] = (),
) -> str:
    if dataset_ref.row_count <= 0:
        return '=""'

    masks = _filter_masks(
        spec,
        binding,
        dataset_ref,
        parameter_system,
        comparison_parameter=comparison_parameter,
        dimension_overrides=dimension_overrides,
        skip_dimensions=skip_dimensions,
    )
    match_count = _match_count_expr(dataset_ref, masks)

    def component(name: str) -> str:
        try:
            column_code = binding.components[name]
        except KeyError as exc:
            raise DashboardWriteError(
                "dashboard.missing_metric_component",
                f"Metrica {metric.code!r} exige componente {name!r}.",
            ) from exc
        return _sum_expr(_numeric_column_range(dataset_ref, column_code), masks)

    aggregation = metric.aggregation
    if aggregation in {MetricAggregation.VALUE, MetricAggregation.SUM}:
        if not binding.value_column:
            raise DashboardWriteError(
                "dashboard.value_column_required",
                f"Metrica {metric.code!r} exige value_column.",
            )
        calculation = _sum_expr(_numeric_column_range(dataset_ref, binding.value_column), masks)
    elif aggregation == MetricAggregation.COUNT:
        calculation = match_count
    elif aggregation == MetricAggregation.RATIO:
        calculation = f"({component('numerator')}/{component('denominator')})"
    elif aggregation == MetricAggregation.AVERAGE:
        calculation = f"({component('sum')}/{component('count')})"
    elif aggregation == MetricAggregation.WEIGHTED_AVERAGE:
        calculation = f"({component('weighted_sum')}/{component('weight')})"
    elif aggregation == MetricAggregation.NPS:
        respondents = component("respondents")
        calculation = f"((({component('promoters')})/{respondents})-(({component('detractors')})/{respondents}))*100"
    else:  # pragma: no cover - enum keeps this defensive branch unreachable
        raise DashboardWriteError(
            "dashboard.unsupported_aggregation",
            f"Agregacao nao suportada no Dashboard V1: {aggregation}.",
        )

    invalidating = tuple(component(name) for name in metric.invalid_when_positive_components)
    if invalidating:
        tests = ",".join(f"({expression}>0)" for expression in invalidating)
        return f'=IF({match_count}=0,"",IF(OR({tests}),"",IFERROR({calculation},"")))'
    return f'=IF({match_count}=0,"",IFERROR({calculation},""))'


def _metric_number_format(metric: MetricSpec, theme: ExcelOfficialTheme) -> str:
    if metric.number_format:
        return metric.number_format
    formats = theme.number_formats
    precision = max(0, min(metric.display_precision, 6))
    decimal = "0" if precision == 0 else "0." + ("0" * precision)
    if metric.unit == MetricUnit.COUNT:
        return formats.integer
    if metric.unit == MetricUnit.PERCENT:
        return "0%" if precision == 0 else "0." + ("0" * precision) + "%"
    if metric.unit == MetricUnit.NPS:
        return decimal
    if metric.unit == MetricUnit.SCORE_0_10:
        return decimal
    if metric.unit == MetricUnit.CURRENCY_BRL:
        return formats.currency_brl
    return decimal


def _delta_number_format(metric: MetricSpec, theme: ExcelOfficialTheme) -> str:
    if metric.unit == MetricUnit.PERCENT:
        precision = max(0, min(metric.display_precision, 6))
        zeros = "" if precision == 0 else "." + ("0" * precision)
        return f"+0{zeros}%;-0{zeros}%;0{zeros}%"
    if metric.unit == MetricUnit.CURRENCY_BRL:
        return theme.number_formats.currency_brl
    precision = max(0, min(metric.display_precision, 6))
    zeros = "" if precision == 0 else "." + ("0" * precision)
    return f"+0{zeros};-0{zeros};0{zeros}"


def _preflight_dashboard(
    workbook: Workbook,
    spec: WorkbookSpec,
    dataset_refs: Mapping[str, DatasetRef],
    parameter_system: ParameterSystemRef,
) -> None:
    assert_valid_workbook_spec(spec)
    if any(name.casefold() == DASHBOARD_SHEET.casefold() for name in workbook.sheetnames):
        raise DashboardWriteError("dashboard.sheet_exists", f"Aba {DASHBOARD_SHEET!r} ja existe.")

    metrics = _metric_by_code(spec)
    bindings = _binding_by_metric(spec)
    parameters = _parameter_by_code(spec)

    for parameter in spec.parameters:
        if parameter.code not in parameter_system.parameters:
            raise DashboardWriteError(
                "dashboard.missing_parameter_ref",
                f"Parametro {parameter.code!r} nao foi materializado pelo sistema de PARAMETROS.",
            )

    for kpi in spec.dashboard.kpis:
        metric = metrics.get(kpi.metric_code)
        if metric is None:
            raise DashboardWriteError(
                "dashboard.metric_spec_required",
                f"KPI {kpi.metric_code!r} precisa de MetricSpec executavel no 01E.",
            )
        binding = bindings.get(kpi.metric_code)
        if binding is None:
            raise DashboardWriteError(
                "dashboard.metric_binding_required",
                f"KPI {kpi.metric_code!r} precisa de MetricBinding.",
            )
        dataset_ref = dataset_refs.get(binding.dataset_code)
        if dataset_ref is None:
            raise DashboardWriteError(
                "dashboard.missing_dataset_ref",
                f"Dataset {binding.dataset_code!r} da metrica {kpi.metric_code!r} nao foi materializado.",
            )
        if kpi.comparison and kpi.comparison not in parameters:
            raise DashboardWriteError(
                "dashboard.unknown_comparison_parameter",
                f"Parametro de comparacao inexistente: {kpi.comparison!r}.",
            )
        # Build both expressions during preflight. This catches missing columns/components
        # before the PAINEL sheet is created.
        _metric_expression(spec, metric, binding, dataset_ref, parameter_system)
        if kpi.comparison:
            _metric_expression(
                spec,
                metric,
                binding,
                dataset_ref,
                parameter_system,
                comparison_parameter=kpi.comparison,
            )


def _set_outline_border(ws, min_row: int, max_row: int, min_col: int, max_col: int, color: str) -> None:
    side = Side(style="thin", color=color)
    for row in range(min_row, max_row + 1):
        for col in range(min_col, max_col + 1):
            cell = ws.cell(row, col)
            cell.border = Border(
                left=side if col == min_col else cell.border.left,
                right=side if col == max_col else cell.border.right,
                top=side if row == min_row else cell.border.top,
                bottom=side if row == max_row else cell.border.bottom,
            )


def _write_context(
    ws,
    spec: WorkbookSpec,
    parameter_system: ParameterSystemRef,
    *,
    start_row: int,
    theme: ExcelOfficialTheme,
) -> tuple[dict[str, str], int]:
    params = sorted(enumerate(spec.parameters), key=lambda item: (item[1].display_order, item[0]))
    if not params:
        return {}, start_row

    total_cols = theme.dashboard.total_columns
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=total_cols)
    ws.cell(start_row, 1, "CONTEXTO ATUAL")
    apply_cell_role(ws.cell(start_row, 1), CellRole.SECTION, theme=theme)
    set_section_row_height(ws, start_row, theme)

    refs: dict[str, str] = {}
    row = start_row + 1
    per_row = theme.dashboard.context_items_per_row
    width = theme.dashboard.context_columns_per_item
    for index, (_, parameter) in enumerate(params):
        slot = index % per_row
        block_row = row + (index // per_row) * 2
        col1 = 1 + slot * width
        col2 = min(total_cols, col1 + width - 1)
        ws.merge_cells(start_row=block_row, start_column=col1, end_row=block_row, end_column=col2)
        label_cell = ws.cell(block_row, col1, parameter.label)
        apply_cell_role(label_cell, CellRole.CONTROL, theme=theme)
        ws.merge_cells(start_row=block_row + 1, start_column=col1, end_row=block_row + 1, end_column=col2)
        value_cell = ws.cell(block_row + 1, col1, f"={parameter_system.parameters[parameter.code].defined_name}")
        apply_cell_role(value_cell, CellRole.VISUALIZATION, theme=theme)
        refs[parameter.code] = value_cell.coordinate

    rows_used = ((len(params) - 1) // per_row + 1) * 2
    return refs, row + rows_used


def _write_kpi_card(
    ws,
    *,
    spec: WorkbookSpec,
    kpi: KpiSpec,
    metric: MetricSpec,
    binding: MetricBinding,
    dataset_ref: DatasetRef,
    dataset_refs: Mapping[str, DatasetRef],
    parameter_system: ParameterSystemRef,
    row: int,
    col: int,
    theme: ExcelOfficialTheme,
) -> KpiRef:
    width = theme.dashboard.kpi_columns
    height = theme.dashboard.kpi_rows
    end_col = col + width - 1
    end_row = row + height - 1
    p = theme.palette

    # Fill/protect the complete card before merging regions.
    for rr in range(row, end_row + 1):
        for cc in range(col, end_col + 1):
            cell = ws.cell(rr, cc)
            cell.fill = PatternFill(fill_type="solid", fgColor=p.green_light)
            cell.protection = Protection(locked=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=end_col)
    ws.merge_cells(start_row=row + 1, start_column=col, end_row=row + 2, end_column=end_col)

    label = kpi.label_override or metric.label
    label_cell = ws.cell(row, col, label)
    apply_cell_role(label_cell, CellRole.KPI_LABEL, theme=theme)

    value_formula = _metric_expression(spec, metric, binding, dataset_ref, parameter_system)
    value_cell = ws.cell(row + 1, col, value_formula)
    apply_cell_role(
        value_cell,
        CellRole.KPI_VALUE,
        number_format=_metric_number_format(metric, theme),
        theme=theme,
    )

    comparison_cell_ref: str | None = None
    delta_cell_ref: str | None = None
    if kpi.comparison:
        compare_formula = _metric_expression(
            spec,
            metric,
            binding,
            dataset_ref,
            parameter_system,
            comparison_parameter=kpi.comparison,
        )
        compare_label = ws.cell(row + 3, col, "Comparação")
        apply_cell_role(compare_label, CellRole.NOTE, theme=theme)
        compare_cell = ws.cell(row + 3, col + 1, compare_formula)
        apply_cell_role(
            compare_cell,
            CellRole.DERIVED,
            number_format=_metric_number_format(metric, theme),
            theme=theme,
        )
        delta_cell = ws.cell(
            row + 3,
            col + 2,
            f'=IF(OR({value_cell.coordinate}="",{compare_cell.coordinate}=""),"",{value_cell.coordinate}-{compare_cell.coordinate})',
        )
        apply_cell_role(
            delta_cell,
            CellRole.DERIVED,
            number_format=_delta_number_format(metric, theme),
            theme=theme,
        )
        comparison_cell_ref = compare_cell.coordinate
        delta_cell_ref = delta_cell.coordinate
    else:
        ws.merge_cells(start_row=row + 3, start_column=col, end_row=row + 3, end_column=end_col)
        note = ws.cell(row + 3, col, "Atualiza automaticamente pelos filtros do arquivo")
        apply_cell_role(note, CellRole.NOTE, theme=theme)
        note.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    target_cell_ref: str | None = None
    dynamic_target = target_expression(spec, metric.code, dataset_refs, parameter_system)
    if kpi.target is not None or dynamic_target is not None:
        target_label = ws.cell(row + 4, col, "Meta")
        apply_cell_role(target_label, CellRole.NOTE, theme=theme)
        ws.merge_cells(start_row=row + 4, start_column=col + 1, end_row=row + 4, end_column=end_col)
        target_value = kpi.target if kpi.target is not None else dynamic_target
        target_cell = ws.cell(row + 4, col + 1, target_value)
        apply_cell_role(
            target_cell,
            CellRole.DERIVED if dynamic_target is not None and kpi.target is None else CellRole.STATIC,
            number_format=_metric_number_format(metric, theme),
            theme=theme,
        )
        target_cell.alignment = Alignment(horizontal="center", vertical="center")
        target_cell_ref = target_cell.coordinate
    else:
        ws.merge_cells(start_row=row + 4, start_column=col, end_row=row + 4, end_column=end_col)
        provenance = ws.cell(row + 4, col, "Data UNIVC · snapshot offline")
        apply_cell_role(provenance, CellRole.NOTE, theme=theme)
        provenance.alignment = Alignment(horizontal="center", vertical="center")

    _set_outline_border(ws, row, end_row, col, end_col, p.gray_300)
    return KpiRef(
        metric_code=kpi.metric_code,
        label_cell=label_cell.coordinate,
        value_cell=value_cell.coordinate,
        card_range=f"{get_column_letter(col)}{row}:{get_column_letter(end_col)}{end_row}",
        comparison_cell=comparison_cell_ref,
        delta_cell=delta_cell_ref,
        target_cell=target_cell_ref,
    )


def _chart_anchors(spec: WorkbookSpec, *, start_row: int, theme: ExcelOfficialTheme) -> tuple[dict[str, str], int]:
    anchors: dict[str, str] = {}
    if not spec.dashboard.charts:
        return anchors, start_row
    for index, chart in enumerate(spec.dashboard.charts):
        row = start_row + (index // 2) * theme.dashboard.chart_rows
        col = 1 if index % 2 == 0 else 1 + theme.dashboard.chart_columns
        anchors[chart.code] = f"{get_column_letter(col)}{row}"
    rows = ((len(spec.dashboard.charts) - 1) // 2 + 1) * theme.dashboard.chart_rows
    return anchors, start_row + rows


def write_dashboard_sheet(
    workbook: Workbook,
    spec: WorkbookSpec,
    dataset_refs: Mapping[str, DatasetRef],
    parameter_system: ParameterSystemRef,
    *,
    theme: ExcelOfficialTheme = DEFAULT_THEME,
) -> DashboardRef:
    """Materialize the institutional PAINEL shell and KPI cards.

    Phase 01E intentionally creates no chart objects. It reserves stable chart
    anchors for 01F and returns them in DashboardRef.
    """
    _preflight_dashboard(workbook, spec, dataset_refs, parameter_system)

    ws = workbook.create_sheet(DASHBOARD_SHEET)
    apply_sheet_defaults(ws, theme=theme, zoom=85, freeze_panes="A4", tab_color=theme.palette.green)
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"

    total_cols = theme.dashboard.total_columns
    last_letter = get_column_letter(total_cols)
    ws.merge_cells(f"A1:{last_letter}1")
    ws["A1"] = spec.dashboard.title or spec.identity.workbook_title
    apply_cell_role(ws["A1"], CellRole.TITLE, theme=theme)
    set_title_row_height(ws, 1, theme)

    ws.merge_cells(f"A2:{last_letter}2")
    ws["A2"] = (
        f"{spec.identity.directorate_label} · Snapshot {spec.snapshot.export_id} · "
        "ajuste os filtros na aba PARAMETROS"
    )
    apply_cell_role(ws["A2"], CellRole.SUBTITLE, theme=theme)
    set_subtitle_row_height(ws, 2, theme)

    context_cells, next_row = _write_context(
        ws,
        spec,
        parameter_system,
        start_row=4,
        theme=theme,
    )
    next_row += 1

    kpi_refs: list[KpiRef] = []
    if spec.dashboard.kpis:
        ws.merge_cells(start_row=next_row, start_column=1, end_row=next_row, end_column=total_cols)
        ws.cell(next_row, 1, "RESUMO EXECUTIVO DOS KPIs")
        apply_cell_role(ws.cell(next_row, 1), CellRole.SECTION, theme=theme)
        set_section_row_height(ws, next_row, theme)
        card_start = next_row + 2

        metrics = _metric_by_code(spec)
        bindings = _binding_by_metric(spec)
        for index, kpi in enumerate(sorted(spec.dashboard.kpis, key=lambda item: item.priority)):
            card_row = card_start + (index // theme.dashboard.kpis_per_row) * (
                theme.dashboard.kpi_rows + theme.dashboard.kpi_gap_rows
            )
            card_col = 1 + (index % theme.dashboard.kpis_per_row) * theme.dashboard.kpi_columns
            binding = bindings[kpi.metric_code]
            ref = _write_kpi_card(
                ws,
                spec=spec,
                kpi=kpi,
                metric=metrics[kpi.metric_code],
                binding=binding,
                dataset_ref=dataset_refs[binding.dataset_code],
                dataset_refs=dataset_refs,
                parameter_system=parameter_system,
                row=card_row,
                col=card_col,
                theme=theme,
            )
            kpi_refs.append(ref)

        kpi_rows = ((len(spec.dashboard.kpis) - 1) // theme.dashboard.kpis_per_row + 1) * (
            theme.dashboard.kpi_rows + theme.dashboard.kpi_gap_rows
        )
        next_row = card_start + kpi_rows

    chart_anchors: dict[str, str] = {}
    if spec.dashboard.charts:
        ws.merge_cells(start_row=next_row, start_column=1, end_row=next_row, end_column=total_cols)
        ws.cell(next_row, 1, "VISUALIZAÇÕES")
        apply_cell_role(ws.cell(next_row, 1), CellRole.SECTION, theme=theme)
        set_section_row_height(ws, next_row, theme)
        chart_anchors, next_row = _chart_anchors(spec, start_row=next_row + 2, theme=theme)

    if spec.dashboard.attention_blocks:
        attention_row = next_row + 1
        ws.merge_cells(start_row=attention_row, start_column=1, end_row=attention_row, end_column=total_cols)
        ws.cell(attention_row, 1, "ATENÇÕES / GOVERNANÇA")
        apply_cell_role(ws.cell(attention_row, 1), CellRole.SECTION, theme=theme)
        set_section_row_height(ws, attention_row, theme)
        for offset, text in enumerate(spec.dashboard.attention_blocks, start=1):
            ws.merge_cells(
                start_row=attention_row + offset,
                start_column=1,
                end_row=attention_row + offset,
                end_column=total_cols,
            )
            ws.cell(attention_row + offset, 1, text)
            apply_cell_role(ws.cell(attention_row + offset, 1), CellRole.NOTE, theme=theme)
        next_row = attention_row + len(spec.dashboard.attention_blocks) + 1

    for column in range(1, total_cols + 1):
        ws.column_dimensions[get_column_letter(column)].width = 13.5

    protect_sheet(ws)
    return DashboardRef(
        sheet_name=DASHBOARD_SHEET,
        context_cells=context_cells,
        kpis=tuple(kpi_refs),
        chart_anchors=chart_anchors,
        next_row=next_row,
    )


__all__ = [
    "DASHBOARD_SHEET",
    "DashboardWriteError",
    "write_dashboard_sheet",
]
