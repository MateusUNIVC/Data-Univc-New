from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from ..constants import TECHNICAL_LAYER_NOTICE
from ..contract import ColumnDataType, ColumnSpec, DatasetRef, DatasetSpec
from ..protection import protect_sheet
from ..styles import (
    apply_cell_role,
    apply_sheet_defaults,
    set_subtitle_row_height,
    set_table_header_row_height,
    set_title_row_height,
)
from ..theme import CellRole, DEFAULT_THEME, ExcelOfficialTheme


@dataclass(frozen=True, slots=True)
class DatasetWriteError(ValueError):
    code: str
    dataset_code: str
    message: str
    row_index: int | None = None
    column_code: str | None = None

    def __str__(self) -> str:
        location: list[str] = [self.dataset_code]
        if self.row_index is not None:
            location.append(f"row={self.row_index}")
        if self.column_code:
            location.append(f"column={self.column_code}")
        return f"{self.code} [{' / '.join(location)}]: {self.message}"


def _workbook_table_names(workbook: Workbook) -> set[str]:
    names: set[str] = set()
    for worksheet in workbook.worksheets:
        names.update(name.casefold() for name in worksheet.tables.keys())
    return names


def _expected_python_types(data_type: ColumnDataType) -> tuple[type, ...]:
    if data_type is ColumnDataType.TEXT:
        return (str,)
    if data_type is ColumnDataType.INTEGER:
        return (int,)
    if data_type is ColumnDataType.DECIMAL:
        return (int, float, Decimal)
    if data_type is ColumnDataType.DATE:
        return (date,)
    if data_type is ColumnDataType.DATETIME:
        return (datetime,)
    if data_type is ColumnDataType.BOOLEAN:
        return (bool,)
    return (object,)


def _validate_value(dataset: DatasetSpec, column: ColumnSpec, value: Any, row_index: int) -> None:
    if value is None:
        if not column.nullable:
            raise DatasetWriteError(
                code="dataset.null_not_allowed",
                dataset_code=dataset.code,
                row_index=row_index,
                column_code=column.code,
                message="Valor nulo em coluna declarada como nullable=False.",
            )
        return

    expected = _expected_python_types(column.data_type)
    valid = isinstance(value, expected)

    # bool is an int subclass in Python, but must not leak into numeric columns.
    if column.data_type in {ColumnDataType.INTEGER, ColumnDataType.DECIMAL} and isinstance(value, bool):
        valid = False
    # datetime is a date subclass, but a DATE column must remain date-only.
    if column.data_type is ColumnDataType.DATE and isinstance(value, datetime):
        valid = False

    if not valid:
        expected_text = ", ".join(item.__name__ for item in expected)
        raise DatasetWriteError(
            code="dataset.invalid_value_type",
            dataset_code=dataset.code,
            row_index=row_index,
            column_code=column.code,
            message=f"Esperado {column.data_type.value} ({expected_text}), recebido {type(value).__name__}.",
        )


def _validate_dataset_rows(dataset: DatasetSpec) -> None:
    columns = {column.code: column for column in dataset.columns}
    for row_index, row in enumerate(dataset.rows, start=1):
        unknown = set(row) - set(columns)
        if unknown:
            raise DatasetWriteError(
                code="dataset.unknown_column",
                dataset_code=dataset.code,
                row_index=row_index,
                message=f"Colunas nao declaradas: {', '.join(sorted(unknown))}.",
            )
        for column in dataset.columns:
            _validate_value(dataset, column, row.get(column.code), row_index)


def _preflight_dataset(workbook: Workbook, dataset: DatasetSpec) -> None:
    if not dataset.columns:
        raise DatasetWriteError(
            code="dataset.no_columns",
            dataset_code=dataset.code,
            message="Dataset precisa declarar ao menos uma coluna.",
        )
    labels = [column.label.casefold() for column in dataset.columns]
    if len(labels) != len(set(labels)):
        raise DatasetWriteError(
            code="dataset.duplicate_header",
            dataset_code=dataset.code,
            message="Excel Table nao aceita cabecalhos duplicados no contrato oficial.",
        )
    _validate_dataset_rows(dataset)
    if dataset.table_name.casefold() in _workbook_table_names(workbook):
        raise DatasetWriteError(
            code="dataset.table_name_collision",
            dataset_code=dataset.code,
            message=f"Ja existe uma Excel Table chamada {dataset.table_name!r} no workbook.",
        )


def _number_format_for_column(column: ColumnSpec, theme: ExcelOfficialTheme) -> str | None:
    if column.number_format:
        return column.number_format

    semantic = (column.semantic_type or "").casefold()
    formats = theme.number_formats

    if semantic in {"currency", "currency_brl", "brl"}:
        return formats.currency_brl
    if semantic in {"percent", "percentage"}:
        return formats.percentage
    if semantic in {"percentage_point", "percentage_points", "pp"}:
        return formats.percentage_points
    if semantic in {"nps", "score_0_10", "score"}:
        return formats.decimal_1
    if semantic in {"count", "integer"}:
        return formats.integer

    if column.data_type is ColumnDataType.INTEGER:
        return formats.integer
    if column.data_type is ColumnDataType.DECIMAL:
        return formats.decimal_2
    if column.data_type is ColumnDataType.DATE:
        return formats.date
    if column.data_type is ColumnDataType.DATETIME:
        return formats.datetime
    return None


def _write_literal_value(cell: Cell, value: Any, column: ColumnSpec) -> None:
    cell.value = value
    # Imported TEXT is data, never an Excel formula. Explicitly force string
    # storage so source values such as "=1+1" survive as literal text.
    if value is not None and column.data_type is ColumnDataType.TEXT:
        cell.data_type = "s"


def _column_width(dataset: DatasetSpec, column: ColumnSpec, theme: ExcelOfficialTheme) -> float:
    layout = theme.layout
    longest = len(column.label)
    for row in dataset.rows[: layout.table_width_scan_rows]:
        value = row.get(column.code)
        if value is not None:
            longest = max(longest, len(str(value)))
    proposed = longest + layout.table_column_padding
    return float(max(layout.table_min_column_width, min(proposed, layout.table_max_column_width)))


def _target_range_is_empty(
    worksheet: Worksheet,
    *,
    start_row: int,
    start_column: int,
    end_row: int,
    end_column: int,
) -> bool:
    for row in worksheet.iter_rows(
        min_row=start_row,
        max_row=end_row,
        min_col=start_column,
        max_col=end_column,
    ):
        if any(cell.value is not None for cell in row):
            return False
    return True


def write_dataset_table(
    worksheet: Worksheet,
    dataset: DatasetSpec,
    *,
    start_row: int = 1,
    start_column: int = 1,
    theme: ExcelOfficialTheme = DEFAULT_THEME,
    freeze_below_header: bool = False,
    protect: bool = False,
) -> DatasetRef:
    """Write one DatasetSpec as an Excel Table and return its structural ref.

    This is the low-level table primitive. It never creates a worksheet and it
    refuses to overwrite existing cells or reuse a table name elsewhere in the
    workbook. Rows are validated before any worksheet mutation.
    """
    if start_row < 1 or start_column < 1:
        raise DatasetWriteError(
            code="dataset.invalid_anchor",
            dataset_code=dataset.code,
            message="start_row e start_column devem ser >= 1.",
        )
    workbook = worksheet.parent
    _preflight_dataset(workbook, dataset)

    header_row = start_row
    first_data_row = header_row + 1
    data_row_count = len(dataset.rows)
    # OOXML/Excel interoperability is more robust with one structural empty data
    # row when the snapshot contains zero records. DatasetRef keeps row_count=0,
    # so downstream components never treat that structural row as a fact.
    physical_data_rows = max(1, data_row_count)
    last_data_row = first_data_row + physical_data_rows - 1
    end_column = start_column + len(dataset.columns) - 1

    if not _target_range_is_empty(
        worksheet,
        start_row=header_row,
        start_column=start_column,
        end_row=last_data_row,
        end_column=end_column,
    ):
        raise DatasetWriteError(
            code="dataset.target_not_empty",
            dataset_code=dataset.code,
            message="Area de destino possui conteudo e nao sera sobrescrita.",
        )

    header_cells: list[Cell] = []
    for offset, column in enumerate(dataset.columns):
        cell = worksheet.cell(header_row, start_column + offset, column.label)
        apply_cell_role(cell, CellRole.TABLE_HEADER, theme=theme)
        header_cells.append(cell)
    set_table_header_row_height(worksheet, header_row, theme)

    for row_offset, row in enumerate(dataset.rows):
        excel_row = first_data_row + row_offset
        for column_offset, column in enumerate(dataset.columns):
            cell = worksheet.cell(excel_row, start_column + column_offset)
            value = row.get(column.code)
            _write_literal_value(cell, value, column)
            role = CellRole.TECHNICAL if dataset.technical or column.technical else CellRole.IMPORTED
            apply_cell_role(
                cell,
                role,
                number_format=_number_format_for_column(column, theme),
                theme=theme,
            )

    if data_row_count == 0:
        for column_offset, column in enumerate(dataset.columns):
            cell = worksheet.cell(first_data_row, start_column + column_offset)
            role = CellRole.TECHNICAL if dataset.technical or column.technical else CellRole.IMPORTED
            apply_cell_role(
                cell,
                role,
                number_format=_number_format_for_column(column, theme),
                theme=theme,
            )

    start_letter = get_column_letter(start_column)
    end_letter = get_column_letter(end_column)
    table_ref = f"{start_letter}{header_row}:{end_letter}{last_data_row}"
    table = Table(displayName=dataset.table_name, ref=table_ref)
    table.tableStyleInfo = TableStyleInfo(
        name=theme.tables.style_name,
        showFirstColumn=theme.tables.show_first_column,
        showLastColumn=theme.tables.show_last_column,
        showRowStripes=theme.tables.show_row_stripes,
        showColumnStripes=theme.tables.show_column_stripes,
    )
    worksheet.add_table(table)

    column_letters: dict[str, str] = {}
    for offset, column in enumerate(dataset.columns):
        column_index = start_column + offset
        letter = get_column_letter(column_index)
        column_letters[column.code] = letter
        worksheet.column_dimensions[letter].width = _column_width(dataset, column, theme)

    if freeze_below_header:
        worksheet.freeze_panes = f"{start_letter}{first_data_row}"
    if protect:
        protect_sheet(worksheet)

    logical_last_data_row = header_row if data_row_count == 0 else header_row + data_row_count
    return DatasetRef(
        dataset_code=dataset.code,
        sheet_name=worksheet.title,
        table_name=dataset.table_name,
        table_ref=table_ref,
        header_row=header_row,
        first_data_row=first_data_row,
        last_data_row=logical_last_data_row,
        physical_last_row=last_data_row,
        first_column=start_column,
        last_column=end_column,
        row_count=data_row_count,
        column_letters=column_letters,
        has_structural_empty_row=data_row_count == 0,
    )


def write_dataset_sheet(
    workbook: Workbook,
    dataset: DatasetSpec,
    *,
    theme: ExcelOfficialTheme = DEFAULT_THEME,
    protect: bool = True,
) -> DatasetRef:
    """Create a standardized standalone dataset sheet.

    Row 1 is the dataset title, row 2 documents its provenance/technical role,
    and row 4 begins the Excel Table. This layout intentionally mirrors the
    strongest parts of the existing Academic golden-master without reusing its
    legacy builder implementation.
    """
    if any(ws.title.casefold() == dataset.sheet_name.casefold() for ws in workbook.worksheets):
        raise DatasetWriteError(
            code="dataset.sheet_name_collision",
            dataset_code=dataset.code,
            message=f"Ja existe uma aba chamada {dataset.sheet_name!r} no workbook.",
        )
    # Validate all data and workbook-wide table identity before mutating the workbook.
    _preflight_dataset(workbook, dataset)

    worksheet = workbook.create_sheet(dataset.sheet_name)
    tab_color = theme.palette.gray_600 if dataset.technical else theme.palette.green
    apply_sheet_defaults(
        worksheet,
        theme=theme,
        freeze_panes="A5",
        tab_color=tab_color,
    )

    end_column = max(1, len(dataset.columns))
    end_letter = get_column_letter(end_column)
    worksheet.merge_cells(f"A1:{end_letter}1")
    worksheet["A1"] = dataset.label
    apply_cell_role(worksheet["A1"], CellRole.TITLE, theme=theme)
    set_title_row_height(worksheet, 1, theme)

    worksheet.merge_cells(f"A2:{end_letter}2")
    if dataset.technical:
        subtitle = TECHNICAL_LAYER_NOTICE
        role = CellRole.TECHNICAL
    else:
        subtitle = dataset.source.strip() or "Snapshot importado do Data UNIVC. Dados protegidos contra edicao acidental."
        role = CellRole.SUBTITLE
    worksheet["A2"] = subtitle
    apply_cell_role(worksheet["A2"], role, theme=theme)
    set_subtitle_row_height(worksheet, 2, theme)

    ref = write_dataset_table(
        worksheet,
        dataset,
        start_row=4,
        start_column=1,
        theme=theme,
        freeze_below_header=True,
        protect=False,
    )
    if protect:
        protect_sheet(worksheet)
    return ref


__all__ = [
    "DatasetWriteError",
    "write_dataset_sheet",
    "write_dataset_table",
]
