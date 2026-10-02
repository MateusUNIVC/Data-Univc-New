from __future__ import annotations

from copy import copy
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from typing import Any, Iterable, Mapping

from openpyxl import Workbook
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from ..constants import TECHNICAL_LAYER_NOTICE
from ..contract import (
    ColumnDataType,
    DatasetSpec,
    DimensionSpec,
    ParameterRef,
    ParameterSpec,
    ParameterSystemRef,
    Scalar,
    SupportListRef,
    WorkbookSpec,
)
from ..protection import protect_sheet
from ..styles import (
    apply_cell_role,
    apply_sheet_defaults,
    set_section_row_height,
    set_subtitle_row_height,
    set_title_row_height,
)
from ..theme import CellRole, DEFAULT_THEME, ExcelOfficialTheme
from ..validation import assert_valid_workbook_spec
from .names import add_defined_name, parameter_defined_name, support_list_defined_name

PARAMETERS_SHEET = "PARAMETROS"
SUPPORT_LISTS_SHEET = "LISTAS DE APOIO"


@dataclass(frozen=True, slots=True)
class ParameterWriteError(ValueError):
    code: str
    parameter_code: str
    message: str

    def __str__(self) -> str:
        return f"{self.code} [{self.parameter_code}]: {self.message}"


@dataclass(frozen=True, slots=True)
class _DimensionMember:
    key: Scalar
    label: Scalar
    sort_value: Any
    parent_key: Scalar = None


def _sorted_parameters(parameters: Iterable[ParameterSpec]) -> tuple[ParameterSpec, ...]:
    indexed = list(enumerate(parameters))
    indexed.sort(key=lambda item: (item[1].display_order, item[0]))
    return tuple(parameter for _, parameter in indexed)


def _dataset_by_code(spec: WorkbookSpec) -> dict[str, DatasetSpec]:
    return {dataset.code: dataset for dataset in spec.datasets}


def _dimension_by_code(spec: WorkbookSpec) -> dict[str, DimensionSpec]:
    return {dimension.code: dimension for dimension in spec.dimensions}


def _parameter_by_code(spec: WorkbookSpec) -> dict[str, ParameterSpec]:
    return {parameter.code: parameter for parameter in spec.parameters}


def _sort_key(value: Any, fallback: Any) -> tuple[int, Any, str]:
    if value is None:
        value = fallback
    if isinstance(value, (int, float, Decimal, date, datetime)) and not isinstance(value, bool):
        return (0, value, str(fallback).casefold())
    return (1, str(value).casefold(), str(fallback).casefold())


def _dimension_members(spec: WorkbookSpec, dimension: DimensionSpec) -> tuple[_DimensionMember, ...]:
    dataset = _dataset_by_code(spec)[dimension.dataset_code]
    unique: dict[Any, _DimensionMember] = {}
    parent_col = dimension.parent_key_column or dimension.parent_dimension

    for row in dataset.rows:
        key = row.get(dimension.key_column)
        label = row.get(dimension.label_column)
        if key is None or label is None:
            continue
        parent_key = row.get(parent_col) if parent_col else None
        member = _DimensionMember(
            key=key,
            label=label,
            sort_value=row.get(dimension.sort_order_column) if dimension.sort_order_column else None,
            parent_key=parent_key,
        )
        # A child member may legitimately repeat under different parents.
        unique_key = (parent_key, key) if dimension.parent_dimension else key
        existing = unique.get(unique_key)
        if existing is not None and str(existing.label) != str(label):
            raise ParameterWriteError(
                code="dimension.ambiguous_key",
                parameter_code=dimension.code,
                message=f"Chave {key!r} possui mais de um rotulo na dimensao {dimension.code!r}.",
            )
        unique[unique_key] = member

    members = list(unique.values())
    members.sort(key=lambda item: _sort_key(item.sort_value, item.label))
    return tuple(members)


def _direct_dataset_values(dataset: DatasetSpec, parameter: ParameterSpec) -> tuple[Scalar, ...]:
    if parameter.values_column:
        column_code = parameter.values_column
    else:
        visible = [column.code for column in dataset.columns if column.visible and not column.technical]
        if len(visible) != 1:
            raise ParameterWriteError(
                code="parameter.dataset_values_column_required",
                parameter_code=parameter.code,
                message=(
                    f"Dataset {dataset.code!r} possui {len(visible)} colunas visiveis; "
                    "declare values_column para usar o dataset diretamente como lista."
                ),
            )
        column_code = visible[0]

    declared = {column.code for column in dataset.columns}
    if column_code not in declared:
        raise ParameterWriteError(
            code="parameter.unknown_values_column",
            parameter_code=parameter.code,
            message=f"Coluna {column_code!r} nao existe em {dataset.code!r}.",
        )

    seen: set[tuple[type, Any]] = set()
    values: list[Scalar] = []
    for row in dataset.rows:
        value = row.get(column_code)
        if value is None:
            continue
        marker = (type(value), value)
        if marker not in seen:
            seen.add(marker)
            values.append(value)
    return tuple(values)


def _independent_values(spec: WorkbookSpec, parameter: ParameterSpec) -> tuple[Scalar, ...]:
    dimensions = _dimension_by_code(spec)
    datasets = _dataset_by_code(spec)
    if parameter.values_source in dimensions:
        return tuple(member.label for member in _dimension_members(spec, dimensions[parameter.values_source]))
    if parameter.values_source in datasets:
        return _direct_dataset_values(datasets[parameter.values_source], parameter)
    return ()


def _with_empty_option(values: Iterable[Scalar], empty_option: str | None) -> tuple[Scalar, ...]:
    output: list[Scalar] = []
    if empty_option is not None:
        output.append(empty_option)
    for value in values:
        if value not in output:
            output.append(value)
    return tuple(output)


def _resolve_initial_value(spec: WorkbookSpec, parameter: ParameterSpec) -> Scalar:
    if parameter.code in spec.initial_state.values:
        return spec.initial_state.values[parameter.code]
    if parameter.initial_value is not None:
        return parameter.initial_value
    if parameter.code in spec.snapshot.initial_scope:
        return spec.snapshot.initial_scope[parameter.code]
    return None


def _validate_parameter_scalar_type(parameter: ParameterSpec, value: Scalar) -> None:
    if value is None or parameter.data_type is None:
        return
    expected: dict[ColumnDataType, tuple[type, ...]] = {
        ColumnDataType.TEXT: (str,),
        ColumnDataType.INTEGER: (int,),
        ColumnDataType.DECIMAL: (int, float, Decimal),
        ColumnDataType.DATE: (date,),
        ColumnDataType.DATETIME: (datetime,),
        ColumnDataType.BOOLEAN: (bool,),
    }
    valid = isinstance(value, expected[parameter.data_type])
    if parameter.data_type in {ColumnDataType.INTEGER, ColumnDataType.DECIMAL} and isinstance(value, bool):
        valid = False
    if parameter.data_type is ColumnDataType.DATE and isinstance(value, datetime):
        valid = False
    if not valid:
        raise ParameterWriteError(
            code="parameter.invalid_value_type",
            parameter_code=parameter.code,
            message=f"Valor inicial {value!r} nao corresponde a {parameter.data_type.value}.",
        )


def _generated_defined_names(spec: WorkbookSpec) -> tuple[str, ...]:
    names: list[str] = []
    for parameter in spec.parameters:
        names.append(parameter_defined_name(parameter.code))
        if parameter.values_source:
            names.append(support_list_defined_name(parameter.code))
    return tuple(names)


def _preflight_parameter_system(workbook: Workbook, spec: WorkbookSpec) -> None:
    assert_valid_workbook_spec(spec)
    existing_sheets = {worksheet.title.casefold() for worksheet in workbook.worksheets}
    for sheet_name in (PARAMETERS_SHEET, SUPPORT_LISTS_SHEET):
        if sheet_name.casefold() in existing_sheets:
            raise ParameterWriteError(
                code="parameter.sheet_collision",
                parameter_code="__system__",
                message=f"Aba {sheet_name!r} ja existe no workbook.",
            )

    generated = _generated_defined_names(spec)
    folded = [name.casefold() for name in generated]
    if len(folded) != len(set(folded)):
        raise ParameterWriteError(
            code="parameter.generated_name_collision",
            parameter_code="__system__",
            message="Codigos de parametros geram nomes definidos Excel duplicados apos normalizacao.",
        )
    existing_names = {str(name).casefold() for name in workbook.defined_names.keys()}
    conflict = sorted(name for name in generated if name.casefold() in existing_names)
    if conflict:
        raise ParameterWriteError(
            code="parameter.defined_name_collision",
            parameter_code="__system__",
            message=f"Nomes definidos ja existentes: {', '.join(conflict)}.",
        )

    for parameter in _sorted_parameters(spec.parameters):
        value = _resolve_initial_value(spec, parameter)
        _validate_parameter_scalar_type(parameter, value)
        if parameter.values_source and not parameter.depends_on:
            values = _with_empty_option(_independent_values(spec, parameter), parameter.empty_option)
            if not values:
                raise ParameterWriteError(
                    code="parameter.empty_values_source",
                    parameter_code=parameter.code,
                    message="Fonte de valores do dropdown nao possui opcoes.",
                )
        if parameter.values_source and parameter.depends_on:
            dimension = _dimension_by_code(spec)[parameter.values_source]
            members = _dimension_members(spec, dimension)
            if not members:
                raise ParameterWriteError(
                    code="parameter.empty_values_source",
                    parameter_code=parameter.code,
                    message="Fonte de valores do dropdown dependente nao possui opcoes.",
                )
            parent_labels, _ = _parent_label_maps(spec, dimension.parent_dimension or "")
            unknown = {str(member.parent_key) for member in members if member.parent_key not in parent_labels}
            if unknown:
                raise ParameterWriteError(
                    code="parameter.unknown_parent_member",
                    parameter_code=parameter.code,
                    message=f"Valores pai inexistentes: {', '.join(sorted(unknown, key=str.casefold))}.",
                )
        _validate_initial_value(spec, parameter, None, value)


def _parameter_number_format(parameter: ParameterSpec, value: Scalar, theme: ExcelOfficialTheme) -> str | None:
    if parameter.number_format:
        return parameter.number_format
    data_type = parameter.data_type
    if data_type is ColumnDataType.DATE or (data_type is None and isinstance(value, date) and not isinstance(value, datetime)):
        return theme.number_formats.date
    if data_type is ColumnDataType.DATETIME or (data_type is None and isinstance(value, datetime)):
        return theme.number_formats.datetime
    if data_type is ColumnDataType.INTEGER:
        return theme.number_formats.integer
    if data_type is ColumnDataType.DECIMAL:
        return theme.number_formats.decimal_2
    return None


def _ensure_new_sheet(workbook: Workbook, sheet_name: str, *, parameter_code: str) -> Worksheet:
    if any(ws.title.casefold() == sheet_name.casefold() for ws in workbook.worksheets):
        raise ParameterWriteError(
            code="parameter.sheet_collision",
            parameter_code=parameter_code,
            message=f"Aba {sheet_name!r} ja existe no workbook.",
        )
    return workbook.create_sheet(sheet_name)


def _write_support_header(
    worksheet: Worksheet,
    *,
    start_column: int,
    headers: tuple[str, ...],
    theme: ExcelOfficialTheme,
) -> None:
    for offset, header in enumerate(headers):
        cell = worksheet.cell(4, start_column + offset, header)
        apply_cell_role(cell, CellRole.TABLE_HEADER, theme=theme)
        worksheet.column_dimensions[get_column_letter(start_column + offset)].width = max(18.0, min(42.0, len(header) + 4.0))


def _write_support_value(cell, value: Scalar, *, theme: ExcelOfficialTheme) -> None:
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"
    apply_cell_role(cell, CellRole.TECHNICAL, theme=theme)


def _parent_label_maps(spec: WorkbookSpec, parent_dimension_code: str) -> tuple[dict[Any, Scalar], dict[Any, int]]:
    dimensions = _dimension_by_code(spec)
    parent = dimensions[parent_dimension_code]
    members = _dimension_members(spec, parent)
    labels: dict[Any, Scalar] = {}
    order: dict[Any, int] = {}
    for index, member in enumerate(members):
        labels[member.key] = member.label
        labels.setdefault(member.label, member.label)
        order[member.key] = index
        order.setdefault(member.label, index)
    return labels, order


def write_support_lists(
    workbook: Workbook,
    spec: WorkbookSpec,
    *,
    theme: ExcelOfficialTheme = DEFAULT_THEME,
) -> Mapping[str, SupportListRef]:
    """Materialize dropdown sources in the visible technical support sheet.

    Independent lists become static defined ranges. A one-parent dependent list
    becomes a non-volatile dynamic defined range based on INDEX/MATCH/COUNTIF.
    """
    assert_valid_workbook_spec(spec)
    worksheet = _ensure_new_sheet(workbook, SUPPORT_LISTS_SHEET, parameter_code="__support_lists__")
    apply_sheet_defaults(worksheet, theme=theme, freeze_panes="A5", tab_color=theme.palette.gray_600)
    worksheet.merge_cells("A1:H1")
    worksheet["A1"] = SUPPORT_LISTS_SHEET
    apply_cell_role(worksheet["A1"], CellRole.TITLE, theme=theme)
    set_title_row_height(worksheet, 1, theme)
    worksheet.merge_cells("A2:H2")
    worksheet["A2"] = TECHNICAL_LAYER_NOTICE
    apply_cell_role(worksheet["A2"], CellRole.TECHNICAL, theme=theme)
    set_subtitle_row_height(worksheet, 2, theme)

    dimensions = _dimension_by_code(spec)
    parameters = _parameter_by_code(spec)
    refs: dict[str, SupportListRef] = {}
    current_col = 1
    qsheet = quote_sheetname(SUPPORT_LISTS_SHEET)

    for parameter in _sorted_parameters(spec.parameters):
        if not parameter.values_source:
            continue
        list_name = support_list_defined_name(parameter.code)
        if len(parameter.depends_on) > 1:
            raise ParameterWriteError(
                code="parameter.multiple_dependencies_not_supported_v1",
                parameter_code=parameter.code,
                message="Contract V1 suporta no maximo uma dependencia de dropdown por parametro.",
            )

        if parameter.depends_on:
            dependency_code = parameter.depends_on[0]
            dependency = parameters[dependency_code]
            dimension = dimensions.get(parameter.values_source)
            if dimension is None or not dimension.parent_dimension:
                raise ParameterWriteError(
                    code="parameter.dependency_requires_child_dimension",
                    parameter_code=parameter.code,
                    message="Dropdown dependente precisa usar DimensionSpec com parent_dimension.",
                )
            if dependency.values_source != dimension.parent_dimension:
                raise ParameterWriteError(
                    code="parameter.dependency_dimension_mismatch",
                    parameter_code=parameter.code,
                    message=(
                        f"Parametro pai {dependency_code!r} usa {dependency.values_source!r}, "
                        f"mas a dimensao filha declara parent_dimension={dimension.parent_dimension!r}."
                    ),
                )

            parent_labels, parent_order = _parent_label_maps(spec, dimension.parent_dimension)
            members = list(_dimension_members(spec, dimension))
            unknown_parents = sorted(
                {str(member.parent_key) for member in members if member.parent_key not in parent_labels},
                key=str.casefold,
            )
            if unknown_parents:
                raise ParameterWriteError(
                    code="parameter.unknown_parent_member",
                    parameter_code=parameter.code,
                    message=f"Valores pai inexistentes na dimensao {dimension.parent_dimension!r}: {', '.join(unknown_parents)}.",
                )
            members.sort(
                key=lambda member: (
                    parent_order.get(member.parent_key, 10**9),
                    _sort_key(member.sort_value, member.label),
                )
            )

            _write_support_header(
                worksheet,
                start_column=current_col,
                headers=(dependency.label, parameter.label),
                theme=theme,
            )
            rows: list[tuple[Scalar, Scalar]] = []
            if dependency.empty_option is not None:
                global_children = sorted({member.label for member in members}, key=lambda value: str(value).casefold())
                all_children = _with_empty_option(global_children, parameter.empty_option)
                rows.extend((dependency.empty_option, child) for child in all_children)
            last_parent: Any = object()
            for member in members:
                parent_label = parent_labels[member.parent_key]
                if parent_label != last_parent and parameter.empty_option is not None:
                    rows.append((parent_label, parameter.empty_option))
                rows.append((parent_label, member.label))
                last_parent = parent_label

            if not rows:
                raise ParameterWriteError(
                    code="parameter.empty_values_source",
                    parameter_code=parameter.code,
                    message="Fonte de valores do dropdown dependente nao possui opcoes.",
                )

            first_row = 5
            last_row = first_row + len(rows) - 1
            for offset, (parent_value, child_value) in enumerate(rows):
                row = first_row + offset
                _write_support_value(worksheet.cell(row, current_col), parent_value, theme=theme)
                _write_support_value(worksheet.cell(row, current_col + 1), child_value, theme=theme)

            parent_col = get_column_letter(current_col)
            value_col = get_column_letter(current_col + 1)
            parent_range = f"{qsheet}!${parent_col}${first_row}:${parent_col}${last_row}"
            value_range = f"{qsheet}!${value_col}${first_row}:${value_col}${last_row}"
            parent_name = parameter_defined_name(dependency_code)
            dynamic_formula = (
                f"INDEX({value_range},MATCH({parent_name},{parent_range},0)):"
                f"INDEX({value_range},MATCH({parent_name},{parent_range},0)+COUNTIF({parent_range},{parent_name})-1)"
            )
            try:
                add_defined_name(workbook, list_name, dynamic_formula)
            except ValueError as exc:
                raise ParameterWriteError("parameter.defined_name_collision", parameter.code, str(exc)) from exc

            refs[parameter.code] = SupportListRef(
                parameter_code=parameter.code,
                defined_name=list_name,
                sheet_name=SUPPORT_LISTS_SHEET,
                value_range=value_range,
                row_count=len(rows),
                parent_parameter_code=dependency_code,
                parent_range=parent_range,
                dynamic=True,
            )
            current_col += 3
            continue

        values = _with_empty_option(_independent_values(spec, parameter), parameter.empty_option)
        if not values:
            raise ParameterWriteError(
                code="parameter.empty_values_source",
                parameter_code=parameter.code,
                message="Fonte de valores do dropdown nao possui opcoes.",
            )
        _write_support_header(worksheet, start_column=current_col, headers=(parameter.label,), theme=theme)
        first_row = 5
        last_row = first_row + len(values) - 1
        for offset, value in enumerate(values):
            _write_support_value(worksheet.cell(first_row + offset, current_col), value, theme=theme)
        value_col = get_column_letter(current_col)
        value_range = f"{qsheet}!${value_col}${first_row}:${value_col}${last_row}"
        try:
            add_defined_name(workbook, list_name, value_range)
        except ValueError as exc:
            raise ParameterWriteError("parameter.defined_name_collision", parameter.code, str(exc)) from exc
        refs[parameter.code] = SupportListRef(
            parameter_code=parameter.code,
            defined_name=list_name,
            sheet_name=SUPPORT_LISTS_SHEET,
            value_range=value_range,
            row_count=len(values),
        )
        current_col += 2

    protect_sheet(worksheet)
    return refs


def _allowed_values_for_initial(
    spec: WorkbookSpec,
    parameter: ParameterSpec,
    support_ref: SupportListRef | None,
) -> tuple[Scalar, ...] | None:
    if not parameter.values_source:
        return None
    if not parameter.depends_on:
        return _with_empty_option(_independent_values(spec, parameter), parameter.empty_option)

    dimensions = _dimension_by_code(spec)
    parameters = _parameter_by_code(spec)
    dimension = dimensions[parameter.values_source]
    dependency_code = parameter.depends_on[0]
    dependency = parameters[dependency_code]
    parent_value = _resolve_initial_value(spec, dependency)
    if parent_value is None:
        return ()

    parent_labels, _ = _parent_label_maps(spec, dimension.parent_dimension or "")
    members = _dimension_members(spec, dimension)
    if dependency.empty_option is not None and parent_value == dependency.empty_option:
        values = [member.label for member in members]
    else:
        values = [member.label for member in members if parent_labels.get(member.parent_key) == parent_value]
    return _with_empty_option(values, parameter.empty_option)


def _validate_initial_value(
    spec: WorkbookSpec,
    parameter: ParameterSpec,
    support_ref: SupportListRef | None,
    value: Scalar,
) -> None:
    if value is None:
        if parameter.required:
            raise ParameterWriteError(
                code="parameter.required_without_value",
                parameter_code=parameter.code,
                message="Parametro obrigatorio nao possui valor inicial.",
            )
        return
    allowed = _allowed_values_for_initial(spec, parameter, support_ref)
    if allowed is not None and value not in allowed:
        raise ParameterWriteError(
            code="parameter.initial_value_out_of_scope",
            parameter_code=parameter.code,
            message=f"Valor inicial {value!r} nao pertence ao escopo autorizado/lista disponivel.",
        )


def _add_list_validation(
    worksheet: Worksheet,
    cell_reference: str,
    parameter: ParameterSpec,
    support_ref: SupportListRef,
) -> None:
    validation = DataValidation(
        type="list",
        formula1=f"={support_ref.defined_name}",
        allow_blank=not parameter.required,
    )
    validation.error = "Selecione um valor da lista autorizada."
    validation.errorTitle = "Valor invalido"
    validation.promptTitle = parameter.label
    if parameter.depends_on:
        validation.prompt = "A lista depende do controle anterior. Ao altera-lo, confirme novamente este valor."
    else:
        validation.prompt = parameter.description or "Escolha um valor da lista."
    validation.showInputMessage = True
    validation.showErrorMessage = True
    validation.errorStyle = "stop"
    worksheet.add_data_validation(validation)
    validation.add(worksheet[cell_reference])


def write_parameters_sheet(
    workbook: Workbook,
    spec: WorkbookSpec,
    support_lists: Mapping[str, SupportListRef],
    *,
    theme: ExcelOfficialTheme = DEFAULT_THEME,
) -> Mapping[str, ParameterRef]:
    assert_valid_workbook_spec(spec)
    worksheet = _ensure_new_sheet(workbook, PARAMETERS_SHEET, parameter_code="__parameters__")
    apply_sheet_defaults(worksheet, theme=theme, freeze_panes="A6", tab_color=theme.palette.green)

    worksheet.merge_cells("A1:F1")
    worksheet["A1"] = "PARÂMETROS DO PAINEL"
    apply_cell_role(worksheet["A1"], CellRole.TITLE, theme=theme)
    set_title_row_height(worksheet, 1, theme)

    worksheet.merge_cells("A2:F2")
    worksheet["A2"] = (
        "Os filtros do site apenas inicializam estes controles. "
        "Depois do download, o próprio Excel controla os recortes offline autorizados."
    )
    apply_cell_role(worksheet["A2"], CellRole.SUBTITLE, theme=theme)
    set_subtitle_row_height(worksheet, 2, theme)

    worksheet.merge_cells("A4:F4")
    worksheet["A4"] = "CONTROLES INTERATIVOS"
    apply_cell_role(worksheet["A4"], CellRole.SECTION, theme=theme)
    set_section_row_height(worksheet, 4, theme)

    worksheet.column_dimensions["A"].width = 30
    worksheet.column_dimensions["B"].width = 28
    worksheet.column_dimensions["C"].width = 3
    worksheet.column_dimensions["D"].width = 30
    worksheet.column_dimensions["E"].width = 24
    worksheet.column_dimensions["F"].width = 18

    refs: dict[str, ParameterRef] = {}
    for offset, parameter in enumerate(_sorted_parameters(spec.parameters)):
        row = 6 + offset
        label_cell = worksheet.cell(row, 1, parameter.label)
        apply_cell_role(label_cell, CellRole.STATIC, theme=theme)
        label_font = copy(label_cell.font)
        label_font.bold = True
        label_cell.font = label_font

        value = _resolve_initial_value(spec, parameter)
        support_ref = support_lists.get(parameter.code)
        _validate_initial_value(spec, parameter, support_ref, value)

        value_cell = worksheet.cell(row, 2, value)
        if isinstance(value, str):
            value_cell.data_type = "s"
        role = CellRole.INPUT if parameter.editable else CellRole.CONTROL
        apply_cell_role(
            value_cell,
            role,
            number_format=_parameter_number_format(parameter, value, theme),
            theme=theme,
        )

        worksheet.merge_cells(start_row=row, start_column=4, end_row=row, end_column=6)
        note_cell = worksheet.cell(row, 4, parameter.description)
        apply_cell_role(note_cell, CellRole.NOTE, theme=theme)
        if parameter.depends_on:
            dependency_labels = ", ".join(_parameter_by_code(spec)[code].label for code in parameter.depends_on)
            suffix = f" Depende de: {dependency_labels}."
            note_cell.value = f"{parameter.description}{suffix}".strip()

        cell_reference = f"B{row}"
        defined_name = parameter_defined_name(parameter.code)
        qsheet = quote_sheetname(PARAMETERS_SHEET)
        try:
            add_defined_name(workbook, defined_name, f"{qsheet}!$B${row}")
        except ValueError as exc:
            raise ParameterWriteError("parameter.defined_name_collision", parameter.code, str(exc)) from exc

        if parameter.editable and support_ref is not None:
            _add_list_validation(worksheet, cell_reference, parameter, support_ref)

        refs[parameter.code] = ParameterRef(
            parameter_code=parameter.code,
            sheet_name=PARAMETERS_SHEET,
            cell_reference=cell_reference,
            defined_name=defined_name,
            support_list_name=support_ref.defined_name if support_ref else None,
            editable=parameter.editable,
            dependency_codes=parameter.depends_on,
        )

    note_row = 7 + len(spec.parameters)
    worksheet.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=6)
    worksheet.cell(note_row, 1, "Células amarelas são editáveis. Alterações locais não são sincronizadas com o Data UNIVC.")
    apply_cell_role(worksheet.cell(note_row, 1), CellRole.NOTE, theme=theme)

    protect_sheet(worksheet)
    worksheet.protection.selectUnlockedCells = True
    worksheet.protection.selectLockedCells = True
    return refs


def write_parameter_system(
    workbook: Workbook,
    spec: WorkbookSpec,
    *,
    theme: ExcelOfficialTheme = DEFAULT_THEME,
) -> ParameterSystemRef:
    """Create LISTAS DE APOIO + PARAMETROS and all V1 defined names/validations."""
    _preflight_parameter_system(workbook, spec)
    support_lists = write_support_lists(workbook, spec, theme=theme)
    parameters = write_parameters_sheet(workbook, spec, support_lists, theme=theme)
    return ParameterSystemRef(
        sheet_name=PARAMETERS_SHEET,
        support_sheet_name=SUPPORT_LISTS_SHEET,
        parameters=parameters,
        support_lists=support_lists,
    )


__all__ = [
    "PARAMETERS_SHEET",
    "SUPPORT_LISTS_SHEET",
    "ParameterWriteError",
    "write_parameter_system",
    "write_parameters_sheet",
    "write_support_lists",
]
