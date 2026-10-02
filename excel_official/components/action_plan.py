from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from typing import Mapping
import unicodedata

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, PatternFill, Protection
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from ..constants import REQUIRED_INSTITUTIONAL_SHEETS
from ..contract import ActionPlanRef, ActionPlanRowSource, ActionPlanSpec, WorkbookSpec
from ..protection import protect_sheet
from ..styles import (
    apply_cell_role,
    apply_sheet_defaults,
    set_section_row_height,
    set_subtitle_row_height,
    set_table_header_row_height,
    set_title_row_height,
)
from ..theme import CellRole, DEFAULT_THEME, ExcelOfficialTheme
from ..validation import assert_valid_workbook_spec
from .names import add_defined_name
from .parameters import SUPPORT_LISTS_SHEET

ACTION_PLAN_SHEET = REQUIRED_INSTITUTIONAL_SHEETS[5]
ACTION_STATUS_DEFINED_NAME = "LST_ACTION_STATUS"
ACTION_INDICATOR_DEFINED_NAME = "LST_ACTION_INDICATOR"
ACTION_TABLE_NAME = "TblPlanoAcaoOfficial"

_FIELD_LABELS = {
    "indicator": "Indicador",
    "problem": "Problema",
    "diagnosis": "Diagn\u00f3stico",
    "action": "A\u00e7\u00e3o",
    "owner": "Respons\u00e1vel",
    "deadline": "Prazo",
    "status": "Status",
    "days_remaining": "Dias restantes",
}

_FIELD_WIDTHS = {
    "indicator": 24,
    "problem": 34,
    "diagnosis": 34,
    "action": 42,
    "owner": 24,
    "deadline": 14,
    "status": 18,
    "days_remaining": 16,
}


@dataclass(frozen=True, slots=True)
class ActionPlanWriteError(ValueError):
    code: str
    message: str

    def __str__(self) -> str:
        return f"{self.code}: {self.message}"


def _workbook_table_names(workbook: Workbook) -> set[str]:
    names: set[str] = set()
    for worksheet in workbook.worksheets:
        names.update(str(name).casefold() for name in worksheet.tables.keys())
    return names


def _field_label(plan: ActionPlanSpec, field_code: str) -> str:
    return plan.field_labels.get(field_code, _FIELD_LABELS.get(field_code, field_code.replace("_", " ").title()))


def _defined_name_exists(workbook: Workbook, name: str) -> bool:
    return name.casefold() in {str(item).casefold() for item in workbook.defined_names.keys()}


def _preflight_action_plan(workbook: Workbook, spec: WorkbookSpec) -> None:
    assert_valid_workbook_spec(spec)
    plan = spec.action_plan
    if not plan.enabled:
        return

    if any(ws.title.casefold() == ACTION_PLAN_SHEET.casefold() for ws in workbook.worksheets):
        raise ActionPlanWriteError("action_plan.sheet_collision", f"Aba {ACTION_PLAN_SHEET!r} ja existe no workbook.")
    if ACTION_TABLE_NAME.casefold() in _workbook_table_names(workbook):
        raise ActionPlanWriteError("action_plan.table_collision", f"Excel Table {ACTION_TABLE_NAME!r} ja existe no workbook.")

    if plan.status_field in plan.local_editable_fields:
        if SUPPORT_LISTS_SHEET not in workbook.sheetnames:
            raise ActionPlanWriteError(
                "action_plan.support_sheet_required",
                "Status editavel exige LISTAS DE APOIO materializada antes do PLANO_DE_ACAO.",
            )
        if _defined_name_exists(workbook, ACTION_STATUS_DEFINED_NAME):
            raise ActionPlanWriteError(
                "action_plan.status_name_collision",
                f"Nome definido {ACTION_STATUS_DEFINED_NAME!r} ja existe.",
            )
    if plan.indicator_field in plan.local_editable_fields and spec.metrics:
        if SUPPORT_LISTS_SHEET not in workbook.sheetnames:
            raise ActionPlanWriteError(
                "action_plan.support_sheet_required",
                "Indicador editavel exige LISTAS DE APOIO materializada antes do PLANO_DE_ACAO.",
            )
        if _defined_name_exists(workbook, ACTION_INDICATOR_DEFINED_NAME):
            raise ActionPlanWriteError(
                "action_plan.indicator_name_collision",
                f"Nome definido {ACTION_INDICATOR_DEFINED_NAME!r} ja existe.",
            )


def _append_support_list(
    workbook: Workbook,
    *,
    header_label: str,
    values: tuple[str, ...],
    defined_name: str,
    theme: ExcelOfficialTheme,
) -> str:
    ws = workbook[SUPPORT_LISTS_SHEET]
    start_col = max(1, ws.max_column + 2)
    header = ws.cell(4, start_col, header_label)
    apply_cell_role(header, CellRole.TABLE_HEADER, theme=theme)
    ws.column_dimensions[get_column_letter(start_col)].width = 28

    start_row = 5
    for offset, value in enumerate(values):
        cell = ws.cell(start_row + offset, start_col, value)
        apply_cell_role(cell, CellRole.TECHNICAL, theme=theme)
        cell.protection = Protection(locked=True)

    end_row = start_row + len(values) - 1
    qsheet = quote_sheetname(SUPPORT_LISTS_SHEET)
    letter = get_column_letter(start_col)
    attr_text = f"{qsheet}!${letter}${start_row}:${letter}${end_row}"
    add_defined_name(workbook, defined_name, attr_text)
    protect_sheet(ws)
    return defined_name


def _write_status_support_list(
    workbook: Workbook,
    plan: ActionPlanSpec,
    *,
    theme: ExcelOfficialTheme,
) -> str | None:
    if plan.status_field not in plan.local_editable_fields:
        return None
    return _append_support_list(
        workbook,
        header_label="Status do plano de acao",
        values=plan.status_options,
        defined_name=ACTION_STATUS_DEFINED_NAME,
        theme=theme,
    )


def _write_indicator_support_list(
    workbook: Workbook,
    spec: WorkbookSpec,
    *,
    theme: ExcelOfficialTheme,
) -> str | None:
    plan = spec.action_plan
    if plan.indicator_field not in plan.local_editable_fields or not spec.metrics:
        return None
    labels = tuple(metric.label for metric in spec.metrics)
    return _append_support_list(
        workbook,
        header_label="Indicadores do plano de acao",
        values=labels,
        defined_name=ACTION_INDICATOR_DEFINED_NAME,
        theme=theme,
    )


def _days_formula(plan: ActionPlanSpec, row: int, field_columns: Mapping[str, int]) -> str | None:
    if not plan.days_remaining_field or plan.days_remaining_field not in field_columns:
        return None
    if not plan.deadline_field or plan.deadline_field not in field_columns:
        return None
    if not plan.status_field or plan.status_field not in field_columns:
        return None

    deadline = f"{get_column_letter(field_columns[plan.deadline_field])}{row}"
    status = f"{get_column_letter(field_columns[plan.status_field])}{row}"
    completed = [f'{status}="{value.replace(chr(34), chr(34) * 2)}"' for value in plan.completed_statuses]
    completed_expr = ",".join(completed)
    if completed_expr:
        return f'=IF(OR({deadline}="",OR({completed_expr})),"",{deadline}-TODAY())'
    return f'=IF({deadline}="","",{deadline}-TODAY())'


def _ascii_key(value: str) -> str:
    return unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii").casefold()


def _apply_status_conditional_formatting(
    ws,
    cell_range: str,
    plan: ActionPlanSpec,
    theme: ExcelOfficialTheme,
) -> None:
    if not cell_range:
        return
    start_cell = cell_range.split(":", 1)[0]
    col = "".join(ch for ch in start_cell if ch.isalpha())
    row = "".join(ch for ch in start_cell if ch.isdigit())
    completed = set(plan.completed_statuses)
    for status in plan.status_options:
        key = _ascii_key(status)
        if status in completed:
            fill = PatternFill("solid", fgColor=theme.palette.green_pale)
        elif "andamento" in key:
            fill = PatternFill("solid", fgColor=theme.palette.warning_fill)
        elif "suspens" in key:
            fill = PatternFill("solid", fgColor=theme.palette.gray_100)
        else:
            fill = PatternFill("solid", fgColor=theme.palette.info_fill)
        safe_status = status.replace('"', '""')
        ws.conditional_formatting.add(
            cell_range,
            FormulaRule(formula=[f'${col}{row}="{safe_status}"'], fill=fill),
        )


def _write_source_cell(cell, source: ActionPlanRowSource, *, theme: ExcelOfficialTheme) -> None:
    if source is ActionPlanRowSource.OFFICIAL:
        cell.value = "OFICIAL"
        apply_cell_role(cell, CellRole.IMPORTED, theme=theme)
    else:
        cell.value = "LOCAL"
        apply_cell_role(cell, CellRole.CONTROL, theme=theme)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.protection = Protection(locked=True)


def _write_value_cell(
    cell,
    *,
    field_code: str,
    value,
    source: ActionPlanRowSource,
    plan: ActionPlanSpec,
    theme: ExcelOfficialTheme,
) -> None:
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"

    is_local_editable = source is ActionPlanRowSource.LOCAL and field_code in plan.local_editable_fields
    if field_code == plan.days_remaining_field:
        apply_cell_role(cell, CellRole.DERIVED, number_format=theme.number_formats.integer, theme=theme)
        cell.protection = Protection(locked=True)
        return
    if is_local_editable:
        number_format = theme.number_formats.date if field_code == plan.deadline_field else None
        apply_cell_role(cell, CellRole.INPUT, number_format=number_format, theme=theme)
        cell.protection = Protection(locked=False)
        return

    role = CellRole.IMPORTED if source is ActionPlanRowSource.OFFICIAL else CellRole.STATIC
    number_format = theme.number_formats.date if field_code == plan.deadline_field else None
    apply_cell_role(cell, role, number_format=number_format, theme=theme)
    cell.protection = Protection(locked=True)


def write_action_plan_sheet(
    workbook: Workbook,
    spec: WorkbookSpec,
    *,
    theme: ExcelOfficialTheme = DEFAULT_THEME,
) -> ActionPlanRef | None:
    """Materialize the institutional action-plan surface.

    Official rows are immutable snapshot data. LOCAL rows expose only the fields
    declared in ``local_editable_fields``. Those cells are intentionally local
    to the XLSX and never imply synchronization back to Data UNIVC.
    """

    _preflight_action_plan(workbook, spec)
    plan = spec.action_plan
    if not plan.enabled:
        return None

    status_list_name = _write_status_support_list(workbook, plan, theme=theme)
    indicator_list_name = _write_indicator_support_list(workbook, spec, theme=theme)

    official_rows = [row for row in plan.rows if row.source is ActionPlanRowSource.OFFICIAL]
    local_rows = [row for row in plan.rows if row.source is ActionPlanRowSource.LOCAL]
    materialized_rows: list[tuple[ActionPlanRowSource, Mapping[str, object]]] = [
        (ActionPlanRowSource.OFFICIAL, row.values) for row in official_rows
    ]
    materialized_rows.extend((ActionPlanRowSource.LOCAL, row.values) for row in local_rows)
    materialized_rows.extend((ActionPlanRowSource.LOCAL, {}) for _ in range(plan.local_blank_rows))

    ws = workbook.create_sheet(ACTION_PLAN_SHEET)
    apply_sheet_defaults(ws, theme=theme, freeze_panes="B7", tab_color=theme.palette.green_mid)
    ws.sheet_state = "visible"

    total_columns = 1 + len(plan.fields)
    end_letter = get_column_letter(total_columns)
    ws.merge_cells(f"A1:{end_letter}1")
    ws["A1"] = "PLANO DE A\u00c7\u00c3O"
    apply_cell_role(ws["A1"], CellRole.TITLE, theme=theme)
    set_title_row_height(ws, 1, theme)

    ws.merge_cells(f"A2:{end_letter}2")
    ws["A2"] = plan.local_notice
    apply_cell_role(ws["A2"], CellRole.NOTE, theme=theme)
    set_subtitle_row_height(ws, 2, theme)

    ws.merge_cells(f"A4:{end_letter}4")
    ws["A4"] = "A\u00c7\u00d5ES OFICIAIS (somente leitura) + ACOMPANHAMENTO LOCAL (campos amarelos edit\u00e1veis)"
    apply_cell_role(ws["A4"], CellRole.SECTION, theme=theme)
    set_section_row_height(ws, 4, theme)

    header_row = 6
    ws.cell(header_row, 1, "Origem")
    apply_cell_role(ws.cell(header_row, 1), CellRole.TABLE_HEADER, theme=theme)
    field_columns: dict[str, int] = {}
    for index, field_code in enumerate(plan.fields, start=2):
        field_columns[field_code] = index
        cell = ws.cell(header_row, index, _field_label(plan, field_code))
        apply_cell_role(cell, CellRole.TABLE_HEADER, theme=theme)
    set_table_header_row_height(ws, header_row, theme)

    ws.column_dimensions["A"].width = 12
    for field_code, col_index in field_columns.items():
        ws.column_dimensions[get_column_letter(col_index)].width = _FIELD_WIDTHS.get(field_code, 22)

    first_data_row = header_row + 1
    if not materialized_rows:
        # Keep one locked structural row so the Excel Table remains valid.
        materialized_rows.append((ActionPlanRowSource.OFFICIAL, {}))
    last_data_row = first_data_row + len(materialized_rows) - 1

    local_row_numbers: list[int] = []
    official_row_count = 0
    actual_local_row_count = 0
    for offset, (source, values) in enumerate(materialized_rows):
        excel_row = first_data_row + offset
        _write_source_cell(ws.cell(excel_row, 1), source, theme=theme)
        if source is ActionPlanRowSource.OFFICIAL:
            official_row_count += 1
        else:
            actual_local_row_count += 1
            local_row_numbers.append(excel_row)

        for field_code, col_index in field_columns.items():
            cell = ws.cell(excel_row, col_index)
            if field_code == plan.days_remaining_field:
                formula = _days_formula(plan, excel_row, field_columns)
                _write_value_cell(
                    cell,
                    field_code=field_code,
                    value=formula or "",
                    source=source,
                    plan=plan,
                    theme=theme,
                )
                if formula:
                    cell.data_type = "f"
            else:
                value = values.get(field_code)
                _write_value_cell(
                    cell,
                    field_code=field_code,
                    value=value,
                    source=source,
                    plan=plan,
                    theme=theme,
                )

    table_ref = f"A{header_row}:{end_letter}{last_data_row}"
    table = Table(displayName=ACTION_TABLE_NAME, ref=table_ref)
    table.tableStyleInfo = TableStyleInfo(
        name=theme.tables.style_name,
        showFirstColumn=theme.tables.show_first_column,
        showLastColumn=theme.tables.show_last_column,
        showRowStripes=theme.tables.show_row_stripes,
        showColumnStripes=theme.tables.show_column_stripes,
    )
    ws.add_table(table)

    editable_ranges: dict[str, str] = {}
    if local_row_numbers:
        first_local = min(local_row_numbers)
        last_local = max(local_row_numbers)
        for field_code in plan.local_editable_fields:
            col = field_columns.get(field_code)
            if col is None:
                continue
            editable_ranges[field_code] = f"{get_column_letter(col)}{first_local}:{get_column_letter(col)}{last_local}"

    indicator_validation_range: str | None = None
    if indicator_list_name and plan.indicator_field in field_columns and plan.indicator_field in plan.local_editable_fields and local_row_numbers:
        col = field_columns[plan.indicator_field]
        first_local = min(local_row_numbers)
        last_local = max(local_row_numbers)
        indicator_validation_range = f"{get_column_letter(col)}{first_local}:{get_column_letter(col)}{last_local}"
        validation = DataValidation(type="list", formula1=f"={indicator_list_name}", allow_blank=True)
        validation.error = "Selecione um indicador valido da lista institucional."
        validation.errorTitle = "Indicador invalido"
        validation.showErrorMessage = True
        ws.add_data_validation(validation)
        validation.add(indicator_validation_range)

    status_validation_range: str | None = None
    if status_list_name and plan.status_field in field_columns and plan.status_field in plan.local_editable_fields and local_row_numbers:
        col = field_columns[plan.status_field]
        first_local = min(local_row_numbers)
        last_local = max(local_row_numbers)
        status_validation_range = f"{get_column_letter(col)}{first_local}:{get_column_letter(col)}{last_local}"
        validation = DataValidation(type="list", formula1=f"={status_list_name}", allow_blank=True)
        validation.error = "Selecione um status valido da lista institucional."
        validation.errorTitle = "Status invalido"
        validation.prompt = "Este status e local ao arquivo e nao sincroniza automaticamente com o Data UNIVC."
        validation.promptTitle = "Acompanhamento local"
        validation.showErrorMessage = True
        validation.showInputMessage = True
        ws.add_data_validation(validation)
        validation.add(status_validation_range)
        _apply_status_conditional_formatting(ws, status_validation_range, plan, theme)

    if plan.days_remaining_field in field_columns:
        col = get_column_letter(field_columns[plan.days_remaining_field])
        days_range = f"{col}{first_data_row}:{col}{last_data_row}"
        ws.conditional_formatting.add(
            days_range,
            CellIsRule(operator="lessThan", formula=["0"], fill=PatternFill("solid", fgColor=theme.palette.error_fill)),
        )

    # Lock everything first, then restore explicitly local input cells. This
    # prevents accidental editability introduced by future style changes.
    for row_cells in ws.iter_rows(min_row=1, max_row=last_data_row, min_col=1, max_col=total_columns):
        for cell in row_cells:
            cell.protection = Protection(locked=True)
    for cell_range in editable_ranges.values():
        for row_cells in ws[cell_range]:
            for cell in row_cells:
                cell.protection = Protection(locked=False)
                apply_cell_role(cell, CellRole.INPUT, number_format=cell.number_format, theme=theme)

    protect_sheet(ws)
    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True

    return ActionPlanRef(
        sheet_name=ACTION_PLAN_SHEET,
        header_row=header_row,
        first_data_row=first_data_row,
        last_data_row=last_data_row,
        table_range=table_ref,
        official_row_count=official_row_count,
        local_row_count=actual_local_row_count,
        local_editable_ranges=editable_ranges,
        status_validation_range=status_validation_range,
        indicator_validation_range=indicator_validation_range,
        notice_cell="A2",
    )


__all__ = [
    "ACTION_PLAN_SHEET",
    "ACTION_STATUS_DEFINED_NAME",
    "ACTION_INDICATOR_DEFINED_NAME",
    "ACTION_TABLE_NAME",
    "ActionPlanWriteError",
    "write_action_plan_sheet",
]
