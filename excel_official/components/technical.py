from __future__ import annotations

from dataclasses import dataclass

from openpyxl import Workbook
from openpyxl.styles import Protection
from openpyxl.worksheet.table import Table, TableStyleInfo

from ..constants import TECHNICAL_LAYER_NOTICE
from ..contract import MetricSpec, MetricUnit, WorkbookSpec
from ..protection import protect_sheet
from ..styles import apply_cell_role, apply_sheet_defaults, set_table_header_row_height, set_title_row_height
from ..theme import CellRole, DEFAULT_THEME, ExcelOfficialTheme

TARGETS_SHEET = "METAS"
INDICATORS_SHEET = "INDICADORES"
CALC_SHEET = "CALC"
SUPPORT_SHEET = "LISTAS DE APOIO"
PERIOD_DIMENSION_SHEET = "DIM_PERIODO"
MONTH_DIMENSION_SHEET = "DIM_MES"


@dataclass(frozen=True, slots=True)
class TechnicalRegistryRef:
    targets_sheet: str | None = None
    indicators_sheet: str | None = None


def _technical_header(workbook: Workbook, sheet_name: str, title: str, theme: ExcelOfficialTheme):
    ws = workbook.create_sheet(sheet_name)
    apply_sheet_defaults(ws, theme=theme, freeze_panes="A5", tab_color=theme.palette.gray_600)
    ws.sheet_state = "visible"
    ws.merge_cells("A1:G1")
    ws["A1"] = title
    apply_cell_role(ws["A1"], CellRole.TITLE, theme=theme)
    set_title_row_height(ws, 1, theme)
    ws.merge_cells("A2:G2")
    ws["A2"] = TECHNICAL_LAYER_NOTICE
    apply_cell_role(ws["A2"], CellRole.TECHNICAL, theme=theme)
    return ws


def _format_for_metric(metric: MetricSpec, theme: ExcelOfficialTheme) -> str:
    if metric.number_format:
        return metric.number_format
    precision = max(0, min(metric.display_precision, 6))
    decimal = "0" if precision == 0 else "0." + ("0" * precision)
    if metric.unit is MetricUnit.PERCENT:
        return "0%" if precision == 0 else "0." + ("0" * precision) + "%"
    if metric.unit is MetricUnit.CURRENCY_BRL:
        return theme.number_formats.currency_brl
    if metric.unit is MetricUnit.COUNT:
        return theme.number_formats.integer
    return decimal


def write_indicator_registry(workbook: Workbook, spec: WorkbookSpec, *, theme: ExcelOfficialTheme = DEFAULT_THEME) -> str:
    if INDICATORS_SHEET in workbook.sheetnames:
        return INDICATORS_SHEET
    ws = _technical_header(workbook, INDICATORS_SHEET, "REGISTRO DE INDICADORES", theme)
    headers = ("Codigo", "Indicador", "Agregacao", "Unidade", "Dimensoes permitidas", "Recorte offline", "Faixa valida")
    for col, label in enumerate(headers, 1):
        ws.cell(4, col, label)
        apply_cell_role(ws.cell(4, col), CellRole.TABLE_HEADER, theme=theme)
    set_table_header_row_height(ws, 4, theme)
    row = 5
    for metric in spec.metrics:
        valid_range = ""
        if metric.valid_min is not None or metric.valid_max is not None:
            valid_range = f"{metric.valid_min if metric.valid_min is not None else '-inf'} .. {metric.valid_max if metric.valid_max is not None else '+inf'}"
        values = (
            metric.code,
            metric.label,
            metric.aggregation.value,
            metric.unit.value,
            ", ".join(metric.allowed_dimensions),
            "SIM" if metric.offline_recut else "NAO",
            valid_range,
        )
        for col, value in enumerate(values, 1):
            ws.cell(row, col, value)
            apply_cell_role(ws.cell(row, col), CellRole.TECHNICAL, theme=theme)
        row += 1
    if spec.metrics:
        table = Table(displayName="TblExcelOfficialIndicators", ref=f"A4:G{row - 1}")
        table.tableStyleInfo = TableStyleInfo(name=theme.tables.style_name, showRowStripes=True, showFirstColumn=False, showLastColumn=False, showColumnStripes=False)
        ws.add_table(table)
    else:
        ws["A5"] = "Nenhum MetricSpec declarado."
        apply_cell_role(ws["A5"], CellRole.NOTE, theme=theme)
    widths = (32, 34, 22, 18, 38, 18, 22)
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + idx)].width = width
    for cells in ws.iter_rows():
        for cell in cells:
            cell.protection = Protection(locked=True)
    protect_sheet(ws)
    return INDICATORS_SHEET


def write_targets_registry(workbook: Workbook, spec: WorkbookSpec, *, theme: ExcelOfficialTheme = DEFAULT_THEME) -> str:
    if TARGETS_SHEET in workbook.sheetnames:
        return TARGETS_SHEET
    ws = _technical_header(workbook, TARGETS_SHEET, "REGISTRO DE METAS", theme)
    headers = ("Codigo", "Indicador", "Meta", "Unidade", "Origem")
    for col, label in enumerate(headers, 1):
        ws.cell(4, col, label)
        apply_cell_role(ws.cell(4, col), CellRole.TABLE_HEADER, theme=theme)
    set_table_header_row_height(ws, 4, theme)
    metrics = {item.code: item for item in spec.metrics}
    targets: list[tuple[str, object, str]] = []
    for kpi in spec.dashboard.kpis:
        if kpi.target is not None:
            targets.append((kpi.metric_code, kpi.target, "PAINEL · estática"))
    if spec.matrix is not None and spec.matrix.target is not None:
        targets.append((spec.matrix.metric_code, spec.matrix.target, "MATRIZ · estática"))
    static_codes = {code for code, _, _ in targets}
    for binding in spec.target_bindings:
        if binding.metric_code not in static_codes:
            targets.append((binding.metric_code, "dinâmica", f"{binding.dataset_code} · PARAMETROS"))
    row = 5
    for code, target, source in targets:
        metric = metrics.get(code)
        values = (code, metric.label if metric else code, target, metric.unit.value if metric else "", source)
        for col, value in enumerate(values, 1):
            ws.cell(row, col, value)
            apply_cell_role(ws.cell(row, col), CellRole.TECHNICAL, theme=theme)
        if metric and isinstance(target, (int, float)):
            ws.cell(row, 3).number_format = _format_for_metric(metric, theme)
        row += 1
    if targets:
        table = Table(displayName="TblExcelOfficialTargets", ref=f"A4:E{row - 1}")
        table.tableStyleInfo = TableStyleInfo(name=theme.tables.style_name, showRowStripes=True, showFirstColumn=False, showLastColumn=False, showColumnStripes=False)
        ws.add_table(table)
    else:
        ws["A5"] = "Nenhuma meta declarada neste snapshot."
        apply_cell_role(ws["A5"], CellRole.NOTE, theme=theme)
    widths = (32, 34, 18, 18, 18)
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + idx)].width = width
    for cells in ws.iter_rows():
        for cell in cells:
            cell.protection = Protection(locked=True)
    protect_sheet(ws)
    return TARGETS_SHEET


def ensure_technical_placeholder(workbook: Workbook, sheet_name: str, *, theme: ExcelOfficialTheme = DEFAULT_THEME) -> None:
    if sheet_name in workbook.sheetnames:
        ws = workbook[sheet_name]
        ws.sheet_state = "visible"
        protect_sheet(ws)
        return
    ws = _technical_header(workbook, sheet_name, sheet_name, theme)
    ws["A4"] = "Nenhum conteudo adicional materializado para este snapshot."
    apply_cell_role(ws["A4"], CellRole.NOTE, theme=theme)
    ws.column_dimensions["A"].width = 54
    for cells in ws.iter_rows():
        for cell in cells:
            cell.protection = Protection(locked=True)
    protect_sheet(ws)


def ensure_technical_layers(workbook: Workbook, spec: WorkbookSpec, *, theme: ExcelOfficialTheme = DEFAULT_THEME) -> TechnicalRegistryRef:
    targets = write_targets_registry(workbook, spec, theme=theme) if spec.technical.include_targets else None
    indicators = write_indicator_registry(workbook, spec, theme=theme) if spec.technical.include_indicators else None
    if spec.technical.include_calc:
        ensure_technical_placeholder(workbook, CALC_SHEET, theme=theme)
    if spec.technical.include_support_lists:
        ensure_technical_placeholder(workbook, SUPPORT_SHEET, theme=theme)
    if spec.technical.include_period_dimension:
        ensure_technical_placeholder(workbook, PERIOD_DIMENSION_SHEET, theme=theme)
    if spec.technical.include_month_dimension:
        ensure_technical_placeholder(workbook, MONTH_DIMENSION_SHEET, theme=theme)
    return TechnicalRegistryRef(targets_sheet=targets, indicators_sheet=indicators)


__all__ = [
    "CALC_SHEET",
    "INDICATORS_SHEET",
    "MONTH_DIMENSION_SHEET",
    "PERIOD_DIMENSION_SHEET",
    "SUPPORT_SHEET",
    "TARGETS_SHEET",
    "TechnicalRegistryRef",
    "ensure_technical_layers",
    "ensure_technical_placeholder",
    "write_indicator_registry",
    "write_targets_registry",
]
