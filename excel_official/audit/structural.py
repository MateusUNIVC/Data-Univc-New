from __future__ import annotations

import re
from collections.abc import Mapping

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries

from ..constants import OPTIONAL_TECHNICAL_SHEETS, REQUIRED_INSTITUTIONAL_SHEETS, STANDARD_TECHNICAL_SHEETS
from ..contract import ActionPlanRef, ParameterSystemRef, QualitySeverity, WorkbookSpec
from ..metadata import CUSTOM_PROPERTY_NAMES
from .models import AuditFinding, AuditResult
from .quality import blocking_quality_findings

_FORMULA_ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!")
_UNSUPPORTED_FORMULAS = (
    "FILTER(",
    "XLOOKUP(",
    "SORT(",
    "UNIQUE(",
    "LET(",
    "LAMBDA(",
    "OFFSET(",
    "INDIRECT(",
    "SEQUENCE(",
)
_EXTERNAL_REF = re.compile(r"\[[^\]]+\.(?:xlsx|xlsm|xls|xlsb)\]", re.IGNORECASE)


def _finding(code: str, message: str, *, severity: QualitySeverity = QualitySeverity.BLOCKING, sheet: str | None = None, cell: str | None = None) -> AuditFinding:
    return AuditFinding(code=code, severity=severity, message=message, sheet_name=sheet, cell_reference=cell)


def _technical_requirements(spec: WorkbookSpec) -> tuple[str, ...]:
    output: list[str] = []
    flags = spec.technical
    if flags.include_targets:
        output.append("METAS")
    if flags.include_indicators:
        output.append("INDICADORES")
    if flags.include_calc:
        output.append("CALC")
    if flags.include_support_lists:
        output.append("LISTAS DE APOIO")
    if flags.include_period_dimension:
        output.append("DIM_PERIODO")
    if flags.include_month_dimension:
        output.append("DIM_MES")
    return tuple(output)


def _allowed_unlocked_cells(parameter_system: ParameterSystemRef | None, action_plan_ref: ActionPlanRef | None) -> dict[str, set[str]]:
    allowed: dict[str, set[str]] = {}
    if parameter_system is not None:
        target = allowed.setdefault(parameter_system.sheet_name, set())
        target.update(item.cell_reference for item in parameter_system.parameters.values() if item.editable)
    if action_plan_ref is not None:
        target = allowed.setdefault(action_plan_ref.sheet_name, set())
        for cell_range in action_plan_ref.local_editable_ranges.values():
            min_col, min_row, max_col, max_row = range_boundaries(cell_range)
            ws_cells = {
                f"{get_column_letter(col)}{row}"
                for row in range(min_row, max_row + 1)
                for col in range(min_col, max_col + 1)
            }
            target.update(ws_cells)
    return allowed


def _chart_formulas(workbook: Workbook):
    for ws in workbook.worksheets:
        for chart in getattr(ws, "_charts", ()):
            for series in getattr(chart, "ser", ()):
                candidates = []
                val = getattr(series, "val", None)
                cat = getattr(series, "cat", None)
                if val is not None:
                    num_ref = getattr(val, "numRef", None)
                    if num_ref is not None:
                        candidates.append(getattr(num_ref, "f", None))
                if cat is not None:
                    for attr in ("strRef", "numRef"):
                        ref = getattr(cat, attr, None)
                        if ref is not None:
                            candidates.append(getattr(ref, "f", None))
                for formula in candidates:
                    if formula:
                        yield ws.title, str(formula)


def audit_workbook(
    workbook: Workbook,
    spec: WorkbookSpec,
    *,
    parameter_system: ParameterSystemRef | None = None,
    action_plan_ref: ActionPlanRef | None = None,
) -> AuditResult:
    findings: list[AuditFinding] = list(blocking_quality_findings(spec))
    checks_run = len(spec.quality.checks)

    checks_run += 1
    missing = [name for name in REQUIRED_INSTITUTIONAL_SHEETS if name not in workbook.sheetnames]
    if missing:
        findings.append(_finding("structure.required_sheets", f"Abas institucionais ausentes: {', '.join(missing)}."))

    checks_run += 1
    if tuple(workbook.sheetnames[: len(REQUIRED_INSTITUTIONAL_SHEETS)]) != REQUIRED_INSTITUTIONAL_SHEETS:
        findings.append(_finding("structure.sheet_order", "As seis abas institucionais nao ocupam a ordem oficial do Contract V1."))

    for sheet_name in _technical_requirements(spec):
        checks_run += 1
        if sheet_name not in workbook.sheetnames:
            findings.append(_finding("structure.technical_sheet_missing", f"Aba tecnica requerida ausente: {sheet_name}."))
            continue
        ws = workbook[sheet_name]
        if ws.sheet_state != "visible":
            findings.append(_finding("structure.technical_sheet_hidden", f"Aba tecnica precisa permanecer visivel: {sheet_name}.", sheet=sheet_name))
        if not ws.protection.sheet:
            findings.append(_finding("structure.technical_sheet_unprotected", f"Aba tecnica precisa estar protegida: {sheet_name}.", sheet=sheet_name))

    for sheet_name in REQUIRED_INSTITUTIONAL_SHEETS:
        if sheet_name not in workbook.sheetnames:
            continue
        checks_run += 1
        if not workbook[sheet_name].protection.sheet:
            findings.append(_finding("structure.institutional_sheet_unprotected", f"Aba institucional nao protegida: {sheet_name}.", sheet=sheet_name))

    checks_run += 1
    if getattr(workbook, "_external_links", None):
        findings.append(_finding("security.external_links", "Workbook contem links externos."))

    checks_run += 1
    if getattr(workbook, "vba_archive", None) is not None:
        findings.append(_finding("security.macros", "Workbook contem pacote VBA/macros."))

    allowed_unlocked = _allowed_unlocked_cells(parameter_system, action_plan_ref)
    for ws in workbook.worksheets:
        allowed = allowed_unlocked.get(ws.title, set())
        for row in ws.iter_rows():
            for cell in row:
                if cell.coordinate in allowed:
                    continue
                if cell.protection.locked is False:
                    findings.append(_finding("security.unexpected_unlocked_cell", "Celula desbloqueada fora das superficies de input permitidas.", sheet=ws.title, cell=cell.coordinate))
                    break
            else:
                continue
            break
        checks_run += 1

    table_names: set[str] = set()
    for ws in workbook.worksheets:
        for table in ws.tables.values():
            checks_run += 1
            folded = table.name.casefold()
            if folded in table_names:
                findings.append(_finding("structure.duplicate_table", f"Nome de tabela duplicado: {table.name}.", sheet=ws.title))
            table_names.add(folded)
            if "#REF!" in str(table.ref).upper():
                findings.append(_finding("structure.table_ref", f"Tabela {table.name} possui referencia invalida.", sheet=ws.title))

    for ws in workbook.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if not (isinstance(value, str) and value.startswith("=")):
                    continue
                checks_run += 1
                upper = value.upper()
                if any(token in upper for token in _FORMULA_ERRORS):
                    findings.append(_finding("formula.error_token", "Formula contem token de erro.", sheet=ws.title, cell=cell.coordinate))
                unsupported = next((token for token in _UNSUPPORTED_FORMULAS if token in upper), None)
                if unsupported:
                    findings.append(_finding("formula.unsupported_function", f"Formula usa funcao fora do Contract V1: {unsupported[:-1]}.", sheet=ws.title, cell=cell.coordinate))
                if _EXTERNAL_REF.search(value):
                    findings.append(_finding("formula.external_reference", "Formula referencia outro arquivo Excel.", sheet=ws.title, cell=cell.coordinate))

    for sheet_name, formula in _chart_formulas(workbook):
        checks_run += 1
        if "#REF!" in formula.upper():
            findings.append(_finding("chart.broken_reference", "Grafico contem referencia #REF!.", sheet=sheet_name))
        if _EXTERNAL_REF.search(formula):
            findings.append(_finding("chart.external_reference", "Grafico referencia outro arquivo Excel.", sheet=sheet_name))

    checks_run += 1
    defined_names_text = "\n".join(str(item.attr_text) for item in workbook.defined_names.values())
    if "#REF!" in defined_names_text.upper():
        findings.append(_finding("structure.defined_name_ref", "Named range contem referencia #REF!."))

    checks_run += 1
    custom_names = set(workbook.custom_doc_props.names)
    missing_properties = [name for name in CUSTOM_PROPERTY_NAMES if name not in custom_names]
    if missing_properties:
        findings.append(_finding("metadata.custom_properties", f"Metadados customizados ausentes: {', '.join(missing_properties)}."))

    return AuditResult(findings=tuple(findings), checks_run=checks_run)


__all__ = ["audit_workbook"]
