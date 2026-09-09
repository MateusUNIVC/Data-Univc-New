from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
import logging
import os
import re
from io import BytesIO
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

from academic_analytics import (
    _current_calendar_periods,
    build_academic_dashboard,
    summarize_nps_semesters,
    to_semester,
)
from academic_catalog import DCS_COURSES, DTNH_COURSES
from excel_errors import ExcelExportLimitError


# Visual language intentionally follows the user's original DTNH workbook.
GREEN_DARK = "045233"
GREEN = "0B7A54"
GREEN_LIGHT = "E2F0D9"
TEAL_LIGHT = "DDEFE9"
YELLOW = "FFF6DC"
GRAY = "F5F6F5"
GRAY_TEXT = "5A5A5A"
BORDER = "D9E1DD"
RED_LIGHT = "FDE7E7"
ORANGE_LIGHT = "FFF4D8"
WHITE = "FFFFFF"
BLACK = "000000"
LINKED_GREEN = "008000"
INPUT_BLUE = "0000FF"
PURPLE = "7030A0"
PALE_GREEN = "F3F8F5"
PALE_TEAL = "E8F3F0"
PALE_BLUE = "EAF2FA"
PALE_GOLD = "FFF4D6"
SOFT_GRAY = "EEF1EF"
CHART_META = "7F8C8D"
CHART_PROMOTER = "0B7A54"
CHART_NEUTRAL = "D9B44A"
CHART_DETRACTOR = "C74B50"

DATA_START = 5
SERIES_START = 15
EXCEL_MAX_ROW = 1_048_576
DEFAULT_INPUT_RESERVE = max(25, int(os.getenv("ACADEMIC_EXCEL_INPUT_RESERVE_ROWS", "250")))
LOGGER = logging.getLogger(__name__)


def _semester_key(period: str | None) -> int | None:
    sem = to_semester(str(period or "").strip().upper())
    if not sem:
        return None
    return int(sem[:4]) * 10 + int(sem[-1])


def _month_key(period: str | None) -> int | None:
    text = str(period or "").strip()
    if len(text) != 7 or text[4] != "-":
        return None
    try:
        return int(text[:4]) * 12 + int(text[5:7])
    except Exception:
        return None


def _safe_date(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return value
    if value in (None, ""):
        return None
    text = str(value)
    try:
        return datetime.fromisoformat(text[:19])
    except Exception:
        return text




def _extract_year(value: Any) -> int | None:
    match = re.search(r"(?:^|\D)(\d{4})(?:\D|$)", str(value or ""))
    if not match:
        return None
    year = int(match.group(1))
    return year if 1900 <= year <= 9999 else None


def _dynamic_row_end(
    label: str,
    count: int,
    *,
    start_row: int = DATA_START,
    minimum_rows: int = 50,
    reserve: int | None = None,
) -> int:
    """Return an Excel row boundary that always includes every source row.

    Extra blank rows are only an editing convenience. Existing data is never
    truncated; if the physical Excel row limit is reached, export fails with an
    explicit message instead of producing an incomplete workbook.
    """
    actual_end = start_row - 1 + max(0, count)
    if actual_end > EXCEL_MAX_ROW:
        raise ExcelExportLimitError(
            f"A base {label} possui {count:,} linhas e excede o limite físico do Excel. "
            "A exportação foi interrompida para evitar perda silenciosa de dados."
        )
    reserve_rows = DEFAULT_INPUT_RESERVE if reserve is None else max(0, int(reserve))
    reserve_rows = max(reserve_rows, min(5000, max(0, count // 20)))
    desired_rows = max(minimum_rows, count + reserve_rows)
    # When the dataset is close to Excel's physical limit, shrink only the
    # optional blank reserve. Every real source row remains present.
    return min(start_row - 1 + desired_rows, EXCEL_MAX_ROW)

def _copy_font(font: Font, *, color: str | None = None, bold: bool | None = None, size: float | None = None) -> Font:
    return Font(
        name=font.name or "Arial",
        size=size if size is not None else font.sz,
        bold=font.bold if bold is None else bold,
        italic=font.italic,
        color=color if color is not None else (font.color.rgb[-6:] if font.color and font.color.type == "rgb" and font.color.rgb else BLACK),
    )


@dataclass
class AcademicExcelContext:
    directorate: str
    granularity: str = "semestral"
    reference: str | None = None
    comparison: str | None = None
    course: str | None = None
    discipline: str | None = None
    window_periods: int | str | None = None


class AcademicWorkbookBuilder:
    """Builds the academic workbook from a blank OOXML package.

    Deliberately does *not* clone sheet XML, workbook views, drawings or tables
    from previous templates. This avoids carrying invalid sheetViews and stale
    relationships that made Microsoft Excel enter repair mode.
    """

    def __init__(self, *, context: AcademicExcelContext, data: dict[str, Any]):
        self.context = context
        self.data = data
        self.code = context.directorate.upper()
        if self.code not in {"DTNH", "DCS"}:
            raise ValueError("AcademicWorkbookBuilder supports DTNH or DCS")
        self.prefix = self.code
        self.kpi_codes = [f"{self.prefix}-01", f"{self.prefix}-02", f"{self.prefix}-03"]
        self.canonical_courses = list(DTNH_COURSES if self.code == "DTNH" else DCS_COURSES)

        # Materialize source collections once. Every workbook range below is
        # derived from the real volume instead of hard-coded row ceilings.
        self.nps_rows = list(self.data.get("nps") or [])
        self.teacher_rows = list(self.data.get("teacher") or [])
        self.result_rows = list(self.data.get("results") or [])
        self.goal_rows = list(self.data.get("goals") or [])
        self.action_rows = list(self.data.get("actions") or [])
        self.discipline_rows = list(self.data.get("disciplines") or [])
        self.nps_summary_rows = summarize_nps_semesters(self.nps_rows)

        course_names: list[str] = []
        sources = [
            *self.canonical_courses,
            *[row.get("curso") for row in self.data.get("courses", []) if isinstance(row, dict)],
            *[row.get("curso") for row in self.discipline_rows if isinstance(row, dict)],
            *[row.get("curso") for row in self.nps_rows if isinstance(row, dict)],
            *[row.get("curso") for row in self.teacher_rows if isinstance(row, dict)],
            *[row.get("curso") for row in self.result_rows if isinstance(row, dict)],
        ]
        for value in sources:
            name = str(value or "").strip()
            if name and name not in course_names:
                course_names.append(name)
        self.course_names = course_names

        self.nps_end = _dynamic_row_end("NPS DISCENTES", len(self.nps_rows), minimum_rows=100)
        self.nps_summary_end = _dynamic_row_end("NPS SEMESTRAL", len(self.nps_summary_rows), minimum_rows=50, reserve=25)
        self.teacher_end = _dynamic_row_end("AVALIAÇÃO DOCENTE", len(self.teacher_rows), minimum_rows=100)
        self.result_end = _dynamic_row_end("RESULTADOS ACADÊMICOS", len(self.result_rows), minimum_rows=100)
        self.goal_end = _dynamic_row_end("METAS", len(self.goal_rows), minimum_rows=50)
        self.action_end = _dynamic_row_end(
            "PLANO DE AÇÃO", len(self.action_rows), start_row=8, minimum_rows=50
        )
        self.discipline_end = _dynamic_row_end("DISCIPLINAS", len(self.discipline_rows), minimum_rows=100)
        self.course_end = _dynamic_row_end("CURSOS", len(self.course_names), minimum_rows=25)

        current_month, current_semester = _current_calendar_periods()
        years = {int(current_month[:4])}
        for value in (context.reference, context.comparison, current_semester):
            year = _extract_year(value)
            if year:
                years.add(year)
        for collection, field in (
            (self.nps_rows, "periodo"),
            (self.teacher_rows, "periodo"),
            (self.result_rows, "periodo"),
            (self.goal_rows, "vigencia"),
        ):
            for row in collection:
                year = _extract_year(row.get(field) if isinstance(row, dict) else None)
                if year:
                    years.add(year)
        current_year = int(current_month[:4])
        self.calendar_start_year = min(min(years), current_year - 5)
        self.calendar_end_year = max(max(years), current_year + 2)
        self.semesters = [
            f"{year}-SEM{semester}"
            for year in range(self.calendar_start_year, self.calendar_end_year + 1)
            for semester in (1, 2)
        ]
        self.months = [
            f"{year}-{month:02d}"
            for year in range(self.calendar_start_year, self.calendar_end_year + 1)
            for month in range(1, 13)
        ]
        # Earliest periods with actual institutional data. The calendar keeps a
        # small future/input buffer, but "Todo histórico" must begin at the real
        # first observation rather than displaying years of empty categories.
        data_semesters = [
            sem
            for collection, field in (
                (self.nps_summary_rows, "periodo"),
                (self.teacher_rows, "periodo"),
                (self.result_rows, "periodo"),
            )
            for row in collection
            for sem in [to_semester(row.get(field) if isinstance(row, dict) else None)]
            if sem
        ]
        reference_semester = to_semester(context.reference) or current_semester
        if not data_semesters:
            data_semesters = [reference_semester]
        first_data_semester = min(data_semesters, key=lambda value: (int(value[:4]), int(value[-1])))
        first_year = int(first_data_semester[:4])
        first_semester_number = int(first_data_semester[-1])
        self.data_min_semester_key = first_year * 2 + first_semester_number
        self.data_min_month_key = first_year * 12 + (1 if first_semester_number == 1 else 7)

        self.dim_period_end = DATA_START - 1 + len(self.semesters)
        self.dim_month_end = DATA_START - 1 + len(self.months)
        self.period_list_end = DATA_START - 1 + len(self.months) + len(self.semesters)
        self.semester_list_end = DATA_START - 1 + len(self.semesters)
        self.month_list_end = DATA_START - 1 + len(self.months)
        self.course_list_end = DATA_START + len(self.course_names)
        discipline_names = sorted({
            str(row.get("disciplina") or "").strip()
            for row in self.discipline_rows if isinstance(row, dict) and str(row.get("disciplina") or "").strip()
        })
        self.discipline_names = discipline_names
        self.discipline_list_end = DATA_START + len(discipline_names)
        self.series_end = SERIES_START + max(24, len(self.months), len(self.semesters)) - 1
        self.nps_series_end = SERIES_START + max(1, len(self.semesters)) - 1
        # The source bases can span decades, but an executive chart must remain
        # legible. The panel therefore uses compact rolling windows while the
        # complete history remains available in NPS SEMESTRAL and CALC.
        self.chart_period_points = max(6, int(os.getenv("ACADEMIC_EXCEL_CHART_PERIODS", "12")))
        self.chart_nps_points = max(4, int(os.getenv("ACADEMIC_EXCEL_CHART_NPS_SEMESTERS", "8")))
        self.chart_period_end = SERIES_START + self.chart_period_points - 1
        self.chart_nps_end = SERIES_START + self.chart_nps_points - 1

        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.ws: dict[str, Any] = {}

        self.title_font = Font(name="Arial", size=15, bold=True, color=GREEN_DARK)
        self.section_font = Font(name="Arial", size=10, bold=True, color=WHITE)
        self.header_font = Font(name="Arial", size=9, bold=True, color=WHITE)
        self.body_font = Font(name="Arial", size=9, color=BLACK)
        self.static_font = Font(name="Arial", size=9, color=GRAY_TEXT)
        self.formula_font = Font(name="Arial", size=9, color=BLACK)
        self.linked_font = Font(name="Arial", size=9, color=LINKED_GREEN)
        self.input_font = Font(name="Arial", size=9, color=INPUT_BLUE)
        self.control_font = Font(name="Arial", size=9, bold=True, color=PURPLE)
        self.thin_bottom = Border(bottom=Side(style="thin", color=BORDER))

    # ------------------------------------------------------------------
    # Styling helpers
    # ------------------------------------------------------------------
    def _new_sheet(self, title: str):
        ws = self.wb.create_sheet(title)
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.outlinePr.summaryBelow = True
        self.ws[title] = ws
        return ws

    def _title(self, ws, text: str, subtitle: str | None = None, *, end_col: int = 8):
        # The reference workbook uses a calm white header, institutional green
        # typography and generous spacing. We reproduce that visual language
        # without copying any OOXML from the old template.
        if end_col > 1:
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
        ws["A1"] = text
        ws["A1"].font = self.title_font
        ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
        ws["A1"].border = Border(bottom=Side(style="medium", color=GREEN))
        ws.row_dimensions[1].height = 28
        if subtitle:
            if end_col > 1:
                ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_col)
            ws["A2"] = subtitle
            ws["A2"].font = self.static_font
            ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[2].height = 32

    def _section(self, ws, row: int, text: str, *, start_col: int = 1, end_col: int = 8):
        ws.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=end_col)
        cell = ws.cell(row, start_col)
        cell.value = text
        cell.fill = PatternFill("solid", fgColor=GREEN_DARK)
        cell.font = Font(name="Arial", size=11, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[row].height = 24

    def _headers(self, ws, row: int, labels: Iterable[str], *, start_col: int = 1):
        for offset, label in enumerate(labels):
            cell = ws.cell(row, start_col + offset)
            cell.value = label
            cell.fill = PatternFill("solid", fgColor=GREEN)
            cell.font = self.header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=Side(style="thin", color=GREEN_DARK))
        ws.row_dimensions[row].height = 30

    def _body_range(self, ws, min_row: int, max_row: int, min_col: int, max_col: int, *, font: Font | None = None):
        f = font or self.body_font
        for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
            for c in row:
                c.font = f
                c.alignment = Alignment(vertical="center")

    def _setup_base_sheet(self, ws, *, widths: dict[str, float], freeze: str = "A5"):
        for col, width in widths.items():
            ws.column_dimensions[col].width = width
        # A5 creates only a bottom-left pane. It never creates stale top-right
        # selections, which was the exact source of Excel's sheet7-9 repair.
        ws.freeze_panes = freeze
        ws.auto_filter.ref = f"A4:{get_column_letter(ws.max_column or 1)}4"

    def _add_validation(self, ws, cell_range: str, formula1: str, *, prompt: str | None = None):
        dv = DataValidation(type="list", formula1=formula1, allow_blank=True)
        if prompt:
            dv.promptTitle = "Data UNIVC"
            dv.prompt = prompt
            dv.showInputMessage = True
        ws.add_data_validation(dv)
        dv.add(cell_range)

    def _apply_status_cf(self, ws, rng: str):
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("Dentro",{rng.split(":")[0]}))'], fill=PatternFill("solid", fgColor=GREEN_LIGHT)))

    @staticmethod
    def _paint_range(ws, min_row: int, max_row: int, min_col: int, max_col: int, *, fill: str | None = None,
                     border: Border | None = None):
        for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
            for cell in row:
                if fill:
                    cell.fill = PatternFill("solid", fgColor=fill)
                if border:
                    cell.border = border

    def _kpi_card(self, ws, *, start_col: int, end_col: int, label_row: int, value_row: int,
                  label: str, formula: str, number_format: str = "General", fill: str = PALE_GREEN,
                  font_size: float = 16, value_color: str = BLACK):
        border = Border(
            left=Side(style="thin", color=BORDER), right=Side(style="thin", color=BORDER),
            top=Side(style="thin", color=BORDER), bottom=Side(style="thin", color=BORDER),
        )
        self._paint_range(ws, label_row, value_row + 1, start_col, end_col, fill=fill, border=border)
        ws.merge_cells(start_row=label_row, start_column=start_col, end_row=label_row, end_column=end_col)
        ws.merge_cells(start_row=value_row, start_column=start_col, end_row=value_row + 1, end_column=end_col)
        label_cell = ws.cell(label_row, start_col)
        label_cell.value = label
        label_cell.font = Font(name="Arial", size=8, bold=True, color=GRAY_TEXT)
        label_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        value_cell = ws.cell(value_row, start_col)
        value_cell.value = formula
        value_cell.font = Font(name="Arial", size=font_size, bold=True, color=value_color)
        value_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        value_cell.number_format = number_format
        ws.row_dimensions[label_row].height = 22
        ws.row_dimensions[value_row].height = 24
        ws.row_dimensions[value_row + 1].height = 24

    @staticmethod
    def _chart_base(chart, *, title: str, x_title: str, y_title: str, width: float = 25.5,
                    height: float = 9.0):
        chart.title = title
        chart.style = 13
        chart.width = width
        chart.height = height
        chart.x_axis.title = x_title
        chart.y_axis.title = y_title
        chart.legend.position = "b"
        chart.display_blanks = "gap"
        # All chart source ranges live in CALC helper columns. Excel normally
        # ignores hidden cells unless this flag is disabled, which was one of
        # the reasons the former panel opened with apparently empty charts.
        chart.visible_cells_only = False
        chart.y_axis.majorGridlines = None
        try:
            chart.x_axis.tickLblSkip = 1
            chart.x_axis.tickMarkSkip = 1
        except Exception:
            pass

    @staticmethod
    def _style_line_chart(chart, *, value_color: str = GREEN, goal_color: str = CHART_META,
                          value_labels: bool = True, number_format: str = "0.0"):
        if len(chart.series) >= 1:
            observed = chart.series[0]
            observed.graphicalProperties.line.solidFill = value_color
            observed.graphicalProperties.line.width = 28575
            observed.marker.symbol = "circle"
            observed.marker.size = 7
            observed.marker.graphicalProperties.solidFill = value_color
            observed.marker.graphicalProperties.line.solidFill = value_color
            if value_labels:
                labels = DataLabelList()
                labels.showVal = True
                labels.showLegendKey = False
                labels.showCatName = False
                labels.showSerName = False
                labels.dLblPos = "t"
                labels.numFmt = number_format
                observed.dLbls = labels
        if len(chart.series) >= 2:
            goal = chart.series[1]
            goal.graphicalProperties.line.solidFill = goal_color
            goal.graphicalProperties.line.width = 19050
            try:
                goal.graphicalProperties.line.prstDash = "dash"
            except Exception:
                pass
            goal.marker.symbol = "none"

    @staticmethod
    def _style_stacked_chart(chart):
        colors = [CHART_PROMOTER, CHART_NEUTRAL, CHART_DETRACTOR]
        for series, color in zip(chart.series, colors):
            series.graphicalProperties.solidFill = color
            series.graphicalProperties.line.solidFill = color
        labels = DataLabelList()
        labels.showVal = True
        labels.showLegendKey = False
        labels.showCatName = False
        labels.showSerName = False
        labels.dLblPos = "ctr"
        labels.numFmt = "#,##0"
        chart.dLbls = labels

    # ------------------------------------------------------------------
    # Workbook construction
    # ------------------------------------------------------------------
    def build(self) -> Workbook:
        self._resolve_initial_context()
        for name in [
            "LEIA-ME", "PARAMETROS", "PAINEL", "QUALIDADE E GOVERNANÇA", "MATRIZ",
            "PLANO_DE_ACAO", "NPS DISCENTES", "NPS SEMESTRAL", "AVALIAÇÃO DOCENTE", "RESULTADOS ACADÊMICOS",
            "METAS", "CURSOS", "DISCIPLINAS", "INDICADORES", "CALC", "LISTAS DE APOIO",
            "DIM_PERIODO", "DIM_MES",
        ]:
            self._new_sheet(name)

        self._build_lists_and_dimensions()
        self._build_courses()
        self._build_disciplines()
        self._build_indicators()
        self._build_nps()
        self._build_nps_semester()
        self._build_teacher_eval()
        self._build_results()
        self._build_goals()
        self._build_actions()
        self._build_parameters()
        self._build_calc()
        self._build_matrix()
        self._build_quality()
        self._build_panel()
        self._build_readme()
        self._defined_names()
        self._finish()
        return self.wb

    def _resolve_initial_context(self):
        gran = "mensal" if str(self.context.granularity or "").casefold() in {"mensal", "mes", "mês", "month", "monthly"} else "semestral"
        self.context.granularity = gran
        snapshot = {
            "courses": [x.get("curso") if isinstance(x, dict) else x for x in self.data.get("courses", [])],
            "disciplines": self.data.get("disciplines_map", {}),
            "metas": self.data.get("goals", []),
            "nps": self.data.get("nps", []),
            "avaliacao_docente": self.data.get("teacher", []),
            "resultados": self.data.get("results", []),
        }
        try:
            dash = build_academic_dashboard(
                snapshot,
                course=self.context.course or "(todos)",
                discipline=self.context.discipline or "(todas)",
                reference=self.context.reference,
                comparison=self.context.comparison,
                directorate_code=self.code,
                granularity=gran,
                window_semesters=None if isinstance(self.context.window_periods, str) else self.context.window_periods,
            )
            ctx = dash.get("contexto", {}) or {}
        except Exception as exc:
            LOGGER.exception("Falha ao resolver o contexto do Excel acadêmico", exc_info=exc)
            raise RuntimeError("Não foi possível calcular o contexto do painel acadêmico para a exportação.") from exc
        self.reference = self.context.reference or ctx.get("referencia")
        self.comparison = self.context.comparison or ctx.get("comparacao")
        current_month, current_semester = _current_calendar_periods()
        if not self.reference:
            self.reference = current_month if gran == "mensal" else current_semester
        if not self.comparison:
            if gran == "mensal":
                y, m = int(self.reference[:4]), int(self.reference[-2:])
                m -= 1
                if m == 0:
                    y -= 1; m = 12
                self.comparison = f"{y:04d}-{m:02d}"
            else:
                y = int(self.reference[:4]); s = int(self.reference[-1])
                self.comparison = f"{y-1:04d}-SEM2" if s == 1 else f"{y:04d}-SEM1"
        self.selected_course = self.context.course if self.context.course and self.context.course != "(todos)" else "(todos)"
        self.selected_discipline = self.context.discipline if self.selected_course != "(todos)" and self.context.discipline and self.context.discipline != "(todas)" else "(todas)"
        self.window = self.context.window_periods if self.context.window_periods is not None else (12 if gran == "mensal" else 4)

    def _build_lists_and_dimensions(self):
        ws = self.ws["LISTAS DE APOIO"]
        self._title(
            ws,
            "LISTAS DE APOIO",
            f"Calendário dinâmico de {self.calendar_start_year} a {self.calendar_end_year}; faixas usadas nas validações do painel.",
            end_col=11,
        )
        self._headers(ws, 4, ["Visualização", "Janela", "Desempenho", "Períodos", "Semestres", "Meses", "Status", "Cursos", "Disciplinas", "Situação", "Motivo"])
        for i, value in enumerate(["Semestral", "Mensal"], DATA_START):
            ws.cell(i, 1).value = value
        for i, value in enumerate([2, 4, 6, 12, 24, 36, 60, "Todo histórico"], DATA_START):
            ws.cell(i, 2).value = value
        for i, value in enumerate(["Semestre", "Geral"], DATA_START):
            ws.cell(i, 3).value = value
        for i, value in enumerate(["Não iniciado", "Em andamento", "Concluído", "Atrasado", "Cancelado"], DATA_START):
            ws.cell(i, 7).value = value
        for i, value in enumerate(["Aprovado", "Reprovado", "Cursando"], DATA_START):
            ws.cell(i, 10).value = value
        for i, value in enumerate(["nota", "falta", "outro"], DATA_START):
            ws.cell(i, 11).value = value

        wsp = self.ws["DIM_PERIODO"]
        self._title(wsp, "DIMENSÃO DE SEMESTRES", "Calendário acadêmico gerado a partir do histórico real e do período corrente.", end_col=6)
        self._headers(wsp, 4, ["Período", "Ano", "Semestre", "Início", "Fim", "Chave"])
        for r, label in enumerate(self.semesters, DATA_START):
            year = int(label[:4]); semester = int(label[-1])
            start = date(year, 1 if semester == 1 else 7, 1)
            end = date(year, 6, 30) if semester == 1 else date(year, 12, 31)
            values = [label, year, semester, start, end, year * 2 + semester]
            for c, value in enumerate(values, 1):
                wsp.cell(r, c).value = value
            wsp.cell(r, 4).number_format = "dd/mm/yyyy"
            wsp.cell(r, 5).number_format = "dd/mm/yyyy"
        wsp.freeze_panes = "A5"
        for col, width in {"A": 16, "B": 10, "C": 10, "D": 14, "E": 14, "F": 12}.items():
            wsp.column_dimensions[col].width = width
        wsp.auto_filter.ref = f"A4:F{self.dim_period_end}"

        wsm = self.ws["DIM_MES"]
        self._title(wsm, "DIMENSÃO DE MESES", "Cada mês aponta para o fechamento semestral acadêmico correspondente.", end_col=8)
        self._headers(wsm, 4, ["Mês", "Data", "Ano", "Nº mês", "Semestre", "Início semestre", "Rótulo", "Chave"])
        for r, label in enumerate(self.months, DATA_START):
            year = int(label[:4]); month = int(label[-2:])
            semester = f"{year}-SEM{1 if month <= 6 else 2}"
            values = [label, date(year, month, 1), year, month, semester, date(year, 1 if month <= 6 else 7, 1), f"{month:02d}/{year}", year * 12 + month]
            for c, value in enumerate(values, 1):
                wsm.cell(r, c).value = value
            wsm.cell(r, 2).number_format = "mm/yyyy"
            wsm.cell(r, 6).number_format = "dd/mm/yyyy"
        wsm.freeze_panes = "A5"
        for col, width in {"A": 12, "B": 12, "C": 10, "D": 10, "E": 16, "F": 16, "G": 12, "H": 12}.items():
            wsm.column_dimensions[col].width = width
        wsm.auto_filter.ref = f"A4:H{self.dim_month_end}"

        all_periods = self.months + self.semesters
        for i, value in enumerate(all_periods, DATA_START):
            ws.cell(i, 4).value = value
        for i, value in enumerate(self.semesters, DATA_START):
            ws.cell(i, 5).value = value
        for i, value in enumerate(self.months, DATA_START):
            ws.cell(i, 6).value = value

        ws.cell(DATA_START, 8).value = "(todos)"
        for i, value in enumerate(self.course_names, DATA_START + 1):
            ws.cell(i, 8).value = value
        ws.cell(DATA_START, 9).value = "(todas)"
        for i, value in enumerate(self.discipline_names, DATA_START + 1):
            ws.cell(i, 9).value = value
        for c in range(1, 12):
            ws.column_dimensions[get_column_letter(c)].width = 18
        ws.column_dimensions["D"].width = 16
        ws.column_dimensions["H"].width = 46
        ws.column_dimensions["I"].width = 48
        ws.freeze_panes = "A5"

    def _build_courses(self):
        ws = self.ws["CURSOS"]
        self._title(ws, f"CURSOS — {self.code}", "Catálogo acadêmico usado pelo painel.", end_col=4)
        self._headers(ws, 4, ["Curso", "Ativo", "Vigência início", "Vigência fim"])
        by_name = {str(x.get("curso")): x for x in self.data.get("courses", []) if isinstance(x, dict)}
        names = list(self.course_names)
        for r, name in enumerate(names, DATA_START):
            item = by_name.get(name, {})
            vals = [name, "S" if item.get("ativo", True) else "N", item.get("vigencia_inicio") or "", item.get("vigencia_fim") or ""]
            for c, v in enumerate(vals, 1):
                ws.cell(r, c).value = v
                ws.cell(r, c).font = self.linked_font
        self._body_range(ws, 5, self.course_end, 1, 4)
        for col, width in {"A":46,"B":10,"C":18,"D":18}.items(): ws.column_dimensions[col].width=width
        ws.freeze_panes = "A5"; ws.auto_filter.ref=f"A4:D{self.course_end}"

    def _build_disciplines(self):
        ws = self.ws["DISCIPLINAS"]
        self._title(ws, "DISCIPLINAS", "Cadastro vinculado ao curso. A lista pode ser ampliada pelas importações do SEI.", end_col=5)
        self._headers(ws, 4, ["Curso", "Disciplina", "Ativa", "Vigência início", "Vigência fim"])
        for r, item in enumerate(self.discipline_rows, DATA_START):
            vals = [item.get("curso"), item.get("disciplina"), "S" if item.get("ativo", True) else "N", item.get("vigencia_inicio") or "", item.get("vigencia_fim") or ""]
            for c, v in enumerate(vals, 1): ws.cell(r, c).value=v; ws.cell(r,c).font=self.linked_font
        self._body_range(ws, 5, self.discipline_end, 1, 5)
        for col, width in {"A":42,"B":48,"C":10,"D":18,"E":18}.items(): ws.column_dimensions[col].width=width
        ws.freeze_panes="A5"; ws.auto_filter.ref=f"A4:E{self.discipline_end}"

    def _build_indicators(self):
        ws=self.ws["INDICADORES"]
        self._title(ws, f"INDICADORES — {self.code}", "Definições usadas no painel e nas metas.", end_col=8)
        self._headers(ws,4,["Código","Indicador","Unidade","Direção","Periodicidade","Fonte","Responsável","Regra"])
        rows=[
            [self.kpi_codes[0],"NPS Discente","pontos","Maior é melhor","Semestral","NPS","Acadêmico","Fechamento semestral por curso; mensal apenas como fallback quando não houver fechamento"],
            [self.kpi_codes[1],"Avaliação Docente pelo Aluno","nota 0–10","Maior é melhor","Semestral","Avaliação docente","Acadêmico","Média ponderada pelo número de respondentes"],
            [self.kpi_codes[2],"Taxa de Aprovação e Desempenho Acadêmico","%","Maior é melhor","Semestral","Mapa de notas/SEI","Acadêmico","Aprovados / resultados finalizados"],
        ]
        for r, vals in enumerate(rows,5):
            for c,v in enumerate(vals,1): ws.cell(r,c).value=v; ws.cell(r,c).font=self.static_font
        for col,width in {"A":14,"B":40,"C":16,"D":18,"E":18,"F":22,"G":18,"H":60}.items(): ws.column_dimensions[col].width=width

    def _build_nps(self):
        ws=self.ws["NPS DISCENTES"]
        self._title(ws, f"ENTRADA — {self.code}-01 · NPS Discente", "Aceita AAAA-MM e AAAA-SEM1/SEM2. O fechamento semestral é oficial; lançamentos mensais servem apenas como fallback quando o curso não possui fechamento no semestre.", end_col=12)
        headers=["Período","Tipo período","Semestre","Curso","Respondentes","Promotores","Neutros","Detratores","NPS","Data de lançamento","Lançado por","Validação da linha"]
        self._headers(ws,4,headers)
        rows=list(self.nps_rows)
        for r in range(DATA_START,self.nps_end+1):
            # Formulas are intentionally simple and standard Excel formulas.
            ws.cell(r,2).value=f'=IF(A{r}="","",IF(LEN(A{r})=7,"Mensal",IF(OR(RIGHT(A{r},4)="SEM1",RIGHT(A{r},4)="SEM2"),"Semestral","Inválido")))'
            ws.cell(r,3).value=f'=IF(A{r}="","",IF(LEN(A{r})=7,LEFT(A{r},4)&"-SEM"&IF(VALUE(RIGHT(A{r},2))<=6,1,2),A{r}))'
            ws.cell(r,9).value=f'=IF(E{r}>0,100*(F{r}-H{r})/E{r},"")'
            ws.cell(r,12).value=f'=IF(COUNTA(A{r},D{r}:H{r})=0,"",IF(OR(A{r}="",D{r}="",E{r}<0,F{r}<0,G{r}<0,H{r}<0,F{r}+G{r}+H{r}<>E{r}),"ERRO","OK"))'
            for c in (2,3,9,12): ws.cell(r,c).font=self.formula_font
            if r <= 4+len(rows):
                item=rows[r-5]
                vals=[item.get("periodo"),None,None,item.get("curso"),item.get("respondentes"),item.get("promotores"),item.get("neutros"),item.get("detratores"),None,_safe_date(item.get("data_lancamento")),item.get("lancado_por"),None]
                for c,v in enumerate(vals,1):
                    if c in (2,3,9,12): continue
                    ws.cell(r,c).value=v; ws.cell(r,c).font=self.linked_font
        ws.column_dimensions["A"].width=16; ws.column_dimensions["B"].width=14; ws.column_dimensions["C"].width=16; ws.column_dimensions["D"].width=46
        for c in range(5,10): ws.column_dimensions[get_column_letter(c)].width=14
        ws.column_dimensions["J"].width=20; ws.column_dimensions["K"].width=28; ws.column_dimensions["L"].width=20
        for r in range(5,self.nps_end+1):
            for c in (5,6,7,8): ws.cell(r,c).number_format='0'
            ws.cell(r,9).number_format='0.0'
            ws.cell(r,10).number_format='dd/mm/yyyy hh:mm'
        self._body_range(ws,5,self.nps_end,1,12)
        ws.freeze_panes="A5"; ws.auto_filter.ref=f"A4:L{self.nps_end}"
        self._add_validation(ws,f"A5:A{self.nps_end}","LST_PERIODOS")
        self._add_validation(ws,f"D5:D{self.nps_end}","LST_CURSOS_BASE")

    def _build_nps_semester(self):
        ws = self.ws["NPS SEMESTRAL"]
        self._title(
            ws,
            f"CONSOLIDADO OFICIAL — {self.code}-01 · NPS por semestre",
            "Uma linha por semestre e curso. O fechamento semestral prevalece; registros mensais são agregados apenas quando o curso não possui fechamento no semestre.",
            end_col=11,
        )
        headers = [
            "Semestre", "Curso", "Fonte oficial", "Linhas fonte", "Respondentes",
            "Promotores", "Neutros", "Detratores", "NPS", "Último lançamento", "Validação",
        ]
        self._headers(ws, 4, headers)
        for r in range(DATA_START, self.nps_summary_end + 1):
            ws.cell(r, 9).value = f'=IF(E{r}>0,100*(F{r}-H{r})/E{r},"")'
            ws.cell(r, 11).value = (
                f'=IF(COUNTA(A{r}:H{r})=0,"",IF(OR(A{r}="",B{r}="",E{r}<0,F{r}<0,G{r}<0,H{r}<0,'
                f'F{r}+G{r}+H{r}<>E{r}),"ERRO","OK"))'
            )
            ws.cell(r, 9).font = self.formula_font
            ws.cell(r, 11).font = self.formula_font
            if r <= DATA_START - 1 + len(self.nps_summary_rows):
                item = self.nps_summary_rows[r - DATA_START]
                values = [
                    item.get("periodo"), item.get("curso"), item.get("fonte_nps"),
                    item.get("linhas_fonte"), item.get("respondentes"), item.get("promotores"),
                    item.get("neutros"), item.get("detratores"), None,
                    _safe_date(item.get("data_lancamento")), None,
                ]
                for c, value in enumerate(values, 1):
                    if c in (9, 11):
                        continue
                    ws.cell(r, c).value = value
                    ws.cell(r, c).font = self.linked_font
        widths = {
            "A": 16, "B": 46, "C": 30, "D": 14, "E": 16, "F": 14,
            "G": 14, "H": 14, "I": 12, "J": 20, "K": 18,
        }
        for col, width in widths.items():
            ws.column_dimensions[col].width = width
        for r in range(DATA_START, self.nps_summary_end + 1):
            for c in (4, 5, 6, 7, 8):
                ws.cell(r, c).number_format = "0"
            ws.cell(r, 9).number_format = "0.0"
            ws.cell(r, 10).number_format = "dd/mm/yyyy hh:mm"
        self._body_range(ws, DATA_START, self.nps_summary_end, 1, 11)
        ws.freeze_panes = "A5"
        ws.auto_filter.ref = f"A4:K{self.nps_summary_end}"

    def _build_teacher_eval(self):
        ws=self.ws["AVALIAÇÃO DOCENTE"]
        self._title(ws, f"ENTRADA — {self.code}-02 · Avaliação Docente pelo Aluno", "Base semestral. Na visão mensal o mesmo resultado é apresentado de janeiro a junho ou de julho a dezembro.", end_col=10)
        headers=["Período","Curso","Disciplina","Professor","Respondentes","Nota média","Soma ponderada","Data de lançamento","Lançado por","Validação da linha"]
        self._headers(ws,4,headers)
        rows=list(self.teacher_rows)
        for r in range(DATA_START,self.teacher_end+1):
            ws.cell(r,7).value=f'=IF(OR(E{r}="",F{r}=""),"",E{r}*F{r})'
            ws.cell(r,10).value=f'=IF(COUNTA(A{r}:F{r})=0,"",IF(OR(A{r}="",B{r}="",C{r}="",D{r}="",E{r}<0,F{r}<0,F{r}>10),"ERRO","OK"))'
            ws.cell(r,7).font=self.formula_font; ws.cell(r,10).font=self.formula_font
            if r <= 4+len(rows):
                item=rows[r-5]
                vals=[item.get("periodo"),item.get("curso"),item.get("disciplina"),item.get("professor"),item.get("respondentes"),item.get("nota_media"),None,_safe_date(item.get("data_lancamento")),item.get("lancado_por"),None]
                for c,v in enumerate(vals,1):
                    if c in (7,10): continue
                    ws.cell(r,c).value=v; ws.cell(r,c).font=self.linked_font
        widths={"A":16,"B":42,"C":48,"D":34,"E":14,"F":14,"G":16,"H":20,"I":28,"J":20}
        for col,w in widths.items(): ws.column_dimensions[col].width=w
        for r in range(5,self.teacher_end+1):
            ws.cell(r,5).number_format='0'; ws.cell(r,6).number_format='0.00'; ws.cell(r,7).number_format='0.00'; ws.cell(r,8).number_format='dd/mm/yyyy hh:mm'
        self._body_range(ws,5,self.teacher_end,1,10)
        ws.freeze_panes="A5"; ws.auto_filter.ref=f"A4:J{self.teacher_end}"
        self._add_validation(ws,f"A5:A{self.teacher_end}","LST_SEMESTRES")
        self._add_validation(ws,f"B5:B{self.teacher_end}","LST_CURSOS_BASE")
        self._add_validation(ws,f"C5:C{self.teacher_end}","LST_DISCIPLINAS_BASE")

    def _build_results(self):
        ws=self.ws["RESULTADOS ACADÊMICOS"]
        self._title(ws, f"ENTRADA AGREGADA — {self.code}-03 · Aprovação e Notas", "Uma linha por semestre + curso + disciplina. Registros individuais permanecem no sistema, não no painel Excel.", end_col=16)
        headers=["Período","Curso","Disciplina","Total","Finalizados","Aprovados","Reprovados","Reprovados por nota","Reprovados por falta","Reprovados outros","Em andamento","Quantidade de notas","Soma das notas","Média","Taxa de aprovação","Validação da linha"]
        self._headers(ws,4,headers)
        rows=list(self.result_rows)
        for r in range(DATA_START,self.result_end+1):
            ws.cell(r,7).value=f'=IF(COUNTA(A{r}:C{r})=0,"",SUM(H{r}:J{r}))'
            ws.cell(r,14).value=f'=IF(L{r}>0,M{r}/L{r},"")'
            # IMPORTANT: F/E is stored as a fraction and formatted as %, never *100.
            ws.cell(r,15).value=f'=IF(E{r}>0,F{r}/E{r},"")'
            ws.cell(r,16).value=(
                f'=IF(COUNTA(A{r}:C{r})=0,"",IF(OR(D{r}<0,E{r}<0,F{r}<0,H{r}<0,I{r}<0,J{r}<0,K{r}<0,L{r}<0),'
                f'"ERRO: valores inválidos",IF(D{r}<>E{r}+K{r},"ERRO: total inconsistente",IF(E{r}<>F{r}+G{r},'
                f'"ERRO: finalizados inconsistentes",IF(G{r}<>SUM(H{r}:J{r}),"ERRO: reprovações inconsistentes","OK")))))'
            )
            for c in (7,14,15,16): ws.cell(r,c).font=self.formula_font
            if r <= 4+len(rows):
                item=rows[r-5]
                vals=[item.get("periodo"),item.get("curso"),item.get("disciplina"),item.get("total_registros"),item.get("finalizados"),item.get("aprovados"),None,item.get("reprovados_nota"),item.get("reprovados_falta"),item.get("reprovados_outro"),item.get("em_andamento"),item.get("notas_contagem"),item.get("soma_notas"),None,None,None]
                for c,v in enumerate(vals,1):
                    if c in (7,14,15,16): continue
                    ws.cell(r,c).value=v; ws.cell(r,c).font=self.linked_font
        widths={"A":16,"B":42,"C":48,"D":12,"E":12,"F":12,"G":12,"H":18,"I":18,"J":18,"K":14,"L":18,"M":16,"N":12,"O":18,"P":28}
        for col,w in widths.items(): ws.column_dimensions[col].width=w
        for r in range(5,self.result_end+1):
            for c in range(4,13): ws.cell(r,c).number_format='0'
            ws.cell(r,13).number_format='0.00'; ws.cell(r,14).number_format='0.00'; ws.cell(r,15).number_format='0.0%'
        self._body_range(ws,5,self.result_end,1,16)
        ws.freeze_panes="A5"; ws.auto_filter.ref=f"A4:P{self.result_end}"
        self._add_validation(ws,f"A5:A{self.result_end}","LST_SEMESTRES")
        self._add_validation(ws,f"B5:B{self.result_end}","LST_CURSOS_BASE")
        self._add_validation(ws,f"C5:C{self.result_end}","LST_DISCIPLINAS_BASE")

    def _build_goals(self):
        ws=self.ws["METAS"]
        self._title(ws, "METAS", "Vigência acadêmica: SEM1 começa em janeiro e SEM2 começa em julho. A chave numérica evita interpretações por mês final do semestre.", end_col=9)
        self._headers(ws,4,["Indicador","Recorte","Vigência","Chave vigência","Meta","Atenção","Limite superior","Justificativa","Validação"])
        rows=self.goal_rows
        for r in range(DATA_START,self.goal_end+1):
            ws.cell(r,4).value=f'=IF(C{r}="","",VALUE(LEFT(C{r},4))*10+VALUE(RIGHT(C{r},1)))'
            ws.cell(r,9).value=f'=IF(COUNTA(A{r}:C{r},E{r}:F{r})=0,"",IF(OR(A{r}="",B{r}="",C{r}="",E{r}="",F{r}=""),"ERRO","OK"))'
            ws.cell(r,4).font=self.formula_font; ws.cell(r,9).font=self.formula_font
            if r <= 4+len(rows):
                item=rows[r-5]
                vals=[item.get("indicador"),item.get("recorte"),item.get("vigencia"),None,item.get("meta"),item.get("atencao"),item.get("limite_superior"),item.get("justificativa"),None]
                for c,v in enumerate(vals,1):
                    if c in (4,9): continue
                    ws.cell(r,c).value=v; ws.cell(r,c).font=self.linked_font
        widths={"A":14,"B":52,"C":16,"D":16,"E":14,"F":14,"G":16,"H":54,"I":18}
        for col,w in widths.items(): ws.column_dimensions[col].width=w
        self._body_range(ws,5,self.goal_end,1,9)
        ws.freeze_panes="A5"; ws.auto_filter.ref=f"A4:I{self.goal_end}"
        self._add_validation(ws,f"A5:A{self.goal_end}","LST_INDICADORES")
        self._add_validation(ws,f"C5:C{self.goal_end}","LST_SEMESTRES")

    def _build_actions(self):
        ws=self.ws["PLANO_DE_ACAO"]
        self._title(ws, "PLANO DE AÇÃO", "Ações associadas a indicadores e recortes que exigem acompanhamento.", end_col=14)
        self._headers(ws,7,["Nº","Período","Indicador","Recorte","Resultado","Meta","Problema","Causa provável","Ação corretiva","Responsável","Prazo","Meta da ação","Status","Dias p/ prazo"])
        rows=self.action_rows
        for r,item in enumerate(rows,8):
            vals=[item.get("numero"),item.get("mes"),item.get("indicador"),item.get("recorte"),item.get("resultado"),item.get("meta"),item.get("problema"),item.get("causa"),item.get("acao"),item.get("responsavel"),item.get("prazo"),item.get("meta_acao"),item.get("status"),item.get("dias_prazo")]
            for c,v in enumerate(vals,1): ws.cell(r,c).value=v; ws.cell(r,c).font=self.linked_font
        for r in range(8,self.action_end+1):
            if ws.cell(r,14).value is None:
                ws.cell(r,14).value=f'=IF(OR(K{r}="",M{r}="Concluído"),"",K{r}-TODAY())'
                ws.cell(r,14).font=self.formula_font
        widths=[8,14,14,38,12,12,38,38,44,26,14,24,18,14]
        for c,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(c)].width=w
        self._body_range(ws,8,self.action_end,1,14)
        ws.freeze_panes="A8"; ws.auto_filter.ref=f"A7:N{self.action_end}"
        self._add_validation(ws,f"C8:C{self.action_end}","LST_INDICADORES")
        self._add_validation(ws,f"M8:M{self.action_end}","LST_STATUS")

    def _build_parameters(self):
        ws=self.ws["PARAMETROS"]
        self._title(ws, "PARÂMETROS DO PAINEL", "Seleções válidas para DTNH/DCS. As células amarelas são controles do usuário.", end_col=6)
        self._section(ws,4,"SELEÇÃO",end_col=6)
        labels=[
            (5,"Visualização", "Mensal ou Semestral."),
            (7,"Período de referência", "Período analisado."),
            (8,"Período de comparação", "Mantido mesmo quando estiver além da janela mínima."),
            (9,"Janela mínima", "Quantidade mínima de períodos mostrados; a comparação nunca é descartada."),
            (11,"Curso em foco", "(todos) para consolidado."),
            (12,"Disciplina em foco", "(todas) para consolidado."),
            (14,"Desempenho entre cursos", "Geral ou um semestre específico."),
            (15,"Semestre do desempenho", "Usado quando o modo é Semestre."),
        ]
        current_semester = _current_calendar_periods()[1]
        initial={5:"Mensal" if self.context.granularity=="mensal" else "Semestral",7:self.reference,8:self.comparison,9:self.window,11:self.selected_course,12:self.selected_discipline,14:"Semestre",15:to_semester(self.reference) or current_semester}
        for r,label,note in labels:
            ws.cell(r,1).value=label; ws.cell(r,1).font=Font(name="Arial",size=10,bold=True,color=BLACK)
            ws.cell(r,2).value=initial[r]; ws.cell(r,2).fill=PatternFill("solid",fgColor=YELLOW); ws.cell(r,2).font=self.control_font; ws.cell(r,2).alignment=Alignment(horizontal="center")
            ws.cell(r,4).value=note; ws.cell(r,4).font=self.static_font
            ws.merge_cells(start_row=r,start_column=4,end_row=r,end_column=6)
        self._section(ws,18,"DIAGNÓSTICO / CHAVES",end_col=6)
        diagnostics=[
            (19,"Semestre da referência", '=IF($B$5="Mensal",LEFT($B$7,4)&"-SEM"&IF(VALUE(RIGHT($B$7,2))<=6,1,2),$B$7)'),
            (20,"Chave da referência", '=IF($B$5="Mensal",VALUE(LEFT($B$7,4))*12+VALUE(RIGHT($B$7,2)),VALUE(LEFT($B$7,4))*2+VALUE(RIGHT($B$7,1)))'),
            (21,"Chave da comparação", '=IF($B$5="Mensal",VALUE(LEFT($B$8,4))*12+VALUE(RIGHT($B$8,2)),VALUE(LEFT($B$8,4))*2+VALUE(RIGHT($B$8,1)))'),
            (22,"Início efetivo da janela", '=MIN($B$21,IF($B$9="Todo histórico",IF($B$5="Mensal",$B$26,$B$27),$B$20-IFERROR(VALUE($B$9),4)+1))'),
            (23,"Fim efetivo da janela", '=$B$20'),
            (24,"Semestre da comparação", '=IF($B$5="Mensal",LEFT($B$8,4)&"-SEM"&IF(VALUE(RIGHT($B$8,2))<=6,1,2),$B$8)'),
            (25,"Semestre anterior para NPS", '=IF($B$24<>$B$19,$B$24,IF(RIGHT($B$19,1)="1",(VALUE(LEFT($B$19,4))-1)&"-SEM2",LEFT($B$19,4)&"-SEM1"))'),
            (26,"Primeira chave mensal com dados", self.data_min_month_key),
            (27,"Primeira chave semestral com dados", self.data_min_semester_key),
        ]
        for r,label,formula in diagnostics:
            ws.cell(r,1).value=label; ws.cell(r,1).font=self.static_font
            ws.cell(r,2).value=formula; ws.cell(r,2).font=self.formula_font; ws.cell(r,2).fill=PatternFill("solid",fgColor=GRAY)
        for col,w in {"A":30,"B":24,"C":4,"D":32,"E":18,"F":18}.items(): ws.column_dimensions[col].width=w
        self._add_validation(ws,"B5","LST_VISUALIZACAO")
        self._add_validation(ws,"B7","LST_PERIODOS")
        self._add_validation(ws,"B8","LST_PERIODOS")
        self._add_validation(ws,"B9","LST_JANELA")
        self._add_validation(ws,"B11","LST_CURSOS")
        self._add_validation(ws,"B12","LST_DISCIPLINAS")
        self._add_validation(ws,"B14","LST_DESEMPENHO")
        self._add_validation(ws,"B15","LST_SEMESTRES")

    def _sumifs(self, sum_range: str, criteria: list[tuple[str,str]]) -> str:
        args=[sum_range]
        for rng,crit in criteria:
            args.extend([rng,crit])
        return "SUMIFS("+",".join(args)+")"

    def _build_calc(self):
        ws = self.ws["CALC"]
        self._title(
            ws,
            "CAMADA DE PROCESSAMENTO — NÃO EDITAR",
            "Cálculos dinâmicos do painel. O NPS usa exclusivamente o consolidado semestral oficial; colunas auxiliares ficam ocultas.",
            end_col=18,
        )
        self._section(ws, 3, "CONTEXTO RESOLVIDO", end_col=18)
        context_labels = [
            ("A4", "Visualização", "B4", "=PARAMETROS!B5"),
            ("D4", "Referência", "E4", "=PARAMETROS!B7"),
            ("G4", "Comparação", "H4", "=PARAMETROS!B8"),
            ("J4", "Chave referência", "K4", "=PARAMETROS!B20"),
            ("A5", "Chave comparação", "B5", "=PARAMETROS!B21"),
            ("D5", "Início efetivo", "E5", "=PARAMETROS!B22"),
            ("G5", "Fim efetivo", "H5", "=PARAMETROS!B23"),
            ("J5", "Escopo", "K5", '=IF(PARAMETROS!B12<>"(todas)",IF(PARAMETROS!B11="(todos)","TOTAL",PARAMETROS!B11&" » "&PARAMETROS!B12),IF(PARAMETROS!B11<>"(todos)",PARAMETROS!B11,"TOTAL"))'),
        ]
        for label_cell, label, value_cell, formula in context_labels:
            ws[label_cell] = label
            ws[label_cell].font = self.static_font
            ws[value_cell] = formula
            ws[value_cell].font = self.formula_font
            ws[value_cell].fill = PatternFill("solid", fgColor=GRAY)

        self._section(ws, 7, "RESUMO EXECUTIVO", end_col=18)
        self._headers(ws, 8, [
            "KPI", "Referência", "Comparação", "Variação", "Meta vigente", "Atenção", "Status",
            "Aprovados", "Rep. nota", "Rep. falta", "Média nota", "Escopo da meta", "Vigência da meta",
            "Respondentes NPS", "Promotores", "Neutros", "Detratores", "Fonte NPS",
        ])
        labels = [
            ("NPS Discente", self.kpi_codes[0], "D", "E", "AC", "AD"),
            ("Avaliação Docente", self.kpi_codes[1], "F", "G", "AE", "AF"),
            ("Taxa de Aprovação", self.kpi_codes[2], "H", "I", "AG", "AH"),
        ]
        for idx, (label, code, value_col, meta_col, scope_col, key_col) in enumerate(labels, 9):
            ws.cell(idx, 1).value = label
            ws.cell(idx, 2).value = f'=IFERROR(INDEX(${value_col}${SERIES_START}:${value_col}${self.series_end},MATCH(PARAMETROS!$B$7,$A${SERIES_START}:$A${self.series_end},0)),"")'
            ws.cell(idx, 3).value = f'=IFERROR(INDEX(${value_col}${SERIES_START}:${value_col}${self.series_end},MATCH(PARAMETROS!$B$8,$A${SERIES_START}:$A${self.series_end},0)),"")'
            ws.cell(idx, 4).value = f'=IF(OR(B{idx}="",C{idx}=""),"",B{idx}-C{idx})'
            ws.cell(idx, 5).value = f'=IFERROR(INDEX(${meta_col}${SERIES_START}:${meta_col}${self.series_end},MATCH(PARAMETROS!$B$7,$A${SERIES_START}:$A${self.series_end},0)),"")'
            match = f'MATCH(PARAMETROS!$B$7,$A${SERIES_START}:$A${self.series_end},0)'
            scope = f'INDEX(${scope_col}${SERIES_START}:${scope_col}${self.series_end},{match})'
            key = f'INDEX(${key_col}${SERIES_START}:${key_col}${self.series_end},{match})'
            ws.cell(idx, 6).value = f'=IFERROR(SUMIFS(METAS!$F$5:$F${self.goal_end},METAS!$A$5:$A${self.goal_end},"{code}",METAS!$B$5:$B${self.goal_end},{scope},METAS!$D$5:$D${self.goal_end},{key}),"")'
            ws.cell(idx, 7).value = f'=IF(B{idx}="","Sem dados",IF(E{idx}="","Sem meta",IF(B{idx}>=E{idx},"Dentro da meta",IF(AND(F{idx}<>"",B{idx}>=F{idx}),"Atenção","Fora da meta"))))'
            if idx == 9:
                for column, series_column in ((14, "N"), (15, "O"), (16, "P"), (17, "Q"), (18, "R")):
                    ws.cell(idx, column).value = f'=IFERROR(INDEX(${series_column}${SERIES_START}:${series_column}${self.series_end},{match}),"")'
            if idx == 11:
                for column, series_column in ((8, "V"), (9, "W"), (10, "X"), (11, "M")):
                    ws.cell(idx, column).value = f'=IFERROR(INDEX(${series_column}${SERIES_START}:${series_column}${self.series_end},{match}),"")'
            ws.cell(idx, 12).value = f'=IFERROR({scope},"")'
            ws.cell(idx, 13).value = f'=IFERROR(INT({key}/10)&"-SEM"&MOD({key},10),"")'

        for r in range(9, 12):
            for c in range(2, 19):
                ws.cell(r, c).font = self.formula_font
        for c in (2, 3, 4, 5, 6):
            ws.cell(9, c).number_format = "0.0"
            ws.cell(10, c).number_format = "0.00"
            ws.cell(11, c).number_format = '0.0"%"'
        for c in (8, 9, 10, 14, 15, 16, 17):
            ws.cell(9, c).number_format = "0"
            ws.cell(11, c).number_format = "0"
        ws.cell(11, 11).number_format = "0.00"

        self._section(ws, 13, "SÉRIE DO GRÁFICO", end_col=18)
        self._headers(ws, 14, [
            "Período", "Chave", "Semestre", "NPS", "Meta NPS", "Avaliação", "Meta avaliação",
            "Aprovação %", "Meta aprovação", "Aprovados", "Rep. nota", "Rep. falta", "Média notas",
            "Respondentes", "Promotores", "Neutros", "Detratores", "Fonte NPS",
        ])
        helpers = [
            "Aval resp", "Aval soma", "Finalizados", "Aprovados", "Rep nota", "Rep falta",
            "Qtd notas", "Soma notas", "Chave sem", "Escopo desejado", "Escopo NPS", "Key NPS",
            "Escopo Aval", "Key Aval", "Escopo Apr", "Key Apr",
        ]
        for i, header in enumerate(helpers, 19):
            ws.cell(14, i).value = header
            ws.cell(14, i).font = self.header_font
            ws.cell(14, i).fill = PatternFill("solid", fgColor=GREEN_DARK)

        for r in range(SERIES_START, self.series_end + 1):
            offset = r - SERIES_START
            ws.cell(r, 2).value = f'=IF(PARAMETROS!$B$22+{offset}>PARAMETROS!$B$23,"",PARAMETROS!$B$22+{offset})'
            ws.cell(r, 1).value = (
                f'=IF(B{r}="","",IF($B$4="Mensal",'
                f'IFERROR(INDEX(DIM_MES!$A$5:$A${self.dim_month_end},MATCH(B{r},DIM_MES!$H$5:$H${self.dim_month_end},0)),""),'
                f'IFERROR(INDEX(DIM_PERIODO!$A$5:$A${self.dim_period_end},MATCH(B{r},DIM_PERIODO!$F$5:$F${self.dim_period_end},0)),"")))'
            )
            ws.cell(r, 3).value = (
                f'=IF(A{r}="","",IF($B$4="Mensal",'
                f'IFERROR(INDEX(DIM_MES!$E$5:$E${self.dim_month_end},MATCH(B{r},DIM_MES!$H$5:$H${self.dim_month_end},0)),""),A{r}))'
            )
            course = 'IF(PARAMETROS!$B$11="(todos)","*",PARAMETROS!$B$11)'
            discipline = 'IF(PARAMETROS!$B$12="(todas)","*",PARAMETROS!$B$12)'

            nps_criteria = [
                (f"'NPS SEMESTRAL'!$A$5:$A${self.nps_summary_end}", f'$C{r}'),
                (f"'NPS SEMESTRAL'!$B$5:$B${self.nps_summary_end}", course),
                (f"'NPS SEMESTRAL'!$K$5:$K${self.nps_summary_end}", '"OK"'),
            ]
            for column, source in ((14, "E"), (15, "F"), (16, "G"), (17, "H")):
                ws.cell(r, column).value = '=' + self._sumifs(
                    f"'NPS SEMESTRAL'!${source}$5:${source}${self.nps_summary_end}", nps_criteria
                )
            ws.cell(r, 4).value = f'=IF(N{r}>0,100*(O{r}-Q{r})/N{r},"")'
            semester_count = (
                f'COUNTIFS(\'NPS SEMESTRAL\'!$A$5:$A${self.nps_summary_end},$C{r},'
                f'\'NPS SEMESTRAL\'!$B$5:$B${self.nps_summary_end},{course},'
                f'\'NPS SEMESTRAL\'!$C$5:$C${self.nps_summary_end},"Fechamento semestral",'
                f'\'NPS SEMESTRAL\'!$K$5:$K${self.nps_summary_end},"OK")'
            )
            fallback_count = semester_count.replace('"Fechamento semestral"', '"Agregado mensal (fallback)"')
            ws.cell(r, 18).value = (
                f'=IF(N{r}=0,"",IF(AND({semester_count}>0,{fallback_count}>0),"Misto",'
                f'IF({semester_count}>0,"Fechamento semestral","Agregado mensal (fallback)")))'
            )

            teacher_criteria = [
                (f"'AVALIAÇÃO DOCENTE'!$A$5:$A${self.teacher_end}", f'$C{r}'),
                (f"'AVALIAÇÃO DOCENTE'!$B$5:$B${self.teacher_end}", course),
                (f"'AVALIAÇÃO DOCENTE'!$C$5:$C${self.teacher_end}", discipline),
                (f"'AVALIAÇÃO DOCENTE'!$J$5:$J${self.teacher_end}", '"OK"'),
            ]
            ws.cell(r, 19).value = '=' + self._sumifs(f"'AVALIAÇÃO DOCENTE'!$E$5:$E${self.teacher_end}", teacher_criteria)
            ws.cell(r, 20).value = '=' + self._sumifs(f"'AVALIAÇÃO DOCENTE'!$G$5:$G${self.teacher_end}", teacher_criteria)
            ws.cell(r, 6).value = f'=IF(S{r}>0,T{r}/S{r},"")'

            result_criteria = [
                (f"'RESULTADOS ACADÊMICOS'!$A$5:$A${self.result_end}", f'$C{r}'),
                (f"'RESULTADOS ACADÊMICOS'!$B$5:$B${self.result_end}", course),
                (f"'RESULTADOS ACADÊMICOS'!$C$5:$C${self.result_end}", discipline),
                (f"'RESULTADOS ACADÊMICOS'!$P$5:$P${self.result_end}", '"OK"'),
            ]
            for column, source in ((21, "E"), (22, "F"), (23, "H"), (24, "I"), (25, "L"), (26, "M")):
                ws.cell(r, column).value = '=' + self._sumifs(
                    f"'RESULTADOS ACADÊMICOS'!${source}$5:${source}${self.result_end}", result_criteria
                )
            ws.cell(r, 8).value = f'=IF(U{r}>0,V{r}/U{r}*100,"")'
            ws.cell(r, 10).value = f'=IF(A{r}="","",V{r})'
            ws.cell(r, 11).value = f'=IF(A{r}="","",W{r})'
            ws.cell(r, 12).value = f'=IF(A{r}="","",X{r})'
            ws.cell(r, 13).value = f'=IF(Y{r}>0,Z{r}/Y{r},"")'

            ws.cell(r, 27).value = f'=IF(C{r}="","",VALUE(LEFT(C{r},4))*10+VALUE(RIGHT(C{r},1)))'
            ws.cell(r, 28).value = '=$K$5'
            for code, scope_col, key_col, visible_meta_col in (
                (self.kpi_codes[0], 29, 30, 5),
                (self.kpi_codes[1], 31, 32, 7),
                (self.kpi_codes[2], 33, 34, 9),
            ):
                scope_letter = get_column_letter(scope_col)
                key_letter = get_column_letter(key_col)
                ws.cell(r, scope_col).value = (
                    f'=IF(COUNTIFS(METAS!$A$5:$A${self.goal_end},"{code}",METAS!$B$5:$B${self.goal_end},$AB{r},'
                    f'METAS!$D$5:$D${self.goal_end},"<="&$AA{r})>0,$AB{r},"TOTAL")'
                )
                ws.cell(r, key_col).value = (
                    f'=IFERROR(AGGREGATE(14,6,METAS!$D$5:$D${self.goal_end}/'
                    f'((METAS!$A$5:$A${self.goal_end}="{code}")*(METAS!$B$5:$B${self.goal_end}=${scope_letter}{r})*'
                    f'(METAS!$D$5:$D${self.goal_end}<=$AA{r})),1),"")'
                )
                ws.cell(r, visible_meta_col).value = (
                    f'=IF(${key_letter}{r}="","",SUMIFS(METAS!$E$5:$E${self.goal_end},METAS!$A$5:$A${self.goal_end},"{code}",'
                    f'METAS!$B$5:$B${self.goal_end},${scope_letter}{r},METAS!$D$5:$D${self.goal_end},${key_letter}{r}))'
                )

            for c in range(1, 35):
                ws.cell(r, c).font = self.formula_font
            for c in (4, 5):
                ws.cell(r, c).number_format = "0.0"
            for c in (6, 7, 13):
                ws.cell(r, c).number_format = "0.00"
            for c in (8, 9):
                ws.cell(r, c).number_format = '0.0"%"'
            for c in (10, 11, 12, 14, 15, 16, 17):
                ws.cell(r, c).number_format = "0"

        # Dedicated semester-only NPS series. It is intentionally independent
        # from the monthly/semester toggle so the executive chart never repeats
        # the same semester six times or mixes monthly and semester sources.
        semester_headers = [
            "Semestre NPS", "NPS semestral", "Meta NPS semestral", "Respondentes NPS",
            "Promotores NPS", "Neutros NPS", "Detratores NPS", "Fonte NPS semestral",
            "Chave semestre", "Escopo meta NPS", "Chave meta NPS",
        ]
        for column, header in enumerate(semester_headers, 35):
            ws.cell(14, column).value = header
            ws.cell(14, column).font = self.header_font
            ws.cell(14, column).fill = PatternFill("solid", fgColor=GREEN_DARK)
            ws.cell(14, column).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        semester_min_key = 'PARAMETROS!$B$27'
        selected_semester_key = 'VALUE(LEFT(PARAMETROS!$B$19,4))*2+VALUE(RIGHT(PARAMETROS!$B$19,1))'
        comparison_semester_key = 'VALUE(LEFT(PARAMETROS!$B$25,4))*2+VALUE(RIGHT(PARAMETROS!$B$25,1))'
        semester_window = (
            f'IF(PARAMETROS!$B$9="Todo histórico",{selected_semester_key}-{semester_min_key}+1,'
            f'IF(PARAMETROS!$B$5="Mensal",ROUNDUP(IFERROR(VALUE(PARAMETROS!$B$9),12)/6,0),'
            f'IFERROR(VALUE(PARAMETROS!$B$9),4)))'
        )
        semester_start_key = (
            f'MAX({semester_min_key},MIN({comparison_semester_key},'
            f'{selected_semester_key}-({semester_window})+1))'
        )
        course = 'IF(PARAMETROS!$B$11="(todos)","*",PARAMETROS!$B$11)'
        desired_scope = 'IF(PARAMETROS!$B$11="(todos)","TOTAL",PARAMETROS!$B$11)'
        for r in range(SERIES_START, self.nps_series_end + 1):
            offset = r - SERIES_START
            ws.cell(r, 43).value = f'=IF(({semester_start_key})+{offset}>{selected_semester_key},"",({semester_start_key})+{offset})'
            ws.cell(r, 35).value = (
                f'=IF(AQ{r}="","",IFERROR(INDEX(DIM_PERIODO!$A$5:$A${self.dim_period_end},'
                f'MATCH(AQ{r},DIM_PERIODO!$F$5:$F${self.dim_period_end},0)),""))'
            )
            nps_criteria = [
                (f"'NPS SEMESTRAL'!$A$5:$A${self.nps_summary_end}", f'$AI{r}'),
                (f"'NPS SEMESTRAL'!$B$5:$B${self.nps_summary_end}", course),
                (f"'NPS SEMESTRAL'!$K$5:$K${self.nps_summary_end}", '"OK"'),
            ]
            for column, source in ((38, "E"), (39, "F"), (40, "G"), (41, "H")):
                ws.cell(r, column).value = '=' + self._sumifs(
                    f"'NPS SEMESTRAL'!${source}$5:${source}${self.nps_summary_end}", nps_criteria
                )
            ws.cell(r, 36).value = f'=IF(AL{r}>0,100*(AM{r}-AO{r})/AL{r},"")'
            semester_count = (
                f"COUNTIFS('NPS SEMESTRAL'!$A$5:$A${self.nps_summary_end},$AI{r},"
                f"'NPS SEMESTRAL'!$B$5:$B${self.nps_summary_end},{course},"
                f"'NPS SEMESTRAL'!$C$5:$C${self.nps_summary_end},\"Fechamento semestral\","
                f"'NPS SEMESTRAL'!$K$5:$K${self.nps_summary_end},\"OK\")"
            )
            fallback_count = semester_count.replace('"Fechamento semestral"', '"Agregado mensal (fallback)"')
            ws.cell(r, 42).value = (
                f'=IF(AL{r}=0,"",IF(AND({semester_count}>0,{fallback_count}>0),"Misto",'
                f'IF({semester_count}>0,"Fechamento semestral","Agregado mensal (fallback)")))'
            )
            ws.cell(r, 44).value = (
                f'=IF(COUNTIFS(METAS!$A$5:$A${self.goal_end},"{self.kpi_codes[0]}",'
                f'METAS!$B$5:$B${self.goal_end},{desired_scope},METAS!$D$5:$D${self.goal_end},"<="&$AQ{r})>0,'
                f'{desired_scope},"TOTAL")'
            )
            ws.cell(r, 45).value = (
                f'=IFERROR(AGGREGATE(14,6,METAS!$D$5:$D${self.goal_end}/'
                f'((METAS!$A$5:$A${self.goal_end}="{self.kpi_codes[0]}")*'
                f'(METAS!$B$5:$B${self.goal_end}=$AR{r})*(METAS!$D$5:$D${self.goal_end}<=$AQ{r})),1),"")'
            )
            ws.cell(r, 37).value = (
                f'=IF($AS{r}="","",SUMIFS(METAS!$E$5:$E${self.goal_end},'
                f'METAS!$A$5:$A${self.goal_end},"{self.kpi_codes[0]}",'
                f'METAS!$B$5:$B${self.goal_end},$AR{r},METAS!$D$5:$D${self.goal_end},$AS{r}))'
            )
            for c in range(35, 46):
                ws.cell(r, c).font = self.formula_font
            for c in (36, 37):
                ws.cell(r, c).number_format = "0.0"
            for c in (38, 39, 40, 41):
                ws.cell(r, c).number_format = "0"

        # The executive NPS card is also semester-native. In monthly mode its
        # comparison is the previous semester, not the prior month from the same
        # closing. This block deliberately overrides the generic row formulas.
        nps_ref_match = f'MATCH(PARAMETROS!$B$19,$AI${SERIES_START}:$AI${self.nps_series_end},0)'
        nps_cmp_match = f'MATCH(PARAMETROS!$B$25,$AI${SERIES_START}:$AI${self.nps_series_end},0)'
        ws["B9"] = f'=IFERROR(INDEX($AJ${SERIES_START}:$AJ${self.nps_series_end},{nps_ref_match}),"")'
        ws["C9"] = f'=IFERROR(INDEX($AJ${SERIES_START}:$AJ${self.nps_series_end},{nps_cmp_match}),"")'
        ws["D9"] = '=IF(OR(B9="",C9=""),"",B9-C9)'
        ws["E9"] = f'=IFERROR(INDEX($AK${SERIES_START}:$AK${self.nps_series_end},{nps_ref_match}),"")'
        nps_scope = f'INDEX($AR${SERIES_START}:$AR${self.nps_series_end},{nps_ref_match})'
        nps_goal_key = f'INDEX($AS${SERIES_START}:$AS${self.nps_series_end},{nps_ref_match})'
        ws["F9"] = (
            f'=IFERROR(SUMIFS(METAS!$F$5:$F${self.goal_end},METAS!$A$5:$A${self.goal_end},'
            f'"{self.kpi_codes[0]}",METAS!$B$5:$B${self.goal_end},{nps_scope},'
            f'METAS!$D$5:$D${self.goal_end},{nps_goal_key}),"")'
        )
        ws["G9"] = '=IF(B9="","Sem dados",IF(E9="","Sem meta",IF(B9>=E9,"Dentro da meta",IF(AND(F9<>"",B9>=F9),"Atenção","Fora da meta"))))'
        for target_col, source_col in ((14, "AL"), (15, "AM"), (16, "AN"), (17, "AO"), (18, "AP")):
            ws.cell(9, target_col).value = f'=IFERROR(INDEX(${source_col}${SERIES_START}:${source_col}${self.nps_series_end},{nps_ref_match}),"")'
        ws["L9"] = f'=IFERROR({nps_scope},"")'
        ws["M9"] = f'=IFERROR(INT({nps_goal_key}/10)&"-SEM"&MOD({nps_goal_key},10),"")'

        # Compact rolling datasets used exclusively by the charts. The full
        # series above remains intact for calculations and audit, while these
        # ranges keep axis labels readable even after many years of data.
        compact_headers = [
            "Período", "Chave", "Avaliação docente", "Meta",
            "Taxa de aprovação", "Meta", "Aprovados", "Reprovados por nota",
            "Reprovados por falta", "Média das notas", "Semestre", "NPS",
            "Meta", "Respondentes", "Promotores", "Neutros",
            "Detratores", "Fonte",
        ]
        for column, header in enumerate(compact_headers, 47):
            ws.cell(14, column).value = header
            ws.cell(14, column).font = self.header_font
            ws.cell(14, column).fill = PatternFill("solid", fgColor=GREEN_DARK)
            ws.cell(14, column).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        compact_start_key = f'MAX(PARAMETROS!$B$22,PARAMETROS!$B$23-{self.chart_period_points - 1})'
        nps_selected_key = 'VALUE(LEFT(PARAMETROS!$B$19,4))*2+VALUE(RIGHT(PARAMETROS!$B$19,1))'
        nps_compact_start = f'MAX(PARAMETROS!$B$27,({semester_start_key}),({nps_selected_key})-{self.chart_nps_points - 1})'
        compact_rows = max(self.chart_period_points, self.chart_nps_points)
        for r in range(SERIES_START, SERIES_START + compact_rows):
            offset = r - SERIES_START
            # AU:BD — selected visualization, always the most recent periods.
            if offset < self.chart_period_points:
                ws.cell(r, 48).value = f'=IF(({compact_start_key})+{offset}>PARAMETROS!$B$23,"",({compact_start_key})+{offset})'
                ws.cell(r, 47).value = (
                    f'=IF(AV{r}="","",IF($B$4="Mensal",'
                    f'IFERROR(INDEX(DIM_MES!$A$5:$A${self.dim_month_end},MATCH(AV{r},DIM_MES!$H$5:$H${self.dim_month_end},0)),""),'
                    f'IFERROR(INDEX(DIM_PERIODO!$A$5:$A${self.dim_period_end},MATCH(AV{r},DIM_PERIODO!$F$5:$F${self.dim_period_end},0)),"")))'
                )
                for target_col, source_col in ((49, "F"), (50, "G"), (51, "H"), (52, "I"),
                                                (53, "J"), (54, "K"), (55, "L"), (56, "M")):
                    ws.cell(r, target_col).value = (
                        f'=IF($AU{r}="","",IFERROR(INDEX(${source_col}${SERIES_START}:${source_col}${self.series_end},'
                        f'MATCH($AU{r},$A${SERIES_START}:$A${self.series_end},0)),""))'
                    )
            # BE:BL — eight most recent official NPS semesters.
            if offset < self.chart_nps_points:
                ws.cell(r, 57).value = (
                    f'=IF(({nps_compact_start})+{offset}>{nps_selected_key},"",'
                    f'IFERROR(INDEX(DIM_PERIODO!$A$5:$A${self.dim_period_end},'
                    f'MATCH(({nps_compact_start})+{offset},DIM_PERIODO!$F$5:$F${self.dim_period_end},0)),""))'
                )
                for target_col, source_col in ((58, "AJ"), (59, "AK"), (60, "AL"), (61, "AM"),
                                                (62, "AN"), (63, "AO"), (64, "AP")):
                    ws.cell(r, target_col).value = (
                        f'=IF($BE{r}="","",IFERROR(INDEX(${source_col}${SERIES_START}:${source_col}${self.nps_series_end},'
                        f'MATCH($BE{r},$AI${SERIES_START}:$AI${self.nps_series_end},0)),""))'
                    )
            for c in range(47, 65):
                ws.cell(r, c).font = self.formula_font
            for c in (49, 50, 56, 58, 59):
                ws.cell(r, c).number_format = "0.0"
            ws.cell(r, 51).number_format = '0.0"%"'
            ws.cell(r, 52).number_format = '0.0"%"'
            for c in (53, 54, 55, 60, 61, 62, 63):
                ws.cell(r, c).number_format = "0"

        widths = {
            "A": 16, "B": 12, "C": 16, "D": 13, "E": 13, "F": 15, "G": 15,
            "H": 14, "I": 14, "J": 12, "K": 12, "L": 12, "M": 12,
            "N": 15, "O": 13, "P": 13, "Q": 13, "R": 30,
        }
        for col, width in widths.items():
            ws.column_dimensions[col].width = width
        for c in range(19, 65):
            ws.column_dimensions[get_column_letter(c)].hidden = True
        ws.freeze_panes = "A15"

    def _course_nps_formula(self, course_cell: str, *, general: bool) -> str:
        """Weighted NPS from the official semester consolidation.

        The NPS SEMESTRAL sheet already contains exactly one official source per
        course and semester, so this formula cannot double count monthly rows.
        """
        criteria = [
            (f"'NPS SEMESTRAL'!$B$5:$B${self.nps_summary_end}", course_cell),
            (f"'NPS SEMESTRAL'!$K$5:$K${self.nps_summary_end}", '"OK"'),
        ]
        if not general:
            criteria.insert(0, (f"'NPS SEMESTRAL'!$A$5:$A${self.nps_summary_end}", 'PARAMETROS!$B$15'))
        resp = self._sumifs(f"'NPS SEMESTRAL'!$E$5:$E${self.nps_summary_end}", criteria)
        prom = self._sumifs(f"'NPS SEMESTRAL'!$F$5:$F${self.nps_summary_end}", criteria)
        det = self._sumifs(f"'NPS SEMESTRAL'!$H$5:$H${self.nps_summary_end}", criteria)
        return f'=IF({resp}>0,100*({prom}-{det})/{resp},"")'

    def _course_nps_count_formula(self, course_cell: str, *, general: bool) -> str:
        criteria = [
            (f"'NPS SEMESTRAL'!$B$5:$B${self.nps_summary_end}", course_cell),
            (f"'NPS SEMESTRAL'!$K$5:$K${self.nps_summary_end}", '"OK"'),
        ]
        if not general:
            criteria.insert(0, (f"'NPS SEMESTRAL'!$A$5:$A${self.nps_summary_end}", 'PARAMETROS!$B$15'))
        return '=' + self._sumifs(f"'NPS SEMESTRAL'!$E$5:$E${self.nps_summary_end}", criteria)

    def _course_nps_source_formula(self, course_cell: str, *, general: bool) -> str:
        extra = '' if general else f"'NPS SEMESTRAL'!$A$5:$A${self.nps_summary_end},PARAMETROS!$B$15,"
        base = (
            f"'NPS SEMESTRAL'!$B$5:$B${self.nps_summary_end},{course_cell},"
            f"'NPS SEMESTRAL'!$K$5:$K${self.nps_summary_end},\"OK\","
        )
        sem = f"COUNTIFS({extra}{base}'NPS SEMESTRAL'!$C$5:$C${self.nps_summary_end},\"Fechamento semestral\")"
        fallback = f"COUNTIFS({extra}{base}'NPS SEMESTRAL'!$C$5:$C${self.nps_summary_end},\"Agregado mensal (fallback)\")"
        count_formula = self._course_nps_count_formula(course_cell, general=general)[1:]
        return (
            f'=IF({count_formula}=0,"",IF(AND({sem}>0,{fallback}>0),"Misto",'
            f'IF({sem}>0,"Fechamento semestral","Agregado mensal (fallback)")))'
        )

    def _build_matrix(self):
        ws = self.ws["MATRIZ"]
        self._title(
            ws,
            "MATRIZ DE DESEMPENHO DOS CURSOS",
            "NPS calculado somente pelo consolidado oficial de cada semestre; alterna entre visão geral e o semestre escolhido em PARAMETROS.",
            end_col=16,
        )
        headers = [
            "Curso", "NPS", "Respondentes NPS", "Fonte NPS", "Avaliação docente",
            "Aprovação", "Média notas", "Aprovados", "Rep. nota", "Rep. falta",
            "Finalizados", "Qtd notas", "Período", "Leitura", "Indicador crítico", "Observação",
        ]
        self._headers(ws, 4, headers)
        for idx, name in enumerate(self.course_names, DATA_START):
            ws.cell(idx, 1).value = name
            ws.cell(idx, 1).font = self.static_font
            general_nps = self._course_nps_formula(f'$A{idx}', general=True)[1:]
            semester_nps = self._course_nps_formula(f'$A{idx}', general=False)[1:]
            general_count = self._course_nps_count_formula(f'$A{idx}', general=True)[1:]
            semester_count = self._course_nps_count_formula(f'$A{idx}', general=False)[1:]
            general_source = self._course_nps_source_formula(f'$A{idx}', general=True)[1:]
            semester_source = self._course_nps_source_formula(f'$A{idx}', general=False)[1:]
            ws.cell(idx, 2).value = f'=IF(PARAMETROS!$B$14="Geral",{general_nps},{semester_nps})'
            ws.cell(idx, 3).value = f'=IF(PARAMETROS!$B$14="Geral",{general_count},{semester_count})'
            ws.cell(idx, 4).value = f'=IF(PARAMETROS!$B$14="Geral",{general_source},{semester_source})'

            teacher_resp_all = (
                f"SUMIFS('AVALIAÇÃO DOCENTE'!$E$5:$E${self.teacher_end},"
                f"'AVALIAÇÃO DOCENTE'!$B$5:$B${self.teacher_end},$A{idx},"
                f"'AVALIAÇÃO DOCENTE'!$J$5:$J${self.teacher_end},\"OK\")"
            )
            teacher_sum_all = teacher_resp_all.replace("$E$5:$E$", "$G$5:$G$")
            teacher_resp_sem = (
                f"SUMIFS('AVALIAÇÃO DOCENTE'!$E$5:$E${self.teacher_end},"
                f"'AVALIAÇÃO DOCENTE'!$A$5:$A${self.teacher_end},PARAMETROS!$B$15,"
                f"'AVALIAÇÃO DOCENTE'!$B$5:$B${self.teacher_end},$A{idx},"
                f"'AVALIAÇÃO DOCENTE'!$J$5:$J${self.teacher_end},\"OK\")"
            )
            teacher_sum_sem = teacher_resp_sem.replace("$E$5:$E$", "$G$5:$G$")
            ws.cell(idx, 5).value = (
                f'=IF(PARAMETROS!$B$14="Geral",IF({teacher_resp_all}>0,{teacher_sum_all}/{teacher_resp_all},""),'
                f'IF({teacher_resp_sem}>0,{teacher_sum_sem}/{teacher_resp_sem},""))'
            )

            def result_sum(column: str, *, semester: bool) -> str:
                period_criteria = (
                    f"'RESULTADOS ACADÊMICOS'!$A$5:$A${self.result_end},PARAMETROS!$B$15,"
                    if semester else ""
                )
                return (
                    f"SUMIFS('RESULTADOS ACADÊMICOS'!${column}$5:${column}${self.result_end},"
                    f"{period_criteria}'RESULTADOS ACADÊMICOS'!$B$5:$B${self.result_end},$A{idx},"
                    f"'RESULTADOS ACADÊMICOS'!$P$5:$P${self.result_end},\"OK\")"
                )

            result_columns = {8: "F", 9: "H", 10: "I", 11: "E", 12: "L"}
            for target_column, source_column in result_columns.items():
                ws.cell(idx, target_column).value = (
                    f'=IF(PARAMETROS!$B$14="Geral",{result_sum(source_column, semester=False)},'
                    f'{result_sum(source_column, semester=True)})'
                )
            notes_all = result_sum("M", semester=False)
            notes_sem = result_sum("M", semester=True)
            ws.cell(idx, 6).value = f'=IF(K{idx}>0,H{idx}/K{idx}*100,"")'
            ws.cell(idx, 7).value = f'=IF(L{idx}>0,IF(PARAMETROS!$B$14="Geral",{notes_all},{notes_sem})/L{idx},"")'
            ws.cell(idx, 13).value = '=IF(PARAMETROS!$B$14="Geral","Todo histórico",PARAMETROS!$B$15)'
            ws.cell(idx, 14).value = (
                f'=IF(COUNTA(B{idx},E{idx},F{idx})=0,"Sem dados",'
                f'IF(AND(OR(B{idx}="",B{idx}>=50),OR(E{idx}="",E{idx}>=8),OR(F{idx}="",F{idx}>=80)),"Bom",'
                f'IF(AND(OR(B{idx}="",B{idx}>=0),OR(E{idx}="",E{idx}>=7),OR(F{idx}="",F{idx}>=70)),"Atenção","Crítico")))'
            )
            ws.cell(idx, 15).value = (
                f'=IF(COUNTA(B{idx},E{idx},F{idx})=0,"",IF(F{idx}=MIN(IF(B{idx}="",999,B{idx}),'
                f'IF(E{idx}="",999,E{idx}*10),IF(F{idx}="",999,F{idx})),"Aprovação",'
                f'IF(B{idx}=MIN(IF(B{idx}="",999,B{idx}),IF(E{idx}="",999,E{idx}*10),'
                f'IF(F{idx}="",999,F{idx})),"NPS","Avaliação")))'
            )
            ws.cell(idx, 16).value = ''
            for c in range(2, 16):
                ws.cell(idx, c).font = self.formula_font
            ws.cell(idx, 2).number_format = '0.0'
            ws.cell(idx, 3).number_format = '0'
            ws.cell(idx, 5).number_format = '0.00'
            ws.cell(idx, 6).number_format = '0.0"%"'
            ws.cell(idx, 7).number_format = '0.00'
            for c in (8, 9, 10, 11, 12):
                ws.cell(idx, c).number_format = '0'

        widths = [44, 12, 18, 30, 18, 14, 14, 12, 12, 12, 12, 12, 18, 16, 18, 34]
        for c, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = width
        last_row = max(DATA_START, DATA_START - 1 + len(self.course_names))
        ws.freeze_panes = 'A5'
        ws.auto_filter.ref = f"A4:P{last_row}"
        ws.conditional_formatting.add(
            f"N5:N{last_row}",
            FormulaRule(formula=['N5="Crítico"'], fill=PatternFill('solid', fgColor=RED_LIGHT)),
        )
        ws.conditional_formatting.add(
            f"N5:N{last_row}",
            FormulaRule(formula=['N5="Atenção"'], fill=PatternFill('solid', fgColor=ORANGE_LIGHT)),
        )
        ws.conditional_formatting.add(
            f"N5:N{last_row}",
            FormulaRule(formula=['N5="Bom"'], fill=PatternFill('solid', fgColor=GREEN_LIGHT)),
        )

    def _build_quality(self):
        ws = self.ws["QUALIDADE E GOVERNANÇA"]
        self._title(
            ws,
            "QUALIDADE E GOVERNANÇA",
            "Checagens de integridade, completude da exportação e coerência das seleções.",
            end_col=7,
        )
        self._headers(ws, 4, ["Teste", "Resultado", "Esperado", "Status", "Ação recomendada", "Fonte", "Observação"])
        checks = [
            ("Linhas NPS com erro", f'=COUNTIF(\'NPS DISCENTES\'!$L$5:$L${self.nps_end},"ERRO*")', 0, "Corrigir a composição de respondentes."),
            ("Consolidados NPS com erro", f'=COUNTIF(\'NPS SEMESTRAL\'!$K$5:$K${self.nps_summary_end},"ERRO*")', 0, "Revisar os fechamentos semestrais/fallbacks."),
            ("Linhas avaliação docente com erro", f'=COUNTIF(\'AVALIAÇÃO DOCENTE\'!$J$5:$J${self.teacher_end},"ERRO*")', 0, "Corrigir a base de avaliação docente."),
            ("Linhas resultados com erro", f'=COUNTIF(\'RESULTADOS ACADÊMICOS\'!$P$5:$P${self.result_end},"ERRO*")', 0, "Corrigir totais/finalizados/reprovações."),
            ("Metas com erro", f'=COUNTIF(METAS!$I$5:$I${self.goal_end},"ERRO*")', 0, "Corrigir vigência ou valores da meta."),
            ("Linhas NPS exportadas", f'=COUNTA(\'NPS DISCENTES\'!$A$5:$A${self.nps_end})', len(self.nps_rows), "A exportação deve conter toda a base, sem cortes."),
            ("Consolidados NPS exportados", f'=COUNTA(\'NPS SEMESTRAL\'!$A$5:$A${self.nps_summary_end})', len(self.nps_summary_rows), "A exportação deve conter todos os semestres/cursos."),
            ("Avaliações exportadas", f'=COUNTA(\'AVALIAÇÃO DOCENTE\'!$A$5:$A${self.teacher_end})', len(self.teacher_rows), "A exportação deve conter toda a base."),
            ("Resultados exportados", f'=COUNTA(\'RESULTADOS ACADÊMICOS\'!$A$5:$A${self.result_end})', len(self.result_rows), "A exportação deve conter toda a base agregada."),
            ("Disciplinas exportadas", f'=COUNTA(DISCIPLINAS!$B$5:$B${self.discipline_end})', len(self.discipline_rows), "A exportação deve conter todo o catálogo."),
            ("Referência compatível com visualização", '=IF(PARAMETROS!$B$5="Mensal",IF(LEN(PARAMETROS!$B$7)=7,1,0),IF(ISNUMBER(SEARCH("SEM",PARAMETROS!$B$7)),1,0))', 1, "Escolher uma referência compatível."),
            ("Comparação não posterior à referência", '=IF(PARAMETROS!$B$21<=PARAMETROS!$B$20,1,0)', 1, "A comparação deve ser anterior ou igual à referência."),
        ]
        for row, (label, formula, expected, action) in enumerate(checks, DATA_START):
            ws.cell(row, 1).value = label
            ws.cell(row, 2).value = formula
            ws.cell(row, 3).value = expected
            ws.cell(row, 4).value = f'=IF(B{row}=C{row},"OK","REVISAR")'
            ws.cell(row, 5).value = action
            ws.cell(row, 6).value = "Workbook gerado"
            ws.cell(row, 7).value = "Nenhuma base é truncada silenciosamente."
            for c in range(1, 8):
                ws.cell(row, c).font = self.formula_font if c in (2, 4) else self.static_font
        last = DATA_START - 1 + len(checks)
        ws.conditional_formatting.add(f"D5:D{last}", FormulaRule(formula=['D5="REVISAR"'], fill=PatternFill('solid', fgColor=RED_LIGHT)))
        ws.conditional_formatting.add(f"D5:D{last}", FormulaRule(formula=['D5="OK"'], fill=PatternFill('solid', fgColor=GREEN_LIGHT)))
        for col, width in {"A": 42, "B": 16, "C": 14, "D": 14, "E": 50, "F": 22, "G": 42}.items():
            ws.column_dimensions[col].width = width
        ws.freeze_panes = 'A5'
        ws.auto_filter.ref = f"A4:G{last}"

    @staticmethod
    def _add_value_labels(chart, *, series_index: int = 0):
        if len(chart.series) <= series_index:
            return
        labels = DataLabelList()
        labels.showVal = True
        labels.showLegendKey = False
        labels.showCatName = False
        labels.showSerName = False
        labels.showPercent = False
        chart.series[series_index].dLbls = labels

    def _build_panel(self):
        ws = self.ws["PAINEL"]
        self._title(
            ws,
            f"PAINEL DE GESTÃO — {self.code}",
            "Leitura executiva inspirada no painel institucional de referência: cartões objetivos, séries compactas e detalhamento semestral sem poluição visual.",
            end_col=12,
        )

        # Context cards follow the same rhythm as the reference workbook.
        context = [
            (1, 2, "Visualização", "=PARAMETROS!B5"),
            (3, 4, "Referência", "=PARAMETROS!B7"),
            (5, 6, "Comparação", "=PARAMETROS!B8"),
            (7, 8, "Janela", "=PARAMETROS!B9"),
            (9, 10, "Curso em foco", "=PARAMETROS!B11"),
            (11, 12, "Disciplina", "=PARAMETROS!B12"),
        ]
        for start_col, end_col, label, formula in context:
            self._paint_range(ws, 4, 5, start_col, end_col, fill=PALE_TEAL,
                              border=Border(bottom=Side(style="thin", color=BORDER)))
            ws.merge_cells(start_row=4, start_column=start_col, end_row=4, end_column=end_col)
            ws.merge_cells(start_row=5, start_column=start_col, end_row=5, end_column=end_col)
            ws.cell(4, start_col).value = label
            ws.cell(4, start_col).font = Font(name="Arial", size=8, bold=True, color=GRAY_TEXT)
            ws.cell(4, start_col).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(5, start_col).value = formula
            ws.cell(5, start_col).font = Font(name="Arial", size=11, bold=True, color=BLACK)
            ws.cell(5, start_col).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[4].height = 20
        ws.row_dimensions[5].height = 28

        # ------------------------------------------------------------------
        # NPS — executive cards, rolling semester chart, detailed table and
        # composition chart. Full history remains in NPS SEMESTRAL.
        # ------------------------------------------------------------------
        self._section(ws, 7, f"{self.prefix}-01 · NPS DISCENTE — FECHAMENTO SEMESTRAL OFICIAL", end_col=12)
        cards = [
            (1, 2, "NPS referência", "=CALC!B9", "0.0", PALE_TEAL, 18),
            (3, 4, "NPS anterior", "=CALC!C9", "0.0", PALE_GREEN, 16),
            (5, 6, "Variação", "=CALC!D9", '+0.0;-0.0;0.0', PALE_GREEN, 16),
            (7, 8, "Meta vigente", "=CALC!E9", "0.0", PALE_BLUE, 16),
            (9, 10, "Respondentes", "=CALC!N9", "#,##0", PALE_GOLD, 16),
            (11, 12, "Status", "=CALC!G9", "General", PALE_GREEN, 12),
        ]
        for start_col, end_col, label, formula, fmt, fill, size in cards:
            self._kpi_card(ws, start_col=start_col, end_col=end_col, label_row=8, value_row=9,
                           label=label, formula=formula, number_format=fmt, fill=fill, font_size=size)
        for formula, color in (( 'ISNUMBER(SEARCH("Dentro",$K$9))', GREEN_LIGHT),
                               ( 'ISNUMBER(SEARCH("Atenção",$K$9))', ORANGE_LIGHT),
                               ( 'ISNUMBER(SEARCH("Fora",$K$9))', RED_LIGHT)):
            ws.conditional_formatting.add("K9:L10", FormulaRule(formula=[formula], fill=PatternFill("solid", fgColor=color)))
        ws.merge_cells("A12:L12")
        ws["A12"] = '="Fonte utilizada: "&IF(CALC!R9="","sem dados",CALC!R9)&"   •   Escopo da meta: "&IF(CALC!L9="","sem meta específica",CALC!L9)&IF(CALC!M9="","","   •   Vigência: "&CALC!M9)'
        ws["A12"].font = Font(name="Arial", size=9, italic=True, color=GRAY_TEXT)
        ws["A12"].fill = PatternFill("solid", fgColor=PALE_GREEN)
        ws["A12"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[12].height = 24

        nps_chart = LineChart()
        self._chart_base(nps_chart, title=f"NPS — até {self.chart_nps_points} semestres mais recentes",
                         x_title="Semestre", y_title="NPS", width=33.5, height=8.8)
        nps_chart.y_axis.scaling.min = -100
        nps_chart.y_axis.scaling.max = 100
        nps_chart.add_data(Reference(self.ws["CALC"], min_col=58, max_col=59, min_row=14,
                                     max_row=self.chart_nps_end), titles_from_data=True)
        nps_chart.set_categories(Reference(self.ws["CALC"], min_col=57, min_row=SERIES_START,
                                           max_row=self.chart_nps_end))
        self._style_line_chart(nps_chart, number_format="0.0")
        ws.add_chart(nps_chart, "A14")

        detail_header_row = 31
        self._headers(ws, detail_header_row,
                      ["Semestre", "NPS", "Meta", "Respondentes", "Promotores", "Neutros", "Detratores", "Fonte"],
                      start_col=1)
        for offset in range(self.chart_nps_points):
            row = detail_header_row + 1 + offset
            source_row = SERIES_START + offset
            for target_col, source_col in enumerate(range(57, 65), 1):
                ws.cell(row, target_col).value = f'=CALC!{get_column_letter(source_col)}{source_row}'
                ws.cell(row, target_col).font = self.formula_font
                ws.cell(row, target_col).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                if offset % 2:
                    ws.cell(row, target_col).fill = PatternFill("solid", fgColor=PALE_GREEN)
            for c in (2, 3):
                ws.cell(row, c).number_format = "0.0"
            for c in (4, 5, 6, 7):
                ws.cell(row, c).number_format = "#,##0"
            ws.row_dimensions[row].height = 22

        # Current semester interpretation block to the right of the detailed table.
        self._section(ws, detail_header_row, "LEITURA DO SEMESTRE", start_col=9, end_col=12)
        interpretation = [
            (detail_header_row + 1, "Promotores", "=CALC!O9", "#,##0"),
            (detail_header_row + 3, "Neutros", "=CALC!P9", "#,##0"),
            (detail_header_row + 5, "Detratores", "=CALC!Q9", "#,##0"),
            (detail_header_row + 7, "Total", "=CALC!N9", "#,##0"),
        ]
        for row, label, formula, fmt in interpretation:
            self._paint_range(ws, row, row + 1, 9, 12, fill=PALE_TEAL,
                              border=Border(bottom=Side(style="thin", color=BORDER)))
            ws.merge_cells(start_row=row, start_column=9, end_row=row, end_column=10)
            ws.merge_cells(start_row=row, start_column=11, end_row=row + 1, end_column=12)
            ws.cell(row, 9).value = label
            ws.cell(row, 9).font = Font(name="Arial", size=9, bold=True, color=GRAY_TEXT)
            ws.cell(row, 9).alignment = Alignment(horizontal="left", vertical="center")
            ws.cell(row, 11).value = formula
            ws.cell(row, 11).font = Font(name="Arial", size=14, bold=True, color=GREEN_DARK)
            ws.cell(row, 11).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row, 11).number_format = fmt

        composition_row = detail_header_row + self.chart_nps_points + 3
        self._section(ws, composition_row, "COMPOSIÇÃO DAS RESPOSTAS — QUANTIDADES POR SEMESTRE", end_col=12)
        composition = BarChart()
        composition.type = "col"
        composition.grouping = "stacked"
        composition.overlap = 100
        self._chart_base(composition, title="Promotores, neutros e detratores",
                         x_title="Semestre", y_title="Respondentes", width=33.5, height=8.8)
        composition.add_data(Reference(self.ws["CALC"], min_col=61, max_col=63, min_row=14,
                                       max_row=self.chart_nps_end), titles_from_data=True)
        composition.set_categories(Reference(self.ws["CALC"], min_col=57, min_row=SERIES_START,
                                              max_row=self.chart_nps_end))
        self._style_stacked_chart(composition)
        ws.add_chart(composition, f"A{composition_row + 2}")

        # ------------------------------------------------------------------
        # Teacher evaluation — compact rolling chart with values on each point.
        # ------------------------------------------------------------------
        teacher_row = composition_row + 20
        self._section(ws, teacher_row, f"{self.prefix}-02 · AVALIAÇÃO DOCENTE PELO ALUNO", end_col=12)
        teacher_cards = [
            (1, 2, "Valor referência", "=CALC!B10", "0.00", PALE_TEAL, 18),
            (3, 4, "Comparação", "=CALC!C10", "0.00", PALE_GREEN, 16),
            (5, 6, "Variação", "=CALC!D10", '+0.00;-0.00;0.00', PALE_GREEN, 16),
            (7, 8, "Meta", "=CALC!E10", "0.00", PALE_BLUE, 16),
            (9, 10, "Atenção", "=CALC!F10", "0.00", PALE_GOLD, 16),
            (11, 12, "Status", "=CALC!G10", "General", PALE_GREEN, 12),
        ]
        for start_col, end_col, label, formula, fmt, fill, size in teacher_cards:
            self._kpi_card(ws, start_col=start_col, end_col=end_col, label_row=teacher_row + 1,
                           value_row=teacher_row + 2, label=label, formula=formula,
                           number_format=fmt, fill=fill, font_size=size)
        teacher_status = f"K{teacher_row + 2}"
        for formula, color in ((f'ISNUMBER(SEARCH("Dentro",${teacher_status}))', GREEN_LIGHT),
                               (f'ISNUMBER(SEARCH("Atenção",${teacher_status}))', ORANGE_LIGHT),
                               (f'ISNUMBER(SEARCH("Fora",${teacher_status}))', RED_LIGHT)):
            ws.conditional_formatting.add(f"K{teacher_row + 2}:L{teacher_row + 3}", FormulaRule(formula=[formula], fill=PatternFill("solid", fgColor=color)))
        teacher_chart = LineChart()
        self._chart_base(teacher_chart, title=f"Avaliação docente — até {self.chart_period_points} períodos mais recentes",
                         x_title="Período", y_title="Nota média", width=33.5, height=8.8)
        teacher_chart.y_axis.scaling.min = 0
        teacher_chart.y_axis.scaling.max = 10
        teacher_chart.add_data(Reference(self.ws["CALC"], min_col=49, max_col=50, min_row=14,
                                         max_row=self.chart_period_end), titles_from_data=True)
        teacher_chart.set_categories(Reference(self.ws["CALC"], min_col=47, min_row=SERIES_START,
                                                max_row=self.chart_period_end))
        self._style_line_chart(teacher_chart, number_format="0.00")
        ws.add_chart(teacher_chart, f"A{teacher_row + 6}")

        # ------------------------------------------------------------------
        # Approval — cards include the actual counts, not only the percentage.
        # ------------------------------------------------------------------
        approval_row = teacher_row + 25
        self._section(ws, approval_row, f"{self.prefix}-03 · APROVAÇÃO, REPROVAÇÃO E NOTAS", end_col=12)
        approval_cards = [
            (1, 2, "Aprovação", "=CALC!B11", '0.0"%"', PALE_TEAL, 18),
            (3, 4, "Comparação", "=CALC!C11", '0.0"%"', PALE_GREEN, 16),
            (5, 6, "Meta", "=CALC!E11", '0.0"%"', PALE_BLUE, 16),
            (7, 8, "Aprovados", "=CALC!H11", "#,##0", PALE_GOLD, 16),
            (9, 10, "Rep. nota / falta", '=TEXT(CALC!I11,"#,##0")&" / "&TEXT(CALC!J11,"#,##0")', "General", PALE_GOLD, 14),
            (11, 12, "Status", "=CALC!G11", "General", PALE_GREEN, 12),
        ]
        for start_col, end_col, label, formula, fmt, fill, size in approval_cards:
            self._kpi_card(ws, start_col=start_col, end_col=end_col, label_row=approval_row + 1,
                           value_row=approval_row + 2, label=label, formula=formula,
                           number_format=fmt, fill=fill, font_size=size)
        approval_status = f"K{approval_row + 2}"
        for formula, color in ((f'ISNUMBER(SEARCH("Dentro",${approval_status}))', GREEN_LIGHT),
                               (f'ISNUMBER(SEARCH("Atenção",${approval_status}))', ORANGE_LIGHT),
                               (f'ISNUMBER(SEARCH("Fora",${approval_status}))', RED_LIGHT)):
            ws.conditional_formatting.add(f"K{approval_row + 2}:L{approval_row + 3}", FormulaRule(formula=[formula], fill=PatternFill("solid", fgColor=color)))
        approval_chart = LineChart()
        self._chart_base(approval_chart, title=f"Taxa de aprovação — até {self.chart_period_points} períodos mais recentes",
                         x_title="Período", y_title="Aprovação (%)", width=33.5, height=8.8)
        approval_chart.y_axis.scaling.min = 0
        approval_chart.y_axis.scaling.max = 100
        approval_chart.add_data(Reference(self.ws["CALC"], min_col=51, max_col=52, min_row=14,
                                          max_row=self.chart_period_end), titles_from_data=True)
        approval_chart.set_categories(Reference(self.ws["CALC"], min_col=47, min_row=SERIES_START,
                                                 max_row=self.chart_period_end))
        self._style_line_chart(approval_chart, number_format='0.0"%"')
        ws.add_chart(approval_chart, f"A{approval_row + 6}")

        # ------------------------------------------------------------------
        # Course comparison — chart plus auditable table, following the base
        # workbook's pattern of graph first and detail immediately below.
        # ------------------------------------------------------------------
        course_row = approval_row + 25
        self._section(ws, course_row, "DESEMPENHO ENTRE CURSOS — NPS OFICIAL", end_col=12)
        course_last_row = max(DATA_START, DATA_START - 1 + len(self.course_names))
        course_chart = BarChart()
        course_chart.type = "bar"
        course_chart.grouping = "clustered"
        self._chart_base(course_chart, title="NPS por curso", x_title="NPS", y_title="Curso",
                         width=33.5, height=max(8.5, min(13.0, 6.0 + len(self.course_names) * 0.55)))
        course_chart.legend = None
        course_chart.x_axis.scaling.min = -100
        course_chart.x_axis.scaling.max = 100
        course_chart.add_data(Reference(self.ws["MATRIZ"], min_col=2, max_col=2, min_row=4,
                                        max_row=course_last_row), titles_from_data=True)
        course_chart.set_categories(Reference(self.ws["MATRIZ"], min_col=1, min_row=DATA_START,
                                               max_row=course_last_row))
        if course_chart.series:
            series = course_chart.series[0]
            series.graphicalProperties.solidFill = GREEN
            series.graphicalProperties.line.solidFill = GREEN_DARK
            labels = DataLabelList()
            labels.showVal = True
            labels.showLegendKey = False
            labels.dLblPos = "outEnd"
            labels.numFmt = "0.0"
            series.dLbls = labels
        ws.add_chart(course_chart, f"A{course_row + 2}")

        course_table_row = course_row + 22
        self._headers(ws, course_table_row,
                      ["Curso", "NPS", "Respondentes", "Fonte", "Avaliação", "Aprovação", "Média notas", "Leitura"],
                      start_col=1)
        for offset in range(max(1, len(self.course_names))):
            row = course_table_row + 1 + offset
            source_row = DATA_START + offset
            mapping = [(1, 1), (2, 2), (3, 3), (4, 4), (5, 5), (6, 6), (7, 7), (8, 14)]
            for target_col, source_col in mapping:
                ws.cell(row, target_col).value = f'=MATRIZ!{get_column_letter(source_col)}{source_row}'
                ws.cell(row, target_col).font = self.formula_font
                ws.cell(row, target_col).alignment = Alignment(horizontal="center" if target_col > 1 else "left",
                                                                vertical="center", wrap_text=True)
                if offset % 2:
                    ws.cell(row, target_col).fill = PatternFill("solid", fgColor=PALE_GREEN)
            ws.cell(row, 2).number_format = "0.0"
            ws.cell(row, 3).number_format = "#,##0"
            ws.cell(row, 5).number_format = "0.00"
            ws.cell(row, 6).number_format = '0.0"%"'
            ws.cell(row, 7).number_format = "0.00"
            ws.row_dimensions[row].height = 24

        final_row = course_table_row + 1 + max(1, len(self.course_names)) + 2
        ws.merge_cells(start_row=final_row, start_column=1, end_row=final_row, end_column=12)
        ws.cell(final_row, 1).value = (
            "Antes da reunião gerencial, confirme que QUALIDADE E GOVERNANÇA apresenta zero erros. "
            "O painel exibe janelas compactas; o histórico integral permanece nas bases e em NPS SEMESTRAL."
        )
        ws.cell(final_row, 1).fill = PatternFill("solid", fgColor=YELLOW)
        ws.cell(final_row, 1).font = Font(name="Arial", size=10, bold=True, color=BLACK)
        ws.cell(final_row, 1).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[final_row].height = 36

        widths = {
            "A": 23, "B": 13, "C": 16, "D": 14, "E": 15, "F": 14,
            "G": 15, "H": 18, "I": 16, "J": 16, "K": 16, "L": 24,
        }
        for col, width in widths.items():
            ws.column_dimensions[col].width = width
        ws.sheet_view.zoomScale = 85
        ws.freeze_panes = "A7"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_margins.left = 0.25
        ws.page_margins.right = 0.25
        ws.page_margins.top = 0.4
        ws.page_margins.bottom = 0.4
        ws.print_title_rows = "1:5"
        ws.print_area = f"A1:L{final_row}"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0

    def _build_readme(self):
        ws = self.ws["LEIA-ME"]
        self._title(
            ws,
            f"DATA UNIVC — PAINEL {self.code}",
            "Workbook reconstruído do zero, com NPS oficial semestral, faixas dinâmicas e painel visual inspirado no modelo institucional fornecido.",
            end_col=2,
        )
        self._section(ws, 4, "COMO USAR", end_col=2)
        lines = [
            "1. Use PARAMETROS para escolher referência, comparação, janela, curso e disciplina.",
            "2. O NPS é sempre interpretado por semestre. O fechamento semestral de cada curso prevalece sobre qualquer lançamento mensal.",
            "3. Quando um curso não possui fechamento semestral, os registros mensais daquele curso são agregados como fallback do semestre, sem dupla contagem.",
            "4. A aba NPS SEMESTRAL documenta, para cada curso e semestre, respondentes, composição, NPS e fonte oficial utilizada.",
            "5. O PAINEL usa séries compactas: oito semestres de NPS e doze períodos para avaliação/aprovação, sempre com rótulos de valor. O histórico integral permanece nas bases.",
            "6. Avaliação docente e resultados acadêmicos são semestrais. Na visão mensal, seus fechamentos são repetidos nos meses correspondentes apenas para navegação.",
            "7. RESULTADOS ACADÊMICOS é agregado por semestre + curso + disciplina; nomes e matrículas de alunos não são exportados.",
            "8. Todas as faixas, calendários, cursos e disciplinas são dimensionados pelo volume real. Nenhum registro é descartado silenciosamente.",
            "9. Se a base exceder o limite físico de 1.048.576 linhas do Excel, a exportação é interrompida com erro explícito em vez de gerar arquivo incompleto.",
            "10. METAS usa chave YYYY*10+SEM. SEM1 vale desde janeiro; SEM2 vale desde julho.",
            "11. CALC contém as fórmulas do painel e possui colunas auxiliares ocultas; não editar.",
            "12. O arquivo não usa macros, consultas externas, tabelas OOXML copiadas, nem views/drawings herdados.",
        ]
        for row, line in enumerate(lines, 6):
            ws.cell(row, 1).value = line
            ws.cell(row, 1).font = Font(name='Arial', size=10, color=BLACK)
            ws.cell(row, 1).alignment = Alignment(wrap_text=True, vertical='top')
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
            ws.row_dimensions[row].height = 34
        colors_row = 6 + len(lines) + 2
        self._section(ws, colors_row, "CORES", end_col=2)
        colors = [
            ("Amarelo", "Controle do usuário", YELLOW),
            ("Verde", "Dado vinculado/importado", GREEN_LIGHT),
            ("Azul-esverdeado", "KPI/volume em destaque", TEAL_LIGHT),
            ("Cinza", "Fórmula/resultado", GRAY),
            ("Vermelho claro", "Erro/revisão", RED_LIGHT),
        ]
        for row, (label, description, color) in enumerate(colors, colors_row + 2):
            ws.cell(row, 1).value = label
            ws.cell(row, 1).fill = PatternFill('solid', fgColor=color)
            ws.cell(row, 1).font = self.body_font
            ws.cell(row, 2).value = description
            ws.cell(row, 2).font = self.body_font
        ws.column_dimensions['A'].width = 86
        ws.column_dimensions['B'].width = 44

    def _defined_names(self):
        names = {
            'LST_VISUALIZACAO': "'LISTAS DE APOIO'!$A$5:$A$6",
            'LST_JANELA': "'LISTAS DE APOIO'!$B$5:$B$12",
            'LST_DESEMPENHO': "'LISTAS DE APOIO'!$C$5:$C$6",
            'LST_PERIODOS': f"'LISTAS DE APOIO'!$D$5:$D${self.period_list_end}",
            'LST_SEMESTRES': f"'LISTAS DE APOIO'!$E$5:$E${self.semester_list_end}",
            'LST_MESES': f"'LISTAS DE APOIO'!$F$5:$F${self.month_list_end}",
            'LST_STATUS': "'LISTAS DE APOIO'!$G$5:$G$9",
            'LST_CURSOS': f"'LISTAS DE APOIO'!$H$5:$H${self.course_list_end}",
            'LST_DISCIPLINAS': f"'LISTAS DE APOIO'!$I$5:$I${self.discipline_list_end}",
            'LST_CURSOS_BASE': f"'CURSOS'!$A$5:$A${self.course_end}",
            'LST_DISCIPLINAS_BASE': f"'DISCIPLINAS'!$B$5:$B${self.discipline_end}",
            'LST_INDICADORES': "'INDICADORES'!$A$5:$A$7",
            'BASE_NPS_SEMESTRAL': f"'NPS SEMESTRAL'!$A$5:$K${self.nps_summary_end}",
        }
        for name, reference in names.items():
            self.wb.defined_names.add(DefinedName(name, attr_text=reference))

    def _finish(self):
        # Blank reserve cells are editable, but source rows are never cut to fit
        # these ranges. All capacities were computed from the real collections.
        for sheet, max_row, input_cols in [
            ("NPS DISCENTES", self.nps_end, (1, 4, 5, 6, 7, 8, 10, 11)),
            ("AVALIAÇÃO DOCENTE", self.teacher_end, (1, 2, 3, 4, 5, 6, 8, 9)),
            ("RESULTADOS ACADÊMICOS", self.result_end, (1, 2, 3, 4, 5, 6, 8, 9, 10, 11, 12, 13)),
            ("METAS", self.goal_end, (1, 2, 3, 5, 6, 7, 8)),
        ]:
            ws = self.ws[sheet]
            for row in range(DATA_START, max_row + 1):
                for column in input_cols:
                    cell = ws.cell(row, column)
                    if cell.value is None:
                        cell.font = self.input_font

        # Subtle banding is applied only to rows that contain real exported
        # records. The editable reserve stays white, which makes it obvious
        # where the current base ends without imposing any artificial limit.
        zebra_blocks = [
            ("NPS DISCENTES", DATA_START, len(self.nps_rows), 12),
            ("NPS SEMESTRAL", DATA_START, len(self.nps_summary_rows), 11),
            ("AVALIAÇÃO DOCENTE", DATA_START, len(self.teacher_rows), 10),
            ("RESULTADOS ACADÊMICOS", DATA_START, len(self.result_rows), 16),
            ("METAS", DATA_START, len(self.goal_rows), 9),
            ("CURSOS", DATA_START, len(self.course_names), 4),
            ("DISCIPLINAS", DATA_START, len(self.discipline_rows), 5),
            ("PLANO_DE_ACAO", 8, len(self.action_rows), 14),
        ]
        for sheet_name, first_row, count, last_column in zebra_blocks:
            ws = self.ws[sheet_name]
            for offset in range(max(0, count)):
                row = first_row + offset
                ws.row_dimensions[row].height = max(ws.row_dimensions[row].height or 15, 20)
                if offset % 2:
                    for cell in ws.iter_rows(min_row=row, max_row=row, min_col=1, max_col=last_column):
                        for item in cell:
                            if item.fill.fill_type is None:
                                item.fill = PatternFill("solid", fgColor=PALE_GREEN)

        validation_columns = [
            ("NPS DISCENTES", 12, self.nps_end),
            ("NPS SEMESTRAL", 11, self.nps_summary_end),
            ("AVALIAÇÃO DOCENTE", 10, self.teacher_end),
            ("RESULTADOS ACADÊMICOS", 16, self.result_end),
            ("METAS", 9, self.goal_end),
        ]
        for sheet, column, max_row in validation_columns:
            ws = self.ws[sheet]
            letter = get_column_letter(column)
            ws.conditional_formatting.add(
                f"{letter}5:{letter}{max_row}",
                FormulaRule(formula=[f'LEFT({letter}5,4)="ERRO"'], fill=PatternFill('solid', fgColor=RED_LIGHT)),
            )
            ws.conditional_formatting.add(
                f"{letter}5:{letter}{max_row}",
                FormulaRule(formula=[f'{letter}5="OK"'], fill=PatternFill('solid', fgColor=GREEN_LIGHT)),
            )

        # Consistent navigation and tab language across the whole workbook.
        tab_colors = {
            "LEIA-ME": "4F9D8C", "PARAMETROS": "D9B44A", "PAINEL": GREEN_DARK,
            "QUALIDADE E GOVERNANÇA": "C97A40", "MATRIZ": GREEN,
            "PLANO_DE_ACAO": "7F8C8D", "NPS DISCENTES": "3E8E6C",
            "NPS SEMESTRAL": "2F6F57", "AVALIAÇÃO DOCENTE": "5D8FB4",
            "RESULTADOS ACADÊMICOS": "6A8E3A", "METAS": "A98B3A",
            "CURSOS": "7B9E87", "DISCIPLINAS": "7B9E87", "INDICADORES": "7B9E87",
            "CALC": "B8B8B8", "LISTAS DE APOIO": "B8B8B8",
            "DIM_PERIODO": "B8B8B8", "DIM_MES": "B8B8B8",
        }
        for sheet_name, color in tab_colors.items():
            if sheet_name in self.ws:
                self.ws[sheet_name].sheet_properties.tabColor = color
        self.wb.active = self.wb.sheetnames.index("PAINEL")
        try:
            self.wb.calculation.calcMode = "auto"
            self.wb.calculation.fullCalcOnLoad = True
            self.wb.calculation.forceFullCalc = True
        except Exception:
            LOGGER.warning("Não foi possível configurar o recálculo automático do workbook", exc_info=True)
        for ws in self.wb.worksheets:
            ws.sheet_view.showGridLines = False
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            if ws.title != "PAINEL":
                ws.sheet_view.zoomScale = 90 if ws.title not in {"CALC", "DIM_PERIODO", "DIM_MES", "LISTAS DE APOIO"} else 80

def data_from_repository(repo) -> dict[str, Any]:
    """Collect only the data required by the academic workbook."""
    courses=repo.list_courses()
    disciplines=repo.list_disciplines()
    dmap: dict[str,list[str]]={}
    for item in disciplines:
        dmap.setdefault(str(item.get('curso') or ''),[]).append(str(item.get('disciplina') or ''))
    return {
        'courses': courses,
        'disciplines': disciplines,
        'disciplines_map': dmap,
        'nps': list(reversed(repo.list_records('nps'))),
        'teacher': list(reversed(repo.list_records('avaliacao_docente'))),
        'results': list(reversed(repo.academic_result_summary())),
        'goals': repo.list_goals(),
        'actions': repo.list_actions(),
    }


def build_academic_workbook_bytes(repo, *, granularity: str | None = None, reference: str | None = None,
                                  comparison: str | None = None, course: str | None = None,
                                  discipline: str | None = None, window_periods: int | str | None = None) -> BytesIO:
    context=AcademicExcelContext(
        directorate=repo.directorate_code,
        granularity=granularity or 'semestral',
        reference=reference,
        comparison=comparison,
        course=course,
        discipline=discipline,
        window_periods=window_periods,
    )
    builder=AcademicWorkbookBuilder(context=context,data=data_from_repository(repo))
    wb=builder.build()
    output=BytesIO(); wb.save(output); wb.close(); output.seek(0)
    return output


def build_standalone_academic_workbook(*, directorate: str, output_path: str, data: dict[str,Any] | None = None,
                                       granularity: str = 'semestral', reference: str | None = None,
                                       comparison: str | None = None, window_periods: int | str | None = None) -> str:
    data=data or {'courses':[], 'disciplines':[], 'disciplines_map':{}, 'nps':[], 'teacher':[], 'results':[], 'goals':[], 'actions':[]}
    context=AcademicExcelContext(directorate=directorate,granularity=granularity,reference=reference,comparison=comparison,window_periods=window_periods)
    wb=AcademicWorkbookBuilder(context=context,data=data).build(); wb.save(output_path); wb.close(); return output_path
