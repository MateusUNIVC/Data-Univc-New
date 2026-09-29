#!/usr/bin/env python3
from __future__ import annotations

import re
import sys
import tempfile
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import Workbook
from academic_excel_parser import inspecionar_relatorio, iterar_registros

with tempfile.TemporaryDirectory() as temp:
    path = Path(temp) / "sei_bad_dimension.xlsx"
    wb = Workbook(); ws = wb.active
    ws["A1"]="Ano/Semestre:"; ws["C1"]="2026/1"
    ws["A3"]="Unidade Ensino:"; ws["A4"]="Curso:"; ws["C4"]="Administração"
    ws["A5"]="Disciplina:"; ws["C5"]="Gestão"; ws["A6"]="Turma:"; ws["C6"]="ADM1"; ws["N6"]="1º"
    ws["A7"]="Matrícula"; ws["C7"]="Nome"; ws["K7"]="Média"; ws["R7"]="Situação"
    ws["A8"]="1"; ws["C8"]="Aluno"; ws["K8"]=8; ws["R8"]="Aprovado"; ws["A9"]="Qtd de alunos:"; ws["C9"]=1
    wb.save(path); wb.close()
    extracted=Path(temp)/"x"; extracted.mkdir()
    with zipfile.ZipFile(path) as z: z.extractall(extracted)
    sheet=extracted/"xl/worksheets/sheet1.xml"
    xml=sheet.read_text(encoding="utf-8")
    xml,n=re.subn(r'<dimension ref="[^"]+"\s*/>', '<dimension ref="A1:A1"/>', xml, count=1)
    assert n==1
    sheet.write_text(xml,encoding="utf-8")
    rebuilt=Path(temp)/"rebuilt.xlsx"
    with zipfile.ZipFile(rebuilt,"w",zipfile.ZIP_DEFLATED) as z:
        for f in extracted.rglob("*"):
            if f.is_file(): z.write(f,f.relative_to(extracted))
    rebuilt.replace(path)
    meta,warnings=inspecionar_relatorio(path)
    assert meta["total_registros_aluno_disciplina"]==1
    assert sum(1 for _ in iterar_registros(path))==1
    assert warnings==[]
print("OK v0.11.6.6: XLSX SEI com dimensão subestimada é lido em streaming.")
