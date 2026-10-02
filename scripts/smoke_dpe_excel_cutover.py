from __future__ import annotations

import argparse
import hashlib
import json
import os
import sys
import urllib.parse
import urllib.request
from io import BytesIO
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook
from release_info import APP_VERSION, SCHEMA_VERSION

REQUIRED_SHEETS = ["LEIA-ME", "PARAMETROS", "PAINEL", "QUALIDADE E GOVERNANÇA", "MATRIZ", "PLANO_DE_ACAO"]


def main() -> int:
    parser = argparse.ArgumentParser(description="Smoke test do Excel Oficial da DPE apos o cutover.")
    parser.add_argument("--base-url", required=True)
    parser.add_argument("--cookie", default=os.getenv("DATA_UNIVC_DPE_COOKIE", ""), help="Cookie de sessão; padrão: DATA_UNIVC_DPE_COOKIE")
    parser.add_argument("--period-id", default="")
    parser.add_argument("--reference", default="")
    parser.add_argument("--output", default="")
    args = parser.parse_args()
    if not args.cookie:
        print("ERROR: informe --cookie ou DATA_UNIVC_DPE_COOKIE", file=sys.stderr)
        return 2

    query = {"diretoria": "DPE"}
    if args.period_id:
        query["period_id"] = args.period_id
    if args.reference:
        query["referencia"] = args.reference
    url = args.base_url.rstrip("/") + "/api/dpe/excel?" + urllib.parse.urlencode(query)
    request = urllib.request.Request(url, headers={"Cookie": args.cookie, "Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"})
    try:
        with urllib.request.urlopen(request, timeout=180) as response:
            raw = response.read()
            engine = response.headers.get("X-Data-UNIVC-Excel-Engine", "")
            content_type = response.headers.get("Content-Type", "")
            status = response.status
    except Exception as exc:
        print(f"ERROR: {type(exc).__name__}: {exc}", file=sys.stderr)
        return 2

    issues: list[str] = []
    if status != 200:
        issues.append(f"HTTP {status}")
    if engine != "excel_official":
        issues.append(f"engine={engine or 'blank'}")
    if "spreadsheetml.sheet" not in content_type:
        issues.append(f"content_type={content_type}")

    workbook_info = {"sheets": [], "external_links": 0, "vba": False, "directorate": None, "system_version": None, "schema_version": None}
    try:
        wb = load_workbook(BytesIO(raw), read_only=False, data_only=False, keep_links=True)
        workbook_info["sheets"] = wb.sheetnames
        workbook_info["external_links"] = len(getattr(wb, "_external_links", []))
        workbook_info["vba"] = wb.vba_archive is not None
        props = getattr(wb, "custom_doc_props", None)
        if props is not None:
            values = {item.name: item.value for item in props}
            workbook_info["directorate"] = values.get("DataUNIVC.DirectorateCode")
            workbook_info["system_version"] = values.get("DataUNIVC.SystemVersion")
            workbook_info["schema_version"] = str(values.get("DataUNIVC.SchemaVersion") or "")
        wb.close()
    except Exception as exc:
        issues.append(f"xlsx={type(exc).__name__}: {exc}")

    if workbook_info["sheets"][:6] != REQUIRED_SHEETS:
        issues.append("institutional_sheets")
    if workbook_info["external_links"]:
        issues.append("external_links")
    if workbook_info["vba"]:
        issues.append("vba")
    if workbook_info["directorate"] not in (None, "DPE"):
        issues.append(f"directorate={workbook_info['directorate']}")
    if workbook_info["system_version"] not in (None, APP_VERSION):
        issues.append(f"app_version={workbook_info['system_version']}")
    if workbook_info["schema_version"] not in ("", str(SCHEMA_VERSION)):
        issues.append(f"schema_version={workbook_info['schema_version']}")

    result = {
        "status": "PASS" if not issues else "FAIL",
        "url": url,
        "engine": engine,
        "bytes": len(raw),
        "sha256": hashlib.sha256(raw).hexdigest(),
        "app_version": APP_VERSION,
        "schema_version": str(SCHEMA_VERSION),
        "workbook": workbook_info,
        "issues": issues,
    }
    rendered = json.dumps(result, ensure_ascii=False, indent=2)
    if args.output:
        Path(args.output).write_text(rendered + "\n", encoding="utf-8")
    print(rendered)
    return 0 if not issues else 2


if __name__ == "__main__":
    raise SystemExit(main())
