from __future__ import annotations

import sys
import unittest
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from academic_excel_v2_builder import build_academic_report_payload, build_academic_report_workbook
from academic_excel_v3_builder import build_academic_interactive_payload, build_academic_interactive_workbook
from tests.test_faculty_student_v0115 import FacultyStudentV0115Tests


def main() -> int:
    suite = unittest.defaultTestLoader.discover(str(ROOT / "tests"), pattern="test_faculty_student_v011*.py")
    result = unittest.TextTestRunner(verbosity=1).run(suite)
    if not result.wasSuccessful():
        return 1

    case = FacultyStudentV0115Tests()
    db, repo, survey, _ = case._repos()
    try:
        case._import_official(survey)
        repo.create_goal({
            "indicador": "DCS-02", "recorte": "TOTAL", "vigencia": "2026-SEM1",
            "meta": 85, "atencao": 75, "limite_superior": 100,
            "justificativa": "Verificação v0.11.5",
        })

        payload_v2 = build_academic_report_payload(repo, reference="2026-SEM1", window_periods="all")
        if payload_v2["teacher"][0]["favorabilidade"] != 80.0:
            raise AssertionError("Favorabilidade V2 divergente.")
        wb_v2 = load_workbook(BytesIO(build_academic_report_workbook(payload_v2).getvalue()), data_only=False)
        if wb_v2["Avaliação Docente"]["K5"].value != 0.8:
            raise AssertionError("Excel V2 não materializou 80% como percentual.")

        payload_v3 = build_academic_interactive_payload(repo, reference="2026-SEM1", window_periods="all")
        wb_v3 = load_workbook(BytesIO(build_academic_interactive_workbook(payload_v3).getvalue()), data_only=False)
        formula = str(wb_v3["CALC"]["G5"].value or "")
        if "SUMIFS" not in formula or "*100" not in formula:
            raise AssertionError("Excel V3 não está calculando favorabilidade por contagens.")
        print("v0.11.5 verification OK: KPI 02 percentual + V2/V3 coerentes")
        return 0
    finally:
        db.close()


if __name__ == "__main__":
    raise SystemExit(main())
