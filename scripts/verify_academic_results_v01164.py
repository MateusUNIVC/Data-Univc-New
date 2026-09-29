from __future__ import annotations

import sys
import tempfile
import uuid
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import Workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import sessionmaker

from academic_excel_parser import inspecionar_relatorio
from database import Base
from models import AcademicResult, Directorate
from repository import ACADEMIC_RESULT_IMPORT_BATCH_SIZE, DatabaseRepository
from security import AuthorizationContext


def build_report(path: Path, rows: int) -> None:
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Ano/Semestre:"
    ws["C1"] = "2026/1"
    r = 3
    for label, value in (("Unidade Ensino:", None), ("Curso:", "Administração"), ("Disciplina:", "Gestão")):
        ws.cell(r, 1, label)
        if value is not None:
            ws.cell(r, 3, value)
        r += 1
    ws.cell(r, 1, "Turma:"); ws.cell(r, 3, "ADM1"); ws.cell(r, 14, "1º"); r += 1
    ws.cell(r, 1, "Matrícula"); ws.cell(r, 3, "Nome"); ws.cell(r, 11, "Média"); ws.cell(r, 18, "Situação"); r += 1
    for index in range(rows):
        ws.cell(r, 1, f"{index:07d}")
        ws.cell(r, 3, f"Aluno {index}")
        ws.cell(r, 11, 8.0)
        ws.cell(r, 18, "Aprovado")
        r += 1
    ws.cell(r, 1, "Qtd de alunos:"); ws.cell(r, 3, rows)
    wb.save(path)
    wb.close()


def main() -> int:
    rows = ACADEMIC_RESULT_IMPORT_BATCH_SIZE * 2 + 7
    with tempfile.TemporaryDirectory() as temp:
        path = Path(temp) / "resultado.xlsx"
        build_report(path, rows)
        metadata, warnings = inspecionar_relatorio(path)
        assert metadata["parser_mode"] == "read_only_streaming"
        assert metadata["total_registros_aluno_disciplina"] == rows
        assert not warnings

        engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(engine)
        Session = sessionmaker(bind=engine, expire_on_commit=False)
        db = Session()
        directorate = Directorate(code="DTNH", name="DTNH", active=True)
        db.add(directorate); db.flush(); db.commit()
        user = AuthorizationContext(
            user_id=str(uuid.uuid4()), email="verify@univc.edu.br", full_name="Verify",
            role="editor", directorate_id=directorate.id, directorate_code="DTNH", directorate_name="DTNH",
        )
        repo = DatabaseRepository(db, user)
        first = repo.import_sei_report_file(path, expected_course="Administração")
        second = repo.import_sei_report_file(path, expected_course="Administração")
        count = db.scalar(select(func.count(AcademicResult.id)))
        assert first["inseridos"] == rows
        assert first["lotes_processados"] == 3
        assert second["inseridos"] == 0
        assert second["ignorados"] == rows
        assert count == rows
        db.close(); engine.dispose()

    print(f"OK v0.11.6.4: {rows} registros, 3 lotes, reimportação idempotente.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
