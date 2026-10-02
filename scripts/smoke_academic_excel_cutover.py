from __future__ import annotations

import argparse
import hashlib
import json
import os
import urllib.error
import urllib.parse
import urllib.request
from io import BytesIO
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
from typing import Any

from openpyxl import load_workbook

from excel_official.constants import REQUIRED_INSTITUTIONAL_SHEETS
from release_info import APP_VERSION, SCHEMA_VERSION


EXPECTED_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _get(url: str, *, cookie: str | None = None, timeout: float = 60.0) -> tuple[int, Any, bytes]:
    headers = {"User-Agent": "Data-UNIVC-Academic-Excel-Cutover-Smoke/1"}
    if cookie:
        headers["Cookie"] = cookie
    request = urllib.request.Request(url, headers=headers, method="GET")
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            return response.status, response.headers, response.read()
    except urllib.error.HTTPError as exc:
        body = exc.read()
        raise RuntimeError(f"HTTP {exc.code} for {url}: {body[:500]!r}") from exc


def _custom_properties(workbook) -> dict[str, object]:
    return {item.name: item.value for item in workbook.custom_doc_props.props}


def main() -> int:
    parser = argparse.ArgumentParser(description="Authenticated smoke test for the Academic Excel Official cutover route.")
    parser.add_argument("--base-url", default="http://127.0.0.1:8000")
    parser.add_argument("--directorate", choices=("DTNH", "DCS"), required=True)
    parser.add_argument(
        "--cookie-env",
        default="DATA_UNIVC_SMOKE_COOKIE",
        help="Environment variable containing the full Cookie header for an authorized session. The value is never printed.",
    )
    parser.add_argument("--reference")
    parser.add_argument("--course")
    parser.add_argument("--discipline")
    parser.add_argument("--window", default="4")
    parser.add_argument("--output", type=Path)
    parser.add_argument("--report", type=Path, help="Write non-secret JSON evidence for finalization.")
    parser.add_argument("--timeout", type=float, default=90.0)
    args = parser.parse_args()

    cookie = os.getenv(args.cookie_env, "").strip()
    if not cookie:
        raise SystemExit(f"Missing authorized session cookie in environment variable {args.cookie_env}.")

    base = args.base_url.rstrip("/")
    status, _headers, health_body = _get(f"{base}/api/health/ready", timeout=args.timeout)
    if status != 200:
        raise SystemExit(f"Health readiness returned HTTP {status}.")
    health = json.loads(health_body.decode("utf-8"))
    if not health.get("ok"):
        raise SystemExit(f"Health readiness is not OK: {health!r}")
    if str(health.get("version")) != APP_VERSION:
        raise SystemExit(f"Server version {health.get('version')!r} does not match local release {APP_VERSION}.")

    query = {"janela": args.window}
    if args.reference:
        query["referencia"] = args.reference
    if args.course:
        query["curso"] = args.course
    if args.discipline:
        query["disciplina"] = args.discipline
    url = f"{base}/api/excel-interativo?{urllib.parse.urlencode(query)}"
    status, headers, raw = _get(url, cookie=cookie, timeout=args.timeout)
    if status != 200:
        raise SystemExit(f"Excel route returned HTTP {status}.")
    content_type = str(headers.get("Content-Type") or "").split(";", 1)[0].strip().lower()
    if content_type != EXPECTED_CONTENT_TYPE:
        raise SystemExit(f"Unexpected Excel content type: {content_type!r}")
    engine = str(headers.get("X-Data-UNIVC-Excel-Engine") or "").strip()
    if engine != "excel_official":
        raise SystemExit(f"Route is not serving Excel Official (engine header={engine!r}).")
    if not raw.startswith(b"PK"):
        raise SystemExit("Excel response is not an XLSX/ZIP payload.")

    wb = load_workbook(BytesIO(raw), data_only=False, read_only=False)
    try:
        missing = [name for name in REQUIRED_INSTITUTIONAL_SHEETS if name not in wb.sheetnames]
        if missing:
            raise SystemExit("Excel Official smoke workbook is missing institutional sheets: " + ", ".join(missing))
        props = _custom_properties(wb)
        if str(props.get("DataUNIVC.DirectorateCode") or "").upper() != args.directorate:
            raise SystemExit(
                f"Workbook directorate {props.get('DataUNIVC.DirectorateCode')!r} does not match expected {args.directorate}."
            )
        if str(props.get("DataUNIVC.SystemVersion") or "") != APP_VERSION:
            raise SystemExit("Workbook system version does not match the deployed application.")
        if str(props.get("DataUNIVC.SchemaVersion") or "") != str(SCHEMA_VERSION):
            raise SystemExit("Workbook schema version does not match the deployed application.")
        external_links = len(getattr(wb, "_external_links", None) or ())
        has_vba = getattr(wb, "vba_archive", None) is not None
        if external_links:
            raise SystemExit("Smoke workbook contains external links.")
        if has_vba:
            raise SystemExit("Smoke workbook contains VBA/macros.")
        sheetnames = list(wb.sheetnames)
    finally:
        wb.close()

    output_path = None
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_bytes(raw)
        output_path = str(args.output.resolve())
    evidence = {
        "report_version": 1,
        "status": "PASS",
        "directorate": args.directorate,
        "engine": engine,
        "app_version": APP_VERSION,
        "schema_version": str(SCHEMA_VERSION),
        "bytes": len(raw),
        "sha256": hashlib.sha256(raw).hexdigest(),
        "sheets": sheetnames,
        "external_links": external_links,
        "has_vba": has_vba,
        "workbook_path": output_path,
    }
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(evidence, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(
        f"PASS {args.directorate}: engine=excel_official, bytes={len(raw)}, "
        f"version={APP_VERSION}, schema={SCHEMA_VERSION}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
