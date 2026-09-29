from __future__ import annotations

import re
import unicodedata
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


class DPECostExcelError(ValueError):
    def __init__(self, message: str, errors: list[dict[str, Any]] | None = None):
        super().__init__(message)
        self.errors = errors or []


def _norm(value: Any) -> str:
    raw = unicodedata.normalize("NFKD", str(value or ""))
    raw = "".join(ch for ch in raw if not unicodedata.combining(ch))
    raw = raw.casefold().strip()
    raw = re.sub(r"[^a-z0-9]+", " ", raw)
    return " ".join(raw.split())


ALIASES = {
    "description": {"descricao", "despesa", "historico"},
    "amount": {"valor", "valor da despesa", "valor total"},
    "expense_date": {"data", "data da despesa", "competencia data"},
    "category": {"categoria", "categoria da despesa"},
    "cost_center": {"setor", "centro de custo", "origem setor"},
    "expense_kind": {"tipo", "tipo de despesa"},
    "expense_scope": {"tratamento", "tratamento da despesa", "destino economico", "destino econômico"},
    "counterparty_name": {"fornecedor", "beneficiario", "fornecedor beneficiario"},
    "document_number": {"documento", "numero documento", "nota fiscal", "nf"},
    "source_reference": {"referencia", "fonte referencia", "origem referencia"},
    "external_key": {"chave externa", "id externo", "identificador externo"},
    "notes": {"observacoes", "observacao", "notas"},
}
REQUIRED = ("description", "amount", "category")


def _header_map(ws) -> dict[str, int]:
    aliases = {alias: key for key, values in ALIASES.items() for alias in values}
    result: dict[str, int] = {}
    for cell in ws[1]:
        normalized = _norm(cell.value)
        key = aliases.get(normalized)
        if key and key not in result:
            result[key] = cell.column
    return result


def _cell_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return value


def parse_expense_workbook(path: str | Path) -> list[dict[str, Any]]:
    try:
        wb = load_workbook(path, data_only=False, read_only=True)
    except Exception as exc:
        raise DPECostExcelError("Nao foi possivel abrir o arquivo Excel.") from exc
    if not wb.sheetnames:
        raise DPECostExcelError("O arquivo nao possui abas.")
    ws = wb["Despesas"] if "Despesas" in wb.sheetnames else wb[wb.sheetnames[0]]
    headers = _header_map(ws)
    missing = [key for key in REQUIRED if key not in headers]
    if missing:
        labels = {"description": "Descricao", "amount": "Valor", "category": "Categoria"}
        raise DPECostExcelError(
            "O arquivo nao possui todas as colunas obrigatorias.",
            [{"row": 1, "field": labels.get(key, key), "error": "Cabecalho ausente."} for key in missing],
        )
    rows: list[dict[str, Any]] = []
    for row_number in range(2, ws.max_row + 1):
        values: dict[str, Any] = {}
        has_value = False
        for key, column in headers.items():
            value = ws.cell(row_number, column).value
            if isinstance(value, str) and value.startswith("="):
                raise DPECostExcelError(
                    "A importacao nao aceita formulas nos campos de origem.",
                    [{"row": row_number, "field": key, "error": "Substitua a formula pelo valor final."}],
                )
            value = _cell_value(value)
            values[key] = value
            if value not in (None, ""):
                has_value = True
        if not has_value:
            continue
        rows.append({"row_number": row_number, "raw_data": values})
    if not rows:
        raise DPECostExcelError("O arquivo nao possui linhas de despesas para importar.")
    return rows


def build_expense_import_template() -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Despesas"
    headers = [
        "Descricao", "Valor", "Data", "Categoria", "Tratamento", "Setor", "Tipo",
        "Fornecedor / beneficiario", "Documento", "Referencia", "Chave externa", "Observacoes",
    ]
    ws.append(headers)
    ws.append([
        "Exemplo: Software institucional", 12000, "2026-09-15", "SOFTWARE", "INSTITUCIONAL", "TI", "GENERAL",
        "Fornecedor exemplo", "NF-0001", "Contrato anual", "SOFT-2026-09", "Remova esta linha antes de importar",
    ])
    fill = PatternFill("solid", fgColor="1F6F43")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
    widths = [34, 16, 15, 24, 20, 24, 16, 30, 20, 28, 24, 38]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + index)].width = width
    ws.freeze_panes = "A2"
    notes = wb.create_sheet("LEIA-ME")
    notes["A1"] = "Importacao de despesas - DPE"
    notes["A1"].font = Font(bold=True, size=14)
    notes["A3"] = "Obrigatorio: Descricao, Valor e Categoria."
    notes["A4"] = "Categoria e Setor aceitam codigo ou nome cadastrado no DPE."
    notes["A5"] = "Tratamento aceita COMPARTILHADA ou INSTITUCIONAL. Se vazio, sera COMPARTILHADA."
    notes["A6"] = "Despesa DIRETA exige escolher um curso/contexto e, por seguranca, deve ser registrada pela tela da DPE."
    notes["A7"] = "Tipo aceita GENERAL ou PAYROLL. Se vazio, sera GENERAL."
    notes["A8"] = "A importacao sempre gera uma previa e nao grava despesas antes da confirmacao."
    notes["A9"] = "Linhas duplicadas ou invalidas bloqueiam a confirmacao do lote."
    notes.column_dimensions["A"].width = 110
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
