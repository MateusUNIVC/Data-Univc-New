from __future__ import annotations

import json
import math
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

from openpyxl import load_workbook

from academic_excel_official import (
    build_academic_excel_official_artifact_from_payload,
    build_academic_snapshot_context,
)
from academic_excel_v3_builder import (
    build_academic_interactive_payload,
    build_academic_interactive_workbook,
)
from excel_official.adapters.academic import ACADEMIC_METRICS, AcademicAdapter
from excel_official.audit.quality import evaluate_metric
from excel_official.context import AdapterInput


PARITY_REPORT_VERSION = 1


@dataclass(frozen=True, slots=True)
class AcademicParityCase:
    metric_code: str
    label: str
    period: str
    course: str = "(todos)"
    discipline: str = "(todas)"
    expected: float | None = None
    actual: float | None = None
    tolerance: float = 0.0
    status: str = "PASS"
    delta: float | None = None
    source: str = "backend_vs_core"
    note: str = ""


@dataclass(frozen=True, slots=True)
class WorkbookEvidence:
    engine: str
    build_ok: bool
    bytes_size: int = 0
    sheets: int = 0
    formulas: int = 0
    charts: int = 0
    error: str | None = None


@dataclass(frozen=True, slots=True)
class AcademicParityReport:
    directorate: str
    source_kind: str
    generated_at: str
    system_version: str
    schema_version: str
    payload_hash: str | None
    cases: tuple[AcademicParityCase, ...] = ()
    legacy_workbook: WorkbookEvidence | None = None
    new_workbook: WorkbookEvidence | None = None
    new_release_allowed: bool = False
    source_quality_ok: bool = True
    source_quality_notes: tuple[str, ...] = ()
    report_version: int = PARITY_REPORT_VERSION
    notes: tuple[str, ...] = ()

    @property
    def failures(self) -> tuple[AcademicParityCase, ...]:
        return tuple(item for item in self.cases if item.status == "FAIL")

    @property
    def warnings(self) -> tuple[AcademicParityCase, ...]:
        return tuple(item for item in self.cases if item.status == "WARN")

    @property
    def semantic_passed(self) -> bool:
        return not self.failures

    @property
    def cutover_status(self) -> str:
        builds_ok = bool(
            self.legacy_workbook
            and self.legacy_workbook.build_ok
            and self.new_workbook
            and self.new_workbook.build_ok
        )
        if not self.semantic_passed or not builds_ok or not self.new_release_allowed or not self.source_quality_ok:
            return "BLOCKED"
        if self.source_kind == "production":
            return "READY"
        return "CANDIDATE_PASS"

    def to_dict(self) -> dict[str, Any]:
        data = asdict(self)
        data.update(
            {
                "summary": {
                    "cases": len(self.cases),
                    "passed": sum(1 for item in self.cases if item.status == "PASS"),
                    "warnings": len(self.warnings),
                    "failures": len(self.failures),
                    "semantic_passed": self.semantic_passed,
                    "source_quality_ok": self.source_quality_ok,
                    "new_release_allowed": self.new_release_allowed,
                    "cutover_status": self.cutover_status,
                }
            }
        )
        return data

    def to_json(self) -> str:
        return json.dumps(self.to_dict(), ensure_ascii=False, indent=2, sort_keys=False)

    def to_markdown(self) -> str:
        lines = [
            f"# Academic Excel Parity Report — {self.directorate}",
            "",
            f"- Source kind: **{self.source_kind}**",
            f"- Generated at: `{self.generated_at}`",
            f"- System version: `{self.system_version}`",
            f"- Schema version: `{self.schema_version}`",
            f"- Payload hash: `{self.payload_hash or 'n/a'}`",
            f"- Cases: **{len(self.cases)}**",
            f"- Failures: **{len(self.failures)}**",
            f"- Warnings: **{len(self.warnings)}**",
            f"- Source quality: **{'PASS' if self.source_quality_ok else 'FAIL'}**",
            f"- New workbook release audit: **{'PASS' if self.new_release_allowed else 'FAIL'}**",
            f"- Cutover status: **{self.cutover_status}**",
            "",
            "## Workbook build evidence",
            "",
        ]
        for evidence in (self.legacy_workbook, self.new_workbook):
            if evidence is None:
                continue
            lines.append(
                f"- {evidence.engine}: {'PASS' if evidence.build_ok else 'FAIL'} · "
                f"{evidence.bytes_size} bytes · {evidence.sheets} sheets · "
                f"{evidence.formulas} formulas · {evidence.charts} charts"
                + (f" · {evidence.error}" if evidence.error else "")
            )
        if self.source_quality_notes:
            lines.extend(["", "## Source quality", ""])
            lines.extend(f"- {item}" for item in self.source_quality_notes)
        if self.failures or self.warnings:
            lines.extend(["", "## Non-passing cases", ""])
            for item in (*self.failures, *self.warnings):
                lines.append(
                    f"- **{item.status}** `{item.metric_code}` · {item.period} · "
                    f"{item.course} · {item.discipline}: expected={item.expected!r}, "
                    f"actual={item.actual!r}, delta={item.delta!r}, tolerance={item.tolerance}"
                    + (f" · {item.note}" if item.note else "")
                )
        if self.notes:
            lines.extend(["", "## Notes", ""])
            lines.extend(f"- {item}" for item in self.notes)
        lines.append("")
        return "\n".join(lines)


@dataclass(frozen=True, slots=True)
class AcademicCutoverReadiness:
    reports: tuple[AcademicParityReport, ...]
    required_directorates: tuple[str, ...] = ("DTNH", "DCS")

    @property
    def status(self) -> str:
        by_code = {item.directorate: item for item in self.reports}
        if any(code not in by_code for code in self.required_directorates):
            return "BLOCKED"
        if any(by_code[code].cutover_status == "BLOCKED" for code in self.required_directorates):
            return "BLOCKED"
        if all(by_code[code].cutover_status == "READY" for code in self.required_directorates):
            versions = {(by_code[code].system_version, by_code[code].schema_version) for code in self.required_directorates}
            if len(versions) == 1:
                return "READY"
            return "BLOCKED"
        return "CANDIDATE_PASS"

    def to_markdown(self) -> str:
        lines = [
            "# Academic Excel Cutover Readiness",
            "",
            f"Overall status: **{self.status}**",
            "",
            "| Diretoria | Origem | Casos | Falhas | Qualidade | Release | Status |",
            "|---|---|---:|---:|---|---|---|",
        ]
        by_code = {item.directorate: item for item in self.reports}
        for code in self.required_directorates:
            report = by_code.get(code)
            if report is None:
                lines.append(f"| {code} | — | 0 | — | — | — | MISSING |")
                continue
            lines.append(
                f"| {code} | {report.source_kind} | {len(report.cases)} | {len(report.failures)} | "
                f"{'PASS' if report.source_quality_ok else 'FAIL'} | "
                f"{'PASS' if report.new_release_allowed else 'FAIL'} | {report.cutover_status} |"
            )
        lines.extend(
            [
                "",
                "`READY` is only possible when both DTNH and DCS were audited from production snapshots, "
                "all semantic cases passed, source quality passed, both workbooks built, the new release audit passed, "
                "and both reports use the same application/schema version.",
                "",
            ]
        )
        return "\n".join(lines)


def _clean(value: Any) -> str:
    return str(value or "").strip()


def _rows(payload: Mapping[str, Any], key: str, *, period: str, course: str | None = None, discipline: str | None = None) -> list[dict[str, Any]]:
    output: list[dict[str, Any]] = []
    for raw in payload.get(key, ()):
        if _clean(raw.get("periodo")) != period:
            continue
        if course not in (None, "(todos)") and _clean(raw.get("curso")) != course:
            continue
        if discipline not in (None, "(todas)") and _clean(raw.get("disciplina")) != discipline:
            continue
        output.append(dict(raw))
    return output


def _legacy_nps(rows: Sequence[Mapping[str, Any]]) -> float | None:
    respondents = sum(int(item.get("respondentes") or 0) for item in rows)
    if respondents <= 0:
        return None
    promoters = sum(int(item.get("promotores") or 0) for item in rows)
    detractors = sum(int(item.get("detratores") or 0) for item in rows)
    return round((promoters - detractors) / respondents * 100, 2)


def _legacy_teacher(rows: Sequence[Mapping[str, Any]]) -> float | None:
    unmapped = sum(int(item.get("nao_mapeados") or 0) for item in rows)
    classified = sum(int(item.get("classificados") or 0) for item in rows)
    favorable = sum(int(item.get("favoraveis") or 0) for item in rows)
    if unmapped > 0 or classified <= 0:
        return None
    return round(favorable / classified, 4)


def _legacy_approval(rows: Sequence[Mapping[str, Any]]) -> float | None:
    finalized = sum(int(item.get("finalizados") or 0) for item in rows)
    if finalized <= 0:
        return None
    approved = sum(int(item.get("aprovados") or 0) for item in rows)
    return round(approved / finalized, 4)


def _legacy_average_grade(rows: Sequence[Mapping[str, Any]]) -> float | None:
    count = sum(int(item.get("notas_contagem") or 0) for item in rows)
    if count <= 0:
        return None
    total = 0.0
    for item in rows:
        row_count = int(item.get("notas_contagem") or 0)
        if item.get("soma_notas") not in (None, ""):
            total += float(item.get("soma_notas") or 0.0)
        elif item.get("media_notas") not in (None, ""):
            total += float(item.get("media_notas") or 0.0) * row_count
    return round(total / count, 2)


def _compare(expected: float | None, actual: float | None, tolerance: float) -> tuple[str, float | None]:
    if expected is None and actual is None:
        return "PASS", None
    if expected is None or actual is None:
        return "FAIL", None
    if not (math.isfinite(float(expected)) and math.isfinite(float(actual))):
        return ("PASS", 0.0) if expected == actual else ("FAIL", None)
    delta = float(actual) - float(expected)
    return ("PASS" if abs(delta) <= tolerance else "FAIL"), delta


def _workbook_evidence(engine: str, builder) -> WorkbookEvidence:
    try:
        built = builder()
        if hasattr(built, "to_bytes"):
            buffer = built.to_bytes()
        elif isinstance(built, BytesIO):
            buffer = built
        else:
            buffer = built
        if hasattr(buffer, "getvalue"):
            raw = buffer.getvalue()
        elif isinstance(buffer, bytes):
            raw = buffer
        else:
            raw = bytes(buffer)
        wb = load_workbook(BytesIO(raw), data_only=False, read_only=False)
        try:
            formulas = sum(
                1
                for ws in wb.worksheets
                for row in ws.iter_rows()
                for cell in row
                if isinstance(cell.value, str) and cell.value.startswith("=")
            )
            charts = sum(len(ws._charts) for ws in wb.worksheets)
            return WorkbookEvidence(
                engine=engine,
                build_ok=True,
                bytes_size=len(raw),
                sheets=len(wb.sheetnames),
                formulas=formulas,
                charts=charts,
            )
        finally:
            wb.close()
    except Exception as exc:  # readiness evidence must record rather than hide build failures
        return WorkbookEvidence(engine=engine, build_ok=False, error=f"{type(exc).__name__}: {exc}")


def _case(
    *,
    spec,
    metric_code: str,
    label: str,
    period: str,
    course: str,
    discipline: str,
    expected: float | None,
    tolerance: float,
    source: str = "backend_vs_core",
    note: str = "",
) -> AcademicParityCase:
    actual = evaluate_metric(spec, metric_code)
    status, delta = _compare(expected, actual, tolerance)
    return AcademicParityCase(
        metric_code=metric_code,
        label=label,
        period=period,
        course=course,
        discipline=discipline,
        expected=expected,
        actual=actual,
        tolerance=tolerance,
        status=status,
        delta=delta,
        source=source,
        note=note,
    )


def _source_quality(payload: Mapping[str, Any]) -> tuple[bool, tuple[str, ...]]:
    quality = payload.get("quality", {}) or {}
    notes: list[str] = []
    nps_bad = int(quality.get("nps_inconsistent_rows") or 0)
    result_bad = int(quality.get("result_inconsistent_rows") or 0)
    if nps_bad:
        notes.append(f"NPS consistency failures reported by current Academic payload: {nps_bad}")
    if result_bad:
        notes.append(f"Academic-result consistency failures reported by current Academic payload: {result_bad}")
    if not payload.get("semesters"):
        notes.append("No academic periods are available in the authorized snapshot.")
    if not payload.get("courses_catalog"):
        notes.append("No active courses are available in the authorized snapshot.")
    return not notes, tuple(notes)


def audit_academic_payload_parity(
    payload: Mapping[str, Any],
    *,
    directorate: str | None = None,
    source_kind: str = "fixture",
    generated_by: str = "parity-audit",
    build_workbooks: bool = True,
) -> AcademicParityReport:
    directorate_code = _clean(directorate or payload.get("directorate")).upper()
    if directorate_code not in {"DTNH", "DCS"}:
        raise ValueError("Academic parity audit supports only DTNH/DCS snapshots.")
    if source_kind not in {"fixture", "local", "production"}:
        raise ValueError("source_kind must be fixture, local or production.")

    context = build_academic_snapshot_context(
        payload,
        generated_by=generated_by,
        directorate=directorate_code,
        export_id=f"parity-{directorate_code.lower()}",
    )
    adapter = AcademicAdapter()
    spec_cache: dict[tuple[tuple[str, Any], ...], Any] = {}

    def spec_for(period: str, course: str = "(todos)", discipline: str = "(todas)"):
        state = {
            "reference_period": period,
            "course": course,
            "discipline": discipline,
        }
        key = tuple(sorted(state.items()))
        if key not in spec_cache:
            spec_cache[key] = adapter.build_spec(
                AdapterInput(
                    snapshot_context=context,
                    authorized_data=payload,
                    initial_state=state,
                )
            )
        return spec_cache[key]

    cases: list[AcademicParityCase] = []
    periods = tuple(_clean(item) for item in payload.get("semesters", ()) if _clean(item))
    courses = tuple(
        _clean(item.get("curso"))
        for item in payload.get("courses_catalog", ())
        if _clean(item.get("curso")) and item.get("ativo") is not False
    )
    discipline_pairs = tuple(
        (_clean(item.get("curso")), _clean(item.get("disciplina")))
        for item in payload.get("disciplines_catalog", ())
        if _clean(item.get("curso")) and _clean(item.get("disciplina")) and item.get("ativo") is not False
    )

    for period in periods:
        spec = spec_for(period)
        cases.append(_case(
            spec=spec,
            metric_code=ACADEMIC_METRICS["nps_institution"],
            label="01A · NPS Instituição · Alunos",
            period=period,
            course="(todos)",
            discipline="(todas)",
            expected=_legacy_nps(_rows(payload, "institution_students", period=period)),
            tolerance=0.011,
        ))
        cases.append(_case(
            spec=spec,
            metric_code=ACADEMIC_METRICS["nps_faculty"],
            label="01C · NPS Instituição · Docentes",
            period=period,
            course="(todos)",
            discipline="(todas)",
            expected=_legacy_nps(_rows(payload, "faculty_nps", period=period)),
            tolerance=0.011,
        ))

        # Natural course-level and rolled-up NPS parity.
        for course in courses:
            course_spec = spec_for(period, course)
            course_rows = _rows(payload, "course_nps", period=period, course=course)
            if course_rows:
                cases.append(_case(
                    spec=course_spec,
                    metric_code=ACADEMIC_METRICS["nps_course"],
                    label="01B · NPS do Curso",
                    period=period,
                    course=course,
                    discipline="(todas)",
                    expected=_legacy_nps(course_rows),
                    tolerance=0.011,
                ))
            inst_course_rows = _rows(payload, "institution_students_by_course", period=period, course=course)
            if inst_course_rows:
                cases.append(_case(
                    spec=course_spec,
                    metric_code=ACADEMIC_METRICS["nps_institution_course"],
                    label="01A · NPS Instituição · Alunos · por Curso",
                    period=period,
                    course=course,
                    discipline="(todas)",
                    expected=_legacy_nps(inst_course_rows),
                    tolerance=0.011,
                ))

        # Total and course rollups for teacher/results verify weighted recomposition.
        rollups: list[tuple[str, str]] = [("(todos)", "(todas)")]
        rollups.extend((course, "(todas)") for course in courses)
        rollups.extend((course, discipline) for course, discipline in discipline_pairs)
        seen_rollups: set[tuple[str, str]] = set()
        for course, discipline in rollups:
            if (course, discipline) in seen_rollups:
                continue
            seen_rollups.add((course, discipline))
            teacher_rows = _rows(payload, "teacher", period=period, course=course, discipline=discipline)
            result_rows = _rows(payload, "results", period=period, course=course, discipline=discipline)
            state_spec = spec_for(period, course, discipline)
            if teacher_rows:
                cases.append(_case(
                    spec=state_spec,
                    metric_code=ACADEMIC_METRICS["teacher"],
                    label="02 · Avaliação Docente",
                    period=period,
                    course=course,
                    discipline=discipline,
                    expected=_legacy_teacher(teacher_rows),
                    tolerance=0.00011,
                ))
            if result_rows:
                cases.append(_case(
                    spec=state_spec,
                    metric_code=ACADEMIC_METRICS["approval"],
                    label="03 · Aprovação",
                    period=period,
                    course=course,
                    discipline=discipline,
                    expected=_legacy_approval(result_rows),
                    tolerance=0.00011,
                ))
                cases.append(_case(
                    spec=state_spec,
                    metric_code=ACADEMIC_METRICS["average_grade"],
                    label="Média das Notas",
                    period=period,
                    course=course,
                    discipline=discipline,
                    expected=_legacy_average_grade(result_rows),
                    tolerance=0.011,
                ))

    # Cross-check values already reported by the current payload against their
    # exported sufficient statistics. This catches drift inside the legacy/backend
    # path itself instead of only comparing two recomputations of the same facts.
    reported_checks: list[tuple[str, str, str, str, str, float | None, float | None, float]] = []
    for raw in payload.get("institution_students", ()):
        period = _clean(raw.get("periodo"))
        if raw.get("valor") not in (None, ""):
            reported_checks.append((ACADEMIC_METRICS["nps_institution"], "01A · payload reported", period, "(todos)", "(todas)", float(raw["valor"]), _legacy_nps([raw]), 0.011))
    for raw in payload.get("institution_students_by_course", ()):
        period, course = _clean(raw.get("periodo")), _clean(raw.get("curso"))
        if raw.get("valor") not in (None, ""):
            reported_checks.append((ACADEMIC_METRICS["nps_institution_course"], "01A por curso · payload reported", period, course, "(todas)", float(raw["valor"]), _legacy_nps([raw]), 0.011))
    for raw in payload.get("course_nps", ()):
        period, course = _clean(raw.get("periodo")), _clean(raw.get("curso"))
        if raw.get("valor") not in (None, ""):
            reported_checks.append((ACADEMIC_METRICS["nps_course"], "01B · payload reported", period, course, "(todas)", float(raw["valor"]), _legacy_nps([raw]), 0.011))
    for raw in payload.get("faculty_nps", ()):
        period = _clean(raw.get("periodo"))
        if raw.get("valor") not in (None, ""):
            reported_checks.append((ACADEMIC_METRICS["nps_faculty"], "01C · payload reported", period, "(todos)", "(todas)", float(raw["valor"]), _legacy_nps([raw]), 0.011))
    for raw in payload.get("teacher", ()):
        period, course, discipline = _clean(raw.get("periodo")), _clean(raw.get("curso")), _clean(raw.get("disciplina"))
        reported = raw.get("favorabilidade")
        expected = None if reported in (None, "") else float(reported) / 100.0
        actual = _legacy_teacher([raw])
        status, delta = _compare(expected, actual, 0.00011)
        cases.append(AcademicParityCase(
            metric_code=ACADEMIC_METRICS["teacher"], label="02 · payload reported",
            period=period, course=course, discipline=discipline, expected=expected, actual=actual,
            tolerance=0.00011, status=status, delta=delta, source="payload_reported_vs_components",
        ))
    for raw in payload.get("results", ()):
        period, course, discipline = _clean(raw.get("periodo")), _clean(raw.get("curso")), _clean(raw.get("disciplina"))
        if raw.get("media_notas") not in (None, ""):
            expected = float(raw["media_notas"])
            actual = _legacy_average_grade([raw])
            status, delta = _compare(expected, actual, 0.011)
            cases.append(AcademicParityCase(
                metric_code=ACADEMIC_METRICS["average_grade"], label="Média · payload reported",
                period=period, course=course, discipline=discipline, expected=expected, actual=actual,
                tolerance=0.011, status=status, delta=delta, source="payload_reported_vs_components",
            ))
    for metric_code, label, period, course, discipline, expected, actual, tolerance in reported_checks:
        status, delta = _compare(expected, actual, tolerance)
        cases.append(AcademicParityCase(
            metric_code=metric_code, label=label, period=period, course=course, discipline=discipline,
            expected=expected, actual=actual, tolerance=tolerance, status=status, delta=delta,
            source="payload_reported_vs_components",
        ))

    source_quality_ok, source_quality_notes = _source_quality(payload)

    legacy_evidence: WorkbookEvidence | None = None
    new_evidence: WorkbookEvidence | None = None
    new_release_allowed = False
    if build_workbooks:
        legacy_evidence = _workbook_evidence(
            "academic_v3_legacy",
            lambda: build_academic_interactive_workbook(dict(payload)),
        )
        try:
            artifact = build_academic_excel_official_artifact_from_payload(
                payload,
                generated_by=generated_by,
                directorate=directorate_code,
                export_id=f"parity-build-{directorate_code.lower()}",
            )
            new_release_allowed = artifact.audit.release_allowed
            new_evidence = _workbook_evidence("excel_official_core", lambda: artifact)
        except Exception as exc:
            new_evidence = WorkbookEvidence(
                engine="excel_official_core",
                build_ok=False,
                error=f"{type(exc).__name__}: {exc}",
            )
    else:
        # Semantic-only mode is useful for large snapshot batches, but cannot declare cutover readiness.
        new_release_allowed = False

    notes: list[str] = []
    if source_kind != "production":
        notes.append("This report is not production evidence; a successful result is CANDIDATE_PASS, not READY.")
    if not build_workbooks:
        notes.append("Workbook build evidence was skipped; cutover status cannot be READY.")

    return AcademicParityReport(
        directorate=directorate_code,
        source_kind=source_kind,
        generated_at=datetime.now(timezone.utc).isoformat(),
        system_version=str(payload.get("app_version") or context.system_version),
        schema_version=str(context.schema_version),
        payload_hash=context.payload_hash,
        cases=tuple(cases),
        legacy_workbook=legacy_evidence,
        new_workbook=new_evidence,
        new_release_allowed=new_release_allowed,
        source_quality_ok=source_quality_ok,
        source_quality_notes=source_quality_notes,
        notes=tuple(notes),
    )


def audit_academic_repository_parity(
    repo,
    *,
    reference: str | None = None,
    comparison: str | None = None,
    course: str | None = None,
    discipline: str | None = None,
    window_periods: int | str | None = "all",
    build_workbooks: bool = True,
) -> AcademicParityReport:
    """Run production parity from an already-authorized request repository.

    Authorization stays in the caller/repository. The audit consumes the same
    full-history payload as the current Excel V3 and never bypasses scope.
    """
    payload = build_academic_interactive_payload(
        repo,
        reference=reference,
        comparison=comparison,
        course=course,
        discipline=discipline,
        window_periods=window_periods,
    )
    return audit_academic_payload_parity(
        payload,
        directorate=repo.directorate_code,
        source_kind="production",
        generated_by=getattr(repo.ctx, "email", None) or getattr(repo.ctx, "full_name", None) or "production-audit",
        build_workbooks=build_workbooks,
    )


def write_parity_report(report: AcademicParityReport, output_dir: str | Path) -> tuple[Path, Path]:
    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=True)
    stem = f"academic_parity_{report.directorate.lower()}"
    json_path = output / f"{stem}.json"
    md_path = output / f"{stem}.md"
    json_path.write_text(report.to_json() + "\n", encoding="utf-8")
    md_path.write_text(report.to_markdown(), encoding="utf-8")
    return json_path, md_path


def write_cutover_readiness(reports: Iterable[AcademicParityReport], output_path: str | Path) -> Path:
    readiness = AcademicCutoverReadiness(tuple(reports))
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(readiness.to_markdown(), encoding="utf-8")
    return path


__all__ = [
    "AcademicCutoverReadiness",
    "AcademicParityCase",
    "AcademicParityReport",
    "WorkbookEvidence",
    "audit_academic_payload_parity",
    "audit_academic_repository_parity",
    "write_cutover_readiness",
    "write_parity_report",
]
