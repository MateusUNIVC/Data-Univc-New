from __future__ import annotations

from collections import Counter
from datetime import date, datetime
from io import BytesIO
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import SeriesLabel
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from dm_analytics import build_dm_dashboard
from dm_catalog import DM_AREAS
from excel_errors import ExcelExportLimitError
from release_info import APP_VERSION

GREEN_DARK = "045233"
GREEN = "0B7A54"
GREEN_LIGHT = "E2F0D9"
PALE_GREEN = "F3F8F5"
PALE_TEAL = "E8F3F0"
PALE_BLUE = "EAF2FA"
PALE_GOLD = "FFF4D6"
PALE_RED = "FDE7E7"
GRAY = "F5F6F5"
SOFT_GRAY = "EEF1EF"
GRAY_TEXT = "5A5A5A"
BORDER = "D9E1DD"
WHITE = "FFFFFF"
BLACK = "000000"
LINKED_GREEN = "008000"
CHART_GREEN = "0B7A54"
CHART_TEAL = "398C78"
CHART_BLUE = "4679A6"
CHART_GOLD = "D9B44A"
CHART_RED = "C74B50"
CHART_GRAY = "87978E"
EXCEL_MAX_ROW = 1_048_576

INTEGER_FMT = '#,##0;[Red](#,##0);-'
DECIMAL_FMT = '0.0;[Red](0.0);-'
PERCENT_FMT = '0.0%;[Red](0.0%);-'
DATE_FMT = 'dd/mm/yyyy'
DATETIME_FMT = 'dd/mm/yyyy hh:mm'


def _to_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if not value:
        return None
    try:
        return date.fromisoformat(str(value)[:10])
    except ValueError:
        return None


def _to_datetime(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    if not value:
        return None
    text = str(value).strip()
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")[:25])
    except ValueError:
        try:
            return datetime.fromisoformat(text[:19])
        except ValueError:
            return None


def _status_fill(status: str | None) -> str:
    text = str(status or "")
    if text in {"Dentro da meta", "Encerrada", "Titulado", "Concluído", "Sucesso"}:
        return GREEN_LIGHT
    if text in {"Atenção", "Aberta", "Em formação", "Em acompanhamento", "Dados pendentes"}:
        return PALE_GOLD
    if text in {"Fora da meta", "Desligado", "Falha", "Erro"}:
        return PALE_RED
    if text in {"Em andamento", "Ativo"}:
        return PALE_BLUE
    return GRAY


def _semester_for_date(value: date) -> str:
    return f"{value.year:04d}-SEM{1 if value.month <= 6 else 2}"


def _effective_target(
    targets: list[dict[str, Any]], indicator_code: str, metric_key: str, as_of: date,
    *, default: float | None = None,
) -> float | None:
    reference = _semester_for_date(as_of)
    rows = [
        row for row in targets
        if str(row.get("indicator_code") or "").upper() == indicator_code
        and str(row.get("metric_key") or "") == metric_key
        and row.get("dimension_key") in (None, "", "TOTAL")
        and str(row.get("valid_from") or "") <= reference
        and (not row.get("valid_to") or str(row.get("valid_to")) >= reference)
    ]
    rows.sort(key=lambda row: str(row.get("valid_from") or ""), reverse=True)
    if not rows:
        return default
    value = rows[0].get("target")
    return float(value) if value is not None else default


class DMExcelV2Builder:
    """Gerencial DM workbook aligned with the current web dashboard contract.

    No technical CALC/list/matrix sheets are created. The workbook intentionally
    contains only management-facing tabs and the operational student trajectory.
    """

    SHEETS = (
        "Resumo Executivo",
        "DM-01 Evolução",
        "DM-02 Defesas",
        "Turmas",
        "Alunos e Defesas",
        "Metas",
        "Integração SEI",
        "Parâmetros",
    )

    def __init__(
        self,
        cohorts: Iterable[dict[str, Any]],
        students: Iterable[dict[str, Any]],
        *,
        targets: Iterable[dict[str, Any]] | None = None,
        sync_runs: Iterable[dict[str, Any]] | None = None,
        area_code: str | None = None,
        cohort_id: int | None = None,
        as_of: date | None = None,
    ) -> None:
        self.all_cohorts = [dict(row) for row in cohorts]
        self.all_students = [dict(row) for row in students]
        self.targets = [dict(row) for row in (targets or [])]
        self.sync_runs = [dict(row) for row in (sync_runs or [])]
        self.as_of = as_of or date.today()
        self.area_code = str(area_code or "").strip().upper() or None
        self.cohort_id = int(cohort_id) if cohort_id not in (None, "") else None

        self.dashboard = build_dm_dashboard(
            self.all_cohorts,
            self.all_students,
            area_code=self.area_code,
            cohort_id=self.cohort_id,
            as_of=self.as_of,
            targets=self.targets,
        )
        self.cohort_rows = [
            dict(row)
            for area in (self.dashboard.get("visible_areas") or [])
            for row in (area.get("cohorts") or [])
        ]
        cohort_ids = {int(row["id"]) for row in self.cohort_rows if row.get("id") is not None}
        self.student_rows = [
            dict(row) for row in self.all_students
            if row.get("cohort_id") is not None and int(row.get("cohort_id")) in cohort_ids
        ]
        self.cohort_source = {
            int(row["id"]): dict(row)
            for row in self.all_cohorts if row.get("id") is not None and int(row["id"]) in cohort_ids
        }
        if len(self.student_rows) + 20 >= EXCEL_MAX_ROW:
            raise ExcelExportLimitError(
                "A base de alunos do recorte excede o limite físico do Excel. A exportação foi cancelada sem truncamento."
            )
        if len(self.cohort_rows) + 20 >= EXCEL_MAX_ROW:
            raise ExcelExportLimitError(
                "A base de turmas do recorte excede o limite físico do Excel. A exportação foi cancelada sem truncamento."
            )

        self.dm01_target = _effective_target(self.targets, "DM-01", "cohort_members", self.as_of)
        self.dm02_target = _effective_target(self.targets, "DM-02", "average_months_to_defense", self.as_of, default=24)
        self.dm02_chart_data_start: int | None = None
        self.dm02_chart_data_end: int | None = None

        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.ws: dict[str, Any] = {}
        self.title_font = Font(name="Arial", size=16, bold=True, color=GREEN_DARK)
        self.subtitle_font = Font(name="Arial", size=9, color=GRAY_TEXT)
        self.section_font = Font(name="Arial", size=10, bold=True, color=WHITE)
        self.header_font = Font(name="Arial", size=9, bold=True, color=WHITE)
        self.body_font = Font(name="Arial", size=9, color=BLACK)
        self.linked_font = Font(name="Arial", size=9, color=LINKED_GREEN)
        self.static_font = Font(name="Arial", size=9, color=GRAY_TEXT)
        self.formula_font = Font(name="Arial", size=9, color=BLACK)
        self.border = Border(
            left=Side(style="thin", color=BORDER),
            right=Side(style="thin", color=BORDER),
            top=Side(style="thin", color=BORDER),
            bottom=Side(style="thin", color=BORDER),
        )

    # ---------- shared UI ----------
    def _new_sheet(self, name: str):
        ws = self.wb.create_sheet(name)
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 85
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_margins.left = 0.3
        ws.page_margins.right = 0.3
        ws.page_margins.top = 0.45
        ws.page_margins.bottom = 0.45
        self.ws[name] = ws
        return ws

    def _title(self, ws, title: str, subtitle: str, end_col: int) -> None:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
        ws.cell(1, 1).value = title
        ws.cell(1, 1).font = self.title_font
        ws.cell(1, 1).alignment = Alignment(vertical="center")
        ws.cell(1, 1).border = Border(bottom=Side(style="medium", color=GREEN))
        ws.row_dimensions[1].height = 28
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_col)
        ws.cell(2, 1).value = subtitle
        ws.cell(2, 1).font = self.subtitle_font
        ws.cell(2, 1).alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[2].height = 30

    def _section(self, ws, row: int, label: str, end_col: int) -> None:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=end_col)
        cell = ws.cell(row, 1)
        cell.value = label
        cell.font = self.section_font
        cell.fill = PatternFill("solid", fgColor=GREEN_DARK)
        cell.alignment = Alignment(vertical="center")
        ws.row_dimensions[row].height = 23

    def _headers(self, ws, row: int, labels: list[str]) -> None:
        for col, label in enumerate(labels, 1):
            cell = ws.cell(row, col)
            cell.value = label
            cell.font = self.header_font
            cell.fill = PatternFill("solid", fgColor=GREEN)
            cell.border = self.border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 30

    def _table(self, ws, *, start_row: int, end_row: int, end_col: int, name: str) -> None:
        if end_row <= start_row:
            return
        ref = f"A{start_row}:{get_column_letter(end_col)}{end_row}"
        table = Table(displayName=name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium4", showFirstColumn=False, showLastColumn=False,
            showRowStripes=True, showColumnStripes=False,
        )
        ws.add_table(table)

    def _set_widths(self, ws, widths: list[float]) -> None:
        for col, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(col)].width = width

    def _body_cell(self, cell, value: Any, *, linked: bool = True, wrap: bool = False) -> None:
        cell.value = value
        cell.font = self.linked_font if linked else self.body_font
        cell.alignment = Alignment(vertical="center", wrap_text=wrap)
        cell.border = Border(bottom=Side(style="hair", color=BORDER))

    def _status_cell(self, cell, status: str | None) -> None:
        cell.fill = PatternFill("solid", fgColor=_status_fill(status))
        cell.font = Font(name="Arial", size=9, bold=True, color=GREEN_DARK if status != "Fora da meta" else "9C1C24")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def _kpi_card(self, ws, *, c1: int, c2: int, row: int, label: str, value: Any, note: str = "", fmt: str = "General", fill: str = PALE_GREEN) -> None:
        for r in range(row, row + 4):
            for c in range(c1, c2 + 1):
                ws.cell(r, c).fill = PatternFill("solid", fgColor=fill)
                ws.cell(r, c).border = self.border
        ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
        ws.merge_cells(start_row=row + 1, start_column=c1, end_row=row + 2, end_column=c2)
        ws.merge_cells(start_row=row + 3, start_column=c1, end_row=row + 3, end_column=c2)
        ws.cell(row, c1).value = label
        ws.cell(row, c1).font = Font(name="Arial", size=8, bold=True, color=GRAY_TEXT)
        ws.cell(row, c1).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.cell(row + 1, c1).value = value
        ws.cell(row + 1, c1).number_format = fmt
        ws.cell(row + 1, c1).font = Font(name="Arial", size=15, bold=True, color=GREEN_DARK)
        ws.cell(row + 1, c1).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row + 3, c1).value = note
        ws.cell(row + 3, c1).font = Font(name="Arial", size=7.5, color=GRAY_TEXT)
        ws.cell(row + 3, c1).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def _chart_defaults(self, chart, title: str, y_title: str, *, width: float = 13.5, height: float = 7.2) -> None:
        chart.title = title
        chart.y_axis.title = y_title
        chart.width = width
        chart.height = height
        chart.legend.position = "b"
        chart.style = 13
        chart.display_blanks = "gap"
        labels = DataLabelList()
        labels.showVal = False
        chart.dLbls = labels

    @staticmethod
    def _style_series(chart, colors: list[str], *, line: bool = False) -> None:
        for series, color in zip(chart.series, colors):
            if line:
                series.graphicalProperties.line.solidFill = color
                series.graphicalProperties.line.width = 22000
                series.marker.symbol = "circle"
                series.marker.size = 6
                series.marker.graphicalProperties.solidFill = color
                series.marker.graphicalProperties.line.solidFill = color
            else:
                series.graphicalProperties.solidFill = color
                series.graphicalProperties.line.solidFill = color

    @staticmethod
    def _series_labels(chart, labels: list[str]) -> None:
        for series, label in zip(chart.series, labels):
            series.tx = SeriesLabel(v=label)

    def _scope_text(self) -> str:
        if self.cohort_id and self.cohort_rows:
            row = self.cohort_rows[0]
            return f"{row.get('area_name') or row.get('area_code')} · Turma {row.get('cohort_number')}"
        if self.area_code:
            return DM_AREAS.get(self.area_code, self.area_code)
        return "Todas as áreas e turmas"

    # ---------- build ----------
    def build(self) -> Workbook:
        for name in self.SHEETS:
            self._new_sheet(name)
        self._build_students()
        self._build_cohorts()
        self._build_dm01()
        self._build_dm02()
        self._build_targets()
        self._build_sei()
        self._build_parameters()
        self._build_summary()
        self._finish()
        return self.wb

    def _build_students(self) -> None:
        ws = self.ws["Alunos e Defesas"]
        self._title(
            ws,
            "ALUNOS E DEFESAS · DM",
            f"Trajetória individual do recorte: {self._scope_text()}. Defesa registrada confirma Titulado; Data de Titulação não faz parte do modelo ativo.",
            20,
        )
        ws.column_dimensions["T"].width = 3
        headers = [
            "Área", "Turma", "Matrícula", "Aluno", "Ingresso", "Qualificação", "Defesa", "Defesa marcada",
            "Status", "Saída", "Orientador", "Linha de pesquisa", "Meses até a defesa", "Qualidade do ingresso",
            "Origem", "Situação bruta no SEI", "Última consulta acadêmica SEI", "Última leitura SEI", "Observações",
        ]
        self._headers(ws, 4, headers)
        rows = sorted(self.student_rows, key=lambda x: (str(x.get("area_code") or ""), int(x.get("cohort_number") or 0), str(x.get("student_name") or "")))
        data_end = 4 + len(rows)
        for r, row in enumerate(rows, 5):
            quality = row.get("data_quality") or (
                "Ingresso pendente" if not row.get("entry_date") else ("Ingresso provisório" if row.get("entry_date_estimated") else "Ingresso confirmado")
            )
            values = [
                row.get("area_code"), f"Turma {row.get('cohort_number')}" if row.get("cohort_number") is not None else None,
                row.get("student_code"), row.get("student_name"), _to_date(row.get("entry_date")), _to_date(row.get("qualification_date")),
                _to_date(row.get("defense_date")), _to_date(row.get("defense_scheduled_date")), row.get("status"), _to_date(row.get("exit_date")),
                row.get("advisor"), row.get("research_line"), None, quality, row.get("source_system"), row.get("sei_raw_status"),
                _to_datetime(row.get("last_course_dates_sei_at")), _to_datetime(row.get("last_seen_sei_at")), row.get("notes"),
            ]
            for c, value in enumerate(values, 1):
                if c == 13:
                    continue
                self._body_cell(ws.cell(r, c), value, linked=True, wrap=c in {4, 11, 12, 16, 19})
            # Only confirmed individual entry dates are used by DM-02.
            ws.cell(r, 13).value = (
                f'=IF(OR(E{r}="",G{r}="",N{r}="Ingresso provisório",N{r}="Ingresso pendente"),"",'
                f'ROUND((YEAR(G{r})-YEAR(E{r}))*12+(MONTH(G{r})-MONTH(E{r}))+(DAY(G{r})-DAY(E{r}))/30.4375,1))'
            )
            ws.cell(r, 13).font = self.formula_font
            ws.cell(r, 13).number_format = DECIMAL_FMT
            ws.cell(r, 13).border = Border(bottom=Side(style="hair", color=BORDER))
            for c in (5, 6, 7, 8, 10):
                ws.cell(r, c).number_format = DATE_FMT
            for c in (17, 18):
                ws.cell(r, c).number_format = DATETIME_FMT
            self._status_cell(ws.cell(r, 9), row.get("status"))
        self._set_widths(ws, [10, 13, 16, 34, 14, 15, 14, 17, 14, 14, 25, 28, 18, 24, 16, 23, 22, 20, 40])
        ws.freeze_panes = "A5"
        if rows:
            self._table(ws, start_row=4, end_row=data_end, end_col=len(headers), name="tblDmAlunosDefesas")
        else:
            ws["A5"] = "Nenhum aluno no recorte atual."
            ws["A5"].font = self.static_font

    def _build_cohorts(self) -> None:
        ws = self.ws["Turmas"]
        self._title(ws, "TURMAS · DM", f"Cadastro e situação das turmas no recorte: {self._scope_text()}.", 14)
        ws.column_dimensions["N"].width = 3
        headers = [
            "Área", "Turma", "Abertura", "Vagas autorizadas", "Status", "Total", "Ativos", "Titulados", "Desligados",
            "Ocupação", "Origem", "Última leitura SEI", "Observações",
        ]
        self._headers(ws, 4, headers)
        rows = sorted(self.cohort_rows, key=lambda x: (str(x.get("area_code") or ""), int(x.get("cohort_number") or 0)))
        student_end = max(5, 4 + len(self.student_rows))
        for r, summary in enumerate(rows, 5):
            source = self.cohort_source.get(int(summary.get("id") or 0), {})
            base = [
                summary.get("area_code"), f"Turma {summary.get('cohort_number')}", _to_date(summary.get("opening_date")),
                summary.get("vacancies_authorized"), summary.get("status"), None, None, None, None, None,
                source.get("source_system"), _to_datetime(source.get("last_seen_sei_at")), source.get("notes"),
            ]
            for c, value in enumerate(base, 1):
                if c in {6, 7, 8, 9, 10}:
                    continue
                self._body_cell(ws.cell(r, c), value, linked=True, wrap=c in {13})
            for c, status in [(5, summary.get("status"))]:
                self._status_cell(ws.cell(r, c), status)
            # Derived directly from the visible individual base.
            scope_a = f"'Alunos e Defesas'!$A$5:$A${student_end}"
            scope_b = f"'Alunos e Defesas'!$B$5:$B${student_end}"
            scope_i = f"'Alunos e Defesas'!$I$5:$I${student_end}"
            ws.cell(r, 6).value = f'=COUNTIFS({scope_a},A{r},{scope_b},B{r})'
            ws.cell(r, 7).value = f'=COUNTIFS({scope_a},A{r},{scope_b},B{r},{scope_i},"Ativo")'
            ws.cell(r, 8).value = f'=COUNTIFS({scope_a},A{r},{scope_b},B{r},{scope_i},"Titulado")'
            ws.cell(r, 9).value = f'=COUNTIFS({scope_a},A{r},{scope_b},B{r},{scope_i},"Desligado")'
            ws.cell(r, 10).value = f'=IF(D{r}>0,F{r}/D{r},"")'
            for c in (6, 7, 8, 9):
                ws.cell(r, c).font = self.formula_font
                ws.cell(r, c).number_format = INTEGER_FMT
            ws.cell(r, 10).font = self.formula_font
            ws.cell(r, 10).number_format = PERCENT_FMT
            ws.cell(r, 3).number_format = DATE_FMT
            ws.cell(r, 12).number_format = DATETIME_FMT
        self._set_widths(ws, [10, 13, 14, 18, 17, 12, 12, 12, 12, 14, 16, 20, 42])
        ws.freeze_panes = "A5"
        if rows:
            self._table(ws, start_row=4, end_row=4 + len(rows), end_col=len(headers), name="tblDmTurmas")
        else:
            ws["A5"] = "Nenhuma turma no recorte atual."
            ws["A5"].font = self.static_font

    def _build_dm01(self) -> None:
        ws = self.ws["DM-01 Evolução"]
        self._title(ws, "DM-01 · EVOLUÇÃO POR TURMA", "Número absoluto de membros e composição de cada turma. A meta é institucional e aplicada turma a turma.", 12)
        headers = ["Área", "Turma", "Abertura", "Status turma", "Vagas", "Total", "Ativos", "Titulados", "Desligados", "Ocupação", "Meta membros", "Status DM-01"]
        self._headers(ws, 4, headers)
        rows = sorted(self.cohort_rows, key=lambda x: (str(x.get("area_code") or ""), int(x.get("cohort_number") or 0)))
        for r, row in enumerate(rows, 5):
            values = [
                row.get("area_code"), f"{row.get('area_code')} · T{row.get('cohort_number')}", _to_date(row.get("opening_date")), row.get("status"),
                row.get("vacancies_authorized"), row.get("total_students"), row.get("active_students"), row.get("graduated_students"),
                row.get("dropped_students"), None, self.dm01_target, None,
            ]
            for c, value in enumerate(values, 1):
                if c in {10, 12}:
                    continue
                self._body_cell(ws.cell(r, c), value, linked=True)
            ws.cell(r, 3).number_format = DATE_FMT
            for c in (5, 6, 7, 8, 9, 11):
                ws.cell(r, c).number_format = INTEGER_FMT
            ws.cell(r, 10).value = f'=IF(E{r}>0,F{r}/E{r},"")'
            ws.cell(r, 10).font = self.formula_font
            ws.cell(r, 10).number_format = PERCENT_FMT
            ws.cell(r, 12).value = (
                f'=IF(OR(D{r}="Planejada",D{r}="Aberta"),"Em formação",IF(F{r}=0,"Sem dados",'
                f'IF(K{r}="","Sem meta",IF(F{r}>=K{r},"Dentro da meta","Fora da meta"))))'
            )
            ws.cell(r, 12).font = self.formula_font
            ws.cell(r, 12).border = Border(bottom=Side(style="hair", color=BORDER))
            self._status_cell(ws.cell(r, 4), row.get("status"))
            self._status_cell(ws.cell(r, 12), row.get("dm01_status"))
        self._set_widths(ws, [10, 13, 14, 17, 12, 12, 12, 12, 12, 14, 15, 18])
        ws.freeze_panes = "A5"
        end = 4 + len(rows)
        if rows:
            self._table(ws, start_row=4, end_row=end, end_col=12, name="tblDm01Evolucao")
            chart = BarChart()
            chart.type = "col"
            chart.grouping = "clustered"
            chart.add_data(Reference(ws, min_col=6, max_col=9, min_row=4, max_row=end), titles_from_data=True)
            chart.set_categories(Reference(ws, min_col=2, min_row=5, max_row=end))
            self._chart_defaults(chart, "Composição das turmas", "Alunos")
            self._style_series(chart, [CHART_GRAY, CHART_GREEN, CHART_BLUE, CHART_RED])
            self._series_labels(chart, ["Total", "Ativos", "Titulados", "Desligados"])
            chart_row = end + 3
            ws.add_chart(chart, f"A{chart_row}")
            line = LineChart()
            line.add_data(Reference(ws, min_col=6, max_col=6, min_row=4, max_row=end), titles_from_data=True)
            if self.dm01_target is not None:
                line.add_data(Reference(ws, min_col=11, max_col=11, min_row=4, max_row=end), titles_from_data=True)
            line.set_categories(Reference(ws, min_col=2, min_row=5, max_row=end))
            self._chart_defaults(line, "Total de membros e meta", "Alunos")
            self._style_series(line, [CHART_GREEN, CHART_GOLD], line=True)
            self._series_labels(line, ["Total de membros", "Meta DM-01"] if self.dm01_target is not None else ["Total de membros"])
            ws.add_chart(line, f"G{chart_row}")

    def _build_dm02(self) -> None:
        ws = self.ws["DM-02 Defesas"]
        self._title(ws, "DM-02 · TEMPO ATÉ A DEFESA", "Ingresso individual confirmado → defesa. A abertura da turma não substitui o ingresso do aluno.", 16)
        ws.column_dimensions["P"].width = 3
        headers = [
            "Área", "Turma", "Status turma", "Alunos", "Ingressos confirmados", "Defesas", "Média meses", "Mediana meses",
            "Até 24 meses", "Risco >18m sem qualificação", "Risco >24m sem defesa marcada", "Risco >30m sem defesa",
            "Meta meses", "Status DM-02", "Última defesa",
        ]
        self._headers(ws, 4, headers)
        rows = sorted(self.cohort_rows, key=lambda x: (str(x.get("area_code") or ""), int(x.get("cohort_number") or 0)))
        for r, row in enumerate(rows, 5):
            values = [
                row.get("area_code"), f"{row.get('area_code')} · T{row.get('cohort_number')}", row.get("status"), row.get("total_students"),
                row.get("confirmed_entry_dates"), row.get("defenses_count"), row.get("average_months_to_defense"), row.get("median_months_to_defense"),
                (float(row.get("on_time_graduation_pct")) / 100) if row.get("on_time_graduation_pct") is not None else None,
                row.get("active_over_18_no_qualification"), row.get("active_over_24_no_scheduled_defense"), row.get("active_over_30_no_defense"),
                self.dm02_target, row.get("dm02_status"), _to_date(row.get("last_defense_date")),
            ]
            for c, value in enumerate(values, 1):
                self._body_cell(ws.cell(r, c), value, linked=True)
            for c in (4, 5, 6, 10, 11, 12):
                ws.cell(r, c).number_format = INTEGER_FMT
            for c in (7, 8, 13):
                ws.cell(r, c).number_format = DECIMAL_FMT
            ws.cell(r, 9).number_format = PERCENT_FMT
            ws.cell(r, 15).number_format = DATE_FMT
            self._status_cell(ws.cell(r, 3), row.get("status"))
            self._status_cell(ws.cell(r, 14), row.get("dm02_status"))
        self._set_widths(ws, [10, 15, 17, 11, 19, 11, 14, 14, 15, 21, 24, 20, 14, 20, 19])
        ws.freeze_panes = "A5"
        end = 4 + len(rows)
        if rows:
            self._table(ws, start_row=4, end_row=end, end_col=15, name="tblDm02Defesas")
            valid_rows = [row for row in rows if row.get("average_months_to_defense") is not None]
            compact_section = end + 3
            self._section(ws, compact_section, "TURMAS COM TEMPO CALCULADO", 5)
            self._headers(ws, compact_section + 1, ["Turma", "Média meses", "Meta meses", "Defesas", "Até 24 meses"])
            for rr, row in enumerate(valid_rows, compact_section + 2):
                vals = [
                    f"{row.get('area_code')} · T{row.get('cohort_number')}",
                    row.get("average_months_to_defense"), self.dm02_target, row.get("defenses_count"),
                    (float(row.get("on_time_graduation_pct")) / 100) if row.get("on_time_graduation_pct") is not None else None,
                ]
                for cc, value in enumerate(vals, 1):
                    self._body_cell(ws.cell(rr, cc), value, linked=True)
                ws.cell(rr, 2).number_format = DECIMAL_FMT
                ws.cell(rr, 3).number_format = DECIMAL_FMT
                ws.cell(rr, 4).number_format = INTEGER_FMT
                ws.cell(rr, 5).number_format = PERCENT_FMT
            if valid_rows:
                compact_end = compact_section + 1 + len(valid_rows)
                self._table(ws, start_row=compact_section + 1, end_row=compact_end, end_col=5, name="tblDm02TempoCalculado")
                self.dm02_chart_data_start = compact_section + 2
                self.dm02_chart_data_end = compact_end
                line = LineChart()
                line.add_data(Reference(ws, min_col=2, max_col=3, min_row=compact_section + 1, max_row=compact_end), titles_from_data=True)
                line.set_categories(Reference(ws, min_col=1, min_row=compact_section + 2, max_row=compact_end))
                self._chart_defaults(line, "Tempo médio por turma x meta", "Meses")
                self._style_series(line, [CHART_GREEN, CHART_GOLD], line=True)
                self._series_labels(line, ["Média de meses", "Meta DM-02"])
                chart_row = compact_end + 3
            else:
                ws.cell(compact_section + 2, 1).value = "Ainda não há turma com tempo calculável no recorte."
                ws.cell(compact_section + 2, 1).font = self.static_font
                chart_row = compact_section + 5
            risk = BarChart()
            risk.type = "col"
            risk.grouping = "clustered"
            risk.add_data(Reference(ws, min_col=10, max_col=12, min_row=4, max_row=end), titles_from_data=True)
            risk.set_categories(Reference(ws, min_col=2, min_row=5, max_row=end))
            self._chart_defaults(risk, "Riscos de prazo entre alunos ativos", "Alunos")
            self._style_series(risk, [CHART_GOLD, CHART_TEAL, CHART_RED])
            self._series_labels(risk, [">18m sem qualificação", ">24m sem defesa marcada", ">30m sem defesa"])
            if valid_rows:
                ws.add_chart(line, f"A{chart_row}")
            ws.add_chart(risk, f"H{chart_row}")

    def _build_targets(self) -> None:
        ws = self.ws["Metas"]
        self._title(ws, "METAS · DM", "A Diretoria de Mestrado possui somente duas metas institucionais: membros por turma e tempo médio até a defesa.", 10)
        ws.column_dimensions["J"].width = 3
        self._kpi_card(ws, c1=1, c2=3, row=4, label="DM-01 · Membros por turma", value=self.dm01_target if self.dm01_target is not None else "—", note="Meta efetiva na data de corte", fmt=INTEGER_FMT)
        self._kpi_card(ws, c1=4, c2=6, row=4, label="DM-02 · Tempo até a defesa", value=self.dm02_target if self.dm02_target is not None else "—", note="Meses · referência padrão 24 quando não há meta cadastrada", fmt=DECIMAL_FMT, fill=PALE_TEAL)
        self._section(ws, 9, "HISTÓRICO DE METAS CADASTRADAS", 9)
        headers = ["Indicador", "Métrica", "Vigência inicial", "Vigência final", "Meta", "Unidade", "Situação na data de corte", "Justificativa", "Cadastrado por"]
        self._headers(ws, 10, headers)
        metric_labels = {
            ("DM-01", "cohort_members"): ("Membros por turma", "alunos"),
            ("DM-02", "average_months_to_defense"): ("Tempo médio até a defesa", "meses"),
        }
        rows = [row for row in self.targets if (str(row.get("indicator_code") or "").upper(), str(row.get("metric_key") or "")) in metric_labels]
        for r, row in enumerate(rows, 11):
            key = (str(row.get("indicator_code") or "").upper(), str(row.get("metric_key") or ""))
            label, unit = metric_labels[key]
            current_sem = _semester_for_date(self.as_of)
            valid_from = str(row.get("valid_from") or "")
            valid_to = str(row.get("valid_to") or "")
            situation = "Futura" if valid_from > current_sem else ("Encerrada" if valid_to and valid_to < current_sem else "Ativa")
            values = [key[0], label, valid_from, valid_to or None, row.get("target"), unit, situation, row.get("justification"), row.get("inserted_by")]
            for c, value in enumerate(values, 1):
                self._body_cell(ws.cell(r, c), value, linked=True, wrap=c in {8})
            ws.cell(r, 5).number_format = DECIMAL_FMT
            self._status_cell(ws.cell(r, 7), "Dentro da meta" if situation == "Ativa" else None)
        self._set_widths(ws, [13, 30, 18, 18, 14, 14, 20, 48, 30])
        ws.freeze_panes = "A11"
        if rows:
            self._table(ws, start_row=10, end_row=10 + len(rows), end_col=9, name="tblDmMetas")
        else:
            ws["A11"] = "Nenhuma meta personalizada cadastrada. DM-02 usa 24 meses como referência padrão."
            ws["A11"].font = self.static_font

    def _scope_quality(self) -> dict[str, int]:
        rows = self.student_rows
        total = len(rows)
        confirmed = sum(1 for row in rows if row.get("entry_date") and not row.get("entry_date_estimated"))
        provisional = sum(1 for row in rows if row.get("entry_date") and row.get("entry_date_estimated"))
        missing = sum(1 for row in rows if not row.get("entry_date"))
        sei_seen = sum(1 for row in rows if "SEI" in str(row.get("source_system") or "").upper())
        dates_checked = sum(1 for row in rows if row.get("last_course_dates_sei_at"))
        defenses_from_sei = sum(1 for row in rows if row.get("defense_date") and row.get("last_course_dates_sei_at") and row.get("status") == "Titulado")
        return {
            "total": total,
            "confirmed": confirmed,
            "provisional": provisional,
            "missing": missing,
            "sei_seen": sei_seen,
            "dates_checked": dates_checked,
            "defenses_from_sei": defenses_from_sei,
        }

    def _build_sei(self) -> None:
        ws = self.ws["Integração SEI"]
        self._title(ws, "INTEGRAÇÃO SEI · DM", f"Qualidade do recorte {self._scope_text()} e histórico institucional das sincronizações.", 14)
        ws.column_dimensions["N"].width = 3
        quality = self._scope_quality()
        self._kpi_card(ws, c1=1, c2=2, row=4, label="Alunos no recorte", value=quality["total"], fmt=INTEGER_FMT)
        self._kpi_card(ws, c1=3, c2=4, row=4, label="Ingressos confirmados", value=quality["confirmed"], fmt=INTEGER_FMT, fill=PALE_TEAL)
        self._kpi_card(ws, c1=5, c2=6, row=4, label="Ingressos pendentes", value=quality["missing"], fmt=INTEGER_FMT, fill=PALE_GOLD)
        self._kpi_card(ws, c1=7, c2=8, row=4, label="Reconhecidos pelo SEI", value=quality["sei_seen"], fmt=INTEGER_FMT, fill=PALE_BLUE)
        self._kpi_card(ws, c1=9, c2=10, row=4, label="Datas consultadas", value=quality["dates_checked"], fmt=INTEGER_FMT, fill=PALE_BLUE)
        self._kpi_card(ws, c1=11, c2=13, row=4, label="Defesas confirmadas pelo SEI", value=quality["defenses_from_sei"], fmt=INTEGER_FMT)
        self._section(ws, 9, "HISTÓRICO INSTITUCIONAL DE SINCRONIZAÇÕES", 13)
        headers = [
            "Execução", "Origem", "Status", "Turmas detectadas", "Turmas criadas", "Turmas atualizadas",
            "Alunos detectados", "Alunos criados", "Alunos atualizados", "Alunos inalterados", "Alunos não vistos",
            "Observações", "Concluída em",
        ]
        self._headers(ws, 10, headers)
        rows = self.sync_runs[:50]
        for r, row in enumerate(rows, 11):
            warnings = row.get("warnings") or []
            raw_status = str(row.get("status") or "")
            status_map = {
                "completed": "Concluída", "success": "Concluída", "sucesso": "Concluída", "concluido": "Concluída", "concluído": "Concluída",
                "failed": "Falha", "error": "Falha", "erro": "Falha", "running": "Em andamento", "processing": "Em andamento",
            }
            display_status = status_map.get(raw_status.casefold(), raw_status or "—")
            values = [
                row.get("id"), row.get("source_name") or row.get("source_type"), display_status, row.get("cohorts_detected"),
                row.get("cohorts_created"), row.get("cohorts_updated"), row.get("students_detected"), row.get("students_created"),
                row.get("students_updated"), row.get("students_unchanged"), row.get("students_not_seen"), " | ".join(map(str, warnings)),
                _to_datetime(row.get("completed_at")),
            ]
            for c, value in enumerate(values, 1):
                self._body_cell(ws.cell(r, c), value, linked=True, wrap=c in {2, 12})
            for c in range(4, 12):
                ws.cell(r, c).number_format = INTEGER_FMT
            ws.cell(r, 13).number_format = DATETIME_FMT
            visual = "Concluído" if display_status == "Concluída" else display_status
            self._status_cell(ws.cell(r, 3), visual)
        self._set_widths(ws, [12, 32, 16, 17, 15, 17, 18, 15, 17, 18, 17, 52, 20])
        ws.freeze_panes = "A11"
        if rows:
            self._table(ws, start_row=10, end_row=10 + len(rows), end_col=13, name="tblDmSeiHistorico")
        else:
            ws["A11"] = "Nenhuma sincronização SEI registrada."
            ws["A11"].font = self.static_font
        note_row = 13 + len(rows)
        ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row + 1, end_column=13)
        ws.cell(note_row, 1).value = "Observação: o histórico de execuções SEI é institucional e pode conter sincronizações de turmas fora do recorte atual. Os cartões acima usam somente os alunos do filtro exportado."
        ws.cell(note_row, 1).font = self.static_font
        ws.cell(note_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(note_row, 1).fill = PatternFill("solid", fgColor=GRAY)
        ws.row_dimensions[note_row].height = 28

    def _build_parameters(self) -> None:
        ws = self.ws["Parâmetros"]
        self._title(ws, "PARÂMETROS · DM", "Contexto exato usado para gerar este relatório.", 7)
        ws.column_dimensions["G"].width = 3
        selected_cohort = self.cohort_rows[0] if self.cohort_id and self.cohort_rows else None
        rows = [
            ("Versão do Data UNIVC", APP_VERSION),
            ("Diretoria", "DM · Diretoria de Mestrado"),
            ("Área", DM_AREAS.get(self.area_code, "Todas as áreas") if self.area_code else "Todas as áreas"),
            ("Turma", f"Turma {selected_cohort.get('cohort_number')}" if selected_cohort else "Todas as turmas"),
            ("Data de corte", self.as_of),
            ("Unidade de análise", "Turma"),
            ("Status ativos de aluno", "Ativo · Titulado · Desligado"),
            ("Regra de titulação", "Data de defesa registrada → Titulado"),
            ("DM-01", "Número absoluto de membros por turma"),
            ("DM-02", "Tempo médio entre ingresso individual confirmado e defesa"),
            ("Meta DM-01 efetiva", self.dm01_target if self.dm01_target is not None else "Sem meta cadastrada"),
            ("Meta DM-02 efetiva", self.dm02_target if self.dm02_target is not None else "Sem meta cadastrada"),
            ("Fonte", "Banco Data UNIVC + integrações SEI registradas"),
            ("Gerado em", datetime.now()),
        ]
        self._section(ws, 4, "CONTEXTO DA EXPORTAÇÃO", 6)
        for r, (label, value) in enumerate(rows, 5):
            ws.cell(r, 1).value = label
            ws.cell(r, 1).font = Font(name="Arial", size=9, bold=True, color=GRAY_TEXT)
            ws.cell(r, 1).fill = PatternFill("solid", fgColor=GRAY)
            ws.cell(r, 1).border = self.border
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
            ws.cell(r, 2).value = value
            ws.cell(r, 2).font = self.linked_font if r not in {11, 12} else self.body_font
            ws.cell(r, 2).fill = PatternFill("solid", fgColor=PALE_GREEN)
            ws.cell(r, 2).border = self.border
            ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="center")
            if isinstance(value, date) and not isinstance(value, datetime):
                ws.cell(r, 2).number_format = DATE_FMT
            if isinstance(value, datetime):
                ws.cell(r, 2).number_format = DATETIME_FMT
        self._set_widths(ws, [30, 28, 18, 18, 18, 18])

    def _build_summary(self) -> None:
        ws = self.ws["Resumo Executivo"]
        overall = self.dashboard.get("overall") or {}
        self._title(ws, "RESUMO EXECUTIVO · DM", f"Diretoria de Mestrado · {self._scope_text()} · corte em {self.as_of.strftime('%d/%m/%Y')}", 15)
        ws.column_dimensions["O"].width = 3
        self._kpi_card(ws, c1=1, c2=2, row=4, label="Turmas", value=overall.get("cohort_count", 0), fmt=INTEGER_FMT)
        self._kpi_card(ws, c1=3, c2=4, row=4, label="Alunos", value=overall.get("total_students", 0), fmt=INTEGER_FMT, fill=PALE_TEAL)
        self._kpi_card(ws, c1=5, c2=6, row=4, label="Ativos", value=overall.get("active_students", 0), fmt=INTEGER_FMT, fill=PALE_BLUE)
        self._kpi_card(ws, c1=7, c2=8, row=4, label="Titulados", value=overall.get("graduated_students", 0), fmt=INTEGER_FMT)
        self._kpi_card(ws, c1=9, c2=10, row=4, label="Desligados", value=overall.get("dropped_students", 0), fmt=INTEGER_FMT, fill=PALE_RED)
        self._kpi_card(ws, c1=11, c2=12, row=4, label="Tempo médio até defesa", value=overall.get("average_months_to_defense") if overall.get("average_months_to_defense") is not None else "—", note="meses", fmt=DECIMAL_FMT, fill=PALE_TEAL)
        occupancy = (float(overall.get("occupancy_pct")) / 100) if overall.get("occupancy_pct") is not None else None
        self._kpi_card(ws, c1=13, c2=14, row=4, label="Ocupação", value=occupancy if occupancy is not None else "—", fmt=PERCENT_FMT, fill=PALE_GOLD)

        self._section(ws, 9, "INDICADORES INSTITUCIONAIS", 14)
        self._headers(ws, 10, ["Indicador", "Métrica", "Resultado", "Meta", "Situação", "Leitura"])
        dm01_total = len(self.cohort_rows)
        dm01_good = sum(1 for row in self.cohort_rows if row.get("dm01_status") == "Dentro da meta")
        dm01_status = "Sem meta" if self.dm01_target is None else ("Dentro da meta" if dm01_total and dm01_good == dm01_total else "Fora da meta")
        dm02_avg = overall.get("average_months_to_defense")
        dm02_status = "Sem dados" if dm02_avg is None else ("Dentro da meta" if self.dm02_target is not None and float(dm02_avg) <= self.dm02_target else "Fora da meta")
        rows = [
            ("DM-01", "Membros por turma", f"{dm01_good} de {dm01_total} turma(s) dentro da meta" if dm01_total else "Sem turmas", self.dm01_target, dm01_status, "O acompanhamento oficial é turma a turma."),
            ("DM-02", "Tempo médio até a defesa", dm02_avg, self.dm02_target, dm02_status, "Ingresso individual confirmado → defesa."),
        ]
        for r, values in enumerate(rows, 11):
            for c, value in enumerate(values, 1):
                self._body_cell(ws.cell(r, c), value, linked=c in {3, 4}, wrap=c in {2, 6})
            if r == 12:
                ws.cell(r, 3).number_format = DECIMAL_FMT
                ws.cell(r, 4).number_format = DECIMAL_FMT
            else:
                ws.cell(r, 4).number_format = INTEGER_FMT
            self._status_cell(ws.cell(r, 5), values[4])
        self._set_widths(ws, [12, 26, 24, 14, 18, 48, 13, 13, 13, 13, 13, 13, 13, 13])

        self._section(ws, 15, "LEITURA RÁPIDA POR TURMA", 14)
        headers = ["Área", "Turma", "Ativos", "Titulados", "Desligados", "Total", "Tempo médio", "Meta DM-02"]
        self._headers(ws, 16, headers)
        cohort_rows = sorted(self.cohort_rows, key=lambda x: (str(x.get("area_code") or ""), int(x.get("cohort_number") or 0)))
        for r, row in enumerate(cohort_rows, 17):
            values = [row.get("area_code"), f"{row.get('area_code')} · T{row.get('cohort_number')}", row.get("active_students"), row.get("graduated_students"), row.get("dropped_students"), row.get("total_students"), row.get("average_months_to_defense"), self.dm02_target]
            for c, value in enumerate(values, 1):
                self._body_cell(ws.cell(r, c), value, linked=True)
            for c in (3, 4, 5, 6):
                ws.cell(r, c).number_format = INTEGER_FMT
            for c in (7, 8):
                ws.cell(r, c).number_format = DECIMAL_FMT
        end = 16 + len(cohort_rows)
        if cohort_rows:
            self._table(ws, start_row=16, end_row=end, end_col=8, name="tblDmResumoTurmas")
            chart = BarChart()
            chart.type = "col"
            chart.grouping = "stacked"
            chart.add_data(Reference(ws, min_col=3, max_col=5, min_row=16, max_row=end), titles_from_data=True)
            chart.set_categories(Reference(ws, min_col=2, min_row=17, max_row=end))
            self._chart_defaults(chart, "Composição das turmas", "Alunos", width=13.3, height=7.0)
            self._style_series(chart, [CHART_GREEN, CHART_BLUE, CHART_RED])
            self._series_labels(chart, ["Ativos", "Titulados", "Desligados"])
            chart_row = end + 3
            ws.add_chart(chart, f"A{chart_row}")
            if self.dm02_chart_data_start and self.dm02_chart_data_end:
                source = self.ws["DM-02 Defesas"]
                line = LineChart()
                header = self.dm02_chart_data_start - 1
                line.add_data(Reference(source, min_col=2, max_col=3, min_row=header, max_row=self.dm02_chart_data_end), titles_from_data=True)
                line.set_categories(Reference(source, min_col=1, min_row=self.dm02_chart_data_start, max_row=self.dm02_chart_data_end))
                self._chart_defaults(line, "Tempo médio até a defesa", "Meses", width=13.3, height=7.0)
                self._style_series(line, [CHART_GREEN, CHART_GOLD], line=True)
                self._series_labels(line, ["Média de meses", "Meta DM-02"])
                ws.add_chart(line, f"H{chart_row}")

    def _finish(self) -> None:
        self.wb.active = self.wb.sheetnames.index("Resumo Executivo")
        self.wb.calculation.fullCalcOnLoad = True
        self.wb.calculation.forceFullCalc = True
        self.wb.calculation.calcMode = "auto"
        tab_colors = {
            "Resumo Executivo": GREEN_DARK,
            "DM-01 Evolução": GREEN,
            "DM-02 Defesas": GREEN,
            "Turmas": CHART_BLUE,
            "Alunos e Defesas": CHART_BLUE,
            "Metas": CHART_GOLD,
            "Integração SEI": CHART_TEAL,
            "Parâmetros": CHART_GRAY,
        }
        for name, color in tab_colors.items():
            self.ws[name].sheet_properties.tabColor = color


def build_dm_v2_workbook(
    cohorts: Iterable[dict[str, Any]],
    students: Iterable[dict[str, Any]],
    *,
    targets: Iterable[dict[str, Any]] | None = None,
    sync_runs: Iterable[dict[str, Any]] | None = None,
    area_code: str | None = None,
    cohort_id: int | None = None,
    as_of: date | None = None,
) -> BytesIO:
    builder = DMExcelV2Builder(
        cohorts,
        students,
        targets=targets,
        sync_runs=sync_runs,
        area_code=area_code,
        cohort_id=cohort_id,
        as_of=as_of,
    )
    workbook = builder.build()
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
