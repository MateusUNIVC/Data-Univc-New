from __future__ import annotations

from io import BytesIO
from copy import copy, deepcopy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill

from analytics import active_goal
from dadm_analytics import build_dadm_dashboard
from dpe_analytics import build_dpe_dashboard
from repository import DatabaseRepository
from schemas import KPI_META
from excel_errors import ExcelExportLimitError


ROOT = Path(__file__).resolve().parent
DTNH_TEMPLATE = ROOT / "assets" / "paineis" / "Painel_DTNH_template.xlsx"
DTNH_STYLE_REFERENCE = ROOT / "assets" / "referencias" / "Painel_DTNH_padrao_v4.xlsx"
DCS_TEMPLATE = ROOT / "assets" / "paineis" / "Painel_DCS_template.xlsx"
DADM_TEMPLATE = ROOT / "assets" / "paineis" / "Painel_DADM_template.xlsx"
DPE_TEMPLATE = ROOT / "assets" / "paineis" / "Painel_DPE_template.xlsx"


def _export_academic(
    repo: DatabaseRepository,
    *,
    granularity: str | None = None,
    reference: str | None = None,
    comparison: str | None = None,
    course: str | None = None,
    discipline: str | None = None,
    window_periods: int | str | None = None,
) -> BytesIO:
    """Exporta DTNH/DCS no relatório acadêmico V2 orientado ao dashboard.

    A rota oficial não expõe mais o workbook operacional/legado com abas de
    cálculo e bases editáveis. O relatório usa apenas agregações gerenciais e
    preserva o mesmo contexto selecionado na interface acadêmica.
    """
    from academic_excel_v2_builder import build_academic_report_workbook_bytes

    return build_academic_report_workbook_bytes(
        repo,
        granularity=granularity,
        reference=reference,
        comparison=comparison,
        course=course,
        discipline=discipline,
        window_periods=window_periods,
    )

def _clear_values(ws, start_row: int, end_row: int, start_col: int, end_col: int) -> None:
    for row in range(start_row, end_row + 1):
        for col in range(start_col, end_col + 1):
            ws.cell(row, col).value = None


EXCEL_MAX_ROW = 1_048_576


def _dynamic_end_row(label: str, start_row: int, item_count: int, template_end_row: int) -> int:
    """Return a safe sheet end row without ever discarding source records."""
    actual_end = start_row + max(0, int(item_count)) - 1
    if actual_end > EXCEL_MAX_ROW:
        raise ExcelExportLimitError(
            f"{label} possui {item_count:,} registros e ultrapassa o limite físico "
            f"de {EXCEL_MAX_ROW:,} linhas do Excel. Nenhum registro foi exportado parcialmente."
        )
    return max(template_end_row, actual_end)


def _ensure_sheet_rows(
    ws,
    end_row: int,
    *,
    style_row: int,
    start_col: int,
    end_col: int,
) -> None:
    """Extend a legacy template while preserving the last designed row style."""
    if end_row <= ws.max_row:
        return
    source_height = ws.row_dimensions[style_row].height
    for row in range(ws.max_row + 1, end_row + 1):
        if source_height is not None:
            ws.row_dimensions[row].height = source_height
        for col in range(start_col, end_col + 1):
            source = ws.cell(style_row, col)
            target = ws.cell(row, col)
            if source.has_style:
                target._style = copy(source._style)
            if source.alignment is not None:
                target.alignment = copy(source.alignment)
            if source.protection is not None:
                target.protection = copy(source.protection)


def _set_validation_ref(ws, old_ref: str, new_ref: str) -> None:
    for validation in ws.data_validations.dataValidation:
        if str(validation.sqref) == old_ref:
            validation.sqref = new_ref


def _expand_tables_to_row(ws, end_row: int) -> None:
    """Keep Excel table/autofilter ranges aligned with dynamically written data."""
    from openpyxl.utils.cell import get_column_letter, range_boundaries

    for name in ws.tables:
        table = ws.tables[name]
        min_col, min_row, max_col, max_row = range_boundaries(table.ref)
        if end_row <= max_row:
            continue
        new_ref = (
            f"{get_column_letter(min_col)}{min_row}:"
            f"{get_column_letter(max_col)}{end_row}"
        )
        table.ref = new_ref
        if table.autoFilter is not None:
            table.autoFilter.ref = new_ref


def _status_fill(status: str | None) -> PatternFill:
    text = str(status or '').lower()
    if 'dentro' in text:
        return PatternFill('solid', fgColor='E8F1ED')
    if 'atenção' in text or 'atencao' in text:
        return PatternFill('solid', fgColor='FFF4D8')
    if 'fora' in text:
        return PatternFill('solid', fgColor='FDE7E7')
    return PatternFill('solid', fgColor='F5F6F5')


def _goal_for(snapshot: dict, code: str, period: str | None, scope: str | None = None):
    if not period:
        return None
    return active_goal(snapshot.get('metas', []), code, period, scope)


def _goal_value(goal: dict | None, field: str = 'meta'):
    if not goal or goal.get(field) in (None, ''):
        return None
    try:
        return float(goal[field])
    except (TypeError, ValueError):
        return None


def _lower_status(value: float | None, goal: dict | None) -> str:
    if value is None:
        return 'Sem dados'
    target = _goal_value(goal, 'meta')
    if target is None:
        return 'Sem meta'
    attention = _goal_value(goal, 'atencao')
    if attention is None or attention < target:
        attention = target
    if value <= target:
        return 'Dentro da meta'
    if value <= attention:
        return 'Atenção'
    return 'Fora da meta'


def _lower_situation(value: float | None, goal: dict | None) -> str:
    status = _lower_status(value, goal)
    if status == 'Dentro da meta':
        return 'Melhor ou igual à meta'
    if status == 'Atenção':
        return 'Entre a meta e o limiar'
    if status == 'Fora da meta':
        return 'Acima do limiar de atenção'
    return status


def _set_status_cell(cell, status: str | None) -> None:
    cell.value = status or ''
    cell.fill = _status_fill(status)
    cell.font = Font(name='Arial', size=10, bold=True, color='000000')
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)


def _valid_rows(rows, period: str | None = None, *, course_id=None, modality=None, center=None):
    out = []
    for row in rows or []:
        if row.get('validacao') not in (None, '', 'OK'):
            continue
        if period is not None and str(row.get('periodo') or '') != period:
            continue
        if course_id is not None and row.get('course_id') != course_id:
            continue
        if modality is not None and str(row.get('modalidade') or '') != modality:
            continue
        if center is not None and str(row.get('centro_custo') or '') != center:
            continue
        out.append(row)
    return out


def _active_students(snapshot: dict, period: str | None, *, course_id=None, modality=None) -> int | None:
    rows = _valid_rows(snapshot.get('alunos', []), period, course_id=course_id, modality=modality)
    if not rows:
        return None
    return sum(int(row.get('alunos_ativos') or 0) for row in rows)


def _attrition(snapshot: dict, period: str | None, *, course_id=None, modality=None) -> tuple[float | None, int, int]:
    rows = _valid_rows(snapshot.get('evasao', []), period, course_id=course_id, modality=modality)
    base = sum(int(row.get('alunos_inicio') or 0) for row in rows)
    departures = sum(int(row.get('desligamentos') or 0) for row in rows)
    value = round(departures / base * 100, 2) if base else None
    return value, base, departures


def _cost_metrics(snapshot: dict, period: str | None, *, modality=None, center=None) -> tuple[float | None, float, int | None]:
    costs = _valid_rows(snapshot.get('custos', []), period, modality=modality, center=center)
    expense = sum(float(row.get('despesa') or 0) for row in costs)
    students = _active_students(snapshot, period, modality=modality)
    value = round(expense / students, 2) if expense and students else None
    return value, round(expense, 2), students


def _infra_metrics(snapshot: dict, period: str | None) -> dict:
    infra = _valid_rows(snapshot.get('infraestrutura', []), period)
    expense = sum(float(row.get('despesa') or 0) for row in infra)
    area = sum(float(row.get('area_m2') or 0) for row in infra)
    students = _active_students(snapshot, period)
    return {
        'despesa': round(expense, 2),
        'area_m2': round(area, 2),
        'alunos': students,
        'custo_m2': round(expense / area, 2) if expense and area else None,
        'custo_aluno': round(expense / students, 2) if expense and students else None,
    }


def _pct_change(current: float | None, previous: float | None) -> float | None:
    if current is None or previous in (None, 0):
        return None
    return round((float(current) - float(previous)) / float(previous) * 100, 2)


def _periods_for_export(dashboard: dict) -> list[str]:
    # Match the DTNH panel principle: the displayed window starts when an actual KPI
    # has a value. A population-only month (used only as a denominator baseline) must
    # not create an empty first bar in all three executive charts.
    periods = set()
    specs = {
        'evasao': 'valor',
        'custo_administrativo': 'valor',
        'infraestrutura': 'custo_m2',
    }
    for series_name, value_key in specs.items():
        for item in dashboard.get('series', {}).get(series_name, []) or []:
            period = str(item.get('periodo') or '')
            if period and item.get(value_key) is not None:
                periods.add(period)
    return sorted(periods)[-12:]


def _series_item(dashboard: dict, series_name: str, period: str | None) -> dict:
    return next((x for x in dashboard.get('series', {}).get(series_name, []) if x.get('periodo') == period), {})


def _write_panel_context(wb, dashboard: dict, snapshot: dict) -> None:
    ctx = dashboard.get('contexto', {})
    periods = _periods_for_export(dashboard)
    params = wb['PARAMETROS']
    params['B5'] = 'DADM'
    params['B6'] = periods[0] if periods else (ctx.get('inicio') or '')
    params['B7'] = ctx.get('referencia') or ''
    params['B8'] = ctx.get('comparacao') or ''
    params['B9'] = periods[0] if periods else (ctx.get('inicio') or '')
    params['B10'] = periods[-1] if periods else (ctx.get('fim') or '')
    params['B11'] = ctx.get('curso_label') or '(todos)'
    params['B12'] = ctx.get('modalidade') or '(todas)'
    params['B13'] = ctx.get('setor') or '(todos)'

    # Lists used by the same data validations as the DTNH workbook. They grow
    # with the actual catalog instead of stopping silently at row 34.
    courses = list(snapshot.get('cursos_academicos', []) or [])
    centers = sorted({
        str(item.get('centro_custo') or '').strip()
        for item in snapshot.get('custos', []) or []
        if str(item.get('centro_custo') or '').strip()
    })
    course_end = _dynamic_end_row('Lista de cursos da DADM', 6, len(courses), 34)
    center_end = _dynamic_end_row('Lista de centros de custo da DADM', 6, len(centers), 34)
    support_end = max(course_end, center_end)

    ws = wb['LISTAS DE APOIO']
    _ensure_sheet_rows(ws, support_end, style_row=min(14, ws.max_row), start_col=1, end_col=9)
    _clear_values(ws, 5, support_end, 8, 9)
    ws['H5'] = '(todos)'
    for row, course in enumerate(courses, 6):
        ws.cell(row, 8).value = course.get('label') or course.get('curso')
    ws['I5'] = '(todos)'
    for row, center in enumerate(centers, 6):
        ws.cell(row, 9).value = center

    for validation in params.data_validations.dataValidation:
        target = str(validation.sqref)
        if target == 'B11':
            validation.formula1 = f"'LISTAS DE APOIO'!$H$5:$H${course_end}"
        elif target == 'B13':
            validation.formula1 = f"'LISTAS DE APOIO'!$I$5:$I${center_end}"


def _populate_calc(wb, dashboard: dict, snapshot: dict) -> None:
    ws = wb['CALC']
    # Keep the headings and styling that came directly from DTNH; clear only data blocks.
    for start, end in ((8,17),(22,33),(37,37),(41,50),(55,66),(70,70),(74,83),(88,99),(103,103),(107,115)):
        _clear_values(ws, start, end, 1, 15)

    ctx = dashboard.get('contexto', {})
    ref = ctx.get('referencia')
    comp = ctx.get('comparacao')
    periods = _periods_for_export(dashboard)
    ws['B4'] = ctx.get('inicio') or (periods[0] if periods else '')
    ws['E4'] = ctx.get('fim') or (periods[-1] if periods else '')
    ws['H4'] = ref or ''
    ws['K4'] = comp or ''
    ws['N4'] = ' · '.join(x for x in [ctx.get('curso_label'), ctx.get('modalidade'), ctx.get('setor')] if x) or 'Consolidado'

    # Goal inheritance table: TOTAL in first row, then course-specific attrition goals when present.
    scopes = ['TOTAL']
    scopes.extend(sorted({str(g.get('recorte')) for g in snapshot.get('metas', []) if g.get('indicador') == 'DADM-01' and str(g.get('recorte') or '') not in ('', 'TOTAL')}))
    for row, scope in enumerate(scopes[:10], 8):
        ws.cell(row, 1).value = scope
        g1 = _goal_for(snapshot, 'DADM-01', ref, None if scope == 'TOTAL' else scope)
        g9 = _goal_for(snapshot, 'DADM-09', ref)
        g10 = _goal_for(snapshot, 'DADM-10', ref)
        pairs = [(g1,2,3),(g9,4,5),(g10,6,7)]
        for goal, c_meta, c_attn in pairs:
            ws.cell(row, c_meta).value = _goal_value(goal, 'meta')
            ws.cell(row, c_attn).value = _goal_value(goal, 'atencao')
        ws.cell(row,2).number_format = ws.cell(row,3).number_format = '0.00"%"'
        for c in (4,5,6,7): ws.cell(row,c).number_format = 'R$ #,##0.00'

    # DADM-01 monthly block (rows 22:33).
    for row, period in enumerate(periods, 22):
        value, base, departures = _attrition(snapshot, period)
        goal = _goal_for(snapshot, 'DADM-01', period)
        vals = [period, base, departures, value, _goal_value(goal), _goal_value(goal,'atencao'), _lower_status(value,goal), len(_valid_rows(snapshot.get('evasao',[]),period))]
        for c,v in enumerate(vals,1): ws.cell(row,c).value=v
        for c in (4,5,6): ws.cell(row,c).number_format='0.00"%"'
        ws.cell(row,2).number_format=ws.cell(row,3).number_format='#,##0'

    ev = dashboard.get('cards', {}).get('evasao', {})
    ev_goal = _goal_for(snapshot, 'DADM-01', ev.get('periodo') or ref)
    ev_value = ev.get('valor'); ev_comp = ev.get('comparacao')
    summary = [
        'DADM-01', ev_value, ev_comp, ev.get('variacao'), _goal_value(ev_goal), _goal_value(ev_goal,'atencao'),
        None if ev_value is None or _goal_value(ev_goal) is None else round(_goal_value(ev_goal)-float(ev_value),2),
        None if ev_value is None or _goal_value(ev_goal,'atencao') is None else round(_goal_value(ev_goal,'atencao')-float(ev_value),2),
        None if ev_value in (None,0) or _goal_value(ev_goal) is None else _goal_value(ev_goal)/float(ev_value),
        _lower_status(ev_value,ev_goal), _lower_situation(ev_value,ev_goal), len(_valid_rows(snapshot.get('evasao',[]),ev.get('periodo') or ref)),
        None, ev.get('periodo') or ref, ev.get('comparacao_periodo') or comp,
    ]
    for c,v in enumerate(summary,1): ws.cell(37,c).value=v
    for c in (2,3,5,6,7,8): ws.cell(37,c).number_format='0.00"%"'
    ws.cell(37,4).number_format='+0.00" p.p.";-0.00" p.p.";0.00" p.p."'
    ws.cell(37,9).number_format='0%'

    courses = snapshot.get('cursos_academicos', [])
    for row, course in enumerate(courses[:10], 41):
        cid = course.get('id'); name = course.get('curso')
        rv, rbase, rdep = _attrition(snapshot, ref, course_id=cid)
        cv, _, _ = _attrition(snapshot, comp, course_id=cid)
        goal = _goal_for(snapshot, 'DADM-01', ref, name)
        vals=[name,course.get('diretoria'),course.get('modalidade'),rbase,rdep,rv,cv,None if rv is None or cv is None else round(rv-cv,2),_goal_value(goal),_goal_value(goal,'atencao'),_lower_status(rv,goal)]
        for c,v in enumerate(vals,1): ws.cell(row,c).value=v
        for c in (6,7,8,9,10): ws.cell(row,c).number_format='0.00"%"'
        ws.cell(row,4).number_format=ws.cell(row,5).number_format='#,##0'

    # DADM-09 monthly block.
    for row, period in enumerate(periods, 55):
        value, expense, students = _cost_metrics(snapshot, period)
        goal=_goal_for(snapshot,'DADM-09',period)
        vals=[period,expense,students,value,_goal_value(goal),_goal_value(goal,'atencao'),_lower_status(value,goal),len(_valid_rows(snapshot.get('custos',[]),period))]
        for c,v in enumerate(vals,1): ws.cell(row,c).value=v
        for c in (2,4,5,6): ws.cell(row,c).number_format='R$ #,##0.00'
        ws.cell(row,3).number_format='#,##0'

    cost = dashboard.get('cards',{}).get('custo_administrativo',{})
    cost_goal=_goal_for(snapshot,'DADM-09',cost.get('periodo') or ref)
    cv=cost.get('valor'); cc=cost.get('comparacao')
    summary=[
        'DADM-09',cv,cc,cost.get('variacao'),_goal_value(cost_goal),_goal_value(cost_goal,'atencao'),
        None if cv is None or _goal_value(cost_goal) is None else round(_goal_value(cost_goal)-float(cv),2),
        None if cv is None or _goal_value(cost_goal,'atencao') is None else round(_goal_value(cost_goal,'atencao')-float(cv),2),
        None if cv in (None,0) or _goal_value(cost_goal) is None else _goal_value(cost_goal)/float(cv),
        _lower_status(cv,cost_goal),_lower_situation(cv,cost_goal),len(_valid_rows(snapshot.get('custos',[]),cost.get('periodo') or ref)),None,cost.get('periodo') or ref,cost.get('comparacao_periodo') or comp,
    ]
    for c,v in enumerate(summary,1): ws.cell(70,c).value=v
    for c in (2,3,5,6,7,8): ws.cell(70,c).number_format='R$ #,##0.00'
    ws.cell(70,4).number_format='+0.00"%";-0.00"%";0.00"%"'
    ws.cell(70,9).number_format='0%'

    centers=sorted({str(x.get('centro_custo') or '') for x in snapshot.get('custos',[]) if str(x.get('centro_custo') or '')})
    total_ref=sum(float(x.get('despesa') or 0) for x in _valid_rows(snapshot.get('custos',[]),ref))
    center_rows=[]
    for center in centers:
        ref_rows=_valid_rows(snapshot.get('custos',[]),ref,center=center)
        comp_rows=_valid_rows(snapshot.get('custos',[]),comp,center=center)
        rexp=sum(float(x.get('despesa') or 0) for x in ref_rows)
        cexp=sum(float(x.get('despesa') or 0) for x in comp_rows)
        mods=sorted({str(x.get('modalidade') or '') for x in ref_rows})
        reading='Sem lançamento'
        var=_pct_change(rexp if ref_rows else None,cexp if comp_rows else None)
        if ref_rows:
            reading='Maior participação' if total_ref and rexp/total_ref >= 0.20 else ('Despesa em alta' if var is not None and var > 5 else 'Acompanhar')
        center_rows.append((rexp,center,mods,cexp,var,reading))
    for row, item in enumerate(sorted(center_rows,reverse=True)[:10],74):
        rexp,center,mods,cexp,var,reading=item
        vals=[center,rexp,(', '.join(mods) if len(mods)>1 else (mods[0] if mods else '')),rexp,cexp if cexp else None,var,(rexp/total_ref if total_ref else None),reading]
        for c,v in enumerate(vals,1): ws.cell(row,c).value=v
        for c in (2,4,5): ws.cell(row,c).number_format='R$ #,##0.00'
        ws.cell(row,6).number_format='+0.00%;-0.00%;0.00%'
        ws.cell(row,7).number_format='0.0%'

    # DADM-10 monthly block.
    for row, period in enumerate(periods, 88):
        m=_infra_metrics(snapshot,period); goal=_goal_for(snapshot,'DADM-10',period)
        vals=[period,m['despesa'],m['area_m2'],m['alunos'],m['custo_m2'],_goal_value(goal),_goal_value(goal,'atencao'),m['custo_aluno'],_lower_status(m['custo_m2'],goal),len(_valid_rows(snapshot.get('infraestrutura',[]),period))]
        for c,v in enumerate(vals,1): ws.cell(row,c).value=v
        for c in (2,5,6,7,8): ws.cell(row,c).number_format='R$ #,##0.00'
        ws.cell(row,3).number_format='#,##0.00'; ws.cell(row,4).number_format='#,##0'

    infra=dashboard.get('cards',{}).get('infraestrutura',{})
    infra_goal=_goal_for(snapshot,'DADM-10',infra.get('periodo') or ref)
    iv=infra.get('custo_m2'); ic=infra.get('comparacao_custo_m2')
    summary=[
        'DADM-10',iv,ic,infra.get('variacao_custo_m2'),_goal_value(infra_goal),_goal_value(infra_goal,'atencao'),
        None if iv is None or _goal_value(infra_goal) is None else round(_goal_value(infra_goal)-float(iv),2),
        None if iv is None or _goal_value(infra_goal,'atencao') is None else round(_goal_value(infra_goal,'atencao')-float(iv),2),
        None if iv in (None,0) or _goal_value(infra_goal) is None else _goal_value(infra_goal)/float(iv),
        _lower_status(iv,infra_goal),_lower_situation(iv,infra_goal),len(_valid_rows(snapshot.get('infraestrutura',[]),infra.get('periodo') or ref)),None,infra.get('periodo') or ref,infra.get('comparacao_periodo') or comp,
    ]
    for c,v in enumerate(summary,1): ws.cell(103,c).value=v
    for c in (2,3,5,6,7,8): ws.cell(103,c).number_format='R$ #,##0.00'
    ws.cell(103,4).number_format='+0.00"%";-0.00"%";0.00"%"'
    ws.cell(103,9).number_format='0%'

    rm=_infra_metrics(snapshot,ref); cm=_infra_metrics(snapshot,comp)
    detail=[
        ('R$/m²',rm['custo_m2'],cm['custo_m2'],_pct_change(rm['custo_m2'],cm['custo_m2']),_goal_value(infra_goal),_goal_value(infra_goal,'atencao'),_lower_status(rm['custo_m2'],infra_goal)),
        ('R$/aluno',rm['custo_aluno'],cm['custo_aluno'],_pct_change(rm['custo_aluno'],cm['custo_aluno']),None,None,'Leitura auxiliar'),
        ('Despesa de infraestrutura',rm['despesa'],cm['despesa'],_pct_change(rm['despesa'],cm['despesa']),None,None,'Leitura auxiliar'),
        ('Área em uso (m²)',rm['area_m2'],cm['area_m2'],_pct_change(rm['area_m2'],cm['area_m2']),None,None,'Leitura auxiliar'),
        ('Alunos ativos',rm['alunos'],cm['alunos'],_pct_change(rm['alunos'],cm['alunos']),None,None,'Leitura auxiliar'),
    ]
    for row,(label,rv,cv,var,meta,attn,status) in enumerate(detail,107):
        vals=[label,'Institucional','Principal' if label=='R$/m²' else 'Complementar',rv,cv,var,meta,attn,status]
        for c,v in enumerate(vals,1): ws.cell(row,c).value=v
        if label in ('R$/m²','R$/aluno','Despesa de infraestrutura'):
            for c in (4,5,7,8): ws.cell(row,c).number_format='R$ #,##0.00'
        elif label=='Área em uso (m²)':
            ws.cell(row,4).number_format=ws.cell(row,5).number_format='#,##0.00'
        else:
            ws.cell(row,4).number_format=ws.cell(row,5).number_format='#,##0'
        ws.cell(row,6).number_format='+0.00"%";-0.00"%";0.00"%"'


def _add_dtnh_style_chart(
    ws_panel,
    reference_chart,
    title: str,
    data_col_letter: str,
    header_row: int,
    start_row: int,
    end_row: int,
    y_format: str,
) -> None:
    """Clone the DTNH chart itself, changing only DADM data/title.

    This intentionally preserves the reference workbook's anchor, plot geometry,
    legend, gap width, fill, axes and typography instead of recreating an
    approximation with a second chart style.
    """
    if end_row < start_row:
        return
    chart = deepcopy(reference_chart)
    chart.title = title
    chart.y_axis.number_format = y_format
    series = chart.series[0]
    series.val.numRef.f = f"CALC!${data_col_letter}${start_row}:${data_col_letter}${end_row}"
    if series.cat and series.cat.strRef:
        series.cat.strRef.f = f"CALC!$A${start_row}:$A${end_row}"
    elif series.cat and series.cat.numRef:
        series.cat.numRef.f = f"CALC!$A${start_row}:$A${end_row}"
    if series.tx and series.tx.strRef:
        series.tx.strRef.f = f"CALC!{data_col_letter}{header_row}"
    ws_panel._charts.append(chart)

def _populate_dadm_bases(wb, snapshot: dict) -> None:
    evasao = list(snapshot.get('evasao', []) or [])
    ws = wb['DADM-01 EVASAO']
    end_row = _dynamic_end_row('Base DADM-01 Evasão', 5, len(evasao), 1200)
    _ensure_sheet_rows(ws, end_row, style_row=1200, start_col=1, end_col=10)
    _clear_values(ws, 5, end_row, 1, 10)
    _expand_tables_to_row(ws, end_row)
    for row, item in enumerate(reversed(evasao), 5):
        goal = _goal_for(snapshot, 'DADM-01', item.get('periodo'), item.get('curso'))
        value = item.get('valor')
        status = _lower_status(value, goal)
        vals = [
            item.get('periodo'), item.get('diretoria_academica'), item.get('curso'),
            item.get('modalidade'), item.get('alunos_inicio'), item.get('desligamentos'),
            value, _goal_value(goal), status, item.get('lancado_por'),
        ]
        for col, val in enumerate(vals, 1):
            ws.cell(row, col).value = val
        ws.cell(row, 5).number_format = ws.cell(row, 6).number_format = '#,##0'
        ws.cell(row, 7).number_format = ws.cell(row, 8).number_format = '0.00"%"'
        _set_status_cell(ws.cell(row, 9), status)

    custos = list(snapshot.get('custos', []) or [])
    ws = wb['DADM-09 CUSTOS']
    end_row = _dynamic_end_row('Base DADM-09 Custos', 5, len(custos), 1200)
    _ensure_sheet_rows(ws, end_row, style_row=1200, start_col=1, end_col=5)
    _clear_values(ws, 5, end_row, 1, 5)
    _expand_tables_to_row(ws, end_row)
    for row, item in enumerate(reversed(custos), 5):
        vals = [
            item.get('periodo'), item.get('centro_custo'), item.get('modalidade'),
            item.get('despesa'), item.get('lancado_por'),
        ]
        for col, val in enumerate(vals, 1):
            ws.cell(row, col).value = val
        ws.cell(row, 4).number_format = 'R$ #,##0.00'

    alunos = list(snapshot.get('alunos', []) or [])
    ws = wb['DADM-09 ALUNOS']
    end_row = _dynamic_end_row('Base DADM-09 Alunos', 5, len(alunos), 1200)
    _ensure_sheet_rows(ws, end_row, style_row=1200, start_col=1, end_col=7)
    _clear_values(ws, 5, end_row, 1, 7)
    _expand_tables_to_row(ws, end_row)
    for row, item in enumerate(reversed(alunos), 5):
        vals = [
            item.get('periodo'), item.get('diretoria_academica'), item.get('curso'),
            item.get('modalidade'), item.get('alunos_ativos'), item.get('origem'),
            'Sim' if item.get('temporario') else 'Não',
        ]
        for col, val in enumerate(vals, 1):
            ws.cell(row, col).value = val
        ws.cell(row, 5).number_format = '#,##0'

    infraestrutura = list(snapshot.get('infraestrutura', []) or [])
    ws = wb['DADM-10 INFRA']
    end_row = _dynamic_end_row('Base DADM-10 Infraestrutura', 5, len(infraestrutura), 1200)
    _ensure_sheet_rows(ws, end_row, style_row=1200, start_col=1, end_col=9)
    _clear_values(ws, 5, end_row, 1, 9)
    _expand_tables_to_row(ws, end_row)
    for row, item in enumerate(reversed(infraestrutura), 5):
        period = item.get('periodo')
        metrics = _infra_metrics(snapshot, period)
        goal = _goal_for(snapshot, 'DADM-10', period)
        status = _lower_status(metrics.get('custo_m2'), goal)
        vals = [
            period, metrics['despesa'], metrics['area_m2'], metrics['custo_m2'],
            metrics['alunos'], metrics['custo_aluno'], _goal_value(goal), status,
            item.get('lancado_por'),
        ]
        for col, val in enumerate(vals, 1):
            ws.cell(row, col).value = val
        for col in (2, 4, 6, 7):
            ws.cell(row, col).number_format = 'R$ #,##0.00'
        ws.cell(row, 3).number_format = '#,##0.00'
        ws.cell(row, 5).number_format = '#,##0'
        _set_status_cell(ws.cell(row, 8), status)


def _populate_matrix_and_courses(wb, dashboard: dict, snapshot: dict) -> None:
    ws=wb['MATRIZ']
    periods=_periods_for_export(dashboard)
    ref=dashboard.get('contexto',{}).get('referencia')
    comp=dashboard.get('contexto',{}).get('comparacao')

    # Clear only the four data regions; section labels/styles remain exactly where DTNH puts them.
    for start,end in ((5,14),(18,27),(31,40),(44,52)):
        _clear_values(ws,start,end,1,18 if start!=44 else 14)

    # DADM-01 by course.
    for idx,period in enumerate(periods[:12],2): ws.cell(5,idx).value=period
    tail=['Referência','Comparação','Variação','Meta','Status']
    for c,v in enumerate(tail,14): ws.cell(5,c).value=v
    for row,course in enumerate(snapshot.get('cursos_academicos',[])[:9],6):
        ws.cell(row,1).value=f"{course.get('diretoria')} · {course.get('curso')}"
        for idx,period in enumerate(periods[:12],2):
            value,_,_=_attrition(snapshot,period,course_id=course.get('id'))
            ws.cell(row,idx).value=value; ws.cell(row,idx).number_format='0.00"%"'
        rv,_,_=_attrition(snapshot,ref,course_id=course.get('id')); cv,_,_=_attrition(snapshot,comp,course_id=course.get('id'))
        goal=_goal_for(snapshot,'DADM-01',ref,course.get('curso'))
        vals=[rv,cv,None if rv is None or cv is None else round(rv-cv,2),_goal_value(goal),_lower_status(rv,goal)]
        for c,v in enumerate(vals,14): ws.cell(row,c).value=v
        for c in (14,15,16,17): ws.cell(row,c).number_format='0.00"%"'
        _set_status_cell(ws.cell(row,18),vals[-1])

    # DADM-09 by modality; TOTAL uses the official goal and modalities are analytical slices.
    for idx,period in enumerate(periods[:12],2): ws.cell(18,idx).value=period
    for c,v in enumerate(tail,14): ws.cell(18,c).value=v
    modalities=[('TOTAL',None),('Presencial','Presencial'),('EAD','EAD'),('Semipresencial','Semipresencial')]
    for row,(label,mod) in enumerate(modalities,19):
        ws.cell(row,1).value=label
        for idx,period in enumerate(periods[:12],2):
            value,_,_=_cost_metrics(snapshot,period,modality=mod)
            ws.cell(row,idx).value=value; ws.cell(row,idx).number_format='R$ #,##0.00'
        rv,_,_=_cost_metrics(snapshot,ref,modality=mod); cv,_,_=_cost_metrics(snapshot,comp,modality=mod)
        goal=_goal_for(snapshot,'DADM-09',ref) if mod is None else None
        status=_lower_status(rv,goal) if mod is None else ('Sem dados' if rv is None else 'Recorte analítico')
        vals=[rv,cv,_pct_change(rv,cv),_goal_value(goal),status]
        for c,v in enumerate(vals,14): ws.cell(row,c).value=v
        ws.cell(row,14).number_format=ws.cell(row,15).number_format=ws.cell(row,17).number_format='R$ #,##0.00'
        ws.cell(row,16).number_format='+0.00"%";-0.00"%";0.00"%"'
        _set_status_cell(ws.cell(row,18),status)

    # DADM-10 multi-measure matrix; only R$/m² carries the official status.
    for idx,period in enumerate(periods[:12],2): ws.cell(31,idx).value=period
    for c,v in enumerate(tail,14): ws.cell(31,c).value=v
    measures=[('R$/m²','custo_m2','currency'),('R$/aluno','custo_aluno','currency'),('Despesa','despesa','currency'),('Área em uso (m²)','area_m2','number2'),('Alunos ativos','alunos','integer')]
    for row,(label,key,fmt) in enumerate(measures,32):
        ws.cell(row,1).value=label
        for idx,period in enumerate(periods[:12],2):
            value=_infra_metrics(snapshot,period).get(key); ws.cell(row,idx).value=value
            ws.cell(row,idx).number_format={'currency':'R$ #,##0.00','number2':'#,##0.00','integer':'#,##0'}[fmt]
        rm=_infra_metrics(snapshot,ref); cm=_infra_metrics(snapshot,comp); rv=rm.get(key); cv=cm.get(key)
        goal=_goal_for(snapshot,'DADM-10',ref) if key=='custo_m2' else None
        status=_lower_status(rv,goal) if key=='custo_m2' else ('Sem dados' if rv is None else 'Leitura auxiliar')
        vals=[rv,cv,_pct_change(rv,cv),_goal_value(goal),status]
        for c,v in enumerate(vals,14): ws.cell(row,c).value=v
        ws.cell(row,14).number_format=ws.cell(row,15).number_format=ws.cell(row,17).number_format={'currency':'R$ #,##0.00','number2':'#,##0.00','integer':'#,##0'}[fmt]
        ws.cell(row,16).number_format='+0.00"%";-0.00"%";0.00"%"'
        _set_status_cell(ws.cell(row,18),status)

    # Population by course, matching the DTNH lower matrix block.
    for idx,period in enumerate(periods[:12],3): ws.cell(44,idx).value=period
    for row,course in enumerate(snapshot.get('cursos_academicos',[])[:8],45):
        ws.cell(row,1).value=f"{course.get('diretoria')} · {course.get('curso')}"
        ws.cell(row,2).value=course.get('modalidade')
        for idx,period in enumerate(periods[:12],3):
            ws.cell(row,idx).value=_active_students(snapshot,period,course_id=course.get('id'))
            ws.cell(row,idx).number_format='#,##0'

    course_rows = list(snapshot.get('cursos_academicos', []) or [])
    ws = wb['CURSOS']
    course_end = _dynamic_end_row('Catálogo de cursos da DADM', 5, len(course_rows), 404)
    _ensure_sheet_rows(ws, course_end, style_row=404, start_col=1, end_col=6)
    _clear_values(ws, 5, course_end, 1, 6)
    _set_validation_ref(ws, 'B5:B13', f'B5:B{course_end}')
    for row, item in enumerate(course_rows, 5):
        vals = [
            item.get('diretoria'), item.get('curso'), item.get('modalidade'), 'S',
            'População + evasão',
            'A população é produzida pela diretoria acadêmica e consumida pela DADM.',
        ]
        for col, val in enumerate(vals, 1):
            ws.cell(row, col).value = val
        ws.row_dimensions[row].height = 30


def _populate_governance(wb, repo: DatabaseRepository, snapshot: dict) -> None:
    goals = list(repo.list_goals() or [])
    ws = wb['METAS']
    goal_end = _dynamic_end_row('Metas da DADM', 5, len(goals), 404)
    _ensure_sheet_rows(ws, goal_end, style_row=404, start_col=1, end_col=8)
    _clear_values(ws, 5, goal_end, 1, 8)
    _set_validation_ref(ws, 'A5:A404', f'A5:A{goal_end}')
    for row in range(5, goal_end + 1):
        ws.cell(row, 4).value = f'=IFERROR(VALUE(SUBSTITUTE($C{row},"-","")),0)'
    for row, item in enumerate(goals, 5):
        vals = [
            item['indicador'], item['recorte'], item['vigencia'], None, item['meta'],
            item['atencao'], item['limite_superior'], item['justificativa'],
        ]
        for col, val in enumerate(vals, 1):
            if col != 4:
                ws.cell(row, col).value = val
        unit = item.get('unidade', '')
        if unit == '%':
            for col in (5, 6, 7):
                ws.cell(row, col).number_format = '0.00"%"'
        elif str(unit).startswith('R$'):
            for col in (5, 6, 7):
                ws.cell(row, col).number_format = 'R$ #,##0.00'

    actions = list(repo.list_actions() or [])
    ws = wb['PLANO_DE_ACAO']
    action_end = _dynamic_end_row('Planos de ação da DADM', 9, len(actions), 308)
    _ensure_sheet_rows(ws, action_end, style_row=308, start_col=1, end_col=14)
    _clear_values(ws, 9, action_end, 1, 14)
    _set_validation_ref(ws, 'C9:C308', f'C9:C{action_end}')
    _set_validation_ref(ws, 'M9:M308', f'M9:M{action_end}')
    ws['D5'] = f'=COUNTIFS($B$9:$B${action_end},PARAMETROS!$B$7)'
    fields = [
        'numero', 'mes', 'indicador', 'recorte', 'resultado', 'meta', 'problema',
        'causa', 'acao', 'responsavel', 'prazo', 'meta_acao', 'status',
    ]
    for row in range(9, action_end + 1):
        ws.cell(row, 14).value = (
            f'=IF($K{row}="","",IF(OR($M{row}="Concluído",'
            f'$M{row}="Cancelado"),"—",DATEVALUE($K{row})-TODAY()))'
        )
    for row, item in enumerate(actions, 9):
        for col, field in enumerate(fields, 1):
            ws.cell(row, col).value = item.get(field)
        _set_status_cell(ws.cell(row, 13), item.get('status'))

    # The definition sheet is static and already follows the DTNH 10-column schema.
    ws = wb['QUALIDADE E GOVERNANÇA']
    checks = [('evasao', 6), ('custos', 7), ('alunos', 8), ('infraestrutura', 9)]
    for key, row in checks:
        errors = sum(
            1 for item in snapshot.get(key, [])
            if item.get('validacao') not in (None, '', 'OK')
        )
        ws.cell(row, 3).value = errors
        ws.cell(row, 3).alignment = Alignment(horizontal='center', vertical='center')
        ws.cell(row, 3).fill = PatternFill(
            'solid', fgColor='E8F1ED' if errors == 0 else 'FDE7E7'
        )


def _export_dadm(repo: DatabaseRepository) -> BytesIO:
    if not DADM_TEMPLATE.exists():
        raise FileNotFoundError('Template do Excel da DADM não encontrado.')
    snapshot=repo.snapshot_dadm()
    # Resolve the latest closing first, then make the prior operational month the
    # explicit comparison. This mirrors the ready-to-use DTNH export instead of
    # leaving the comparison selector blank on first opening.
    preliminary=build_dadm_dashboard(snapshot,window_months=12)
    operational_periods=_periods_for_export(preliminary)
    reference=operational_periods[-1] if operational_periods else preliminary.get('contexto',{}).get('referencia')
    comparison=operational_periods[-2] if len(operational_periods)>1 else None
    dashboard=build_dadm_dashboard(
        snapshot,
        referencia=reference,
        comparacao=comparison,
        inicio=operational_periods[0] if operational_periods else None,
        fim=reference,
    )
    wb=load_workbook(DADM_TEMPLATE)

    _write_panel_context(wb,dashboard,snapshot)
    _populate_dadm_bases(wb,snapshot)
    _populate_calc(wb,dashboard,snapshot)
    _populate_matrix_and_courses(wb,dashboard,snapshot)
    _populate_governance(wb,repo,snapshot)

    periods=_periods_for_export(dashboard)
    count=min(12,len(periods))
    ws_panel=wb['PAINEL']; ws_panel._charts=[]
    if count:
        ref_wb=load_workbook(DTNH_STYLE_REFERENCE if DTNH_STYLE_REFERENCE.exists() else DTNH_TEMPLATE, data_only=False)
        try:
            reference_charts=ref_wb['PAINEL']._charts
            _add_dtnh_style_chart(ws_panel,reference_charts[0],'DADM-01 — mês a mês na janela','D',21,22,21+count,'0.00"%"')
            _add_dtnh_style_chart(ws_panel,reference_charts[1],'DADM-09 — custo administrativo por aluno mês a mês','D',54,55,54+count,'R$ #,##0')
            _add_dtnh_style_chart(ws_panel,reference_charts[2],'DADM-10 — infraestrutura por m² mês a mês','E',87,88,87+count,'R$ #,##0')
        finally:
            ref_wb.close()

    try:
        wb.calculation.fullCalcOnLoad=True
        wb.calculation.forceFullCalc=True
        wb.calculation.calcMode='auto'
    except Exception:
        pass
    buffer=BytesIO(); wb.save(buffer); wb.close(); buffer.seek(0); return buffer


def _dpe_status_fill(status: str | None) -> PatternFill:
    return _status_fill(status)


def _populate_dpe_base_sheet(ws, rows: list[dict], kind: str) -> None:
    # Preserve the template design while extending the source base to every real row.
    # The workbook is aborted explicitly only at Excel's physical row limit.
    rows = list(rows or [])
    labels = {
        'resultado': 'Base DPE-01 Resultado Operacional',
        'orcamento': 'Base DPE-04 Execução Orçamentária',
        'caixa': 'Base DPE-05 Caixa',
    }
    end_col = 10 if kind in {'resultado', 'orcamento'} else 9
    end_row = _dynamic_end_row(labels.get(kind, 'Base DPE'), 5, len(rows), 1200)
    _ensure_sheet_rows(ws, end_row, style_row=1200, start_col=1, end_col=end_col)

    if kind == 'resultado':
        input_cols = (1, 2, 3, 4, 5, 9, 10)
        _set_validation_ref(ws, 'B5:B1200', f'B5:B{end_row}')
        for row in range(5, end_row + 1):
            for col in input_cols:
                ws.cell(row, col).value = None
        for row, item in enumerate(reversed(rows), 5):
            ws.cell(row, 1).value = item.get('periodo')
            ws.cell(row, 2).value = item.get('tipo_recorte')
            ws.cell(row, 3).value = item.get('recorte')
            ws.cell(row, 4).value = item.get('receita_liquida')
            ws.cell(row, 5).value = item.get('despesa_total')
            ws.cell(row, 6).value = f'=IF(OR(D{row}="",E{row}=""),"",D{row}-E{row})'
            ws.cell(row, 7).value = f'=IFERROR(F{row}/D{row},"")'
            ws.cell(row, 8).value = (
                f'=IF(AND(A{row}<>"",B{row}<>"",C{row}<>"",D{row}>0,E{row}>=0),'
                '"OK","REVISAR")'
            )
            ws.cell(row, 9).value = item.get('data_lancamento')
            ws.cell(row, 10).value = item.get('lancado_por')
            for col in (4, 5, 6):
                ws.cell(row, col).number_format = 'R$ #,##0.00'
            ws.cell(row, 7).number_format = '0.00%'
    elif kind == 'orcamento':
        input_cols = (1, 2, 3, 4, 5, 9, 10)
        for row in range(5, end_row + 1):
            for col in input_cols:
                ws.cell(row, col).value = None
        for row, item in enumerate(reversed(rows), 5):
            ws.cell(row, 1).value = item.get('periodo')
            ws.cell(row, 2).value = item.get('unidade')
            ws.cell(row, 3).value = item.get('centro_custo')
            ws.cell(row, 4).value = item.get('despesa_orcada')
            ws.cell(row, 5).value = item.get('despesa_realizada')
            ws.cell(row, 6).value = f'=IFERROR(E{row}/D{row},"")'
            ws.cell(row, 7).value = f'=IF(OR(D{row}="",E{row}=""),"",E{row}-D{row})'
            ws.cell(row, 8).value = (
                f'=IF(AND(A{row}<>"",B{row}<>"",C{row}<>"",D{row}>0,E{row}>=0),'
                '"OK","REVISAR")'
            )
            ws.cell(row, 9).value = item.get('data_lancamento')
            ws.cell(row, 10).value = item.get('lancado_por')
            for col in (4, 5, 7):
                ws.cell(row, col).number_format = 'R$ #,##0.00'
            ws.cell(row, 6).number_format = '0.00%'
    else:
        input_cols = (1, 2, 3, 4, 5, 8, 9)
        _set_validation_ref(ws, 'C5:C1200', f'C5:C{end_row}')
        for row in range(5, end_row + 1):
            for col in input_cols:
                ws.cell(row, col).value = None
        for row, item in enumerate(reversed(rows), 5):
            ws.cell(row, 1).value = item.get('periodo')
            ws.cell(row, 2).value = item.get('conta')
            ws.cell(row, 3).value = item.get('tipo_movimento')
            ws.cell(row, 4).value = item.get('natureza')
            ws.cell(row, 5).value = item.get('valor')
            ws.cell(row, 6).value = (
                f'=IF(E{row}="","",IF(C{row}="Entrada",E{row},-E{row}))'
            )
            ws.cell(row, 7).value = (
                f'=IF(AND(A{row}<>"",B{row}<>"",OR(C{row}="Entrada",'
                f'C{row}="Saída"),D{row}<>"",E{row}>0),"OK","REVISAR")'
            )
            ws.cell(row, 8).value = item.get('data_lancamento')
            ws.cell(row, 9).value = item.get('lancado_por')
            for col in (5, 6):
                ws.cell(row, col).number_format = 'R$ #,##0.00'


def _export_dpe(repo: DatabaseRepository) -> BytesIO:
    if not DPE_TEMPLATE.exists():
        raise FileNotFoundError('Template do Excel da DPE não encontrado.')
    snapshot=repo.snapshot_dpe()
    prelim=build_dpe_dashboard(snapshot,window_months=12)
    periods=prelim.get('periodos',{}).get('mensais',[]) or []
    reference=periods[-1] if periods else None
    comparison=periods[-2] if len(periods)>1 else None
    dashboard=build_dpe_dashboard(snapshot,referencia=reference,comparacao=comparison,inicio=periods[0] if periods else None,fim=reference)
    wb=load_workbook(DPE_TEMPLATE)

    # Parameters reflect the snapshot exported, while remaining editable by the recipient.
    ws=wb['PARAMETROS']
    values={5:'DPE',6:reference or '',7:comparison or '',8:(periods[0] if periods else ''),9:(reference or ''),10:'Institucional',11:'TOTAL',12:'(todas)',13:'(todos)',14:'(todas)',15:'(todas)'}
    for r,v in values.items(): ws.cell(r,2).value=v

    _populate_dpe_base_sheet(wb['DPE-01 RESULTADO'],snapshot.get('resultado',[]),'resultado')
    _populate_dpe_base_sheet(wb['DPE-04 ORCAMENTO'],snapshot.get('orcamento',[]),'orcamento')
    _populate_dpe_base_sheet(wb['DPE-05 CAIXA'],snapshot.get('caixa',[]),'caixa')

    # CALC is filled with values calculated server-side, so charts work immediately even
    # when the spreadsheet engine has not yet recalculated formula cells in the raw bases.
    ws=wb['CALC']; _clear_values(ws,5,40,1,12)
    series_result={x.get('periodo'):x for x in dashboard.get('series',{}).get('resultado',[])}
    series_budget={x.get('periodo'):x for x in dashboard.get('series',{}).get('orcamento',[])}
    series_cash={x.get('periodo'):x for x in dashboard.get('series',{}).get('caixa',[])}
    for row,period in enumerate(periods[-12:],5):
        rr=series_result.get(period,{}) ; bb=series_budget.get(period,{}) ; cc=series_cash.get(period,{})
        vals=[period,rr.get('valor'),rr.get('meta'),rr.get('status'),bb.get('valor'),bb.get('meta'),bb.get('limite_superior'),bb.get('status'),cc.get('valor'),cc.get('saldo_acumulado'),cc.get('meta'),cc.get('status')]
        for c,v in enumerate(vals,1): ws.cell(row,c).value=v
        for c in (2,3,5,6,7): ws.cell(row,c).number_format='0.00"%"'
        for c in (9,10,11): ws.cell(row,c).number_format='R$ #,##0.00'
        for c in (4,8,12): _set_status_cell(ws.cell(row,c),ws.cell(row,c).value)

    # Executive cards and details.
    ws=wb['PAINEL']
    cards=dashboard.get('cards',{})
    card_specs=[
        ('resultado',7,['valor','comparacao','meta','status','resultado','receita'],['0.00"%"','0.00"%"','0.00"%"',None,'R$ #,##0.00','R$ #,##0.00']),
        ('orcamento',37,['valor','comparacao','meta','status','orcado','realizado'],['0.00"%"','0.00"%"','0.00"%"',None,'R$ #,##0.00','R$ #,##0.00']),
        ('caixa',67,['valor','comparacao','meta','status','entradas','saidas'],['R$ #,##0.00','R$ #,##0.00','R$ #,##0.00',None,'R$ #,##0.00','R$ #,##0.00']),
    ]
    for key,row,fields,formats in card_specs:
        card=cards.get(key,{})
        for c,(field,fmt) in enumerate(zip(fields,formats),1):
            ws.cell(row,c).value=card.get(field)
            if fmt: ws.cell(row,c).number_format=fmt
        _set_status_cell(ws.cell(row,4),card.get('status'))
    # Detail tables
    for r in range(24,31): _clear_values(ws,r,r,1,4)
    for r,item in enumerate(dashboard.get('comparacoes',{}).get('resultado_por_recorte',[])[:7],24):
        vals=[item.get('label'),item.get('valor'),'', 'Recorte analítico']
        for c,v in enumerate(vals,1): ws.cell(r,c).value=v
        ws.cell(r,2).number_format='0.00"%"'
    for r in range(54,61): _clear_values(ws,r,r,1,4)
    for r,item in enumerate(dashboard.get('comparacoes',{}).get('orcamento_por_diretoria',[])[:7],54):
        vals=[item.get('label'),item.get('valor'),item.get('realizado'),'']
        for c,v in enumerate(vals,1): ws.cell(r,c).value=v
        ws.cell(r,2).number_format='0.00"%"';ws.cell(r,3).number_format='R$ #,##0.00'
    for r in range(84,91): _clear_values(ws,r,r,1,4)
    for r,item in enumerate(dashboard.get('comparacoes',{}).get('caixa_por_conta',[])[:7],84):
        vals=[item.get('label'),item.get('valor'),item.get('saidas'),'']
        for c,v in enumerate(vals,1): ws.cell(r,c).value=v
        ws.cell(r,2).number_format=ws.cell(r,3).number_format='R$ #,##0.00'

    # Matrix: TOTAL/selected default for each KPI plus useful detailed rows.
    ws=wb['MATRIZ']
    for header_row in (5,21,37):
        for idx,p in enumerate(periods[-12:],2): ws.cell(header_row,idx).value=p
        for c,v in enumerate(['Referência','Comparação','Variação','Meta','Status'],14): ws.cell(header_row,c).value=v
    # DPE-01 current series in first line, then recortes latest.
    result_series=dashboard.get('series',{}).get('resultado',[])
    ws.cell(6,1).value='Institucional · TOTAL'
    byp={x.get('periodo'):x for x in result_series}
    for idx,p in enumerate(periods[-12:],2): ws.cell(6,idx).value=byp.get(p,{}).get('valor'); ws.cell(6,idx).number_format='0.00"%"'
    rc=cards.get('resultado',{}); vals=[rc.get('valor'),rc.get('comparacao'),None if rc.get('valor') is None or rc.get('comparacao') is None else round(rc['valor']-rc['comparacao'],2),rc.get('meta'),rc.get('status')]
    for c,v in enumerate(vals,14): ws.cell(6,c).value=v
    for c in (14,15,16,17): ws.cell(6,c).number_format='0.00"%"'
    _set_status_cell(ws.cell(6,18),rc.get('status'))
    for r,item in enumerate(dashboard.get('comparacoes',{}).get('resultado_por_recorte',[])[:7],7): ws.cell(r,1).value=item.get('label'); ws.cell(r,14).value=item.get('valor'); ws.cell(r,14).number_format='0.00"%"'; ws.cell(r,18).value='Recorte analítico'

    budget_series=dashboard.get('series',{}).get('orcamento',[]); byp={x.get('periodo'):x for x in budget_series}; ws.cell(22,1).value='TOTAL'
    for idx,p in enumerate(periods[-12:],2): ws.cell(22,idx).value=byp.get(p,{}).get('valor'); ws.cell(22,idx).number_format='0.00"%"'
    bc=cards.get('orcamento',{}); vals=[bc.get('valor'),bc.get('comparacao'),None if bc.get('valor') is None or bc.get('comparacao') is None else round(bc['valor']-bc['comparacao'],2),bc.get('meta'),bc.get('status')]
    for c,v in enumerate(vals,14): ws.cell(22,c).value=v
    for c in (14,15,16,17): ws.cell(22,c).number_format='0.00"%"'
    _set_status_cell(ws.cell(22,18),bc.get('status'))
    for r,item in enumerate(dashboard.get('comparacoes',{}).get('orcamento_por_diretoria',[])[:7],23): ws.cell(r,1).value=item.get('label'); ws.cell(r,14).value=item.get('valor'); ws.cell(r,14).number_format='0.00"%"'; ws.cell(r,18).value='Recorte analítico'

    cash_series=dashboard.get('series',{}).get('caixa',[]); byp={x.get('periodo'):x for x in cash_series}; ws.cell(38,1).value='Saldo mensal'
    for idx,p in enumerate(periods[-12:],2): ws.cell(38,idx).value=byp.get(p,{}).get('valor'); ws.cell(38,idx).number_format='R$ #,##0.00'
    cc=cards.get('caixa',{}); vals=[cc.get('valor'),cc.get('comparacao'),None if cc.get('valor') is None or cc.get('comparacao') is None else round(cc['valor']-cc['comparacao'],2),cc.get('meta'),cc.get('status')]
    for c,v in enumerate(vals,14): ws.cell(38,c).value=v
    for c in (14,15,16,17): ws.cell(38,c).number_format='R$ #,##0.00'
    _set_status_cell(ws.cell(38,18),cc.get('status'))
    for r,(label,key) in enumerate([('Entradas','entradas'),('Saídas','saidas'),('Saldo acumulado','saldo_acumulado')],39):
        ws.cell(r,1).value=label
        for idx,p in enumerate(periods[-12:],2): ws.cell(r,idx).value=byp.get(p,{}).get(key); ws.cell(r,idx).number_format='R$ #,##0.00'

    # Goals and action plans. All records are exported; legacy row counts are
    # expanded dynamically instead of acting as invisible caps.
    goals = list(repo.list_goals() or [])
    ws = wb['METAS']
    goal_end = _dynamic_end_row('Metas da DPE', 5, len(goals), 404)
    _ensure_sheet_rows(ws, goal_end, style_row=404, start_col=1, end_col=8)
    _clear_values(ws, 5, goal_end, 1, 8)
    for row, item in enumerate(goals, 5):
        vals = [
            item.get('indicador'), item.get('recorte'), item.get('vigencia'), None,
            item.get('meta'), item.get('atencao'), item.get('limite_superior'),
            item.get('justificativa'),
        ]
        for col, val in enumerate(vals, 1):
            if col != 4:
                ws.cell(row, col).value = val
        if item.get('unidade') == '%':
            for col in (5, 6, 7):
                ws.cell(row, col).number_format = '0.00"%"'
        elif str(item.get('unidade') or '').startswith('R$'):
            for col in (5, 6, 7):
                ws.cell(row, col).number_format = 'R$ #,##0.00'

    actions = list(repo.list_actions() or [])
    ws = wb['PLANO_DE_ACAO']
    action_end = _dynamic_end_row('Planos de ação da DPE', 9, len(actions), 308)
    _ensure_sheet_rows(ws, action_end, style_row=308, start_col=1, end_col=13)
    _clear_values(ws, 9, action_end, 1, 13)
    fields = [
        'numero', 'mes', 'indicador', 'recorte', 'resultado', 'meta', 'problema',
        'causa', 'acao', 'responsavel', 'prazo', 'meta_acao', 'status',
    ]
    for row, item in enumerate(actions, 9):
        for col, field in enumerate(fields, 1):
            ws.cell(row, col).value = item.get(field)
        _set_status_cell(ws.cell(row, 13), item.get('status'))

    # Support dimensions / quality.
    refs = repo.dpe_reference_options()
    support_specs = [
        (3, 'diretorias'), (4, 'centros_custo'), (5, 'contas'),
        (6, 'naturezas'), (7, 'recortes'),
    ]
    support_lengths = [len(refs.get(key, []) or []) for _, key in support_specs]
    support_end = _dynamic_end_row(
        'Listas de apoio da DPE', 5, max(support_lengths, default=0), 100
    )
    ws = wb['LISTAS DE APOIO']
    _ensure_sheet_rows(ws, support_end, style_row=100, start_col=1, end_col=8)
    _clear_values(ws, 5, support_end, 3, 8)
    for row, directorate in enumerate(refs.get('diretorias', []) or [], 5):
        ws.cell(row, 3).value = directorate.get('code')
    for col, key in support_specs[1:]:
        for row, value in enumerate(refs.get(key, []) or [], 5):
            ws.cell(row, col).value = value

    dimension_rows = []
    for kind, key in [
        ('Tipo de recorte', 'tipos_recorte'), ('Diretoria', 'diretorias'),
        ('Centro de custo', 'centros_custo'), ('Conta', 'contas'),
        ('Natureza', 'naturezas'),
    ]:
        for value in refs.get(key, []) or []:
            label = value.get('code') if isinstance(value, dict) else value
            description = value.get('name') if isinstance(value, dict) else ''
            dimension_rows.append((kind, label, description, 'Banco/Data UNIVC', 'S'))
    ws = wb['DIMENSOES']
    dimension_end = _dynamic_end_row(
        'Dimensões da DPE', 5, len(dimension_rows), 404
    )
    _ensure_sheet_rows(ws, dimension_end, style_row=404, start_col=1, end_col=5)
    _clear_values(ws, 5, dimension_end, 1, 5)
    for row, values in enumerate(dimension_rows, 5):
        for col, value in enumerate(values, 1):
            ws.cell(row, col).value = value

    ws = wb['QUALIDADE E GOVERNANÇA']
    quality = cards.get('qualidade', {})
    for row, key in [(6, 'resultado'), (7, 'orcamento'), (8, 'caixa')]:
        ws.cell(row, 3).value = quality.get(key, 0)
        ws.cell(row, 3).fill = PatternFill(
            'solid', fgColor='E8F1ED' if not quality.get(key) else 'FDE7E7'
        )
    ws.cell(9, 3).value = (
        0 if any(
            goal.get('indicador') == 'DPE-04'
            and goal.get('limite_superior') is not None
            for goal in goals
        ) else 1
    )
    ws.cell(10, 3).value = 1  # pending integration with payroll reserve condition.

    try:
        wb.calculation.fullCalcOnLoad=True;wb.calculation.forceFullCalc=True;wb.calculation.calcMode='auto'
    except Exception: pass
    buffer=BytesIO();wb.save(buffer);wb.close();buffer.seek(0);return buffer


def export_academic_interactive_excel(
    repo: DatabaseRepository,
    *,
    reference: str | None = None,
    comparison: str | None = None,
    course: str | None = None,
    discipline: str | None = None,
    window_periods: int | str | None = None,
) -> BytesIO:
    """Generate the academic Excel Interativo V3 beta with full authorized history."""
    from academic_excel_v3_builder import build_academic_interactive_workbook_bytes

    return build_academic_interactive_workbook_bytes(
        repo,
        reference=reference,
        comparison=comparison,
        course=course,
        discipline=discipline,
        window_periods=window_periods,
    )

def export_formatted_excel(
    repo: DatabaseRepository,
    *,
    granularity: str | None = None,
    reference: str | None = None,
    comparison: str | None = None,
    course: str | None = None,
    discipline: str | None = None,
    window_periods: int | str | None = None,
) -> BytesIO:
    """Generate the selected directorate snapshot entirely in memory.

    For DTNH/DCS, the optional arguments preserve the same academic context
    selected in the web dashboard when exporting the institutional workbook.
    """
    if repo.directorate_code == 'DADM':
        return _export_dadm(repo)
    if repo.directorate_code == 'DPE':
        return _export_dpe(repo)
    return _export_academic(
        repo,
        granularity=granularity,
        reference=reference,
        comparison=comparison,
        course=course,
        discipline=discipline,
        window_periods=window_periods,
    )
