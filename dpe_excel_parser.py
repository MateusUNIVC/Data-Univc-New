from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from management_catalog import indicator_spec, validate_period

BASE_HEADER_ROW = 4
BASE_DATA_START = 5


class DPEExcelImportError(ValueError):
    def __init__(self, message: str, errors: list[dict[str, Any]] | None = None):
        super().__init__(message)
        self.errors = errors or []


@dataclass(frozen=True)
class SheetLayout:
    indicator_code: str
    sheet_name: str
    dimensions: dict[str, str]
    fields: dict[str, str]
    notes: str = "observacoes"
    source: str = "fonte referencia"
    validated: str = "validado"


LAYOUTS: dict[str, SheetLayout] = {
    "DPE-01": SheetLayout(
        "DPE-01",
        "BASE DPE-01",
        {
            "course": "curso",
            "academic_directorate": "diretoria academica",
        },
        {
            "net_revenue": "receita liquida",
            "faculty_cost": "custo docente",
            "coordination_cost": "custo de coordenacao",
            "other_direct_cost": "outros custos diretos",
            "indirect_cost": "rateio indireto",
            "weekly_teaching_hours": "horas aula semanais",
            "teacher_count": "n de docentes",
        },
    ),
    "DPE-02": SheetLayout(
        "DPE-02",
        "BASE DPE-02",
        {
            "cost_center": "centro de custo",
            "expense_nature": "natureza da despesa",
        },
        {
            "net_revenue": "receita liquida",
            "personnel_expense": "despesa de pessoal",
            "operational_expense": "despesa operacional",
            "administrative_expense": "despesa administrativa",
            "financial_expense": "despesa financeira",
            "capex": "investimento em imobilizado",
        },
    ),
    "DPE-03": SheetLayout(
        "DPE-03",
        "BASE DPE-03",
        {
            "category": "categoria",
            "nature": "natureza",
        },
        {
            "faculty_payroll": "folha docente completa",
            "administrative_payroll": "folha administrativa completa",
            "gross_salaries": "salarios brutos",
            "charges": "encargos",
            "provisions": "provisoes",
            "net_revenue": "receita liquida do mes",
        },
    ),
}


def _norm(value: Any) -> str:
    raw = str(value or "").replace("º", "").replace("ª", "").replace("°", "")
    text = unicodedata.normalize("NFKD", raw)
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = text.casefold().strip()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def _bool(value: Any) -> bool:
    normalized = _norm(value)
    if normalized in {"", "sim", "s", "1", "true", "verdadeiro", "validado"}:
        return True
    if normalized in {"nao", "n", "0", "false", "falso", "pendente"}:
        return False
    raise ValueError("Use Sim ou Não na coluna Validado.")


def _number(value: Any, *, required: bool, label: str) -> float | int | None:
    if value in (None, ""):
        if required:
            raise ValueError(f"Informe {label}.")
        return None
    if isinstance(value, str) and value.startswith("="):
        raise ValueError(f"{label} precisa ser um valor, não uma fórmula.")
    if isinstance(value, str):
        text = value.strip().replace("R$", "").replace(" ", "")
        if "," in text and "." in text:
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", ".")
        value = text
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{label} possui valor numérico inválido.") from exc
    if number != number or abs(number) == float("inf"):
        raise ValueError(f"{label} possui valor numérico inválido.")
    return number


def _headers(ws) -> dict[str, int]:
    result: dict[str, int] = {}
    for cell in ws[BASE_HEADER_ROW]:
        normalized = _norm(cell.value)
        if normalized:
            result[normalized] = cell.column
    return result


def _required_column(headers: dict[str, int], name: str, *, sheet: str) -> int:
    normalized = _norm(name)
    if normalized not in headers:
        raise DPEExcelImportError(
            f'A aba "{sheet}" não possui a coluna obrigatória "{name}".',
            [{"sheet": sheet, "row": BASE_HEADER_ROW, "field": name, "error": "Cabeçalho ausente."}],
        )
    return headers[normalized]


def parse_dpe_workbook(path: str | Path, *, indicator_code: str | None = None) -> dict[str, Any]:
    """Lê uma planilha DPE e devolve payloads prontos para validação do backend.

    O parser ignora as colunas calculadas e sempre recalcula os indicadores no
    servidor a partir dos componentes. Nenhuma correção silenciosa é aplicada.
    """
    workbook = load_workbook(path, data_only=False, read_only=False)
    codes: Iterable[str]
    if indicator_code:
        code = indicator_spec("DPE", indicator_code)["code"]
        codes = [code]
    else:
        codes = LAYOUTS.keys()

    payloads: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    sheets_read: list[str] = []

    for code in codes:
        layout = LAYOUTS[code]
        sheet_name = layout.sheet_name
        if indicator_code and "DADOS" in workbook.sheetnames:
            sheet_name = "DADOS"
        if sheet_name not in workbook.sheetnames:
            if indicator_code:
                errors.append({"sheet": sheet_name, "row": None, "field": None, "error": "Aba não encontrada."})
            continue
        ws = workbook[sheet_name]
        sheets_read.append(sheet_name)
        headers = _headers(ws)
        try:
            period_col = _required_column(headers, "Período", sheet=sheet_name)
            dimension_cols = {}
            for key, label in layout.dimensions.items():
                if code == "DPE-01" and key == "academic_directorate":
                    # O curso é a entrada oficial. A diretoria é inferida pelo catálogo DTNH/DCS no router.
                    dimension_cols[key] = headers.get(_norm(label))
                else:
                    dimension_cols[key] = _required_column(headers, label, sheet=sheet_name)
            field_cols = {
                key: _required_column(headers, label, sheet=sheet_name)
                for key, label in layout.fields.items()
            }
            notes_col = headers.get(_norm("Observações"))
            source_col = headers.get(_norm("Fonte / referência"))
            validated_col = headers.get(_norm("Validado"))
        except DPEExcelImportError as exc:
            errors.extend(exc.errors)
            continue

        spec = indicator_spec("DPE", code)
        required = {field["key"]: bool(field.get("required")) for field in spec.get("fields") or []}
        labels = {field["key"]: field["label"] for field in spec.get("fields") or []}

        for row_index in range(BASE_DATA_START, ws.max_row + 1):
            raw_period = ws.cell(row_index, period_col).value
            has_any_component = any(ws.cell(row_index, col).value not in (None, "") for col in field_cols.values())
            has_any_dimension = any(col is not None and ws.cell(row_index, col).value not in (None, "") for col in dimension_cols.values())
            if raw_period in (None, "") and not has_any_component and not has_any_dimension:
                continue
            row_errors: list[dict[str, Any]] = []
            try:
                period = validate_period("DPE", raw_period)
            except Exception as exc:
                row_errors.append({"sheet": sheet_name, "row": row_index, "field": "Período", "error": str(exc)})
                period = None

            dimensions = {
                key: str(ws.cell(row_index, col).value or "").strip()
                for key, col in dimension_cols.items()
                if col is not None and str(ws.cell(row_index, col).value or "").strip()
            }
            if code == "DPE-01" and not dimensions.get("course"):
                row_errors.append({"sheet": sheet_name, "row": row_index, "field": "Curso", "error": "O curso é obrigatório no DPE-01."})

            values: dict[str, float | int | None] = {}
            for key, col in field_cols.items():
                try:
                    values[key] = _number(
                        ws.cell(row_index, col).value,
                        required=required.get(key, False),
                        label=labels.get(key, key),
                    )
                except ValueError as exc:
                    row_errors.append({"sheet": sheet_name, "row": row_index, "field": labels.get(key, key), "error": str(exc)})

            try:
                validated = _bool(ws.cell(row_index, validated_col).value) if validated_col else True
            except ValueError as exc:
                row_errors.append({"sheet": sheet_name, "row": row_index, "field": "Validado", "error": str(exc)})
                validated = True

            if row_errors:
                errors.extend(row_errors)
                continue
            payloads.append({
                "indicator_code": code,
                "period": period,
                "dimensions": dimensions,
                "values": values,
                "notes": str(ws.cell(row_index, notes_col).value or "").strip() if notes_col else None,
                "source_reference": str(ws.cell(row_index, source_col).value or "").strip() if source_col else None,
                "validated": validated,
                "_source": {"sheet": sheet_name, "row": row_index},
            })

    if not sheets_read:
        expected = ", ".join(LAYOUTS[code].sheet_name for code in codes)
        raise DPEExcelImportError(
            f"Nenhuma base DPE foi encontrada. Abas esperadas: {expected}.",
            [{"sheet": None, "row": None, "field": None, "error": "Estrutura DPE não encontrada."}],
        )
    if errors:
        raise DPEExcelImportError(
            f"A importação possui {len(errors)} inconsistência(s). Nenhuma linha foi gravada.",
            errors,
        )
    if not payloads:
        raise DPEExcelImportError(
            "A planilha não contém linhas preenchidas para importação.",
            [{"sheet": ", ".join(sheets_read), "row": None, "field": None, "error": "Nenhuma linha preenchida."}],
        )
    return {"payloads": payloads, "sheets": sheets_read, "total": len(payloads)}
