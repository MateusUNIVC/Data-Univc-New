from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import DataPoint
from openpyxl.chart.series import SeriesLabel
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

from academic_analytics import (
    _academic_goal,
    build_academic_dashboard,
    semester_sort_key,
    summarize_nps_semesters,
    to_semester,
)
from analytics import goal_info, status_for
from release_info import APP_VERSION
from security import DirectorateScope
from survey_repository import SurveyRepository


# Data UNIVC academic reporting palette. Kept deliberately close to the shared
# DTNH/DCS product UI instead of the historical workbook template.
BRAND_950 = "082F23"
BRAND_900 = "0B3D2D"
BRAND_800 = "0E5940"
BRAND_700 = "13704F"
BRAND_600 = "1B825D"
BRAND_100 = "DCEEE5"
BRAND_50 = "F0F8F4"
ACCENT = "A9CB55"
INK_950 = "14211B"
INK_800 = "314139"
INK_600 = "66756E"
LINE = "DCE5E0"
SURFACE_MUTED = "F7FAF8"
WHITE = "FFFFFF"
GREEN = "2E7D32"
YELLOW = "F5A623"
RED = "C62828"
NEUTRAL = "8A9690"
IMPORT_GREEN = "008000"
STATIC_GRAY = "666666"
TEAL = "008080"

THIN_LINE = Side(style="thin", color=LINE)
INTEGER_FMT = "#,##0"
DECIMAL_FMT = "0.00"
NPS_FMT = "0.0"
PCT_FMT = "0.0%"

STATUS_COLORS = {
    "Dentro da meta": "DCEFE2",
    "Atenção": "FFF1CC",
    "Fora da meta": "F9DADA",
    "Sem meta": "ECEFED",
    "Sem dados": "F3F5F4",
}
STATUS_CHART_COLORS = {
    "Dentro da meta": GREEN,
    "Atenção": YELLOW,
    "Fora da meta": RED,
    "Sem meta": NEUTRAL,
    "Sem dados": NEUTRAL,
}


@dataclass
class AcademicReportContext:
    directorate: str
    granularity: str = "semestral"
    reference: str | None = None
    comparison: str | None = None
    course: str | None = None
    discipline: str | None = None
    window_periods: int | str | None = None


def _clean_filter(value: str | None, all_label: str) -> str | None:
    text = str(value or "").strip()
    if not text or text in {"(todos)", "(todas)", all_label}:
        return None
    return text


def _scope_for_repo(repo) -> DirectorateScope:
    return DirectorateScope(
        user=repo.ctx,
        directorate_id=repo.directorate_id,
        directorate_code=repo.directorate_code,
        directorate_name=repo.directorate_name,
        can_write=False,
        is_home=repo.ctx.directorate_id == repo.directorate_id,
    )


def _goal_fields(goal: dict[str, Any] | None, code: str, *, course: str | None = None, discipline: str | None = None) -> dict[str, Any]:
    info = goal_info(goal, code, course, discipline)
    if not info:
        return {"meta": None, "atencao": None, "meta_vigencia": None, "meta_recorte": None}
    return {
        "meta": info.get("meta"),
        "atencao": info.get("atencao"),
        "meta_vigencia": info.get("vigencia"),
        "meta_recorte": info.get("recorte"),
    }


def _visible_semesters(dashboard: dict[str, Any]) -> list[str]:
    rows = dashboard.get("series", {}).get("nps_course_semestral") or dashboard.get("series", {}).get("nps_semestral") or []
    values = [str(row.get("periodo")) for row in rows if row.get("periodo")]
    if values:
        return sorted(set(values), key=semester_sort_key)
    context = dashboard.get("contexto", {})
    refs = [to_semester(context.get("referencia")), to_semester(context.get("comparacao"))]
    return sorted({value for value in refs if value}, key=semester_sort_key)


def _filter_semesters(rows: list[dict[str, Any]], semesters: set[str]) -> list[dict[str, Any]]:
    if not semesters:
        return list(rows)
    return [row for row in rows if to_semester(row.get("periodo")) in semesters]


def _course_nps_rows(snapshot: dict[str, Any], dashboard: dict[str, Any], code: str, course_filter: str | None) -> list[dict[str, Any]]:
    semesters = set(_visible_semesters(dashboard))
    out: list[dict[str, Any]] = []
    for row in summarize_nps_semesters(snapshot.get("nps", [])):
        period = to_semester(row.get("periodo"))
        course = str(row.get("curso") or "").strip()
        if semesters and period not in semesters:
            continue
        if course_filter and course != course_filter:
            continue
        goal = _academic_goal(snapshot.get("metas", []), code, period or "", course, None)
        out.append({
            "periodo": period,
            "curso": course,
            "respondentes": int(row.get("respondentes") or 0),
            "promotores": int(row.get("promotores") or 0),
            "neutros": int(row.get("neutros") or 0),
            "detratores": int(row.get("detratores") or 0),
            "fonte": row.get("fonte_nps") or row.get("fonte") or "Base acadêmica",
            **_goal_fields(goal, code, course=course),
            "status": status_for(row.get("nps"), goal, code),
        })
    return sorted(out, key=lambda item: (semester_sort_key(item["periodo"]), item["curso"]))


def _teacher_rows(snapshot: dict[str, Any], dashboard: dict[str, Any], code: str, course_filter: str | None, discipline_filter: str | None) -> list[dict[str, Any]]:
    semesters = set(_visible_semesters(dashboard))
    out: list[dict[str, Any]] = []
    for row in snapshot.get("avaliacao_docente", []):
        period = to_semester(row.get("periodo"))
        course = str(row.get("curso") or "").strip()
        discipline = str(row.get("disciplina") or "").strip()
        if semesters and period not in semesters:
            continue
        if course_filter and course != course_filter:
            continue
        if discipline_filter and discipline != discipline_filter:
            continue
        goal = _academic_goal(snapshot.get("metas", []), code, period or "", course, discipline)
        value = row.get("nota_media") if row.get("nota_media") is not None else row.get("valor")
        out.append({
            "periodo": period,
            "curso": course,
            "disciplina": discipline,
            "respondentes": int(row.get("respondentes") or 0),
            "nota_media": float(value) if value is not None else None,
            **_goal_fields(goal, code, course=course, discipline=discipline),
            "status": status_for(value, goal, code),
        })
    return sorted(out, key=lambda item: (semester_sort_key(item["periodo"]), item["curso"], item["disciplina"]))


def _result_rows(snapshot: dict[str, Any], dashboard: dict[str, Any], code: str, course_filter: str | None, discipline_filter: str | None) -> list[dict[str, Any]]:
    semesters = set(_visible_semesters(dashboard))
    out: list[dict[str, Any]] = []
    for row in snapshot.get("resultados", []):
        period = to_semester(row.get("periodo"))
        course = str(row.get("curso") or "").strip()
        discipline = str(row.get("disciplina") or "").strip()
        if semesters and period not in semesters:
            continue
        if course_filter and course != course_filter:
            continue
        if discipline_filter and discipline != discipline_filter:
            continue
        goal = _academic_goal(snapshot.get("metas", []), code, period or "", course, discipline)
        out.append({
            "periodo": period,
            "curso": course,
            "disciplina": discipline,
            "total_registros": int(row.get("total_registros") or 0),
            "finalizados": int(row.get("finalizados") or 0),
            "aprovados": int(row.get("aprovados") or 0),
            "reprovados_nota": int(row.get("reprovados_nota") or 0),
            "reprovados_falta": int(row.get("reprovados_falta") or 0),
            "reprovados_outro": int(row.get("reprovados_outro") or 0),
            "em_andamento": int(row.get("em_andamento") or 0),
            "media_notas": row.get("media_notas"),
            **_goal_fields(goal, code, course=course, discipline=discipline),
            "status": status_for(row.get("taxa_aprovacao"), goal, code),
        })
    return sorted(out, key=lambda item: (semester_sort_key(item["periodo"]), item["curso"], item["disciplina"]))


def _comparison_rows(snapshot: dict[str, Any], dashboard: dict[str, Any], directorate: str) -> list[dict[str, Any]]:
    reference = dashboard.get("contexto", {}).get("referencia") or ""
    nps_reference = dashboard.get("contexto", {}).get("nps_referencia") or to_semester(reference) or reference
    nps_map = {str(row.get("curso")): row for row in dashboard.get("comparacoes", {}).get("nps", []) if row.get("curso")}
    teacher_map = {str(row.get("curso")): row for row in dashboard.get("comparacoes", {}).get("avaliacao_docente", []) if row.get("curso")}
    approval_map = {str(row.get("curso")): row for row in dashboard.get("comparacoes", {}).get("aprovacao", []) if row.get("curso")}

    course_nps_detail = {
        str(row.get("curso")): row
        for row in summarize_nps_semesters(snapshot.get("nps", []))
        if to_semester(row.get("periodo")) == to_semester(nps_reference)
    }
    teacher_source = [row for row in snapshot.get("avaliacao_docente", []) if to_semester(row.get("periodo")) == to_semester(reference)]
    result_source = [row for row in snapshot.get("resultados", []) if to_semester(row.get("periodo")) == to_semester(reference)]
    courses = sorted(set(nps_map) | set(teacher_map) | set(approval_map))
    out: list[dict[str, Any]] = []
    for course in courses:
        nps_row = course_nps_detail.get(course, {})
        nps_goal = _academic_goal(snapshot.get("metas", []), f"{directorate}-01B", nps_reference, course, None)
        teacher_goal = _academic_goal(snapshot.get("metas", []), f"{directorate}-02", reference, course, None)
        approval_goal = _academic_goal(snapshot.get("metas", []), f"{directorate}-03", reference, course, None)

        teacher_parts = [row for row in teacher_source if str(row.get("curso") or "") == course]
        teacher_resp = sum(int(row.get("respondentes") or 0) for row in teacher_parts)
        teacher_weighted = sum(float(row.get("nota_media") or row.get("valor") or 0) * int(row.get("respondentes") or 0) for row in teacher_parts)

        result_parts = [row for row in result_source if str(row.get("curso") or "") == course]
        finalized = sum(int(row.get("finalizados") or 0) for row in result_parts)
        approved = sum(int(row.get("aprovados") or 0) for row in result_parts)
        out.append({
            "curso": course,
            "nps_respondentes": int(nps_row.get("respondentes") or 0),
            "nps_promotores": int(nps_row.get("promotores") or 0),
            "nps_detratores": int(nps_row.get("detratores") or 0),
            "nps_valor": nps_map.get(course, {}).get("valor"),
            "nps_meta": _goal_fields(nps_goal, f"{directorate}-01B", course=course)["meta"],
            "nps_atencao": _goal_fields(nps_goal, f"{directorate}-01B", course=course)["atencao"],
            "nps_status": status_for(nps_map.get(course, {}).get("valor"), nps_goal, f"{directorate}-01B"),
            "teacher_respondentes": teacher_resp,
            "teacher_weighted_sum": teacher_weighted,
            "teacher_valor": teacher_map.get(course, {}).get("valor"),
            "teacher_meta": _goal_fields(teacher_goal, f"{directorate}-02", course=course)["meta"],
            "teacher_atencao": _goal_fields(teacher_goal, f"{directorate}-02", course=course)["atencao"],
            "teacher_status": status_for(teacher_map.get(course, {}).get("valor"), teacher_goal, f"{directorate}-02"),
            "approval_finalizados": finalized,
            "approval_aprovados": approved,
            "approval_valor": approval_map.get(course, {}).get("valor"),
            "approval_meta": _goal_fields(approval_goal, f"{directorate}-03", course=course)["meta"],
            "approval_atencao": _goal_fields(approval_goal, f"{directorate}-03", course=course)["atencao"],
            "approval_status": status_for(approval_map.get(course, {}).get("valor"), approval_goal, f"{directorate}-03"),
        })
    return out


def build_academic_report_payload(repo, *, granularity: str | None = None, reference: str | None = None,
                                  comparison: str | None = None, course: str | None = None,
                                  discipline: str | None = None, window_periods: int | str | None = None) -> dict[str, Any]:
    if repo.directorate_code not in {"DTNH", "DCS"}:
        raise ValueError("O relatório acadêmico V2 está disponível apenas para DTNH/DCS.")

    course_filter = _clean_filter(course, "(todos)")
    discipline_filter = _clean_filter(discipline, "(todas)")
    snapshot = repo.academic_dashboard_snapshot()
    snapshot["actions"] = repo.list_actions()
    dashboard = build_academic_dashboard(
        snapshot,
        course=course_filter,
        discipline=discipline_filter,
        reference=reference,
        comparison=comparison,
        window_semesters=window_periods if isinstance(window_periods, int) else None,
        directorate_code=repo.directorate_code,
        granularity=granularity,
    )
    if window_periods == "all":
        dashboard = build_academic_dashboard(
            snapshot,
            course=course_filter,
            discipline=discipline_filter,
            reference=reference,
            comparison=comparison,
            window_semesters=None,
            directorate_code=repo.directorate_code,
            granularity=granularity,
        )

    survey = SurveyRepository(repo.db, _scope_for_repo(repo))
    visible_semesters = _visible_semesters(dashboard)
    semester_set = set(visible_semesters)
    institution_history = _filter_semesters(survey.institution_nps_history(), semester_set)
    institution_by_course = _filter_semesters(survey.institution_nps_course_breakdown(), semester_set)
    if course_filter:
        institution_by_course = [row for row in institution_by_course if row.get("curso") == course_filter]
    faculty_history = _filter_semesters(survey.faculty_nps_history(), semester_set)
    institution_history = sorted(institution_history, key=lambda item: semester_sort_key(str(item.get("periodo") or "0000-SEM1")))
    institution_by_course = sorted(institution_by_course, key=lambda item: (semester_sort_key(str(item.get("periodo") or "0000-SEM1")), str(item.get("curso") or "")))
    faculty_history = sorted(faculty_history, key=lambda item: semester_sort_key(str(item.get("periodo") or "0000-SEM1")))

    codes = {
        "nps_institution": f"{repo.directorate_code}-01A",
        "nps_course": f"{repo.directorate_code}-01B",
        "nps_faculty": f"{repo.directorate_code}-01C",
        "teacher": f"{repo.directorate_code}-02",
        "approval": f"{repo.directorate_code}-03",
    }

    goals = [row for row in repo.list_goals() if row.get("indicador") in set(codes.values())]
    actions = [row for row in repo.list_actions() if row.get("indicador") in set(codes.values())]

    # 01C possui uma workspace própria e não é um card executivo no frontend.
    # O relatório, porém, apresenta os cinco KPIs acadêmicos atuais em conjunto.
    nps_reference = dashboard.get("contexto", {}).get("nps_referencia")
    nps_comparison = dashboard.get("contexto", {}).get("nps_comparacao")
    faculty_now = next((row for row in faculty_history if row.get("periodo") == nps_reference), None) or {}
    faculty_prev = next((row for row in faculty_history if row.get("periodo") == nps_comparison), None) or {}
    faculty_goal = _academic_goal(snapshot.get("metas", []), codes["nps_faculty"], nps_reference or "")
    faculty_goal_info = goal_info(faculty_goal, codes["nps_faculty"])
    faculty_value = faculty_now.get("valor")
    faculty_previous = faculty_prev.get("valor")
    report_cards = dict(dashboard.get("cards", {}))
    report_cards["nps_faculty"] = {
        "valor": faculty_value,
        "comparacao": faculty_previous,
        "variacao": (float(faculty_value) - float(faculty_previous)) if faculty_value is not None and faculty_previous is not None else None,
        "status": status_for(faculty_value, faculty_goal, codes["nps_faculty"]),
        "meta": faculty_goal_info.get("meta") if faculty_goal_info else None,
        "meta_info": faculty_goal_info,
        "respondentes": int(faculty_now.get("respondentes") or 0),
        "promotores": int(faculty_now.get("promotores") or 0),
        "neutros": int(faculty_now.get("neutros") or 0),
        "detratores": int(faculty_now.get("detratores") or 0),
    }

    return {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "app_version": APP_VERSION,
        "directorate": repo.directorate_code,
        "directorate_name": repo.directorate_name,
        "context": dashboard.get("contexto", {}),
        "visible_semesters": visible_semesters,
        "dashboard": dashboard,
        "report_cards": report_cards,
        "institution_students": institution_history,
        "institution_students_by_course": institution_by_course,
        "course_nps": _course_nps_rows(snapshot, dashboard, codes["nps_course"], course_filter),
        "faculty_nps": faculty_history,
        "teacher": _teacher_rows(snapshot, dashboard, codes["teacher"], course_filter, discipline_filter),
        "results": _result_rows(snapshot, dashboard, codes["approval"], course_filter, discipline_filter),
        "course_comparison": _comparison_rows(snapshot, dashboard, repo.directorate_code),
        "goals": goals,
        "actions": actions,
        "codes": codes,
        "aggregation_contract": {
            "student_rows_exported": False,
            "survey_raw_answers_exported": False,
            "teacher_identity_in_faculty_nps_exported": False,
            "institution_student_nps_ignores_course_filter_for_official_total": True,
            "faculty_nps_ignores_course_filter": True,
            "course_and_discipline_filters_apply_to_01B_02_03": True,
            "comparison_sheet_keeps_all_courses_in_active_directorate": True,
        },
    }


def _setup_sheet(ws, *, freeze: str | None = None, zoom: int = 90) -> None:
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = zoom
    if freeze:
        ws.freeze_panes = freeze


def _title(ws, title: str, subtitle: str | None = None, *, end_col: int = 10) -> int:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    c = ws.cell(1, 1, title)
    c.font = Font(name="Aptos Display", size=18, bold=True, color=BRAND_950)
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 30
    row = 2
    if subtitle:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_col)
        c = ws.cell(2, 1, subtitle)
        c.font = Font(name="Aptos", size=10, color=INK_600)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[2].height = 32
        row = 3
    return row


def _table_header(cell) -> None:
    cell.font = Font(name="Aptos", size=9, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=BRAND_800)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _body_cell(cell, *, imported: bool = True) -> None:
    cell.font = Font(name="Aptos", size=9, color=IMPORT_GREEN if imported else INK_800)
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    cell.border = Border(bottom=Side(style="hair", color=LINE))


def _status_fill(cell, status: str | None) -> None:
    cell.fill = PatternFill("solid", fgColor=STATUS_COLORS.get(str(status or ""), STATUS_COLORS["Sem meta"]))
    cell.font = Font(name="Aptos", size=9, bold=True, color=INK_950)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _add_table(ws, name: str, start_row: int, end_row: int, end_col: int) -> None:
    if end_row <= start_row:
        return
    table = Table(displayName=name, ref=f"A{start_row}:{get_column_letter(end_col)}{end_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium4",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def _auto_width(ws, *, min_width: int = 10, max_width: int = 44) -> None:
    for column_cells in ws.columns:
        letter = get_column_letter(column_cells[0].column)
        width = min_width
        for cell in column_cells:
            if cell.value is None:
                continue
            width = max(width, min(max_width, len(str(cell.value)) + 2))
        ws.column_dimensions[letter].width = width


def _status_formula(value_cell: str, target_cell: str, attention_cell: str) -> str:
    return (
        f'=IF({value_cell}="","Sem dados",IF({target_cell}="","Sem meta",'
        f'IF({value_cell}>={target_cell},"Dentro da meta",'
        f'IF(AND({attention_cell}<>"",{value_cell}>={attention_cell}),"Atenção","Fora da meta"))))'
    )


def _metric_value_to_excel(key: str, value: Any) -> Any:
    if value is None:
        return None
    if key == "approval":
        return float(value) / 100.0
    return float(value)


def _goal_value_to_excel(key: str, value: Any) -> Any:
    if value in (None, ""):
        return None
    numeric = float(value)
    return numeric / 100.0 if key == "approval" else numeric


def _direction_label(value: Any) -> str:
    direction = str(value or "").strip().lower()
    if direction in {"higher", "higher_is_better", "up", "maximize"}:
        return "Maior é melhor"
    if direction in {"lower", "lower_is_better", "down", "minimize"}:
        return "Menor é melhor"
    return str(value or "—")


def _write_summary(wb: Workbook, payload: dict[str, Any]) -> None:
    ws = wb.create_sheet("Resumo Executivo", 0)
    _setup_sheet(ws, zoom=85)
    context = payload["context"]
    subtitle = (
        f"{payload['directorate']} · referência {context.get('referencia') or '—'} · comparação {context.get('comparacao') or '—'} · "
        f"curso {context.get('curso') or '(todos)'} · disciplina {context.get('disciplina') or '(todas)'}"
    )
    _title(ws, f"{payload['directorate']} · Relatório Acadêmico", subtitle, end_col=12)

    cards = payload.get("report_cards") or payload["dashboard"].get("cards", {})
    codes = payload["codes"]
    metrics = [
        (codes["nps_institution"], "NPS da Instituição · Alunos", "nps_institution", "nps"),
        (codes["nps_course"], "NPS do Curso", "nps_course", "nps"),
        (codes["nps_faculty"], "NPS da Instituição · Docentes", "nps_faculty", "nps"),
        (codes["teacher"], "Avaliação Docente pelo Aluno", "avaliacao_docente", "teacher"),
        (codes["approval"], "Taxa de Aprovação", "aprovacao", "approval"),
    ]
    header_row = 5
    headers = ["KPI", "Indicador", "Resultado", "Meta", "Atenção", "Status", "Comparação", "Variação", "Base/observação"]
    for col, header in enumerate(headers, 1):
        _table_header(ws.cell(header_row, col, header))
    for row_idx, (code, label, card_key, kind) in enumerate(metrics, header_row + 1):
        card = cards.get(card_key, {}) or {}
        meta_info = card.get("meta_info") or {}
        result = _metric_value_to_excel(kind, card.get("valor"))
        comparison = _metric_value_to_excel(kind, card.get("comparacao"))
        meta = _goal_value_to_excel(kind, card.get("meta"))
        attention = _goal_value_to_excel(kind, meta_info.get("atencao"))
        observation = ""
        if card_key == "nps_institution":
            observation = f"{int(card.get('respondentes') or 0):,} respondentes · cobertura {card.get('coverage') or 0}/{card.get('coverage_total') or 0}"
        elif card_key == "nps_course":
            observation = f"{int(card.get('respondentes') or 0):,} respondentes"
        elif card_key == "nps_faculty":
            faculty = next((row for row in payload.get("faculty_nps", []) if row.get("periodo") == context.get("nps_referencia")), None) or {}
            observation = f"{int(faculty.get('respondentes') or 0):,} respondentes · população anônima"
        elif card_key == "aprovacao":
            observation = f"{int(card.get('finalizados') or 0):,} resultados finalizados"
        else:
            observation = "Média ponderada pelos respondentes"
        values = [code, label, result, meta, attention, card.get("status"), comparison, None, observation]
        for col, value in enumerate(values, 1):
            ws.cell(row_idx, col, value)
            _body_cell(ws.cell(row_idx, col), imported=(col not in {8}))
        if result is not None and comparison is not None:
            ws.cell(row_idx, 8, f"=C{row_idx}-G{row_idx}")
            ws.cell(row_idx, 8).font = Font(name="Aptos", size=9, color=INK_800)
        number_format = PCT_FMT if kind == "approval" else DECIMAL_FMT if kind == "teacher" else NPS_FMT
        for col in (3, 4, 5, 7, 8):
            ws.cell(row_idx, col).number_format = number_format
        _status_fill(ws.cell(row_idx, 6), card.get("status"))
    _add_table(ws, "TblAcademicExecutive", header_row, header_row + len(metrics), len(headers))

    # The evolution table mirrors the exact visible dashboard window and is the
    # source for the native charts below. It stays visible so the workbook is
    # auditable without hidden calculation sheets.
    start = 14
    ws.cell(start, 1, "Evolução do recorte selecionado").font = Font(name="Aptos", size=11, bold=True, color=BRAND_800)
    start += 1
    evolution_headers = [
        "Período", "01A NPS Instituição", "Meta 01A", "01B NPS Curso", "Meta 01B",
        "01C NPS Docentes", "Meta 01C", "02 Avaliação Docente", "Meta 02", "03 Aprovação", "Meta 03",
    ]
    for col, header in enumerate(evolution_headers, 1):
        _table_header(ws.cell(start, col, header))
    series = payload["dashboard"].get("series", {})
    periods = payload.get("visible_semesters") or []
    maps = {
        "inst": {row.get("periodo"): row for row in series.get("nps_institution_semestral", [])},
        "course": {row.get("periodo"): row for row in series.get("nps_course_semestral", [])},
        "faculty": {row.get("periodo"): row for row in series.get("nps_faculty_semestral", [])},
        "teacher": {to_semester(row.get("periodo")): row for row in series.get("avaliacao_docente", [])},
        "approval": {to_semester(row.get("periodo")): row for row in series.get("aprovacao", [])},
    }
    # The dashboard window can legitimately include a current/future semester
    # with targets already registered but no factual result yet. Keep the full
    # window in Parâmetros, but do not fabricate zero points in the report
    # evolution table/charts for periods that have no factual metric at all.
    def _has_factual_value(period: str) -> bool:
        candidates = (
            maps["inst"].get(period, {}).get("valor"),
            maps["course"].get(period, {}).get("valor"),
            maps["faculty"].get(period, {}).get("valor"),
            maps["teacher"].get(period, {}).get("valor"),
            maps["approval"].get(period, {}).get("valor"),
        )
        return any(value is not None for value in candidates)

    periods = [period for period in periods if _has_factual_value(period)]
    for idx, period in enumerate(periods, start + 1):
        inst, course, faculty = maps["inst"].get(period, {}), maps["course"].get(period, {}), maps["faculty"].get(period, {})
        teacher, approval = maps["teacher"].get(period, {}), maps["approval"].get(period, {})
        values = [
            period,
            inst.get("valor"), inst.get("meta"),
            course.get("valor"), course.get("meta"),
            faculty.get("valor"), faculty.get("meta"),
            teacher.get("valor"), teacher.get("meta"),
            _metric_value_to_excel("approval", approval.get("valor")), _goal_value_to_excel("approval", approval.get("meta")),
        ]
        for col, value in enumerate(values, 1):
            ws.cell(idx, col, value)
            _body_cell(ws.cell(idx, col))
        for col in range(2, 8):
            ws.cell(idx, col).number_format = NPS_FMT
        ws.cell(idx, 8).number_format = DECIMAL_FMT
        ws.cell(idx, 9).number_format = DECIMAL_FMT
        ws.cell(idx, 10).number_format = PCT_FMT
        ws.cell(idx, 11).number_format = PCT_FMT
    if periods:
        _add_table(ws, "TblAcademicEvolution", start, start + len(periods), len(evolution_headers))
        first_data = start + 1
        last_data = start + len(periods)
        cats = Reference(ws, min_col=1, min_row=first_data, max_row=last_data)

        nps_chart = LineChart()
        nps_chart.title = "NPS · evolução semestral"
        nps_chart.y_axis.title = "NPS"
        nps_chart.x_axis.title = "Semestre"
        nps_chart.height = 7.0
        nps_chart.width = 15.0
        nps_chart.legend.position = "b"
        series_titles = {
            2: "01A NPS Instituição",
            4: "01B NPS Curso",
            6: "01C NPS Docentes",
        }
        for col in (2, 4, 6):
            data = Reference(ws, min_col=col, min_row=start, max_row=last_data)
            nps_chart.add_data(data, titles_from_data=True)
            nps_chart.series[-1].title = SeriesLabel(v=series_titles[col])
        nps_chart.set_categories(cats)
        nps_chart.y_axis.scaling.min = -100
        nps_chart.y_axis.scaling.max = 100
        ws.add_chart(nps_chart, "M5")

        teacher_chart = LineChart()
        teacher_chart.title = "Avaliação Docente pelo Aluno"
        teacher_chart.y_axis.title = "Nota"
        teacher_chart.x_axis.title = "Semestre"
        teacher_chart.height = 7.0
        teacher_chart.width = 15.0
        teacher_chart.legend = None
        data = Reference(ws, min_col=8, min_row=start, max_row=last_data)
        teacher_chart.add_data(data, titles_from_data=True)
        teacher_chart.set_categories(cats)
        teacher_chart.y_axis.scaling.min = 0
        teacher_chart.y_axis.scaling.max = 10
        ws.add_chart(teacher_chart, "M20")

        approval_chart = LineChart()
        approval_chart.title = "Taxa de Aprovação"
        approval_chart.y_axis.title = "%"
        approval_chart.x_axis.title = "Semestre"
        approval_chart.height = 7.0
        approval_chart.width = 15.0
        approval_chart.legend = None
        data = Reference(ws, min_col=10, min_row=start, max_row=last_data)
        approval_chart.add_data(data, titles_from_data=True)
        approval_chart.set_categories(cats)
        approval_chart.y_axis.numFmt = PCT_FMT
        approval_chart.y_axis.scaling.min = 0
        approval_chart.y_axis.scaling.max = 1
        ws.add_chart(approval_chart, "M35")

    for col in range(1, 12):
        ws.column_dimensions[get_column_letter(col)].width = 18 if col > 1 else 15
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["I"].width = 34


def _write_nps_institution_students(wb: Workbook, payload: dict[str, Any]) -> None:
    ws = wb.create_sheet("NPS Instituição Alunos")
    _setup_sheet(ws, freeze="A5")
    _title(ws, "NPS da Instituição · Alunos (01A)", "O NPS oficial da UNIVC agrega DTNH + DCS. O detalhamento por curso abaixo permanece restrito à diretoria em visualização.", end_col=12)
    headers = ["Semestre", "NPS", "Respondentes", "Promotores", "Neutros", "Detratores", "Meta", "Atenção", "Status", "Cobertura", "Pergunta", "Questionário/Fonte"]
    header_row = 4
    for col, header in enumerate(headers, 1):
        _table_header(ws.cell(header_row, col, header))
    code = payload["codes"]["nps_institution"]
    for row_idx, item in enumerate(payload.get("institution_students", []), header_row + 1):
        period = item.get("periodo")
        # The dashboard series already contains the effective 01A goal for each visible semester.
        series_goal = next((row for row in payload["dashboard"].get("series", {}).get("nps_institution_semestral", []) if row.get("periodo") == period), {})
        meta = series_goal.get("meta")
        attention = None
        # Reconstruct attention from registered goals for auditability.
        candidates = [g for g in payload.get("goals", []) if g.get("indicador") == code and g.get("recorte") == "TOTAL" and str(g.get("vigencia") or "") <= str(period or "")]
        if candidates:
            candidates.sort(key=lambda g: str(g.get("vigencia") or ""))
            attention = candidates[-1].get("atencao")
        values = [
            period, None, int(item.get("respondentes") or 0), int(item.get("promotores") or 0), int(item.get("neutros") or 0), int(item.get("detratores") or 0),
            meta, attention, None, f"{item.get('coverage') or 0}/{item.get('coverage_total') or 0}", item.get("question") or "Pergunta oficial", item.get("questionnaire") or item.get("fonte") or "SEI",
        ]
        for col, value in enumerate(values, 1):
            ws.cell(row_idx, col, value)
            _body_cell(ws.cell(row_idx, col), imported=col not in {2, 9})
        ws.cell(row_idx, 2, f'=IF(C{row_idx}=0,"",(D{row_idx}-F{row_idx})/C{row_idx}*100)')
        ws.cell(row_idx, 9, _status_formula(f"B{row_idx}", f"G{row_idx}", f"H{row_idx}"))
        ws.cell(row_idx, 2).number_format = NPS_FMT
        ws.cell(row_idx, 7).number_format = NPS_FMT
        ws.cell(row_idx, 8).number_format = NPS_FMT
        factual = float(item.get("valor")) if item.get("valor") is not None else None
        effective_goal = candidates[-1] if candidates else None
        _status_fill(ws.cell(row_idx, 9), status_for(factual, effective_goal, code))
    rows = payload.get("institution_students", [])
    if rows:
        _add_table(ws, "TblNPSInstitutionStudents", header_row, header_row + len(rows), len(headers))
        chart = LineChart()
        chart.title = "Evolução do NPS da Instituição"
        chart.y_axis.title = "NPS"
        chart.x_axis.title = "Semestre"
        chart.height = 7.5
        chart.width = 15.0
        data = Reference(ws, min_col=2, min_row=header_row, max_row=header_row + len(rows))
        cats = Reference(ws, min_col=1, min_row=header_row + 1, max_row=header_row + len(rows))
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.y_axis.scaling.min = -100
        chart.y_axis.scaling.max = 100
        chart.legend = None
        ws.add_chart(chart, "N4")

    breakdown_start = header_row + max(1, len(rows)) + 4
    ws.cell(breakdown_start, 1, f"Detalhamento por curso · {payload['directorate']}").font = Font(name="Aptos", size=11, bold=True, color=BRAND_800)
    breakdown_start += 1
    breakdown_headers = ["Semestre", "Curso", "NPS institucional", "Respondentes", "Promotores", "Neutros", "Detratores", "Meta 01A", "Atenção", "Status", "Fonte"]
    for col, header in enumerate(breakdown_headers, 1):
        _table_header(ws.cell(breakdown_start, col, header))
    breakdown = payload.get("institution_students_by_course", [])
    for row_idx, item in enumerate(breakdown, breakdown_start + 1):
        values = [item.get("periodo"), item.get("curso"), None, int(item.get("respondentes") or 0), int(item.get("promotores") or 0), int(item.get("neutros") or 0), int(item.get("detratores") or 0), item.get("meta"), item.get("atencao"), item.get("status"), item.get("questionnaire") or item.get("fonte")]
        for col, value in enumerate(values, 1):
            ws.cell(row_idx, col, value)
            _body_cell(ws.cell(row_idx, col), imported=col not in {3, 10})
        ws.cell(row_idx, 3, f'=IF(D{row_idx}=0,"",(E{row_idx}-G{row_idx})/D{row_idx}*100)')
        ws.cell(row_idx, 3).number_format = NPS_FMT
        ws.cell(row_idx, 8).number_format = NPS_FMT
        ws.cell(row_idx, 9).number_format = NPS_FMT
        _status_fill(ws.cell(row_idx, 10), item.get("status"))
    if breakdown:
        _add_table(ws, "TblNPSInstitutionCourseBreakdown", breakdown_start, breakdown_start + len(breakdown), len(breakdown_headers))
    _auto_width(ws)


def _write_nps_course(wb: Workbook, payload: dict[str, Any]) -> None:
    ws = wb.create_sheet("NPS Cursos")
    _setup_sheet(ws, freeze="A5")
    _title(ws, "NPS do Curso (01B)", "Uma linha por curso e semestre oficial. Fechamento semestral prevalece sobre o fallback mensal, exatamente como no painel.", end_col=11)
    headers = ["Semestre", "Curso", "NPS", "Respondentes", "Promotores", "Neutros", "Detratores", "Meta", "Atenção", "Status", "Fonte"]
    header_row = 4
    for col, header in enumerate(headers, 1):
        _table_header(ws.cell(header_row, col, header))
    rows = payload.get("course_nps", [])
    for row_idx, item in enumerate(rows, header_row + 1):
        values = [item.get("periodo"), item.get("curso"), None, item.get("respondentes"), item.get("promotores"), item.get("neutros"), item.get("detratores"), item.get("meta"), item.get("atencao"), item.get("status"), item.get("fonte")]
        for col, value in enumerate(values, 1):
            ws.cell(row_idx, col, value)
            _body_cell(ws.cell(row_idx, col), imported=col not in {3, 10})
        ws.cell(row_idx, 3, f'=IF(D{row_idx}=0,"",(E{row_idx}-G{row_idx})/D{row_idx}*100)')
        ws.cell(row_idx, 3).number_format = NPS_FMT
        ws.cell(row_idx, 8).number_format = NPS_FMT
        ws.cell(row_idx, 9).number_format = NPS_FMT
        _status_fill(ws.cell(row_idx, 10), item.get("status"))
    if rows:
        _add_table(ws, "TblNPSCourses", header_row, header_row + len(rows), len(headers))
    _auto_width(ws)


def _write_nps_faculty(wb: Workbook, payload: dict[str, Any]) -> None:
    ws = wb.create_sheet("NPS Instituição Docentes")
    _setup_sheet(ws, freeze="A5")
    _title(ws, "NPS da Instituição · Docentes (01C)", "População institucional anônima. Não há curso, disciplina ou identificação individual do professor na fonte.", end_col=13)
    headers = ["Semestre", "NPS", "Respondentes", "Promotores", "Neutros", "Detratores", "Meta", "Atenção", "Status", "Pergunta", "Questionário", "Origem", "Arquivo/Fonte"]
    header_row = 4
    for col, header in enumerate(headers, 1):
        _table_header(ws.cell(header_row, col, header))
    code = payload["codes"]["nps_faculty"]
    rows = payload.get("faculty_nps", [])
    for row_idx, item in enumerate(rows, header_row + 1):
        period = item.get("periodo")
        candidates = [g for g in payload.get("goals", []) if g.get("indicador") == code and g.get("recorte") == "TOTAL" and str(g.get("vigencia") or "") <= str(period or "")]
        candidates.sort(key=lambda g: str(g.get("vigencia") or ""))
        goal = candidates[-1] if candidates else {}
        meta, attention = goal.get("meta"), goal.get("atencao")
        values = [period, None, int(item.get("respondentes") or 0), int(item.get("promotores") or 0), int(item.get("neutros") or 0), int(item.get("detratores") or 0), meta, attention, None, item.get("question") or "Pergunta NPS 0–10", item.get("questionnaire") or "—", item.get("origin") or "SEI", item.get("source_filename") or item.get("fonte") or "SEI"]
        for col, value in enumerate(values, 1):
            ws.cell(row_idx, col, value)
            _body_cell(ws.cell(row_idx, col), imported=col not in {2, 9})
        ws.cell(row_idx, 2, f'=IF(C{row_idx}=0,"",(D{row_idx}-F{row_idx})/C{row_idx}*100)')
        ws.cell(row_idx, 9, _status_formula(f"B{row_idx}", f"G{row_idx}", f"H{row_idx}"))
        ws.cell(row_idx, 2).number_format = NPS_FMT
        ws.cell(row_idx, 7).number_format = NPS_FMT
        ws.cell(row_idx, 8).number_format = NPS_FMT
        factual = (float(item.get("valor")) if item.get("valor") is not None else None)
        status = status_for(factual, goal or None, code)
        _status_fill(ws.cell(row_idx, 9), status)
    if rows:
        _add_table(ws, "TblNPSFaculty", header_row, header_row + len(rows), len(headers))
        chart = LineChart()
        chart.title = "Evolução do NPS dos docentes"
        chart.y_axis.title = "NPS"
        chart.x_axis.title = "Semestre"
        chart.height = 7.5
        chart.width = 15.0
        data = Reference(ws, min_col=2, min_row=header_row, max_row=header_row + len(rows))
        cats = Reference(ws, min_col=1, min_row=header_row + 1, max_row=header_row + len(rows))
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.y_axis.scaling.min = -100
        chart.y_axis.scaling.max = 100
        chart.legend = None
        ws.add_chart(chart, "O4")
    _auto_width(ws)


def _write_teacher(wb: Workbook, payload: dict[str, Any]) -> None:
    ws = wb.create_sheet("Avaliação Docente")
    _setup_sheet(ws, freeze="A5")
    _title(ws, "Avaliação Docente pelo Aluno (02)", "Dados agregados por semestre, curso e disciplina; linhas individuais de avaliação não são exportadas.", end_col=9)
    headers = ["Semestre", "Curso", "Disciplina", "Respondentes", "Nota média", "Meta", "Atenção", "Status", "Vigência da meta"]
    header_row = 4
    for col, header in enumerate(headers, 1):
        _table_header(ws.cell(header_row, col, header))
    rows = payload.get("teacher", [])
    for row_idx, item in enumerate(rows, header_row + 1):
        values = [item.get("periodo"), item.get("curso"), item.get("disciplina"), item.get("respondentes"), item.get("nota_media"), item.get("meta"), item.get("atencao"), item.get("status"), item.get("meta_vigencia")]
        for col, value in enumerate(values, 1):
            ws.cell(row_idx, col, value)
            _body_cell(ws.cell(row_idx, col), imported=col != 8)
        for col in (5, 6, 7):
            ws.cell(row_idx, col).number_format = DECIMAL_FMT
        _status_fill(ws.cell(row_idx, 8), item.get("status"))
    if rows:
        _add_table(ws, "TblTeacherEvaluation", header_row, header_row + len(rows), len(headers))
    _auto_width(ws)


def _write_results(wb: Workbook, payload: dict[str, Any]) -> None:
    ws = wb.create_sheet("Aprovação e Resultados")
    _setup_sheet(ws, freeze="A5")
    _title(ws, "Aprovação e Resultados Acadêmicos (03)", "Agregação por semestre, curso e disciplina. Nenhum nome ou matrícula de aluno é exportado.", end_col=15)
    headers = ["Semestre", "Curso", "Disciplina", "Registros", "Finalizados", "Aprovados", "Reprov. nota", "Reprov. falta", "Reprov. outros", "Em andamento", "Média notas", "Taxa aprovação", "Meta", "Atenção", "Status"]
    header_row = 4
    for col, header in enumerate(headers, 1):
        _table_header(ws.cell(header_row, col, header))
    rows = payload.get("results", [])
    for row_idx, item in enumerate(rows, header_row + 1):
        values = [item.get("periodo"), item.get("curso"), item.get("disciplina"), item.get("total_registros"), item.get("finalizados"), item.get("aprovados"), item.get("reprovados_nota"), item.get("reprovados_falta"), item.get("reprovados_outro"), item.get("em_andamento"), item.get("media_notas"), None, _goal_value_to_excel("approval", item.get("meta")), _goal_value_to_excel("approval", item.get("atencao")), item.get("status")]
        for col, value in enumerate(values, 1):
            ws.cell(row_idx, col, value)
            _body_cell(ws.cell(row_idx, col), imported=col not in {12, 15})
        ws.cell(row_idx, 12, f'=IF(E{row_idx}=0,"",F{row_idx}/E{row_idx})')
        ws.cell(row_idx, 11).number_format = DECIMAL_FMT
        for col in (12, 13, 14):
            ws.cell(row_idx, col).number_format = PCT_FMT
        _status_fill(ws.cell(row_idx, 15), item.get("status"))
    if rows:
        _add_table(ws, "TblAcademicResults", header_row, header_row + len(rows), len(headers))
    _auto_width(ws)


def _apply_chart_point_colors(series, statuses: list[str]) -> None:
    points = []
    for idx, status in enumerate(statuses):
        point = DataPoint(idx=idx)
        point.graphicalProperties.solidFill = STATUS_CHART_COLORS.get(status, NEUTRAL)
        point.graphicalProperties.line.solidFill = STATUS_CHART_COLORS.get(status, NEUTRAL)
        points.append(point)
    series.dPt = points


def _write_course_comparison(wb: Workbook, payload: dict[str, Any]) -> None:
    ws = wb.create_sheet("Comparação Cursos")
    _setup_sheet(ws, freeze="A5", zoom=85)
    context = payload["context"]
    _title(ws, "Comparação entre cursos", f"Referência: {context.get('referencia') or '—'} · comparação sempre restrita aos cursos da diretoria {payload['directorate']}", end_col=16)
    headers = [
        "Curso", "NPS 01B", "Meta 01B", "Atenção 01B", "Status 01B",
        "Respondentes NPS", "Avaliação 02", "Meta 02", "Atenção 02", "Status 02",
        "Taxa aprovação 03", "Meta 03", "Atenção 03", "Status 03", "Finalizados", "Aprovados",
    ]
    header_row = 4
    for col, header in enumerate(headers, 1):
        _table_header(ws.cell(header_row, col, header))
    rows = payload.get("course_comparison", [])
    for row_idx, item in enumerate(rows, header_row + 1):
        values = [
            item.get("curso"), None, item.get("nps_meta"), item.get("nps_atencao"), item.get("nps_status"), item.get("nps_respondentes"),
            None, item.get("teacher_meta"), item.get("teacher_atencao"), item.get("teacher_status"),
            None, _goal_value_to_excel("approval", item.get("approval_meta")), _goal_value_to_excel("approval", item.get("approval_atencao")), item.get("approval_status"), item.get("approval_finalizados"), item.get("approval_aprovados"),
        ]
        for col, value in enumerate(values, 1):
            ws.cell(row_idx, col, value)
            _body_cell(ws.cell(row_idx, col), imported=col not in {2, 7, 11, 5, 10, 14})
        if int(item.get("nps_respondentes") or 0):
            ws.cell(row_idx, 2, f'=({int(item.get("nps_promotores") or 0)}-{int(item.get("nps_detratores") or 0)})/{int(item.get("nps_respondentes") or 0)}*100')
        if int(item.get("teacher_respondentes") or 0):
            ws.cell(row_idx, 7, f'={float(item.get("teacher_weighted_sum") or 0)}/{int(item.get("teacher_respondentes") or 0)}')
        if int(item.get("approval_finalizados") or 0):
            ws.cell(row_idx, 11, f'=P{row_idx}/O{row_idx}')
        for col in (2, 3, 4):
            ws.cell(row_idx, col).number_format = NPS_FMT
        for col in (7, 8, 9):
            ws.cell(row_idx, col).number_format = DECIMAL_FMT
        for col in (11, 12, 13):
            ws.cell(row_idx, col).number_format = PCT_FMT
        _status_fill(ws.cell(row_idx, 5), item.get("nps_status"))
        _status_fill(ws.cell(row_idx, 10), item.get("teacher_status"))
        _status_fill(ws.cell(row_idx, 14), item.get("approval_status"))
    if rows:
        _add_table(ws, "TblCourseComparison", header_row, header_row + len(rows), len(headers))
        chart_count = min(len(rows), 20)
        chart_end = header_row + chart_count
        cats = Reference(ws, min_col=1, min_row=header_row + 1, max_row=chart_end)
        specs = [
            ("NPS do Curso · status pela meta", 2, "NPS", [row.get("nps_status") for row in rows[:chart_count]], "R4"),
            ("Avaliação Docente · status pela meta", 7, "Nota", [row.get("teacher_status") for row in rows[:chart_count]], "R20"),
            ("Aprovação · status pela meta", 11, "%", [row.get("approval_status") for row in rows[:chart_count]], "R36"),
        ]
        for title, col, ytitle, statuses, anchor in specs:
            chart = BarChart()
            chart.type = "bar"
            chart.style = 10
            chart.title = title
            chart.y_axis.title = "Curso"
            chart.x_axis.title = ytitle
            chart.height = 7.0
            chart.width = 15.0
            data = Reference(ws, min_col=col, min_row=header_row, max_row=chart_end)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            chart.legend = None
            chart.dLbls = DataLabelList()
            chart.dLbls.showVal = True
            if chart.series:
                _apply_chart_point_colors(chart.series[0], statuses)
            if col == 11:
                chart.x_axis.numFmt = PCT_FMT
                chart.x_axis.scaling.min = 0
                chart.x_axis.scaling.max = 1
            elif col == 7:
                chart.x_axis.scaling.min = 0
                chart.x_axis.scaling.max = 10
            else:
                chart.x_axis.scaling.min = -100
                chart.x_axis.scaling.max = 100
            ws.add_chart(chart, anchor)
    _auto_width(ws)


def _write_management(wb: Workbook, payload: dict[str, Any]) -> None:
    ws = wb.create_sheet("Metas e Planos")
    _setup_sheet(ws, freeze="A5")
    _title(ws, "Metas e Planos de Ação", "Governança dos cinco KPIs acadêmicos atuais da diretoria. Metas históricas permanecem disponíveis para auditoria.", end_col=12)
    goals = payload.get("goals", [])
    headers = ["Indicador", "Recorte", "Vigência", "Meta", "Atenção", "Limite superior", "Unidade", "Direção", "Justificativa"]
    header_row = 4
    for col, header in enumerate(headers, 1):
        _table_header(ws.cell(header_row, col, header))
    for row_idx, item in enumerate(goals, header_row + 1):
        values = [item.get("indicador"), item.get("recorte"), item.get("vigencia"), item.get("meta"), item.get("atencao"), item.get("limite_superior"), item.get("unidade"), _direction_label(item.get("direcao")), item.get("justificativa")]
        for col, value in enumerate(values, 1):
            ws.cell(row_idx, col, value)
            _body_cell(ws.cell(row_idx, col))
        if str(item.get("indicador") or "").endswith("-03"):
            for col in (4, 5, 6):
                if ws.cell(row_idx, col).value not in (None, ""):
                    ws.cell(row_idx, col).value = float(ws.cell(row_idx, col).value) / 100.0
                    ws.cell(row_idx, col).number_format = PCT_FMT
        else:
            for col in (4, 5, 6):
                ws.cell(row_idx, col).number_format = DECIMAL_FMT
    if goals:
        _add_table(ws, "TblAcademicGoals", header_row, header_row + len(goals), len(headers))

    start = header_row + max(1, len(goals)) + 4
    ws.cell(start, 1, "Planos de ação").font = Font(name="Aptos", size=11, bold=True, color=BRAND_800)
    start += 1
    action_headers = ["#", "Período", "Indicador", "Recorte", "Resultado", "Meta", "Problema", "Causa", "Ação", "Responsável", "Prazo", "Meta da ação", "Status", "Dias para o prazo"]
    for col, header in enumerate(action_headers, 1):
        _table_header(ws.cell(start, col, header))
    actions = payload.get("actions", [])
    for row_idx, item in enumerate(actions, start + 1):
        values = [item.get("numero"), item.get("mes"), item.get("indicador"), item.get("recorte"), item.get("resultado"), item.get("meta"), item.get("problema"), item.get("causa"), item.get("acao"), item.get("responsavel"), item.get("prazo"), item.get("meta_acao"), item.get("status"), item.get("dias_prazo")]
        for col, value in enumerate(values, 1):
            ws.cell(row_idx, col, value)
            _body_cell(ws.cell(row_idx, col))
        _status_fill(ws.cell(row_idx, 13), item.get("status") if item.get("status") in STATUS_COLORS else "Sem meta")
    if actions:
        _add_table(ws, "TblAcademicActions", start, start + len(actions), len(action_headers))
    _auto_width(ws)


def _write_parameters(wb: Workbook, payload: dict[str, Any]) -> None:
    ws = wb.create_sheet("Parâmetros")
    _setup_sheet(ws)
    _title(ws, "Parâmetros do Relatório", "Contexto do dashboard usado na geração e contrato de agregação do arquivo.", end_col=4)
    context = payload["context"]
    contract = payload["aggregation_contract"]
    rows = [
        ("Versão Data UNIVC", payload.get("app_version")),
        ("Gerado em (UTC)", payload.get("generated_at")),
        ("Diretoria", f"{payload.get('directorate')} · {payload.get('directorate_name')}"),
        ("Granularidade selecionada", context.get("granularidade")),
        ("Referência", context.get("referencia")),
        ("Comparação", context.get("comparacao")),
        ("Início da janela", context.get("inicio")),
        ("Fim da janela", context.get("fim")),
        ("Curso", context.get("curso")),
        ("Disciplina", context.get("disciplina")),
        ("Referência oficial de NPS", context.get("nps_referencia")),
        ("Comparação oficial de NPS", context.get("nps_comparacao")),
        ("Janela de semestres selecionada", ", ".join(payload.get("visible_semesters") or [])),
        ("Linhas individuais de alunos exportadas", "Não" if not contract["student_rows_exported"] else "Sim"),
        ("Respostas brutas de pesquisas exportadas", "Não" if not contract["survey_raw_answers_exported"] else "Sim"),
        ("Identidade do docente no NPS 01C exportada", "Não" if not contract["teacher_identity_in_faculty_nps_exported"] else "Sim"),
        ("01A oficial respeita filtro de curso", "Não — permanece institucional UNIVC"),
        ("01C respeita filtro de curso", "Não — população docente é anônima e institucional"),
        ("01B/02/03 respeitam curso/disciplina", "Sim"),
        ("Comparação de cursos", f"Todos os cursos ativos/observados da {payload.get('directorate')}"),
    ]
    header_row = 4
    _table_header(ws.cell(header_row, 1, "Parâmetro"))
    _table_header(ws.cell(header_row, 2, "Valor"))
    for row_idx, (label, value) in enumerate(rows, header_row + 1):
        ws.cell(row_idx, 1, label)
        ws.cell(row_idx, 2, value)
        _body_cell(ws.cell(row_idx, 1), imported=False)
        _body_cell(ws.cell(row_idx, 2), imported=False)
        ws.cell(row_idx, 1).font = Font(name="Aptos", size=9, color=STATIC_GRAY)
        ws.cell(row_idx, 2).font = Font(name="Aptos", size=9, color=STATIC_GRAY)
    _add_table(ws, "TblAcademicParameters", header_row, header_row + len(rows), 2)
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 64


def build_academic_report_workbook(payload: dict[str, Any]) -> BytesIO:
    wb = Workbook()
    placeholder = wb.active
    placeholder.title = "_placeholder"

    _write_nps_institution_students(wb, payload)
    _write_nps_course(wb, payload)
    _write_nps_faculty(wb, payload)
    _write_teacher(wb, payload)
    _write_results(wb, payload)
    _write_course_comparison(wb, payload)
    _write_management(wb, payload)
    _write_parameters(wb, payload)
    wb.remove(placeholder)
    _write_summary(wb, payload)

    tab_colors = {
        "Resumo Executivo": BRAND_900,
        "NPS Instituição Alunos": BRAND_700,
        "NPS Cursos": BRAND_600,
        "NPS Instituição Docentes": "4F7A66",
        "Avaliação Docente": "5D8FB4",
        "Aprovação e Resultados": "6A8E3A",
        "Comparação Cursos": ACCENT,
        "Metas e Planos": "A98B3A",
        "Parâmetros": "7F8C8D",
    }
    for ws in wb.worksheets:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.sheet_properties.tabColor = tab_colors.get(ws.title)
        ws.sheet_view.showGridLines = False
    wb.active = 0
    try:
        wb.calculation.calcMode = "auto"
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
    except Exception:
        pass

    output = BytesIO()
    wb.save(output)
    wb.close()
    output.seek(0)
    return output


def build_academic_report_workbook_bytes(repo, *, granularity: str | None = None, reference: str | None = None,
                                           comparison: str | None = None, course: str | None = None,
                                           discipline: str | None = None, window_periods: int | str | None = None) -> BytesIO:
    payload = build_academic_report_payload(
        repo,
        granularity=granularity,
        reference=reference,
        comparison=comparison,
        course=course,
        discipline=discipline,
        window_periods=window_periods,
    )
    return build_academic_report_workbook(payload)
