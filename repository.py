from __future__ import annotations

import json
import unicodedata
from decimal import Decimal, InvalidOperation
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import and_, case, func, insert, or_, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, joinedload

from models import (
    AcademicResult, AcademicStudent, ActionPlan, Attendance, AuditLog, Course, DadmActiveStudents, DadmAdministrativeCost,
    DadmAttrition, DadmInfrastructure, Discipline, Directorate, DpeBudgetExecution, DpeCashMovement,
    DpeOperatingResult, Enrollment, Goal, NpsInstitution, NpsInstitutionFaculty, NpsStudent, Profile, AppUser, TeacherEvaluation,
)
from schemas import (
    ACADEMIC_DIRECTORATES, ACADEMIC_UI_DATASETS, DADM_HEADER_ALIASES, DADM_MODALITIES, DATASETS, DISCIPLINE_HEADER_ALIASES,
    DPE_CASH_MOVEMENT_TYPES, DPE_HEADER_ALIASES, DPE_RESULT_SCOPE_TYPES, HEADER_ALIASES,
    IMPLEMENTED_KPIS, KPI_META, MONTH_RE, SEMESTER_RE, RepositoryError, ValidationError, norm_header,
)
from auth.objects import require_object_for_directorate
from security import UserContext
from academic_excel_parser import classificacao_aprovacao, ler_registros, motivo_reprovacao
from academic_catalog import canonical_course_name, course_name_matches, is_ambiguous_course_name
from analytics import active_goal, goal_info, status_for


DATASET_MODEL = {"nps": NpsStudent, "avaliacao_docente": TeacherEvaluation, "resultados": AcademicResult, "matriculas": Enrollment, "frequencia": Attendance}

# Política de identidade de Educação Física: a habilitação é determinada exclusivamente
# pelo campo Curso: do XLSX. Nome/código da turma e turno do seletor do SEI nunca
# participam dessa decisão. Este identificador também é exposto no diagnóstico do build.
EDUCACAO_FISICA_IDENTITY_POLICY = "xlsx-course-field-only-v3"


class DatabaseRepository:
    def __init__(
        self,
        db: Session,
        ctx: UserContext,
        directorate_id: int | None = None,
        directorate_code: str | None = None,
        directorate_name: str | None = None,
    ):
        self.db = db
        self.ctx = ctx
        self._directorate_id = directorate_id or ctx.directorate_id
        self.directorate_code = directorate_code or ctx.directorate_code
        self.directorate_name = directorate_name or ctx.directorate_name
        # Caches por requisição: importações do SEI podem ter milhares de linhas.
        # Evitam repetir consultas de catálogo para o mesmo curso/disciplina.
        self._course_name_cache: dict[str, Course | None] = {}
        self._discipline_name_cache: dict[tuple[int, str], Discipline | None] = {}
        self._student_registration_cache: dict[str, AcademicStudent | None] = {}

    @property
    def directorate_id(self) -> int:
        return self._directorate_id

    @staticmethod
    def _text_key(value: Any) -> str:
        """Comparação Unicode estável entre SQLite e PostgreSQL.

        O lower() nativo do SQLite é essencialmente ASCII e não converte, por
        exemplo, Á/Ç/Ã em nomes acadêmicos. Isso fazia a segunda linha de uma
        disciplina em caixa alta acentuada parecer um cadastro novo.
        """
        text = unicodedata.normalize("NFKC", str(value or "")).casefold()
        return " ".join(text.split())

    def _course_by_name(self, name: str, *, active_only: bool = False) -> Course | None:
        key = self._text_key(name)
        if not key:
            return None
        if key not in self._course_name_cache:
            rows = self.db.scalars(select(Course).where(Course.directorate_id == self.directorate_id)).all()
            # Preenche o cache inteiro de uma vez; normalmente são poucos cursos.
            for row in rows:
                self._course_name_cache[self._text_key(row.name)] = row
            self._course_name_cache.setdefault(key, None)
        row = self._course_name_cache.get(key)
        if active_only and row is not None and not row.active:
            return None
        return row

    def _discipline_by_name(self, course_id: int, name: str, *, active_only: bool = False) -> Discipline | None:
        key = self._text_key(name)
        if not key:
            return None
        cache_key = (course_id, key)
        if cache_key not in self._discipline_name_cache:
            rows = self.db.scalars(select(Discipline).where(Discipline.course_id == course_id)).all()
            for row in rows:
                self._discipline_name_cache[(course_id, self._text_key(row.name))] = row
            self._discipline_name_cache.setdefault(cache_key, None)
        row = self._discipline_name_cache.get(cache_key)
        if active_only and row is not None and not row.active:
            return None
        return row

    def _commit(self):
        try:
            self.db.commit()
        except IntegrityError as exc:
            self.db.rollback()
            raise ValidationError("Já existe um registro com essa combinação de período e recorte.", {"periodo": "Registro duplicado."}) from exc

    def _audit(self, action: str, entity: str, entity_id: Any = None, details: Any = None):
        self.db.add(AuditLog(
            directorate_id=self.directorate_id,
            user_id=self.ctx.user_id,
            user_email=self.ctx.email,
            action=action,
            entity=entity,
            entity_id=str(entity_id) if entity_id is not None else None,
            details=json.dumps(details, ensure_ascii=False, default=str) if details is not None else None,
        ))

    def get_config(self) -> dict[str, Any]:
        return {
            "responsavel": self.ctx.full_name,
            "diretoria_exibicao": self.directorate_code,
            "diretoria_nome": self.directorate_name,
            "diretoria_usuario": self.ctx.directorate_code,
            "modo": "cloud",
            "usuario": self.ctx.email,
            "papel": self.ctx.role,
        }

    def save_config(self, payload: dict[str, Any]) -> dict[str, Any]:
        profile = self.db.get(Profile, self.ctx.user_id)
        if profile and payload.get("responsavel"):
            name = str(payload["responsavel"]).strip()[:120]
            profile.full_name = name
            app_user = self.db.get(AppUser, self.ctx.user_id)
            if app_user:
                # app_users is the canonical Identity & Access V2 source. Keep its
                # display name synchronized while profiles remains in migration.
                app_user.name = name
            self._audit("update", "profile", profile.id, {"full_name": name})
            self._commit()
            self.ctx.full_name = name or self.ctx.full_name
        return self.get_config()

    def list_courses(self, include_inactive: bool = True) -> list[dict[str, Any]]:
        q = select(Course).where(Course.directorate_id == self.directorate_id)
        if not include_inactive:
            q = q.where(Course.active.is_(True))
        rows = self.db.scalars(q.order_by(Course.name)).all()
        return [{
            "id": r.id,
            "curso": r.name,
            "modalidade": r.modality or "Presencial",
            "ativo": r.active,
            "vigencia_inicio": r.valid_from or "",
            "vigencia_fim": r.valid_to or "",
        } for r in rows]

    def list_disciplines(self, include_inactive: bool = True) -> list[dict[str, Any]]:
        q = select(Discipline).options(joinedload(Discipline.course)).join(Course).where(Course.directorate_id == self.directorate_id)
        if not include_inactive:
            q = q.where(Discipline.active.is_(True), Course.active.is_(True))
        rows = self.db.scalars(q.order_by(Course.name, Discipline.name)).all()
        return [{"id": r.id, "curso": r.course.name, "disciplina": r.name, "ativo": r.active, "vigencia_inicio": r.valid_from or "", "vigencia_fim": r.valid_to or "", "recorte_metas": f"{r.course.name} » {r.name}"} for r in rows]

    def course_maps(self, include_inactive: bool = False):
        courses = self.list_courses(include_inactive=include_inactive)
        disciplines = self.list_disciplines(include_inactive=include_inactive)
        names = [x["curso"] for x in courses if include_inactive or x["ativo"]]
        dmap = {name: [] for name in names}
        for d in disciplines:
            if include_inactive or d["ativo"]:
                dmap.setdefault(d["curso"], []).append(d["disciplina"])
        return names, dmap

    def create_course(self, payload: dict[str, Any]) -> dict[str, Any]:
        raw_name = str(payload.get("curso") or "").strip()
        name = canonical_course_name(raw_name, self.directorate_code)
        start = str(payload.get("vigencia_inicio") or "").strip()
        modality = self._normalize_modality(payload.get("modalidade") or "Presencial")
        errors = {}
        if is_ambiguous_course_name(raw_name, self.directorate_code):
            errors["curso"] = "Cadastre Educação Física separadamente como Bacharelado ou Licenciatura."
        if len(name) < 2: errors["curso"] = "Informe o nome do curso."
        if not MONTH_RE.match(start): errors["vigencia_inicio"] = "Use o formato AAAA-MM."
        if modality not in DADM_MODALITIES: errors["modalidade"] = "Use Presencial, EAD ou Semipresencial."
        if self._course_by_name(name):
            errors["curso"] = "Curso já cadastrado."
        if errors: raise ValidationError("Revise os campos destacados.", errors)
        row = Course(directorate_id=self.directorate_id, name=name, modality=modality, active=True, valid_from=start)
        self.db.add(row); self.db.flush(); self._audit("create", "course", row.id, payload); self._commit()
        return next(x for x in self.list_courses() if x["id"] == row.id)

    def set_course_active(self, row_id: int, active: bool, end_period: str | None = None) -> dict[str, Any]:
        row = require_object_for_directorate(self.db, Course, row_id, self.directorate_id, label="Curso")
        if not active and not MONTH_RE.match(str(end_period or "")): raise ValidationError("Revise os campos destacados.", {"vigencia_fim": "Use o formato AAAA-MM."})
        row.active = active; row.valid_to = None if active else str(end_period)
        self._audit("status", "course", row.id, {"active": active, "valid_to": row.valid_to}); self._commit()
        return next(x for x in self.list_courses() if x["id"] == row.id)

    def create_discipline(self, payload: dict[str, Any]) -> dict[str, Any]:
        course_name = str(payload.get("curso") or "").strip(); name = str(payload.get("disciplina") or "").strip(); start = str(payload.get("vigencia_inicio") or "").strip(); errors = {}
        course = self.db.scalar(select(Course).where(Course.directorate_id == self.directorate_id, Course.name == course_name, Course.active.is_(True)))
        if not course: errors["curso"] = "Selecione um curso ativo."
        if len(name) < 2: errors["disciplina"] = "Informe o nome da disciplina."
        if not MONTH_RE.match(start): errors["vigencia_inicio"] = "Use o formato AAAA-MM."
        if course and self._discipline_by_name(course.id, name): errors["disciplina"] = "Disciplina já cadastrada."
        if errors: raise ValidationError("Revise os campos destacados.", errors)
        row = Discipline(course_id=course.id, name=name, active=True, valid_from=start)
        self.db.add(row); self.db.flush(); self._audit("create", "discipline", row.id, payload); self._commit()
        return next(x for x in self.list_disciplines() if x["id"] == row.id)

    def set_discipline_active(self, row_id: int, active: bool, end_period: str | None = None) -> dict[str, Any]:
        row = require_object_for_directorate(self.db, Discipline, row_id, self.directorate_id, label="Disciplina")
        if not active and not MONTH_RE.match(str(end_period or "")): raise ValidationError("Revise os campos destacados.", {"vigencia_fim": "Use o formato AAAA-MM."})
        row.active = active; row.valid_to = None if active else str(end_period)
        self._audit("status", "discipline", row.id, {"active": active, "valid_to": row.valid_to}); self._commit()
        return next(x for x in self.list_disciplines() if x["id"] == row.id)

    def _to_int(self, value: Any, field: str, errors: dict, allow_zero: bool = True):
        try: number = int(value)
        except (TypeError, ValueError): errors[field] = "Informe um número inteiro."; return None
        if number < 0 or (not allow_zero and number == 0): errors[field] = "O valor deve ser positivo." if not allow_zero else "O valor não pode ser negativo."
        return number

    @staticmethod
    def _semester_from_period(value: Any) -> str:
        text = str(value or "").strip().upper()
        if SEMESTER_RE.match(text):
            return text
        m = MONTH_RE.match(text)
        if m:
            year, month = text.split("-")
            return f"{year}-SEM{1 if int(month) <= 6 else 2}"
        match = __import__("re").match(r"^(\d{4})\s*[/\-]\s*([12])$", text)
        if match:
            return f"{match.group(1)}-SEM{match.group(2)}"
        return text

    @staticmethod
    def _semester_start_month(period: str) -> str:
        semester = DatabaseRepository._semester_from_period(period)
        if SEMESTER_RE.match(semester):
            return f"{semester[:4]}-{'01' if semester.endswith('1') else '07'}"
        if MONTH_RE.match(period):
            return period
        return f"{date.today().year}-01"

    def _get_or_create_course(self, name: str, period: str, modality: str = "Presencial") -> Course:
        raw_course_name = str(name or "").strip()
        if is_ambiguous_course_name(raw_course_name, self.directorate_code):
            raise ValidationError(
                "Revise os campos destacados.",
                {"curso": "Educação Física precisa ser Bacharelado ou Licenciatura."},
            )
        course_name = canonical_course_name(raw_course_name, self.directorate_code)
        if len(course_name) < 2:
            raise ValidationError("Revise os campos destacados.", {"curso": "Informe o nome do curso."})
        row = self._course_by_name(course_name)
        if row:
            if not row.active:
                row.active = True
                row.valid_to = None
            return row
        row = Course(
            directorate_id=self.directorate_id,
            name=course_name,
            modality=self._normalize_modality(modality),
            active=True,
            valid_from=self._semester_start_month(period),
        )
        self.db.add(row)
        self.db.flush()
        self._course_name_cache[self._text_key(course_name)] = row
        return row

    def _get_or_create_discipline(self, course: Course, name: str, period: str) -> Discipline:
        discipline_name = str(name or "").strip()
        if len(discipline_name) < 2:
            raise ValidationError("Revise os campos destacados.", {"disciplina": "Informe o nome da disciplina."})
        row = self._discipline_by_name(course.id, discipline_name)
        if row:
            if not row.active:
                row.active = True
                row.valid_to = None
            return row
        row = Discipline(
            course_id=course.id,
            name=discipline_name,
            active=True,
            valid_from=self._semester_start_month(period),
        )
        self.db.add(row)
        self.db.flush()
        self._discipline_name_cache[(course.id, self._text_key(discipline_name))] = row
        return row

    def _prefetch_students(self, registrations: list[str] | set[str], *, chunk_size: int = 400) -> None:
        pending = [
            str(registration).strip()
            for registration in registrations
            if str(registration or "").strip() and str(registration).strip() not in self._student_registration_cache
        ]
        if not pending:
            return
        unique = list(dict.fromkeys(pending))
        for offset in range(0, len(unique), chunk_size):
            chunk = unique[offset:offset + chunk_size]
            rows = self.db.scalars(
                select(AcademicStudent).where(
                    AcademicStudent.directorate_id == self.directorate_id,
                    AcademicStudent.registration.in_(chunk),
                )
            ).all()
            found = {row.registration: row for row in rows}
            for registration in chunk:
                self._student_registration_cache[registration] = found.get(registration)

    def _reset_import_caches(self) -> None:
        self._course_name_cache = {}
        self._discipline_name_cache = {}
        self._student_registration_cache = {}

    def _get_or_create_student(
        self,
        registration: str,
        name: str,
        course: Course,
        *,
        flush: bool = True,
    ) -> AcademicStudent:
        registration = str(registration or "").strip()
        student_name = str(name or "").strip()
        errors: dict[str, str] = {}
        if not registration:
            errors["matricula"] = "Informe a matrícula do aluno."
        if len(student_name) < 2:
            errors["aluno"] = "Informe o nome do aluno."
        if errors:
            raise ValidationError("Revise os campos destacados.", errors)

        if registration not in self._student_registration_cache:
            self._prefetch_students([registration])
        row = self._student_registration_cache.get(registration)
        if row:
            # A matrícula é a identidade do aluno. Nome e curso podem ser atualizados
            # sem alterar os resultados históricos, que guardam o curso em cada linha.
            if student_name and row.name != student_name:
                row.name = student_name
            if row.course_id != course.id:
                row.course_id = course.id
            if not row.active:
                row.active = True
            return row

        row = AcademicStudent(
            directorate_id=self.directorate_id,
            registration=registration,
            name=student_name,
            course_id=course.id,
            active=True,
            inserted_by=self.ctx.full_name,
        )
        self.db.add(row)
        if flush:
            self.db.flush()
        self._student_registration_cache[registration] = row
        return row

    def _normalize_result_payload(self, payload: dict[str, Any]) -> dict[str, Any]:
        """Validate Resultados Acadêmicos without performing database writes.

        The old validator resolved/created course, discipline and student before
        validating all scalar fields. For large imports that forced one transaction
        savepoint and multiple SELECTs per row. This pure phase preserves the same
        field rules while allowing entity resolution to be batched afterwards.
        """
        errors: dict[str, str] = {}
        period_raw = str(payload.get("periodo") or "").strip().upper()
        period = self._semester_from_period(period_raw)
        if not SEMESTER_RE.match(period):
            errors["periodo"] = "Use o formato AAAA-SEM1 ou AAAA-SEM2."

        raw_course_name = str(payload.get("curso") or "").strip()
        course_name = canonical_course_name(raw_course_name, self.directorate_code)
        if is_ambiguous_course_name(raw_course_name, self.directorate_code):
            errors["curso"] = (
                "Educação Física precisa ser identificada como Bacharelado ou Licenciatura. "
                "Use o curso específico informado pelo SEI."
            )
        if not course_name:
            errors["curso"] = "Informe o curso."

        discipline_name = str(payload.get("disciplina") or "").strip()
        if len(discipline_name) < 2:
            errors["disciplina"] = "Informe a disciplina."
        class_group = str(payload.get("turma") or "").strip()
        if not class_group:
            errors["turma"] = "Informe a turma."
        registration = str(payload.get("matricula") or "").strip()
        student_name = str(payload.get("aluno") or "").strip()
        if not registration:
            errors["matricula"] = "Informe a matrícula do aluno."
        if len(student_name) < 2:
            errors["aluno"] = "Informe o nome do aluno."

        raw_average = payload.get("media")
        average = None
        if raw_average not in (None, "", "--", "-", "—"):
            try:
                score = Decimal(str(raw_average).replace(",", "."))
                if score < 0 or score > 10:
                    errors["media"] = "A média deve ficar entre 0 e 10."
                average = score
            except (InvalidOperation, TypeError, ValueError):
                errors["media"] = "Informe uma média válida entre 0 e 10."

        official_status = str(payload.get("situacao") or "").strip() or None
        approved = self._parse_boolish(payload.get("aprovado"))
        if approved is None and official_status:
            approved = classificacao_aprovacao(official_status)

        reason = str(payload.get("motivo_reprovacao") or "").strip().casefold() or None
        if reason in {"reprovado falta", "frequencia", "frequência"}:
            reason = "falta"
        elif reason in {"reprovado", "nota/media", "media", "média"}:
            reason = "nota"
        if reason is None and official_status:
            reason = motivo_reprovacao(official_status)
        if approved is True:
            reason = None
        if reason not in {None, "nota", "falta", "outro"}:
            reason = "outro"

        if errors:
            raise ValidationError("Revise os campos destacados.", errors)
        return {
            "periodo": period,
            "curso": course_name,
            "disciplina": discipline_name,
            "turma": class_group,
            "matricula": registration,
            "aluno": student_name,
            "periodo_curricular": str(payload.get("periodo_curricular") or payload.get("periodo_matriz") or "").strip() or None,
            "media": average,
            "situacao": official_status,
            "aprovado": approved,
            "motivo_reprovacao": reason,
            "fonte": str(payload.get("fonte") or "manual").strip()[:30] or "manual",
        }

    def _resolve_result_entities(self, clean: dict[str, Any]) -> dict[str, Any]:
        resolved = dict(clean)
        course = self._get_or_create_course(clean["curso"], clean["periodo"])
        discipline = self._get_or_create_discipline(course, clean["disciplina"], clean["periodo"])
        student = self._get_or_create_student(clean["matricula"], clean["aluno"], course)
        resolved["course_id"] = course.id
        resolved["discipline_id"] = discipline.id
        resolved["student_id"] = student.id
        return resolved

    @staticmethod
    def _parse_boolish(value: Any) -> bool | None:
        if value is None or value == "":
            return None
        if isinstance(value, bool):
            return value
        text = str(value).strip().casefold()
        if text in {"1", "sim", "s", "true", "aprovado", "aprovada"}:
            return True
        if text in {"0", "nao", "não", "n", "false", "reprovado", "reprovada"}:
            return False
        return None

    def validate_payload(self, dataset: str, payload: dict[str, Any], ignore_id: int | None = None) -> dict[str, Any]:
        if dataset not in DATASETS:
            raise ValidationError("Base desconhecida.")

        if dataset == "resultados":
            return self._resolve_result_entities(self._normalize_result_payload(payload))

        period_raw = str(payload.get("periodo") or "").strip().upper()
        raw_course_name = str(payload.get("curso") or "").strip()
        course_name = canonical_course_name(raw_course_name, self.directorate_code)
        clean: dict[str, Any] = {"periodo": period_raw, "curso": course_name}
        errors: dict[str, str] = {}
        if is_ambiguous_course_name(raw_course_name, self.directorate_code):
            errors["curso"] = (
                "Educação Física precisa ser identificada como Bacharelado ou Licenciatura. "
                "Use o curso específico informado pelo SEI."
            )

        if dataset == "avaliacao_docente":
            clean["periodo"] = self._semester_from_period(period_raw)
            if not SEMESTER_RE.match(clean["periodo"]):
                errors["periodo"] = "Use o formato AAAA-SEM1 ou AAAA-SEM2."
        elif dataset == "nps":
            clean["periodo"] = self._semester_from_period(period_raw) if SEMESTER_RE.match(period_raw) else period_raw
            if not SEMESTER_RE.match(clean["periodo"]):
                errors["periodo"] = "O NPS acadêmico é semestral. Use AAAA-SEM1 ou AAAA-SEM2."
        elif dataset in {"matriculas", "frequencia"} and not MONTH_RE.match(period_raw):
            errors["periodo"] = "Selecione uma competência mensal válida (AAAA-MM)."

        course = None
        if course_name:
            # Nos novos fluxos acadêmicos, o catálogo nasce junto com o dado.
            # Isso permite iniciar um ambiente vazio e deixar o SEI/Excel descobrir os cursos.
            if dataset in {"nps", "avaliacao_docente", "resultados"}:
                try:
                    course = self._get_or_create_course(course_name, clean["periodo"] or period_raw)
                except ValidationError as exc:
                    errors.update(exc.field_errors)
            else:
                course = self._course_by_name(course_name, active_only=True)
                if not course:
                    errors["curso"] = "Selecione um curso cadastrado e ativo."
        else:
            errors["curso"] = "Informe o curso."
        clean["course_id"] = course.id if course else None

        if dataset == "nps":
            for field in ("respondentes", "promotores", "neutros", "detratores"):
                clean[field] = self._to_int(payload.get(field), field, errors)
            if clean.get("respondentes") == 0:
                errors["respondentes"] = "O total de respondentes deve ser maior que zero."
            relevant = ("respondentes", "promotores", "neutros", "detratores")
            if not any(field in errors for field in relevant):
                if clean["promotores"] + clean["neutros"] + clean["detratores"] != clean["respondentes"]:
                    errors["respondentes"] = "Promotores + neutros + detratores deve ser igual aos respondentes."

        elif dataset == "avaliacao_docente":
            clean["disciplina"] = str(payload.get("disciplina") or "").strip()
            clean["professor"] = str(payload.get("professor") or "").strip()
            if len(clean["professor"]) < 2:
                errors["professor"] = "Informe o nome do professor."
            discipline = None
            if course and clean["disciplina"]:
                try:
                    discipline = self._get_or_create_discipline(course, clean["disciplina"], clean["periodo"])
                except ValidationError as exc:
                    errors.update(exc.field_errors)
            elif not clean["disciplina"]:
                errors["disciplina"] = "Informe a disciplina."
            clean["discipline_id"] = discipline.id if discipline else None
            clean["respondentes"] = self._to_int(payload.get("respondentes"), "respondentes", errors, allow_zero=False)
            try:
                score = Decimal(str(payload.get("nota_media")).replace(",", "."))
                if score < 0 or score > 10:
                    errors["nota_media"] = "A nota média deve ficar entre 0 e 10."
                clean["nota_media"] = score
            except (InvalidOperation, TypeError, ValueError):
                errors["nota_media"] = "Informe uma nota média válida entre 0 e 10."
                clean["nota_media"] = None

        elif dataset == "matriculas":
            clean["matriculas"] = self._to_int(payload.get("matriculas"), "matriculas", errors)

        elif dataset == "frequencia":
            clean["disciplina"] = str(payload.get("disciplina") or "").strip()
            discipline = self._discipline_by_name(course.id, clean["disciplina"], active_only=True) if course and clean["disciplina"] else None
            if not discipline:
                errors["disciplina"] = "Selecione uma disciplina vinculada ao curso."
            clean["discipline_id"] = discipline.id if discipline else None
            clean["presencas_previstas"] = self._to_int(payload.get("presencas_previstas"), "presencas_previstas", errors, allow_zero=False)
            clean["presencas_registradas"] = self._to_int(payload.get("presencas_registradas"), "presencas_registradas", errors)
            if not errors.get("presencas_previstas") and not errors.get("presencas_registradas"):
                if clean["presencas_registradas"] > clean["presencas_previstas"]:
                    errors["presencas_registradas"] = "As presenças registradas não podem superar as previstas."

        if errors:
            raise ValidationError("Revise os campos destacados.", errors)
        return clean

    def _record_to_dict(self, dataset: str, row):
        common = {
            "id": row.id,
            "periodo": row.period,
            "validacao": "OK",
            "data_lancamento": str(row.inserted_at.date()) if row.inserted_at else "",
            "lancado_por": row.inserted_by or "",
        }
        if dataset == "nps":
            value = round((row.promoters - row.detractors) / row.respondents * 100, 2) if row.respondents else None
            source_type = str(getattr(row, "source_type", None) or "manual")
            source_label = "SEI · Questionário" if source_type == "SEI_SURVEY" else "Manual / Excel"
            return {
                **common,
                "curso": row.course.name,
                "respondentes": row.respondents,
                "promotores": row.promoters,
                "neutros": row.neutrals,
                "detratores": row.detractors,
                "valor": value,
                "fonte": source_label,
                "source_type": source_type,
                "survey_run_id": getattr(row, "survey_run_id", None),
                "survey_question_id": getattr(row, "survey_question_id", None),
            }
        if dataset == "avaliacao_docente":
            return {
                **common, "curso": row.course.name, "disciplina": row.discipline.name,
                "professor": row.teacher_name, "respondentes": row.respondents,
                "nota_media": float(row.average_score), "valor": float(row.average_score),
            }
        if dataset == "resultados":
            return {
                **common, "curso": row.course.name, "disciplina": row.discipline.name,
                "turma": row.class_group, "periodo_curricular": row.curriculum_period or "",
                "matricula": row.student.registration, "aluno": row.student.name,
                "media": float(row.final_average) if row.final_average is not None else None,
                "situacao": row.official_status or "", "aprovado": row.approved,
                "motivo_reprovacao": row.failure_reason or "", "fonte": row.source,
                "valor": float(row.final_average) if row.final_average is not None else None,
            }
        if dataset == "matriculas":
            return {**common, "curso": row.course.name, "matriculas": row.active_enrollments, "valor": row.active_enrollments}
        value = round(row.recorded_attendance / row.expected_attendance * 100, 2) if row.expected_attendance else None
        return {**common, "curso": row.course.name, "disciplina": row.discipline.name, "presencas_previstas": row.expected_attendance, "presencas_registradas": row.recorded_attendance, "valor": value}

    def list_records(self, dataset: str, filters: dict[str, str] | None = None):
        if dataset not in DATASETS:
            raise RepositoryError("Base desconhecida.")
        model = DATASET_MODEL[dataset]
        q = select(model).where(model.directorate_id == self.directorate_id)
        if dataset == "nps":
            q = q.options(joinedload(NpsStudent.course))
        elif dataset == "avaliacao_docente":
            q = q.options(joinedload(TeacherEvaluation.course), joinedload(TeacherEvaluation.discipline))
        elif dataset == "resultados":
            q = q.options(joinedload(AcademicResult.course), joinedload(AcademicResult.discipline), joinedload(AcademicResult.student))
        elif dataset == "matriculas":
            q = q.options(joinedload(Enrollment.course))
        elif dataset == "frequencia":
            q = q.options(joinedload(Attendance.course), joinedload(Attendance.discipline))
        filters = filters or {}
        if filters.get("periodo"):
            period_filter = str(filters["periodo"]).upper()
            if dataset in {"avaliacao_docente", "resultados"}:
                period_filter = self._semester_from_period(period_filter)
            q = q.where(model.period == period_filter)
        rows = self.db.scalars(q.order_by(model.period.desc(), model.id.desc())).all()
        result = [self._record_to_dict(dataset, row) for row in rows]
        if dataset == "nps" and result:
            indicator = f"{self.directorate_code}-01B"
            legacy_indicator = f"{self.directorate_code}-01"
            goal_rows = self.db.scalars(
                select(Goal).where(
                    Goal.directorate_id == self.directorate_id,
                    Goal.indicator_code.in_([indicator, legacy_indicator]),
                )
            ).all()
            goals = [
                {
                    "indicador": row.indicator_code,
                    "recorte": row.scope_label,
                    "vigencia": row.valid_from,
                    "meta": row.target,
                    "atencao": row.attention,
                    "limite_superior": row.upper_limit,
                    "justificativa": row.justification or "",
                }
                for row in goal_rows
            ]
            for item in result:
                goal = active_goal(goals, indicator, item.get("periodo") or "", item.get("curso"))
                if not goal:
                    goal = active_goal(goals, legacy_indicator, item.get("periodo") or "", item.get("curso"))
                info = goal_info(goal, indicator, item.get("curso"))
                item.update({
                    "meta": info.get("meta") if info else None,
                    "atencao": info.get("atencao") if info else None,
                    "limite_superior": info.get("limite_superior") if info else None,
                    "meta_vigencia": info.get("vigencia") if info else None,
                    "meta_recorte": info.get("recorte") if info else None,
                    "meta_origem": info.get("origem") if info else None,
                    "status": status_for(item.get("valor"), goal, indicator),
                })
        if filters.get("curso"):
            result = [row for row in result if row.get("curso") == filters["curso"]]
        if filters.get("disciplina"):
            result = [row for row in result if row.get("disciplina") == filters["disciplina"]]
        if filters.get("busca"):
            term = filters["busca"].casefold()
            result = [row for row in result if term in " ".join(str(value) for value in row.values()).casefold()]
        return result

    def _teacher_evaluation_conditions(
        self,
        filters: dict[str, str] | None = None,
        *,
        include_period: bool = True,
        include_search: bool = False,
    ) -> list[Any]:
        """Build SQL predicates for Faculty Evaluation without materializing rows."""
        filters = filters or {}
        conditions: list[Any] = [TeacherEvaluation.directorate_id == self.directorate_id]
        if include_period and filters.get("periodo"):
            conditions.append(TeacherEvaluation.period == self._semester_from_period(filters["periodo"]))
        if filters.get("curso"):
            conditions.append(Course.name == str(filters["curso"]).strip())
        if filters.get("disciplina"):
            conditions.append(Discipline.name == str(filters["disciplina"]).strip())
        if filters.get("professor"):
            conditions.append(TeacherEvaluation.teacher_name == str(filters["professor"]).strip())
        if include_search and filters.get("busca"):
            like = f"%{str(filters['busca']).strip()}%"
            conditions.append(or_(
                TeacherEvaluation.period.ilike(like),
                TeacherEvaluation.teacher_name.ilike(like),
                Course.name.ilike(like),
                Discipline.name.ilike(like),
                TeacherEvaluation.inserted_by.ilike(like),
            ))
        return conditions

    def list_teacher_evaluations_page(
        self,
        filters: dict[str, str] | None = None,
        *,
        page: int = 1,
        page_size: int = 50,
    ) -> dict[str, Any]:
        """Server-side Faculty Evaluation table.

        This dataset can grow every semester. The browser therefore receives only
        the requested page; filtering, searching and counting happen in SQL.
        """
        filters = filters or {}
        page = max(1, int(page or 1))
        page_size = min(100, max(10, int(page_size or 50)))
        conditions = self._teacher_evaluation_conditions(filters, include_period=True, include_search=True)

        total_q = (
            select(func.count(TeacherEvaluation.id))
            .join(Course, TeacherEvaluation.course_id == Course.id)
            .join(Discipline, TeacherEvaluation.discipline_id == Discipline.id)
            .where(*conditions)
        )
        total = int(self.db.scalar(total_q) or 0)
        pages = max(1, (total + page_size - 1) // page_size)
        page = min(page, pages)

        rows = self.db.scalars(
            select(TeacherEvaluation)
            .join(Course, TeacherEvaluation.course_id == Course.id)
            .join(Discipline, TeacherEvaluation.discipline_id == Discipline.id)
            .options(joinedload(TeacherEvaluation.course), joinedload(TeacherEvaluation.discipline))
            .where(*conditions)
            .order_by(TeacherEvaluation.period.desc(), TeacherEvaluation.id.desc())
            .offset((page - 1) * page_size)
            .limit(page_size)
        ).all()
        return {
            "items": [self._record_to_dict("avaliacao_docente", row) for row in rows],
            "total": total,
            "page": page,
            "page_size": page_size,
            "pages": pages,
        }

    def teacher_evaluation_filter_options(self, filters: dict[str, str] | None = None) -> dict[str, list[str]]:
        """Return cascade filter options with DISTINCT SQL queries, never raw rows."""
        filters = filters or {}
        base = [TeacherEvaluation.directorate_id == self.directorate_id]
        periods = list(self.db.scalars(
            select(TeacherEvaluation.period)
            .where(*base)
            .distinct()
            .order_by(TeacherEvaluation.period.desc())
        ).all())
        courses = list(self.db.scalars(
            select(Course.name)
            .join(TeacherEvaluation, TeacherEvaluation.course_id == Course.id)
            .where(*base)
            .distinct()
            .order_by(Course.name)
        ).all())

        discipline_conditions = list(base)
        if filters.get("curso"):
            discipline_conditions.append(Course.name == str(filters["curso"]).strip())
        disciplines = list(self.db.scalars(
            select(Discipline.name)
            .join(TeacherEvaluation, TeacherEvaluation.discipline_id == Discipline.id)
            .join(Course, TeacherEvaluation.course_id == Course.id)
            .where(*discipline_conditions)
            .distinct()
            .order_by(Discipline.name)
        ).all())

        professor_conditions = list(discipline_conditions)
        if filters.get("disciplina"):
            professor_conditions.append(Discipline.name == str(filters["disciplina"]).strip())
        professors = list(self.db.scalars(
            select(TeacherEvaluation.teacher_name)
            .join(Course, TeacherEvaluation.course_id == Course.id)
            .join(Discipline, TeacherEvaluation.discipline_id == Discipline.id)
            .where(*professor_conditions)
            .distinct()
            .order_by(TeacherEvaluation.teacher_name)
        ).all())
        return {
            "periodos": periods,
            "cursos": courses,
            "disciplinas": disciplines,
            "professores": professors,
        }

    @staticmethod
    def _teacher_weighted_value(weighted_sum: Any, respondents: Any) -> float | None:
        total = int(respondents or 0)
        if total <= 0 or weighted_sum is None:
            return None
        return round(float(weighted_sum) / total, 4)

    def teacher_evaluation_analysis(self, filters: dict[str, str] | None = None) -> dict[str, Any]:
        """Faculty Evaluation KPI cards/trend/comparison computed in SQL."""
        filters = filters or {}
        selected_conditions = self._teacher_evaluation_conditions(filters, include_period=True)
        history_conditions = self._teacher_evaluation_conditions(filters, include_period=False)

        def aggregate(conditions: list[Any]) -> dict[str, Any]:
            row = self.db.execute(
                select(
                    func.sum(TeacherEvaluation.average_score * TeacherEvaluation.respondents),
                    func.sum(TeacherEvaluation.respondents),
                    func.count(func.distinct(TeacherEvaluation.teacher_name)),
                    func.count(func.distinct(TeacherEvaluation.discipline_id)),
                    func.count(TeacherEvaluation.id),
                )
                .join(Course, TeacherEvaluation.course_id == Course.id)
                .join(Discipline, TeacherEvaluation.discipline_id == Discipline.id)
                .where(*conditions)
            ).one()
            return {
                "valor": self._teacher_weighted_value(row[0], row[1]),
                "respondentes": int(row[1] or 0),
                "professores": int(row[2] or 0),
                "disciplinas": int(row[3] or 0),
                "registros": int(row[4] or 0),
            }

        trend_rows = self.db.execute(
            select(
                TeacherEvaluation.period,
                func.sum(TeacherEvaluation.average_score * TeacherEvaluation.respondents),
                func.sum(TeacherEvaluation.respondents),
                func.count(func.distinct(TeacherEvaluation.teacher_name)),
                func.count(func.distinct(TeacherEvaluation.discipline_id)),
            )
            .join(Course, TeacherEvaluation.course_id == Course.id)
            .join(Discipline, TeacherEvaluation.discipline_id == Discipline.id)
            .where(*history_conditions)
            .group_by(TeacherEvaluation.period)
            .order_by(TeacherEvaluation.period)
        ).all()
        trend = [
            {
                "periodo": period,
                "valor": self._teacher_weighted_value(weighted, respondents),
                "respondentes": int(respondents or 0),
                "professores": int(teachers or 0),
                "disciplinas": int(disciplines or 0),
            }
            for period, weighted, respondents, teachers, disciplines in trend_rows
        ]

        comparison_conditions = selected_conditions if filters.get("periodo") else history_conditions
        group_expr = Discipline.name if filters.get("professor") else TeacherEvaluation.teacher_name
        comparison_weighted = func.sum(TeacherEvaluation.average_score * TeacherEvaluation.respondents)
        comparison_respondents = func.sum(TeacherEvaluation.respondents)
        comparison_rows = self.db.execute(
            select(group_expr, comparison_weighted, comparison_respondents)
            .join(Course, TeacherEvaluation.course_id == Course.id)
            .join(Discipline, TeacherEvaluation.discipline_id == Discipline.id)
            .where(*comparison_conditions)
            .group_by(group_expr)
            .order_by((comparison_weighted / func.nullif(comparison_respondents, 0)).desc())
            .limit(20)
        ).all()
        comparison = [
            {
                "curso": str(label),
                "valor": self._teacher_weighted_value(weighted, respondents),
                "respondentes": int(respondents or 0),
                "status": "informativo",
            }
            for label, weighted, respondents in comparison_rows
            if label
        ]
        comparison = [row for row in comparison if row["valor"] is not None]

        return {
            "summary": aggregate(selected_conditions),
            "trend": trend,
            "comparison": comparison,
            "comparison_dimension": "disciplina" if filters.get("professor") else "professor",
        }

    def list_results_page(
        self,
        filters: dict[str, str] | None = None,
        *,
        page: int = 1,
        page_size: int = 50,
    ) -> dict[str, Any]:
        """Lista resultados acadêmicos com filtros e paginação no banco.

        A base aluno-disciplina pode crescer para dezenas de milhares de linhas.
        Por isso, esta consulta nunca carrega toda a tabela para o navegador.
        """
        filters = filters or {}
        page = max(1, int(page or 1))
        page_size = min(100, max(10, int(page_size or 50)))

        conditions = [AcademicResult.directorate_id == self.directorate_id]
        if filters.get("periodo"):
            conditions.append(AcademicResult.period == self._semester_from_period(filters["periodo"]))
        if filters.get("curso"):
            conditions.append(Course.name == str(filters["curso"]).strip())
        if filters.get("disciplina"):
            conditions.append(Discipline.name == str(filters["disciplina"]).strip())
        if filters.get("busca"):
            like = f"%{str(filters['busca']).strip()}%"
            conditions.append(or_(
                AcademicStudent.name.ilike(like),
                AcademicStudent.registration.ilike(like),
                Course.name.ilike(like),
                Discipline.name.ilike(like),
                AcademicResult.class_group.ilike(like),
                AcademicResult.official_status.ilike(like),
                AcademicResult.failure_reason.ilike(like),
            ))

        base_join = (
            select(AcademicResult)
            .join(Course, AcademicResult.course_id == Course.id)
            .join(Discipline, AcademicResult.discipline_id == Discipline.id)
            .join(AcademicStudent, AcademicResult.student_id == AcademicStudent.id)
            .where(*conditions)
        )
        total_q = (
            select(func.count(AcademicResult.id))
            .join(Course, AcademicResult.course_id == Course.id)
            .join(Discipline, AcademicResult.discipline_id == Discipline.id)
            .join(AcademicStudent, AcademicResult.student_id == AcademicStudent.id)
            .where(*conditions)
        )
        total = int(self.db.scalar(total_q) or 0)
        pages = max(1, (total + page_size - 1) // page_size)
        page = min(page, pages)
        offset = (page - 1) * page_size

        q = (
            base_join
            .options(
                joinedload(AcademicResult.course),
                joinedload(AcademicResult.discipline),
                joinedload(AcademicResult.student),
            )
            .order_by(
                AcademicResult.period.desc(),
                Course.name.asc(),
                Discipline.name.asc(),
                AcademicStudent.name.asc(),
                AcademicResult.id.desc(),
            )
            .offset(offset)
            .limit(page_size)
        )
        rows = self.db.scalars(q).all()
        return {
            "items": [self._record_to_dict("resultados", row) for row in rows],
            "total": total,
            "page": page,
            "page_size": page_size,
            "pages": pages,
        }

    def academic_result_summary(self, filters: dict[str, str] | None = None) -> list[dict[str, Any]]:
        """Consolida resultados no banco por semestre, curso e disciplina.

        O dashboard e a visão-resumo usam somente estes agregados. Nenhum nome,
        matrícula ou linha individual de aluno é transferido para esses cálculos.
        """
        filters = filters or {}
        conditions = [AcademicResult.directorate_id == self.directorate_id]
        if filters.get("periodo"):
            conditions.append(AcademicResult.period == self._semester_from_period(filters["periodo"]))
        if filters.get("curso"):
            conditions.append(Course.name == str(filters["curso"]).strip())
        if filters.get("disciplina"):
            conditions.append(Discipline.name == str(filters["disciplina"]).strip())

        finalized_case = case((AcademicResult.approved.is_not(None), 1), else_=0)
        approved_case = case((AcademicResult.approved.is_(True), 1), else_=0)
        failed_grade_case = case((and_(AcademicResult.approved.is_(False), AcademicResult.failure_reason == "nota"), 1), else_=0)
        failed_absence_case = case((and_(AcademicResult.approved.is_(False), AcademicResult.failure_reason == "falta"), 1), else_=0)
        failed_other_case = case((and_(
            AcademicResult.approved.is_(False),
            or_(AcademicResult.failure_reason.is_(None), ~AcademicResult.failure_reason.in_(["nota", "falta"])),
        ), 1), else_=0)

        q = (
            select(
                AcademicResult.period.label("periodo"),
                Course.name.label("curso"),
                Discipline.name.label("disciplina"),
                func.count(AcademicResult.id).label("total_registros"),
                func.sum(finalized_case).label("finalizados"),
                func.sum(approved_case).label("aprovados"),
                func.sum(failed_grade_case).label("reprovados_nota"),
                func.sum(failed_absence_case).label("reprovados_falta"),
                func.sum(failed_other_case).label("reprovados_outro"),
                func.count(AcademicResult.final_average).label("notas_contagem"),
                func.sum(AcademicResult.final_average).label("soma_notas"),
                func.avg(AcademicResult.final_average).label("media_notas"),
            )
            .join(Course, AcademicResult.course_id == Course.id)
            .join(Discipline, AcademicResult.discipline_id == Discipline.id)
            .where(*conditions)
            .group_by(AcademicResult.period, Course.name, Discipline.name)
            .order_by(AcademicResult.period.desc(), Course.name.asc(), Discipline.name.asc())
        )
        rows = self.db.execute(q).all()
        result: list[dict[str, Any]] = []
        for row in rows:
            total = int(row.total_registros or 0)
            finalized = int(row.finalizados or 0)
            approved = int(row.aprovados or 0)
            by_grade = int(row.reprovados_nota or 0)
            by_absence = int(row.reprovados_falta or 0)
            other = int(row.reprovados_outro or 0)
            result.append({
                "agregado": True,
                "periodo": row.periodo,
                "curso": row.curso,
                "disciplina": row.disciplina,
                "total_registros": total,
                "finalizados": finalized,
                "aprovados": approved,
                "reprovados": by_grade + by_absence + other,
                "reprovados_nota": by_grade,
                "reprovados_falta": by_absence,
                "reprovados_outro": other,
                "em_andamento": max(0, total - finalized),
                "notas_contagem": int(row.notas_contagem or 0),
                "soma_notas": float(row.soma_notas) if row.soma_notas is not None else 0.0,
                "media_notas": round(float(row.media_notas), 2) if row.media_notas is not None else None,
                "taxa_aprovacao": round(approved / finalized * 100, 2) if finalized else None,
                "validacao": "OK",
            })
        return result

    def academic_result_trend(self, filters: dict[str, str] | None = None) -> list[dict[str, Any]]:
        """Série semestral de aprovação no recorte acadêmico.

        Diferentemente do resumo por disciplina, esta consulta agrega diretamente
        por semestre. Isso permite contar estudantes distintos sem somar o mesmo
        aluno várias vezes quando o recorte inclui mais de uma disciplina.
        """
        filters = filters or {}
        conditions = [AcademicResult.directorate_id == self.directorate_id]
        if filters.get("curso"):
            conditions.append(Course.name == str(filters["curso"]).strip())
        if filters.get("disciplina"):
            conditions.append(Discipline.name == str(filters["disciplina"]).strip())

        finalized_case = case((AcademicResult.approved.is_not(None), 1), else_=0)
        approved_case = case((AcademicResult.approved.is_(True), 1), else_=0)
        approved_student_case = case(
            (AcademicResult.approved.is_(True), AcademicResult.student_id),
            else_=None,
        )
        finalized_student_case = case(
            (AcademicResult.approved.is_not(None), AcademicResult.student_id),
            else_=None,
        )

        q = (
            select(
                AcademicResult.period.label("periodo"),
                func.sum(finalized_case).label("finalizados"),
                func.sum(approved_case).label("aprovacoes"),
                func.count(func.distinct(approved_student_case)).label("alunos_aprovados"),
                func.count(func.distinct(finalized_student_case)).label("alunos_finalizados"),
            )
            .join(Course, AcademicResult.course_id == Course.id)
            .join(Discipline, AcademicResult.discipline_id == Discipline.id)
            .where(*conditions)
            .group_by(AcademicResult.period)
            .order_by(AcademicResult.period.asc())
        )
        rows = self.db.execute(q).all()
        result: list[dict[str, Any]] = []
        for row in rows:
            finalized = int(row.finalizados or 0)
            approvals = int(row.aprovacoes or 0)
            result.append({
                "periodo": row.periodo,
                "finalizados": finalized,
                "aprovados": approvals,
                "reprovados": max(0, finalized - approvals),
                "taxa_aprovacao": round(approvals / finalized * 100, 2) if finalized else None,
                "alunos_aprovados": int(row.alunos_aprovados or 0),
                "alunos_finalizados": int(row.alunos_finalizados or 0),
            })
        return result

    def teacher_evaluation_dashboard_summary(self) -> list[dict[str, Any]]:
        """Aggregate Faculty Evaluation for the executive dashboard in SQL.

        The dashboard never needs professor-level rows; it needs weighted scores by
        semester/course/discipline. Collapsing at this grain preserves every current
        dashboard calculation while keeping raw evaluation volume out of Python.
        """
        weighted = func.sum(TeacherEvaluation.average_score * TeacherEvaluation.respondents)
        respondents = func.sum(TeacherEvaluation.respondents)
        rows = self.db.execute(
            select(
                TeacherEvaluation.period,
                Course.name,
                Discipline.name,
                weighted,
                respondents,
            )
            .join(Course, TeacherEvaluation.course_id == Course.id)
            .join(Discipline, TeacherEvaluation.discipline_id == Discipline.id)
            .where(TeacherEvaluation.directorate_id == self.directorate_id)
            .group_by(TeacherEvaluation.period, Course.name, Discipline.name)
            .order_by(TeacherEvaluation.period.desc(), Course.name, Discipline.name)
        ).all()
        return [
            {
                "agregado": True,
                "periodo": period,
                "curso": course_name,
                "disciplina": discipline_name,
                "professor": "",
                "respondentes": int(total_respondents or 0),
                "nota_media": self._teacher_weighted_value(weighted_sum, total_respondents),
                "valor": self._teacher_weighted_value(weighted_sum, total_respondents),
                "validacao": "OK",
            }
            for period, course_name, discipline_name, weighted_sum, total_respondents in rows
        ]

    def academic_institution_nps_history(self) -> list[dict[str, Any]]:
        """NPS institucional dos alunos, agregado entre DTNH e DCS por semestre.

        O valor institucional representa a UNIVC e, por isso, não acompanha o
        filtro de curso da diretoria em visualização. A agregação usa as
        contagens reais de promotores/neutros/detratores das diretorias
        acadêmicas ativas, evitando média simples entre NPS já calculados.
        """
        academic_codes = ("DTNH", "DCS")
        directorates = list(self.db.execute(
            select(Directorate.id, Directorate.code)
            .where(Directorate.code.in_(academic_codes), Directorate.active.is_(True))
            .order_by(Directorate.code)
        ).all())
        directorate_ids = [row.id for row in directorates]
        if not directorate_ids:
            return []
        rows = list(self.db.execute(
            select(NpsInstitution, Directorate.code)
            .join(Directorate, Directorate.id == NpsInstitution.directorate_id)
            .where(NpsInstitution.directorate_id.in_(directorate_ids))
            .order_by(NpsInstitution.period, Directorate.code)
        ).all())
        grouped: dict[str, list[tuple[NpsInstitution, str]]] = {}
        for row, code in rows:
            grouped.setdefault(row.period, []).append((row, code))
        out: list[dict[str, Any]] = []
        total_directorates = len(directorates)
        for period, parts in grouped.items():
            respondents = sum(int(row.respondents or 0) for row, _ in parts)
            promoters = sum(int(row.promoters or 0) for row, _ in parts)
            neutrals = sum(int(row.neutrals or 0) for row, _ in parts)
            detractors = sum(int(row.detractors or 0) for row, _ in parts)
            value = round((promoters - detractors) / respondents * 100, 2) if respondents else None
            present_codes = sorted({code for _, code in parts})
            out.append({
                "periodo": period,
                "valor": value,
                "respondentes": respondents,
                "promotores": promoters,
                "neutros": neutrals,
                "detratores": detractors,
                "diretorias": present_codes,
                "coverage": len(present_codes),
                "coverage_total": total_directorates,
                "complete": len(present_codes) == total_directorates,
                "fonte": "SEI · NPS institucional",
            })
        return sorted(out, key=lambda item: item["periodo"])

    def academic_faculty_institution_nps_history(self) -> list[dict[str, Any]]:
        """NPS institucional respondido pelos docentes.

        A fonte e global para a UNIVC: o relatorio e anonimo e nao oferece
        curso/diretoria do respondente. DTNH e DCS consultam a mesma serie
        factual e aplicam suas proprias metas de gestao (01C).
        """
        rows = list(self.db.scalars(
            select(NpsInstitutionFaculty).order_by(NpsInstitutionFaculty.period)
        ).all())
        out: list[dict[str, Any]] = []
        for row in rows:
            respondents = int(row.respondents or 0)
            promoters = int(row.promoters or 0)
            neutrals = int(row.neutrals or 0)
            detractors = int(row.detractors or 0)
            out.append({
                "periodo": row.period,
                "valor": round((promoters - detractors) / respondents * 100, 2) if respondents else None,
                "respondentes": respondents,
                "promotores": promoters,
                "neutros": neutrals,
                "detratores": detractors,
                "fonte": "SEI · NPS institucional dos docentes",
                "anonymous_population": True,
            })
        return out

    def academic_dashboard_snapshot(self) -> dict[str, Any]:
        """Snapshot leve para DTNH/DCS, sem transportar linhas individuais de alunos."""
        courses, disciplines = self.course_maps(include_inactive=False)
        total_results = int(self.db.scalar(
            select(func.count(AcademicResult.id)).where(AcademicResult.directorate_id == self.directorate_id)
        ) or 0)
        return {
            "courses": courses,
            "disciplines": disciplines,
            "metas": self.get_metas(),
            "nps": self.list_records("nps"),
            "nps_institution": self.academic_institution_nps_history(),
            "nps_institution_faculty": self.academic_faculty_institution_nps_history(),
            "avaliacao_docente": self.teacher_evaluation_dashboard_summary(),
            "resultados": self.academic_result_summary(),
            "resultados_total": total_results,
        }

    def get_record(self, dataset: str, row_id: int):
        model = DATASET_MODEL.get(dataset)
        if not model:
            raise RepositoryError("Base de dados não encontrada.")
        row = require_object_for_directorate(self.db, model, row_id, self.directorate_id, label="Registro")
        return self._record_to_dict(dataset, row)

    def _find_existing_academic(self, dataset: str, clean: dict[str, Any], ignore_id: int | None = None):
        model = DATASET_MODEL[dataset]
        conditions = [model.directorate_id == self.directorate_id, model.period == clean["periodo"]]
        if dataset in {"nps", "matriculas"}:
            conditions.append(model.course_id == clean["course_id"])
        elif dataset == "avaliacao_docente":
            conditions.extend([
                model.course_id == clean["course_id"], model.discipline_id == clean["discipline_id"],
            ])
            if ignore_id is not None:
                conditions.append(model.id != ignore_id)
            candidates = self.db.scalars(select(model).where(*conditions)).all()
            professor_key = self._text_key(clean["professor"])
            return next((row for row in candidates if self._text_key(row.teacher_name) == professor_key), None)
        elif dataset == "resultados":
            # A duplicidade acadêmica deve representar o mesmo lançamento real.
            # Incluímos curso e turma na chave para não colapsar registros legítimos
            # quando o mesmo aluno aparece em mais de uma turma/oferta da disciplina.
            conditions.extend([
                model.course_id == clean["course_id"],
                model.discipline_id == clean["discipline_id"],
                model.student_id == clean["student_id"],
                model.class_group == clean["turma"],
            ])
        elif dataset == "frequencia":
            conditions.extend([model.course_id == clean["course_id"], model.discipline_id == clean["discipline_id"]])
        if ignore_id is not None:
            conditions.append(model.id != ignore_id)
        return self.db.scalar(select(model).where(*conditions))

    def create_record(self, dataset: str, payload: dict[str, Any]):
        clean = self.validate_payload(dataset, payload)
        if self._find_existing_academic(dataset, clean):
            unique_fields = ", ".join(DATASETS[dataset].unique_fields)
            raise ValidationError(
                "Este registro já existe e não foi duplicado.",
                {"periodo": f"Duplicata detectada pela chave: {unique_fields}."},
            )
        if dataset == "nps":
            row = NpsStudent(directorate_id=self.directorate_id, period=clean["periodo"], course_id=clean["course_id"], respondents=clean["respondentes"], promoters=clean["promotores"], neutrals=clean["neutros"], detractors=clean["detratores"], inserted_by=self.ctx.full_name)
        elif dataset == "avaliacao_docente":
            row = TeacherEvaluation(directorate_id=self.directorate_id, period=clean["periodo"], course_id=clean["course_id"], discipline_id=clean["discipline_id"], teacher_name=clean["professor"], respondents=clean["respondentes"], average_score=clean["nota_media"], inserted_by=self.ctx.full_name)
        elif dataset == "resultados":
            row = AcademicResult(
                directorate_id=self.directorate_id, period=clean["periodo"], course_id=clean["course_id"],
                discipline_id=clean["discipline_id"], student_id=clean["student_id"], class_group=clean["turma"],
                curriculum_period=clean["periodo_curricular"], final_average=clean["media"], official_status=clean["situacao"],
                approved=clean["aprovado"], failure_reason=clean["motivo_reprovacao"], source=clean["fonte"],
                inserted_by=self.ctx.full_name,
            )
        elif dataset == "matriculas":
            row = Enrollment(directorate_id=self.directorate_id, period=clean["periodo"], course_id=clean["course_id"], active_enrollments=clean["matriculas"], inserted_by=self.ctx.full_name)
        else:
            row = Attendance(directorate_id=self.directorate_id, period=clean["periodo"], course_id=clean["course_id"], discipline_id=clean["discipline_id"], expected_attendance=clean["presencas_previstas"], recorded_attendance=clean["presencas_registradas"], inserted_by=self.ctx.full_name)
        self.db.add(row)
        self.db.flush()
        self._audit("create", dataset, row.id, {key: value for key, value in payload.items() if key not in {"senha", "password"}})
        self._commit()
        return self.get_record(dataset, row.id)

    def update_record(self, dataset: str, row_id: int, payload: dict[str, Any]):
        model = DATASET_MODEL.get(dataset)
        if not model:
            raise RepositoryError("Base de dados não encontrada.")
        row = require_object_for_directorate(self.db, model, row_id, self.directorate_id, label="Registro")
        clean = self.validate_payload(dataset, payload, ignore_id=row_id)
        if self._find_existing_academic(dataset, clean, ignore_id=row_id):
            raise ValidationError("A alteração criaria um registro duplicado.", {"periodo": "Já existe outro registro com esta chave."})
        row.period = clean["periodo"]
        row.course_id = clean["course_id"]
        if dataset == "nps":
            row.respondents, row.promoters, row.neutrals, row.detractors = clean["respondentes"], clean["promotores"], clean["neutros"], clean["detratores"]
        elif dataset == "avaliacao_docente":
            row.discipline_id = clean["discipline_id"]
            row.teacher_name = clean["professor"]
            row.respondents = clean["respondentes"]
            row.average_score = clean["nota_media"]
        elif dataset == "resultados":
            row.discipline_id = clean["discipline_id"]
            row.student_id = clean["student_id"]
            row.class_group = clean["turma"]
            row.curriculum_period = clean["periodo_curricular"]
            row.final_average = clean["media"]
            row.official_status = clean["situacao"]
            row.approved = clean["aprovado"]
            row.failure_reason = clean["motivo_reprovacao"]
            row.source = clean["fonte"]
        elif dataset == "matriculas":
            row.active_enrollments = clean["matriculas"]
        else:
            row.discipline_id, row.expected_attendance, row.recorded_attendance = clean["discipline_id"], clean["presencas_previstas"], clean["presencas_registradas"]
        row.inserted_by = self.ctx.full_name
        self._audit("update", dataset, row.id, {key: value for key, value in payload.items() if key not in {"senha", "password"}})
        self._commit()
        return self.get_record(dataset, row.id)

    def delete_record(self, dataset: str, row_id: int):
        model = DATASET_MODEL.get(dataset)
        if not model:
            raise RepositoryError("Base de dados não encontrada.")
        row = require_object_for_directorate(self.db, model, row_id, self.directorate_id, label="Registro")
        self.db.delete(row)
        self._audit("delete", dataset, row_id)
        self._commit()

    def list_goals(self):
        rows = self.db.scalars(
            select(Goal)
            .where(Goal.directorate_id == self.directorate_id)
            .order_by(Goal.indicator_code, Goal.scope_label, Goal.valid_from)
        ).all()
        out = []
        for r in rows:
            meta = KPI_META.get(r.indicator_code, {})
            unit = meta.get("unit") or ("%" if r.indicator_code.endswith(("-02", "-03")) else "pontos")
            out.append({
                "id": r.id, "indicador": r.indicator_code, "recorte": r.scope_label,
                "vigencia": r.valid_from, "meta": r.target, "atencao": r.attention,
                "limite_superior": r.upper_limit, "justificativa": r.justification or "",
                "unidade": unit, "direcao": meta.get("direction", "higher"),
            })
        return out

    def get_metas(self):
        return [{k: v for k, v in x.items() if k != "id" and k != "unidade"} for x in self.list_goals()]

    def _validate_goal(self, payload, ignore_id=None):
        code = str(payload.get("indicador") or "").strip(); recorte = str(payload.get("recorte") or "").strip(); vigencia = str(payload.get("vigencia") or "").strip(); errors = {}
        if not code: errors["indicador"] = "Selecione um KPI."
        elif code not in IMPLEMENTED_KPIS.get(self.directorate_code, ()): errors["indicador"] = "Selecione um KPI implementado para esta diretoria."
        if not recorte:
            errors["recorte"] = "Informe o recorte."
        if self.directorate_code in {"DTNH", "DCS"} and recorte:
            if code.endswith(("-01A", "-01C")) and recorte != "TOTAL":
                label = "NPS da Instituição" if code.endswith("-01A") else "NPS da Instituição · Docentes"
                errors["recorte"] = f"O {label} usa somente o recorte TOTAL."
            elif code.endswith("-01B"):
                courses, _ = self.course_maps(include_inactive=False)
                if recorte != "TOTAL" and recorte not in set(courses):
                    errors["recorte"] = "O NPS do Curso aceita TOTAL ou um curso ativo desta diretoria."
        if not (MONTH_RE.match(vigencia) or (self.directorate_code in {"DTNH", "DCS"} and SEMESTER_RE.match(vigencia.upper()))): errors["vigencia"] = "Use AAAA-MM ou, para DTNH/DCS, AAAA-SEM1/AAAA-SEM2."
        clean = {"indicador": code, "recorte": recorte, "vigencia": vigencia, "justificativa": str(payload.get("justificativa") or "").strip()}
        for field in ("meta", "atencao", "limite_superior"):
            val = payload.get(field)
            if val in (None, ""):
                clean[field] = None
                if field != "limite_superior": errors[field] = "Campo obrigatório."
            else:
                try: clean[field] = float(val)
                except (TypeError, ValueError): errors[field] = "Informe um número."
        if code == "DPE-04" and clean.get("limite_superior") is None:
            errors["limite_superior"] = "Para DPE-04, informe o limite superior da faixa (ex.: 105)."
        if code == "DPE-04" and clean.get("meta") is not None and clean.get("limite_superior") is not None and clean["limite_superior"] <= clean["meta"]:
            errors["limite_superior"] = "O limite superior deve ser maior que o limite inferior da meta."
        if errors: raise ValidationError("Revise os campos destacados.", errors)
        return clean

    def create_goal(self, payload):
        c = self._validate_goal(payload); row = Goal(directorate_id=self.directorate_id, indicator_code=c["indicador"], scope_label=c["recorte"], valid_from=c["vigencia"], target=c["meta"], attention=c["atencao"], upper_limit=c["limite_superior"], justification=c["justificativa"])
        self.db.add(row); self.db.flush(); self._audit("create", "goal", row.id, payload); self._commit(); return next(x for x in self.list_goals() if x["id"] == row.id)

    def update_goal(self, row_id, payload):
        row = require_object_for_directorate(self.db, Goal, row_id, self.directorate_id, label="Meta")
        c = self._validate_goal(payload, row_id); row.indicator_code=c["indicador"]; row.scope_label=c["recorte"]; row.valid_from=c["vigencia"]; row.target=c["meta"]; row.attention=c["atencao"]; row.upper_limit=c["limite_superior"]; row.justification=c["justificativa"]
        self._audit("update", "goal", row.id, payload); self._commit(); return next(x for x in self.list_goals() if x["id"] == row.id)

    def delete_goal(self, row_id):
        row = require_object_for_directorate(self.db, Goal, row_id, self.directorate_id, label="Meta")
        self.db.delete(row); self._audit("delete", "goal", row_id); self._commit()

    def list_actions(self):
        rows = self.db.scalars(select(ActionPlan).where(ActionPlan.directorate_id == self.directorate_id).order_by(ActionPlan.status, ActionPlan.due_date)).all(); out=[]
        for r in rows:
            dias = None if r.status == "Concluído" else (r.due_date - date.today()).days
            out.append({"id":r.id,"numero":r.number,"mes":r.month,"indicador":r.indicator_code,"recorte":r.scope_label,"resultado":r.result_value,"meta":r.target_value,"problema":r.problem,"causa":r.probable_cause or "","acao":r.corrective_action,"responsavel":r.responsible,"prazo":r.due_date.isoformat(),"meta_acao":r.action_target or "","status":r.status,"dias_prazo":dias})
        return out

    def _validate_action(self,payload):
        clean={k:(str(payload.get(k)).strip() if isinstance(payload.get(k),str) else payload.get(k)) for k in ("mes","indicador","recorte","resultado","meta","problema","causa","acao","responsavel","prazo","meta_acao","status")}; errors={}
        for f in ("mes","indicador","recorte","problema","acao","responsavel","prazo","status"):
            if not clean.get(f): errors[f]="Campo obrigatório."
        if clean.get("indicador") and clean["indicador"] not in IMPLEMENTED_KPIS.get(self.directorate_code, ()): errors["indicador"]="Selecione um KPI implementado para esta diretoria."
        if clean.get("mes") and not MONTH_RE.match(str(clean["mes"])): errors["mes"]="Use o formato AAAA-MM."
        if clean.get("status") not in {"Não iniciado","Em andamento","Concluído","Atrasado","Cancelado"}: errors["status"]="Status inválido."
        try: clean["prazo_date"] = datetime.strptime(str(clean.get("prazo"))[:10], "%Y-%m-%d").date()
        except Exception: errors["prazo"]="Use uma data válida."
        for f in ("resultado","meta"):
            if clean.get(f) in (None,""): clean[f]=None
            else:
                try: clean[f]=float(clean[f])
                except Exception: errors[f]="Informe um número."
        if errors: raise ValidationError("Revise os campos destacados.",errors)
        return clean

    def create_action(self,payload):
        c=self._validate_action(payload); maxn=self.db.scalar(select(func.max(ActionPlan.number)).where(ActionPlan.directorate_id==self.directorate_id)) or 0
        r=ActionPlan(directorate_id=self.directorate_id,number=maxn+1,month=c["mes"],indicator_code=c["indicador"],scope_label=c["recorte"],result_value=c["resultado"],target_value=c["meta"],problem=c["problema"],probable_cause=c["causa"],corrective_action=c["acao"],responsible=c["responsavel"],due_date=c["prazo_date"],action_target=c["meta_acao"],status=c["status"])
        self.db.add(r); self.db.flush(); self._audit("create","action_plan",r.id,payload); self._commit(); return next(x for x in self.list_actions() if x["id"]==r.id)

    def update_action(self,row_id,payload):
        r=require_object_for_directorate(self.db, ActionPlan, row_id, self.directorate_id, label="Plano de ação")
        c=self._validate_action(payload); r.month=c["mes"];r.indicator_code=c["indicador"];r.scope_label=c["recorte"];r.result_value=c["resultado"];r.target_value=c["meta"];r.problem=c["problema"];r.probable_cause=c["causa"];r.corrective_action=c["acao"];r.responsible=c["responsavel"];r.due_date=c["prazo_date"];r.action_target=c["meta_acao"];r.status=c["status"]
        self._audit("update","action_plan",r.id,payload);self._commit();return next(x for x in self.list_actions() if x["id"]==r.id)

    def delete_action(self,row_id):
        r=require_object_for_directorate(self.db, ActionPlan, row_id, self.directorate_id, label="Plano de ação")
        self.db.delete(r);self._audit("delete","action_plan",row_id);self._commit()

    # ------------------------------------------------------------------
    # DADM: operational data for the three first administrative KPIs.
    # These methods intentionally stay explicit instead of forcing the
    # administrative structures into the academic DATASETS abstraction.
    # ------------------------------------------------------------------
    def _ensure_dadm(self):
        if self.directorate_code != "DADM":
            raise RepositoryError("Este conjunto de dados pertence à Diretoria Administrativa.")

    def _to_decimal(self, value: Any, field: str, errors: dict, allow_zero: bool = True) -> Decimal | None:
        try:
            number = Decimal(str(value).replace("R$", "").replace(" ", "").replace(".", "").replace(",", ".") if isinstance(value, str) and "," in value else str(value))
        except (InvalidOperation, ValueError, TypeError):
            errors[field] = "Informe um número válido."
            return None
        if number < 0 or (not allow_zero and number == 0):
            errors[field] = "O valor deve ser maior que zero." if not allow_zero else "O valor não pode ser negativo."
        return number

    def _normalize_modality(self, value: Any) -> str:
        text = str(value or "").strip()
        folded = norm_header(text)
        aliases = {
            "presencial": "Presencial",
            "semipresencial": "Semipresencial",
            "semi presencial": "Semipresencial",
            "ead": "EAD",
            "educacao a distancia": "EAD",
            "ensino a distancia": "EAD",
        }
        return aliases.get(folded, text[:80])

    @staticmethod
    def _previous_month(period: str) -> str | None:
        """Return the immediately previous calendar month for an AAAA-MM value."""
        if not MONTH_RE.match(str(period or "")):
            return None
        year, month = map(int, str(period).split("-"))
        if month == 1:
            return f"{year - 1:04d}-12"
        return f"{year:04d}-{month - 1:02d}"

    def list_academic_courses_for_dadm(self) -> list[dict[str, Any]]:
        rows = self.db.execute(
            select(Course, Directorate)
            .join(Directorate, Directorate.id == Course.directorate_id)
            .where(Course.active.is_(True), Directorate.active.is_(True), Directorate.code.in_(ACADEMIC_DIRECTORATES))
            .order_by(Directorate.code, Course.name)
        ).all()
        return [
            {
                "id": course.id,
                "curso": course.name,
                "diretoria": directorate.code,
                "diretoria_nome": directorate.name,
                "modalidade": course.modality or "Presencial",
                "label": f"{directorate.code} · {course.name}",
            }
            for course, directorate in rows
        ]

    def _academic_enrollment_population(self) -> list[dict[str, Any]]:
        """Build the administrative student base from academic enrollment closings.

        The academic directorates remain the owners of enrollment data. DADM consumes
        those closings instead of asking users to re-enter the same population.
        """
        rows = self.db.execute(
            select(Enrollment, Course, Directorate)
            .join(Course, Course.id == Enrollment.course_id)
            .join(Directorate, Directorate.id == Course.directorate_id)
            .where(
                Directorate.code.in_(ACADEMIC_DIRECTORATES),
                Directorate.active.is_(True),
                Course.active.is_(True),
            )
            .order_by(Enrollment.period.desc(), Directorate.code, Course.name)
        ).all()
        return [
            {
                "id": f"academic-{enrollment.id}",
                "periodo": enrollment.period,
                "course_id": course.id,
                "diretoria_academica": directorate.code,
                "curso": course.name,
                "modalidade": course.modality or "Presencial",
                "alunos_ativos": enrollment.active_enrollments,
                "valor": enrollment.active_enrollments,
                "origem": f"{directorate.code} · Matrículas Ativas",
                "temporario": False,
                "validacao": "OK",
                "data_lancamento": str(enrollment.inserted_at.date()) if enrollment.inserted_at else "",
                "lancado_por": enrollment.inserted_by or "",
            }
            for enrollment, course, directorate in rows
        ]

    def _fallback_population(self) -> list[dict[str, Any]]:
        """Return temporary modality totals used only while a modality has no academic source.

        `dadm_active_students` is kept for backward compatibility and demonstration data,
        but it is no longer an operational DADM input. As soon as academic enrollment
        rows exist for a period/modality, the fallback is ignored automatically.
        """
        rows = self.db.scalars(
            select(DadmActiveStudents)
            .where(DadmActiveStudents.directorate_id == self.directorate_id)
            .order_by(DadmActiveStudents.period.desc(), DadmActiveStudents.modality)
        ).all()
        return [
            {
                "id": f"fallback-{row.id}",
                "periodo": row.period,
                "course_id": None,
                "diretoria_academica": "DEMO",
                "curso": "Base temporária",
                "modalidade": row.modality,
                "alunos_ativos": row.active_students,
                "valor": row.active_students,
                "origem": "Base demonstrativa temporária",
                "temporario": True,
                "validacao": "OK",
                "data_lancamento": str(row.inserted_at.date()) if row.inserted_at else "",
                "lancado_por": row.inserted_by or "",
            }
            for row in rows
        ]

    def list_dadm_population(self) -> list[dict[str, Any]]:
        """Unified population consumed by DADM-01/09/10.

        Academic enrollment rows win. Temporary DADM rows are included only for a
        period/modality pair that still has no academic enrollment source (currently
        useful to demonstrate EAD and Semipresencial before DEAD is operational).
        """
        academic = self._academic_enrollment_population()
        covered = {
            (str(row.get("periodo") or ""), self._normalize_modality(row.get("modalidade")))
            for row in academic
        }
        fallback = [
            row for row in self._fallback_population()
            if (str(row.get("periodo") or ""), self._normalize_modality(row.get("modalidade"))) not in covered
        ]
        return sorted(
            [*academic, *fallback],
            key=lambda row: (
                str(row.get("periodo") or ""),
                str(row.get("diretoria_academica") or ""),
                str(row.get("curso") or ""),
            ),
            reverse=True,
        )

    def dadm_student_base(self, period: str, course_id: int) -> dict[str, Any]:
        """Resolve DADM-01 denominator from the selected course's previous month closing."""
        self._ensure_dadm()
        previous = self._previous_month(period)
        if not previous:
            raise ValidationError("Revise os campos destacados.", {"periodo": "Selecione uma competência mensal válida."})
        course = self.db.get(Course, int(course_id)) if course_id else None
        if not course or not course.active:
            raise ValidationError("Revise os campos destacados.", {"course_id": "Selecione um curso acadêmico ativo."})
        directorate = self.db.get(Directorate, course.directorate_id)
        if not directorate or directorate.code not in ACADEMIC_DIRECTORATES or not directorate.active:
            raise ValidationError("Revise os campos destacados.", {"course_id": "Selecione um curso de uma diretoria acadêmica ativa."})
        enrollment = self.db.scalar(
            select(Enrollment).where(
                Enrollment.course_id == course.id,
                Enrollment.directorate_id == course.directorate_id,
                Enrollment.period == previous,
            )
        )
        if not enrollment:
            raise ValidationError(
                "Não há base acadêmica para calcular a evasão.",
                {"course_id": f"Cadastre Matrículas Ativas de {previous} para este curso antes do fechamento de {period}."},
            )
        return {
            "periodo_base": previous,
            "alunos_inicio": int(enrollment.active_enrollments),
            "course_id": course.id,
            "curso": course.name,
            "diretoria_academica": directorate.code,
            "modalidade": course.modality or "Presencial",
            "origem": f"{directorate.code} · Matrículas Ativas ({previous})",
        }

    def dadm_reference_options(self) -> dict[str, Any]:
        self._ensure_dadm()
        modalities = set(DADM_MODALITIES)
        for row in self.list_dadm_population():
            if row.get("modalidade"):
                modalities.add(str(row["modalidade"]))
        for row in self.db.scalars(
            select(DadmAdministrativeCost).where(DadmAdministrativeCost.directorate_id == self.directorate_id)
        ).all():
            if row.modality:
                modalities.add(row.modality)
        cost_centers = self.db.scalars(
            select(DadmAdministrativeCost.cost_center)
            .where(DadmAdministrativeCost.directorate_id == self.directorate_id)
            .distinct()
            .order_by(DadmAdministrativeCost.cost_center)
        ).all()
        return {
            "cursos_academicos": self.list_academic_courses_for_dadm(),
            "modalidades": sorted(modalities),
            "centros_custo": [x for x in cost_centers if x],
        }

    def _validate_dadm_attrition(self, payload: dict[str, Any]) -> dict[str, Any]:
        self._ensure_dadm()
        errors: dict[str, str] = {}
        period = str(payload.get("periodo") or "").strip()
        if not MONTH_RE.match(period):
            errors["periodo"] = "Selecione uma competência mensal válida (AAAA-MM)."

        course = None
        course_id = payload.get("course_id")
        if course_id not in (None, ""):
            try:
                course = self.db.get(Course, int(course_id))
            except (TypeError, ValueError):
                course = None
        else:
            course_name = str(payload.get("curso") or "").strip()
            directorate_code = str(payload.get("diretoria_academica") or "").strip().upper()
            q = select(Course).join(Directorate).where(
                Course.active.is_(True),
                func.lower(Course.name) == course_name.lower(),
                Directorate.code.in_(ACADEMIC_DIRECTORATES),
            )
            if directorate_code:
                q = q.where(Directorate.code == directorate_code)
            matches = self.db.scalars(q).all()
            if len(matches) == 1:
                course = matches[0]
            elif len(matches) > 1:
                errors["diretoria_academica"] = "Há cursos com o mesmo nome em mais de uma diretoria; informe a diretoria acadêmica."

        if not course or not course.active:
            errors["course_id"] = "Selecione um curso acadêmico ativo."
        elif course.directorate_id == self.directorate_id:
            errors["course_id"] = "A evasão deve ser vinculada a um curso de uma diretoria acadêmica."

        departures = self._to_int(payload.get("desligamentos"), "desligamentos", errors)
        if errors:
            raise ValidationError("Revise os campos destacados.", errors)
        base = self.dadm_student_base(period, course.id)
        students = base["alunos_inicio"]
        if departures is not None and departures > students:
            raise ValidationError(
                "Revise os campos destacados.",
                {"desligamentos": "Os desligamentos não podem superar a base automática de alunos do início do mês."},
            )
        return {
            "periodo": period,
            "course_id": course.id,
            "alunos_inicio": students,
            "desligamentos": departures,
            "periodo_base": base["periodo_base"],
            "modalidade": base["modalidade"],
        }

    def _dadm_attrition_to_dict(self, row: DadmAttrition) -> dict[str, Any]:
        directorate = self.db.get(Directorate, row.course.directorate_id)
        value = round(row.departures / row.active_students_start * 100, 2) if row.active_students_start else None
        temporary = bool(not directorate or not directorate.active or str(row.course.name).startswith("Base demonstrativa"))
        return {
            "id": row.id,
            "periodo": row.period,
            "course_id": row.course_id,
            "diretoria_academica": directorate.code if directorate else "",
            "curso": row.course.name,
            "modalidade": row.course.modality or "Presencial",
            "temporario": temporary,
            "origem": "Base demonstrativa temporária de evasão" if temporary else f"{directorate.code} · desligamentos",
            "periodo_base": self._previous_month(row.period),
            "alunos_inicio": row.active_students_start,
            "desligamentos": row.departures,
            "valor": value,
            "validacao": "OK",
            "data_lancamento": str(row.inserted_at.date()) if row.inserted_at else "",
            "lancado_por": row.inserted_by or "",
        }

    def list_dadm_attrition(self) -> list[dict[str, Any]]:
        self._ensure_dadm()
        rows = self.db.scalars(
            select(DadmAttrition)
            .options(joinedload(DadmAttrition.course))
            .join(Course)
            .where(DadmAttrition.directorate_id == self.directorate_id)
            .order_by(DadmAttrition.period.desc(), Course.name)
        ).all()
        return [self._dadm_attrition_to_dict(row) for row in rows]

    def create_dadm_attrition(self, payload: dict[str, Any]) -> dict[str, Any]:
        c = self._validate_dadm_attrition(payload)
        row = DadmAttrition(
            directorate_id=self.directorate_id,
            period=c["periodo"],
            course_id=c["course_id"],
            active_students_start=c["alunos_inicio"],
            departures=c["desligamentos"],
            inserted_by=self.ctx.full_name,
        )
        self.db.add(row); self.db.flush(); self._audit("create", "dadm_attrition", row.id, payload); self._commit()
        return self._dadm_attrition_to_dict(row)

    def update_dadm_attrition(self, row_id: int, payload: dict[str, Any]) -> dict[str, Any]:
        row = require_object_for_directorate(self.db, DadmAttrition, row_id, self.directorate_id, label="Registro de evasão")
        c = self._validate_dadm_attrition(payload)
        row.period = c["periodo"]; row.course_id = c["course_id"]
        row.active_students_start = c["alunos_inicio"]; row.departures = c["desligamentos"]
        row.inserted_by = self.ctx.full_name
        self._audit("update", "dadm_attrition", row.id, payload); self._commit()
        return self._dadm_attrition_to_dict(row)

    def delete_dadm_attrition(self, row_id: int):
        row = require_object_for_directorate(self.db, DadmAttrition, row_id, self.directorate_id, label="Registro de evasão")
        self.db.delete(row); self._audit("delete", "dadm_attrition", row_id); self._commit()

    def _validate_dadm_students(self, payload: dict[str, Any]) -> dict[str, Any]:
        self._ensure_dadm(); errors: dict[str, str] = {}
        period = str(payload.get("periodo") or "").strip()
        modality = self._normalize_modality(payload.get("modalidade"))
        if not MONTH_RE.match(period): errors["periodo"] = "Selecione uma competência mensal válida (AAAA-MM)."
        if len(modality) < 2: errors["modalidade"] = "Informe a modalidade."
        students = self._to_int(payload.get("alunos_ativos"), "alunos_ativos", errors)
        if errors: raise ValidationError("Revise os campos destacados.", errors)
        return {"periodo": period, "modalidade": modality, "alunos_ativos": students}

    def _dadm_students_to_dict(self, row: DadmActiveStudents) -> dict[str, Any]:
        return {
            "id": row.id, "periodo": row.period, "modalidade": row.modality,
            "alunos_ativos": row.active_students, "valor": row.active_students, "validacao": "OK",
            "data_lancamento": str(row.inserted_at.date()) if row.inserted_at else "", "lancado_por": row.inserted_by or "",
        }

    def list_dadm_active_students(self) -> list[dict[str, Any]]:
        self._ensure_dadm()
        return self.list_dadm_population()

    def create_dadm_active_students(self, payload: dict[str, Any]) -> dict[str, Any]:
        c = self._validate_dadm_students(payload)
        row = DadmActiveStudents(
            directorate_id=self.directorate_id, period=c["periodo"], modality=c["modalidade"],
            active_students=c["alunos_ativos"], inserted_by=self.ctx.full_name,
        )
        self.db.add(row); self.db.flush(); self._audit("create", "dadm_active_students", row.id, payload); self._commit()
        return self._dadm_students_to_dict(row)

    def update_dadm_active_students(self, row_id: int, payload: dict[str, Any]) -> dict[str, Any]:
        row = require_object_for_directorate(self.db, DadmActiveStudents, row_id, self.directorate_id, label="População mensal")
        c = self._validate_dadm_students(payload)
        row.period = c["periodo"]; row.modality = c["modalidade"]; row.active_students = c["alunos_ativos"]; row.inserted_by = self.ctx.full_name
        self._audit("update", "dadm_active_students", row.id, payload); self._commit(); return self._dadm_students_to_dict(row)

    def delete_dadm_active_students(self, row_id: int):
        row = require_object_for_directorate(self.db, DadmActiveStudents, row_id, self.directorate_id, label="População mensal")
        self.db.delete(row); self._audit("delete", "dadm_active_students", row_id); self._commit()

    def _validate_dadm_cost(self, payload: dict[str, Any]) -> dict[str, Any]:
        self._ensure_dadm(); errors: dict[str, str] = {}
        period = str(payload.get("periodo") or "").strip()
        center = str(payload.get("centro_custo") or "").strip()[:160]
        modality = self._normalize_modality(payload.get("modalidade"))
        if not MONTH_RE.match(period): errors["periodo"] = "Selecione uma competência mensal válida (AAAA-MM)."
        if len(center) < 2: errors["centro_custo"] = "Informe o centro de custo administrativo."
        if len(modality) < 2: errors["modalidade"] = "Informe a modalidade."
        expense = self._to_decimal(payload.get("despesa"), "despesa", errors)
        if errors: raise ValidationError("Revise os campos destacados.", errors)
        return {"periodo": period, "centro_custo": center, "modalidade": modality, "despesa": expense}

    def _dadm_cost_to_dict(self, row: DadmAdministrativeCost) -> dict[str, Any]:
        return {
            "id": row.id, "periodo": row.period, "centro_custo": row.cost_center, "modalidade": row.modality,
            "despesa": float(row.expense_amount), "valor": float(row.expense_amount), "validacao": "OK",
            "data_lancamento": str(row.inserted_at.date()) if row.inserted_at else "", "lancado_por": row.inserted_by or "",
        }

    def list_dadm_admin_costs(self) -> list[dict[str, Any]]:
        self._ensure_dadm()
        rows = self.db.scalars(
            select(DadmAdministrativeCost).where(DadmAdministrativeCost.directorate_id == self.directorate_id)
            .order_by(DadmAdministrativeCost.period.desc(), DadmAdministrativeCost.cost_center, DadmAdministrativeCost.modality)
        ).all()
        return [self._dadm_cost_to_dict(row) for row in rows]

    def create_dadm_admin_cost(self, payload: dict[str, Any]) -> dict[str, Any]:
        c = self._validate_dadm_cost(payload)
        row = DadmAdministrativeCost(
            directorate_id=self.directorate_id, period=c["periodo"], cost_center=c["centro_custo"],
            modality=c["modalidade"], expense_amount=c["despesa"], inserted_by=self.ctx.full_name,
        )
        self.db.add(row); self.db.flush(); self._audit("create", "dadm_administrative_cost", row.id, payload); self._commit()
        return self._dadm_cost_to_dict(row)

    def update_dadm_admin_cost(self, row_id: int, payload: dict[str, Any]) -> dict[str, Any]:
        row = require_object_for_directorate(self.db, DadmAdministrativeCost, row_id, self.directorate_id, label="Despesa administrativa")
        c = self._validate_dadm_cost(payload)
        row.period = c["periodo"]; row.cost_center = c["centro_custo"]; row.modality = c["modalidade"]
        row.expense_amount = c["despesa"]; row.inserted_by = self.ctx.full_name
        self._audit("update", "dadm_administrative_cost", row.id, payload); self._commit(); return self._dadm_cost_to_dict(row)

    def delete_dadm_admin_cost(self, row_id: int):
        row = require_object_for_directorate(self.db, DadmAdministrativeCost, row_id, self.directorate_id, label="Despesa administrativa")
        self.db.delete(row); self._audit("delete", "dadm_administrative_cost", row_id); self._commit()

    def _validate_dadm_infrastructure(self, payload: dict[str, Any]) -> dict[str, Any]:
        self._ensure_dadm(); errors: dict[str, str] = {}
        period = str(payload.get("periodo") or "").strip()
        if not MONTH_RE.match(period): errors["periodo"] = "Selecione uma competência mensal válida (AAAA-MM)."
        expense = self._to_decimal(payload.get("despesa"), "despesa", errors)
        area = self._to_decimal(payload.get("area_m2"), "area_m2", errors, allow_zero=False)
        if errors: raise ValidationError("Revise os campos destacados.", errors)
        return {"periodo": period, "despesa": expense, "area_m2": area}

    def _dadm_infrastructure_to_dict(self, row: DadmInfrastructure) -> dict[str, Any]:
        expense = float(row.infrastructure_expense)
        area = float(row.area_in_use_m2)
        return {
            "id": row.id, "periodo": row.period, "despesa": expense, "area_m2": area,
            "custo_m2": round(expense / area, 2) if area else None, "validacao": "OK",
            "data_lancamento": str(row.inserted_at.date()) if row.inserted_at else "", "lancado_por": row.inserted_by or "",
        }

    def list_dadm_infrastructure(self) -> list[dict[str, Any]]:
        self._ensure_dadm()
        rows = self.db.scalars(
            select(DadmInfrastructure).where(DadmInfrastructure.directorate_id == self.directorate_id)
            .order_by(DadmInfrastructure.period.desc())
        ).all()
        return [self._dadm_infrastructure_to_dict(row) for row in rows]

    def create_dadm_infrastructure(self, payload: dict[str, Any]) -> dict[str, Any]:
        c = self._validate_dadm_infrastructure(payload)
        row = DadmInfrastructure(
            directorate_id=self.directorate_id, period=c["periodo"], infrastructure_expense=c["despesa"],
            area_in_use_m2=c["area_m2"], inserted_by=self.ctx.full_name,
        )
        self.db.add(row); self.db.flush(); self._audit("create", "dadm_infrastructure", row.id, payload); self._commit()
        return self._dadm_infrastructure_to_dict(row)

    def update_dadm_infrastructure(self, row_id: int, payload: dict[str, Any]) -> dict[str, Any]:
        row = require_object_for_directorate(self.db, DadmInfrastructure, row_id, self.directorate_id, label="Registro de infraestrutura")
        c = self._validate_dadm_infrastructure(payload)
        row.period = c["periodo"]; row.infrastructure_expense = c["despesa"]; row.area_in_use_m2 = c["area_m2"]
        row.inserted_by = self.ctx.full_name
        self._audit("update", "dadm_infrastructure", row.id, payload); self._commit(); return self._dadm_infrastructure_to_dict(row)

    def delete_dadm_infrastructure(self, row_id: int):
        row = require_object_for_directorate(self.db, DadmInfrastructure, row_id, self.directorate_id, label="Registro de infraestrutura")
        self.db.delete(row); self._audit("delete", "dadm_infrastructure", row_id); self._commit()

    def snapshot_dadm(self) -> dict[str, Any]:
        self._ensure_dadm()
        return {
            "evasao": self.list_dadm_attrition(),
            "alunos": self.list_dadm_active_students(),
            "custos": self.list_dadm_admin_costs(),
            "infraestrutura": self.list_dadm_infrastructure(),
            "cursos_academicos": self.list_academic_courses_for_dadm(),
            "metas": self.get_metas(),
            "actions": self.list_actions(),
        }

    def _find_existing_dadm(self, dataset: str, clean: dict[str, Any]):
        if dataset == "dadm-evasao":
            return self.db.scalar(select(DadmAttrition).where(
                DadmAttrition.directorate_id == self.directorate_id,
                DadmAttrition.period == clean["periodo"], DadmAttrition.course_id == clean["course_id"],
            ))
        if dataset == "dadm-alunos":
            return self.db.scalar(select(DadmActiveStudents).where(
                DadmActiveStudents.directorate_id == self.directorate_id,
                DadmActiveStudents.period == clean["periodo"], DadmActiveStudents.modality == clean["modalidade"],
            ))
        if dataset == "dadm-custos":
            return self.db.scalar(select(DadmAdministrativeCost).where(
                DadmAdministrativeCost.directorate_id == self.directorate_id,
                DadmAdministrativeCost.period == clean["periodo"],
                DadmAdministrativeCost.cost_center == clean["centro_custo"],
                DadmAdministrativeCost.modality == clean["modalidade"],
            ))
        if dataset == "dadm-infraestrutura":
            return self.db.scalar(select(DadmInfrastructure).where(
                DadmInfrastructure.directorate_id == self.directorate_id,
                DadmInfrastructure.period == clean["periodo"],
            ))
        raise RepositoryError("Base administrativa desconhecida.")

    def _validate_dadm_import_payload(self, dataset: str, payload: dict[str, Any]) -> dict[str, Any]:
        if dataset == "dadm-evasao": return self._validate_dadm_attrition(payload)
        if dataset == "dadm-alunos": return self._validate_dadm_students(payload)
        if dataset == "dadm-custos": return self._validate_dadm_cost(payload)
        if dataset == "dadm-infraestrutura": return self._validate_dadm_infrastructure(payload)
        raise RepositoryError("Base administrativa desconhecida.")

    def _create_dadm_import_payload(self, dataset: str, payload: dict[str, Any]):
        if dataset == "dadm-evasao": return self.create_dadm_attrition(payload)
        if dataset == "dadm-alunos": return self.create_dadm_active_students(payload)
        if dataset == "dadm-custos": return self.create_dadm_admin_cost(payload)
        if dataset == "dadm-infraestrutura": return self.create_dadm_infrastructure(payload)
        raise RepositoryError("Base administrativa desconhecida.")

    def _update_dadm_import_payload(self, dataset: str, row_id: int, payload: dict[str, Any]):
        if dataset == "dadm-evasao": return self.update_dadm_attrition(row_id, payload)
        if dataset == "dadm-alunos": return self.update_dadm_active_students(row_id, payload)
        if dataset == "dadm-custos": return self.update_dadm_admin_cost(row_id, payload)
        if dataset == "dadm-infraestrutura": return self.update_dadm_infrastructure(row_id, payload)
        raise RepositoryError("Base administrativa desconhecida.")

    def import_dadm_file(self, dataset: str, file_path: Path, mode: str = "add"):
        self._ensure_dadm()
        aliases = DADM_HEADER_ALIASES.get(dataset)
        if not aliases: raise RepositoryError("Base administrativa desconhecida.")
        try:
            wb = load_workbook(file_path, data_only=True, read_only=True)
        except Exception as exc:
            raise ValidationError("O arquivo enviado não é uma planilha Excel válida.") from exc
        try:
            ws = wb["IMPORTACAO"] if "IMPORTACAO" in wb.sheetnames else wb.active
            header_row = None; mapping = {}
            for rr in range(1, min(ws.max_row, 20) + 1):
                normalized = {norm_header(ws.cell(rr, c).value): c for c in range(1, ws.max_column + 1)}
                candidate = {}
                for field, accepted in aliases.items():
                    for header, col in normalized.items():
                        if header in accepted:
                            candidate[field] = col; break
                if len(candidate) == len(aliases):
                    header_row = rr; mapping = candidate; break
            if not header_row:
                raise ValidationError("Não encontrei os cabeçalhos esperados. Use o modelo DADM fornecido pelo sistema.")
            raw = []
            for rr in range(header_row + 1, ws.max_row + 1):
                payload = {field: ws.cell(rr, col).value for field, col in mapping.items()}
                if any(v not in (None, "") for v in payload.values()): raw.append((rr, payload))
        finally:
            wb.close()

        result = {"inseridos": 0, "atualizados": 0, "ignorados": 0, "erros": []}
        for source_row, payload in raw:
            try:
                clean = self._validate_dadm_import_payload(dataset, payload)
                existing = self._find_existing_dadm(dataset, clean)
                if existing:
                    if mode == "update":
                        self._update_dadm_import_payload(dataset, existing.id, payload); result["atualizados"] += 1
                    else:
                        result["ignorados"] += 1
                else:
                    self._create_dadm_import_payload(dataset, payload); result["inseridos"] += 1
            except Exception as exc:
                self.db.rollback()
                result["erros"].append({
                    "linha": source_row, "mensagem": str(exc), "campos": getattr(exc, "field_errors", {}),
                })
        self._audit("import", dataset, None, result); self._commit(); return result

    # ------------------------------------------------------------------
    # DPE: Resultado Operacional, Execução Orçamentária e Saldo de Caixa.
    # Unlike DADM-01/09/10, these flows do not depend on a student base.
    # Their dimensions are financial/managerial: recorte, diretoria/centro
    # de custo, conta e natureza do movimento.
    # ------------------------------------------------------------------
    def _ensure_dpe(self):
        if self.directorate_code != "DPE":
            raise RepositoryError("Este conjunto de dados pertence à DPE.")

    @staticmethod
    def _normalize_choice(value: Any, choices: tuple[str, ...]) -> str:
        text = str(value or "").strip()
        folded = norm_header(text)
        for item in choices:
            if folded == norm_header(item):
                return item
        return text

    def dpe_reference_options(self) -> dict[str, Any]:
        self._ensure_dpe()
        directorates = self.db.scalars(select(Directorate).order_by(Directorate.code)).all()
        result_rows = self.db.scalars(
            select(DpeOperatingResult).where(DpeOperatingResult.directorate_id == self.directorate_id)
        ).all()
        budget_rows = self.db.scalars(
            select(DpeBudgetExecution).where(DpeBudgetExecution.directorate_id == self.directorate_id)
        ).all()
        cash_rows = self.db.scalars(
            select(DpeCashMovement).where(DpeCashMovement.directorate_id == self.directorate_id)
        ).all()
        return {
            "tipos_recorte": list(DPE_RESULT_SCOPE_TYPES),
            "tipos_movimento": list(DPE_CASH_MOVEMENT_TYPES),
            "diretorias": [{"code": d.code, "name": d.name, "active": d.active} for d in directorates],
            "recortes": sorted({r.scope_label for r in result_rows if r.scope_label}),
            "centros_custo": sorted({r.cost_center for r in budget_rows if r.cost_center}),
            "contas": sorted({r.account for r in cash_rows if r.account}),
            "naturezas": sorted({r.nature for r in cash_rows if r.nature}),
        }

    def _validate_dpe_result(self, payload: dict[str, Any]) -> dict[str, Any]:
        self._ensure_dpe(); errors: dict[str, str] = {}
        period = str(payload.get("periodo") or "").strip()
        scope_type = self._normalize_choice(payload.get("tipo_recorte"), DPE_RESULT_SCOPE_TYPES)
        scope_label = str(payload.get("recorte") or "").strip()[:180]
        if not MONTH_RE.match(period): errors["periodo"] = "Selecione uma competência mensal válida (AAAA-MM)."
        if scope_type not in DPE_RESULT_SCOPE_TYPES: errors["tipo_recorte"] = "Selecione um tipo de recorte válido."
        if scope_type == "Institucional": scope_label = "TOTAL"
        elif len(scope_label) < 2: errors["recorte"] = "Informe o recorte correspondente."
        revenue = self._to_decimal(payload.get("receita_liquida"), "receita_liquida", errors, allow_zero=False)
        expense = self._to_decimal(payload.get("despesa_total"), "despesa_total", errors)
        if errors: raise ValidationError("Revise os campos destacados.", errors)
        return {"periodo": period, "tipo_recorte": scope_type, "recorte": scope_label, "receita_liquida": revenue, "despesa_total": expense}

    @staticmethod
    def _dpe_result_to_dict(row: DpeOperatingResult) -> dict[str, Any]:
        revenue = float(row.net_revenue); expense = float(row.total_expense)
        result = revenue - expense
        margin = result / revenue * 100 if revenue else None
        return {
            "id": row.id, "periodo": row.period, "tipo_recorte": row.scope_type, "recorte": row.scope_label,
            "receita_liquida": revenue, "despesa_total": expense, "resultado_operacional": round(result, 2),
            "margem_operacional": round(margin, 2) if margin is not None else None,
            "valor": round(margin, 2) if margin is not None else None, "validacao": "OK",
            "data_lancamento": str(row.inserted_at.date()) if row.inserted_at else "", "lancado_por": row.inserted_by or "",
        }

    def list_dpe_results(self) -> list[dict[str, Any]]:
        self._ensure_dpe()
        rows = self.db.scalars(select(DpeOperatingResult).where(DpeOperatingResult.directorate_id == self.directorate_id).order_by(DpeOperatingResult.period.desc(), DpeOperatingResult.scope_type, DpeOperatingResult.scope_label)).all()
        return [self._dpe_result_to_dict(r) for r in rows]

    def create_dpe_result(self, payload: dict[str, Any]) -> dict[str, Any]:
        c = self._validate_dpe_result(payload)
        row = DpeOperatingResult(directorate_id=self.directorate_id, period=c["periodo"], scope_type=c["tipo_recorte"], scope_label=c["recorte"], net_revenue=c["receita_liquida"], total_expense=c["despesa_total"], inserted_by=self.ctx.full_name)
        self.db.add(row); self.db.flush(); self._audit("create", "dpe_operating_result", row.id, payload); self._commit(); return self._dpe_result_to_dict(row)

    def update_dpe_result(self, row_id: int, payload: dict[str, Any]) -> dict[str, Any]:
        row = require_object_for_directorate(self.db, DpeOperatingResult, row_id, self.directorate_id, label="Resultado operacional")
        c = self._validate_dpe_result(payload)
        row.period=c["periodo"]; row.scope_type=c["tipo_recorte"]; row.scope_label=c["recorte"]; row.net_revenue=c["receita_liquida"]; row.total_expense=c["despesa_total"]; row.inserted_by=self.ctx.full_name
        self._audit("update", "dpe_operating_result", row.id, payload); self._commit(); return self._dpe_result_to_dict(row)

    def delete_dpe_result(self, row_id: int):
        row = require_object_for_directorate(self.db, DpeOperatingResult, row_id, self.directorate_id, label="Resultado operacional")
        self.db.delete(row); self._audit("delete", "dpe_operating_result", row_id); self._commit()

    def _validate_dpe_budget(self, payload: dict[str, Any]) -> dict[str, Any]:
        self._ensure_dpe(); errors: dict[str, str] = {}
        period = str(payload.get("periodo") or "").strip()
        unit = str(payload.get("unidade") or "").strip().upper()[:40]
        center = str(payload.get("centro_custo") or "").strip()[:180]
        if not MONTH_RE.match(period): errors["periodo"] = "Selecione uma competência mensal válida (AAAA-MM)."
        valid_units = set(self.db.scalars(select(Directorate.code)).all())
        if unit not in valid_units: errors["unidade"] = "Selecione uma diretoria cadastrada."
        if len(center) < 2: errors["centro_custo"] = "Informe o centro de custo."
        budgeted = self._to_decimal(payload.get("despesa_orcada"), "despesa_orcada", errors, allow_zero=False)
        actual = self._to_decimal(payload.get("despesa_realizada"), "despesa_realizada", errors)
        if errors: raise ValidationError("Revise os campos destacados.", errors)
        return {"periodo": period, "unidade": unit, "centro_custo": center, "despesa_orcada": budgeted, "despesa_realizada": actual}

    @staticmethod
    def _dpe_budget_to_dict(row: DpeBudgetExecution) -> dict[str, Any]:
        budgeted=float(row.budgeted_expense); actual=float(row.actual_expense); execution=actual/budgeted*100 if budgeted else None
        return {"id":row.id,"periodo":row.period,"unidade":row.budget_directorate,"centro_custo":row.cost_center,"despesa_orcada":budgeted,"despesa_realizada":actual,"execucao":round(execution,2) if execution is not None else None,"desvio":round(actual-budgeted,2),"valor":round(execution,2) if execution is not None else None,"validacao":"OK","data_lancamento":str(row.inserted_at.date()) if row.inserted_at else "","lancado_por":row.inserted_by or ""}

    def list_dpe_budget(self) -> list[dict[str, Any]]:
        self._ensure_dpe(); rows=self.db.scalars(select(DpeBudgetExecution).where(DpeBudgetExecution.directorate_id==self.directorate_id).order_by(DpeBudgetExecution.period.desc(),DpeBudgetExecution.budget_directorate,DpeBudgetExecution.cost_center)).all(); return [self._dpe_budget_to_dict(r) for r in rows]

    def create_dpe_budget(self, payload: dict[str, Any]) -> dict[str, Any]:
        c=self._validate_dpe_budget(payload); row=DpeBudgetExecution(directorate_id=self.directorate_id,period=c["periodo"],budget_directorate=c["unidade"],cost_center=c["centro_custo"],budgeted_expense=c["despesa_orcada"],actual_expense=c["despesa_realizada"],inserted_by=self.ctx.full_name); self.db.add(row);self.db.flush();self._audit("create","dpe_budget_execution",row.id,payload);self._commit();return self._dpe_budget_to_dict(row)

    def update_dpe_budget(self,row_id:int,payload:dict[str,Any])->dict[str,Any]:
        row=require_object_for_directorate(self.db, DpeBudgetExecution, row_id, self.directorate_id, label="Execução orçamentária")
        c=self._validate_dpe_budget(payload);row.period=c["periodo"];row.budget_directorate=c["unidade"];row.cost_center=c["centro_custo"];row.budgeted_expense=c["despesa_orcada"];row.actual_expense=c["despesa_realizada"];row.inserted_by=self.ctx.full_name;self._audit("update","dpe_budget_execution",row.id,payload);self._commit();return self._dpe_budget_to_dict(row)

    def delete_dpe_budget(self,row_id:int):
        row=require_object_for_directorate(self.db, DpeBudgetExecution, row_id, self.directorate_id, label="Execução orçamentária")
        self.db.delete(row);self._audit("delete","dpe_budget_execution",row_id);self._commit()

    def _validate_dpe_cash(self,payload:dict[str,Any])->dict[str,Any]:
        self._ensure_dpe();errors:dict[str,str]={};period=str(payload.get("periodo") or "").strip();account=str(payload.get("conta") or "").strip()[:180];movement=self._normalize_choice(payload.get("tipo_movimento"),DPE_CASH_MOVEMENT_TYPES);nature=str(payload.get("natureza") or "").strip()[:180]
        if not MONTH_RE.match(period):errors["periodo"]="Selecione uma competência mensal válida (AAAA-MM)."
        if len(account)<2:errors["conta"]="Informe a conta."
        if movement not in DPE_CASH_MOVEMENT_TYPES:errors["tipo_movimento"]="Selecione Entrada ou Saída."
        if len(nature)<2:errors["natureza"]="Informe a natureza do movimento."
        amount=self._to_decimal(payload.get("valor"),"valor",errors,allow_zero=False)
        if errors:raise ValidationError("Revise os campos destacados.",errors)
        return {"periodo":period,"conta":account,"tipo_movimento":movement,"natureza":nature,"valor":amount}

    @staticmethod
    def _dpe_cash_to_dict(row:DpeCashMovement)->dict[str,Any]:
        amount=float(row.amount);signed=amount if row.movement_type=="Entrada" else -amount
        return {"id":row.id,"periodo":row.period,"conta":row.account,"tipo_movimento":row.movement_type,"natureza":row.nature,"valor":amount,"valor_assinado":signed,"validacao":"OK","data_lancamento":str(row.inserted_at.date()) if row.inserted_at else "","lancado_por":row.inserted_by or ""}

    def list_dpe_cash(self)->list[dict[str,Any]]:
        self._ensure_dpe();rows=self.db.scalars(select(DpeCashMovement).where(DpeCashMovement.directorate_id==self.directorate_id).order_by(DpeCashMovement.period.desc(),DpeCashMovement.account,DpeCashMovement.movement_type,DpeCashMovement.nature)).all();return [self._dpe_cash_to_dict(r) for r in rows]

    def create_dpe_cash(self,payload:dict[str,Any])->dict[str,Any]:
        c=self._validate_dpe_cash(payload);row=DpeCashMovement(directorate_id=self.directorate_id,period=c["periodo"],account=c["conta"],movement_type=c["tipo_movimento"],nature=c["natureza"],amount=c["valor"],inserted_by=self.ctx.full_name);self.db.add(row);self.db.flush();self._audit("create","dpe_cash_movement",row.id,payload);self._commit();return self._dpe_cash_to_dict(row)

    def update_dpe_cash(self,row_id:int,payload:dict[str,Any])->dict[str,Any]:
        row=require_object_for_directorate(self.db, DpeCashMovement, row_id, self.directorate_id, label="Movimento de caixa")
        c=self._validate_dpe_cash(payload);row.period=c["periodo"];row.account=c["conta"];row.movement_type=c["tipo_movimento"];row.nature=c["natureza"];row.amount=c["valor"];row.inserted_by=self.ctx.full_name;self._audit("update","dpe_cash_movement",row.id,payload);self._commit();return self._dpe_cash_to_dict(row)

    def delete_dpe_cash(self,row_id:int):
        row=require_object_for_directorate(self.db, DpeCashMovement, row_id, self.directorate_id, label="Movimento de caixa")
        self.db.delete(row);self._audit("delete","dpe_cash_movement",row_id);self._commit()

    def snapshot_dpe(self)->dict[str,Any]:
        self._ensure_dpe();return {"resultado":self.list_dpe_results(),"orcamento":self.list_dpe_budget(),"caixa":self.list_dpe_cash(),"metas":self.get_metas(),"actions":self.list_actions()}

    def _find_existing_dpe(self,dataset:str,c:dict[str,Any]):
        if dataset=="dpe-resultado":return self.db.scalar(select(DpeOperatingResult).where(DpeOperatingResult.directorate_id==self.directorate_id,DpeOperatingResult.period==c["periodo"],DpeOperatingResult.scope_type==c["tipo_recorte"],DpeOperatingResult.scope_label==c["recorte"]))
        if dataset=="dpe-orcamento":return self.db.scalar(select(DpeBudgetExecution).where(DpeBudgetExecution.directorate_id==self.directorate_id,DpeBudgetExecution.period==c["periodo"],DpeBudgetExecution.budget_directorate==c["unidade"],DpeBudgetExecution.cost_center==c["centro_custo"]))
        if dataset=="dpe-caixa":return self.db.scalar(select(DpeCashMovement).where(DpeCashMovement.directorate_id==self.directorate_id,DpeCashMovement.period==c["periodo"],DpeCashMovement.account==c["conta"],DpeCashMovement.movement_type==c["tipo_movimento"],DpeCashMovement.nature==c["natureza"]))
        raise RepositoryError("Base DPE desconhecida.")

    def import_dpe_file(self,dataset:str,file_path:Path,mode:str="add"):
        self._ensure_dpe();aliases=DPE_HEADER_ALIASES.get(dataset)
        if not aliases:raise RepositoryError("Base DPE desconhecida.")
        try:wb=load_workbook(file_path,data_only=True,read_only=True)
        except Exception as exc:raise ValidationError("O arquivo enviado não é uma planilha Excel válida.") from exc
        try:
            ws=wb["MODELO"] if "MODELO" in wb.sheetnames else (wb["IMPORTACAO"] if "IMPORTACAO" in wb.sheetnames else wb.active);header_row=None;mapping={}
            for rr in range(1,min(ws.max_row,20)+1):
                normalized={norm_header(ws.cell(rr,c).value):c for c in range(1,ws.max_column+1)};candidate={}
                for field,accepted in aliases.items():
                    for header,col in normalized.items():
                        if header in accepted:candidate[field]=col;break
                if len(candidate)==len(aliases):header_row=rr;mapping=candidate;break
            if not header_row:raise ValidationError("Não encontrei os cabeçalhos esperados. Use o modelo DPE fornecido pelo sistema.")
            raw=[]
            for rr in range(header_row+1,ws.max_row+1):
                payload={field:ws.cell(rr,col).value for field,col in mapping.items()}
                if any(v not in (None,"") for v in payload.values()):raw.append((rr,payload))
        finally:wb.close()
        result={"inseridos":0,"atualizados":0,"ignorados":0,"erros":[]}
        validators={"dpe-resultado":self._validate_dpe_result,"dpe-orcamento":self._validate_dpe_budget,"dpe-caixa":self._validate_dpe_cash};creators={"dpe-resultado":self.create_dpe_result,"dpe-orcamento":self.create_dpe_budget,"dpe-caixa":self.create_dpe_cash};updaters={"dpe-resultado":self.update_dpe_result,"dpe-orcamento":self.update_dpe_budget,"dpe-caixa":self.update_dpe_cash}
        for source_row,payload in raw:
            try:
                clean=validators[dataset](payload);existing=self._find_existing_dpe(dataset,clean)
                if existing:
                    if mode=="update":updaters[dataset](existing.id,payload);result["atualizados"]+=1
                    else:result["ignorados"]+=1
                else:creators[dataset](payload);result["inseridos"]+=1
            except Exception as exc:
                self.db.rollback();result["erros"].append({"linha":source_row,"mensagem":str(exc),"campos":getattr(exc,"field_errors",{})})
        self._audit("import",dataset,None,result);self._commit();return result

    def snapshot(self):
        courses, disciplines = self.course_maps(include_inactive=False)
        return {
            "courses": courses,
            "disciplines": disciplines,
            "metas": self.get_metas(),
            "nps": self.list_records("nps"),
            "avaliacao_docente": self.list_records("avaliacao_docente"),
            "resultados": self.list_records("resultados"),
        }

    def _new_academic_model(self, dataset: str, clean: dict[str, Any]):
        if dataset == "nps":
            return NpsStudent(
                directorate_id=self.directorate_id, period=clean["periodo"], course_id=clean["course_id"],
                respondents=clean["respondentes"], promoters=clean["promotores"], neutrals=clean["neutros"],
                detractors=clean["detratores"], inserted_by=self.ctx.full_name,
            )
        if dataset == "avaliacao_docente":
            return TeacherEvaluation(
                directorate_id=self.directorate_id, period=clean["periodo"], course_id=clean["course_id"],
                discipline_id=clean["discipline_id"], teacher_name=clean["professor"],
                respondents=clean["respondentes"], average_score=clean["nota_media"], inserted_by=self.ctx.full_name,
            )
        if dataset == "resultados":
            return AcademicResult(
                directorate_id=self.directorate_id, period=clean["periodo"], course_id=clean["course_id"],
                discipline_id=clean["discipline_id"], student_id=clean["student_id"], class_group=clean["turma"],
                curriculum_period=clean["periodo_curricular"], final_average=clean["media"],
                official_status=clean["situacao"], approved=clean["aprovado"],
                failure_reason=clean["motivo_reprovacao"], source=clean["fonte"], inserted_by=self.ctx.full_name,
            )
        if dataset == "matriculas":
            return Enrollment(
                directorate_id=self.directorate_id, period=clean["periodo"], course_id=clean["course_id"],
                active_enrollments=clean["matriculas"], inserted_by=self.ctx.full_name,
            )
        return Attendance(
            directorate_id=self.directorate_id, period=clean["periodo"], course_id=clean["course_id"],
            discipline_id=clean["discipline_id"], expected_attendance=clean["presencas_previstas"],
            recorded_attendance=clean["presencas_registradas"], inserted_by=self.ctx.full_name,
        )

    def _apply_academic_update(self, dataset: str, row, clean: dict[str, Any]):
        row.period = clean["periodo"]
        row.course_id = clean["course_id"]
        if dataset == "nps":
            row.respondents = clean["respondentes"]
            row.promoters = clean["promotores"]
            row.neutrals = clean["neutros"]
            row.detractors = clean["detratores"]
        elif dataset == "avaliacao_docente":
            row.discipline_id = clean["discipline_id"]
            row.teacher_name = clean["professor"]
            row.respondents = clean["respondentes"]
            row.average_score = clean["nota_media"]
        elif dataset == "resultados":
            row.discipline_id = clean["discipline_id"]
            row.student_id = clean["student_id"]
            row.class_group = clean["turma"]
            row.curriculum_period = clean["periodo_curricular"]
            row.final_average = clean["media"]
            row.official_status = clean["situacao"]
            row.approved = clean["aprovado"]
            row.failure_reason = clean["motivo_reprovacao"]
            row.source = clean["fonte"]
        elif dataset == "matriculas":
            row.active_enrollments = clean["matriculas"]
        else:
            row.discipline_id = clean["discipline_id"]
            row.expected_attendance = clean["presencas_previstas"]
            row.recorded_attendance = clean["presencas_registradas"]
        row.inserted_by = self.ctx.full_name

    @staticmethod
    def _chunked(values: list[Any], size: int = 400):
        for offset in range(0, len(values), size):
            yield values[offset:offset + size]

    @staticmethod
    def _result_key(clean: dict[str, Any]) -> tuple[Any, ...]:
        return (
            clean["periodo"],
            clean["course_id"],
            clean["discipline_id"],
            clean["student_id"],
            clean["turma"],
        )

    def _prefetch_existing_results(self, prepared: list[tuple[Any, dict[str, Any]]]) -> dict[tuple[Any, ...], AcademicResult]:
        if not prepared:
            return {}
        grouped: dict[tuple[str, int], set[int]] = {}
        for _, clean in prepared:
            grouped.setdefault((clean["periodo"], clean["course_id"]), set()).add(clean["student_id"])
        result: dict[tuple[Any, ...], AcademicResult] = {}
        for (period, course_id), student_ids in grouped.items():
            ordered_ids = sorted(student_ids)
            for chunk in self._chunked(ordered_ids):
                rows = self.db.scalars(
                    select(AcademicResult).where(
                        AcademicResult.directorate_id == self.directorate_id,
                        AcademicResult.period == period,
                        AcademicResult.course_id == course_id,
                        AcademicResult.student_id.in_(chunk),
                    )
                ).all()
                for row in rows:
                    result[(row.period, row.course_id, row.discipline_id, row.student_id, row.class_group)] = row
        return result

    def _import_academic_results_batch(
        self,
        raw: list[tuple[Any, dict[str, Any]]],
        mode: str = "add",
        source: str | None = None,
    ) -> dict[str, Any]:
        result = {"inseridos": 0, "atualizados": 0, "ignorados": 0, "erros": []}
        normalized: list[tuple[Any, dict[str, Any]]] = []

        # Phase 1: pure validation. Invalid rows never create catalog/student side effects.
        for source_row, original_payload in raw:
            payload = dict(original_payload)
            if source:
                payload["fonte"] = source
            try:
                normalized.append((source_row, self._normalize_result_payload(payload)))
            except Exception as exc:
                result["erros"].append({
                    "linha": source_row,
                    "mensagem": str(exc),
                    "campos": getattr(exc, "field_errors", {}),
                })

        if not normalized:
            self._audit("import", "resultados", None, {**result, "fonte": source or "excel", "engine": "batch"})
            self._commit()
            return result

        # Phase 2: resolve the small catalog dimensions once per unique key.
        course_by_name: dict[str, Course] = {}
        discipline_by_key: dict[tuple[int, str], Discipline] = {}
        for _, clean in normalized:
            course = course_by_name.get(clean["curso"])
            if course is None:
                course = self._get_or_create_course(clean["curso"], clean["periodo"])
                course_by_name[clean["curso"]] = course
            clean["course_id"] = course.id
            discipline_key = (course.id, self._text_key(clean["disciplina"]))
            discipline = discipline_by_key.get(discipline_key)
            if discipline is None:
                discipline = self._get_or_create_discipline(course, clean["disciplina"], clean["periodo"])
                discipline_by_key[discipline_key] = discipline
            clean["discipline_id"] = discipline.id

        # Phase 3: prefetch students, update known ones in memory, and create missing
        # registrations with one executemany instead of SELECT+INSERT per line.
        registrations = [clean["matricula"] for _, clean in normalized]
        self._prefetch_students(registrations)
        last_student_state: dict[str, dict[str, Any]] = {}
        for _, clean in normalized:
            last_student_state[clean["matricula"]] = clean
            student = self._student_registration_cache.get(clean["matricula"])
            if student is not None:
                if student.name != clean["aluno"]:
                    student.name = clean["aluno"]
                if student.course_id != clean["course_id"]:
                    student.course_id = clean["course_id"]
                if not student.active:
                    student.active = True

        missing = [registration for registration in dict.fromkeys(registrations) if self._student_registration_cache.get(registration) is None]
        if missing:
            self.db.execute(
                insert(AcademicStudent),
                [
                    {
                        "directorate_id": self.directorate_id,
                        "registration": registration,
                        "name": last_student_state[registration]["aluno"],
                        "course_id": last_student_state[registration]["course_id"],
                        "active": True,
                        "inserted_by": self.ctx.full_name,
                    }
                    for registration in missing
                ],
            )
            # Refresh only the newly-created identities so result foreign keys are known.
            for chunk in self._chunked(missing):
                students = self.db.scalars(
                    select(AcademicStudent).where(
                        AcademicStudent.directorate_id == self.directorate_id,
                        AcademicStudent.registration.in_(chunk),
                    )
                ).all()
                for student in students:
                    self._student_registration_cache[student.registration] = student

        prepared: list[tuple[Any, dict[str, Any]]] = []
        for source_row, clean in normalized:
            student = self._student_registration_cache.get(clean["matricula"])
            if student is None:
                result["erros"].append({
                    "linha": source_row,
                    "mensagem": "Não foi possível resolver a matrícula do aluno após a preparação em lote.",
                    "campos": {"matricula": "Aluno não localizado."},
                })
                continue
            clean["student_id"] = student.id
            prepared.append((source_row, clean))

        # Phase 4: load existing facts in bounded chunks and reproduce sequential
        # duplicate semantics inside the batch.
        existing_by_key = self._prefetch_existing_results(prepared)
        new_by_key: dict[tuple[Any, ...], dict[str, Any]] = {}
        for _source_row, clean in prepared:
            key = self._result_key(clean)
            existing = existing_by_key.get(key)
            pending = new_by_key.get(key)
            if existing is not None:
                if mode == "update":
                    self._apply_academic_update("resultados", existing, clean)
                    result["atualizados"] += 1
                else:
                    result["ignorados"] += 1
                continue
            if pending is not None:
                if mode == "update":
                    pending.update({
                        "curriculum_period": clean["periodo_curricular"],
                        "final_average": clean["media"],
                        "official_status": clean["situacao"],
                        "approved": clean["aprovado"],
                        "failure_reason": clean["motivo_reprovacao"],
                        "source": clean["fonte"],
                        "inserted_by": self.ctx.full_name,
                    })
                    result["atualizados"] += 1
                else:
                    result["ignorados"] += 1
                continue

            new_by_key[key] = {
                "directorate_id": self.directorate_id,
                "period": clean["periodo"],
                "course_id": clean["course_id"],
                "discipline_id": clean["discipline_id"],
                "student_id": clean["student_id"],
                "class_group": clean["turma"],
                "curriculum_period": clean["periodo_curricular"],
                "final_average": clean["media"],
                "official_status": clean["situacao"],
                "approved": clean["aprovado"],
                "failure_reason": clean["motivo_reprovacao"],
                "source": clean["fonte"],
                "inserted_by": self.ctx.full_name,
            }
            result["inseridos"] += 1

        if new_by_key:
            self.db.execute(insert(AcademicResult), list(new_by_key.values()))

        self._audit("import", "resultados", None, {**result, "fonte": source or "excel", "engine": "batch"})
        self._commit()
        return result

    def _import_academic_payloads(
        self,
        dataset: str,
        raw: list[tuple[Any, dict[str, Any]]],
        mode: str = "add",
        source: str | None = None,
    ):
        if dataset != "resultados":
            return self._import_academic_payloads_legacy(dataset, raw, mode=mode, source=source)
        try:
            return self._import_academic_results_batch(raw, mode=mode, source=source)
        except IntegrityError:
            # Concurrent catalog/student/result inserts can still race despite the
            # prefetch. Roll back the batch and fall back to the row-isolated path,
            # preserving the historical error semantics rather than failing the file.
            self.db.rollback()
            self._reset_import_caches()
            return self._import_academic_payloads_legacy(dataset, raw, mode=mode, source=source)

    def _import_academic_payloads_legacy(self, dataset: str, raw: list[tuple[Any, dict[str, Any]]], mode: str = "add", source: str | None = None):
        if dataset not in DATASETS:
            raise RepositoryError("Base desconhecida.")
        result = {"inseridos": 0, "atualizados": 0, "ignorados": 0, "erros": []}
        for source_row, original_payload in raw:
            payload = dict(original_payload)
            if dataset == "resultados" and source:
                payload["fonte"] = source
            try:
                with self.db.begin_nested():
                    clean = self.validate_payload(dataset, payload)
                    existing = self._find_existing_academic(dataset, clean)
                    if existing:
                        if mode == "update":
                            self._apply_academic_update(dataset, existing, clean)
                            self.db.flush()
                            result["atualizados"] += 1
                        else:
                            result["ignorados"] += 1
                    else:
                        row = self._new_academic_model(dataset, clean)
                        self.db.add(row)
                        self.db.flush()
                        result["inseridos"] += 1
            except Exception as exc:
                result["erros"].append({
                    "linha": source_row,
                    "mensagem": str(exc),
                    "campos": getattr(exc, "field_errors", {}),
                })
        self._audit("import", dataset, None, {**result, "fonte": source or "excel"})
        self._commit()
        return result

    def import_sei_report_file(
        self,
        file_path: Path,
        mode: str = "add",
        expected_course: str | None = None,
    ) -> dict[str, Any]:
        """Importa diretamente o XLSX bruto 'Mapa de Nota do Aluno por Turma' do SEI.

        Quando ``expected_course`` é informado pelo fluxo automático, o rótulo
        interno do XLSX precisa corresponder ao curso solicitado antes de qualquer
        gravação. Isso impede cruzar habilitações ou cursos parecidos.
        """
        try:
            metadata, registros, warnings = ler_registros(file_path)
        except Exception as exc:
            raise ValidationError(
                "A planilha não corresponde ao relatório bruto de notas do SEI/UNIVC.",
                {"arquivo": str(exc)},
            ) from exc

        if not metadata.get("ano") or not metadata.get("semestre"):
            raise ValidationError("Não foi possível identificar Ano/Semestre no relatório do SEI.")

        raw_report_course = str(metadata.get("curso") or "").strip()
        canonical_report_course = canonical_course_name(raw_report_course, self.directorate_code)
        if is_ambiguous_course_name(raw_report_course, self.directorate_code):
            raise ValidationError(
                "O relatório usa o nome ambíguo 'Educação Física'. Gere o relatório selecionando Bacharelado ou Licenciatura no SEI.",
                {"curso": raw_report_course},
            )
        if expected_course and not course_name_matches(raw_report_course, expected_course, self.directorate_code):
            raise ValidationError(
                "O XLSX retornado pelo SEI pertence a um curso diferente do solicitado.",
                {"curso": f"Esperado: {expected_course}. Relatório: {raw_report_course or 'não identificado'}."},
            )

        # A identidade do curso vem do campo ``Curso:`` de cada bloco do XLSX,
        # nunca do código/nome da turma. O SEI pode reutilizar convenções históricas
        # de turma (por exemplo, EFB) mesmo em relatórios cuja habilitação selecionada
        # é Licenciatura. Usar EFB/EFL como prova de identidade gerava falso conflito.
        #
        # Para manter a proteção contra mistura de habilitações, validamos TODOS os
        # rótulos de curso encontrados nos registros contra ``expected_course``.
        if expected_course:
            report_course_labels = sorted({
                str(record.get("curso") or "").strip()
                for record in registros
                if str(record.get("curso") or "").strip()
            })
            ambiguous_labels = [
                label
                for label in report_course_labels
                if is_ambiguous_course_name(label, self.directorate_code)
            ]
            if ambiguous_labels:
                raise ValidationError(
                    "O relatório contém bloco(s) com o nome ambíguo 'Educação Física'. Gere o relatório selecionando Bacharelado ou Licenciatura no SEI.",
                    {"curso": "; ".join(ambiguous_labels)},
                )

            mismatched_labels = [
                label
                for label in report_course_labels
                if not course_name_matches(label, expected_course, self.directorate_code)
            ]
            if mismatched_labels:
                raise ValidationError(
                    "O XLSX retornado pelo SEI contém bloco(s) de curso diferente do solicitado.",
                    {
                        "curso": (
                            f"Esperado: {expected_course}. "
                            f"Encontrei: {', '.join(mismatched_labels)}."
                        )
                    },
                )

        period = f"{metadata['ano']}-SEM{metadata['semestre']}"
        metadata["curso_origem_sei"] = raw_report_course
        metadata["curso"] = canonical_report_course
        raw: list[tuple[Any, dict[str, Any]]] = []
        for index, record in enumerate(registros, start=1):
            raw.append((index, {
                "periodo": period,
                "curso": canonical_course_name(
                    record.get("curso") or raw_report_course or metadata.get("curso"),
                    self.directorate_code,
                ),
                "disciplina": record.get("disciplina"),
                "turma": record.get("turma") or "SEM_TURMA",
                "periodo_curricular": record.get("periodo"),
                "matricula": record.get("matricula"),
                "aluno": record.get("nome"),
                "media": record.get("media"),
                "situacao": record.get("situacao"),
                "aprovado": record.get("aprovado"),
                "motivo_reprovacao": record.get("motivo_reprovacao"),
                "fonte": "SEI",
            }))
        result = self._import_academic_payloads("resultados", raw, mode=mode, source="SEI")

        # Diagnóstico explícito da identidade do relatório. A habilitação já foi
        # validada exclusivamente pelo campo Curso: do XLSX. O nome/código da turma
        # não é consultado nem mesmo como heurística auxiliar.
        metadata["curso_solicitado"] = expected_course
        metadata["identidade_curso_validada_por"] = "campo Curso: do XLSX"
        metadata["identidade_curso_validada"] = bool(
            not expected_course
            or course_name_matches(raw_report_course, expected_course, self.directorate_code)
        )

        result["metadados"] = metadata
        result["avisos_parser"] = warnings
        result["registros_lidos"] = len(registros)
        return result

    def import_file(self, dataset: str, file_path: Path, mode: str = "add"):
        if dataset not in DATASETS:
            raise RepositoryError("Base desconhecida.")

        # Para Resultados Acadêmicos aceitamos dois formatos:
        # 1) o XLSX bruto gerado pelo próprio SEI; 2) o modelo padronizado do Data UNIVC.
        if dataset == "resultados":
            try:
                return self.import_sei_report_file(file_path, mode)
            except ValidationError:
                # Se não for o relatório bruto do SEI, tenta o modelo padronizado abaixo.
                pass

        try:
            wb = load_workbook(file_path, data_only=True, read_only=True)
        except Exception as exc:
            raise ValidationError("O arquivo enviado não é uma planilha Excel válida.") from exc
        try:
            ws = wb["IMPORTACAO"] if "IMPORTACAO" in wb.sheetnames else wb.active
            aliases = HEADER_ALIASES[dataset]
            header_row = None
            mapping: dict[str, int] = {}
            for rr in range(1, min(ws.max_row, 30) + 1):
                normalized = {norm_header(ws.cell(rr, c).value): c for c in range(1, ws.max_column + 1)}
                candidate: dict[str, int] = {}
                for field, accepted in aliases.items():
                    for header, col in normalized.items():
                        if header in accepted:
                            candidate[field] = col
                            break
                # 'aprovado' e 'motivo_reprovacao' são opcionais no modelo de resultados:
                # podem ser inferidos da Situação oficial.
                required_fields = set(aliases)
                if dataset == "resultados":
                    required_fields -= {"aprovado", "motivo_reprovacao"}
                if required_fields.issubset(candidate):
                    header_row = rr
                    mapping = candidate
                    break
            if not header_row:
                raise ValidationError("Não encontrei os cabeçalhos esperados. Use o modelo fornecido pelo Data UNIVC ou o relatório bruto do SEI.")
            raw = []
            for rr in range(header_row + 1, ws.max_row + 1):
                payload = {field: ws.cell(rr, col).value for field, col in mapping.items()}
                if any(value not in (None, "") for value in payload.values()):
                    raw.append((rr, payload))
        finally:
            wb.close()
        return self._import_academic_payloads(dataset, raw, mode=mode, source="Excel")

    def import_disciplines_file(self, file_path: Path, mode: str = "add"):
        """Importa o cadastro de disciplinas a partir do modelo institucional.

        Chave de identificação: Curso + Disciplina. O campo "Recorte para METAS"
        é determinístico e não é persistido: o sistema sempre o deriva como
        "Curso » Disciplina" para evitar divergências entre cadastro e metas.
        """
        try:
            wb = load_workbook(file_path, data_only=True, read_only=True)
        except Exception as exc:
            raise ValidationError("O arquivo enviado não é uma planilha Excel válida.") from exc
        try:
            ws = wb["IMPORTACAO"] if "IMPORTACAO" in wb.sheetnames else wb.active
            aliases = DISCIPLINE_HEADER_ALIASES
            header_row = None
            mapping = {}
            for rr in range(1, min(ws.max_row, 20) + 1):
                normalized = {norm_header(ws.cell(rr, c).value): c for c in range(1, ws.max_column + 1)}
                candidate = {}
                for field, accepted in aliases.items():
                    for header, col in normalized.items():
                        if header in accepted:
                            candidate[field] = col
                            break
                if len(candidate) == len(aliases):
                    header_row = rr
                    mapping = candidate
                    break
            if not header_row:
                raise ValidationError(
                    "Não encontrei os cabeçalhos esperados para disciplinas. Use o modelo fornecido."
                )
            raw = []
            for rr in range(header_row + 1, ws.max_row + 1):
                payload = {field: ws.cell(rr, col).value for field, col in mapping.items()}
                if any(v not in (None, "") for v in payload.values()):
                    raw.append((rr, payload))
        finally:
            wb.close()

        def parse_active(value):
            text = str(value or "").strip().lower()
            text = text.translate(str.maketrans("ãáàâéêíóôõúüç", "aaaaeeiooouuc"))
            if text in {"s", "sim", "true", "1", "ativo", "ativa"}:
                return True
            if text in {"n", "nao", "false", "0", "inativo", "inativa"}:
                return False
            return None

        result = {"inseridos": 0, "atualizados": 0, "ignorados": 0, "erros": []}
        for source_row, payload in raw:
            try:
                errors = {}
                course_name = str(payload.get("curso") or "").strip()
                discipline_name = str(payload.get("disciplina") or "").strip()
                valid_from = str(payload.get("vigencia_inicio") or "").strip()
                valid_to = str(payload.get("vigencia_fim") or "").strip()
                supplied_scope = str(payload.get("recorte_metas") or "").strip()
                active = parse_active(payload.get("ativo"))

                course = self._course_by_name(course_name) if course_name else None
                if not course:
                    errors["curso"] = "Curso não cadastrado nesta diretoria. Cadastre o curso antes de importar as disciplinas."
                elif active is True and not course.active:
                    errors["curso"] = "O curso está inativo; não é possível importar uma disciplina ativa para ele."
                if len(discipline_name) < 2:
                    errors["disciplina"] = "Informe o nome da disciplina."
                if active is None:
                    errors["ativo"] = "Use S ou N."
                if not MONTH_RE.match(valid_from):
                    errors["vigencia_inicio"] = "Use o formato AAAA-MM."
                if valid_to and not MONTH_RE.match(valid_to):
                    errors["vigencia_fim"] = "Use o formato AAAA-MM ou deixe em branco."
                if valid_to and MONTH_RE.match(valid_from) and MONTH_RE.match(valid_to) and valid_to < valid_from:
                    errors["vigencia_fim"] = "A vigência fim não pode ser anterior à vigência início."
                if active is False and not valid_to:
                    errors["vigencia_fim"] = "Informe a vigência fim para disciplinas inativas."
                if active is True and valid_to:
                    errors["vigencia_fim"] = "Para disciplina ativa, deixe a vigência fim em branco."

                canonical_course = course.name if course else course_name
                expected_scope = f"{canonical_course} » {discipline_name}"
                if supplied_scope and supplied_scope != expected_scope:
                    errors["recorte_metas"] = f"Use '{expected_scope}' ou deixe o campo em branco; o sistema gera esse recorte automaticamente."

                if errors:
                    raise ValidationError("Revise os campos destacados.", errors)

                existing = self._discipline_by_name(course.id, discipline_name)
                if existing:
                    if mode == "update":
                        existing.name = discipline_name
                        existing.active = bool(active)
                        existing.valid_from = valid_from
                        existing.valid_to = valid_to or None
                        self._audit("update", "discipline", existing.id, {
                            "curso": canonical_course,
                            "disciplina": discipline_name,
                            "ativo": bool(active),
                            "vigencia_inicio": valid_from,
                            "vigencia_fim": valid_to or "",
                            "recorte_metas": expected_scope,
                            "origem": "importacao_excel",
                        })
                        self._commit()
                        result["atualizados"] += 1
                    else:
                        result["ignorados"] += 1
                else:
                    row = Discipline(
                        course_id=course.id,
                        name=discipline_name,
                        active=bool(active),
                        valid_from=valid_from,
                        valid_to=valid_to or None,
                    )
                    self.db.add(row)
                    self.db.flush()
                    self._audit("create", "discipline", row.id, {
                        "curso": canonical_course,
                        "disciplina": discipline_name,
                        "ativo": bool(active),
                        "vigencia_inicio": valid_from,
                        "vigencia_fim": valid_to or "",
                        "recorte_metas": expected_scope,
                        "origem": "importacao_excel",
                    })
                    self._commit()
                    result["inseridos"] += 1
            except Exception as exc:
                self.db.rollback()
                result["erros"].append({
                    "linha": source_row,
                    "mensagem": str(exc),
                    "campos": getattr(exc, "field_errors", {}),
                })

        self._audit("import", "disciplinas", None, result)
        self._commit()
        return result

    def recent_audit(self,limit=50):
        rows=self.db.scalars(select(AuditLog).where(AuditLog.directorate_id==self.directorate_id).order_by(AuditLog.created_at.desc()).limit(limit)).all()
        return [{"id":r.id,"acao":r.action,"entidade":r.entity,"usuario":r.user_email or "","data":r.created_at.isoformat() if r.created_at else "","detalhes":r.details or ""} for r in rows]
