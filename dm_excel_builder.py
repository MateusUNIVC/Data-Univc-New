from __future__ import annotations

from datetime import date
from io import BytesIO
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break

from dm_analytics import build_dm_dashboard
from dm_catalog import COHORT_STATUSES, DIPLOMA_STATUSES, DM_AREAS, STUDENT_STATUSES
from excel_errors import ExcelExportLimitError

GREEN_DARK = "045233"
GREEN = "0B7A54"
GREEN_LIGHT = "E2F0D9"
TEAL_LIGHT = "DDEFE9"
PALE_GREEN = "F3F8F5"
PALE_TEAL = "E8F3F0"
PALE_BLUE = "EAF2FA"
PALE_GOLD = "FFF4D6"
GRAY = "F5F6F5"
GRAY_TEXT = "5A5A5A"
BORDER = "D9E1DD"
WHITE = "FFFFFF"
BLACK = "000000"
INPUT_BLUE = "0000FF"
LINKED_GREEN = "008000"
PURPLE = "7030A0"
CHART_GREEN = "0B7A54"
CHART_TEAL = "398C78"
CHART_GOLD = "D9B44A"
CHART_RED = "C74B50"
CHART_BLUE = "4679A6"
CHART_META = "7F8C8D"
EXCEL_MAX_ROW = 1_048_576

INTEGER_FMT = '#,##0;[Red](#,##0);-'
DECIMAL_FMT = '0.0;[Red](0.0);-'
PERCENT_FMT = '0.0%;[Red](0.0%);-'
DATE_FMT = 'dd/mm/yyyy'


def _to_date(value: Any) -> date | None:
    if isinstance(value, date):
        return value
    if not value:
        return None
    try:
        return date.fromisoformat(str(value)[:10])
    except ValueError:
        return None


def _formula_status_color(status: str) -> str:
    if status == "Dentro da meta":
        return PALE_GREEN
    if status == "Atenção":
        return PALE_GOLD
    if status == "Fora da meta":
        return "FDE7E7"
    return GRAY


class DMWorkbookBuilder:
    def __init__(
        self,
        cohorts: Iterable[dict[str, Any]],
        students: Iterable[dict[str, Any]],
        *,
        area_code: str | None = None,
        cohort_id: int | None = None,
        as_of: date | None = None,
    ) -> None:
        self.cohorts = [dict(row) for row in cohorts]
        self.students = [dict(row) for row in students]
        self.as_of = as_of or date.today()
        self.dashboard = build_dm_dashboard(
            self.cohorts,
            self.students,
            area_code=area_code,
            cohort_id=cohort_id,
            as_of=self.as_of,
        )
        if len(self.students) + 20 >= EXCEL_MAX_ROW:
            raise ExcelExportLimitError(
                "A base de alunos excede o limite físico de linhas do Excel. A exportação foi cancelada sem truncamento."
            )
        if len(self.cohorts) + 20 >= EXCEL_MAX_ROW:
            raise ExcelExportLimitError(
                "A base de turmas excede o limite físico de linhas do Excel. A exportação foi cancelada sem truncamento."
            )
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.ws: dict[str, Any] = {}
        self.title_font = Font(name="Arial", size=15, bold=True, color=GREEN_DARK)
        self.section_font = Font(name="Arial", size=11, bold=True, color=WHITE)
        self.header_font = Font(name="Arial", size=9, bold=True, color=WHITE)
        self.body_font = Font(name="Arial", size=9, color=BLACK)
        self.static_font = Font(name="Arial", size=9, color=GRAY_TEXT)
        self.linked_font = Font(name="Arial", size=9, color=LINKED_GREEN)
        self.input_font = Font(name="Arial", size=9, color=INPUT_BLUE)
        self.border_box = Border(
            left=Side(style="thin", color=BORDER),
            right=Side(style="thin", color=BORDER),
            top=Side(style="thin", color=BORDER),
            bottom=Side(style="thin", color=BORDER),
        )
        self.calc_ranges: dict[str, tuple[int, int]] = {}

    def _new_sheet(self, title: str):
        ws = self.wb.create_sheet(title)
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 85
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_margins.left = 0.3
        ws.page_margins.right = 0.3
        ws.page_margins.top = 0.45
        ws.page_margins.bottom = 0.45
        self.ws[title] = ws
        return ws

    def _title(self, ws, text: str, subtitle: str | None = None, *, end_col: int = 12):
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
        ws.cell(1, 1).value = text
        ws.cell(1, 1).font = self.title_font
        ws.cell(1, 1).alignment = Alignment(vertical="center")
        ws.cell(1, 1).border = Border(bottom=Side(style="medium", color=GREEN))
        ws.row_dimensions[1].height = 28
        if subtitle:
            ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_col)
            ws.cell(2, 1).value = subtitle
            ws.cell(2, 1).font = self.static_font
            ws.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[2].height = 34

    def _section(self, ws, row: int, text: str, *, end_col: int = 12):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=end_col)
        cell = ws.cell(row, 1)
        cell.value = text
        cell.fill = PatternFill("solid", fgColor=GREEN_DARK)
        cell.font = self.section_font
        cell.alignment = Alignment(vertical="center")
        ws.row_dimensions[row].height = 24

    def _headers(self, ws, row: int, labels: list[str], *, start_col: int = 1):
        for offset, label in enumerate(labels):
            cell = ws.cell(row, start_col + offset)
            cell.value = label
            cell.fill = PatternFill("solid", fgColor=GREEN)
            cell.font = self.header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = self.border_box
        ws.row_dimensions[row].height = 30

    def _paint(self, ws, r1: int, r2: int, c1: int, c2: int, *, fill: str):
        for row in ws.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2):
            for cell in row:
                cell.fill = PatternFill("solid", fgColor=fill)
                cell.border = self.border_box

    def _kpi_card(
        self,
        ws,
        *,
        start_col: int,
        end_col: int,
        label_row: int,
        value_row: int,
        label: str,
        value: Any,
        number_format: str = "General",
        fill: str = PALE_GREEN,
        font_size: int = 15,
    ):
        self._paint(ws, label_row, value_row + 1, start_col, end_col, fill=fill)
        ws.merge_cells(start_row=label_row, start_column=start_col, end_row=label_row, end_column=end_col)
        ws.merge_cells(start_row=value_row, start_column=start_col, end_row=value_row + 1, end_column=end_col)
        ws.cell(label_row, start_col).value = label
        ws.cell(label_row, start_col).font = Font(name="Arial", size=8, bold=True, color=GRAY_TEXT)
        ws.cell(label_row, start_col).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.cell(value_row, start_col).value = value
        ws.cell(value_row, start_col).font = Font(name="Arial", size=font_size, bold=True, color=GREEN_DARK)
        ws.cell(value_row, start_col).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(value_row, start_col).number_format = number_format
        ws.row_dimensions[label_row].height = 20
        ws.row_dimensions[value_row].height = 24
        ws.row_dimensions[value_row + 1].height = 16

    def _chart_base(self, chart, *, title: str, x_title: str, y_title: str, width: float, height: float):
        chart.title = title
        chart.x_axis.title = x_title
        chart.y_axis.title = y_title
        chart.width = width
        chart.height = height
        chart.legend.position = "b"
        chart.style = 13
        chart.display_blanks = "gap"
        chart.plotVisOnly = False
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showVal = True
        chart.dataLabels.showLegendKey = False
        chart.dataLabels.showCatName = False
        chart.dataLabels.showSerName = False
        chart.dataLabels.showPercent = False

    def _style_series(self, chart, colors: list[str], *, line: bool = False):
        for series, color in zip(chart.series, colors):
            if line:
                series.graphicalProperties.line.solidFill = color
                series.graphicalProperties.line.width = 22000
                series.marker.symbol = "circle"
                series.marker.size = 6
                series.marker.graphicalProperties.solidFill = color
                series.marker.graphicalProperties.line.solidFill = color
            else:
                series.graphicalProperties.solidFill = color
                series.graphicalProperties.line.solidFill = color

    def build(self) -> Workbook:
        for name in (
            "LEIA-ME", "PARAMETROS", "PAINEL", "DM-01 EVOLUCAO", "DM-02 DEFESAS",
            "TURMAS", "ALUNOS", "QUALIDADE E GOVERNANÇA", "MATRIZ", "LISTAS DE APOIO", "CALC",
        ):
            self._new_sheet(name)
        self._build_readme()
        self._build_parameters()
        self._build_turmas()
        self._build_alunos()
        self._build_lists()
        self._build_calc()
        self._build_panel()
        self._build_dm01()
        self._build_dm02()
        self._build_matrix()
        self._build_quality()
        self._finish()
        return self.wb

    def _build_readme(self):
        ws = self.ws["LEIA-ME"]
        self._title(
            ws,
            "DATA UNIVC · PAINEL DM POR TURMAS",
            "Diretoria de Mestrado · a unidade de análise é a turma/coorte, separada entre as duas áreas do programa.",
            end_col=10,
        )
        self._section(ws, 4, "MODELO DE DADOS", end_col=10)
        items = [
            ("TURMAS", "Cada turma possui área, número, data de abertura, vagas autorizadas e status."),
            ("ALUNOS", "Cada aluno é vinculado a uma turma e possui ingresso, defesa, titulação e acompanhamento do diploma digital."),
            ("DM-01 EVOLUCAO", "Compara ingressos, ativos, titulados, desligados, ocupação e evasão entre turmas da mesma área."),
            ("DM-02 DEFESAS", "Compara tempo até a defesa, titulação em até 24 meses e riscos de prazo entre turmas."),
            ("PAINEL", "Apresenta as duas áreas separadamente. Uma turma só aparece quando tiver sido formalmente aberta."),
        ]
        self._headers(ws, 5, ["Aba", "Finalidade"])
        for row, (name, purpose) in enumerate(items, 6):
            ws.cell(row, 1).value = name
            ws.cell(row, 2).value = purpose
            ws.cell(row, 1).font = Font(name="Arial", size=9, bold=True, color=GREEN_DARK)
            ws.cell(row, 2).font = self.body_font
            ws.cell(row, 2).alignment = Alignment(wrap_text=True)
            for c in (1, 2):
                ws.cell(row, c).border = Border(bottom=Side(style="hair", color=BORDER))
            ws.row_dimensions[row].height = 30
        ws.column_dimensions["A"].width = 25
        ws.column_dimensions["B"].width = 95
        self._section(ws, 13, "REGRA INSTITUCIONAL", end_col=10)
        ws.merge_cells("A14:J18")
        ws["A14"] = (
            "O DM não é acompanhado por semestre ou ano como unidade principal. O eixo é a turma. "
            "Ciência, Tecnologia e Educação e Saúde e Desigualdade Social são áreas independentes: "
            "a existência de uma Turma 21 em uma área não cria automaticamente uma Turma 21 na outra. "
            "A data de abertura da turma é metadado opcional. O tempo até a defesa usa exclusivamente a data individual de ingresso do aluno."
        )
        ws["A14"].font = self.body_font
        ws["A14"].alignment = Alignment(wrap_text=True, vertical="top")
        self._paint(ws, 14, 18, 1, 10, fill=PALE_BLUE)

    def _build_parameters(self):
        ws = self.ws["PARAMETROS"]
        self._title(ws, "PARÂMETROS · DM", "Recorte das páginas executivas", end_col=8)
        self._section(ws, 4, "FILTROS PRINCIPAIS", end_col=8)
        latest = None
        for area in self.dashboard["areas"]:
            candidate = area.get("latest_cohort")
            if candidate:
                candidate_key = (str(candidate.get("opening_date") or ""), int(candidate.get("cohort_number") or 0))
                latest_key = (str(latest.get("opening_date") or ""), int(latest.get("cohort_number") or 0)) if latest else None
                if latest_key is None or candidate_key > latest_key:
                    latest = candidate
        rows = [
            ("Diretoria", "DM"),
            ("Unidade de análise", "Turma"),
            ("Periodicidade", "Por turma — sem semestre ou ano como eixo"),
            ("Área em foco", latest.get("area_name") if latest else DM_AREAS["CTE"]),
            ("Turma em foco", latest.get("cohort_label") if latest else "—"),
            ("Data de abertura", _to_date(latest.get("opening_date")) if latest else None),
            ("Data de corte", self.as_of),
        ]
        for row, (label, value) in enumerate(rows, 6):
            ws.cell(row, 1).value = label
            ws.cell(row, 2).value = value
            ws.cell(row, 1).font = Font(name="Arial", size=9, bold=True, color=GRAY_TEXT)
            ws.cell(row, 2).font = Font(name="Arial", size=10, bold=True, color=GREEN_DARK)
            ws.cell(row, 1).fill = PatternFill("solid", fgColor=GRAY)
            ws.cell(row, 2).fill = PatternFill("solid", fgColor=PALE_GREEN)
            ws.cell(row, 1).border = ws.cell(row, 2).border = self.border_box
            if isinstance(value, date):
                ws.cell(row, 2).number_format = DATE_FMT
        ws.column_dimensions["A"].width = 28
        ws.column_dimensions["B"].width = 52

    def _build_turmas(self):
        ws = self.ws["TURMAS"]
        self._title(
            ws,
            "BASE DE TURMAS · DM",
            "Uma linha por área e turma formalmente aberta. Os campos SEI preservam a origem sem substituir a data institucional de abertura.",
            end_col=13,
        )
        headers = [
            "Área", "Nome da área", "Turma", "Data de abertura", "Vagas autorizadas", "Status",
            "Observações", "Dado demonstrativo", "Total de alunos", "Chave da turma",
            "Origem", "Rótulo original no SEI", "Última leitura no SEI",
        ]
        self._headers(ws, 4, headers)
        sorted_rows = sorted(self.cohorts, key=lambda row: (row.get("area_code", ""), int(row.get("cohort_number") or 0), (row.get("opening_date") or "")))
        for index, row in enumerate(sorted_rows, 5):
            values = [
                row.get("area_code"), row.get("area_name"), row.get("cohort_number"), _to_date(row.get("opening_date")),
                row.get("vacancies_authorized"), row.get("status"), row.get("notes"), "Sim" if row.get("is_demo") else "Não",
                f'=COUNTIFS(ALUNOS!$A:$A,A{index},ALUNOS!$B:$B,C{index})',
                f'=A{index}&"-T"&C{index}',
                row.get("source_system") or ("DEMO" if row.get("is_demo") else "MANUAL"),
                row.get("sei_raw_label"),
                row.get("last_seen_sei_at"),
            ]
            for col, value in enumerate(values, 1):
                ws.cell(index, col).value = value
                ws.cell(index, col).font = self.linked_font if col <= 8 else self.body_font
                ws.cell(index, col).border = Border(bottom=Side(style="hair", color=BORDER))
                ws.cell(index, col).alignment = Alignment(vertical="center", wrap_text=True)
            ws.cell(index, 4).number_format = DATE_FMT
            ws.cell(index, 5).number_format = INTEGER_FMT
            ws.cell(index, 9).number_format = INTEGER_FMT
            ws.cell(index, 13).number_format = "dd/mm/yyyy hh:mm"
        widths = [11, 36, 10, 16, 17, 17, 35, 18, 15, 16, 15, 27, 22]
        for c, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = width
        end = max(5, 4 + len(sorted_rows))
        ws.auto_filter.ref = f"A4:M{end}"
        ws.freeze_panes = "A5"
        ws.print_title_rows = "1:4"

    def _build_alunos(self):
        ws = self.ws["ALUNOS"]
        self._title(
            ws,
            "BASE DE ALUNOS · DM",
            "Vínculo individual com turma, ingresso, qualificação e defesa. Titulação e diploma digital são acrescentados sem alterar a estrutura histórica do painel.",
            end_col=28,
        )
        headers = [
            "Área", "Turma", "Matrícula", "Nome do aluno", "Data de ingresso", "Data de qualificação",
            "Data de defesa", "Data de defesa marcada", "Status", "Data de saída", "Orientador",
            "Linha de pesquisa", "Observações", "Dado demonstrativo", "Meses até a defesa",
            "Risco >18m sem qualificação", "Risco >24m sem defesa marcada", "Risco >30m sem defesa",
            "Defesa em até 24m", "Origem", "Situação bruta no SEI", "Última leitura no SEI",
            "Ingresso provisório", "Última consulta acadêmica no SEI",
            "Data de titulação", "Situação do diploma", "Última atualização de titulação", "Atualizado por",
        ]
        self._headers(ws, 4, headers)
        sorted_rows = sorted(self.students, key=lambda row: (row.get("area_code", ""), int(row.get("cohort_number") or 0), row.get("student_name", "")))
        for index, row in enumerate(sorted_rows, 5):
            raw = [
                row.get("area_code"), row.get("cohort_number"), row.get("student_code"), row.get("student_name"),
                _to_date(row.get("entry_date")), _to_date(row.get("qualification_date")), _to_date(row.get("defense_date")),
                _to_date(row.get("defense_scheduled_date")), row.get("status"), _to_date(row.get("exit_date")),
                row.get("advisor"), row.get("research_line"), row.get("notes"), "Sim" if row.get("is_demo") else "Não",
            ]
            for col, value in enumerate(raw, 1):
                ws.cell(index, col).value = value
                ws.cell(index, col).font = self.linked_font
                ws.cell(index, col).border = Border(bottom=Side(style="hair", color=BORDER))
                ws.cell(index, col).alignment = Alignment(vertical="center", wrap_text=True)
            for c in (5, 6, 7, 8, 10):
                ws.cell(index, c).number_format = DATE_FMT
            ws.cell(index, 20).value = row.get("source_system") or ("DEMO" if row.get("is_demo") else "MANUAL")
            ws.cell(index, 21).value = row.get("sei_raw_status")
            ws.cell(index, 22).value = row.get("last_seen_sei_at")
            for c in range(20, 23):
                ws.cell(index, c).font = self.linked_font
                ws.cell(index, c).border = Border(bottom=Side(style="hair", color=BORDER))
                ws.cell(index, c).alignment = Alignment(vertical="center", wrap_text=True)
            ws.cell(index, 22).number_format = "dd/mm/yyyy hh:mm"
            ws.cell(index, 23).value = "Sim" if row.get("entry_date_estimated") else "Não"
            ws.cell(index, 24).value = row.get("last_course_dates_sei_at")
            for c in range(23, 25):
                ws.cell(index, c).font = self.linked_font
                ws.cell(index, c).border = Border(bottom=Side(style="hair", color=BORDER))
                ws.cell(index, c).alignment = Alignment(vertical="center", wrap_text=True)
            ws.cell(index, 24).number_format = "dd/mm/yyyy hh:mm"

            # New diploma fields are appended after the legacy 24-column layout so
            # formulas and downstream integrations keep their historical positions.
            ws.cell(index, 25).value = _to_date(row.get("graduation_date"))
            ws.cell(index, 26).value = row.get("diploma_status")
            ws.cell(index, 27).value = row.get("graduation_updated_at")
            ws.cell(index, 28).value = row.get("graduation_updated_by")
            for c in range(25, 29):
                ws.cell(index, c).font = self.linked_font
                ws.cell(index, c).border = Border(bottom=Side(style="hair", color=BORDER))
                ws.cell(index, c).alignment = Alignment(vertical="center", wrap_text=True)
            ws.cell(index, 25).number_format = DATE_FMT
            ws.cell(index, 27).number_format = "dd/mm/yyyy hh:mm"

            # DM-02 preserves the legacy column positions O:S.
            ws.cell(index, 15).value = f'=IF(OR(E{index}="",G{index}="",W{index}="Sim"),"",ROUND(YEARFRAC(E{index},G{index})*12,2))'
            ws.cell(index, 16).value = f'=IF(AND(W{index}<>"Sim",I{index}="Ativo",E{index}<>"",F{index}="",YEARFRAC(E{index},PARAMETROS!$B$12)*12>18),1,0)'
            ws.cell(index, 17).value = f'=IF(AND(W{index}<>"Sim",I{index}="Ativo",E{index}<>"",G{index}="",H{index}="",YEARFRAC(E{index},PARAMETROS!$B$12)*12>24),1,0)'
            ws.cell(index, 18).value = f'=IF(AND(W{index}<>"Sim",I{index}="Ativo",E{index}<>"",G{index}="",YEARFRAC(E{index},PARAMETROS!$B$12)*12>30),1,0)'
            ws.cell(index, 19).value = f'=IF(AND(W{index}<>"Sim",I{index}="Titulado",O{index}<>"",O{index}<=24),1,0)'
            for c in range(15, 20):
                ws.cell(index, c).font = self.body_font
                ws.cell(index, c).fill = PatternFill("solid", fgColor=GRAY)
                ws.cell(index, c).border = Border(bottom=Side(style="hair", color=BORDER))
            ws.cell(index, 15).number_format = DECIMAL_FMT
        widths = [10, 10, 16, 30, 16, 18, 16, 21, 13, 15, 24, 26, 35, 18, 18, 22, 25, 22, 20, 15, 20, 22, 18, 24, 16, 18, 24, 28]
        for c, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = width
        end = max(5, 4 + len(sorted_rows))
        ws.auto_filter.ref = f"A4:X{end}"
        ws.freeze_panes = "A5"
        ws.print_title_rows = "1:4"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3

    def _build_lists(self):
        ws = self.ws["LISTAS DE APOIO"]
        self._title(ws, "LISTAS DE APOIO · DM", "Valores institucionais permitidos nos formulários e modelos.", end_col=8)
        self._headers(ws, 4, ["Código da área", "Área", "Status da turma", "Status do aluno", "Situação do diploma"])
        max_len = max(len(DM_AREAS), len(COHORT_STATUSES), len(STUDENT_STATUSES), len(DIPLOMA_STATUSES))
        for i in range(max_len):
            row = 5 + i
            area_items = list(DM_AREAS.items())
            if i < len(area_items):
                ws.cell(row, 1).value = area_items[i][0]
                ws.cell(row, 2).value = area_items[i][1]
            if i < len(COHORT_STATUSES):
                ws.cell(row, 3).value = COHORT_STATUSES[i]
            if i < len(STUDENT_STATUSES):
                ws.cell(row, 4).value = STUDENT_STATUSES[i]
            if i < len(DIPLOMA_STATUSES):
                ws.cell(row, 5).value = DIPLOMA_STATUSES[i]
            for c in range(1, 6):
                ws.cell(row, c).font = self.static_font
                ws.cell(row, c).border = Border(bottom=Side(style="hair", color=BORDER))
        for c, width in enumerate([14, 40, 20, 20, 22], 1):
            ws.column_dimensions[get_column_letter(c)].width = width

    def _build_calc(self):
        ws = self.ws["CALC"]
        headers = [
            "Área", "Nome da área", "Turma", "Rótulo", "Data de abertura", "Vagas", "Total", "Ativos",
            "Titulados", "Desligados", "Trancados", "Ocupação", "Evasão", "Tempo médio", "Tempo mediano",
            "Defesas <=24m", "Titulação no prazo", "Risco qualificação", "Risco defesa marcada", "Risco >30m",
            "Status DM-01", "Status DM-02", "Meta 24m", "Meta titulação", "Meta evasão", "Meta ocupação",
        ]
        self._headers(ws, 4, headers)
        cohorts = sorted(self.cohorts, key=lambda row: (row.get("area_code", ""), int(row.get("cohort_number") or 0), (row.get("opening_date") or "")))
        for index, cohort in enumerate(cohorts, 5):
            area = cohort.get("area_code")
            number = int(cohort.get("cohort_number") or 0)
            ws.cell(index, 1).value = area
            ws.cell(index, 2).value = cohort.get("area_name") or DM_AREAS.get(area, area)
            ws.cell(index, 3).value = number
            ws.cell(index, 4).value = f"Turma {number}"
            ws.cell(index, 5).value = _to_date(cohort.get("opening_date"))
            ws.cell(index, 6).value = cohort.get("vacancies_authorized")
            ws.cell(index, 7).value = f'=COUNTIFS(ALUNOS!$A:$A,A{index},ALUNOS!$B:$B,C{index})'
            ws.cell(index, 8).value = f'=COUNTIFS(ALUNOS!$A:$A,A{index},ALUNOS!$B:$B,C{index},ALUNOS!$I:$I,"Ativo")'
            ws.cell(index, 9).value = f'=COUNTIFS(ALUNOS!$A:$A,A{index},ALUNOS!$B:$B,C{index},ALUNOS!$I:$I,"Titulado")'
            ws.cell(index, 10).value = f'=COUNTIFS(ALUNOS!$A:$A,A{index},ALUNOS!$B:$B,C{index},ALUNOS!$I:$I,"Desligado")'
            ws.cell(index, 11).value = f'=COUNTIFS(ALUNOS!$A:$A,A{index},ALUNOS!$B:$B,C{index},ALUNOS!$I:$I,"Trancado")'
            ws.cell(index, 12).value = f'=IFERROR(G{index}/F{index},"")'
            ws.cell(index, 13).value = f'=IFERROR(J{index}/G{index},"")'
            ws.cell(index, 14).value = f'=IFERROR(AVERAGEIFS(ALUNOS!$O:$O,ALUNOS!$A:$A,A{index},ALUNOS!$B:$B,C{index}),"")'
            summary = next((item for area_item in self.dashboard["areas"] for item in area_item["cohorts"] if int(item["id"]) == int(cohort["id"])), None)
            ws.cell(index, 15).value = summary.get("median_months_to_defense") if summary else None
            ws.cell(index, 16).value = f'=SUMIFS(ALUNOS!$S:$S,ALUNOS!$A:$A,A{index},ALUNOS!$B:$B,C{index})'
            ws.cell(index, 17).value = f'=IFERROR(P{index}/G{index},"")'
            ws.cell(index, 18).value = f'=SUMIFS(ALUNOS!$P:$P,ALUNOS!$A:$A,A{index},ALUNOS!$B:$B,C{index})'
            ws.cell(index, 19).value = f'=SUMIFS(ALUNOS!$Q:$Q,ALUNOS!$A:$A,A{index},ALUNOS!$B:$B,C{index})'
            ws.cell(index, 20).value = f'=SUMIFS(ALUNOS!$R:$R,ALUNOS!$A:$A,A{index},ALUNOS!$B:$B,C{index})'
            ws.cell(index, 21).value = summary.get("dm01_status") if summary else "Sem dados"
            ws.cell(index, 22).value = summary.get("dm02_status") if summary else "Sem dados"
            ws.cell(index, 23).value = 24
            ws.cell(index, 24).value = 0.70
            ws.cell(index, 25).value = 0.05
            ws.cell(index, 26).value = 0.85
            ws.cell(index, 5).number_format = DATE_FMT
            for c in (12, 13, 17, 24, 25, 26):
                ws.cell(index, c).number_format = PERCENT_FMT
            for c in (14, 15, 23):
                ws.cell(index, c).number_format = DECIMAL_FMT
            for c in range(1, 27):
                ws.cell(index, c).font = self.body_font
        for code in DM_AREAS:
            rows = [i for i, cohort in enumerate(cohorts, 5) if cohort.get("area_code") == code]
            self.calc_ranges[code] = (min(rows), max(rows)) if rows else (0, 0)
        ws.sheet_state = "hidden"

    def _panel_setup(self, ws, title: str, subtitle: str):
        self._title(ws, title, subtitle, end_col=12)
        for col in range(1, 13):
            ws.column_dimensions[get_column_letter(col)].width = 14
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.freeze_panes = "A4"

    def _area_section(self, ws, row: int, area: dict[str, Any], *, panel: bool = True) -> int:
        code = area["area_code"]
        self._section(ws, row, f'{code} · {area["area_name"]}')
        totals = area["totals"]
        r = row + 1
        cards = [
            (1, 2, "Turmas cadastradas", totals["cohort_count"], INTEGER_FMT, PALE_GREEN),
            (3, 4, "Alunos cadastrados", totals["total_students"], INTEGER_FMT, PALE_BLUE),
            (5, 6, "Ativos", totals["active_students"], INTEGER_FMT, PALE_TEAL),
            (7, 8, "Titulados", totals["graduated_students"], INTEGER_FMT, PALE_GREEN),
            (9, 10, "Evasão acumulada", (totals["dropout_rate_pct"] or 0) / 100 if totals["dropout_rate_pct"] is not None else None, PERCENT_FMT, PALE_GOLD),
            (11, 12, "Tempo médio até defesa", totals["average_months_to_defense"], DECIMAL_FMT, PALE_TEAL),
        ]
        for start, end, label, value, fmt, fill in cards:
            self._kpi_card(ws, start_col=start, end_col=end, label_row=r, value_row=r + 1, label=label, value=value, number_format=fmt, fill=fill)
        start, end = self.calc_ranges.get(code, (0, 0))
        if start:
            compact_start = max(start, end - 7)
            chart1 = BarChart()
            chart1.type = "col"
            chart1.grouping = "clustered"
            chart1.add_data(Reference(self.ws["CALC"], min_col=8, max_col=10, min_row=4, max_row=end), titles_from_data=True)
            chart1.set_categories(Reference(self.ws["CALC"], min_col=4, min_row=compact_start, max_row=end))
            for idx, series in enumerate(chart1.series):
                source_col = 8 + idx
                series.val.numRef.f = f"'CALC'!${get_column_letter(source_col)}${compact_start}:${get_column_letter(source_col)}${end}"
            self._chart_base(chart1, title="Evolução dos alunos por turma", x_title="Turma", y_title="Alunos", width=16.2, height=8.5)
            self._style_series(chart1, [CHART_GREEN, CHART_TEAL, CHART_RED])
            ws.add_chart(chart1, f"A{r+4}")

            chart2 = LineChart()
            chart2.add_data(Reference(self.ws["CALC"], min_col=14, max_col=15, min_row=4, max_row=end), titles_from_data=True)
            chart2.add_data(Reference(self.ws["CALC"], min_col=23, min_row=4, max_row=end), titles_from_data=True)
            chart2.set_categories(Reference(self.ws["CALC"], min_col=4, min_row=compact_start, max_row=end))
            for idx, series in enumerate(chart2.series):
                source_col = [14, 15, 23][idx]
                series.val.numRef.f = f"'CALC'!${get_column_letter(source_col)}${compact_start}:${get_column_letter(source_col)}${end}"
            self._chart_base(chart2, title="Tempo até a defesa por turma", x_title="Turma", y_title="Meses", width=16.2, height=8.5)
            self._style_series(chart2, [CHART_GREEN, CHART_TEAL, CHART_META], line=True)
            chart2.dataLabels = None
            chart2.series[0].dLbls = DataLabelList(showVal=True, showSerName=False, showCatName=False)
            ws.add_chart(chart2, f"G{r+4}")
        table_row = r + 21
        self._headers(ws, table_row, ["Turma", "Abertura", "Total", "Ativos", "Titulados", "Desligados", "Ocupação", "Evasão", "Média defesa", "No prazo", "Risco", "Situação"])
        cohorts = area["cohorts"][-8:] if panel else area["cohorts"]
        for offset, item in enumerate(cohorts, table_row + 1):
            values = [
                item["cohort_label"], _to_date(item["opening_date"]), item["total_students"], item["active_students"],
                item["graduated_students"], item["dropped_students"],
                (item["occupancy_pct"] or 0) / 100 if item["occupancy_pct"] is not None else None,
                (item["dropout_rate_pct"] or 0) / 100 if item["dropout_rate_pct"] is not None else None,
                item["average_months_to_defense"],
                (item["on_time_graduation_pct"] or 0) / 100 if item["on_time_graduation_pct"] is not None else None,
                item["active_over_18_no_qualification"] + item["active_over_30_no_defense"],
                item["dm02_status"],
            ]
            for col, value in enumerate(values, 1):
                cell = ws.cell(offset, col)
                cell.value = value
                cell.font = self.body_font
                cell.border = Border(bottom=Side(style="hair", color=BORDER))
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.cell(offset, 2).number_format = DATE_FMT
            for col in (7, 8, 10):
                ws.cell(offset, col).number_format = PERCENT_FMT
            ws.cell(offset, 9).number_format = DECIMAL_FMT
            ws.cell(offset, 12).fill = PatternFill("solid", fgColor=_formula_status_color(item["dm02_status"]))
        return table_row + max(len(cohorts), 1) + 2

    def _build_panel(self):
        ws = self.ws["PAINEL"]
        self._panel_setup(
            ws,
            "PAINEL DE GESTÃO — DM",
            "Diretoria de Mestrado · evolução por turma e por área, sem uso de semestre ou ano como eixo.",
        )
        row = 4
        for index, area in enumerate(self.dashboard["areas"]):
            row = self._area_section(ws, row, area, panel=True)
            if index < len(self.dashboard["areas"]) - 1:
                ws.row_breaks.append(Break(id=row - 1))
        ws.print_area = f"A1:L{row}"

    def _build_dm01(self):
        ws = self.ws["DM-01 EVOLUCAO"]
        self._panel_setup(
            ws,
            "DM-01 · EVOLUÇÃO DO NÚMERO DE ALUNOS",
            "Leitura por turma, separada entre Ciência, Tecnologia e Educação e Saúde e Desigualdade Social.",
        )
        row = 4
        for index, area in enumerate(self.dashboard["areas"]):
            self._section(ws, row, f'{area["area_code"]} · {area["area_name"]}')
            start, end = self.calc_ranges.get(area["area_code"], (0, 0))
            if start:
                chart = BarChart()
                chart.type = "col"
                chart.grouping = "clustered"
                chart.add_data(Reference(self.ws["CALC"], min_col=7, max_col=11, min_row=4, max_row=end), titles_from_data=True)
                chart.set_categories(Reference(self.ws["CALC"], min_col=4, min_row=start, max_row=end))
                for idx, series in enumerate(chart.series):
                    source_col = 7 + idx
                    series.val.numRef.f = f"'CALC'!${get_column_letter(source_col)}${start}:${get_column_letter(source_col)}${end}"
                self._chart_base(chart, title="Composição da turma", x_title="Turma", y_title="Alunos", width=33.0, height=9.0)
                self._style_series(chart, [CHART_BLUE, CHART_GREEN, CHART_TEAL, CHART_RED, CHART_GOLD])
                ws.add_chart(chart, f"A{row+1}")
            table_row = row + 20
            self._headers(ws, table_row, ["Turma", "Abertura", "Vagas", "Ingressantes", "Ativos", "Titulados", "Desligados", "Trancados", "Ocupação", "Evasão", "Risco >30m", "Status"])
            for r, item in enumerate(area["cohorts"], table_row + 1):
                vals = [
                    item["cohort_label"], _to_date(item["opening_date"]), item.get("vacancies_authorized"), item["total_students"],
                    item["active_students"], item["graduated_students"], item["dropped_students"], item["locked_students"],
                    (item["occupancy_pct"] or 0) / 100 if item["occupancy_pct"] is not None else None,
                    (item["dropout_rate_pct"] or 0) / 100 if item["dropout_rate_pct"] is not None else None,
                    item["active_over_30_no_defense"], item["dm01_status"],
                ]
                for c, value in enumerate(vals, 1):
                    ws.cell(r, c).value = value
                    ws.cell(r, c).font = self.body_font
                    ws.cell(r, c).border = Border(bottom=Side(style="hair", color=BORDER))
                    ws.cell(r, c).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                ws.cell(r, 2).number_format = DATE_FMT
                ws.cell(r, 9).number_format = PERCENT_FMT
                ws.cell(r, 10).number_format = PERCENT_FMT
                ws.cell(r, 12).fill = PatternFill("solid", fgColor=_formula_status_color(item["dm01_status"]))
            row = table_row + max(1, len(area["cohorts"])) + 3
            if index == 0:
                ws.row_breaks.append(Break(id=row - 1))
        ws.print_area = f"A1:L{row}"

    def _build_dm02(self):
        ws = self.ws["DM-02 DEFESAS"]
        self._panel_setup(
            ws,
            "DM-02 · TEMPO ATÉ A DEFESA",
            "Datas individuais de ingresso e defesa produzem o tempo médio, a mediana e a titulação no prazo por turma.",
        )
        row = 4
        for index, area in enumerate(self.dashboard["areas"]):
            self._section(ws, row, f'{area["area_code"]} · {area["area_name"]}')
            start, end = self.calc_ranges.get(area["area_code"], (0, 0))
            if start:
                chart1 = LineChart()
                chart1.add_data(Reference(self.ws["CALC"], min_col=14, max_col=15, min_row=4, max_row=end), titles_from_data=True)
                chart1.add_data(Reference(self.ws["CALC"], min_col=23, min_row=4, max_row=end), titles_from_data=True)
                chart1.set_categories(Reference(self.ws["CALC"], min_col=4, min_row=start, max_row=end))
                for idx, series in enumerate(chart1.series):
                    source_col = [14, 15, 23][idx]
                    series.val.numRef.f = f"'CALC'!${get_column_letter(source_col)}${start}:${get_column_letter(source_col)}${end}"
                self._chart_base(chart1, title="Tempo médio e mediano até a defesa", x_title="Turma", y_title="Meses", width=16.2, height=8.6)
                self._style_series(chart1, [CHART_GREEN, CHART_TEAL, CHART_META], line=True)
                chart1.dataLabels = None
                chart1.series[0].dLbls = DataLabelList(showVal=True, showSerName=False, showCatName=False)
                ws.add_chart(chart1, f"A{row+1}")

                chart2 = BarChart()
                chart2.type = "col"
                chart2.add_data(Reference(self.ws["CALC"], min_col=17, min_row=4, max_row=end), titles_from_data=True)
                chart2.add_data(Reference(self.ws["CALC"], min_col=24, min_row=4, max_row=end), titles_from_data=True)
                chart2.set_categories(Reference(self.ws["CALC"], min_col=4, min_row=start, max_row=end))
                for idx, series in enumerate(chart2.series):
                    source_col = [17, 24][idx]
                    series.val.numRef.f = f"'CALC'!${get_column_letter(source_col)}${start}:${get_column_letter(source_col)}${end}"
                self._chart_base(chart2, title="Titulação em até 24 meses", x_title="Turma", y_title="Percentual", width=16.2, height=8.6)
                self._style_series(chart2, [CHART_GREEN, CHART_META])
                chart2.y_axis.numFmt = '0%'
                ws.add_chart(chart2, f"G{row+1}")
            table_row = row + 20
            self._headers(ws, table_row, ["Turma", "Defesas", "Média", "Mediana", "No prazo", "Risco qualificação", "Sem defesa marcada", ">30m", "Orientadores", "Orientandos/docente", "Última defesa", "Status"])
            for r, item in enumerate(area["cohorts"], table_row + 1):
                vals = [
                    item["cohort_label"], item["defenses_count"], item["average_months_to_defense"], item["median_months_to_defense"],
                    (item["on_time_graduation_pct"] or 0) / 100 if item["on_time_graduation_pct"] is not None else None,
                    item["active_over_18_no_qualification"], item["active_over_24_no_scheduled_defense"], item["active_over_30_no_defense"],
                    item["active_advisor_count"], item["advisees_per_faculty"], _to_date(item["last_defense_date"]), item["dm02_status"],
                ]
                for c, value in enumerate(vals, 1):
                    ws.cell(r, c).value = value
                    ws.cell(r, c).font = self.body_font
                    ws.cell(r, c).border = Border(bottom=Side(style="hair", color=BORDER))
                    ws.cell(r, c).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                for c in (3, 4, 10):
                    ws.cell(r, c).number_format = DECIMAL_FMT
                ws.cell(r, 5).number_format = PERCENT_FMT
                ws.cell(r, 11).number_format = DATE_FMT
                ws.cell(r, 12).fill = PatternFill("solid", fgColor=_formula_status_color(item["dm02_status"]))
            row = table_row + max(1, len(area["cohorts"])) + 3
            if index == 0:
                ws.row_breaks.append(Break(id=row - 1))
        ws.print_area = f"A1:L{row}"

    def _build_matrix(self):
        ws = self.ws["MATRIZ"]
        self._title(ws, "MATRIZ EXECUTIVA · DM", "Todas as turmas das duas áreas em uma única leitura.", end_col=12)
        headers = ["Área", "Turma", "Abertura", "Alunos", "Ativos", "Titulados", "Evasão", "Ocupação", "Média defesa", "No prazo", "Riscos", "Status"]
        self._headers(ws, 4, headers)
        all_rows = [item for area in self.dashboard["areas"] for item in area["cohorts"]]
        for r, item in enumerate(all_rows, 5):
            vals = [
                item["area_name"], item["cohort_label"], _to_date(item["opening_date"]), item["total_students"], item["active_students"],
                item["graduated_students"], (item["dropout_rate_pct"] or 0) / 100 if item["dropout_rate_pct"] is not None else None,
                (item["occupancy_pct"] or 0) / 100 if item["occupancy_pct"] is not None else None,
                item["average_months_to_defense"], (item["on_time_graduation_pct"] or 0) / 100 if item["on_time_graduation_pct"] is not None else None,
                item["active_over_18_no_qualification"] + item["active_over_30_no_defense"], item["dm02_status"],
            ]
            for c, value in enumerate(vals, 1):
                ws.cell(r, c).value = value
                ws.cell(r, c).font = self.body_font
                ws.cell(r, c).border = Border(bottom=Side(style="hair", color=BORDER))
                ws.cell(r, c).alignment = Alignment(vertical="center", wrap_text=True)
            ws.cell(r, 3).number_format = DATE_FMT
            for c in (7, 8, 10):
                ws.cell(r, c).number_format = PERCENT_FMT
            ws.cell(r, 9).number_format = DECIMAL_FMT
            ws.cell(r, 12).fill = PatternFill("solid", fgColor=_formula_status_color(item["dm02_status"]))
        widths = [36, 12, 15, 11, 11, 11, 12, 12, 14, 12, 10, 18]
        for c, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = width
        ws.freeze_panes = "A5"
        ws.auto_filter.ref = f"A4:L{max(5,4+len(all_rows))}"

    def _build_quality(self):
        ws = self.ws["QUALIDADE E GOVERNANÇA"]
        self._title(ws, "QUALIDADE E GOVERNANÇA · DM", "Controles para que o acompanhamento por turma permaneça auditável.", end_col=10)
        self._section(ws, 4, "CONTROLES OBRIGATÓRIOS", end_col=10)
        rules = [
            "Toda turma deve possuir área e número; a data de abertura é opcional quando a trajetória individual vem do SEI.",
            "Turmas das duas áreas são independentes. Uma numeração existente em CTE não cria a mesma turma em SDS.",
            "A data individual de ingresso pode ser confirmada pelo SEI e é o marco usado pelo DM-02.",
            "A data de conclusão retornada pelo SEI corresponde à defesa aprovada e atualiza diretamente a Data de defesa, confirmando Titulado.",
            "Desligados continuam no denominador da titulação no prazo da turma.",
            "O painel compara turmas dentro da mesma área, ordenadas pelo número da turma; a abertura é metadado opcional.",
            "Os gráficos mostram uma janela visual compacta, mas as bases TURMAS e ALUNOS preservam todo o histórico.",
            "O relatório integral do SEI descobre turmas e matrículas; a consulta individual confirma início e defesa sem inventar qualificação ou saída.",
            "A ausência de um aluno em uma sincronização posterior não exclui nem altera automaticamente o registro histórico.",
        ]
        for row, rule in enumerate(rules, 6):
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
            ws.cell(row, 1).value = f"• {rule}"
            ws.cell(row, 1).font = self.body_font
            ws.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="center")
            ws.cell(row, 1).fill = PatternFill("solid", fgColor=PALE_GREEN if row % 2 == 0 else WHITE)
            ws.row_dimensions[row].height = 30
        quality_row = 6 + len(rules) + 1
        self._section(ws, quality_row, "QUALIDADE DA INTEGRAÇÃO SEI", end_col=10)
        sei_students = sum(1 for row in self.students if "SEI" in str(row.get("source_system") or "").upper())
        quality_items = [
            ("Alunos na base", len(self.students)),
            ("Ingressos individuais confirmados", sum(1 for row in self.students if row.get("entry_date") and not row.get("entry_date_estimated"))),
            ("Ingressos ainda pendentes", sum(1 for row in self.students if not row.get("entry_date"))),
            ("Datas individuais consultadas no SEI", sum(1 for row in self.students if row.get("last_course_dates_sei_at"))),
            ("Defesas confirmadas após consulta SEI", sum(1 for row in self.students if row.get("defense_date") and row.get("last_course_dates_sei_at"))),
            ("Alunos reconhecidos pelo SEI", sei_students),
        ]
        self._headers(ws, quality_row + 1, ["Controle", "Quantidade"])
        for r, (metric, value) in enumerate(quality_items, quality_row + 2):
            ws.cell(r, 1).value = metric
            ws.cell(r, 2).value = value
            ws.cell(r, 1).font = self.body_font
            ws.cell(r, 2).font = Font(name="Arial", size=9, bold=True, color=GREEN_DARK)
            ws.cell(r, 2).number_format = INTEGER_FMT
            ws.cell(r, 1).border = ws.cell(r, 2).border = Border(bottom=Side(style="hair", color=BORDER))

        metas_row = quality_row + 10
        self._section(ws, metas_row, "METAS INICIAIS DO DOCUMENTO", end_col=10)
        metas = [
            ("Membros por turma", "Definida na plataforma por vigência"),
            ("Tempo médio até a defesa", "≤ 24 meses (referência padrão)"),
        ]
        self._headers(ws, metas_row + 1, ["Métrica", "Referência inicial"])
        for r, (metric, target) in enumerate(metas, metas_row + 2):
            ws.cell(r, 1).value = metric
            ws.cell(r, 2).value = target
            ws.cell(r, 1).font = self.body_font
            ws.cell(r, 2).font = Font(name="Arial", size=9, bold=True, color=GREEN_DARK)
            ws.cell(r, 1).border = ws.cell(r, 2).border = Border(bottom=Side(style="hair", color=BORDER))
        ws.column_dimensions["A"].width = 48
        ws.column_dimensions["B"].width = 26

    def _finish(self):
        self.wb.active = self.wb.sheetnames.index("PAINEL")
        self.wb.calculation.fullCalcOnLoad = True
        self.wb.calculation.forceFullCalc = True
        self.wb.calculation.calcMode = "auto"
        tab_colors = {
            "PAINEL": GREEN_DARK,
            "DM-01 EVOLUCAO": GREEN,
            "DM-02 DEFESAS": GREEN,
            "TURMAS": LINKED_GREEN,
            "ALUNOS": LINKED_GREEN,
            "QUALIDADE E GOVERNANÇA": CHART_BLUE,
            "MATRIZ": CHART_GOLD,
        }
        for name, color in tab_colors.items():
            self.ws[name].sheet_properties.tabColor = color
        for ws in self.wb.worksheets:
            ws.sheet_view.showGridLines = False
            ws.sheet_properties.pageSetUpPr.fitToPage = True


def build_dm_workbook(
    cohorts: Iterable[dict[str, Any]],
    students: Iterable[dict[str, Any]],
    *,
    area_code: str | None = None,
    cohort_id: int | None = None,
    as_of: date | None = None,
) -> BytesIO:
    builder = DMWorkbookBuilder(cohorts, students, area_code=area_code, cohort_id=cohort_id, as_of=as_of)
    workbook = builder.build()
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def build_dm_import_template(kind: str) -> BytesIO:
    """Build a deliberately simple, one-sheet import model.

    The template contains no sample records, so a user cannot accidentally import
    demonstration values. Guidance is attached to the header cells as comments and
    the only worksheet is the normal tabular ``DADOS`` sheet.
    """
    normalized = str(kind or "").strip().lower()
    if normalized not in {"turmas", "alunos"}:
        raise ValueError("Modelo DM inválido. Use turmas ou alunos.")

    wb = Workbook()
    ws = wb.active
    ws.title = "DADOS"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"

    if normalized == "turmas":
        headers = ["Área", "Turma", "Data de abertura", "Vagas autorizadas", "Status", "Observações"]
        widths = [40, 12, 19, 20, 18, 48]
        comments = {
            "Área": "Use CTE para Ciência, Tecnologia e Educação ou SDS para Saúde e Desigualdade Social.",
            "Turma": "Número da turma dentro da área. A mesma numeração pode existir de forma independente nas duas áreas.",
            "Data de abertura": "Campo opcional. Data institucional de abertura da turma, quando conhecida. Não é usada como ingresso individual.",
            "Vagas autorizadas": "Quantidade de vagas previstas para a turma. Campo opcional.",
            "Status": "Valores aceitos: Planejada, Aberta, Em andamento ou Encerrada.",
            "Observações": "Campo opcional para informações institucionais sobre a turma.",
        }
    else:
        headers = [
            "Área", "Turma", "Matrícula", "Nome do aluno", "Data de ingresso",
            "Data de qualificação", "Data de defesa", "Data de defesa marcada", "Status", "Data de saída", "Orientador",
            "Linha de pesquisa", "Observações",
        ]
        widths = [40, 12, 18, 30, 19, 20, 18, 23, 15, 17, 25, 28, 48]
        comments = {
            "Área": "Use CTE ou SDS. A área e o número da turma devem corresponder a uma turma já cadastrada.",
            "Turma": "Número da turma na área indicada.",
            "Matrícula": "Código institucional único do aluno.",
            "Nome do aluno": "Nome completo do aluno.",
            "Data de ingresso": "Opcional. Se vazia, permanece pendente até confirmação individual no SEI ou preenchimento manual.",
            "Data de qualificação": "Campo opcional. Formato recomendado: DD/MM/AAAA.",
            "Data de defesa": "Campo opcional. Uma defesa registrada confirma Titulado. A consulta individual ao SEI também atualiza este mesmo campo.",
            "Data de defesa marcada": "Campo opcional para estudantes ativos com defesa agendada.",
            "Status": "Valores aceitos: Ativo, Titulado ou Desligado.",
            "Data de saída": "Informe para alunos Desligados quando a data for conhecida.",
            "Orientador": "Campo opcional.",
            "Linha de pesquisa": "Campo opcional.",
            "Observações": "Campo opcional.",
        }

    for col, header in enumerate(headers, 1):
        cell = ws.cell(1, col, header)
        cell.fill = PatternFill("solid", fgColor=GREEN_DARK)
        cell.font = Font(name="Arial", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color=GREEN))
        cell.comment = Comment(comments.get(header, ""), "Data UNIVC")

    # Keep a practical blank input area with no fake records. Imported values are
    # shown in blue, following the same input convention used by the other panels.
    for row in range(2, 502):
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row, col)
            cell.font = Font(name="Arial", size=9, color=INPUT_BLUE)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = Border(bottom=Side(style="hair", color=BORDER))
        if normalized == "turmas":
            ws.cell(row, 3).number_format = DATE_FMT
        else:
            for col in (5, 6, 7, 8, 10):
                ws.cell(row, col).number_format = DATE_FMT

    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 34
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}501"

    area_dv = DataValidation(type="list", formula1='"CTE,SDS"', allow_blank=False)
    area_dv.error = "Selecione CTE ou SDS."
    area_dv.errorTitle = "Área inválida"
    area_dv.prompt = "CTE = Ciência, Tecnologia e Educação; SDS = Saúde e Desigualdade Social."
    area_dv.promptTitle = "Área do Mestrado"
    area_dv.showInputMessage = True
    area_dv.showErrorMessage = True
    ws.add_data_validation(area_dv)
    area_dv.add("A2:A501")

    status_values = COHORT_STATUSES if normalized == "turmas" else STUDENT_STATUSES
    status_col = 5 if normalized == "turmas" else 9
    status_dv = DataValidation(type="list", formula1='"' + ",".join(status_values) + '"', allow_blank=False)
    status_dv.error = "Selecione um dos status disponíveis na lista."
    status_dv.errorTitle = "Status inválido"
    status_dv.showErrorMessage = True
    ws.add_data_validation(status_dv)
    status_dv.add(f"{get_column_letter(status_col)}2:{get_column_letter(status_col)}501")
    ws.print_title_rows = "1:1"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}501"
    wb.calculation.fullCalcOnLoad = True

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output

