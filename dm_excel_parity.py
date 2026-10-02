from __future__ import annotations

import json
import math
from dataclasses import asdict, dataclass
from datetime import date, datetime, timezone
from io import BytesIO
from typing import Any, Iterable, Mapping

from openpyxl import load_workbook

from dm_analytics import build_dm_dashboard
from dm_excel_official import build_dm_excel_official_artifact_from_payload, build_dm_excel_official_payload, build_dm_snapshot_context
from dm_excel_v2_builder import build_dm_v2_workbook
from excel_official.adapters.dm import DMAdapter, DM_METRICS
from excel_official.audit.quality import evaluate_metric
from excel_official.context import AdapterInput
from release_info import APP_VERSION, SCHEMA_VERSION

PARITY_REPORT_VERSION = 1


@dataclass(frozen=True, slots=True)
class DMParityCase:
    scope: str
    metric_code: str
    expected: float | int | None
    actual: float | int | None
    tolerance: float
    status: str
    delta: float | None = None


@dataclass(frozen=True, slots=True)
class WorkbookEvidence:
    engine: str
    build_ok: bool
    bytes_size: int = 0
    sheets: int = 0
    formulas: int = 0
    charts: int = 0
    error: str | None = None


@dataclass(frozen=True, slots=True)
class DMParityReport:
    source_kind: str
    generated_at: str
    system_version: str
    schema_version: str
    payload_hash: str | None
    cases: tuple[DMParityCase, ...]
    legacy_workbook: WorkbookEvidence | None
    new_workbook: WorkbookEvidence | None
    new_release_allowed: bool
    source_quality_ok: bool
    source_quality_notes: tuple[str, ...] = ()
    report_version: int = PARITY_REPORT_VERSION
    directorate: str = "DM"

    @property
    def failures(self) -> tuple[DMParityCase, ...]:
        return tuple(case for case in self.cases if case.status == "FAIL")

    @property
    def cutover_status(self) -> str:
        builds_ok = bool(
            self.legacy_workbook and self.legacy_workbook.build_ok
            and self.new_workbook and self.new_workbook.build_ok
        )
        if self.failures or not builds_ok or not self.new_release_allowed or not self.source_quality_ok:
            return "BLOCKED"
        return "READY" if self.source_kind == "production" else "CANDIDATE_PASS"

    def to_dict(self) -> dict[str, Any]:
        data = asdict(self)
        data["summary"] = {
            "cases": len(self.cases),
            "passed": sum(1 for case in self.cases if case.status == "PASS"),
            "failures": len(self.failures),
            "semantic_passed": not self.failures,
            "source_quality_ok": self.source_quality_ok,
            "new_release_allowed": self.new_release_allowed,
            "cutover_status": self.cutover_status,
        }
        return data

    def to_json(self) -> str:
        return json.dumps(self.to_dict(), ensure_ascii=False, indent=2)

    def to_markdown(self) -> str:
        lines = [
            "# DM Excel Production Parity",
            "",
            f"- Source kind: **{self.source_kind}**",
            f"- Cases: **{len(self.cases)}**",
            f"- Failures: **{len(self.failures)}**",
            f"- Source quality: **{'PASS' if self.source_quality_ok else 'FAIL'}**",
            f"- Excel Official release audit: **{'PASS' if self.new_release_allowed else 'FAIL'}**",
            f"- Cutover status: **{self.cutover_status}**",
            "",
        ]
        if self.failures:
            lines.extend(["## Failures", ""])
            for case in self.failures:
                lines.append(
                    f"- `{case.scope}` · `{case.metric_code}`: expected={case.expected!r}, actual={case.actual!r}, delta={case.delta!r}"
                )
        if self.source_quality_notes:
            lines.extend(["", "## Source quality", ""])
            lines.extend(f"- {note}" for note in self.source_quality_notes)
        lines.append("")
        return "\n".join(lines)


def _evidence(engine: str, buffer: BytesIO) -> WorkbookEvidence:
    try:
        raw = buffer.getvalue()
        wb = load_workbook(BytesIO(raw), read_only=False, data_only=False, keep_links=True)
        formulas = sum(1 for ws in wb.worksheets for row in ws.iter_rows() for cell in row if cell.data_type == "f")
        charts = sum(len(getattr(ws, "_charts", ())) for ws in wb.worksheets)
        sheets = len(wb.sheetnames)
        wb.close()
        return WorkbookEvidence(engine, True, len(raw), sheets, formulas, charts)
    except Exception as exc:  # pragma: no cover - defensive production evidence
        return WorkbookEvidence(engine, False, error=f"{type(exc).__name__}: {exc}")


def _same(expected: float | int | None, actual: float | int | None, tolerance: float) -> tuple[bool, float | None]:
    if expected is None or actual is None:
        return expected is None and actual is None, None
    delta = float(actual) - float(expected)
    return math.isclose(float(expected), float(actual), abs_tol=tolerance, rel_tol=0.0), delta


def _expected_metrics(overall: Mapping[str, Any]) -> dict[str, float | int | None]:
    def pct(key: str) -> float | None:
        value = overall.get(key)
        return None if value is None else float(value) / 100.0

    return {
        DM_METRICS["members"]: int(overall.get("total_students") or 0),
        DM_METRICS["active"]: int(overall.get("active_students") or 0),
        DM_METRICS["occupancy"]: pct("occupancy_pct"),
        DM_METRICS["avg_defense"]: overall.get("average_months_to_defense"),
        DM_METRICS["on_time"]: pct("on_time_graduation_pct"),
        DM_METRICS["risk30"]: int(overall.get("active_over_30_no_defense") or 0),
    }


def audit_dm_parity(
    cohorts: Iterable[dict[str, Any]],
    students: Iterable[dict[str, Any]],
    *,
    targets: Iterable[dict[str, Any]] | None = None,
    actions: Iterable[dict[str, Any]] | None = None,
    sync_runs: Iterable[dict[str, Any]] | None = None,
    as_of: date | None = None,
    source_kind: str = "fixture",
    generated_by: str | None = None,
    build_workbooks: bool = True,
) -> DMParityReport:
    cohort_rows = [dict(row) for row in cohorts]
    student_rows = [dict(row) for row in students]
    target_rows = [dict(row) for row in (targets or ())]
    action_rows = [dict(row) for row in (actions or ())]
    sync_rows = [dict(row) for row in (sync_runs or ())]
    cutoff = as_of or date.today()

    payload = build_dm_excel_official_payload(
        cohort_rows,
        student_rows,
        targets=target_rows,
        actions=action_rows,
        sync_runs=sync_rows,
        as_of=cutoff,
    )
    context = build_dm_snapshot_context(payload, generated_by=generated_by)

    scopes: list[tuple[str, str | None, int | None, str | None]] = [("DM · todas", None, None, None)]
    area_codes = []
    for cohort in cohort_rows:
        area = str(cohort.get("area_code") or "").strip().upper()
        if area and area not in area_codes:
            area_codes.append(area)
    for area in area_codes:
        scopes.append((f"Área {area}", area, None, None))
    for cohort in cohort_rows:
        cid = cohort.get("id")
        if cid in (None, ""):
            continue
        key = str(cohort.get("cohort_key") or "").strip() or f"{str(cohort.get('area_code') or '').strip().upper()} | Turma {int(cohort.get('cohort_number') or 0)}"
        scopes.append((key, str(cohort.get("area_code") or "").strip().upper() or None, int(cid), key))

    cases: list[DMParityCase] = []
    tolerance_by_metric = {
        DM_METRICS["members"]: 0.0,
        DM_METRICS["active"]: 0.0,
        DM_METRICS["occupancy"]: 0.00011,
        DM_METRICS["avg_defense"]: 0.011,
        DM_METRICS["on_time"]: 0.00011,
        DM_METRICS["risk30"]: 0.0,
    }

    for label, area, cohort_id, cohort_key in scopes:
        backend = build_dm_dashboard(
            cohort_rows,
            student_rows,
            targets=target_rows,
            area_code=area if cohort_id is None else None,
            cohort_id=cohort_id,
            as_of=cutoff,
        )
        expected = _expected_metrics(backend["overall"])
        initial: dict[str, Any] = {"area": area or "(todas)", "cohort": cohort_key or "(todas)"}
        spec = DMAdapter().build_spec(AdapterInput(context, payload, initial))
        for metric_code, expected_value in expected.items():
            actual = evaluate_metric(spec, metric_code)
            tolerance = tolerance_by_metric[metric_code]
            passed, delta = _same(expected_value, actual, tolerance)
            cases.append(DMParityCase(label, metric_code, expected_value, actual, tolerance, "PASS" if passed else "FAIL", delta))

    quality_notes: list[str] = []
    if not cohort_rows:
        quality_notes.append("Base de turmas vazia.")
    if not student_rows:
        quality_notes.append("Base de alunos vazia.")
    source_quality_ok = not quality_notes

    legacy_evidence: WorkbookEvidence | None = None
    new_evidence: WorkbookEvidence | None = None
    new_release_allowed = False
    if build_workbooks:
        try:
            legacy = build_dm_v2_workbook(
                cohort_rows,
                student_rows,
                targets=target_rows,
                sync_runs=sync_rows,
                as_of=cutoff,
            )
            legacy_evidence = _evidence("dm_v2", legacy)
        except Exception as exc:  # pragma: no cover
            legacy_evidence = WorkbookEvidence("dm_v2", False, error=f"{type(exc).__name__}: {exc}")
        try:
            artifact = build_dm_excel_official_artifact_from_payload(payload, generated_by=generated_by)
            new_release_allowed = artifact.audit.release_allowed
            new_evidence = _evidence("excel_official", artifact.to_bytes())
        except Exception as exc:  # pragma: no cover
            new_evidence = WorkbookEvidence("excel_official", False, error=f"{type(exc).__name__}: {exc}")

    return DMParityReport(
        source_kind=str(source_kind or "fixture").strip().lower(),
        generated_at=datetime.now(timezone.utc).isoformat(),
        system_version=APP_VERSION,
        schema_version=str(SCHEMA_VERSION),
        payload_hash=context.payload_hash,
        cases=tuple(cases),
        legacy_workbook=legacy_evidence,
        new_workbook=new_evidence,
        new_release_allowed=new_release_allowed,
        source_quality_ok=source_quality_ok,
        source_quality_notes=tuple(quality_notes),
    )


__all__ = ["DMParityCase", "DMParityReport", "PARITY_REPORT_VERSION", "WorkbookEvidence", "audit_dm_parity"]
