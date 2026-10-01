from pathlib import Path

import pytest
from openpyxl import Workbook

from dm_sei_parser import DMSEIParseError, parse_dm_sei_workbook


def _add_block(ws, row, *, course, label, students, declared=None):
    ws.cell(row, 1).value = "Curso:"
    ws.cell(row, 5).value = course
    ws.cell(row + 1, 1).value = "Turma:"
    ws.cell(row + 1, 5).value = label
    ws.cell(row + 2, 1).value = "Número"
    ws.cell(row + 2, 2).value = "Matrícula"
    ws.cell(row + 2, 7).value = "Nome do Aluno"
    ws.cell(row + 2, 16).value = "Situação"
    current = row + 3
    for ordinal, code, name, status in students:
        ws.cell(current, 1).value = ordinal
        ws.cell(current, 2).value = code
        ws.cell(current, 7).value = name
        ws.cell(current, 16).value = status
        current += 1
    ws.cell(current, 1).value = "Total Matriculados Turma:"
    ws.cell(current, 11).value = len(students) if declared is None else declared
    return current + 2


def _workbook(tmp_path: Path, blocks):
    path = tmp_path / "sei.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "AlunosPorUnidadeCursoTurmaSinte"
    row = 1
    for block in blocks:
        row = _add_block(ws, row, **block)
    wb.save(path)
    return path


def test_split_same_cohort_is_merged_and_students_are_preserved(tmp_path):
    path = _workbook(
        tmp_path,
        [
            {
                "course": "Ciência, Tecnologia e Educação",
                "label": "17-CTE",
                "students": [(1, "100001", "Aluno Um", "Ativa"), (2, "100002", "Aluno Dois", "Ativa")],
            },
            {
                "course": "Ciência, Tecnologia e Educação",
                "label": "17-CTE Mestrado Univc",
                "students": [(3, "100003", "Aluno Três", "Ativa")],
            },
        ],
    )
    report = parse_dm_sei_workbook(path)
    assert len(report.cohorts) == 1
    cohort = report.cohorts[0]
    assert cohort.key == "CTE:17"
    assert cohort.student_count == 3
    assert cohort.declared_total == 3
    assert len(report.students) == 3
    assert {row.student_code for row in report.students} == {"100001", "100002", "100003"}
    assert any("consolid" in warning.casefold() for warning in report.warnings)


def test_identical_student_repeated_inside_split_cohort_is_deduplicated(tmp_path):
    path = _workbook(
        tmp_path,
        [
            {
                "course": "Ciência, Tecnologia e Educação",
                "label": "17-CTE",
                "students": [(1, "100001", "Aluno Um", "Ativa")],
            },
            {
                "course": "Ciência, Tecnologia e Educação",
                "label": "17-CTE Mestrado Univc",
                "students": [(2, "100001", "Aluno Um", "Ativo")],
            },
        ],
    )
    report = parse_dm_sei_workbook(path)
    assert len(report.cohorts) == 1
    assert len(report.students) == 1
    assert report.cohorts[0].student_count == 1
    assert any("repetida" in warning.casefold() for warning in report.warnings)


def test_conflicting_student_in_split_same_cohort_still_blocks_import(tmp_path):
    path = _workbook(
        tmp_path,
        [
            {
                "course": "Ciência, Tecnologia e Educação",
                "label": "17-CTE",
                "students": [(1, "100001", "Aluno Um", "Ativa")],
            },
            {
                "course": "Ciência, Tecnologia e Educação",
                "label": "17-CTE Mestrado Univc",
                "students": [(2, "100001", "Outra Pessoa", "Ativa")],
            },
        ],
    )
    with pytest.raises(DMSEIParseError, match="dados conflitantes"):
        parse_dm_sei_workbook(path)


def test_same_student_in_different_cohorts_still_blocks_import(tmp_path):
    path = _workbook(
        tmp_path,
        [
            {
                "course": "Ciência, Tecnologia e Educação",
                "label": "17-CTE",
                "students": [(1, "100001", "Aluno Um", "Ativa")],
            },
            {
                "course": "Ciência, Tecnologia e Educação",
                "label": "18-CTE",
                "students": [(2, "100001", "Aluno Um", "Ativa")],
            },
        ],
    )
    with pytest.raises(DMSEIParseError, match="repete uma ou mais matrículas"):
        parse_dm_sei_workbook(path)


def test_test_and_real_blocks_with_same_key_are_not_silently_merged(tmp_path):
    path = _workbook(
        tmp_path,
        [
            {
                "course": "Ciência, Tecnologia e Educação",
                "label": "17-CTE_TURMA-TESTE",
                "students": [(1, "100001", "Aluno Teste", "Ativa")],
            },
            {
                "course": "Ciência, Tecnologia e Educação",
                "label": "17-CTE",
                "students": [(2, "100002", "Aluno Real", "Ativa")],
            },
        ],
    )
    with pytest.raises(DMSEIParseError, match="bloco real e um bloco de teste"):
        parse_dm_sei_workbook(path)
