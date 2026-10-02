from __future__ import annotations

from dataclasses import replace
from datetime import date, datetime, timezone
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from excel_official import (
    AdapterCapabilities,
    ColumnDataType,
    ColumnSpec,
    DashboardSpec,
    DatasetSpec,
    DimensionSpec,
    IdentitySpec,
    InitialStateSpec,
    ParameterSpec,
    ParameterWriteError,
    SnapshotSpec,
    WorkbookSpec,
    validate_workbook_spec,
    write_parameter_system,
)
from excel_official.components.names import parameter_defined_name, support_list_defined_name


def _spec(*, course_initial="Psicologia", discipline_initial="(todas)") -> WorkbookSpec:
    periods = DatasetSpec(
        code="periods",
        label="Dimensao de periodos",
        sheet_name="DIM_PERIODO",
        table_name="TblDimPeriodo",
        columns=(
            ColumnSpec("period", "Periodo", ColumnDataType.TEXT),
            ColumnSpec("order", "Ordem", ColumnDataType.INTEGER, technical=True),
        ),
        rows=(
            {"period": "2026-SEM2", "order": 2},
            {"period": "2026-SEM1", "order": 1},
        ),
        technical=True,
    )
    courses = DatasetSpec(
        code="courses",
        label="Cursos",
        sheet_name="CURSOS",
        table_name="TblCursos",
        columns=(
            ColumnSpec("course_key", "Codigo", ColumnDataType.TEXT, technical=True),
            ColumnSpec("course", "Curso", ColumnDataType.TEXT),
            ColumnSpec("order", "Ordem", ColumnDataType.INTEGER, technical=True),
        ),
        rows=(
            {"course_key": "ODO", "course": "Odontologia", "order": 2},
            {"course_key": "PSI", "course": "Psicologia", "order": 1},
        ),
    )
    disciplines = DatasetSpec(
        code="disciplines",
        label="Disciplinas",
        sheet_name="DISCIPLINAS",
        table_name="TblDisciplinas",
        columns=(
            ColumnSpec("discipline_key", "Codigo", ColumnDataType.TEXT, technical=True),
            ColumnSpec("discipline", "Disciplina", ColumnDataType.TEXT),
            ColumnSpec("course_key", "Codigo curso", ColumnDataType.TEXT, technical=True),
            ColumnSpec("order", "Ordem", ColumnDataType.INTEGER, technical=True),
        ),
        rows=(
            {"discipline_key": "ANA", "discipline": "Anatomia", "course_key": "ODO", "order": 1},
            {"discipline_key": "CLI", "discipline": "Clinica", "course_key": "ODO", "order": 2},
            {"discipline_key": "PSO", "discipline": "Psicologia Social", "course_key": "PSI", "order": 2},
            {"discipline_key": "NEU", "discipline": "Neurociencia", "course_key": "PSI", "order": 1},
        ),
    )
    windows = DatasetSpec(
        code="windows",
        label="Janelas",
        sheet_name="JANELAS",
        table_name="TblJanelas",
        columns=(ColumnSpec("window", "Janela", ColumnDataType.INTEGER),),
        rows=({"window": 4}, {"window": 6}, {"window": 8}, {"window": 12}),
    )
    return WorkbookSpec(
        identity=IdentitySpec(
            workbook_title="Painel DCS",
            directorate_code="DCS",
            directorate_label="Diretoria de Ciencias da Saude",
            adapter_code="academic",
            adapter_version=1,
        ),
        snapshot=SnapshotSpec(
            export_id="01d-test",
            generated_at=datetime(2026, 10, 1, 22, 0, tzinfo=timezone.utc),
            generated_by="test-user",
            authorization_scope=("DCS",),
            system_version="0.13.0",
            schema_version=49,
            adapter_version=1,
            initial_scope={"course": "Odontologia"},
        ),
        initial_state=InitialStateSpec(
            values={
                "reference_period": "2026-SEM2",
                "course": course_initial,
                "discipline": discipline_initial,
                "as_of": date(2026, 10, 1),
            }
        ),
        parameters=(
            ParameterSpec(
                "reference_period",
                "Periodo de referencia",
                "period",
                required=True,
                description="Periodo principal analisado.",
                display_order=10,
            ),
            ParameterSpec(
                "course",
                "Curso",
                "course",
                empty_option="(todos)",
                description="Curso em foco.",
                display_order=20,
            ),
            ParameterSpec(
                "discipline",
                "Disciplina",
                "discipline",
                depends_on=("course",),
                empty_option="(todas)",
                description="Disciplina em foco.",
                display_order=30,
            ),
            ParameterSpec(
                "window",
                "Janela historica",
                "windows",
                initial_value=4,
                values_column="window",
                display_order=40,
            ),
            ParameterSpec(
                "as_of",
                "Data de corte",
                "",
                initial_value=date(2026, 9, 30),
                data_type=ColumnDataType.DATE,
                description="Controle livre sem dropdown.",
                display_order=50,
            ),
        ),
        datasets=(periods, courses, disciplines, windows),
        dimensions=(
            DimensionSpec("period", "Periodo", "periods", "period", "period", sort_order_column="order"),
            DimensionSpec("course", "Curso", "courses", "course_key", "course", sort_order_column="order"),
            DimensionSpec(
                "discipline",
                "Disciplina",
                "disciplines",
                "discipline_key",
                "discipline",
                sort_order_column="order",
                parent_dimension="course",
                parent_key_column="course_key",
            ),
        ),
        metric_codes=(),
        metric_bindings=(),
        domain_sheets=(),
        dashboard=DashboardSpec(title="DCS"),
        matrix=None,
        capabilities=AdapterCapabilities(interactive_dimensions=("period", "course", "discipline")),
    )


def _errors(spec: WorkbookSpec) -> set[str]:
    return {item.code for item in validate_workbook_spec(spec) if item.severity == "ERROR"}


def _blank_workbook() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def test_01d_contract_accepts_child_dimension_and_free_date_parameter():
    assert validate_workbook_spec(_spec()) == ()


def test_01d_contract_rejects_missing_child_parent_key_column():
    spec = _spec()
    bad = replace(spec.dimensions[2], parent_key_column="missing")
    spec = replace(spec, dimensions=spec.dimensions[:2] + (bad,))
    assert "dimension.unknown_column" in _errors(spec)
    assert "dimension.missing_parent_key_column" in _errors(spec)


def test_01d_contract_rejects_dependency_mismatch():
    spec = _spec()
    bad = replace(spec.parameters[2], depends_on=("reference_period",))
    spec = replace(spec, parameters=spec.parameters[:2] + (bad,) + spec.parameters[3:])
    assert "parameter.dependency_dimension_mismatch" in _errors(spec)


def test_01d_contract_rejects_more_than_one_dependency_in_v1():
    spec = _spec()
    bad = replace(spec.parameters[2], depends_on=("course", "reference_period"))
    spec = replace(spec, parameters=spec.parameters[:2] + (bad,) + spec.parameters[3:])
    assert "parameter.multiple_dependencies_not_supported_v1" in _errors(spec)


def test_01d_parameter_system_creates_visible_protected_sheets_and_refs():
    wb = _blank_workbook()
    refs = write_parameter_system(wb, _spec())
    assert wb.sheetnames == ["LISTAS DE APOIO", "PARAMETROS"]
    assert wb["LISTAS DE APOIO"].sheet_state == "visible"
    assert wb["PARAMETROS"].sheet_state == "visible"
    assert wb["LISTAS DE APOIO"].protection.sheet is True
    assert wb["PARAMETROS"].protection.sheet is True
    assert refs.parameters["course"].defined_name == "P_COURSE"
    assert refs.support_lists["course"].defined_name == "LST_COURSE"
    assert refs.support_lists["discipline"].dynamic is True


def test_01d_initial_state_overrides_parameter_and_snapshot_seed():
    wb = _blank_workbook()
    refs = write_parameter_system(wb, _spec())
    course_cell = refs.parameters["course"].cell_reference
    as_of_cell = refs.parameters["as_of"].cell_reference
    assert wb["PARAMETROS"][course_cell].value == "Psicologia"
    as_of_value = wb["PARAMETROS"][as_of_cell].value
    if isinstance(as_of_value, datetime):
        as_of_value = as_of_value.date()
    assert as_of_value == date(2026, 10, 1)


def test_01d_independent_list_is_sorted_and_empty_option_is_first():
    wb = _blank_workbook()
    refs = write_parameter_system(wb, _spec())
    ref = refs.support_lists["course"]
    ws = wb[ref.sheet_name]
    col = ref.value_range.split("!")[1].split("$")[1]
    values = [ws[f"{col}{row}"].value for row in range(5, 8)]
    assert values == ["(todos)", "Psicologia", "Odontologia"]


def test_01d_dependent_list_uses_nonvolatile_index_match_countif():
    wb = _blank_workbook()
    refs = write_parameter_system(wb, _spec())
    defined = wb.defined_names[refs.support_lists["discipline"].defined_name]
    formula = defined.attr_text.upper()
    assert "INDEX(" in formula
    assert "MATCH(" in formula
    assert "COUNTIF(" in formula
    assert "P_COURSE" in formula
    assert "INDIRECT(" not in formula
    assert "OFFSET(" not in formula


def test_01d_all_courses_parent_option_exposes_all_disciplines():
    wb = _blank_workbook()
    refs = write_parameter_system(wb, _spec(course_initial="(todos)"))
    ref = refs.support_lists["discipline"]
    ws = wb[ref.sheet_name]
    parent_col = ref.parent_range.split("!")[1].split("$")[1]
    value_col = ref.value_range.split("!")[1].split("$")[1]
    all_values = []
    for row in range(5, 30):
        if ws[f"{parent_col}{row}"].value == "(todos)":
            all_values.append(ws[f"{value_col}{row}"].value)
    assert all_values == ["(todas)", "Anatomia", "Clinica", "Neurociencia", "Psicologia Social"]


def test_01d_dependent_initial_value_must_belong_to_selected_parent():
    wb = _blank_workbook()
    with pytest.raises(ParameterWriteError) as exc:
        write_parameter_system(wb, _spec(course_initial="Psicologia", discipline_initial="Anatomia"))
    assert exc.value.code == "parameter.initial_value_out_of_scope"


def test_01d_input_cells_are_unlocked_and_free_date_has_no_list_validation():
    wb = _blank_workbook()
    refs = write_parameter_system(wb, _spec())
    ws = wb["PARAMETROS"]
    assert ws[refs.parameters["course"].cell_reference].protection.locked is False
    assert ws[refs.parameters["as_of"].cell_reference].protection.locked is False
    formulas = {dv.formula1: str(dv.sqref) for dv in ws.data_validations.dataValidation}
    assert "=LST_COURSE" in formulas
    assert "=LST_DISCIPLINE" in formulas
    assert refs.parameters["as_of"].cell_reference not in " ".join(formulas.values())
    assert ws[refs.parameters["as_of"].cell_reference].number_format == "dd/mm/yyyy"


def test_01d_direct_dataset_list_requires_explicit_column_when_ambiguous():
    spec = _spec()
    courses = spec.datasets[1]
    ambiguous_courses = replace(
        courses,
        columns=(
            replace(courses.columns[0], technical=False),
            courses.columns[1],
            replace(courses.columns[2], technical=False),
        ),
    )
    bad = replace(spec.parameters[3], values_source="courses", values_column=None)
    spec = replace(
        spec,
        parameters=spec.parameters[:3] + (bad,) + spec.parameters[4:],
        datasets=(spec.datasets[0], ambiguous_courses) + spec.datasets[2:],
    )
    wb = _blank_workbook()
    with pytest.raises(ParameterWriteError) as exc:
        write_parameter_system(wb, spec)
    assert exc.value.code == "parameter.dataset_values_column_required"


def test_01d_names_are_stable_excel_safe_and_accent_insensitive():
    assert parameter_defined_name("período de referência") == "P_PERIODO_DE_REFERENCIA"
    assert support_list_defined_name("curso/foco") == "LST_CURSO_FOCO"


def test_01d_workbook_roundtrip_preserves_names_and_validations(tmp_path: Path):
    wb = _blank_workbook()
    refs = write_parameter_system(wb, _spec())
    path = tmp_path / "parameters.xlsx"
    wb.save(path)
    reopened = load_workbook(path)
    assert "P_COURSE" in reopened.defined_names
    assert "LST_COURSE" in reopened.defined_names
    assert "LST_DISCIPLINE" in reopened.defined_names
    dvs = reopened["PARAMETROS"].data_validations.dataValidation
    assert {dv.formula1 for dv in dvs} >= {"=LST_COURSE", "=LST_DISCIPLINE", "=LST_REFERENCE_PERIOD"}
    assert reopened["LISTAS DE APOIO"].sheet_state == "visible"
    assert refs.parameters["course"].cell_reference == "B7"


def test_01d_architecture_keeps_contract_and_adapters_openpyxl_independent():
    root = Path(__file__).resolve().parents[1] / "excel_official"
    architecture_files = [
        root / "contract.py",
        root / "context.py",
        root / "validation.py",
        root / "theme.py",
        root / "adapters" / "base.py",
    ]
    source = "\n".join(path.read_text(encoding="utf-8") for path in architecture_files)
    assert "import openpyxl" not in source
    assert "from openpyxl" not in source


def test_01d_required_parameter_can_use_snapshot_initial_scope_as_fallback():
    spec = _spec()
    period = replace(spec.parameters[0], initial_value=None)
    state = InitialStateSpec(values={key: value for key, value in spec.initial_state.values.items() if key != "reference_period"})
    snapshot = replace(spec.snapshot, initial_scope={**spec.snapshot.initial_scope, "reference_period": "2026-SEM1"})
    spec = replace(spec, parameters=(period,) + spec.parameters[1:], initial_state=state, snapshot=snapshot)
    assert "parameter.required_without_initial_value" not in _errors(spec)
    wb = _blank_workbook()
    refs = write_parameter_system(wb, spec)
    assert wb["PARAMETROS"][refs.parameters["reference_period"].cell_reference].value == "2026-SEM1"


def test_01d_preflight_does_not_mutate_workbook_on_invalid_initial_value():
    wb = _blank_workbook()
    wb.create_sheet("EXISTENTE")
    before = tuple(wb.sheetnames)
    with pytest.raises(ParameterWriteError):
        write_parameter_system(wb, _spec(course_initial="Psicologia", discipline_initial="Anatomia"))
    assert tuple(wb.sheetnames) == before
    assert list(wb.defined_names.keys()) == []


def test_01d_preflight_rejects_normalized_defined_name_collision_without_mutation():
    spec = _spec()
    p1 = ParameterSpec("local-filter", "Filtro local 1", "", initial_value="A", display_order=90)
    p2 = ParameterSpec("local_filter", "Filtro local 2", "", initial_value="B", display_order=91)
    spec = replace(spec, parameters=spec.parameters + (p1, p2))
    wb = _blank_workbook()
    with pytest.raises(ParameterWriteError) as exc:
        write_parameter_system(wb, spec)
    assert exc.value.code == "parameter.generated_name_collision"
    assert wb.sheetnames == []


def test_01d_parameter_declared_date_rejects_datetime_seed_before_mutation():
    spec = _spec()
    state = InitialStateSpec(values={**spec.initial_state.values, "as_of": datetime(2026, 10, 1, 10, 0)})
    spec = replace(spec, initial_state=state)
    wb = _blank_workbook()
    with pytest.raises(ParameterWriteError) as exc:
        write_parameter_system(wb, spec)
    assert exc.value.code == "parameter.invalid_value_type"
    assert wb.sheetnames == []
