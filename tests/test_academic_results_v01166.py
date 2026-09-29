from __future__ import annotations

import tempfile
import unittest
import uuid
import zipfile
from pathlib import Path

from openpyxl import Workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import sessionmaker

from academic_excel_parser import inspecionar_relatorio, iterar_registros
from database import Base
from models import AcademicResult, Directorate
from repository import DatabaseRepository
from security import AuthorizationContext


class AcademicResultsV01166Tests(unittest.TestCase):
    @staticmethod
    def _make_sei_report_with_underreported_dimension(path: Path) -> None:
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Ano/Semestre:"
        ws["C1"] = "2026/1"
        ws["A3"] = "Unidade Ensino:"
        ws["A4"] = "Curso:"
        ws["C4"] = "Administração"
        ws["A5"] = "Disciplina:"
        ws["C5"] = "Gestão"
        ws["A6"] = "Turma:"
        ws["C6"] = "ADM1"
        ws["N6"] = "1º"
        ws["A7"] = "Matrícula"
        ws["C7"] = "Nome"
        ws["K7"] = "Média"
        ws["R7"] = "Situação"
        ws["A8"] = "000001"
        ws["C8"] = "Aluno Um"
        ws["K8"] = 8.5
        ws["R8"] = "Aprovado"
        ws["A9"] = "000002"
        ws["C9"] = "Aluno Dois"
        ws["K9"] = 4.0
        ws["R9"] = "Reprovado"
        ws["A10"] = "Qtd de alunos:"
        ws["C10"] = 2
        wb.save(path)
        wb.close()

        # Simula XLSX de gerador externo/SEI que declara uma dimensão menor que
        # a área real da planilha. O modo normal do openpyxl tolera isso, mas o
        # read-only pode truncar a iteração se a dimensão não for resetada.
        with tempfile.TemporaryDirectory() as extracted_dir:
            extracted = Path(extracted_dir)
            with zipfile.ZipFile(path, "r") as archive:
                archive.extractall(extracted)
            sheet = extracted / "xl" / "worksheets" / "sheet1.xml"
            xml = sheet.read_text(encoding="utf-8")
            import re
            xml, count = re.subn(r'<dimension ref="[^"]+"\s*/>', '<dimension ref="A1:A1"/>', xml, count=1)
            if count != 1:
                raise AssertionError("Não foi possível adulterar a dimensão do fixture XLSX.")
            sheet.write_text(xml, encoding="utf-8")
            rebuilt = path.with_suffix(".rebuilt.xlsx")
            with zipfile.ZipFile(rebuilt, "w", zipfile.ZIP_DEFLATED) as archive:
                for item in extracted.rglob("*"):
                    if item.is_file():
                        archive.write(item, item.relative_to(extracted))
            rebuilt.replace(path)

    def test_streaming_parser_ignores_underreported_xlsx_dimension(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "sei_dimension.xlsx"
            self._make_sei_report_with_underreported_dimension(path)
            metadata, warnings = inspecionar_relatorio(path)
            self.assertEqual(metadata["curso"], "Administração")
            self.assertEqual(metadata["ano"], 2026)
            self.assertEqual(metadata["semestre"], 1)
            self.assertEqual(metadata["total_registros_aluno_disciplina"], 2)
            self.assertEqual(sum(1 for _ in iterar_registros(path)), 2)
            self.assertEqual(warnings, [])

    def test_repository_persists_realistic_sei_report_even_with_bad_dimension(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "sei_dimension.xlsx"
            self._make_sei_report_with_underreported_dimension(path)
            engine = create_engine("sqlite:///:memory:")
            Base.metadata.create_all(engine)
            Session = sessionmaker(bind=engine, expire_on_commit=False)
            db = Session()
            directorate = Directorate(code="DTNH", name="DTNH", active=True)
            db.add(directorate); db.flush(); db.commit()
            user = AuthorizationContext(
                user_id=str(uuid.uuid4()), email="dimension@univc.edu.br", full_name="Dimension Test",
                role="editor", directorate_id=directorate.id, directorate_code="DTNH", directorate_name="DTNH",
            )
            result = DatabaseRepository(db, user).import_sei_report_file(path, expected_course="Administração")
            self.assertEqual(result["registros_lidos"], 2)
            self.assertEqual(result["inseridos"], 2)
            self.assertEqual(db.scalar(select(func.count(AcademicResult.id))), 2)
            db.close(); engine.dispose()

    def test_sei_course_import_uses_inline_progress_without_blocking_overlay(self):
        root = Path(__file__).resolve().parents[1]
        js = (root / "static" / "js" / "app.js").read_text(encoding="utf-8")
        start = js.index("async function submitSeiImport(event)")
        end = js.index("function surveyFilterSelect", start)
        body = js[start:end]
        self.assertIn("window.DataUNIVC?.progress?.render", body)
        self.assertIn("blocking:false", body)
        self.assertIn("O progresso detalhado permanece visível no próprio formulário.", body)


if __name__ == "__main__":
    unittest.main()
