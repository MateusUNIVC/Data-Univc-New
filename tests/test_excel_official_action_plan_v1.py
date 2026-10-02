from __future__ import annotations

from dataclasses import replace
from datetime import date
from pathlib import Path
import sys

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from test_excel_official_dashboard_v1 import _build, _spec  # noqa: E402

from excel_official import (  # noqa: E402
    ACTION_PLAN_SHEET,
    ActionPlanRowSource,
    ActionPlanRowSpec,
    ActionPlanSpec,
    ActionPlanWriteError,
    validate_workbook_spec,
    write_action_plan_sheet,
)


def _action_spec(**kwargs):
    base = _spec()
    plan = ActionPlanSpec(
        official_fields=(),
        local_editable_fields=(
            "indicator",
            "problem",
            "diagnosis",
            "action",
            "owner",
            "deadline",
            "status",
        ),
        rows=(
            ActionPlanRowSpec(
                source=ActionPlanRowSource.OFFICIAL,
                values={
                    "indicator": "academic.nps",
                    "problem": "NPS abaixo da meta",
                    "diagnosis": "Baixa participacao em dois cursos",
                    "action": "Revisar plano de escuta discente",
                    "owner": "Coordenacao",
                    "deadline": date(2026, 10, 20),
                    "status": "Em andamento",
                },
            ),
            ActionPlanRowSpec(
                source=ActionPlanRowSource.LOCAL,
                values={
                    "indicator": "Aprovacao",
                    "status": "N\u00e3o iniciado",
                },
            ),
        ),
        local_blank_rows=2,
    )
    return replace(base, action_plan=replace(plan, **kwargs))


def _errors(spec) -> set[str]:
    return {item.code for item in validate_workbook_spec(spec) if item.severity == "ERROR"}


def test_01i_action_plan_contract_is_valid():
    assert validate_workbook_spec(_action_spec()) == ()


def test_01i_action_plan_materializes_official_and_local_rows():
    spec = _action_spec()
    wb, _, _, _ = _build(spec)
    ref = write_action_plan_sheet(wb, spec)
    ws = wb[ACTION_PLAN_SHEET]

    assert ref is not None
    assert ws.sheet_state == "visible"
    assert ws.protection.sheet is True
    assert ws["A1"].value == "PLANO DE A\u00c7\u00c3O"
    assert "n\u00e3o s\u00e3o sincronizadas" in ws[ref.notice_cell].value
    assert ws["A6"].value == "Origem"
    assert ws["B6"].value == "Indicador"
    assert ws["G6"].value == "Prazo"
    assert ws["H6"].value == "Status"
    assert ws["I6"].value == "Dias restantes"
    assert ref.official_row_count == 1
    assert ref.local_row_count == 3
    assert ws["A7"].value == "OFICIAL"
    assert ws["A8"].value == "LOCAL"
    assert ref.table_range == "A6:I10"


def test_01i_official_rows_are_locked_and_local_fields_are_explicitly_editable():
    spec = _action_spec()
    wb, _, _, _ = _build(spec)
    ref = write_action_plan_sheet(wb, spec)
    ws = wb[ref.sheet_name]

    for col in range(1, 10):
        assert ws.cell(7, col).protection.locked is True
    assert ws["A8"].protection.locked is True
    for col in range(2, 9):
        assert ws.cell(8, col).protection.locked is False
    assert ws["I8"].protection.locked is True
    assert ws["B8"].fill.fgColor.rgb.endswith("FFF6DC")


def test_01i_status_validation_uses_auditable_support_list_and_named_range():
    spec = _action_spec()
    wb, _, _, _ = _build(spec)
    ref = write_action_plan_sheet(wb, spec)
    ws = wb[ref.sheet_name]

    assert ref.status_validation_range == "H8:H10"
    assert "LST_ACTION_STATUS" in wb.defined_names
    assert "LST_ACTION_INDICATOR" in wb.defined_names
    assert ref.indicator_validation_range == "B8:B10"
    validations = list(ws.data_validations.dataValidation)
    assert len(validations) == 2
    by_formula = {item.formula1: str(item.sqref) for item in validations}
    assert by_formula["=LST_ACTION_STATUS"] == "H8:H10"
    assert by_formula["=LST_ACTION_INDICATOR"] == "B8:B10"

    support = wb["LISTAS DE APOIO"]
    values = [cell.value for row in support.iter_rows() for cell in row if cell.value is not None]
    assert "N\u00e3o iniciado" in values
    assert "Conclu\u00eddo" in values


def test_01i_days_remaining_is_derived_locked_and_uses_completed_status_semantics():
    spec = _action_spec()
    wb, _, _, _ = _build(spec)
    ref = write_action_plan_sheet(wb, spec)
    ws = wb[ref.sheet_name]

    formula = str(ws["I7"].value)
    assert formula.startswith("=IF(OR(G7=\"\"")
    assert "TODAY()" in formula
    assert "Conclu\u00eddo" in formula
    assert ws["I7"].data_type == "f"
    assert ws["I7"].protection.locked is True
    assert ws["I8"].protection.locked is True
    assert ws["I8"].number_format == "#,##0"


def test_01i_text_from_official_source_is_literal_not_formula():
    spec = _action_spec(
        rows=(
            ActionPlanRowSpec(
                source=ActionPlanRowSource.OFFICIAL,
                values={"indicator": "=1+1", "status": "Em andamento"},
            ),
        ),
        local_blank_rows=1,
    )
    wb, _, _, _ = _build(spec)
    write_action_plan_sheet(wb, spec)
    cell = wb[ACTION_PLAN_SHEET]["B7"]
    assert cell.value == "=1+1"
    assert cell.data_type == "s"


def test_01i_disabled_action_plan_is_noop():
    spec = _action_spec(enabled=False)
    wb, _, _, _ = _build(spec)
    before = tuple(wb.sheetnames)
    assert write_action_plan_sheet(wb, spec) is None
    assert tuple(wb.sheetnames) == before
    assert ACTION_PLAN_SHEET not in wb.sheetnames


def test_01i_preflight_collision_does_not_mutate_support_lists():
    spec = _action_spec()
    wb, _, _, _ = _build(spec)
    wb.create_sheet(ACTION_PLAN_SHEET)
    before_names = tuple(wb.defined_names.keys())
    support_before = wb["LISTAS DE APOIO"].max_column

    with pytest.raises(ActionPlanWriteError) as exc:
        write_action_plan_sheet(wb, spec)
    assert exc.value.code == "action_plan.sheet_collision"
    assert tuple(wb.defined_names.keys()) == before_names
    assert wb["LISTAS DE APOIO"].max_column == support_before


def test_01i_contract_rejects_invalid_action_plan_semantics():
    base = _action_spec()

    editable_days = replace(
        base,
        action_plan=replace(
            base.action_plan,
            local_editable_fields=base.action_plan.local_editable_fields + ("days_remaining",),
        ),
    )
    assert "action_plan.derived_field_editable" in _errors(editable_days)

    bad_status = replace(
        base,
        action_plan=replace(
            base.action_plan,
            rows=(ActionPlanRowSpec(values={"status": "Inventado"}),),
        ),
    )
    assert "action_plan.invalid_status_value" in _errors(bad_status)

    unknown_completed = replace(
        base,
        action_plan=replace(base.action_plan, completed_statuses=("Finalizado",)),
    )
    assert "action_plan.unknown_completed_status" in _errors(unknown_completed)

    unknown_field = replace(
        base,
        action_plan=replace(
            base.action_plan,
            rows=(ActionPlanRowSpec(values={"unknown": "x"}),),
        ),
    )
    assert "action_plan.row_unknown_field" in _errors(unknown_field)


def test_01i_roundtrip_preserves_table_validation_formulas_and_protection(tmp_path: Path):
    spec = _action_spec()
    wb, _, _, _ = _build(spec)
    ref = write_action_plan_sheet(wb, spec)
    path = tmp_path / "action_plan.xlsx"
    wb.save(path)

    reopened = load_workbook(path, data_only=False)
    ws = reopened[ref.sheet_name]
    assert ws.protection.sheet is True
    assert ws.freeze_panes == "B7"
    assert "TblPlanoAcaoOfficial" in ws.tables
    assert ws["I7"].data_type == "f"
    assert ws["B7"].protection.locked is True
    assert ws["B8"].protection.locked is False
    assert len(ws.data_validations.dataValidation) == 2
    assert reopened.calculation.calcMode == "auto"
    assert reopened.calculation.fullCalcOnLoad is True
    assert reopened.calculation.forceFullCalc is True


def test_01i_formulas_remain_excel_2019_compatible():
    spec = _action_spec()
    wb, _, _, _ = _build(spec)
    write_action_plan_sheet(wb, spec)
    formulas = "\n".join(
        str(cell.value).upper()
        for row in wb[ACTION_PLAN_SHEET].iter_rows()
        for cell in row
        if cell.data_type == "f"
    )
    assert "TODAY()" in formulas
    for unsupported in ("FILTER(", "XLOOKUP(", "SORT(", "UNIQUE(", "INDIRECT(", "OFFSET("):
        assert unsupported not in formulas
