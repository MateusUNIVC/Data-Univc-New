from __future__ import annotations

from dataclasses import replace
from datetime import datetime, timezone
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from excel_official import (
    AdapterCapabilities,
    ChartRole,
    ChartSpec,
    ChartType,
    ColumnDataType,
    ColumnSpec,
    DashboardSpec,
    DashboardWriteError,
    DatasetSpec,
    DimensionSpec,
    IdentitySpec,
    InitialStateSpec,
    KpiSpec,
    MetricAggregation,
    MetricBinding,
    MetricSpec,
    MetricUnit,
    ParameterSpec,
    SnapshotSpec,
    WorkbookSpec,
    validate_workbook_spec,
    write_dashboard_sheet,
    write_dataset_sheet,
    write_parameter_system,
)


def _spec(*, include_chart: bool = False, empty_facts: bool = False) -> WorkbookSpec:
    periods = DatasetSpec(
        code="periods",
        label="Periodos",
        sheet_name="DIM_PERIODO",
        table_name="TblDimPeriodo01E",
        technical=True,
        columns=(
            ColumnSpec("period", "Periodo", ColumnDataType.TEXT, semantic_type="period"),
            ColumnSpec("order", "Ordem", ColumnDataType.INTEGER, technical=True),
        ),
        rows=(
            {"period": "2026-SEM1", "order": 1},
            {"period": "2026-SEM2", "order": 2},
        ),
        grain=("period",),
    )
    courses = DatasetSpec(
        code="courses",
        label="Cursos",
        sheet_name="CURSOS",
        table_name="TblCursos01E",
        columns=(ColumnSpec("course", "Curso", ColumnDataType.TEXT, semantic_type="course"),),
        rows=({"course": "Psicologia"}, {"course": "Odontologia"}),
        grain=("course",),
    )
    fact_rows = () if empty_facts else (
        {
            "period": "2026-SEM1",
            "course": "Psicologia",
            "respondents": 10,
            "promoters": 7,
            "detractors": 1,
            "approved": 8,
            "finalized": 10,
            "grade_sum": 82.0,
            "grade_count": 10,
            "active_students": 120,
        },
        {
            "period": "2026-SEM2",
            "course": "Psicologia",
            "respondents": 20,
            "promoters": 16,
            "detractors": 2,
            "approved": 18,
            "finalized": 20,
            "grade_sum": 174.0,
            "grade_count": 20,
            "active_students": 128,
        },
        {
            "period": "2026-SEM2",
            "course": "Odontologia",
            "respondents": 10,
            "promoters": 5,
            "detractors": 2,
            "approved": 7,
            "finalized": 10,
            "grade_sum": 75.0,
            "grade_count": 10,
            "active_students": 90,
        },
    )
    facts = DatasetSpec(
        code="academic_facts",
        label="Fatos Academicos",
        sheet_name="FATOS ACADEMICOS",
        table_name="TblAcademicFacts01E",
        columns=(
            ColumnSpec("period", "Periodo", ColumnDataType.TEXT, semantic_type="period"),
            ColumnSpec("course", "Curso", ColumnDataType.TEXT, semantic_type="course"),
            ColumnSpec("respondents", "Respondentes", ColumnDataType.INTEGER),
            ColumnSpec("promoters", "Promotores", ColumnDataType.INTEGER),
            ColumnSpec("detractors", "Detratores", ColumnDataType.INTEGER),
            ColumnSpec("approved", "Aprovados", ColumnDataType.INTEGER),
            ColumnSpec("finalized", "Finalizados", ColumnDataType.INTEGER),
            ColumnSpec("grade_sum", "Soma notas", ColumnDataType.DECIMAL),
            ColumnSpec("grade_count", "Qtd notas", ColumnDataType.INTEGER),
            ColumnSpec("active_students", "Alunos ativos", ColumnDataType.INTEGER),
        ),
        rows=fact_rows,
        grain=("period", "course"),
        filter_dimensions=("period", "course"),
        dimension_columns={"period": "period", "course": "course"},
    )
    charts = ()
    if include_chart:
        charts = (
            ChartSpec(
                code="nps_evolution",
                title="Evolucao do NPS",
                metric_code="academic.nps",
                dataset_code="academic_facts",
                dimension_code="period",
                chart_type=ChartType.LINE,
                role=ChartRole.EVOLUTION,
            ),
        )
    return WorkbookSpec(
        identity=IdentitySpec(
            workbook_title="Painel DCS",
            directorate_code="DCS",
            directorate_label="Diretoria de Ciencias da Saude",
            adapter_code="academic",
            adapter_version=1,
        ),
        snapshot=SnapshotSpec(
            export_id="01e-test",
            generated_at=datetime(2026, 10, 1, 22, 0, tzinfo=timezone.utc),
            generated_by="test-user",
            authorization_scope=("DCS",),
            system_version="0.13.0",
            schema_version=49,
            adapter_version=1,
        ),
        initial_state=InitialStateSpec(
            values={
                "reference_period": "2026-SEM2",
                "comparison_period": "2026-SEM1",
                "course": "Psicologia",
            }
        ),
        parameters=(
            ParameterSpec("reference_period", "Periodo de referencia", "period", required=True, display_order=10),
            ParameterSpec("comparison_period", "Comparar com", "period", required=True, display_order=20),
            ParameterSpec("course", "Curso", "course", empty_option="(todos)", display_order=30),
        ),
        datasets=(periods, courses, facts),
        dimensions=(
            DimensionSpec("period", "Periodo", "periods", "period", "period", sort_order_column="order"),
            DimensionSpec("course", "Curso", "courses", "course", "course"),
        ),
        metric_codes=("academic.nps", "academic.approval", "academic.average_grade", "academic.active_students"),
        metric_bindings=(
            MetricBinding(
                "academic.nps",
                "academic_facts",
                components={"respondents": "respondents", "promoters": "promoters", "detractors": "detractors"},
                filter_parameters={"period": "reference_period", "course": "course"},
            ),
            MetricBinding(
                "academic.approval",
                "academic_facts",
                components={"numerator": "approved", "denominator": "finalized"},
                filter_parameters={"period": "reference_period", "course": "course"},
            ),
            MetricBinding(
                "academic.average_grade",
                "academic_facts",
                components={"sum": "grade_sum", "count": "grade_count"},
                filter_parameters={"period": "reference_period", "course": "course"},
            ),
            MetricBinding(
                "academic.active_students",
                "academic_facts",
                value_column="active_students",
                filter_parameters={"period": "reference_period", "course": "course"},
            ),
        ),
        domain_sheets=(),
        dashboard=DashboardSpec(
            title="PAINEL DE GESTAO · DCS",
            kpis=(
                KpiSpec("academic.nps", comparison="comparison_period", target=70, priority=10),
                KpiSpec("academic.approval", comparison="comparison_period", target=0.85, priority=20),
                KpiSpec("academic.average_grade", comparison="comparison_period", target=7.0, priority=30),
                KpiSpec("academic.active_students", priority=40),
            ),
            charts=charts,
        ),
        matrix=None,
        capabilities=AdapterCapabilities(
            supports_comparison=True,
            interactive_dimensions=("period", "course"),
        ),
        metrics=(
            MetricSpec(
                "academic.nps",
                "NPS da Instituicao",
                MetricAggregation.NPS,
                MetricUnit.NPS,
                allowed_dimensions=("period", "course"),
            ),
            MetricSpec(
                "academic.approval",
                "Aprovacao",
                MetricAggregation.RATIO,
                MetricUnit.PERCENT,
                allowed_dimensions=("period", "course"),
            ),
            MetricSpec(
                "academic.average_grade",
                "Media das notas",
                MetricAggregation.AVERAGE,
                MetricUnit.SCORE_0_10,
                allowed_dimensions=("period", "course"),
            ),
            MetricSpec(
                "academic.active_students",
                "Alunos ativos",
                MetricAggregation.SUM,
                MetricUnit.COUNT,
                allowed_dimensions=("period", "course"),
                display_precision=0,
            ),
        ),
    )


def _build(spec: WorkbookSpec):
    wb = Workbook()
    wb.remove(wb.active)
    refs = {dataset.code: write_dataset_sheet(wb, dataset) for dataset in spec.datasets}
    parameters = write_parameter_system(wb, spec)
    dashboard = write_dashboard_sheet(wb, spec, refs, parameters)
    return wb, refs, parameters, dashboard


def _errors(spec: WorkbookSpec) -> set[str]:
    return {item.code for item in validate_workbook_spec(spec) if item.severity == "ERROR"}


def test_01e_metric_contract_and_dashboard_spec_are_valid():
    assert validate_workbook_spec(_spec()) == ()


def test_01e_dashboard_is_visible_protected_and_contains_context_and_kpis():
    wb, _, _, dashboard = _build(_spec())
    ws = wb["PAINEL"]
    assert ws.sheet_state == "visible"
    assert ws.protection.sheet is True
    assert ws["A1"].value == "PAINEL DE GESTAO · DCS"
    assert dashboard.context_cells.keys() == {"reference_period", "comparison_period", "course"}
    assert len(dashboard.kpis) == 4
    assert ws._charts == []


def test_01e_nps_formula_uses_components_filters_and_not_average_of_nps():
    wb, _, _, dashboard = _build(_spec())
    ws = wb["PAINEL"]
    formula = ws[dashboard.kpis[0].value_cell].value.upper()
    assert "SUMPRODUCT(" in formula
    assert "P_REFERENCE_PERIOD" in formula
    assert "P_COURSE" in formula
    assert "PROMOTORES" not in formula  # formulas bind physical ranges, not display labels
    assert "AVERAGE(" not in formula
    assert "*100" in formula


def test_01e_comparison_uses_comparison_period_and_delta_is_numeric_formula():
    wb, _, _, dashboard = _build(_spec())
    ws = wb["PAINEL"]
    kpi = dashboard.kpis[0]
    compare = ws[kpi.comparison_cell].value.upper()
    delta = ws[kpi.delta_cell].value.upper()
    assert "P_COMPARISON_PERIOD" in compare
    assert "P_REFERENCE_PERIOD" not in compare
    assert kpi.value_cell in delta
    assert kpi.comparison_cell in delta
    assert "TEXT(" not in delta


def test_01e_number_formats_follow_metric_units():
    wb, _, _, dashboard = _build(_spec())
    ws = wb["PAINEL"]
    formats = [ws[item.value_cell].number_format for item in dashboard.kpis]
    assert formats == ["0.0", "0.0%", "0.0", "#,##0"]
    assert ws[dashboard.kpis[1].delta_cell].number_format == "+0.0%;-0.0%;0.0%"


def test_01e_targets_remain_numeric_and_use_metric_format():
    wb, _, _, dashboard = _build(_spec())
    ws = wb["PAINEL"]
    assert ws[dashboard.kpis[0].target_cell].value == 70
    assert ws[dashboard.kpis[1].target_cell].value == 0.85
    assert ws[dashboard.kpis[1].target_cell].number_format == "0.0%"


def test_01e_empty_fact_snapshot_renders_blank_kpis_not_false_zero():
    wb, _, _, dashboard = _build(_spec(empty_facts=True))
    ws = wb["PAINEL"]
    assert all(ws[item.value_cell].value == '=""' for item in dashboard.kpis)


def test_01e_chart_specs_only_reserve_anchors_no_chart_objects_yet():
    wb, _, _, dashboard = _build(_spec(include_chart=True))
    ws = wb["PAINEL"]
    assert dashboard.chart_anchors == {"nps_evolution": dashboard.chart_anchors["nps_evolution"]}
    assert dashboard.chart_anchors["nps_evolution"].startswith("A")
    assert ws._charts == []
    assert any(cell.value == "VISUALIZAÇÕES" for row in ws.iter_rows() for cell in row)


def test_01e_preflight_failure_does_not_create_dashboard_sheet():
    spec = _spec()
    wb = Workbook()
    wb.remove(wb.active)
    refs = {dataset.code: write_dataset_sheet(wb, dataset) for dataset in spec.datasets}
    parameters = write_parameter_system(wb, spec)
    refs.pop("academic_facts")
    before = tuple(wb.sheetnames)
    with pytest.raises(DashboardWriteError) as exc:
        write_dashboard_sheet(wb, spec, refs, parameters)
    assert exc.value.code == "dashboard.missing_dataset_ref"
    assert tuple(wb.sheetnames) == before
    assert "PAINEL" not in wb.sheetnames


def test_01e_contract_rejects_missing_nps_component():
    spec = _spec()
    binding = spec.metric_bindings[0]
    bad = replace(binding, components={"respondents": "respondents", "promoters": "promoters"})
    spec = replace(spec, metric_bindings=(bad,) + spec.metric_bindings[1:])
    assert "metric_binding.missing_components" in _errors(spec)


def test_01e_contract_rejects_comparison_when_dimension_is_not_bound():
    spec = _spec()
    binding = replace(spec.metric_bindings[0], filter_parameters={"course": "course"})
    spec = replace(spec, metric_bindings=(binding,) + spec.metric_bindings[1:])
    assert "kpi.comparison_dimension_not_bound" in _errors(spec)


def test_01e_dashboard_roundtrip_preserves_formulas_and_protection(tmp_path: Path):
    wb, _, _, dashboard = _build(_spec())
    path = tmp_path / "dashboard.xlsx"
    wb.save(path)
    reopened = load_workbook(path, data_only=False)
    ws = reopened["PAINEL"]
    assert ws.protection.sheet is True
    assert ws[dashboard.kpis[0].value_cell].data_type == "f"
    assert "SUMPRODUCT(" in ws[dashboard.kpis[0].value_cell].value.upper()
    assert ws._charts == []


def test_01e_contract_layers_remain_openpyxl_independent():
    root = Path(__file__).resolve().parents[1] / "excel_official"
    for rel in ["contract.py", "context.py", "validation.py", "theme.py", "adapters/base.py"]:
        assert "openpyxl" not in (root / rel).read_text(encoding="utf-8")
