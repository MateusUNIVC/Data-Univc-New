from __future__ import annotations

from dataclasses import replace
from pathlib import Path
import sys

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from test_excel_official_dashboard_v1 import _build, _spec  # noqa: E402

from excel_official import (  # noqa: E402
    MATRIX_SHEET,
    MatrixSpec,
    MatrixWriteError,
    validate_workbook_spec,
    write_matrix_sheet,
)


def _with_matrix(**kwargs):
    base = MatrixSpec(
        metric_code="academic.nps",
        row_dimension="course",
        column_dimension="period",
        target=70,
        delta=True,
    )
    return replace(_spec(), matrix=replace(base, **kwargs))


def _errors(spec) -> set[str]:
    return {item.code for item in validate_workbook_spec(spec) if item.severity == "ERROR"}


def test_01g_matrix_materializes_course_by_period_with_target_and_delta():
    spec = _with_matrix()
    wb, refs, params, _ = _build(spec)
    matrix_ref = write_matrix_sheet(wb, spec, refs, params)

    assert matrix_ref is not None
    ws = wb[MATRIX_SHEET]
    assert ws.sheet_state == "visible"
    assert ws.protection.sheet is True
    assert ws["A1"].value == "MATRIZ DE INDICADORES"
    assert ws["A4"].value == "Indicador"
    assert ws["B4"].value == "NPS da Instituicao"
    assert ws["A6"].value == "Curso"
    assert ws["B6"].value == "2026-SEM1"
    assert ws["C6"].value == "2026-SEM2"
    assert ws["D6"].value == "Meta"
    assert ws["E6"].value == "Δ"
    assert matrix_ref.value_range == "B7:C8"
    assert matrix_ref.target_column == 4
    assert matrix_ref.delta_column == 5


def test_01g_matrix_formulas_override_both_axes_and_reuse_metric_engine():
    spec = _with_matrix()
    wb, refs, params, _ = _build(spec)
    write_matrix_sheet(wb, spec, refs, params)
    formula = wb[MATRIX_SHEET]["B7"].value.upper()

    assert "SUMPRODUCT(" in formula
    assert f'"{wb[MATRIX_SHEET]["A7"].value.upper()}"' in formula
    assert '"2026-SEM1"' in formula
    assert "P_COURSE" not in formula
    assert "P_REFERENCE_PERIOD" not in formula
    assert "AVERAGE(" not in formula
    assert "*100" in formula


def test_01g_target_and_delta_use_metric_formats_and_numeric_formulas():
    spec = _with_matrix()
    wb, refs, params, _ = _build(spec)
    write_matrix_sheet(wb, spec, refs, params)
    ws = wb[MATRIX_SHEET]

    assert ws["D7"].value == 70
    assert ws["D7"].number_format == "0.0"
    assert ws["E7"].value == '=IF(OR(C7="",B7=""),"",C7-B7)'
    assert ws["E7"].number_format == "+0.0;-0.0;0.0"


def test_01g_percentage_matrix_preserves_percentage_number_format():
    spec = _with_matrix(metric_code="academic.approval", target=0.85)
    wb, refs, params, _ = _build(spec)
    write_matrix_sheet(wb, spec, refs, params)
    ws = wb[MATRIX_SHEET]

    assert ws["B4"].value == "Aprovacao"
    assert ws["B7"].number_format == "0.0%"
    assert ws["D7"].number_format == "0.0%"
    assert ws["E7"].number_format == "+0.0%;-0.0%;0.0%"


def test_01g_row_sorting_changes_dimension_order_but_columns_keep_declared_order():
    spec = _with_matrix(sorting="label_desc")
    wb, refs, params, _ = _build(spec)
    write_matrix_sheet(wb, spec, refs, params)
    ws = wb[MATRIX_SHEET]

    assert ws["A7"].value == "Psicologia"
    assert ws["A8"].value == "Odontologia"
    assert ws["B6"].value == "2026-SEM1"
    assert ws["C6"].value == "2026-SEM2"


def test_01g_empty_behavior_zero_wraps_blank_metric_result():
    spec = _with_matrix(empty_behavior="zero")
    wb, refs, params, _ = _build(spec)
    write_matrix_sheet(wb, spec, refs, params)
    formula = wb[MATRIX_SHEET]["B7"].value.upper()
    assert formula.startswith("=IF((IF(")
    assert '=\"\",0,' in formula


def test_01g_matrix_none_is_noop_and_does_not_create_sheet():
    spec = _spec()
    wb, refs, params, _ = _build(spec)
    result = write_matrix_sheet(wb, spec, refs, params)
    assert result is None
    assert MATRIX_SHEET not in wb.sheetnames


def test_01g_preflight_failure_does_not_create_matrix_sheet():
    spec = _with_matrix()
    wb, refs, params, _ = _build(spec)
    refs.pop("academic_facts")
    before = tuple(wb.sheetnames)

    with pytest.raises(MatrixWriteError) as exc:
        write_matrix_sheet(wb, spec, refs, params)
    assert exc.value.code == "matrix.missing_dataset_ref"
    assert tuple(wb.sheetnames) == before
    assert MATRIX_SHEET not in wb.sheetnames


def test_01g_contract_rejects_same_dimensions_invalid_sort_and_invalid_empty_behavior():
    assert "matrix.same_dimensions" in _errors(_with_matrix(column_dimension="course"))
    assert "matrix.invalid_sort" in _errors(_with_matrix(sorting="metric_desc"))
    assert "matrix.invalid_empty_behavior" in _errors(_with_matrix(empty_behavior="na"))


def test_01g_contract_rejects_metric_dimension_not_allowed():
    spec = _with_matrix()
    metric = next(item for item in spec.metrics if item.code == "academic.nps")
    metrics = tuple(replace(item, allowed_dimensions=("period",)) if item.code == metric.code else item for item in spec.metrics)
    invalid = replace(spec, metrics=metrics)
    assert "matrix.dimension_not_allowed" in _errors(invalid)


def test_01g_roundtrip_preserves_matrix_formulas_formats_and_protection(tmp_path: Path):
    spec = _with_matrix()
    wb, refs, params, _ = _build(spec)
    write_matrix_sheet(wb, spec, refs, params)
    path = tmp_path / "matrix.xlsx"
    wb.save(path)

    reopened = load_workbook(path, data_only=False)
    ws = reopened[MATRIX_SHEET]
    assert ws.protection.sheet is True
    assert ws.freeze_panes == "B7"
    assert ws["B7"].data_type == "f"
    assert ws["E7"].data_type == "f"
    assert reopened.calculation.calcMode == "auto"
    assert reopened.calculation.fullCalcOnLoad is True
    assert reopened.calculation.forceFullCalc is True


def test_01g_matrix_formulas_remain_excel_2019_compatible():
    spec = _with_matrix()
    wb, refs, params, _ = _build(spec)
    write_matrix_sheet(wb, spec, refs, params)
    formulas = "\n".join(
        str(cell.value).upper()
        for row in wb[MATRIX_SHEET].iter_rows()
        for cell in row
        if cell.data_type == "f"
    )
    for unsupported in ("FILTER(", "XLOOKUP(", "SORT(", "UNIQUE(", "INDIRECT(", "OFFSET("):
        assert unsupported not in formulas
