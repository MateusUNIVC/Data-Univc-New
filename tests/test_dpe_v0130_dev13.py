from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

import release_info
from dpe_excel_modern import build_dpe_operational_workbook

ROOT = Path(__file__).resolve().parents[1]


def _payload():
    return {
        "selected_period": {"id": 1, "period": "2026-09", "status": "CALCULATED"},
        "revenues": {"categories": [{"code": "COURSE_REVENUE", "name": "Receita de curso", "scope": "COURSE"}]},
        "revenue_entries": [
            {"id": 1, "period_offering_id": 10, "offering_label": "Administração", "course_name": "Administração", "category_code": "COURSE_REVENUE", "category_name": "Receita de curso", "category_scope": "COURSE", "description": "Receita do curso", "amount": 1000, "source_type": "MANUAL"},
            {"id": 2, "period_offering_id": None, "offering_label": None, "course_name": None, "category_code": "OTHER", "category_name": "Outras receitas", "category_scope": "INSTITUTIONAL", "description": "Receita institucional", "amount": 100, "source_type": "MANUAL"},
        ],
        "expenses": {"categories": [{"code": "DOCENTE", "name": "Custo docente"}], "expenses": [
            {"id": 1, "period": "2026-09", "description": "Despesa", "amount": 300, "expense_kind": "GENERAL", "expense_scope": "SHARED", "category_code": "DOCENTE", "category_name": "Custo docente", "allocation_rule_code": "EQUAL", "allocation_rule_name": "Divisão igualitária", "allocation_driver": "EQUAL", "source_type": "MANUAL"}
        ]},
        "teaching": {"period_teachers": [], "activities": []},
        "allocation": {"selected_run": {"id": 1, "run_number": 1, "status": "OFFICIAL", "expense_total": 300, "allocated_total": 300, "unallocated_total": 0, "summary": {"result_count": 1}, "results": [
            {"expense_id": 1, "expense_description": "Despesa", "period_offering_id": 10, "offering_label": "Administração", "driver_type": "EQUAL", "allocated_amount": 300, "percentage": 100, "numerator": 1, "denominator": 1, "basis": {"basis_type": "EQUAL"}}
        ]}},
        "analytics": {"courses": [{"course": "Administração", "offering_count": 1}], "limitations": []},
        "management": {"targets": [], "actions": []},
        "closure": {"checklist": {"checks": []}},
        "overview": {"offerings": [{"period_offering_id": 10, "product": "Administração", "label": "Administração", "modality": "PRESENCIAL", "shift": "Noturno", "location": "Sede", "active_students": 10, "students_ready": True, "revenue_ready": True, "cost_available": True, "revenue": 1000, "allocated_cost": 300, "economic_result": 700}]},
    }


def test_release_preserves_dev13_excel_contract_without_schema_regression():
    assert release_info.APP_VERSION.startswith("0.13.0")
    if release_info.APP_VERSION != "0.13.0":
        assert int(release_info.APP_VERSION.rsplit(".", 1)[-1]) >= 13
    assert release_info.SCHEMA_VERSION >= 48


def test_modern_excel_builder_has_visible_database_sheets_and_formulas():
    output = build_dpe_operational_workbook(_payload())
    wb = load_workbook(BytesIO(output.getvalue()), data_only=False)
    assert wb.sheetnames[:9] == ["PAINEL", "RESULTADO", "CURSOS", "RECEITAS", "DESPESAS", "DOCENCIA", "RATEIOS", "METAS_PLANOS", "QUALIDADE"]
    for sheet in ("BASE_CURSOS", "BASE_RECEITAS", "BASE_DESPESAS", "BASE_DOCENCIA", "BASE_RATEIOS"):
        assert wb[sheet].sheet_state == "visible"
        assert wb[sheet].tables
    assert wb["PAINEL"]["A5"].value.startswith("=SUM(")
    assert wb["RESULTADO"]["E5"].value.startswith("=SUMIFS(")
    assert wb["RESULTADO"]["G5"].value == "=E5-F5"
    assert len(wb["PAINEL"]._charts) == 2


def test_full_excel_route_uses_modern_builder_and_old_indicator_export_is_retired():
    router = (ROOT / "dpe_router.py").read_text(encoding="utf-8")
    html = (ROOT / "templates" / "dpe.html").read_text(encoding="utf-8")
    js = (ROOT / "static" / "js" / "dpe_v2.js").read_text(encoding="utf-8")
    assert "DPEExcelExportRepository" in router
    assert "build_dpe_operational_workbook" in router
    assert '@router.get("/api/dpe/excel/{indicator_code}")' not in router
    assert 'id="dpeExportExcel"' in html
    assert "period_id=" in js and "dpeExportExcel" in js


def test_modern_excel_does_not_depend_on_legacy_measurement_builder():
    modern = (ROOT / "dpe_excel_modern.py").read_text(encoding="utf-8")
    export = (ROOT / "dpe_excel_export.py").read_text(encoding="utf-8")
    forbidden = (
        "build_management_dashboard",
        "from management_repository import ManagementRepository",
        "from management_catalog import indicator_spec",
        "from management_catalog import metric_spec",
    )
    for token in forbidden:
        assert token not in modern
        assert token not in export
