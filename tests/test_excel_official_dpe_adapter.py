from __future__ import annotations

from dataclasses import replace
from datetime import datetime, timezone
from io import BytesIO

from openpyxl import load_workbook
import pytest

from dpe_excel_official import build_dpe_excel_official_artifact_from_payload, build_dpe_snapshot_context
from dpe_excel_parity import audit_dpe_payload_parity
from excel_official.adapters.dpe import DPEAdapter, DPEAdapterError, DPE_METRICS
from excel_official.context import AdapterInput


def _payload():
    current = {
        "period_id": 2, "period": "2026-09", "status": "CALCULATED",
        "total_revenue": 210000.0, "course_revenue": 200000.0, "institutional_revenue": 10000.0,
        "expense_total": 110000.0, "economic_result": 100000.0, "operating_margin_percent": 47.619047619,
        "active_students": 180,
    }
    previous = {
        "period_id": 1, "period": "2026-08", "status": "CLOSED",
        "total_revenue": 180000.0, "course_revenue": 170000.0, "institutional_revenue": 10000.0,
        "expense_total": 100000.0, "economic_result": 80000.0, "operating_margin_percent": 44.444444444,
        "active_students": 170,
    }
    return {
        "selected_period": {"id": 2, "period": "2026-09", "status": "CALCULATED"},
        "revenues": {"categories": []},
        "revenue_entries": [
            {"id": 1, "period_offering_id": 10, "offering_label": "Administração · Noturno", "course_name": "Administração", "category_code": "COURSE", "category_name": "Receita de curso", "category_scope": "COURSE", "description": "Mensalidades", "amount": 120000.0, "source_type": "MANUAL"},
            {"id": 2, "period_offering_id": 20, "offering_label": "Farmácia · Noturno", "course_name": "Farmácia", "category_code": "COURSE", "category_name": "Receita de curso", "category_scope": "COURSE", "description": "Mensalidades", "amount": 80000.0, "source_type": "MANUAL"},
            {"id": 3, "period_offering_id": None, "offering_label": None, "course_name": None, "category_code": "OTHER", "category_name": "Outras receitas", "category_scope": "INSTITUTIONAL", "description": "Locação de espaço", "amount": 10000.0, "source_type": "MANUAL"},
        ],
        "expenses": {"expenses": [
            {"id": 1, "period": "2026-09", "description": "Folha docente", "amount": 60000.0, "expense_kind": "PAYROLL", "expense_scope": "SHARED", "cost_center_code": "ADMIN", "cost_center_name": "Administrativo", "category_code": "FOLHA", "category_name": "Folha docente", "allocation_rule_name": "EQUAL", "allocation_driver": "EQUAL", "source_type": "MANUAL"},
            {"id": 2, "period": "2026-09", "description": "Software", "amount": 20000.0, "expense_kind": "GENERAL", "expense_scope": "SHARED", "cost_center_code": "LAB", "cost_center_name": "Laboratórios", "category_code": "SOFT", "category_name": "Software", "allocation_rule_name": "DIRECT", "allocation_driver": "DIRECT", "source_type": "MANUAL"},
            {"id": 3, "period": "2026-09", "description": "Energia", "amount": 30000.0, "expense_kind": "GENERAL", "expense_scope": "INSTITUTIONAL", "cost_center_code": "ADMIN", "cost_center_name": "Administrativo", "category_code": "ENERG", "category_name": "Energia", "allocation_rule_name": "INSTITUTIONAL", "allocation_driver": "INSTITUTIONAL", "source_type": "MANUAL"},
        ]},
        "teaching": {"period_teachers": [{"teacher_id": 1, "teacher": {"display_name": "Docente A"}, "relationship_type": "CLT", "workload_hours": 40.0, "payroll_total": 60000.0, "payroll_count": 1}], "activities": [{"id": 1, "workload_hours": 40.0}], "payroll": [{"match_status": "CONFIRMED"}], "summary": {"period_teacher_count": 1, "activity_count": 1, "total_workload_hours": 40.0, "payroll_count": 1, "payroll_linked_count": 1, "payroll_pending_count": 0, "subject_count": 1}},
        "allocation": {"selected_run": {"id": 5, "run_number": 2, "status": "OFFICIAL", "expense_total": 100000.0, "allocated_total": 100000.0, "unallocated_total": 0.0, "results": [
            {"expense_id": 1, "expense_description": "Folha docente", "period_offering_id": 10, "offering_label": "Administração · Noturno", "driver_type": "EQUAL", "allocated_amount": 40000.0, "expense_snapshot": {"expense_kind": "PAYROLL", "classification": {"cost_center": {"code": "ADMIN", "name": "Administrativo"}, "category": {"code": "FOLHA", "name": "Folha docente"}}}, "offering_snapshot": {"product": {"name": "Administração"}}},
            {"expense_id": 1, "expense_description": "Folha docente", "period_offering_id": 20, "offering_label": "Farmácia · Noturno", "driver_type": "EQUAL", "allocated_amount": 20000.0, "expense_snapshot": {"expense_kind": "PAYROLL", "classification": {"cost_center": {"code": "ADMIN", "name": "Administrativo"}, "category": {"code": "FOLHA", "name": "Folha docente"}}}, "offering_snapshot": {"product": {"name": "Farmácia"}}},
            {"expense_id": 2, "expense_description": "Software", "period_offering_id": 10, "offering_label": "Administração · Noturno", "driver_type": "DIRECT", "allocated_amount": 20000.0, "expense_snapshot": {"expense_kind": "GENERAL", "classification": {"cost_center": {"code": "LAB", "name": "Laboratórios"}, "category": {"code": "SOFT", "name": "Software"}}}, "offering_snapshot": {"product": {"name": "Administração"}}},
            {"expense_id": 3, "expense_description": "Energia", "period_offering_id": 20, "offering_label": "Farmácia · Noturno", "driver_type": "EQUAL", "allocated_amount": 20000.0, "expense_snapshot": {"expense_kind": "GENERAL", "classification": {"cost_center": {"code": "ADMIN", "name": "Administrativo"}, "category": {"code": "ENERG", "name": "Energia"}}}, "offering_snapshot": {"product": {"name": "Farmácia"}}},
        ]}},
        "analytics": {
            "periods": [{"id": 1, "period": "2026-08", "status": "CLOSED"}, {"id": 2, "period": "2026-09", "status": "CALCULATED"}],
            "selected_period": {"id": 2, "period": "2026-09", "status": "CALCULATED"},
            "previous_period": {"id": 1, "period": "2026-08", "status": "CLOSED"},
            "cards": {
                "total_revenue": {"value": 210000.0, "comparison": {"absolute": 30000.0, "percent": 16.666666667, "direction": "up"}},
                "course_revenue": {"value": 200000.0},
                "institutional_revenue": {"value": 10000.0},
                "expense_total": {"value": 110000.0, "comparison": {"absolute": 10000.0, "percent": 10.0, "direction": "up"}},
                "economic_result": {"value": 100000.0, "comparison": {"absolute": 20000.0, "percent": 25.0, "direction": "up"}},
                "operating_margin_percent": {"value": 47.619047619, "comparison": {"absolute": 3.174603175, "percent": 7.142857143, "direction": "up"}},
                "active_students": {"value": 180},
            },
            "trend": [previous, current],
            "expenses_by_sector": [{"key": "Administrativo", "amount": 90000.0}, {"key": "Laboratórios", "amount": 20000.0}],
            "courses": [
                {"course_key": "course:1", "course": "Administração", "offering_count": 1, "active_students": 100, "revenue": 120000.0, "allocated_cost": 60000.0, "economic_result": 60000.0, "margin_percent": 50.0, "teaching_cost": 40000.0, "direct_cost": 20000.0, "shared_cost": 0.0, "revenue_complete": True, "cost_complete": True},
                {"course_key": "course:2", "course": "Farmácia", "offering_count": 1, "active_students": 80, "revenue": 80000.0, "allocated_cost": 40000.0, "economic_result": 40000.0, "margin_percent": 50.0, "teaching_cost": 20000.0, "direct_cost": 0.0, "shared_cost": 20000.0, "revenue_complete": True, "cost_complete": True},
            ],
            "limitations": [],
        },
        "management": {
            "facts": [
                {"indicator_code": "DPE-RESULT", "metric_key": "institutional_result", "value": 100000.0, "dimensions": {}, "dimension_key": "TOTAL", "dimension_label": "TOTAL", "source": "analytics", "available": True},
                {"indicator_code": "DPE-RESULT", "metric_key": "institutional_margin_pct", "value": 47.619047619, "dimensions": {}, "dimension_key": "TOTAL", "dimension_label": "TOTAL", "source": "analytics", "available": True},
                {"indicator_code": "DPE-RESULT", "metric_key": "coverage_index", "value": 1.909090909, "dimensions": {}, "dimension_key": "TOTAL", "dimension_label": "TOTAL", "source": "analytics", "available": True},
                {"indicator_code": "DPE-REVENUE", "metric_key": "total_revenue", "value": 210000.0, "dimensions": {}, "dimension_key": "TOTAL", "dimension_label": "TOTAL", "source": "revenue_ledger", "available": True},
                {"indicator_code": "DPE-REVENUE", "metric_key": "institutional_revenue", "value": 10000.0, "dimensions": {}, "dimension_key": "TOTAL", "dimension_label": "TOTAL", "source": "revenue_ledger", "available": True},
                {"indicator_code": "DPE-EXPENSE", "metric_key": "total_expense", "value": 110000.0, "dimensions": {}, "dimension_key": "TOTAL", "dimension_label": "TOTAL", "source": "expense_ledger", "available": True},
                {"indicator_code": "DPE-EXPENSE", "metric_key": "institutional_expense", "value": 30000.0, "dimensions": {}, "dimension_key": "TOTAL", "dimension_label": "TOTAL", "source": "expense_ledger", "available": True},
                {"indicator_code": "DPE-EXPENSE", "metric_key": "allocatable_expense", "value": 80000.0, "dimensions": {}, "dimension_key": "TOTAL", "dimension_label": "TOTAL", "source": "expense_ledger", "available": True},
                {"indicator_code": "DPE-TEACHING", "metric_key": "teaching_cost", "value": 60000.0, "dimensions": {}, "dimension_key": "TOTAL", "dimension_label": "TOTAL", "source": "allocation_waterfall", "available": True},
                {"indicator_code": "DPE-TEACHING", "metric_key": "total_workload_hours", "value": 40.0, "dimensions": {}, "dimension_key": "TOTAL", "dimension_label": "TOTAL", "source": "teaching_activities", "available": True},
                {"indicator_code": "DPE-TEACHING", "metric_key": "avg_workload_hours_per_teacher", "value": 40.0, "dimensions": {}, "dimension_key": "TOTAL", "dimension_label": "TOTAL", "source": "teaching_activities", "available": True},
                {"indicator_code": "DPE-TEACHING", "metric_key": "teacher_count", "value": 1, "dimensions": {}, "dimension_key": "TOTAL", "dimension_label": "TOTAL", "source": "teaching_activities", "available": True},
                {"indicator_code": "DPE-ALLOCATION", "metric_key": "allocated_expense", "value": 100000.0, "dimensions": {}, "dimension_key": "TOTAL", "dimension_label": "TOTAL", "source": "official_allocation_run", "available": True},
                {"indicator_code": "DPE-ALLOCATION", "metric_key": "unallocated_expense", "value": 0.0, "dimensions": {}, "dimension_key": "TOTAL", "dimension_label": "TOTAL", "source": "official_allocation_run", "available": True},
                {"indicator_code": "DPE-ALLOCATION", "metric_key": "reconciliation_pct", "value": 100.0, "dimensions": {}, "dimension_key": "TOTAL", "dimension_label": "TOTAL", "source": "official_allocation_run", "available": True},
                {"indicator_code": "DPE-RESULT", "metric_key": "course_result", "value": 60000.0, "dimensions": {"course": "Administração", "academic_directorate": "DTNH"}, "dimension_key": "course-adm", "dimension_label": "Curso: Administração | Diretoria acadêmica: DTNH", "source": "analytics", "available": True},
            ],
            "targets": [
                {"indicator_code": "DPE-REVENUE", "metric_key": "total_revenue", "indicator_label": "Receitas", "metric_label": "Receita total", "dimension_key": "TOTAL", "dimension_label": "TOTAL", "valid_from": "2026-01", "valid_to": None, "target": 200000.0, "attention": 180000.0, "target_min": None, "target_max": None, "active_for_period": True, "historical_metric": False, "current_value": 210000.0, "status": {"key": "GOOD", "label": "Dentro da meta"}, "direction": "higher", "fact": {"dimensions": {}, "dimension_key": "TOTAL"}},
                {"indicator_code": "DPE-RESULT", "metric_key": "institutional_margin_pct", "indicator_label": "Resultados", "metric_label": "Margem institucional", "dimension_key": "TOTAL", "dimension_label": "TOTAL", "valid_from": "2026-01", "valid_to": None, "target": 20.0, "attention": 10.0, "target_min": None, "target_max": None, "active_for_period": True, "historical_metric": False, "current_value": 47.619047619, "status": {"key": "GOOD", "label": "Dentro da meta"}, "direction": "higher", "fact": {"dimensions": {}, "dimension_key": "TOTAL"}},
                {"indicator_code": "DPE-RESULT", "metric_key": "course_result", "indicator_label": "Resultados", "metric_label": "Resultado do curso", "dimension_key": "course-adm", "dimension_label": "Curso: Administração | Diretoria acadêmica: DTNH", "valid_from": "2026-01", "valid_to": None, "target": 50000.0, "attention": 40000.0, "target_min": None, "target_max": None, "active_for_period": True, "historical_metric": False, "current_value": 60000.0, "status": {"key": "GOOD", "label": "Dentro da meta"}, "direction": "higher", "fact": {"dimensions": {"course": "Administração", "academic_directorate": "DTNH"}, "dimension_key": "course-adm"}},
                {"indicator_code": "DPE-TEACHING", "metric_key": "avg_workload_hours_per_teacher", "indicator_label": "Docência", "metric_label": "Carga média por docente", "dimension_key": "TOTAL", "dimension_label": "TOTAL", "valid_from": "2026-01", "valid_to": None, "target": None, "attention": None, "target_min": 30.0, "target_max": 50.0, "active_for_period": True, "historical_metric": False, "current_value": 40.0, "status": {"key": "GOOD", "label": "Dentro da meta"}, "direction": "range", "fact": {"dimensions": {}, "dimension_key": "TOTAL"}},
                {"indicator_code": "DPE-ALLOCATION", "metric_key": "reconciliation_pct", "indicator_label": "Distribuição", "metric_label": "Reconciliação", "dimension_key": "TOTAL", "dimension_label": "TOTAL", "valid_from": "2026-01", "valid_to": None, "target": 100.0, "attention": 95.0, "target_min": None, "target_max": None, "active_for_period": True, "historical_metric": False, "current_value": 100.0, "status": {"key": "GOOD", "label": "Dentro da meta"}, "direction": "higher", "fact": {"dimensions": {}, "dimension_key": "TOTAL"}},
            ],
            "actions": [{"indicator_code": "DPE-RESULT", "metric_key": "course_result", "indicator_label": "Resultados", "metric_label": "Resultado do curso", "problem": "Margem abaixo do esperado", "probable_cause": "Custo elevado", "corrective_action": "Revisar composição de custos", "responsible": "DPE", "due_date": "2026-10-31", "status": "Em andamento", "effective_status": "Em andamento"}],
        },
        "closure": {"checklist": {"checks": [{"code": "OFFICIAL_RUN", "label": "Cálculo oficial", "status": "PASS", "detail": "Run oficial atual."}], "summary": {"pass_count": 1, "warning_count": 0, "blocker_count": 0, "review_pending_count": 0, "check_count": 1, "checklist_ready": True, "can_close": True}}, "official_run": {"id": 5, "run_number": 2, "status": "OFFICIAL", "current": True, "expense_total": 100000.0, "allocated_total": 100000.0, "unallocated_total": 0.0}},
        "overview": {
            "summary": {"total_revenue": 210000.0, "expense_total": 110000.0, "allocated_cost": 100000.0, "pending_distribution_total": 0.0, "allocation_coverage_percent": 100.0},
            "offerings": [
                {"period_offering_id": 10, "product": "Administração", "label": "Administração · Noturno", "modality": "PRESENCIAL", "shift": "Noturno", "location": "Sede", "active_students": 100, "revenue": 120000.0, "allocated_cost": 60000.0, "economic_result": 60000.0, "margin_percent": 50.0},
                {"period_offering_id": 20, "product": "Farmácia", "label": "Farmácia · Noturno", "modality": "PRESENCIAL", "shift": "Noturno", "location": "Sede", "active_students": 80, "revenue": 80000.0, "allocated_cost": 40000.0, "economic_result": 40000.0, "margin_percent": 50.0},
            ],
        },
        "app_version": "0.13.0",
        "schema_version": 53,
        "generated_at": datetime(2026, 10, 2, 12, 0, tzinfo=timezone.utc).isoformat(),
    }


def _spec(payload=None):
    payload = payload or _payload()
    ctx = build_dpe_snapshot_context(payload, generated_by="dpe@test.local", export_id="dpe-test-export")
    return DPEAdapter().build_spec(AdapterInput(ctx, payload, {}))


def test_dpe_adapter_builds_fixed_competence_and_interactive_business_dimensions():
    spec = _spec()
    params = {row.code: row for row in spec.parameters}
    assert params["period"].editable is False
    assert params["period"].initial_value == "2026-09"
    assert params["course"].editable is True
    assert params["context"].depends_on == ("course",)
    assert params["cost_center"].editable is True
    assert spec.capabilities.interactive_dimensions == ("course", "context", "cost_center")


def test_dpe_adapter_separates_institutional_and_course_semantics():
    spec = _spec()
    metrics = {row.code: row for row in spec.metrics}
    assert metrics[DPE_METRICS["total_revenue"]].offline_recut is False
    assert metrics[DPE_METRICS["course_result"]].allowed_dimensions == ("course", "context")
    bindings = {row.metric_code: row for row in spec.metric_bindings}
    assert bindings[DPE_METRICS["total_revenue"]].dataset_code == "dpe_period_summary"
    assert bindings[DPE_METRICS["course_result"]].dataset_code == "dpe_contexts"
    assert bindings[DPE_METRICS["selected_expense"]].filter_parameters == {"cost_center": "cost_center"}


def test_dpe_adapter_preserves_financial_bases_and_official_allocation_memory():
    spec = _spec()
    datasets = {row.code: row for row in spec.datasets}
    assert sum(row["amount"] for row in datasets["dpe_revenues"].rows) == 210000.0
    assert sum(row["amount"] for row in datasets["dpe_expenses"].rows) == 110000.0
    assert sum(row["allocated_amount"] for row in datasets["dpe_allocations"].rows) == 100000.0
    assert datasets["dpe_expenses"].sensitivity == "restricted"
    assert datasets["dpe_teaching"].sensitivity == "restricted"


def test_dpe_matrix_is_financial_composition_by_course():
    spec = _spec()
    assert spec.matrix.metric_code == DPE_METRICS["component_amount"]
    assert spec.matrix.row_dimension == "course"
    assert spec.matrix.column_dimension == "component"
    components = next(row for row in spec.datasets if row.code == "dpe_course_components")
    assert {row["component"] for row in components.rows} == {"Receita", "Custo docente", "Custo direto", "Custo compartilhado", "Resultado"}


def test_dpe_excel_official_builds_clean_release_workbook():
    artifact = build_dpe_excel_official_artifact_from_payload(_payload(), generated_by="dpe@test.local", export_id="dpe-test-export")
    assert artifact.audit.release_allowed is True
    assert artifact.audit.findings == ()
    assert artifact.workbook.sheetnames[:6] == ["LEIA-ME", "PARAMETROS", "PAINEL", "QUALIDADE E GOVERNANÇA", "MATRIZ", "PLANO_DE_ACAO"]
    assert len(artifact.refs.dashboard.kpis) == 8
    assert len(artifact.refs.charts.charts) == 5


def test_dpe_workbook_roundtrip_has_no_external_links_or_vba():
    artifact = build_dpe_excel_official_artifact_from_payload(_payload(), generated_by="dpe@test.local")
    wb = load_workbook(BytesIO(artifact.to_bytes().getvalue()), data_only=False, keep_links=True)
    assert wb.vba_archive is None
    assert not wb._external_links
    assert wb["PARAMETROS"].sheet_state == "visible"
    assert wb["CALC"].sheet_state == "visible"
    assert wb["LISTAS DE APOIO"].sheet_state == "visible"


def test_dpe_payload_parity_recomposes_ledgers_and_official_run():
    report = audit_dpe_payload_parity(_payload())
    assert report.status == "CANDIDATE_PASS"
    assert len(report.cases) >= 55
    assert not report.failures


def test_dpe_payload_parity_blocks_divergent_course_cost():
    payload = _payload()
    payload["analytics"]["courses"][0]["allocated_cost"] = 99999.0
    report = audit_dpe_payload_parity(payload)
    assert report.status == "BLOCKED"
    assert any(case.code == "course.allocated_cost" and not case.passed for case in report.cases)


def test_dpe_adapter_rejects_snapshot_without_contexts():
    payload = _payload()
    payload["overview"]["offerings"] = []
    ctx = build_dpe_snapshot_context(payload, generated_by="x")
    with pytest.raises(DPEAdapterError) as exc:
        DPEAdapter().build_spec(AdapterInput(ctx, payload, {}))
    assert exc.value.code == "dpe.no_contexts"


def test_dpe_previous_period_comparison_is_registered():
    spec = _spec()
    kpis = {row.metric_code: row for row in spec.dashboard.kpis}
    assert kpis[DPE_METRICS["total_revenue"]].comparison_metric_code == "dpe.previous.total_revenue"
    previous = next(row for row in spec.datasets if row.code == "dpe_previous_summary")
    assert previous.rows[0]["period"] == "2026-08"
    assert previous.rows[0]["institutional_result"] == 80000.0



def test_dpe_05b_effective_targets_bind_cards_and_preserve_range_targets():
    spec = _spec()
    datasets = {row.code: row for row in spec.datasets}
    targets = datasets["dpe_targets_effective"].rows
    by_metric = {row["metric_code"]: row for row in targets}
    assert by_metric[DPE_METRICS["total_revenue"]]["target"] == 200000.0
    assert by_metric[DPE_METRICS["institutional_margin"]]["target"] == 20.0
    assert by_metric[DPE_METRICS["course_result"]]["course_selector"] == "Administração"
    assert by_metric[DPE_METRICS["course_result"]]["target"] == 50000.0
    range_row = by_metric[DPE_METRICS["avg_workload_hours_per_teacher"]]
    assert range_row["target"] is None
    assert range_row["target_min"] == 30.0
    assert range_row["target_max"] == 50.0

    bindings = {row.metric_code: row for row in spec.target_bindings}
    assert bindings[DPE_METRICS["institutional_margin"]].value_scale == 0.01
    assert bindings[DPE_METRICS["course_result"]].criteria_parameters == {"course_selector": "course"}
    assert DPE_METRICS["avg_workload_hours_per_teacher"] not in bindings


def test_dpe_05b_management_and_closure_summary_recompose_official_sources():
    spec = _spec()
    datasets = {row.code: row for row in spec.datasets}
    management = datasets["dpe_management_summary"].rows[0]
    closure = datasets["dpe_closure_summary"].rows[0]
    assert management["institutional_revenue"] == 10000.0
    assert management["institutional_expense"] == 30000.0
    assert management["allocatable_expense"] == 80000.0
    assert management["teaching_cost"] == 60000.0
    assert management["total_workload_hours"] == 40.0
    assert management["avg_workload_hours_per_teacher"] == 40.0
    assert management["teacher_count"] == 1
    assert management["allocated_expense"] == 100000.0
    assert management["unallocated_expense"] == 0.0
    assert management["reconciliation_pct"] == 1.0
    assert closure["blocker_count"] == 0
    assert closure["payroll_pending_count"] == 0


def test_dpe_05b_action_plan_preserves_probable_cause():
    spec = _spec()
    row = spec.action_plan.rows[0]
    assert row.values["diagnosis"] == "Custo elevado"
    assert row.values["action"] == "Revisar composição de custos"


def test_dpe_05b_workbook_uses_dynamic_targets_and_documents_range_targets():
    artifact = build_dpe_excel_official_artifact_from_payload(_payload(), generated_by="dpe@test.local", export_id="dpe-05b")
    assert artifact.audit.release_allowed is True
    formulas = [
        cell.value
        for row in artifact.workbook["PAINEL"].iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    ]
    assert any("METAS EFETIVAS DPE" in formula and "dpe.total_revenue" in formula for formula in formulas)
    assert any("METAS EFETIVAS DPE" in formula and "dpe.course_result" in formula for formula in formulas)
    assert "METAS EFETIVAS DPE" in artifact.workbook.sheetnames
    assert "GOVERNANCA DPE" in artifact.workbook.sheetnames
    assert any(limit.code == "dpe.range_targets_backend" for limit in artifact.spec.limitations)


def test_dpe_05b_duplicate_active_targets_are_exposed_as_quality_collision():
    payload = _payload()
    duplicate = dict(payload["management"]["targets"][0])
    duplicate["target"] = 250000.0
    payload["management"]["targets"].append(duplicate)
    spec = _spec(payload)
    targets = next(row for row in spec.datasets if row.code == "dpe_targets_effective")
    total_revenue = next(row for row in targets.rows if row["metric_code"] == DPE_METRICS["total_revenue"])
    assert total_revenue["collision_count"] == 2
    assert total_revenue["target"] is None
    summary = next(row for row in spec.datasets if row.code == "dpe_management_summary")
    assert summary.rows[0]["target_collision_count"] == 1



def test_dpe_05b_parity_covers_previous_management_targets_and_closure():
    report = audit_dpe_payload_parity(_payload())
    codes = {case.code for case in report.cases}
    assert "previous.total_revenue.absolute_delta" in codes
    assert "management.DPE-TEACHING.avg_workload_hours_per_teacher" in codes
    assert "target.current_value.DPE-RESULT.course_result" in codes
    assert "closure.official_run.allocated_total" in codes
    assert len(report.cases) >= 55
    assert not report.failures


def test_dpe_05b_parity_blocks_drift_in_management_fact():
    payload = _payload()
    fact = next(row for row in payload["management"]["facts"] if row["indicator_code"] == "DPE-TEACHING" and row["metric_key"] == "avg_workload_hours_per_teacher")
    fact["value"] = 99.0
    report = audit_dpe_payload_parity(payload)
    assert report.status == "BLOCKED"
    assert any(case.code == "management.DPE-TEACHING.avg_workload_hours_per_teacher" and not case.passed for case in report.cases)
