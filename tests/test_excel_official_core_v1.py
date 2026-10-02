from __future__ import annotations

from dataclasses import replace
from pathlib import Path
import sys

import pytest
from openpyxl import load_workbook
from openpyxl.styles import Protection

sys.path.insert(0, str(Path(__file__).resolve().parent))
from test_excel_official_action_plan_v1 import _action_spec  # noqa: E402
from test_excel_official_dashboard_v1 import _spec  # noqa: E402
from test_excel_official_quality_v1 import _quality_spec  # noqa: E402

from excel_official import (  # noqa: E402
    ActionPlanSpec,
    ChartRole,
    ChartSpec,
    ChartType,
    ExcelOfficialCore,
    MatrixSpec,
    REQUIRED_INSTITUTIONAL_SHEETS,
    WorkbookReleaseBlockedError,
    audit_workbook,
)


def _full_spec():
    base = _quality_spec()
    action = _action_spec().action_plan
    chart = ChartSpec(
        code="nps_evolution",
        title="Evolucao do NPS",
        metric_code="academic.nps",
        dataset_code="academic_facts",
        dimension_code="period",
        chart_type=ChartType.LINE,
        role=ChartRole.EVOLUTION,
    )
    dashboard = replace(base.dashboard, charts=(chart,))
    matrix = MatrixSpec(
        metric_code="academic.nps",
        row_dimension="course",
        column_dimension="period",
        target=70,
        delta=True,
    )
    return replace(base, dashboard=dashboard, matrix=matrix, action_plan=action)


def test_01j_core_builds_complete_workbook_in_official_order():
    artifact = ExcelOfficialCore().build(_full_spec())
    wb = artifact.workbook
    assert tuple(wb.sheetnames[:6]) == REQUIRED_INSTITUTIONAL_SHEETS
    for name in REQUIRED_INSTITUTIONAL_SHEETS:
        assert wb[name].sheet_state == "visible"
        assert wb[name].protection.sheet is True
    assert artifact.audit.release_allowed is True
    assert artifact.refs.matrix is not None
    assert artifact.refs.action_plan is not None
    assert artifact.refs.charts.charts


def test_01j_core_materializes_required_technical_layers_visible_and_protected():
    artifact = ExcelOfficialCore().build(_full_spec())
    wb = artifact.workbook
    for name in ("METAS", "INDICADORES", "CALC", "LISTAS DE APOIO", "DIM_PERIODO"):
        assert name in wb.sheetnames
        assert wb[name].sheet_state == "visible"
        assert wb[name].protection.sheet is True
    assert wb["METAS"]["A1"].value == "REGISTRO DE METAS"
    assert wb["INDICADORES"]["A1"].value == "REGISTRO DE INDICADORES"


def test_01j_readme_explains_offline_snapshot_and_local_action_plan():
    artifact = ExcelOfficialCore().build(_full_spec())
    ws = artifact.workbook["LEIA-ME"]
    values = "\n".join(str(cell.value) for row in ws.iter_rows() for cell in row if cell.value is not None)
    assert "Snapshot offline" in values
    assert "nao consulta banco, API ou outro arquivo Excel" in values
    assert "nao sincronizam automaticamente" in values
    assert artifact.spec.snapshot.export_id in values


def test_01j_custom_metadata_contains_snapshot_and_contract_identity():
    artifact = ExcelOfficialCore().build(_full_spec())
    props = {item.name: item.value for item in artifact.workbook.custom_doc_props.props}
    assert props["DataUNIVC.ExportId"] == artifact.spec.snapshot.export_id
    assert props["DataUNIVC.DirectorateCode"] == "DCS"
    assert props["DataUNIVC.SystemVersion"] == "0.13.0"
    assert props["DataUNIVC.SchemaVersion"] == "49"
    assert props["DataUNIVC.ExcelContractVersion"] == 1
    assert props["DataUNIVC.AdapterCode"] == "academic"
    assert props["DataUNIVC.AdapterVersion"] == 1


def test_01j_matrix_none_and_disabled_action_plan_still_keep_institutional_contract():
    spec = replace(_spec(), action_plan=ActionPlanSpec(enabled=False), matrix=None)
    artifact = ExcelOfficialCore().build(spec)
    wb = artifact.workbook
    assert tuple(wb.sheetnames[:6]) == REQUIRED_INSTITUTIONAL_SHEETS
    assert wb["MATRIZ"]["A2"].value == "Nenhuma matriz foi configurada para este snapshot."
    assert wb["PLANO_DE_ACAO"]["A2"].value == "Plano de acao nao habilitado para este snapshot."
    assert artifact.refs.matrix is None
    assert artifact.refs.action_plan is None


def test_01j_blocking_quality_check_refuses_release():
    spec = _quality_spec(empty_facts=True)
    with pytest.raises(WorkbookReleaseBlockedError) as exc:
        ExcelOfficialCore().build(spec)
    assert any(item.code == "quality.facts_non_empty" for item in exc.value.audit.blocking_findings)


def test_01j_blocked_snapshot_can_be_inspected_without_release_for_diagnostics():
    spec = _quality_spec(empty_facts=True)
    artifact = ExcelOfficialCore().build(spec, enforce_release=False)
    assert artifact.audit.release_allowed is False
    assert artifact.audit.blocking_findings
    with pytest.raises(WorkbookReleaseBlockedError):
        artifact.save(Path("/tmp/should-not-exist.xlsx"))


def test_01j_audit_detects_unexpected_unlocked_cell():
    artifact = ExcelOfficialCore().build(_full_spec())
    ws = artifact.workbook["PAINEL"]
    ws["A1"].protection = Protection(locked=False)
    audit = audit_workbook(
        artifact.workbook,
        artifact.spec,
        parameter_system=artifact.refs.parameters,
        action_plan_ref=artifact.refs.action_plan,
    )
    assert any(item.code == "security.unexpected_unlocked_cell" for item in audit.blocking_findings)


def test_01j_audit_detects_unsupported_formula_before_release():
    artifact = ExcelOfficialCore().build(_full_spec())
    ws = artifact.workbook["CALC"]
    ws["Z1"] = "=FILTER(A1:A2,A1:A2<>\"\")"
    audit = audit_workbook(
        artifact.workbook,
        artifact.spec,
        parameter_system=artifact.refs.parameters,
        action_plan_ref=artifact.refs.action_plan,
    )
    assert any(item.code == "formula.unsupported_function" for item in audit.blocking_findings)


def test_01j_roundtrip_preserves_metadata_order_tables_charts_and_recalc(tmp_path: Path):
    artifact = ExcelOfficialCore().build(_full_spec())
    path = artifact.save(tmp_path / "excel_official_01j.xlsx")
    reopened = load_workbook(path, data_only=False)

    assert tuple(reopened.sheetnames[:6]) == REQUIRED_INSTITUTIONAL_SHEETS
    assert reopened.calculation.calcMode == "auto"
    assert reopened.calculation.fullCalcOnLoad is True
    assert reopened.calculation.forceFullCalc is True
    assert len(reopened["PAINEL"]._charts) == 1
    assert reopened["FATOS ACADEMICOS"].tables
    props = {item.name: item.value for item in reopened.custom_doc_props.props}
    assert props["DataUNIVC.ExportId"] == artifact.spec.snapshot.export_id
    assert getattr(reopened, "_external_links", []) == []
    assert reopened.vba_archive is None


def test_01j_all_non_input_cells_remain_locked():
    artifact = ExcelOfficialCore().build(_full_spec())
    allowed = {item.cell_reference for item in artifact.refs.parameters.parameters.values() if item.editable}
    ws = artifact.workbook["PARAMETROS"]
    unlocked = {cell.coordinate for row in ws.iter_rows() for cell in row if cell.protection.locked is False}
    assert unlocked == allowed
