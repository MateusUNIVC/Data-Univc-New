from __future__ import annotations

from dataclasses import replace
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from excel_official import (
    CellRole,
    ColumnDataType,
    ColumnSpec,
    DEFAULT_THEME,
    DatasetRef,
    DatasetSpec,
    DatasetWriteError,
    STYLE_NAMES,
    write_dataset_sheet,
    write_dataset_table,
)


def _dataset(*, technical: bool = False, rows=None) -> DatasetSpec:
    return DatasetSpec(
        code="sample",
        label="Base de Exemplo",
        sheet_name="BASE EXEMPLO" if not technical else "CALC",
        table_name="TblSample",
        source="Data UNIVC / fixture",
        technical=technical,
        columns=(
            ColumnSpec("period", "Periodo", ColumnDataType.TEXT, semantic_type="period", nullable=False),
            ColumnSpec("count", "Quantidade", ColumnDataType.INTEGER, semantic_type="count", nullable=False),
            ColumnSpec("rate", "Taxa", ColumnDataType.DECIMAL, semantic_type="percentage"),
            ColumnSpec("amount", "Valor", ColumnDataType.DECIMAL, semantic_type="currency_brl"),
            ColumnSpec("event_date", "Data", ColumnDataType.DATE),
            ColumnSpec("event_at", "Data Hora", ColumnDataType.DATETIME),
            ColumnSpec("active", "Ativo", ColumnDataType.BOOLEAN),
            ColumnSpec("formula_like", "Texto Literal", ColumnDataType.TEXT),
            ColumnSpec("tech", "Controle", ColumnDataType.TEXT, technical=True),
        ),
        rows=tuple(rows) if rows is not None else (
            {
                "period": "2026-SEM2",
                "count": 12,
                "rate": 0.875,
                "amount": Decimal("1250.50"),
                "event_date": date(2026, 10, 1),
                "event_at": datetime(2026, 10, 1, 19, 30),
                "active": True,
                "formula_like": "=1+1",
                "tech": "source-key",
            },
        ),
    )


def _rgb(color) -> str | None:
    value = getattr(color, "rgb", None)
    if isinstance(value, str):
        return value[-6:]
    return None


def test_dataset_ref_is_neutral_and_resolves_column_letters():
    ref = DatasetRef(
        dataset_code="x",
        sheet_name="BASE",
        table_name="TblX",
        table_ref="B4:D5",
        header_row=4,
        first_data_row=5,
        last_data_row=5,
        physical_last_row=5,
        first_column=2,
        last_column=4,
        row_count=1,
        column_letters={"a": "B", "b": "C", "c": "D"},
    )
    assert ref.column_letter("b") == "C"
    with pytest.raises(KeyError):
        ref.column_letter("missing")


def test_write_dataset_sheet_creates_standardized_title_table_and_ref():
    wb = Workbook()
    wb.remove(wb.active)
    dataset = _dataset()

    ref = write_dataset_sheet(wb, dataset)
    ws = wb[dataset.sheet_name]

    assert ref.dataset_code == dataset.code
    assert ref.sheet_name == dataset.sheet_name
    assert ref.table_name == dataset.table_name
    assert ref.header_row == 4
    assert ref.first_data_row == 5
    assert ref.last_data_row == 5
    assert ref.row_count == 1
    assert ref.table_ref == "A4:I5"
    assert ws["A1"].value == dataset.label
    assert ws["A2"].value == dataset.source
    assert ws.freeze_panes == "A5"
    assert dataset.table_name in ws.tables
    assert ws.tables[dataset.table_name].ref == "A4:I5"
    assert ws.protection.sheet is True
    assert ws.sheet_state == "visible"


def test_table_uses_theme_style_and_header_labels():
    wb = Workbook()
    wb.remove(wb.active)
    dataset = _dataset()
    write_dataset_sheet(wb, dataset)
    ws = wb[dataset.sheet_name]
    table = ws.tables[dataset.table_name]

    assert table.tableStyleInfo.name == DEFAULT_THEME.tables.style_name
    assert table.tableStyleInfo.showRowStripes is True
    assert [ws.cell(4, i).value for i in range(1, 10)] == [column.label for column in dataset.columns]
    assert all(ws.cell(4, i).style == STYLE_NAMES.table_header for i in range(1, 10))


def test_imported_and_technical_columns_use_semantic_styles():
    wb = Workbook()
    wb.remove(wb.active)
    dataset = _dataset()
    write_dataset_sheet(wb, dataset)
    ws = wb[dataset.sheet_name]

    assert ws["A5"].style == STYLE_NAMES.imported
    assert _rgb(ws["A5"].font.color) == DEFAULT_THEME.palette.imported_text
    assert ws["I5"].style == STYLE_NAMES.technical
    assert _rgb(ws["I5"].fill.fgColor) == DEFAULT_THEME.palette.gray_100
    assert ws["A5"].protection.locked is True
    assert ws["I5"].protection.locked is True


def test_technical_dataset_is_visible_protected_and_uses_technical_cells():
    wb = Workbook()
    wb.remove(wb.active)
    dataset = _dataset(technical=True)
    write_dataset_sheet(wb, dataset)
    ws = wb[dataset.sheet_name]

    assert ws.sheet_state == "visible"
    assert ws["A2"].value == "CAMADA TÉCNICA — NÃO ALTERAR"
    assert ws["A2"].style == STYLE_NAMES.technical
    assert ws["A5"].style == STYLE_NAMES.technical
    assert ws.protection.sheet is True


def test_number_formats_are_inferred_from_column_contract():
    wb = Workbook()
    ws = wb.active
    ref = write_dataset_table(ws, _dataset())

    assert ws[f"{ref.column_letter('count')}2"].number_format == DEFAULT_THEME.number_formats.integer
    assert ws[f"{ref.column_letter('rate')}2"].number_format == DEFAULT_THEME.number_formats.percentage
    assert ws[f"{ref.column_letter('amount')}2"].number_format == DEFAULT_THEME.number_formats.currency_brl
    assert ws[f"{ref.column_letter('event_date')}2"].number_format == DEFAULT_THEME.number_formats.date
    assert ws[f"{ref.column_letter('event_at')}2"].number_format == DEFAULT_THEME.number_formats.datetime


def test_explicit_number_format_overrides_semantic_default():
    dataset = _dataset()
    amount = replace(dataset.columns[3], number_format='0.0000')
    dataset = replace(dataset, columns=dataset.columns[:3] + (amount,) + dataset.columns[4:])
    wb = Workbook()
    ws = wb.active
    write_dataset_table(ws, dataset)
    assert ws["D2"].number_format == "0.0000"


def test_formula_like_imported_text_is_not_written_as_excel_formula():
    wb = Workbook()
    ws = wb.active
    write_dataset_table(ws, _dataset())
    cell = ws["H2"]
    assert cell.value == "=1+1"
    assert cell.data_type == "s"

    data = BytesIO()
    wb.save(data)
    data.seek(0)
    reopened = load_workbook(data, data_only=False)
    cell2 = reopened.active["H2"]
    assert cell2.value == "=1+1"
    assert cell2.data_type == "s"


def test_empty_snapshot_keeps_only_one_structural_table_row_but_reports_zero_facts():
    wb = Workbook()
    ws = wb.active
    dataset = _dataset(rows=())
    ref = write_dataset_table(ws, dataset)

    assert ref.row_count == 0
    assert ref.has_structural_empty_row is True
    assert ref.last_data_row == ref.header_row
    assert ref.physical_last_row == ref.first_data_row
    assert ref.table_ref == "A1:I2"
    assert all(ws.cell(2, col).value is None for col in range(1, 10))


def test_writer_refuses_to_overwrite_existing_cells():
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "do not overwrite"
    with pytest.raises(DatasetWriteError) as exc:
        write_dataset_table(ws, _dataset())
    assert exc.value.code == "dataset.target_not_empty"


def test_writer_refuses_sheet_name_collision_case_insensitive():
    wb = Workbook()
    wb.active.title = "base exemplo"
    with pytest.raises(DatasetWriteError) as exc:
        write_dataset_sheet(wb, _dataset())
    assert exc.value.code == "dataset.sheet_name_collision"


def test_writer_refuses_table_name_collision_workbook_wide_case_insensitive():
    wb = Workbook()
    first = replace(_dataset(), sheet_name="BASE 1")
    write_dataset_sheet(wb, first)
    second = replace(_dataset(), code="second", sheet_name="BASE 2", table_name="tblsample")
    before = tuple(wb.sheetnames)
    with pytest.raises(DatasetWriteError) as exc:
        write_dataset_sheet(wb, second)
    assert exc.value.code == "dataset.table_name_collision"
    assert tuple(wb.sheetnames) == before


def test_sheet_writer_rejects_invalid_rows_before_creating_sheet():
    wb = Workbook()
    wb.active.title = "HOME"
    dataset = _dataset(rows=({
        "period": "2026-SEM2",
        "count": "not-an-int",
        "rate": 0.5,
        "amount": 10.0,
        "event_date": None,
        "event_at": None,
        "active": True,
        "formula_like": "ok",
        "tech": "x",
    },))
    before = tuple(wb.sheetnames)
    with pytest.raises(DatasetWriteError) as exc:
        write_dataset_sheet(wb, dataset)
    assert exc.value.code == "dataset.invalid_value_type"
    assert tuple(wb.sheetnames) == before


def test_writer_rejects_wrong_python_type_before_mutating_cells():
    dataset = _dataset(rows=(
        {
            "period": "2026-SEM2",
            "count": "12",
            "rate": 0.5,
            "amount": 10.0,
            "event_date": None,
            "event_at": None,
            "active": True,
            "formula_like": "ok",
            "tech": "x",
        },
    ))
    wb = Workbook()
    ws = wb.active
    with pytest.raises(DatasetWriteError) as exc:
        write_dataset_table(ws, dataset)
    assert exc.value.code == "dataset.invalid_value_type"
    assert exc.value.column_code == "count"
    assert ws["A1"].value is None


def test_writer_rejects_null_in_non_nullable_column_before_mutation():
    dataset = _dataset(rows=(
        {
            "period": None,
            "count": 1,
            "rate": 0.5,
            "amount": 10.0,
            "event_date": None,
            "event_at": None,
            "active": True,
            "formula_like": "ok",
            "tech": "x",
        },
    ))
    wb = Workbook()
    ws = wb.active
    with pytest.raises(DatasetWriteError) as exc:
        write_dataset_table(ws, dataset)
    assert exc.value.code == "dataset.null_not_allowed"
    assert exc.value.column_code == "period"
    assert ws["A1"].value is None


def test_auto_width_respects_institutional_bounds():
    long_text = "X" * 200
    dataset = replace(
        _dataset(),
        rows=(
            {
                "period": long_text,
                "count": 1,
                "rate": 0.5,
                "amount": 10.0,
                "event_date": None,
                "event_at": None,
                "active": True,
                "formula_like": "short",
                "tech": "x",
            },
        ),
    )
    wb = Workbook()
    ws = wb.active
    write_dataset_table(ws, dataset)
    assert ws.column_dimensions["A"].width == DEFAULT_THEME.layout.table_max_column_width
    assert ws.column_dimensions["B"].width >= DEFAULT_THEME.layout.table_min_column_width


def test_dataset_writer_roundtrip_preserves_table_dates_styles_and_protection():
    wb = Workbook()
    wb.remove(wb.active)
    dataset = _dataset()
    write_dataset_sheet(wb, dataset)

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    reopened = load_workbook(stream)
    ws = reopened[dataset.sheet_name]

    assert dataset.table_name in ws.tables
    assert ws.tables[dataset.table_name].ref == "A4:I5"
    assert ws["E5"].value.date() == date(2026, 10, 1)
    assert ws["F5"].value == datetime(2026, 10, 1, 19, 30)
    assert ws["H5"].value == "=1+1"
    assert ws["H5"].data_type == "s"
    assert ws["A5"].style == STYLE_NAMES.imported
    assert ws.protection.sheet is True


def test_architecture_contract_layers_remain_openpyxl_independent_after_dataset_writer():
    root = Path(__file__).resolve().parents[1] / "excel_official"
    architecture_files = [
        root / "contract.py",
        root / "context.py",
        root / "validation.py",
        root / "constants.py",
        root / "exceptions.py",
        root / "theme.py",
        root / "adapters" / "base.py",
    ]
    source = "\n".join(path.read_text(encoding="utf-8") for path in architecture_files)
    assert "import openpyxl" not in source
    assert "from openpyxl" not in source
