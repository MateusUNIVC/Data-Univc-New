from __future__ import annotations

from dataclasses import replace
from pathlib import Path
import sys

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from test_excel_official_dashboard_v1 import _build, _spec  # noqa: E402

from excel_official import (  # noqa: E402
    LimitationSpec,
    QualityCheckSpec,
    QualitySeverity,
    QualitySpec,
    QualityWriteError,
    validate_workbook_spec,
    write_quality_sheet,
)


def _quality_spec(*, empty_facts: bool = False, with_limitation: bool = True):
    base = _spec(empty_facts=empty_facts)
    datasets = []
    for dataset in base.datasets:
        columns = dataset.columns
        if dataset.code == "academic_facts":
            columns = tuple(
                replace(column, nullable=False) if column.code in {"period", "course"} else column
                for column in dataset.columns
            )
        datasets.append(
            replace(
                dataset,
                columns=columns,
                source=f"source:{dataset.code}",
                authorization_scope=("DCS",),
            )
        )
    metrics = tuple(
        replace(metric, valid_min=-100, valid_max=100)
        if metric.code == "academic.nps"
        else metric
        for metric in base.metrics
    )
    quality = QualitySpec(
        dataset_checks=(
            QualityCheckSpec(
                code="facts_non_empty",
                label="Base academica possui registros",
                check_type="dataset_non_empty",
                source_code="academic_facts",
                severity=QualitySeverity.BLOCKING,
                message_ok="Base disponivel",
                message_error="Base academica vazia",
            ),
            QualityCheckSpec(
                code="facts_required",
                label="Campos obrigatorios preenchidos",
                check_type="dataset_required_fields",
                source_code="academic_facts",
                severity=QualitySeverity.ERROR,
                message_ok="Campos obrigatorios consistentes",
                message_error="Existem campos obrigatorios vazios",
            ),
        ),
        metric_checks=(
            QualityCheckSpec(
                code="nps_has_data",
                label="NPS possui dados no recorte",
                check_type="metric_has_data",
                source_code="academic.nps",
                severity=QualitySeverity.WARNING,
                message_ok="NPS calculavel",
                message_error="NPS sem dados no recorte atual",
            ),
            QualityCheckSpec(
                code="nps_range",
                label="NPS dentro da faixa valida",
                check_type="metric_valid_range",
                source_code="academic.nps",
                severity=QualitySeverity.ERROR,
                message_ok="Faixa valida",
                message_error="NPS fora da faixa -100 a 100",
            ),
        ),
        coverage_checks=(
            QualityCheckSpec(
                code="courses_available",
                label="Dimensao Curso disponivel",
                check_type="dimension_non_empty",
                source_code="course",
                severity=QualitySeverity.ERROR,
            ),
        ),
        snapshot_checks=(
            QualityCheckSpec(
                code="export_id_present",
                label="Export ID presente",
                check_type="snapshot_field_present",
                source_code="export_id",
                severity=QualitySeverity.BLOCKING,
            ),
            QualityCheckSpec(
                code="generated_by_present",
                label="Usuario exportador identificado",
                check_type="snapshot_field_present",
                source_code="generated_by",
                severity=QualitySeverity.WARNING,
            ),
        ),
        limitations=(
            LimitationSpec(
                code="detail_limit",
                title="Recorte offline limitado ao grao exportado",
                description="A planilha nao inventa granularidade ausente do snapshot.",
                severity=QualitySeverity.INFO,
                affected_metric="academic.nps",
            ),
        ) if with_limitation else (),
    )
    return replace(
        base,
        datasets=tuple(datasets),
        metrics=metrics,
        quality=quality,
        snapshot=replace(
            base.snapshot,
            minimum_period="2026-SEM1",
            maximum_period="2026-SEM2",
            payload_hash="sha256:test-quality",
        ),
    )


def _errors(spec) -> set[str]:
    return {item.code for item in validate_workbook_spec(spec) if item.severity == "ERROR"}


def test_01h_quality_contract_is_valid():
    assert validate_workbook_spec(_quality_spec()) == ()


def test_01h_quality_sheet_is_visible_protected_and_structured():
    spec = _quality_spec()
    wb, refs, params, _ = _build(spec)
    quality = write_quality_sheet(wb, spec, refs, params)
    ws = wb["QUALIDADE E GOVERNANCA"] if "QUALIDADE E GOVERNANCA" in wb.sheetnames else wb["QUALIDADE E GOVERNAN\u00c7A"]

    assert quality.sheet_name == "QUALIDADE E GOVERNAN\u00c7A"
    assert ws.sheet_state == "visible"
    assert ws.protection.sheet is True
    assert ws["A1"].value == "QUALIDADE E GOVERNAN\u00c7A"
    assert ws[quality.overall_status_cell].data_type == "f"
    assert len(quality.checks) == 7
    assert quality.coverage_range is not None
    assert quality.sources_range is not None
    assert quality.snapshot_range is not None
    assert quality.limitations_range is not None


def test_01h_dataset_checks_are_static_and_auditable():
    spec = _quality_spec()
    wb, refs, params, _ = _build(spec)
    quality = write_quality_sheet(wb, spec, refs, params)
    ws = wb[quality.sheet_name]
    by_code = {item.code: item for item in quality.checks}

    non_empty = by_code["facts_non_empty"]
    required = by_code["facts_required"]
    assert ws[non_empty.evidence_cell].value == 3
    assert ws[non_empty.result_cell].value == "OK"
    assert ws[required.evidence_cell].value == 0
    assert ws[required.result_cell].value == "OK"


def test_01h_metric_checks_reuse_metric_engine_and_current_parameters():
    spec = _quality_spec()
    wb, refs, params, _ = _build(spec)
    quality = write_quality_sheet(wb, spec, refs, params)
    ws = wb[quality.sheet_name]
    by_code = {item.code: item for item in quality.checks}

    has_data = by_code["nps_has_data"]
    range_check = by_code["nps_range"]
    evidence = str(ws[has_data.evidence_cell].value).upper()
    assert "SUMPRODUCT(" in evidence
    assert "P_REFERENCE_PERIOD" in evidence
    assert "P_COURSE" in evidence
    assert "AVERAGE(" not in evidence
    assert ws[has_data.result_cell].value == f'=IF({has_data.evidence_cell}="","WARNING","OK")'
    assert ">=-100" in str(ws[range_check.result_cell].value)
    assert "<=100" in str(ws[range_check.result_cell].value)


def test_01h_empty_dataset_is_blocking_and_metric_checks_do_not_fake_zero():
    spec = _quality_spec(empty_facts=True)
    wb, refs, params, _ = _build(spec)
    quality = write_quality_sheet(wb, spec, refs, params)
    ws = wb[quality.sheet_name]
    by_code = {item.code: item for item in quality.checks}

    assert ws[by_code["facts_non_empty"].evidence_cell].value == 0
    assert ws[by_code["facts_non_empty"].result_cell].value == "BLOCKING"
    assert ws[by_code["nps_has_data"].evidence_cell].value == '=""'
    assert "BLOCKING" in str(ws[quality.overall_status_cell].value)


def test_01h_sources_snapshot_coverage_and_limitations_are_materialized():
    spec = _quality_spec()
    wb, refs, params, _ = _build(spec)
    quality = write_quality_sheet(wb, spec, refs, params)
    ws = wb[quality.sheet_name]

    values = [cell.value for row in ws.iter_rows() for cell in row if cell.value is not None]
    joined = "\n".join(str(value) for value in values)
    assert "2026-SEM1" in joined
    assert "2026-SEM2" in joined
    assert "source:academic_facts" in joined
    assert "sha256:test-quality" in joined
    assert "0.13.0" in joined
    assert "Recorte offline limitado ao grao exportado" in joined
    assert "academic.nps" in joined


def test_01h_overall_status_uses_failed_check_results_and_actual_limitations():
    spec = _quality_spec()
    wb, refs, params, _ = _build(spec)
    quality = write_quality_sheet(wb, spec, refs, params)
    formula = str(wb[quality.sheet_name][quality.overall_status_cell].value)

    assert formula.startswith("=IF(")
    assert '"BLOCKING"' in formula
    assert '"ERROR"' in formula
    assert '"WARNING"' in formula
    assert '"INFO"' in formula
    assert '"NAO UTILIZAR"' in formula
    assert '"APTO COM OBSERVACOES"' in formula


def test_01h_no_declared_checks_or_limitations_is_explicit_not_false_green():
    spec = replace(_spec(), quality=QualitySpec())
    wb, refs, params, _ = _build(spec)
    quality = write_quality_sheet(wb, spec, refs, params)
    assert wb[quality.sheet_name][quality.overall_status_cell].value == "SEM CHECKS DECLARADOS"


def test_01h_preflight_failure_does_not_create_quality_sheet():
    spec = _quality_spec()
    wb, refs, params, _ = _build(spec)
    refs.pop("academic_facts")
    before = tuple(wb.sheetnames)

    with pytest.raises(QualityWriteError) as exc:
        write_quality_sheet(wb, spec, refs, params)
    assert exc.value.code == "quality.missing_dataset_ref"
    assert tuple(wb.sheetnames) == before
    assert "QUALIDADE E GOVERNAN\u00c7A" not in wb.sheetnames


def test_01h_contract_rejects_unsupported_check_unknown_source_and_group_mismatch():
    base = _quality_spec()
    unsupported = replace(
        base,
        quality=replace(
            base.quality,
            dataset_checks=(
                replace(base.quality.dataset_checks[0], check_type="made_up"),
            ),
        ),
    )
    assert "quality.unsupported_check_type" in _errors(unsupported)

    unknown = replace(
        base,
        quality=replace(
            base.quality,
            coverage_checks=(
                replace(base.quality.coverage_checks[0], source_code="missing_dimension"),
            ),
        ),
    )
    assert "quality.unknown_dimension" in _errors(unknown)

    mismatch = replace(
        base,
        quality=replace(
            base.quality,
            dataset_checks=(
                replace(base.quality.dataset_checks[0], check_type="dimension_non_empty", source_code="course"),
            ),
        ),
    )
    assert "quality.check_group_mismatch" in _errors(mismatch)


def test_01h_contract_rejects_metric_range_without_declared_bounds():
    base = _quality_spec()
    metrics = tuple(
        replace(metric, valid_min=None, valid_max=None) if metric.code == "academic.nps" else metric
        for metric in base.metrics
    )
    invalid = replace(base, metrics=metrics)
    assert "quality.metric_range_required" in _errors(invalid)


def test_01h_roundtrip_preserves_formulas_protection_and_visibility(tmp_path: Path):
    spec = _quality_spec()
    wb, refs, params, _ = _build(spec)
    quality = write_quality_sheet(wb, spec, refs, params)
    path = tmp_path / "quality.xlsx"
    wb.save(path)

    reopened = load_workbook(path, data_only=False)
    ws = reopened[quality.sheet_name]
    assert ws.sheet_state == "visible"
    assert ws.protection.sheet is True
    assert ws[quality.overall_status_cell].data_type == "f"
    assert reopened.calculation.calcMode == "auto"
    assert reopened.calculation.fullCalcOnLoad is True
    assert reopened.calculation.forceFullCalc is True


def test_01h_formulas_remain_excel_2019_compatible():
    spec = _quality_spec()
    wb, refs, params, _ = _build(spec)
    quality = write_quality_sheet(wb, spec, refs, params)
    formulas = "\n".join(
        str(cell.value).upper()
        for row in wb[quality.sheet_name].iter_rows()
        for cell in row
        if cell.data_type == "f"
    )
    for unsupported in ("FILTER(", "XLOOKUP(", "SORT(", "UNIQUE(", "INDIRECT(", "OFFSET("):
        assert unsupported not in formulas
