from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook

from excel_official import (
    CellRole,
    DEFAULT_THEME,
    STYLE_NAMES,
    apply_cell_role,
    apply_sheet_defaults,
    protect_sheet,
    register_named_styles,
    unlock_cell,
)


def _rgb(cell_color) -> str | None:
    value = getattr(cell_color, "rgb", None)
    if isinstance(value, str):
        return value[-6:]
    return None


def test_default_theme_freezes_the_univc_golden_master_palette_and_font():
    theme = DEFAULT_THEME
    assert theme.name == "Data UNIVC Excel Official V1"
    assert theme.typography.family == "Arial"
    assert theme.typography.title_size == 15.0
    assert theme.palette.green_dark == "045233"
    assert theme.palette.green == "0B7A54"
    assert theme.palette.input_fill == "FFF6DC"
    assert theme.palette.error_fill == "FDE7E7"
    assert theme.palette.control_text == "7030A0"


def test_table_and_chart_visual_tokens_are_centralized():
    assert DEFAULT_THEME.tables.style_name == "TableStyleMedium4"
    assert DEFAULT_THEME.tables.show_row_stripes is True
    assert DEFAULT_THEME.charts.series_colors[0] == DEFAULT_THEME.palette.green
    assert len(DEFAULT_THEME.charts.series_colors) >= 6


def test_register_named_styles_is_idempotent():
    wb = Workbook()
    first = register_named_styles(wb)
    second = register_named_styles(wb)
    assert first == second
    assert len(first) == len(set(first))
    assert set(first).issubset(set(wb.named_styles))


def test_input_style_is_visually_distinct_and_unlocked():
    wb = Workbook()
    ws = wb.active
    cell = apply_cell_role(ws["A1"], CellRole.INPUT)
    assert cell.style == STYLE_NAMES.input
    assert _rgb(cell.fill.fgColor) == DEFAULT_THEME.palette.input_fill
    assert _rgb(cell.font.color) == DEFAULT_THEME.palette.input_text
    assert cell.font.bold is True
    assert cell.protection.locked is False


def test_imported_derived_static_and_control_roles_have_distinct_semantics():
    wb = Workbook()
    ws = wb.active
    imported = apply_cell_role(ws["A1"], CellRole.IMPORTED)
    derived = apply_cell_role(ws["A2"], CellRole.DERIVED)
    static = apply_cell_role(ws["A3"], CellRole.STATIC)
    control = apply_cell_role(ws["A4"], CellRole.CONTROL)

    assert _rgb(imported.font.color) == DEFAULT_THEME.palette.imported_text
    assert _rgb(derived.font.color) == DEFAULT_THEME.palette.derived_text
    assert _rgb(static.font.color) == DEFAULT_THEME.palette.static_text
    assert _rgb(control.font.color) == DEFAULT_THEME.palette.control_text
    assert _rgb(control.fill.fgColor) == DEFAULT_THEME.palette.gray_50
    assert all(cell.protection.locked for cell in (imported, derived, static, control))


def test_technical_cells_are_visible_style_but_locked():
    wb = Workbook()
    ws = wb.active
    cell = apply_cell_role(ws["A1"], CellRole.TECHNICAL)
    assert _rgb(cell.fill.fgColor) == DEFAULT_THEME.palette.gray_100
    assert _rgb(cell.font.color) == DEFAULT_THEME.palette.static_text
    assert cell.protection.locked is True


def test_title_and_section_styles_follow_golden_master_language():
    wb = Workbook()
    ws = wb.active
    title = apply_cell_role(ws["A1"], CellRole.TITLE)
    section = apply_cell_role(ws["A2"], CellRole.SECTION)

    assert title.font.name == "Arial"
    assert title.font.sz == 15.0
    assert title.font.bold is True
    assert _rgb(title.font.color) == DEFAULT_THEME.palette.green_dark
    assert title.border.bottom.style == "medium"

    assert _rgb(section.fill.fgColor) == DEFAULT_THEME.palette.green_dark
    assert _rgb(section.font.color) == DEFAULT_THEME.palette.white
    assert section.font.bold is True


def test_semantic_status_styles_are_not_ambiguous():
    wb = Workbook()
    ws = wb.active
    success = apply_cell_role(ws["A1"], CellRole.SUCCESS)
    warning = apply_cell_role(ws["A2"], CellRole.WARNING)
    error = apply_cell_role(ws["A3"], CellRole.ERROR)

    assert _rgb(success.fill.fgColor) == DEFAULT_THEME.palette.green_pale
    assert _rgb(warning.fill.fgColor) == DEFAULT_THEME.palette.warning_fill
    assert _rgb(error.fill.fgColor) == DEFAULT_THEME.palette.error_fill
    assert len({_rgb(success.fill.fgColor), _rgb(warning.fill.fgColor), _rgb(error.fill.fgColor)}) == 3


def test_kpi_styles_share_institutional_card_language_without_layout_coordinates():
    wb = Workbook()
    ws = wb.active
    label = apply_cell_role(ws["A1"], CellRole.KPI_LABEL)
    value = apply_cell_role(ws["A2"], CellRole.KPI_VALUE, number_format=DEFAULT_THEME.number_formats.decimal_1)

    assert _rgb(label.fill.fgColor) == DEFAULT_THEME.palette.green_light
    assert _rgb(value.fill.fgColor) == DEFAULT_THEME.palette.green_light
    assert _rgb(value.font.color) == DEFAULT_THEME.palette.green_dark
    assert value.font.sz == DEFAULT_THEME.typography.kpi_value_size
    assert value.number_format == DEFAULT_THEME.number_formats.decimal_1


def test_sheet_defaults_are_centralized():
    wb = Workbook()
    ws = wb.active
    apply_sheet_defaults(ws, freeze_panes="A5", tab_color=DEFAULT_THEME.palette.green)

    assert ws.sheet_view.showGridLines is False
    assert ws.sheet_view.zoomScale == DEFAULT_THEME.layout.default_zoom
    assert ws.freeze_panes == "A5"
    assert ws.page_setup.fitToWidth == 1
    assert ws.page_setup.fitToHeight == 0
    assert ws.page_margins.left == DEFAULT_THEME.layout.margin_left
    assert _rgb(ws.sheet_properties.tabColor) == DEFAULT_THEME.palette.green


def test_sheet_protection_uses_unlocked_cells_as_the_editable_boundary():
    wb = Workbook()
    ws = wb.active
    apply_cell_role(ws["A1"], CellRole.DERIVED)
    apply_cell_role(ws["A2"], CellRole.INPUT)
    unlock_cell(ws["A3"])
    protect_sheet(ws)

    assert ws.protection.sheet is True
    assert ws["A1"].protection.locked is True
    assert ws["A2"].protection.locked is False
    assert ws["A3"].protection.locked is False


def test_styles_survive_xlsx_roundtrip():
    wb = Workbook()
    ws = wb.active
    apply_cell_role(ws["A1"], CellRole.INPUT)
    apply_cell_role(ws["A2"], CellRole.TECHNICAL)
    apply_sheet_defaults(ws)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    reopened = load_workbook(output)
    ws2 = reopened.active

    assert _rgb(ws2["A1"].fill.fgColor) == DEFAULT_THEME.palette.input_fill
    assert ws2["A1"].protection.locked is False
    assert _rgb(ws2["A2"].fill.fgColor) == DEFAULT_THEME.palette.gray_100
    assert ws2.sheet_view.showGridLines is False


def test_contract_and_adapter_layers_remain_openpyxl_independent():
    root = Path(__file__).resolve().parents[1] / "excel_official"
    architecture_files = [
        root / "contract.py",
        root / "context.py",
        root / "validation.py",
        root / "constants.py",
        root / "exceptions.py",
        root / "theme.py",
        root / "adapters" / "base.py",
    ]
    source = "\n".join(path.read_text(encoding="utf-8") for path in architecture_files)
    assert "import openpyxl" not in source
    assert "from openpyxl" not in source
