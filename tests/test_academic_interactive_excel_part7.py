from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from academic_excel_v3_builder import build_academic_interactive_workbook


def _payload(directorate: str = "DTNH") -> dict:
    codes = {
        "nps_institution": f"{directorate}-01A",
        "nps_course": f"{directorate}-01B",
        "nps_faculty": f"{directorate}-01C",
        "teacher": f"{directorate}-02",
        "approval": f"{directorate}-03",
    }
    long_course = "Comunicação Social - Publicidade e Propaganda"
    return {
        "directorate": directorate,
        "generated_at": "2026-09-29T18:00:00+00:00",
        "app_version": "0.13.0",
        "codes": codes,
        "courses_catalog": [
            {"curso": "Administração", "ativo": True, "id": 1},
            {"curso": long_course, "ativo": True, "id": 2},
        ],
        "disciplines_catalog": [
            {"curso": "Administração", "disciplina": "Gestão", "ativo": True, "id": 1},
            {"curso": long_course, "disciplina": "Criação Publicitária", "ativo": True, "id": 2},
        ],
        "semesters": ["2025-SEM2", "2026-SEM1", "2026-SEM2"],
        "initial": {
            "reference": "2026-SEM2",
            "comparison": "2026-SEM1",
            "window": 6,
            "course": "(todos)",
            "discipline": "(todas)",
            "matrix_kpi": "01A · NPS Instituição · Alunos",
        },
        "institution_students_by_course": [
            {"periodo": "2026-SEM1", "curso": "Administração", "respondentes": 100, "promotores": 60, "neutros": 25, "detratores": 15, "valor": 45.0, "meta": 40, "atencao": 35, "status": "Dentro da meta", "fonte": "SEI"},
            {"periodo": "2026-SEM2", "curso": long_course, "respondentes": 80, "promotores": 45, "neutros": 20, "detratores": 15, "valor": 37.5, "meta": 45, "atencao": 40, "status": "Fora da meta", "fonte": "SEI"},
        ],
        "institution_students": [
            {"periodo": "2026-SEM1", "respondentes": 180, "promotores": 105, "neutros": 45, "detratores": 30, "valor": 41.7, "coverage": 2, "coverage_total": 2, "complete": True},
            {"periodo": "2026-SEM2", "respondentes": 220, "promotores": 135, "neutros": 50, "detratores": 35, "valor": 45.5, "coverage": 2, "coverage_total": 2, "complete": True},
        ],
        "course_nps": [
            {"periodo": "2026-SEM2", "curso": "Administração", "respondentes": 110, "promotores": 70, "neutros": 25, "detratores": 15, "meta": 45, "atencao": 40, "status": "Dentro da meta", "fonte": "Questionario"},
        ],
        "faculty_nps": [
            {"periodo": "2026-SEM2", "respondentes": 60, "promotores": 40, "neutros": 12, "detratores": 8, "valor": 53.3, "fonte": "Questionario"},
        ],
        "nps_distribution_rows": [
            {"periodo": "2026-SEM2", "audiencia": "NPS do Curso · Alunos", "escopo": "DTNH", "nota": 6, "respostas": 12, "percentual": 20.0, "respondentes": 60, "media": 7.1, "nps": 35.0},
            {"periodo": "2026-SEM2", "audiencia": "NPS Instituição · Docentes", "escopo": "Todos os docentes", "nota": 10, "respostas": 18, "percentual": 30.0, "respondentes": 60, "media": 7.6, "nps": 53.3},
        ],
        "teacher": [
            {"periodo": "2026-SEM2", "curso": "Administração", "disciplina": "Gestão", "respondentes": 32, "favoraveis": 28, "intermediarias": 3, "desfavoraveis": 1, "classificados": 32, "nao_classificados": 0, "nao_mapeados": 0, "favorabilidade": 87.5, "meta": 80, "atencao": 75, "status": "Dentro da meta"},
        ],
        "results": [
            {"periodo": "2026-SEM2", "curso": "Administração", "disciplina": "Gestão", "total_registros": 110, "finalizados": 100, "aprovados": 90, "reprovados_nota": 5, "reprovados_falta": 3, "reprovados_outro": 2, "em_andamento": 10, "meta": 85, "atencao": 80, "status": "Dentro da meta"},
        ],
        "effective_goals": [],
        "goals": [
            {"indicador": codes["nps_institution"], "recorte": "TOTAL", "vigencia": "2026-SEM2", "meta": 45, "atencao": 40, "limite_superior": None, "unidade": "pontos", "metric_version": "vigente", "justificativa": "Meta"},
        ],
        "actions": [
            {"indicador": codes["nps_institution"], "recorte": "TOTAL", "acao": "Acompanhar NPS", "responsavel": directorate, "prazo": "2026-12-31", "status": "Em andamento", "resultado_esperado": "NPS >= 45"},
        ],
        "quality": {"nps_inconsistent_rows": 0, "result_inconsistent_rows": 0, "periods": 3, "courses": 2, "disciplines": 2},
    }


def _wb(directorate: str = "DTNH"):
    stream = build_academic_interactive_workbook(_payload(directorate))
    return load_workbook(BytesIO(stream.getvalue()), data_only=False)


def test_part7_promotes_interactive_excel_to_official_database_workbook():
    wb = _wb()
    expected_visible = [
        "LEIA-ME", "PARAMETROS", "PAINEL", "QUALIDADE E GOVERNANCA", "MATRIZ", "PLANO_DE_ACAO",
        "NPS DISCENTES", "NPS SEMESTRAL", "NPS DOCENTES", "NPS DISTRIBUICAO", "AVALIACAO DOCENTE", "RESULTADOS ACADEMICOS",
        "METAS", "CURSOS", "DISCIPLINAS", "INDICADORES", "DIM_PERIODO", "DIM_MES",
    ]
    assert [ws.title for ws in wb.worksheets if ws.sheet_state == "visible"] == expected_visible
    assert wb["CALC"].sheet_state == "hidden"
    assert wb["LISTAS DE APOIO"].sheet_state == "hidden"
    assert "beta" not in (wb.properties.description or "").lower()
    assert "oficial" in (wb.properties.description or "").lower()

    expected_tables = {
        "NPS DISCENTES": "TblNpsDiscentes",
        "NPS SEMESTRAL": "TblNpsSemestral",
        "NPS DOCENTES": "TblNpsDocentes",
        "NPS DISTRIBUICAO": "TblNpsDistribuicao",
        "AVALIACAO DOCENTE": "TblAvaliacaoDocente",
        "RESULTADOS ACADEMICOS": "TblResultadosAcademicos",
        "METAS": "TblMetas",
        "CURSOS": "TblCursos",
        "DISCIPLINAS": "TblDisciplinas",
        "INDICADORES": "TblIndicadores",
        "DIM_PERIODO": "TblDimPeriodo",
        "DIM_MES": "TblDimMes",
        "PLANO_DE_ACAO": "TblPlanoAcao",
    }
    for sheet, table in expected_tables.items():
        assert table in wb[sheet].tables
    wb.close()


def test_part7_panel_has_five_executive_charts_and_long_course_labels_are_wrapped():
    wb = _wb()
    assert len(wb["PAINEL"]._charts) == 5
    long_label = wb["CALC"]["J32"].value
    assert "Comunicação Social" in long_label
    assert "Publicidade e Propaganda" in long_label
    assert "\n" in long_label
    assert "…" not in long_label
    wb.close()


def test_part7_keeps_current_faculty_favorability_metric_not_legacy_score():
    wb = _wb()
    ws = wb["INDICADORES"]
    rows = [[ws.cell(r, c).value for c in range(1, 8)] for r in range(5, ws.max_row + 1)]
    teacher = next(row for row in rows if row[0] == "DTNH-02")
    assert teacher[2] == "%"
    assert "Favoráveis / respostas classificadas" in teacher[6]
    assert "0–10" not in " ".join(str(v or "") for v in teacher)
    wb.close()


def test_part7_dcs_uses_same_official_structure():
    wb = _wb("DCS")
    assert wb.properties.description == "Excel Interativo oficial · DCS"
    assert "DCS-02" == wb["INDICADORES"]["A8"].value
    assert len(wb["PAINEL"]._charts) == 5
    wb.close()


def test_part7_ui_and_download_name_no_longer_advertise_beta():
    root = Path(__file__).resolve().parents[1]
    html = (root / "templates" / "index.html").read_text(encoding="utf-8")
    app = (root / "app.py").read_text(encoding="utf-8")
    js = (root / "static" / "js" / "app.js").read_text(encoding="utf-8")
    assert "Excel Interativo (beta)" not in html
    assert "Gerar beta" not in html
    assert "Interativo_beta.xlsx" not in app
    assert "Interativo_beta.xlsx" not in js
    assert 'Painel_{scope.directorate_code}_Interativo.xlsx' in app
