from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from academic_excel_official import build_academic_excel_official_artifact_from_payload
from excel_official.adapters.academic import ACADEMIC_METRICS, AcademicAdapter
from excel_official.audit.quality import evaluate_metric
from excel_official.constants import REQUIRED_INSTITUTIONAL_SHEETS


def _payload(directorate: str = "DTNH", *, teacher_unmapped: int = 0) -> dict:
    codes = {
        "nps_institution": f"{directorate}-01A",
        "nps_course": f"{directorate}-01B",
        "nps_faculty": f"{directorate}-01C",
        "teacher": f"{directorate}-02",
        "approval": f"{directorate}-03",
    }
    return {
        "directorate": directorate,
        "directorate_name": "Diretoria Acadêmica de Teste",
        "generated_at": "2026-10-02T11:00:00+00:00",
        "app_version": "0.13.0",
        "codes": codes,
        "courses_catalog": [
            {"curso": "Administração", "ativo": True, "id": 1},
            {"curso": "Psicologia", "ativo": True, "id": 2},
        ],
        "disciplines_catalog": [
            {"curso": "Administração", "disciplina": "Gestão", "ativo": True, "id": 11},
            {"curso": "Psicologia", "disciplina": "Psicologia Social", "ativo": True, "id": 21},
        ],
        "semesters": ["2026-SEM1", "2026-SEM2"],
        "initial": {
            "reference": "2026-SEM2",
            "comparison": "2026-SEM1",
            "window": 6,
            "course": "(todos)",
            "discipline": "(todas)",
        },
        "institution_students": [
            {"periodo": "2026-SEM1", "respondentes": 180, "promotores": 105, "neutros": 45, "detratores": 30, "valor": 41.6667, "coverage": 2, "coverage_total": 2, "complete": True},
            {"periodo": "2026-SEM2", "respondentes": 220, "promotores": 135, "neutros": 50, "detratores": 35, "valor": 45.4545, "coverage": 2, "coverage_total": 2, "complete": True},
        ],
        "institution_students_by_course": [],
        "course_nps": [
            {"periodo": "2026-SEM1", "curso": "Administração", "respondentes": 80, "promotores": 48, "neutros": 20, "detratores": 12, "valor": 45.0, "fonte": "SEI"},
            {"periodo": "2026-SEM1", "curso": "Psicologia", "respondentes": 100, "promotores": 57, "neutros": 25, "detratores": 18, "valor": 39.0, "fonte": "SEI"},
            {"periodo": "2026-SEM2", "curso": "Administração", "respondentes": 110, "promotores": 70, "neutros": 25, "detratores": 15, "valor": 50.0, "fonte": "SEI"},
            {"periodo": "2026-SEM2", "curso": "Psicologia", "respondentes": 90, "promotores": 50, "neutros": 20, "detratores": 20, "valor": 33.3333, "fonte": "SEI"},
        ],
        "faculty_nps": [
            {"periodo": "2026-SEM1", "respondentes": 50, "promotores": 30, "neutros": 12, "detratores": 8, "valor": 44.0, "fonte": "SEI"},
            {"periodo": "2026-SEM2", "respondentes": 60, "promotores": 40, "neutros": 12, "detratores": 8, "valor": 53.3333, "fonte": "SEI"},
        ],
        "nps_distribution_rows": [
            {"periodo": "2026-SEM2", "audiencia": "NPS do Curso · Alunos", "escopo": directorate, "nota": 10, "respostas": 30, "percentual": 15.0, "respondentes": 200, "media": 7.4, "nps": 42.5},
        ],
        "teacher": [
            {"periodo": "2026-SEM1", "curso": "Administração", "disciplina": "Gestão", "respondentes": 20, "favoraveis": 16, "intermediarias": 3, "desfavoraveis": 1, "classificados": 20, "nao_classificados": 0, "nao_mapeados": 0, "favorabilidade": 80.0},
            {"periodo": "2026-SEM1", "curso": "Psicologia", "disciplina": "Psicologia Social", "respondentes": 20, "favoraveis": 17, "intermediarias": 2, "desfavoraveis": 1, "classificados": 20, "nao_classificados": 0, "nao_mapeados": 0, "favorabilidade": 85.0},
            {"periodo": "2026-SEM2", "curso": "Administração", "disciplina": "Gestão", "respondentes": 32, "favoraveis": 28, "intermediarias": 3, "desfavoraveis": 1, "classificados": 32, "nao_classificados": 0, "nao_mapeados": teacher_unmapped, "favorabilidade": None if teacher_unmapped else 87.5},
            {"periodo": "2026-SEM2", "curso": "Psicologia", "disciplina": "Psicologia Social", "respondentes": 18, "favoraveis": 14, "intermediarias": 3, "desfavoraveis": 1, "classificados": 18, "nao_classificados": 0, "nao_mapeados": 0, "favorabilidade": 77.7778},
        ],
        "results": [
            {"periodo": "2026-SEM1", "curso": "Administração", "disciplina": "Gestão", "total_registros": 90, "finalizados": 80, "aprovados": 68, "reprovados_nota": 6, "reprovados_falta": 4, "reprovados_outro": 2, "em_andamento": 10, "notas_contagem": 80, "soma_notas": 620.0, "media_notas": 7.75},
            {"periodo": "2026-SEM1", "curso": "Psicologia", "disciplina": "Psicologia Social", "total_registros": 75, "finalizados": 70, "aprovados": 60, "reprovados_nota": 5, "reprovados_falta": 3, "reprovados_outro": 2, "em_andamento": 5, "notas_contagem": 70, "soma_notas": 560.0, "media_notas": 8.0},
            {"periodo": "2026-SEM2", "curso": "Administração", "disciplina": "Gestão", "total_registros": 110, "finalizados": 100, "aprovados": 90, "reprovados_nota": 5, "reprovados_falta": 3, "reprovados_outro": 2, "em_andamento": 10, "notas_contagem": 100, "soma_notas": 805.0, "media_notas": 8.05},
            {"periodo": "2026-SEM2", "curso": "Psicologia", "disciplina": "Psicologia Social", "total_registros": 100, "finalizados": 90, "aprovados": 72, "reprovados_nota": 8, "reprovados_falta": 6, "reprovados_outro": 4, "em_andamento": 10, "notas_contagem": 90, "soma_notas": 711.0, "media_notas": 7.9},
        ],
        "effective_goals": [
            {"periodo": "2026-SEM2", "kpi": codes["nps_course"], "curso": "(todos)", "disciplina": "(todas)", "meta": 45.0, "atencao": 40.0, "vigencia": "2026-SEM2", "recorte": "TOTAL"},
        ],
        "goals": [],
        "actions": [
            {"indicador": codes["nps_course"], "acao": "Acompanhar NPS", "responsavel": directorate, "prazo": "2026-12-31", "status": "Em andamento", "resultado_esperado": "Atingir meta"},
        ],
        "quality": {"nps_inconsistent_rows": 0, "result_inconsistent_rows": 0, "periods": 2, "courses": 2, "disciplines": 2},
    }


def _artifact(directorate: str = "DTNH", *, teacher_unmapped: int = 0):
    return build_academic_excel_official_artifact_from_payload(
        _payload(directorate, teacher_unmapped=teacher_unmapped),
        generated_by="qa@univc.br",
        export_id=f"qa-{directorate.lower()}-academic",
    )


def test_academic_adapter_is_declarative_and_supports_dtnh_and_dcs():
    source = Path(__file__).resolve().parents[1] / "excel_official" / "adapters" / "academic.py"
    text = source.read_text(encoding="utf-8")
    assert "from openpyxl" not in text
    assert "DatabaseRepository" not in text
    assert "SurveyRepository" not in text
    assert AcademicAdapter.adapter_code == "academic"

    for directorate in ("DTNH", "DCS"):
        artifact = _artifact(directorate)
        assert artifact.spec.identity.directorate_code == directorate
        assert artifact.spec.identity.adapter_code == "academic"
        assert artifact.audit.release_allowed is True


def test_academic_metric_bindings_preserve_official_scope_semantics():
    spec = _artifact().spec
    bindings = {item.metric_code: item for item in spec.metric_bindings}
    assert bindings[ACADEMIC_METRICS["nps_institution"]].filter_parameters == {"period": "reference_period"}
    assert bindings[ACADEMIC_METRICS["nps_faculty"]].filter_parameters == {"period": "reference_period"}
    assert bindings[ACADEMIC_METRICS["nps_course"]].filter_parameters == {"period": "reference_period", "course": "course"}
    assert bindings[ACADEMIC_METRICS["teacher"]].filter_parameters == {
        "period": "reference_period",
        "course": "course",
        "discipline": "discipline",
    }


def test_academic_metrics_recompose_from_components_without_average_of_averages():
    spec = _artifact().spec
    assert evaluate_metric(spec, ACADEMIC_METRICS["nps_institution"]) == pytest.approx((135 / 220 - 35 / 220) * 100)
    assert evaluate_metric(spec, ACADEMIC_METRICS["nps_course"]) == pytest.approx((120 / 200 - 35 / 200) * 100)
    assert evaluate_metric(spec, ACADEMIC_METRICS["nps_faculty"]) == pytest.approx((40 / 60 - 8 / 60) * 100)
    assert evaluate_metric(spec, ACADEMIC_METRICS["teacher"]) == pytest.approx(42 / 50)
    assert evaluate_metric(spec, ACADEMIC_METRICS["approval"]) == pytest.approx(162 / 190)
    assert evaluate_metric(spec, ACADEMIC_METRICS["average_grade"]) == pytest.approx(1516 / 190)


def test_faculty_favorability_is_unavailable_when_any_unmapped_response_exists():
    artifact = _artifact(teacher_unmapped=2)
    spec = artifact.spec
    assert evaluate_metric(spec, ACADEMIC_METRICS["teacher"]) is None
    metric = next(item for item in spec.metrics if item.code == ACADEMIC_METRICS["teacher"])
    assert metric.invalid_when_positive_components == ("unmapped",)

    formulas = [
        cell.value
        for row in artifact.workbook["PAINEL"].iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    ]
    assert any("unmapped" not in formula and "SUMPRODUCT" in formula and ">0" in formula for formula in formulas)


def test_academic_workbook_has_contract_v1_structure_and_visible_technical_layers():
    artifact = _artifact()
    wb = artifact.workbook
    assert tuple(wb.sheetnames[:6]) == REQUIRED_INSTITUTIONAL_SHEETS
    assert wb["QUALIDADE E GOVERNANÇA"].sheet_state == "visible"
    assert wb["CALC"].sheet_state == "visible"
    assert wb["LISTAS DE APOIO"].sheet_state == "visible"
    assert wb["CALC"].protection.sheet is True
    assert wb["LISTAS DE APOIO"].protection.sheet is True
    assert len(wb["PAINEL"]._charts) == 5
    for sheet in (
        "NPS INSTITUICAO",
        "NPS CURSO",
        "NPS DOCENTES",
        "NPS DISTRIBUICAO",
        "AVALIACAO DOCENTE",
        "RESULTADOS ACADEMICOS",
    ):
        assert sheet in wb.sheetnames


def test_academic_action_plan_keeps_official_rows_locked_and_local_tracking_editable():
    artifact = _artifact()
    ws = artifact.workbook["PLANO_DE_ACAO"]
    # The first table row is official; subsequent preallocated rows are local.
    assert ws["A7"].value == "OFICIAL"
    assert ws["D7"].protection.locked is True
    local_rows = [row for row in range(8, ws.max_row + 1) if ws.cell(row, 1).value == "LOCAL"]
    assert local_rows
    first_local = local_rows[0]
    assert ws.cell(first_local, 4).protection.locked is False


def test_academic_workbook_round_trip_and_release_audit():
    artifact = _artifact()
    stream = artifact.to_bytes()
    wb = load_workbook(BytesIO(stream.getvalue()), data_only=False)
    assert tuple(wb.sheetnames[:6]) == REQUIRED_INSTITUTIONAL_SHEETS
    assert wb.properties.title
    assert len(wb["PAINEL"]._charts) == 5
    assert not getattr(wb, "vba_archive", None)
    wb.close()


def test_academic_adapter_exports_grade_components_required_for_offline_average():
    spec = _artifact().spec
    results = next(item for item in spec.datasets if item.code == "academic_results")
    columns = {item.code for item in results.columns}
    assert {"grade_count", "grade_sum", "reported_average_grade"} <= columns
    assert sum(row["grade_count"] for row in results.rows if row["period"] == "2026-SEM2") == 190
    assert sum(row["grade_sum"] for row in results.rows if row["period"] == "2026-SEM2") == pytest.approx(1516.0)


def test_academic_production_route_defaults_to_excel_official():
    root = Path(__file__).resolve().parents[1]
    service = (root / "excel_service.py").read_text(encoding="utf-8")
    assert "build_academic_interactive_workbook_bytes" in service  # explicit rollback remains available
    assert "ACADEMIC_EXCEL_OFFICIAL_ENABLED" in service
    assert 'os.getenv("ACADEMIC_EXCEL_OFFICIAL_ENABLED", "true")' in service
    assert "build_academic_excel_official_workbook_bytes" in service
