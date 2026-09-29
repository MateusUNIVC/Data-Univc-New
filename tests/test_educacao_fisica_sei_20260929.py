from __future__ import annotations

import tempfile
import unittest
import uuid
from pathlib import Path

from openpyxl import Workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker

from academic_catalog import course_name_matches, course_sei_name_matches
from database import Base
from models import AcademicResult, Course, Directorate
from repository import DatabaseRepository
from schemas import ValidationError
from security import AuthorizationContext
from sei_academic import SEIBot, validar_xlsx_baixado


class EducacaoFisicaSei20260929Tests(unittest.TestCase):
    @staticmethod
    def _course_row(index: int, course: str, turn: str) -> str:
        source = f"formCurso:resultadoConsultaCurso:{index}:j_idt192:j_idt192"
        return f"""
        <tr id="formCurso:resultadoConsultaCurso:{index}">
          <td><a id="formCurso:resultadoConsultaCurso:{index}:j_idt185"
                 onclick="RichFaces.ajax('x',event,{{}});return false;">{course}</a></td>
          <td><a id="formCurso:resultadoConsultaCurso:{index}:j_idt189"
                 onclick="RichFaces.ajax('y',event,{{}});return false;">{turn}</a></td>
          <td>
            <a id="{source}" onclick="RichFaces.ajax('z',event,{{}});return false;"></a>
            <span id="{source}tooltip"><span id="{source}tooltip:content"><label>Selecionar</label></span></span>
          </td>
        </tr>
        """

    @staticmethod
    def _make_current_lic_report(path: Path) -> None:
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Ano/Semestre:"
        ws["C1"] = "2026/2"
        ws["A3"] = "Unidade Ensino:"
        ws["A4"] = "Curso:"
        # Rótulo observado no HAR de 29/09/2026 para a Licenciatura.
        ws["C4"] = "Educação Física"
        ws["A5"] = "Disciplina:"
        ws["C5"] = "Didática"
        ws["A6"] = "Turma:"
        ws["C6"] = "EFL1"
        ws["N6"] = "1º"
        ws["A7"] = "Matrícula"
        ws["C7"] = "Nome"
        ws["K7"] = "Média"
        ws["R7"] = "Situação"
        ws["A8"] = "LIC001"
        ws["C8"] = "Aluno Licenciatura"
        ws["K8"] = 8.0
        ws["R8"] = "Aprovado"
        ws["A9"] = "Qtd de alunos:"
        ws["C9"] = 1
        wb.save(path)
        wb.close()

    def test_generic_educacao_fisica_is_sei_alias_only_for_licenciatura(self):
        # Continua ambíguo no catálogo institucional.
        self.assertFalse(
            course_name_matches("Educação Física", "Educação Física - Licenciatura", "DCS")
        )
        # Mas é aceito quando o fluxo SEI já solicitou explicitamente Licenciatura.
        self.assertTrue(
            course_sei_name_matches("Educação Física", "Educação Física - Licenciatura")
        )
        # Nunca deve virar Bacharelado por esse relaxamento.
        self.assertFalse(
            course_sei_name_matches("Educação Física", "Educação Física - Bacharelado")
        )

    def test_selector_distinguishes_current_licenciatura_from_explicit_bacharelado(self):
        html = "<table><tbody>" + "".join([
            self._course_row(1, "Educação Física (Bac. Presencial)", "INTEGRAL - NOTURNO"),
            self._course_row(5, "Educação Física", "MATUTINO(INATIVO)"),
            self._course_row(6, "Educação Física", "NOTURNO(INATIVO)"),
            self._course_row(7, "Educação Física", "INTEGRAL - NOTURNO"),
            self._course_row(8, "Educação Física", "INTEGRAL - MATUTINO"),
        ]) + "</tbody></table>"

        lic_source, lic_label = SEIBot.encontrar_candidato_do_curso(
            html, "Educação Física - Licenciatura"
        )
        bach_source, bach_label = SEIBot.encontrar_candidato_do_curso(
            html, "Educação Física - Bacharelado"
        )

        self.assertEqual(lic_label, "Educação Física")
        self.assertIn(":7:", lic_source or "")
        self.assertEqual(bach_label, "Educação Física (Bac. Presencial)")
        self.assertIn(":1:", bach_source or "")

    def test_current_generic_xlsx_is_validated_and_persisted_as_licenciatura_with_explicit_context(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "educacao_fisica_atual.xlsx"
            self._make_current_lic_report(path)

            diagnostic = validar_xlsx_baixado(
                path,
                course_name="Educação Física - Licenciatura",
                ano="2026",
                semestre="2",
            )
            self.assertEqual(diagnostic["curso"], "Educação Física")
            self.assertEqual(diagnostic["registros"], 1)

            engine = create_engine("sqlite:///:memory:")
            Base.metadata.create_all(engine)
            Session = sessionmaker(bind=engine, expire_on_commit=False)
            db = Session()
            directorate = Directorate(code="DCS", name="DCS", active=True)
            db.add(directorate)
            db.flush()
            db.commit()
            user = AuthorizationContext(
                user_id=str(uuid.uuid4()),
                email="eflic-current@univc.edu.br",
                full_name="EF Lic Current Test",
                role="editor",
                directorate_id=directorate.id,
                directorate_code="DCS",
                directorate_name="DCS",
            )

            try:
                repo = DatabaseRepository(db, user)
                result = repo.import_sei_report_file(
                    path,
                    expected_course="Educação Física - Licenciatura",
                )
                self.assertEqual(result["registros_lidos"], 1)
                self.assertEqual(result["metadados"]["curso_origem_sei"], "Educação Física")
                self.assertEqual(result["metadados"]["curso"], "Educação Física - Licenciatura")
                stored = db.scalar(
                    select(Course.name)
                    .join(AcademicResult, AcademicResult.course_id == Course.id)
                )
                self.assertEqual(stored, "Educação Física - Licenciatura")
            finally:
                db.close()
                engine.dispose()

    def test_generic_xlsx_without_explicit_course_context_remains_blocked(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "educacao_fisica_ambigua.xlsx"
            self._make_current_lic_report(path)

            engine = create_engine("sqlite:///:memory:")
            Base.metadata.create_all(engine)
            Session = sessionmaker(bind=engine, expire_on_commit=False)
            db = Session()
            directorate = Directorate(code="DCS", name="DCS", active=True)
            db.add(directorate)
            db.flush()
            db.commit()
            user = AuthorizationContext(
                user_id=str(uuid.uuid4()),
                email="ef-ambiguous@univc.edu.br",
                full_name="EF Ambiguous Test",
                role="editor",
                directorate_id=directorate.id,
                directorate_code="DCS",
                directorate_name="DCS",
            )
            try:
                with self.assertRaises(ValidationError):
                    DatabaseRepository(db, user).import_sei_report_file(path)
            finally:
                db.close()
                engine.dispose()


if __name__ == "__main__":
    unittest.main()
