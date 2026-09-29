from __future__ import annotations

import tempfile
import unittest
import uuid
from pathlib import Path

from openpyxl import Workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import sessionmaker

from academic_excel_parser import inspecionar_relatorio, iterar_registros
from database import Base
from models import AcademicResult, Directorate
from repository import ACADEMIC_RESULT_IMPORT_BATCH_SIZE, DatabaseRepository
from security import AuthorizationContext


class AcademicResultsV01164Tests(unittest.TestCase):
    @staticmethod
    def _make_sei_report(path: Path, rows: int = 1200) -> None:
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Ano/Semestre:"
        ws["C1"] = "2026/1"
        row = 3
        ws.cell(row, 1, "Unidade Ensino:"); row += 1
        ws.cell(row, 1, "Curso:"); ws.cell(row, 3, "Administração"); row += 1
        ws.cell(row, 1, "Disciplina:"); ws.cell(row, 3, "Gestão"); row += 1
        ws.cell(row, 1, "Turma:"); ws.cell(row, 3, "ADM1"); ws.cell(row, 14, "1º"); row += 1
        ws.cell(row, 1, "Matrícula"); ws.cell(row, 3, "Nome"); ws.cell(row, 11, "Média"); ws.cell(row, 18, "Situação"); row += 1
        for index in range(rows):
            ws.cell(row, 1, f"{index:07d}")
            ws.cell(row, 3, f"Aluno {index}")
            ws.cell(row, 11, 8.5)
            ws.cell(row, 18, "Aprovado")
            row += 1
        ws.cell(row, 1, "Qtd de alunos:")
        ws.cell(row, 3, rows)
        wb.save(path)
        wb.close()

    def test_parser_streams_without_materializing_workbook(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "sei.xlsx"
            self._make_sei_report(path, 1200)
            metadata, warnings = inspecionar_relatorio(path)
            self.assertEqual(metadata["parser_mode"], "read_only_streaming")
            self.assertEqual(metadata["total_registros_aluno_disciplina"], 1200)
            self.assertEqual(metadata["total_alunos_unicos"], 1200)
            self.assertEqual(warnings, [])
            self.assertEqual(sum(1 for _ in iterar_registros(path)), 1200)

    def test_repository_commits_large_report_in_bounded_batches_and_is_idempotent(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "sei.xlsx"
            self._make_sei_report(path, ACADEMIC_RESULT_IMPORT_BATCH_SIZE * 2 + 7)
            engine = create_engine("sqlite:///:memory:")
            Base.metadata.create_all(engine)
            Session = sessionmaker(bind=engine, expire_on_commit=False)
            db = Session()
            directorate = Directorate(code="DTNH", name="DTNH", active=True)
            db.add(directorate); db.flush(); db.commit()
            user = AuthorizationContext(
                user_id=str(uuid.uuid4()), email="results@univc.edu.br", full_name="Results Test",
                role="editor", directorate_id=directorate.id, directorate_code="DTNH", directorate_name="DTNH",
            )
            repo = DatabaseRepository(db, user)
            total_rows = ACADEMIC_RESULT_IMPORT_BATCH_SIZE * 2 + 7
            first = repo.import_sei_report_file(path, expected_course="Administração")
            self.assertEqual(first["inseridos"], total_rows)
            self.assertEqual(first["lotes_processados"], 3)
            self.assertEqual(first["engine"], "streaming-batch-v1")
            self.assertEqual(db.scalar(select(func.count(AcademicResult.id))), total_rows)
            second = repo.import_sei_report_file(path, expected_course="Administração")
            self.assertEqual(second["inseridos"], 0)
            self.assertEqual(second["ignorados"], total_rows)
            db.close(); engine.dispose()


    def test_standardized_results_model_is_also_chunked(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "modelo.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = "IMPORTACAO"
            ws.append(["Período", "Curso", "Disciplina", "Turma", "Matrícula", "Aluno", "Média", "Situação", "Aprovado", "Motivo da Reprovação"])
            total_rows = ACADEMIC_RESULT_IMPORT_BATCH_SIZE * 2 + 7
            for index in range(total_rows):
                ws.append(["2026-SEM1", "Administração", "Gestão", "ADM1", f"M{index:07d}", f"Aluno {index}", 8.0, "Aprovado", True, None])
            wb.save(path); wb.close()

            engine = create_engine("sqlite:///:memory:")
            Base.metadata.create_all(engine)
            Session = sessionmaker(bind=engine, expire_on_commit=False)
            db = Session()
            directorate = Directorate(code="DTNH", name="DTNH", active=True)
            db.add(directorate); db.flush(); db.commit()
            user = AuthorizationContext(
                user_id=str(uuid.uuid4()), email="model@univc.edu.br", full_name="Model Test",
                role="editor", directorate_id=directorate.id, directorate_code="DTNH", directorate_name="DTNH",
            )
            result = DatabaseRepository(db, user).import_file("resultados", path)
            self.assertEqual(result["inseridos"], total_rows)
            self.assertEqual(result["lotes_processados"], 3)
            db.close(); engine.dispose()

    def test_heavy_import_paths_do_not_block_fastapi_event_loop(self):
        root = Path(__file__).resolve().parents[1]
        app_source = (root / "app.py").read_text(encoding="utf-8")
        self.assertIn('def import_results_from_sei(', app_source)
        self.assertNotIn('async def import_results_from_sei(', app_source)
        self.assertIn('run_in_threadpool(_import_result_file_worker, path, mode, scope)', app_source)

    def test_frontend_processes_selected_courses_one_by_one(self):
        root = Path(__file__).resolve().parents[1]
        js = (root / "static" / "js" / "app.js").read_text(encoding="utf-8")
        start = js.index("async function submitSeiImport(event)")
        end = js.index("function surveyFilterSelect", start)
        body = js[start:end]
        self.assertIn("for(let index=0;index<courses.length;index+=1)", body)
        self.assertIn("cursos:[course]", body)
        self.assertIn("Os cursos concluídos já estão persistidos no banco.", body)
        self.assertIn("window.DataUNIVC?.progress?.render", body)


if __name__ == "__main__":
    unittest.main()
