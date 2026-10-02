from __future__ import annotations

from datetime import date
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

import dm_excel_service
from dm_demo import build_dm_demo_payloads
from dm_excel_parity import audit_dm_parity
from dm_excel_service import export_dm_excel, selected_dm_excel_engine

TARGETS = [
    {"id": 1, "indicator_code": "DM-01", "metric_key": "cohort_members", "target": 15, "valid_from": "2025-SEM1", "valid_to": None, "dimension_key": None},
    {"id": 2, "indicator_code": "DM-02", "metric_key": "average_months_to_defense", "target": 24, "valid_from": "2025-SEM1", "valid_to": None, "dimension_key": None},
]
CUTOFF = date(2026, 8, 28)


def test_03b_dm_flag_defaults_to_excel_official(monkeypatch):
    monkeypatch.delenv("DM_EXCEL_OFFICIAL_ENABLED", raising=False)
    assert selected_dm_excel_engine() == "excel_official"


def test_03b_dm_flag_selects_excel_official(monkeypatch):
    monkeypatch.setenv("DM_EXCEL_OFFICIAL_ENABLED", "true")
    assert selected_dm_excel_engine() == "excel_official"


def test_03b_dm_service_switches_builder_without_changing_route_contract(monkeypatch):
    cohorts, students = build_dm_demo_payloads()
    sentinel = BytesIO(b"official")
    called = {}

    def fake_official(*args, **kwargs):
        called["official"] = kwargs
        return sentinel

    monkeypatch.setenv("DM_EXCEL_OFFICIAL_ENABLED", "true")
    monkeypatch.setattr(dm_excel_service, "build_dm_excel_official_workbook_bytes", fake_official)
    result = export_dm_excel(cohorts, students, targets=TARGETS, as_of=CUTOFF, generated_by="qa@univc")
    assert result is sentinel
    assert called["official"]["generated_by"] == "qa@univc"
    assert called["official"]["as_of"] == CUTOFF


def test_03b_dm_parity_recomposes_all_scopes_without_divergence():
    cohorts, students = build_dm_demo_payloads()
    report = audit_dm_parity(cohorts, students, targets=TARGETS, as_of=CUTOFF, source_kind="fixture", build_workbooks=False)
    assert len(report.cases) == 72
    assert report.failures == ()
    assert report.cutover_status == "BLOCKED"  # workbook evidence intentionally omitted in this focused semantic check


def test_03b_dm_production_candidate_is_ready_when_semantics_and_workbooks_pass():
    cohorts, students = build_dm_demo_payloads()
    report = audit_dm_parity(cohorts, students, targets=TARGETS, as_of=CUTOFF, source_kind="production", build_workbooks=True)
    assert report.failures == ()
    assert report.legacy_workbook and report.legacy_workbook.build_ok
    assert report.new_workbook and report.new_workbook.build_ok
    assert report.new_release_allowed is True
    assert report.cutover_status == "READY"


def test_03b_dm_official_engine_builds_contract_workbook(monkeypatch):
    cohorts, students = build_dm_demo_payloads()
    monkeypatch.setenv("DM_EXCEL_OFFICIAL_ENABLED", "true")
    output = export_dm_excel(cohorts, students, targets=TARGETS, as_of=CUTOFF, generated_by="qa@univc")
    wb = load_workbook(output, read_only=False, data_only=False)
    assert wb.sheetnames[:6] == [
        "LEIA-ME", "PARAMETROS", "PAINEL", "QUALIDADE E GOVERNANÇA", "MATRIZ", "PLANO_DE_ACAO"
    ]
    assert not getattr(wb, "_external_links", [])
    assert wb.vba_archive is None
    wb.close()


def test_03b_dm_ui_exposes_only_one_excel_action():
    root = Path(__file__).resolve().parents[1]
    template = (root / "templates" / "dm.html").read_text(encoding="utf-8")
    js = (root / "static" / "js" / "dm.js").read_text(encoding="utf-8")
    assert "Excel Interativo" not in template
    assert "data-dm-interactive-excel-export" not in template
    assert "data-dm-interactive-excel-export" not in js
    assert "interactiveHref" not in js
    assert template.count('href="/api/dm/excel?diretoria=DM"') >= 3


def test_03b_dm_route_has_engine_header_alias_and_reitoria_parity_gate():
    root = Path(__file__).resolve().parents[1]
    router = (root / "dm_router.py").read_text(encoding="utf-8")
    assert '@router.get("/api/dm/excel")' in router
    assert '"X-Data-UNIVC-Excel-Engine": engine' in router
    assert 'filename = "Excel_Oficial_DM.xlsx" if engine == "excel_official" else "Relatorio_DM.xlsx"' in router
    assert '@router.get("/api/dm/excel-interativo")' in router
    assert "return dm_excel(" in router
    assert '@router.get("/api/admin/excel-official/dm/parity")' in router
    assert "Depends(require_fresh_reitoria)" in router
    assert 'source_kind="production"' in router


def test_03b_dm_excel_official_is_enabled_by_default_in_env_examples():
    root = Path(__file__).resolve().parents[1]
    for name in (".env.example", ".env.production.example"):
        text = (root / name).read_text(encoding="utf-8")
        assert "DM_EXCEL_OFFICIAL_ENABLED=true" in text
