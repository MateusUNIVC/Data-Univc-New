from __future__ import annotations

import io
import re
import warnings
from collections import Counter, OrderedDict
from pathlib import Path
from typing import BinaryIO

from openpyxl import load_workbook

from survey_faculty_institution_models import ParsedFacultyInstitutionWorkbook
from survey_models import OptionAggregate, ParsedQuestion
from survey_parser import (
    clean_text,
    detect_metric_type,
    infer_semester,
    normalize_key,
    parse_count,
    parse_date,
    parse_percentage,
)


PERIOD_RE = re.compile(
    r"Per[ií]odo\s+de\s+(\d{1,2}/\d{1,2}/\d{2,4})\s+at[eé]\s+(\d{1,2}/\d{1,2}/\d{2,4})",
    flags=re.IGNORECASE,
)
GENERATED_RE = re.compile(r"\d{1,2}/\d{1,2}/\d{4}\s+\d{1,2}[.:]\d{2}[.:]\d{2}")


def _load(source: bytes | BinaryIO | str | Path):
    warnings.filterwarnings(
        "ignore",
        message="Workbook contains no default style",
        category=UserWarning,
    )
    if isinstance(source, bytes):
        source = io.BytesIO(source)
    return load_workbook(source, data_only=True, read_only=False)


def _merge_option(question: ParsedQuestion, option: OptionAggregate) -> None:
    key = normalize_key(option.label)
    existing = next((item for item in question.options if normalize_key(item.label) == key), None)
    if existing is None:
        question.options.append(option)
        return

    # Em quebras de página o SEI repete o enunciado e continua as alternativas.
    # Se a alternativa também estiver literalmente repetida, não duplique.
    same_count = existing.count == option.count
    same_pct = (
        existing.source_percentage == option.source_percentage
        or (
            existing.source_percentage is not None
            and option.source_percentage is not None
            and abs(existing.source_percentage - option.source_percentage) < 0.0001
        )
    )
    if same_count and same_pct:
        return

    existing.count += option.count
    if option.source_percentage is not None:
        existing.source_percentage = (existing.source_percentage or 0.0) + option.source_percentage


def _metadata(ws) -> dict[str, str | None]:
    title: str | None = None
    unit: str | None = None
    questionnaire: str | None = None
    period_start: str | None = None
    period_end: str | None = None
    generated_at: str | None = None

    # Layout validado nos relatórios reais 2023.1 e 2025.1. A leitura é por
    # rótulos e não por números fixos de linha, porque o SEI pagina o relatório.
    for row_index in range(1, min(ws.max_row, 100) + 1):
        col_b = clean_text(ws.cell(row_index, 2).value)
        col_e = normalize_key(clean_text(ws.cell(row_index, 5).value))
        col_g = clean_text(ws.cell(row_index, 7).value)
        col_l = clean_text(ws.cell(row_index, 12).value)

        if not generated_at and col_l and GENERATED_RE.fullmatch(col_l):
            generated_at = col_l

        if col_b:
            match = PERIOD_RE.search(col_b)
            if match:
                period_start = parse_date(match.group(1))
                period_end = parse_date(match.group(2))
            elif (
                not title
                and normalize_key(col_b) != "avaliacao institucional"
                and not normalize_key(col_b).startswith("pag")
            ):
                title = col_b

        if col_e == "unidade ensino" and col_g:
            unit = unit or col_g
        elif col_e == "questionario" and col_g:
            questionnaire = questionnaire or col_g

    return {
        "survey_title": title,
        "unit_name": unit,
        "questionnaire_name": questionnaire,
        "period_start": period_start,
        "period_end": period_end,
        "generated_at": generated_at,
    }


def parse_faculty_institution_workbook(
    source: bytes | BinaryIO | str | Path,
    *,
    source_path: str = "",
) -> ParsedFacultyInstitutionWorkbook:
    """Interpreta o relatório agregado/anônimo respondido pelos docentes.

    Layout real do SEI validado em 2023.1 e 2025.1:
      D = pergunta, I = alternativa, M = percentual, P = quantidade;
      E/G = metadados de unidade/questionário; B = título/período; L = geração.

    A ausência de curso/professor é intencional: o público docente é anônimo e
    a unidade analítica é a instituição inteira.
    """

    wb = _load(source)
    ws = wb.active
    meta = _metadata(ws)

    questions: "OrderedDict[str, ParsedQuestion]" = OrderedDict()
    current_key: str | None = None

    for row_index in range(1, ws.max_row + 1):
        question_text = clean_text(ws.cell(row_index, 4).value)  # D
        option_label = clean_text(ws.cell(row_index, 9).value)   # I
        pct_raw = ws.cell(row_index, 13).value                   # M
        count_raw = ws.cell(row_index, 16).value                 # P

        if question_text:
            current_key = normalize_key(question_text)
            if current_key not in questions:
                questions[current_key] = ParsedQuestion(
                    text=question_text,
                    normalized_text=current_key,
                    position=len(questions) + 1,
                )

        if current_key and option_label:
            count = parse_count(count_raw)
            pct = parse_percentage(pct_raw)
            if count is None and pct is None:
                questions[current_key].raw_responses.append(option_label)
                continue
            _merge_option(
                questions[current_key],
                OptionAggregate(
                    label=option_label,
                    count=max(0, int(count or 0)),
                    source_percentage=pct,
                ),
            )

    parsed_questions = list(questions.values())
    if not parsed_questions:
        raise ValueError(
            f"Não foi possível identificar perguntas no relatório docente {source_path or 'XLSX'}."
        )

    for position, question in enumerate(parsed_questions, start=1):
        question.position = position
        metric_type, nps_candidate = detect_metric_type(question.options)
        if not question.options and question.raw_responses:
            metric_type, nps_candidate = "text", False
        question.metric_type = metric_type
        question.nps_candidate = nps_candidate

    totals = [sum(option.count for option in q.options) for q in parsed_questions if q.options]
    if totals:
        respondent_count = Counter(totals).most_common(1)[0][0]
    else:
        raw_totals = [len(q.raw_responses) for q in parsed_questions if q.raw_responses]
        respondent_count = Counter(raw_totals).most_common(1)[0][0] if raw_totals else 0

    period_start = meta.get("period_start")
    title = meta.get("survey_title")
    return ParsedFacultyInstitutionWorkbook(
        source_path=source_path or "relatorio_docente.xlsx",
        unit_name=meta.get("unit_name"),
        survey_title=title,
        questionnaire_name=meta.get("questionnaire_name"),
        period_start=period_start,
        period_end=meta.get("period_end"),
        semester_suggested=infer_semester(title, period_start),
        generated_at=meta.get("generated_at"),
        respondent_count=respondent_count,
        questions=parsed_questions,
    )
