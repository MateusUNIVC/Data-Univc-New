from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from sqlalchemy import case, distinct, func, select
from sqlalchemy.orm import Session

from dadm_tallos_analytics import (
    _conditions,
    _department_labels,
    _rating_distribution,
    _summary,
    _timeline,
    _valid_rating_expr,
)
from dadm_v2_analytics import resolve_month_range
from models import DADMTallosAttendance, DADMTallosSyncRun
from release_info import APP_VERSION

# Spreadsheet palette aligned with the Data UNIVC shared UI foundation.
BRAND_950 = "082F23"
BRAND_900 = "0B3D2D"
BRAND_800 = "0E5940"
BRAND_700 = "13704F"
BRAND_600 = "1B825D"
BRAND_100 = "DCEEE5"
BRAND_50 = "F0F8F4"
ACCENT = "A9CB55"
INK_950 = "14211B"
INK_800 = "314139"
INK_600 = "66756E"
INK_500 = "85918B"
LINE = "DCE5E0"
SURFACE_MUTED = "F7FAF8"
WHITE = "FFFFFF"
IMPORT_GREEN = "008000"
STATIC_GRAY = "666666"
TEAL = "008080"

THIN_LINE = Side(style="thin", color=LINE)
NO_BORDER = Border()
PCT_FMT = "0.0%"
DECIMAL_FMT = "0.00"
INTEGER_FMT = "#,##0"
DURATION_FMT = "[h]:mm:ss"


def _seconds_as_excel_duration(value: Any) -> float | None:
    if value is None:
        return None
    return float(value) / 86400.0



def _text(value: Any, fallback: str = "Todos") -> str:
    text = str(value or "").strip()
    return text or fallback


def _entity_rows(db: Session, conditions: list, *, kind: str, directorate_id: int) -> list[dict[str, Any]]:
    valid_rating = _valid_rating_expr()
    finalized_case = case((DADMTallosAttendance.status == "finalized", 1), else_=0)
    open_case = case((DADMTallosAttendance.status == "open", 1), else_=0)

    if kind == "department":
        key = DADMTallosAttendance.department_key
        label = DADMTallosAttendance.department_name
        rows = db.execute(
            select(
                key.label("entity_id"),
                label.label("entity_name"),
                func.count(DADMTallosAttendance.id).label("attendances"),
                func.count(distinct(DADMTallosAttendance.protocol)).label("protocols"),
                func.count(distinct(DADMTallosAttendance.customer_ref)).label("people"),
                func.count(distinct(DADMTallosAttendance.employee_id)).label("operators"),
                func.avg(DADMTallosAttendance.tme_seconds).label("tme"),
                func.avg(DADMTallosAttendance.tma_seconds).label("tma"),
                func.avg(valid_rating).label("rating"),
                func.count(valid_rating).label("rating_count"),
                func.coalesce(func.sum(finalized_case), 0).label("finalized"),
                func.coalesce(func.sum(open_case), 0).label("open"),
            )
            .where(*conditions)
            .group_by(key, label)
            .order_by(func.count(DADMTallosAttendance.id).desc(), label, key)
        ).all()
        mapped_labels = _department_labels(db, directorate_id)
        out = []
        for row in rows:
            entity_id = str(row.entity_id) if row.entity_id is not None else ""
            name = mapped_labels.get(entity_id) if entity_id else None
            name = name or row.entity_name or entity_id or "Sem departamento"
            out.append(_aggregate_row(row, entity_id=entity_id, entity_name=str(name), operators=int(row.operators or 0)))
        return out

    if kind == "employee":
        key = DADMTallosAttendance.employee_id
        label = DADMTallosAttendance.employee_name
        rows = db.execute(
            select(
                key.label("entity_id"),
                label.label("entity_name"),
                func.count(DADMTallosAttendance.id).label("attendances"),
                func.count(distinct(DADMTallosAttendance.protocol)).label("protocols"),
                func.count(distinct(DADMTallosAttendance.customer_ref)).label("people"),
                func.avg(DADMTallosAttendance.tme_seconds).label("tme"),
                func.avg(DADMTallosAttendance.tma_seconds).label("tma"),
                func.avg(valid_rating).label("rating"),
                func.count(valid_rating).label("rating_count"),
                func.coalesce(func.sum(finalized_case), 0).label("finalized"),
                func.coalesce(func.sum(open_case), 0).label("open"),
            )
            .where(*conditions)
            .group_by(key, label)
            .order_by(func.count(DADMTallosAttendance.id).desc(), label, key)
        ).all()
        out = []
        for row in rows:
            entity_id = str(row.entity_id) if row.entity_id is not None else ""
            name = row.entity_name or entity_id or "Sem operador identificado"
            out.append(_aggregate_row(row, entity_id=entity_id, entity_name=str(name), operators=1 if entity_id else 0))
        return out

    raise ValueError("Tipo de entidade de relatório inválido.")


def _aggregate_row(row, *, entity_id: str, entity_name: str, operators: int) -> dict[str, Any]:
    total = int(row.attendances or 0)
    ratings = int(row.rating_count or 0)
    finalized = int(row.finalized or 0)
    return {
        "entity_id": entity_id,
        "entity_name": entity_name,
        "attendances": total,
        "protocols": int(row.protocols or 0),
        "people": int(row.people or 0),
        "active_operators": operators,
        "tme_avg_seconds": float(row.tme) if row.tme is not None else None,
        "tma_avg_seconds": float(row.tma) if row.tma is not None else None,
        "rating_avg": float(row.rating) if row.rating is not None else None,
        "rating_count": ratings,
        "rating_coverage_pct": (ratings / total * 100.0) if total else None,
        "finalized": finalized,
        "open": int(row.open or 0),
        "finalization_rate_pct": (finalized / total * 100.0) if total else None,
    }


def _entity_month_rows(
    db: Session,
    conditions: list,
    *,
    kind: str,
    directorate_id: int,
) -> list[dict[str, Any]]:
    valid_rating = _valid_rating_expr()
    if kind == "department":
        key = DADMTallosAttendance.department_key
        label = DADMTallosAttendance.department_name
    elif kind == "employee":
        key = DADMTallosAttendance.employee_id
        label = DADMTallosAttendance.employee_name
    else:
        raise ValueError("Tipo de entidade mensal inválido.")

    rows = db.execute(
        select(
            DADMTallosAttendance.month_key.label("period"),
            key.label("entity_id"),
            label.label("entity_name"),
            func.count(DADMTallosAttendance.id).label("attendances"),
            func.count(distinct(DADMTallosAttendance.protocol)).label("protocols"),
            func.count(distinct(DADMTallosAttendance.customer_ref)).label("people"),
            func.avg(DADMTallosAttendance.tme_seconds).label("tme"),
            func.avg(DADMTallosAttendance.tma_seconds).label("tma"),
            func.avg(valid_rating).label("rating"),
            func.count(valid_rating).label("rating_count"),
        )
        .where(*conditions)
        .group_by(DADMTallosAttendance.month_key, key, label)
        .order_by(DADMTallosAttendance.month_key, label, key)
    ).all()

    mapped_labels = _department_labels(db, directorate_id) if kind == "department" else {}
    out = []
    for row in rows:
        entity_id = str(row.entity_id) if row.entity_id is not None else ""
        if kind == "department":
            name = mapped_labels.get(entity_id) if entity_id else None
            name = name or row.entity_name or entity_id or "Sem departamento"
        else:
            name = row.entity_name or entity_id or "Sem operador identificado"
        total = int(row.attendances or 0)
        ratings = int(row.rating_count or 0)
        out.append({
            "period": str(row.period),
            "entity_id": entity_id,
            "entity_name": str(name),
            "attendances": total,
            "protocols": int(row.protocols or 0),
            "people": int(row.people or 0),
            "tme_avg_seconds": float(row.tme) if row.tme is not None else None,
            "tma_avg_seconds": float(row.tma) if row.tma is not None else None,
            "rating_avg": float(row.rating) if row.rating is not None else None,
            "rating_count": ratings,
            "rating_coverage_pct": (ratings / total * 100.0) if total else None,
        })
    return out


def _resolve_filter_labels(
    db: Session,
    directorate_id: int,
    *,
    department: str | None,
    employee: str | None,
    allowed_departments: tuple[str, ...] | None = None,
) -> dict[str, str | None]:
    scope_conditions = [DADMTallosAttendance.directorate_id == directorate_id]
    if allowed_departments is not None:
        allowed_casefold = tuple(str(item).strip().casefold() for item in allowed_departments if str(item).strip())
        if allowed_casefold:
            scope_conditions.append(func.lower(DADMTallosAttendance.department_key).in_(allowed_casefold))
        else:
            scope_conditions.append(DADMTallosAttendance.id == -1)

    department_label = None
    if department:
        labels = _department_labels(db, directorate_id)
        department_label = labels.get(department)
        if not department_label:
            department_label = db.scalar(
                select(DADMTallosAttendance.department_name)
                .where(
                    *scope_conditions,
                    func.lower(DADMTallosAttendance.department_key) == str(department).strip().casefold(),
                )
                .order_by(DADMTallosAttendance.reference_date.desc())
                .limit(1)
            )
        department_label = str(department_label or department)

    employee_label = None
    if employee:
        employee_label = db.scalar(
            select(DADMTallosAttendance.employee_name)
            .where(
                *scope_conditions,
                DADMTallosAttendance.employee_id == employee,
            )
            .order_by(DADMTallosAttendance.reference_date.desc())
            .limit(1)
        )
        employee_label = str(employee_label or employee)

    return {"department": department_label, "employee": employee_label}


def build_report_payload(
    db: Session,
    directorate_id: int,
    from_month: str | None = None,
    to_month: str | None = None,
    *,
    department: str | None = None,
    employee: str | None = None,
    channel: str | None = None,
    status: str | None = None,
    tabulation: str | None = None,
    allowed_departments: tuple[str, ...] | None = None,
) -> dict[str, Any]:
    start_month, end_month, start_date, end_date = resolve_month_range(
        db, directorate_id, from_month, to_month, allowed_departments=allowed_departments
    )
    conditions = _conditions(
        directorate_id,
        start_date,
        end_date,
        department=department,
        employee=employee,
        channel=channel,
        status=status,
        tabulation=tabulation,
        allowed_departments=allowed_departments,
    )
    summary = _summary(db, conditions)
    timeline = _timeline(db, conditions, grain="month")
    ratings = _rating_distribution(db, conditions)
    departments = _entity_rows(db, conditions, kind="department", directorate_id=directorate_id)
    operators = _entity_rows(db, conditions, kind="employee", directorate_id=directorate_id)
    departments_monthly = _entity_month_rows(
        db, conditions, kind="department", directorate_id=directorate_id
    )
    operators_monthly = _entity_month_rows(
        db, conditions, kind="employee", directorate_id=directorate_id
    )
    labels = _resolve_filter_labels(
        db,
        directorate_id,
        department=department,
        employee=employee,
        allowed_departments=allowed_departments,
    )
    last_sync = db.scalar(
        select(DADMTallosSyncRun)
        .where(
            DADMTallosSyncRun.directorate_id == directorate_id,
            DADMTallosSyncRun.status == "completed",
        )
        .order_by(DADMTallosSyncRun.finished_at.desc(), DADMTallosSyncRun.id.desc())
        .limit(1)
    )

    return {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "app_version": APP_VERSION,
        "period": {
            "from_month": start_month,
            "to_month": end_month,
            "start": start_date.isoformat(),
            "end": end_date.isoformat(),
        },
        "filters": {
            "department": department,
            "department_label": labels["department"],
            "employee": employee,
            "employee_label": labels["employee"],
            "channel": channel,
            "status": status,
            "tabulation": tabulation,
        },
        "summary": summary,
        "timeline": timeline,
        "ratings": ratings,
        "departments": departments,
        "operators": operators,
        "departments_monthly": departments_monthly,
        "operators_monthly": operators_monthly,
        "last_sync": {
            "id": last_sync.id,
            "finished_at": last_sync.finished_at.isoformat() if last_sync and last_sync.finished_at else None,
            "end_date": last_sync.end_date.isoformat() if last_sync and last_sync.end_date else None,
        } if last_sync else None,
        "aggregation_contract": {
            "raw_attendance_rows_exported": False,
            "rating_scale": "1-10",
            "satisfaction_threshold_inferred": False,
        },
    }


def _setup_sheet(ws, *, freeze: str | None = None) -> None:
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    if freeze:
        ws.freeze_panes = freeze


def _title(ws, title: str, subtitle: str | None = None, *, end_col: int = 9) -> int:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    c = ws.cell(1, 1, title)
    c.font = Font(name="Aptos Display", size=18, bold=True, color=BRAND_950)
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 30
    row = 2
    if subtitle:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_col)
        c = ws.cell(2, 1, subtitle)
        c.font = Font(name="Aptos", size=10, color=INK_600)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[2].height = 28
        row = 3
    return row


def _table_header(cell) -> None:
    cell.font = Font(name="Aptos", size=9, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=BRAND_800)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _body_cell(cell, *, imported: bool = True) -> None:
    cell.font = Font(name="Aptos", size=9, color=IMPORT_GREEN if imported else INK_800)
    cell.alignment = Alignment(vertical="center")
    cell.border = Border(bottom=Side(style="hair", color=LINE))


def _add_table(ws, name: str, start_row: int, end_row: int, end_col: int) -> None:
    if end_row < start_row:
        return
    ref = f"A{start_row}:{get_column_letter(end_col)}{end_row}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium4",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def _auto_width(ws, *, min_width: int = 10, max_width: int = 34) -> None:
    for column_cells in ws.columns:
        letter = get_column_letter(column_cells[0].column)
        width = min_width
        for cell in column_cells:
            if cell.value is None:
                continue
            width = max(width, min(max_width, len(str(cell.value)) + 2))
        ws.column_dimensions[letter].width = width


def _add_summary_charts(ws, timeline_ws, row_count: int) -> None:
    first = 5
    last = first + row_count - 1

    time_chart = LineChart()
    time_chart.title = "TME e TMA médios"
    time_chart.y_axis.title = "Tempo"
    time_chart.x_axis.title = "Mês"
    time_chart.height = 7.0
    time_chart.width = 13.0
    time_chart.legend.position = "b"
    data = Reference(timeline_ws, min_col=5, max_col=6, min_row=4, max_row=last)
    cats = Reference(timeline_ws, min_col=1, min_row=5, max_row=last)
    time_chart.add_data(data, titles_from_data=True)
    time_chart.set_categories(cats)
    time_chart.y_axis.numFmt = DURATION_FMT
    ws.add_chart(time_chart, "G4")

    rating_chart = LineChart()
    rating_chart.title = "Avaliação média"
    rating_chart.y_axis.title = "Nota (1–10)"
    rating_chart.x_axis.title = "Mês"
    rating_chart.height = 7.0
    rating_chart.width = 13.0
    rating_data = Reference(timeline_ws, min_col=7, min_row=4, max_row=last)
    rating_chart.add_data(rating_data, titles_from_data=True)
    rating_chart.set_categories(cats)
    rating_chart.y_axis.scaling.min = 1
    rating_chart.y_axis.scaling.max = 10
    rating_chart.legend = None
    ws.add_chart(rating_chart, "G19")

    volume_chart = BarChart()
    volume_chart.type = "col"
    volume_chart.title = "Volume de atendimentos"
    volume_chart.y_axis.title = "Atendimentos"
    volume_chart.x_axis.title = "Mês"
    volume_chart.height = 7.0
    volume_chart.width = 13.0
    volume_data = Reference(timeline_ws, min_col=2, min_row=4, max_row=last)
    volume_chart.add_data(volume_data, titles_from_data=True)
    volume_chart.set_categories(cats)
    volume_chart.legend = None
    ws.add_chart(volume_chart, "U4")


def _write_timeline_sheet(wb: Workbook, payload: dict[str, Any]) -> None:
    ws = wb.create_sheet("Evolução mensal")
    _setup_sheet(ws, freeze="A5")
    _title(
        ws,
        "Evolução mensal",
        "Uma linha por mês existente no recorte selecionado. Meses anteriores à existência de uma entidade filtrada não são inventados.",
        end_col=11,
    )
    headers = [
        "Mês", "Atendimentos", "Protocolos", "Pessoas", "TME médio", "TMA médio",
        "Avaliação média (1–10)", "Avaliações", "Cobertura (%)", "Finalizados", "Em aberto",
    ]
    header_row = 4
    for col, header in enumerate(headers, 1):
        _table_header(ws.cell(header_row, col, header))

    for row_idx, item in enumerate(payload["timeline"], header_row + 1):
        values = [
            item.get("period"),
            int(item.get("attendances") or 0),
            int(item.get("protocols") or 0),
            int(item.get("people") or 0),
            _seconds_as_excel_duration(item.get("tme_avg_seconds")),
            _seconds_as_excel_duration(item.get("tma_avg_seconds")),
            item.get("rating_avg"),
            int(item.get("rating_count") or 0),
            None,
            int(item.get("finalized") or 0),
            int(item.get("open") or 0),
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row_idx, col, value)
            _body_cell(cell)
        ws.cell(row_idx, 5).number_format = DURATION_FMT
        ws.cell(row_idx, 6).number_format = DURATION_FMT
        ws.cell(row_idx, 7).number_format = DECIMAL_FMT
        ws.cell(row_idx, 9, f'=IF(B{row_idx}=0,"",H{row_idx}/B{row_idx})')
        ws.cell(row_idx, 9).number_format = PCT_FMT
        ws.cell(row_idx, 9).font = Font(name="Aptos", size=9, color=INK_800)
        for col in (2, 3, 4, 8, 10, 11):
            ws.cell(row_idx, col).number_format = INTEGER_FMT

    end_row = header_row + max(1, len(payload["timeline"]))
    if payload["timeline"]:
        _add_table(ws, "TblDADMEvolucaoMensal", header_row, end_row, len(headers))
    _auto_width(ws)


def _write_entity_sheet(wb: Workbook, payload: dict[str, Any], *, kind: str) -> None:
    if kind == "department":
        title = "Departamentos"
        rows = payload["departments"]
        table_name = "TblDADMDepartamentos"
        entity_header = "Departamento"
        include_operators = True
    else:
        title = "Operadores"
        rows = payload["operators"]
        table_name = "TblDADMOperadores"
        entity_header = "Operador"
        include_operators = False

    ws = wb.create_sheet(title)
    _setup_sheet(ws, freeze="A5")
    _title(
        ws,
        title,
        "Valores agregados no mesmo recorte e com os mesmos filtros usados para gerar o relatório.",
        end_col=12,
    )
    headers = [entity_header, "ID", "Atendimentos", "Protocolos", "Pessoas"]
    if include_operators:
        headers.append("Operadores ativos")
    headers += [
        "TME médio", "TMA médio", "Avaliação média (1–10)", "Avaliações",
        "Cobertura (%)", "Finalizados", "Em aberto", "Taxa de finalização",
    ]
    header_row = 4
    for col, header in enumerate(headers, 1):
        _table_header(ws.cell(header_row, col, header))

    for row_idx, item in enumerate(rows, header_row + 1):
        values = [
            item["entity_name"], item["entity_id"], item["attendances"], item["protocols"], item["people"],
        ]
        if include_operators:
            values.append(item["active_operators"])
        values += [
            _seconds_as_excel_duration(item["tme_avg_seconds"]),
            _seconds_as_excel_duration(item["tma_avg_seconds"]),
            item["rating_avg"],
            item["rating_count"],
            None,
            item["finalized"],
            item["open"],
            None,
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row_idx, col, value)
            _body_cell(cell)
        time_col = 7 if include_operators else 6
        rating_col = time_col + 2
        rating_count_col = rating_col + 1
        coverage_col = rating_count_col + 1
        finalized_col = coverage_col + 1
        open_col = finalized_col + 1
        finalization_col = open_col + 1
        ws.cell(row_idx, time_col).number_format = DURATION_FMT
        ws.cell(row_idx, time_col + 1).number_format = DURATION_FMT
        ws.cell(row_idx, rating_col).number_format = DECIMAL_FMT
        ws.cell(row_idx, coverage_col, f'=IF(C{row_idx}=0,"",{get_column_letter(rating_count_col)}{row_idx}/C{row_idx})')
        ws.cell(row_idx, coverage_col).number_format = PCT_FMT
        ws.cell(row_idx, finalization_col, f'=IF(C{row_idx}=0,"",{get_column_letter(finalized_col)}{row_idx}/C{row_idx})')
        ws.cell(row_idx, finalization_col).number_format = PCT_FMT
        for col in range(3, 7 if include_operators else 6):
            ws.cell(row_idx, col).number_format = INTEGER_FMT
        ws.cell(row_idx, rating_count_col).number_format = INTEGER_FMT
        ws.cell(row_idx, finalized_col).number_format = INTEGER_FMT
        ws.cell(row_idx, open_col).number_format = INTEGER_FMT
        ws.cell(row_idx, coverage_col).font = Font(name="Aptos", size=9, color=INK_800)
        ws.cell(row_idx, finalization_col).font = Font(name="Aptos", size=9, color=INK_800)

    if rows:
        end_row = header_row + len(rows)
        _add_table(ws, table_name, header_row, end_row, len(headers))
        _add_entity_charts(ws, header_row, end_row, include_operators=include_operators)
    _auto_width(ws)


def _add_entity_charts(ws, header_row: int, end_row: int, *, include_operators: bool) -> None:
    # Keep charts readable while retaining every aggregate row in the Excel table.
    chart_end = min(end_row, header_row + 15)
    time_col = 7 if include_operators else 6
    rating_col = time_col + 2
    cats = Reference(ws, min_col=1, min_row=header_row + 1, max_row=chart_end)

    time_chart = BarChart()
    time_chart.type = "bar"
    time_chart.style = 10
    time_chart.title = "Top por volume · TME e TMA"
    time_chart.height = 8.0
    time_chart.width = 14.0
    time_chart.y_axis.title = "Entidade"
    time_chart.x_axis.title = "Tempo"
    data = Reference(ws, min_col=time_col, max_col=time_col + 1, min_row=header_row, max_row=chart_end)
    time_chart.add_data(data, titles_from_data=True)
    time_chart.set_categories(cats)
    time_chart.x_axis.numFmt = DURATION_FMT
    ws.add_chart(time_chart, f"{get_column_letter(len(ws[header_row]) + 2)}4")

    rating_chart = BarChart()
    rating_chart.type = "bar"
    rating_chart.style = 10
    rating_chart.title = "Top por volume · avaliação média"
    rating_chart.height = 8.0
    rating_chart.width = 14.0
    rating_chart.x_axis.title = "Nota (1–10)"
    rating_data = Reference(ws, min_col=rating_col, min_row=header_row, max_row=chart_end)
    rating_chart.add_data(rating_data, titles_from_data=True)
    rating_chart.set_categories(cats)
    rating_chart.x_axis.scaling.min = 1
    rating_chart.x_axis.scaling.max = 10
    rating_chart.legend = None
    ws.add_chart(rating_chart, f"{get_column_letter(len(ws[header_row]) + 2)}21")


def _write_monthly_entity_sheet(wb: Workbook, payload: dict[str, Any], *, kind: str) -> None:
    if kind == "department":
        title = "Departamento x mês"
        rows = payload["departments_monthly"]
        table_name = "TblDADMDepartamentoMes"
        entity_header = "Departamento"
    else:
        title = "Operador x mês"
        rows = payload["operators_monthly"]
        table_name = "TblDADMOperadorMes"
        entity_header = "Operador"

    ws = wb.create_sheet(title)
    _setup_sheet(ws, freeze="A5")
    _title(
        ws,
        title,
        f"Agregação mensal por {entity_header.lower()}; use os filtros da tabela do Excel para explorar o período.",
        end_col=10,
    )
    headers = [
        "Mês", entity_header, "ID", "Atendimentos", "Protocolos", "Pessoas",
        "TME médio", "TMA médio", "Avaliação média (1–10)", "Avaliações", "Cobertura (%)",
    ]
    header_row = 4
    for col, header in enumerate(headers, 1):
        _table_header(ws.cell(header_row, col, header))
    for row_idx, item in enumerate(rows, header_row + 1):
        values = [
            item["period"], item["entity_name"], item["entity_id"], item["attendances"],
            item["protocols"], item["people"], _seconds_as_excel_duration(item["tme_avg_seconds"]),
            _seconds_as_excel_duration(item["tma_avg_seconds"]), item["rating_avg"], item["rating_count"], None,
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row_idx, col, value)
            _body_cell(cell)
        ws.cell(row_idx, 7).number_format = DURATION_FMT
        ws.cell(row_idx, 8).number_format = DURATION_FMT
        ws.cell(row_idx, 9).number_format = DECIMAL_FMT
        ws.cell(row_idx, 11, f'=IF(D{row_idx}=0,"",J{row_idx}/D{row_idx})')
        ws.cell(row_idx, 11).number_format = PCT_FMT
        ws.cell(row_idx, 11).font = Font(name="Aptos", size=9, color=INK_800)
        for col in (4, 5, 6, 10):
            ws.cell(row_idx, col).number_format = INTEGER_FMT
    if rows:
        _add_table(ws, table_name, header_row, header_row + len(rows), len(headers))
    _auto_width(ws)


def _write_ratings_sheet(wb: Workbook, payload: dict[str, Any]) -> None:
    ws = wb.create_sheet("Avaliações")
    _setup_sheet(ws, freeze="A5")
    _title(
        ws,
        "Distribuição das avaliações",
        "Escala TALLOS homologada de 1 a 10. Notas ausentes não entram na média e não são tratadas como zero.",
        end_col=5,
    )
    headers = ["Nota", "Quantidade", "Participação (%)"]
    header_row = 4
    for col, header in enumerate(headers, 1):
        _table_header(ws.cell(header_row, col, header))
    for row_idx, item in enumerate(payload["ratings"], header_row + 1):
        ws.cell(row_idx, 1, int(item["rating"]))
        ws.cell(row_idx, 2, int(item["count"]))
        ws.cell(row_idx, 3, f'=IF(SUM($B$5:$B$14)=0,0,B{row_idx}/SUM($B$5:$B$14))')
        for col in range(1, 4):
            _body_cell(ws.cell(row_idx, col), imported=(col != 3))
        ws.cell(row_idx, 1).number_format = "0"
        ws.cell(row_idx, 2).number_format = INTEGER_FMT
        ws.cell(row_idx, 3).number_format = PCT_FMT
        ws.cell(row_idx, 3).font = Font(name="Aptos", size=9, color=INK_800)
    _add_table(ws, "TblDADMAvaliacoes", header_row, 14, len(headers))

    chart = BarChart()
    chart.type = "col"
    chart.title = "Distribuição das notas"
    chart.y_axis.title = "Avaliações"
    chart.x_axis.title = "Nota"
    chart.height = 8.0
    chart.width = 14.0
    data = Reference(ws, min_col=2, min_row=header_row, max_row=14)
    cats = Reference(ws, min_col=1, min_row=5, max_row=14)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.legend = None
    chart.dLbls = DataLabelList()
    chart.dLbls.showVal = True
    ws.add_chart(chart, "E4")
    _auto_width(ws)


def _write_parameters_sheet(wb: Workbook, payload: dict[str, Any]) -> None:
    ws = wb.create_sheet("Parâmetros")
    _setup_sheet(ws)
    _title(
        ws,
        "Parâmetros do relatório",
        "Contexto usado na geração e informações para auditoria do arquivo.",
        end_col=4,
    )
    filters = payload["filters"]
    period = payload["period"]
    sync = payload.get("last_sync") or {}
    rows = [
        ("Versão Data UNIVC", payload["app_version"]),
        ("Gerado em (UTC)", payload["generated_at"]),
        ("Mês inicial", period["from_month"]),
        ("Mês final", period["to_month"]),
        ("Data inicial", period["start"]),
        ("Data final", period["end"]),
        ("Departamento", _text(filters.get("department_label"))),
        ("ID do departamento", _text(filters.get("department"))),
        ("Operador", _text(filters.get("employee_label"))),
        ("ID do operador", _text(filters.get("employee"))),
        ("Canal", _text(filters.get("channel"))),
        ("Status", _text(filters.get("status"))),
        ("Tabulação", _text(filters.get("tabulation"))),
        ("Última sincronização TALLOS", sync.get("finished_at") or "Não disponível"),
        ("Cobertura da sincronização até", sync.get("end_date") or "Não disponível"),
        ("Linhas de atendimento individuais exportadas", "Não"),
        ("Escala de avaliação", "1 a 10"),
        ("Percentual de satisfação inferido", "Não"),
    ]
    start = 4
    ws.cell(start, 1, "Parâmetro")
    ws.cell(start, 2, "Valor")
    _table_header(ws.cell(start, 1))
    _table_header(ws.cell(start, 2))
    for row_idx, (label, value) in enumerate(rows, start + 1):
        ws.cell(row_idx, 1, label)
        ws.cell(row_idx, 2, value)
        _body_cell(ws.cell(row_idx, 1), imported=False)
        _body_cell(ws.cell(row_idx, 2), imported=False)
        ws.cell(row_idx, 1).font = Font(name="Aptos", size=9, color=STATIC_GRAY)
        ws.cell(row_idx, 2).font = Font(name="Aptos", size=9, color=STATIC_GRAY)
    _add_table(ws, "TblDADMParametros", start, start + len(rows), 2)
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 44


def build_report_workbook(payload: dict[str, Any]) -> BytesIO:
    wb = Workbook()
    # Create data sheets before the summary so cross-sheet chart references exist.
    placeholder = wb.active
    placeholder.title = "_placeholder"
    _write_timeline_sheet(wb, payload)
    _write_entity_sheet(wb, payload, kind="department")
    _write_entity_sheet(wb, payload, kind="employee")
    _write_monthly_entity_sheet(wb, payload, kind="department")
    _write_monthly_entity_sheet(wb, payload, kind="employee")
    _write_ratings_sheet(wb, payload)
    _write_parameters_sheet(wb, payload)

    wb.remove(placeholder)
    # Summary is inserted first after all source sheets exist for cross-sheet chart references.
    summary_ws = wb.create_sheet("Resumo", 0)
    wb.active = 0
    _setup_sheet(summary_ws)
    period = payload["period"]
    filters = payload["filters"]
    subtitle = (
        f"Período {period['from_month']} a {period['to_month']} · "
        f"Departamento: {_text(filters.get('department_label'))} · "
        f"Operador: {_text(filters.get('employee_label'))}"
    )
    row = _title(summary_ws, "DADM · Relatório TALLOS V2", subtitle, end_col=10) + 1
    summary_ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    summary_ws.cell(row, 1, "Indicadores do recorte").font = Font(name="Aptos", size=11, bold=True, color=BRAND_800)
    row += 1
    metrics = [
        ("Atendimentos", "attendances", INTEGER_FMT, False),
        ("Protocolos", "protocols", INTEGER_FMT, False),
        ("Pessoas", "people", INTEGER_FMT, False),
        ("Operadores ativos", "active_operators", INTEGER_FMT, False),
        ("TME médio", "tme_avg_seconds", DURATION_FMT, True),
        ("TME mediano", "tme_median_seconds", DURATION_FMT, True),
        ("TME P90", "tme_p90_seconds", DURATION_FMT, True),
        ("TMA médio", "tma_avg_seconds", DURATION_FMT, True),
        ("TMA mediano", "tma_median_seconds", DURATION_FMT, True),
        ("TMA P90", "tma_p90_seconds", DURATION_FMT, True),
        ("Avaliação média (1–10)", "rating_avg", DECIMAL_FMT, False),
        ("Avaliações", "rating_count", INTEGER_FMT, False),
        ("Cobertura das avaliações", "rating_coverage_pct", PCT_FMT, False),
        ("Finalizados", "finalized", INTEGER_FMT, False),
        ("Em aberto", "open", INTEGER_FMT, False),
        ("Taxa de finalização", "finalization_rate_pct", PCT_FMT, False),
    ]
    summary = payload["summary"]
    for idx, (label, key, number_format, duration) in enumerate(metrics):
        block_col = 1 if idx < 8 else 4
        block_row = row + (idx if idx < 8 else idx - 8)
        label_cell = summary_ws.cell(block_row, block_col, label)
        value_cell = summary_ws.cell(block_row, block_col + 1)
        label_cell.font = Font(name="Aptos", size=9, bold=True, color=INK_600)
        value_cell.font = Font(name="Aptos", size=11, bold=True, color=TEAL)
        value = summary.get(key)
        if key == "rating_coverage_pct":
            value_cell.value = '=IF(B5=0,"",E8/B5)'
        elif key == "finalization_rate_pct":
            value_cell.value = '=IF(B5=0,"",E10/B5)'
        else:
            if duration:
                value = _seconds_as_excel_duration(value)
            value_cell.value = value
        value_cell.number_format = number_format
        label_cell.fill = PatternFill("solid", fgColor=SURFACE_MUTED)
        value_cell.fill = PatternFill("solid", fgColor=BRAND_50)
        label_cell.border = Border(bottom=THIN_LINE)
        value_cell.border = Border(bottom=THIN_LINE)
    summary_ws.column_dimensions["A"].width = 28
    summary_ws.column_dimensions["B"].width = 18
    summary_ws.column_dimensions["C"].width = 4
    summary_ws.column_dimensions["D"].width = 28
    summary_ws.column_dimensions["E"].width = 18
    note_row = row + 9
    summary_ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row + 2, end_column=5)
    note = summary_ws.cell(note_row, 1)
    note.value = (
        "A escala homologada do TALLOS é de 1 a 10. Este relatório não converte notas em percentual de "
        "satisfação e não infere limiar sem regra institucional aprovada. O arquivo contém somente agregados; "
        "nenhum atendimento individual é exportado."
    )
    note.font = Font(name="Aptos", size=9, color=INK_600)
    note.fill = PatternFill("solid", fgColor=BRAND_50)
    note.alignment = Alignment(wrap_text=True, vertical="top")
    if payload["timeline"]:
        _add_summary_charts(summary_ws, wb["Evolução mensal"], len(payload["timeline"]))

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"

    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def build_dadm_v2_report(
    db: Session,
    directorate_id: int,
    from_month: str | None = None,
    to_month: str | None = None,
    *,
    department: str | None = None,
    employee: str | None = None,
    channel: str | None = None,
    status: str | None = None,
    tabulation: str | None = None,
    allowed_departments: tuple[str, ...] | None = None,
) -> tuple[BytesIO, dict[str, Any]]:
    payload = build_report_payload(
        db,
        directorate_id,
        from_month,
        to_month,
        department=department,
        employee=employee,
        channel=channel,
        status=status,
        tabulation=tabulation,
        allowed_departments=allowed_departments,
    )
    return build_report_workbook(payload), payload
