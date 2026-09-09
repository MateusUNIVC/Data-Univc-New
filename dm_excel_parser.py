from __future__ import annotations

import re
import unicodedata
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from dm_catalog import DMValidationError, normalize_area, parse_date, positive_int


class DMExcelImportError(ValueError):
    def __init__(self, message: str, errors: list[dict[str, Any]] | None = None):
        super().__init__(message)
        self.errors = errors or []


def _norm(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = text.replace("º", "").replace("ª", "").replace("°", "")
    text = re.sub(r"[^a-zA-Z0-9]+", " ", text).casefold().strip()
    return re.sub(r"\s+", " ", text)


def _headers(ws, row: int = 1) -> dict[str, int]:
    return {_norm(cell.value): cell.column for cell in ws[row] if _norm(cell.value)}


def _col(headers: dict[str, int], *names: str, required: bool = True) -> int | None:
    for name in names:
        key = _norm(name)
        if key in headers:
            return headers[key]
    if required:
        raise KeyError(names[0])
    return None


def _value(ws, row: int, column: int | None) -> Any:
    return ws.cell(row, column).value if column else None


def parse_dm_cohorts_workbook(path: str | Path) -> list[dict[str, Any]]:
    wb = load_workbook(path, data_only=False, read_only=False)
    ws = wb["DADOS"] if "DADOS" in wb.sheetnames else wb[wb.sheetnames[0]]
    headers = _headers(ws)
    errors: list[dict[str, Any]] = []
    try:
        area_col = _col(headers, "Área", "Area")
        cohort_col = _col(headers, "Turma", "Número da turma")
        opening_col = _col(headers, "Data de abertura", "Abertura", required=False)
        vacancies_col = _col(headers, "Vagas autorizadas", required=False)
        status_col = _col(headers, "Status", required=False)
        notes_col = _col(headers, "Observações", "Observacoes", required=False)
    except KeyError as exc:
        raise DMExcelImportError(
            f'A planilha não possui a coluna obrigatória "{exc.args[0]}".',
            [{"sheet": ws.title, "row": 1, "field": exc.args[0], "error": "Cabeçalho ausente."}],
        ) from exc

    payloads: list[dict[str, Any]] = []
    for row in range(2, ws.max_row + 1):
        raw = [_value(ws, row, col) for col in (area_col, cohort_col, opening_col, vacancies_col, status_col, notes_col)]
        if all(value in (None, "") for value in raw):
            continue
        try:
            area_code, _ = normalize_area(_value(ws, row, area_col))
            cohort_number = positive_int(_value(ws, row, cohort_col), field="turma", required=True)
            opening = parse_date(_value(ws, row, opening_col), field="data_abertura", required=False)
            vacancies = positive_int(_value(ws, row, vacancies_col), field="vagas_autorizadas", required=False) if vacancies_col else None
            payloads.append({
                "area_code": area_code,
                "cohort_number": cohort_number,
                "opening_date": opening.isoformat() if opening else None,
                "vacancies_authorized": vacancies,
                "status": str(_value(ws, row, status_col) or "Em andamento").strip(),
                "notes": str(_value(ws, row, notes_col) or "").strip() or None,
            })
        except DMValidationError as exc:
            for field, message in (exc.field_errors or {"linha": str(exc)}).items():
                errors.append({"sheet": ws.title, "row": row, "field": field, "error": message})
    if errors:
        raise DMExcelImportError(
            f"A importação possui {len(errors)} inconsistência(s). Nenhuma linha foi gravada.",
            errors,
        )
    if not payloads:
        raise DMExcelImportError("A planilha não contém turmas preenchidas.")
    return payloads


def parse_dm_students_workbook(path: str | Path) -> list[dict[str, Any]]:
    wb = load_workbook(path, data_only=False, read_only=False)
    ws = wb["DADOS"] if "DADOS" in wb.sheetnames else wb[wb.sheetnames[0]]
    headers = _headers(ws)
    errors: list[dict[str, Any]] = []
    aliases = {
        "area": ("Área", "Area"),
        "cohort": ("Turma", "Número da turma"),
        "student_code": ("Matrícula", "Matricula", "Código do aluno"),
        "student_name": ("Nome do aluno", "Aluno", "Nome"),
        "entry_date": ("Data de ingresso", "Ingresso"),
        "qualification_date": ("Data de qualificação", "Qualificação"),
        "defense_date": ("Data de defesa", "Defesa", "Data de conclusão no SEI", "Data de conclusão do curso", "Conclusão SEI"),
        "graduation_date": ("Data de titulação", "Titulação", "Data da titulação"),
        "defense_scheduled_date": ("Data de defesa marcada", "Defesa marcada"),
        "status": ("Status",),
        "diploma_status": ("Situação do diploma", "Diploma digital", "Status do diploma"),
        "exit_date": ("Data de saída", "Saída"),
        "advisor": ("Orientador",),
        "research_line": ("Linha de pesquisa",),
        "notes": ("Observações", "Observacoes"),
    }
    columns: dict[str, int | None] = {}
    for key, names in aliases.items():
        required = key in {"area", "cohort", "student_code", "student_name", "status"}
        try:
            columns[key] = _col(headers, *names, required=required)
        except KeyError as exc:
            raise DMExcelImportError(
                f'A planilha não possui a coluna obrigatória "{exc.args[0]}".',
                [{"sheet": ws.title, "row": 1, "field": exc.args[0], "error": "Cabeçalho ausente."}],
            ) from exc

    payloads: list[dict[str, Any]] = []
    required_columns = [columns[key] for key in ("area", "cohort", "student_code", "student_name", "status")]
    for row in range(2, ws.max_row + 1):
        if all(_value(ws, row, col) in (None, "") for col in required_columns):
            continue
        try:
            area_code, _ = normalize_area(_value(ws, row, columns["area"]))
            cohort_number = positive_int(_value(ws, row, columns["cohort"]), field="turma", required=True)
            payload = {
                "area_code": area_code,
                "cohort_number": cohort_number,
                "student_code": str(_value(ws, row, columns["student_code"]) or "").strip(),
                "student_name": str(_value(ws, row, columns["student_name"]) or "").strip(),
                "entry_date": _value(ws, row, columns["entry_date"]),
                "qualification_date": _value(ws, row, columns["qualification_date"]),
                "defense_date": _value(ws, row, columns["defense_date"]),
                "graduation_date": _value(ws, row, columns["graduation_date"]),
                "defense_scheduled_date": _value(ws, row, columns["defense_scheduled_date"]),
                "status": str(_value(ws, row, columns["status"]) or "").strip(),
                "diploma_status": str(_value(ws, row, columns["diploma_status"]) or "").strip() or None,
                "exit_date": _value(ws, row, columns["exit_date"]),
                "advisor": str(_value(ws, row, columns["advisor"]) or "").strip() or None,
                "research_line": str(_value(ws, row, columns["research_line"]) or "").strip() or None,
                "notes": str(_value(ws, row, columns["notes"]) or "").strip() or None,
            }
            for key in ("entry_date", "qualification_date", "defense_date", "graduation_date", "defense_scheduled_date", "exit_date"):
                value = payload[key]
                if value not in (None, ""):
                    payload[key] = parse_date(value, field=key, required=False).isoformat()
            payloads.append(payload)
        except DMValidationError as exc:
            for field, message in (exc.field_errors or {"linha": str(exc)}).items():
                errors.append({"sheet": ws.title, "row": row, "field": field, "error": message})
    if errors:
        raise DMExcelImportError(
            f"A importação possui {len(errors)} inconsistência(s). Nenhuma linha foi gravada.",
            errors,
        )
    if not payloads:
        raise DMExcelImportError("A planilha não contém alunos preenchidos.")
    return payloads
