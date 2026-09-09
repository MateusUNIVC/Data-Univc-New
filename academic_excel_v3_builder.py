from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.defined_name import DefinedName

from academic_analytics import semester_sort_key
from academic_excel_v2_builder import (
    _academic_goal,
    _goal_fields,
    _scope_for_repo,
    build_academic_report_payload,
)
from release_info import APP_VERSION
from survey_repository import SurveyRepository


# Visual language intentionally follows the user's interactive DTNH reference:
# white canvas, institutional green, pale-yellow editable controls and compact
# governance blocks.
GREEN_DARK = "045233"
GREEN = "0B7A54"
GREEN_MID = "13845E"
GREEN_PALE = "E8F1ED"
GREEN_LIGHT = "F2F8F5"
YELLOW_INPUT = "FFF6DC"
YELLOW_WARN = "FFF4D8"
RED_LIGHT = "FDE7E7"
GRAY_50 = "F7F9F8"
GRAY_100 = "EEF2F0"
GRAY_300 = "D8E0DC"
GRAY_600 = "6D7973"
INK = "1D2B25"
WHITE = "FFFFFF"
PURPLE = "6B5CA5"
TEAL = "008A8A"
BLUE = "2F6B9A"
ORANGE = "D98C24"
RED = "C94A4A"

THIN = Side(style="thin", color=GRAY_300)
HAIR = Side(style="hair", color=GRAY_300)
NPS_FMT = '0.0'
SCORE_FMT = '0.00'
PCT_POINT_FMT = '0.0"%"'
INT_FMT = '#,##0'
VAR_FMT = '+0.0;-0.0;0.0'

VISIBLE_SHEETS = [
    "LEIA-ME",
    "PARAMETROS",
    "PAINEL",
    "NPS INSTITUICAO",
    "NPS CURSO",
    "NPS DOCENTES",
    "AVALIACAO DOCENTE",
    "APROVACAO RESULTADOS",
    "MATRIZ",
    "METAS E PLANOS",
    "QUALIDADE E GOVERNANCA",
]


def _sorted_semesters(payload: dict[str, Any]) -> list[str]:
    values: set[str] = set()
    sources = [
        payload.get("institution_students", []),
        payload.get("institution_students_by_course", []),
        payload.get("course_nps", []),
        payload.get("faculty_nps", []),
        payload.get("teacher", []),
        payload.get("results", []),
    ]
    for rows in sources:
        for row in rows or []:
            period = str(row.get("periodo") or "").strip()
            if period:
                values.add(period)
    return sorted(values, key=semester_sort_key)


def _clean_selected(value: str | None, allowed: list[str], fallback: str) -> str:
    text = str(value or "").strip()
    return text if text in allowed else fallback


def build_academic_interactive_payload(
    repo,
    *,
    reference: str | None = None,
    comparison: str | None = None,
    course: str | None = None,
    discipline: str | None = None,
    window_periods: int | str | None = None,
) -> dict[str, Any]:
    """Build the full-history payload used by the academic interactive workbook.

    DTNH and DCS share the same analytical engine. Unlike Excel V2, this export
    deliberately ignores web filtering for the factual bases: web filters only
    seed the initial PARAMETROS values and the workbook owns the analytical
    recut afterwards.
    """
    if repo.directorate_code not in {"DTNH", "DCS"}:
        raise ValueError("O Excel Interativo V3 acadêmico está disponível apenas para DTNH/DCS.")

    base = build_academic_report_payload(repo, window_periods="all")
    survey = SurveyRepository(repo.db, _scope_for_repo(repo))
    base["institution_students"] = survey.institution_nps_history()
    base["institution_students_by_course"] = survey.institution_nps_course_breakdown()
    base["faculty_nps"] = survey.faculty_nps_history()

    courses = [row for row in repo.list_courses(include_inactive=False) if row.get("ativo", True)]
    disciplines = [row for row in repo.list_disciplines(include_inactive=False) if row.get("ativo", True)]
    semesters = _sorted_semesters(base)
    latest = semesters[-1] if semesters else ""
    previous = semesters[-2] if len(semesters) > 1 else ""
    course_names = [str(row.get("curso") or "").strip() for row in courses if row.get("curso")]
    discipline_names = [str(row.get("disciplina") or "").strip() for row in disciplines if row.get("disciplina")]

    if isinstance(window_periods, int) and window_periods in {4, 6, 8, 12}:
        initial_window: int | str = window_periods
    elif str(window_periods or "").lower() == "all":
        initial_window = "Todo histórico"
    else:
        initial_window = 6

    initial_reference = _clean_selected(reference, semesters, latest)
    comparison_fallback = previous if previous != initial_reference else ""
    initial_comparison = _clean_selected(comparison, semesters, comparison_fallback) if comparison else comparison_fallback
    initial_course = course if course in course_names else "(todos)"
    initial_discipline = discipline if discipline in discipline_names else "(todas)"

    snapshot = repo.academic_dashboard_snapshot()
    metas = snapshot.get("metas", [])
    codes = base["codes"]

    # Pre-expand effective goals. This keeps workbook formulas legible while
    # preserving the same hierarchy used by the web app.
    effective_goals: list[dict[str, Any]] = []
    scope_pairs = [("(todos)", "(todas)")]
    scope_pairs.extend((name, "(todas)") for name in course_names)
    scope_pairs.extend(("(todos)", name) for name in discipline_names)
    scope_pairs.extend(
        (str(row.get("curso") or ""), str(row.get("disciplina") or ""))
        for row in disciplines
        if row.get("curso") and row.get("disciplina")
    )
    unique_pairs = list(dict.fromkeys(scope_pairs))

    for period in semesters:
        for key, code in codes.items():
            if key == "nps_faculty":
                pairs = [("(todos)", "(todas)")]
            elif key in {"nps_institution", "nps_course"}:
                pairs = [("(todos)", "(todas)")] + [(name, "(todas)") for name in course_names]
            else:
                pairs = unique_pairs
            for course_name, discipline_name in pairs:
                course_arg = None if course_name == "(todos)" else course_name
                discipline_arg = None if discipline_name == "(todas)" else discipline_name
                goal = _academic_goal(metas, code, period, course_arg, discipline_arg)
                fields = _goal_fields(goal, code, course=course_arg, discipline=discipline_arg)
                effective_goals.append({
                    "periodo": period,
                    "kpi": code,
                    "curso": course_name,
                    "disciplina": discipline_name,
                    "meta": fields.get("meta"),
                    "atencao": fields.get("atencao"),
                    "vigencia": fields.get("meta_vigencia"),
                    "recorte": fields.get("meta_recorte"),
                })

    # Static quality facts for auditability.
    nps_rows = base.get("course_nps", [])
    nps_bad = sum(
        1 for row in nps_rows
        if int(row.get("respondentes") or 0) != int(row.get("promotores") or 0) + int(row.get("neutros") or 0) + int(row.get("detratores") or 0)
    )
    result_rows = base.get("results", [])
    result_bad = sum(
        1 for row in result_rows
        if int(row.get("finalizados") or 0) != int(row.get("aprovados") or 0) + int(row.get("reprovados_nota") or 0) + int(row.get("reprovados_falta") or 0) + int(row.get("reprovados_outro") or 0)
    )

    return {
        **base,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "app_version": APP_VERSION,
        "courses_catalog": courses,
        "disciplines_catalog": disciplines,
        "semesters": semesters,
        "effective_goals": effective_goals,
        "initial": {
            "reference": initial_reference,
            "comparison": initial_comparison,
            "window": initial_window,
            "course": initial_course,
            "discipline": initial_discipline,
            "matrix_kpi": "01A · NPS Instituição · Alunos",
        },
        "quality": {
            "nps_inconsistent_rows": nps_bad,
            "result_inconsistent_rows": result_bad,
            "periods": len(semesters),
            "courses": len(course_names),
            "disciplines": len(discipline_names),
        },
    }


def _setup(ws, *, freeze: str | None = None, zoom: int = 90) -> None:
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = zoom
    if freeze:
        ws.freeze_panes = freeze


def _title(ws, title: str, subtitle: str = "", *, end_col: int = 12) -> int:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    c = ws.cell(1, 1, title)
    c.font = Font(name="Arial", size=15, bold=True, color=GREEN_DARK)
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 28
    if subtitle:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_col)
        c = ws.cell(2, 1, subtitle)
        c.font = Font(name="Arial", size=9, color=GRAY_600)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[2].height = 30
        return 4
    return 3


def _section(ws, row: int, text: str, *, end_col: int = 12) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=end_col)
    c = ws.cell(row, 1, text)
    c.fill = PatternFill("solid", fgColor=GREEN_DARK)
    c.font = Font(name="Arial", size=11, bold=True, color=WHITE)
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = 22


def _header(cell) -> None:
    cell.fill = PatternFill("solid", fgColor=GREEN)
    cell.font = Font(name="Arial", size=8, bold=True, color=WHITE)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = Border(bottom=THIN)


def _body(cell, *, imported: bool = False, formula: bool = False) -> None:
    color = "008000" if imported else ("000000" if formula else GRAY_600)
    cell.font = Font(name="Arial", size=9, color=color)
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    cell.border = Border(bottom=HAIR)


def _status_fill(cell, formula: bool = False) -> None:
    # Status is formula-driven; conditional formatting is added on summary ranges.
    cell.font = Font(name="Arial", size=9, bold=True, color="000000")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if not formula:
        text = str(cell.value or "")
        if text == "Dentro da meta":
            cell.fill = PatternFill("solid", fgColor=GREEN_PALE)
        elif text == "Atenção":
            cell.fill = PatternFill("solid", fgColor=YELLOW_WARN)
        elif text == "Fora da meta":
            cell.fill = PatternFill("solid", fgColor=RED_LIGHT)
        else:
            cell.fill = PatternFill("solid", fgColor=GRAY_100)


def _add_status_cf(ws, cell_range: str) -> None:
    ws.conditional_formatting.add(cell_range, CellIsRule(operator="equal", formula=['"Dentro da meta"'], fill=PatternFill("solid", fgColor=GREEN_PALE)))
    ws.conditional_formatting.add(cell_range, CellIsRule(operator="equal", formula=['"Atenção"'], fill=PatternFill("solid", fgColor=YELLOW_WARN)))
    ws.conditional_formatting.add(cell_range, CellIsRule(operator="equal", formula=['"Fora da meta"'], fill=PatternFill("solid", fgColor=RED_LIGHT)))
    ws.conditional_formatting.add(cell_range, CellIsRule(operator="equal", formula=['"Sem meta"'], fill=PatternFill("solid", fgColor=GRAY_100)))


def _auto_width(ws, *, min_width: int = 10, max_width: int = 34) -> None:
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        width = min_width
        for row in range(1, min(ws.max_row, 250) + 1):
            value = ws.cell(row, col).value
            if value is not None:
                width = max(width, min(max_width, len(str(value)) + 2))
        ws.column_dimensions[letter].width = width


def _add_table(ws, name: str, header_row: int, last_row: int, last_col: int) -> None:
    if last_row <= header_row:
        return
    table = Table(displayName=name, ref=f"A{header_row}:{get_column_letter(last_col)}{last_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium4", showRowStripes=True, showColumnStripes=False,
        showFirstColumn=False, showLastColumn=False,
    )
    ws.add_table(table)


def _write_rows_sheet(wb: Workbook, name: str, headers: list[str], rows: list[list[Any]], *, table_name: str) -> int:
    ws = wb.create_sheet(name)
    _setup(ws, freeze="A5", zoom=80)
    _title(ws, name.replace("_", " "), "Base técnica importada do Data UNIVC. Não editar.", end_col=max(8, len(headers)))
    header_row = 4
    for idx, header in enumerate(headers, 1):
        _header(ws.cell(header_row, idx, header))
    for r_idx, values in enumerate(rows, header_row + 1):
        for c_idx, value in enumerate(values, 1):
            cell = ws.cell(r_idx, c_idx, value)
            _body(cell, imported=True)
    last = header_row + len(rows)
    _add_table(ws, table_name, header_row, last, len(headers))
    _auto_width(ws)
    ws.sheet_state = "hidden"
    return last


def _write_technical_data(wb: Workbook, payload: dict[str, Any]) -> dict[str, int]:
    ends: dict[str, int] = {}
    rows_01a = [[
        r.get("periodo"), r.get("curso"), int(r.get("respondentes") or 0), int(r.get("promotores") or 0),
        int(r.get("neutros") or 0), int(r.get("detratores") or 0), r.get("valor"), r.get("meta"),
        r.get("atencao"), r.get("status"), r.get("fonte"),
    ] for r in payload.get("institution_students_by_course", [])]
    ends["01A_DIR"] = _write_rows_sheet(
        wb, "DADOS_01A_DIR",
        ["Periodo", "Curso", "Respondentes", "Promotores", "Neutros", "Detratores", "NPS", "Meta", "Atencao", "Status", "Fonte"],
        rows_01a, table_name="TblV3_01ADir",
    )

    rows_01a_u = [[
        r.get("periodo"), int(r.get("respondentes") or 0), int(r.get("promotores") or 0), int(r.get("neutros") or 0),
        int(r.get("detratores") or 0), r.get("valor"), int(r.get("coverage") or 0), int(r.get("coverage_total") or 0),
        "Sim" if r.get("complete") else "Não",
    ] for r in payload.get("institution_students", [])]
    ends["01A_UNIVC"] = _write_rows_sheet(
        wb, "DADOS_01A_UNIVC",
        ["Periodo", "Respondentes", "Promotores", "Neutros", "Detratores", "NPS", "Diretorias", "TotalDiretorias", "Completo"],
        rows_01a_u, table_name="TblV3_01AUnivc",
    )

    rows_01b = [[
        r.get("periodo"), r.get("curso"), int(r.get("respondentes") or 0), int(r.get("promotores") or 0),
        int(r.get("neutros") or 0), int(r.get("detratores") or 0),
        ((int(r.get("promotores") or 0) - int(r.get("detratores") or 0)) / int(r.get("respondentes") or 1) * 100) if int(r.get("respondentes") or 0) else None,
        r.get("meta"), r.get("atencao"), r.get("status"), r.get("fonte"),
    ] for r in payload.get("course_nps", [])]
    ends["01B"] = _write_rows_sheet(
        wb, "DADOS_01B",
        ["Periodo", "Curso", "Respondentes", "Promotores", "Neutros", "Detratores", "NPS", "Meta", "Atencao", "Status", "Fonte"],
        rows_01b, table_name="TblV3_01B",
    )

    rows_01c = [[
        r.get("periodo"), int(r.get("respondentes") or 0), int(r.get("promotores") or 0), int(r.get("neutros") or 0),
        int(r.get("detratores") or 0), r.get("valor"), r.get("fonte"),
    ] for r in payload.get("faculty_nps", [])]
    ends["01C"] = _write_rows_sheet(
        wb, "DADOS_01C",
        ["Periodo", "Respondentes", "Promotores", "Neutros", "Detratores", "NPS", "Fonte"],
        rows_01c, table_name="TblV3_01C",
    )

    rows_02 = [[
        r.get("periodo"), r.get("curso"), r.get("disciplina"), int(r.get("respondentes") or 0), r.get("nota_media"),
        (float(r.get("nota_media")) * int(r.get("respondentes") or 0)) if r.get("nota_media") is not None else None,
        r.get("meta"), r.get("atencao"), r.get("status"),
    ] for r in payload.get("teacher", [])]
    ends["02"] = _write_rows_sheet(
        wb, "DADOS_02",
        ["Periodo", "Curso", "Disciplina", "Respondentes", "NotaMedia", "SomaPonderada", "Meta", "Atencao", "Status"],
        rows_02, table_name="TblV3_02",
    )

    rows_03 = [[
        r.get("periodo"), r.get("curso"), r.get("disciplina"), int(r.get("total_registros") or 0), int(r.get("finalizados") or 0),
        int(r.get("aprovados") or 0), int(r.get("reprovados_nota") or 0), int(r.get("reprovados_falta") or 0),
        int(r.get("reprovados_outro") or 0), int(r.get("em_andamento") or 0),
        (int(r.get("aprovados") or 0) / int(r.get("finalizados") or 1) * 100) if int(r.get("finalizados") or 0) else None,
        r.get("meta"), r.get("atencao"), r.get("status"),
    ] for r in payload.get("results", [])]
    ends["03"] = _write_rows_sheet(
        wb, "DADOS_03",
        ["Periodo", "Curso", "Disciplina", "TotalRegistros", "Finalizados", "Aprovados", "ReprovadosNota", "ReprovadosFalta", "ReprovadosOutro", "EmAndamento", "TaxaAprovacao", "Meta", "Atencao", "Status"],
        rows_03, table_name="TblV3_03",
    )

    goal_rows = [[
        r.get("periodo"), r.get("kpi"), r.get("curso"), r.get("disciplina"), r.get("meta"), r.get("atencao"), r.get("vigencia"), r.get("recorte"),
    ] for r in payload.get("effective_goals", [])]
    ends["META"] = _write_rows_sheet(
        wb, "META_EFETIVA",
        ["Periodo", "KPI", "Curso", "Disciplina", "Meta", "Atencao", "Vigencia", "Recorte"],
        goal_rows, table_name="TblV3_Meta",
    )

    raw_goals = [[
        r.get("indicador"), r.get("recorte"), r.get("vigencia"), r.get("meta"), r.get("atencao"), r.get("limite_superior"), r.get("justificativa"),
    ] for r in payload.get("goals", [])]
    ends["METAS_RAW"] = _write_rows_sheet(
        wb, "METAS_RAW",
        ["KPI", "Recorte", "Vigencia", "Meta", "Atencao", "LimiteSuperior", "Justificativa"],
        raw_goals, table_name="TblV3_MetasRaw",
    )

    action_rows = [[
        r.get("indicador"), r.get("recorte") or r.get("escopo"), r.get("acao") or r.get("plano") or r.get("descricao"),
        r.get("responsavel"), r.get("prazo"), r.get("status"), r.get("resultado_esperado"),
    ] for r in payload.get("actions", [])]
    ends["PLANOS"] = _write_rows_sheet(
        wb, "PLANOS_RAW",
        ["KPI", "Recorte", "Acao", "Responsavel", "Prazo", "Status", "ResultadoEsperado"],
        action_rows, table_name="TblV3_Planos",
    )
    return ends


def _named(wb: Workbook, name: str, ref: str) -> None:
    wb.defined_names.add(DefinedName(name, attr_text=ref))


def _write_lists(wb: Workbook, payload: dict[str, Any]) -> None:
    ws = wb.create_sheet("LISTAS")
    _setup(ws)
    lists = {
        "A": ["Semestres"] + payload.get("semesters", []),
        "B": ["Cursos", "(todos)"] + [r.get("curso") for r in payload.get("courses_catalog", []) if r.get("curso")],
        "C": ["Disciplinas", "(todas)"] + list(dict.fromkeys(r.get("disciplina") for r in payload.get("disciplines_catalog", []) if r.get("disciplina"))),
        "D": ["Janelas", 4, 6, 8, 12, "Todo histórico"],
        "E": ["KPI Matriz", "01A · NPS Instituição · Alunos", "01B · NPS do Curso", "02 · Avaliação Docente", "03 · Aprovação"],
    }
    for col, values in lists.items():
        for row, value in enumerate(values, 1):
            ws[f"{col}{row}"] = value
    ws.sheet_state = "hidden"


def _write_parameters(wb: Workbook, payload: dict[str, Any]) -> None:
    ws = wb.create_sheet("PARAMETROS")
    _setup(ws, zoom=95)
    _title(ws, "PARÂMETROS DO PAINEL", "Só as células amarelas são editáveis. Use as listas suspensas; o restante do arquivo responde automaticamente.", end_col=6)
    _section(ws, 4, "SELEÇÃO", end_col=6)
    labels = [
        (5, "Diretoria", payload["directorate"], "Identidade do workbook."),
        (6, "Semestre de referência", payload["initial"]["reference"], "Período principal analisado."),
        (7, "Semestre de comparação", payload["initial"]["comparison"], "Pode ficar vazio."),
        (8, "Janela do gráfico", payload["initial"]["window"], "4, 6, 8, 12 semestres ou todo o histórico (gráfico limitado aos 12 últimos)."),
        (9, "Curso em foco", payload["initial"]["course"], "Afeta os gráficos de diretoria/curso. O benchmark UNIVC ignora este filtro."),
        (10, "Disciplina em foco", payload["initial"]["discipline"], "Usada nos KPIs 02 e 03."),
        (11, "KPI da matriz", payload["initial"]["matrix_kpi"], "Troca o indicador exibido na matriz Curso × Semestre."),
    ]
    for row, label, value, note in labels:
        ws[f"A{row}"] = label
        ws[f"B{row}"] = value
        ws[f"D{row}"] = note
        ws[f"A{row}"].font = Font(name="Arial", size=10, bold=row in {5})
        ws[f"D{row}"].font = Font(name="Arial", size=9, color=GRAY_600)
        ws[f"B{row}"].alignment = Alignment(horizontal="center", vertical="center")
        ws[f"B{row}"].font = Font(name="Arial", size=10, bold=True, color="000000")
        ws[f"B{row}"].fill = PatternFill("solid", fgColor=GRAY_100 if row == 5 else YELLOW_INPUT)
        ws[f"B{row}"].border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
        if row != 5:
            ws[f"B{row}"].protection = Protection(locked=False)

    # Data validations.
    semester_end = max(2, len(payload.get("semesters", [])) + 1)
    course_end = max(2, len(payload.get("courses_catalog", [])) + 2)
    discipline_count = len(list(dict.fromkeys(r.get("disciplina") for r in payload.get("disciplines_catalog", []) if r.get("disciplina"))))
    discipline_end = max(2, discipline_count + 2)
    for cell, formula in [
        ("B6", f"LISTAS!$A$2:$A${semester_end}"),
        ("B7", f"LISTAS!$A$2:$A${semester_end}"),
        ("B8", "LISTAS!$D$2:$D$6"),
        ("B9", f"LISTAS!$B$2:$B${course_end}"),
        ("B10", f"LISTAS!$C$2:$C${discipline_end}"),
        ("B11", "LISTAS!$E$2:$E$5"),
    ]:
        dv = DataValidation(type="list", formula1=formula, allow_blank=(cell in {"B7"}))
        dv.error = "Escolha um valor da lista."
        dv.errorTitle = "Valor inválido"
        ws.add_data_validation(dv)
        dv.add(ws[cell])

    _section(ws, 13, "JANELA EFETIVA", end_col=6)
    ws["A14"] = "Índice da referência"
    ws["B14"] = '=IFERROR(MATCH(P_REF,LISTAS!$A$2:$A$200,0),"")'
    ws["A15"] = "Primeiro índice da janela"
    ws["B15"] = '=IF(P_REF_IDX="","",IF(P_WINDOW="Todo histórico",MAX(1,P_REF_IDX-11),MAX(1,P_REF_IDX-P_WINDOW+1)))'
    ws["A16"] = "Janela efetiva — de"
    ws["B16"] = '=IF(P_START_IDX="","",INDEX(LISTAS!$A$2:$A$200,P_START_IDX))'
    ws["A17"] = "Janela efetiva — até"
    ws["B17"] = '=P_REF'
    ws["A18"] = "Semestres no gráfico"
    ws["B18"] = '=IF(OR(P_REF_IDX="",P_START_IDX=""),0,P_REF_IDX-P_START_IDX+1)'
    for row in range(14, 19):
        ws[f"B{row}"].fill = PatternFill("solid", fgColor=GRAY_100)
        ws[f"B{row}"].font = Font(name="Arial", size=9, bold=True)

    _section(ws, 20, "DIAGNÓSTICO", end_col=6)
    diagnostics = [
        (21, "Referência possui dados?", '=IF(P_REF="","SEM REFERÊNCIA",IF(COUNTIF(DADOS_01B!$A:$A,P_REF)+COUNTIF(DADOS_02!$A:$A,P_REF)+COUNTIF(DADOS_03!$A:$A,P_REF)>0,"OK — há dados acadêmicos","SEM DADOS no período"))'),
        (22, "Comparação possui dados?", '=IF(P_COMP="","(sem comparação)",IF(COUNTIF(DADOS_01B!$A:$A,P_COMP)+COUNTIF(DADOS_02!$A:$A,P_COMP)+COUNTIF(DADOS_03!$A:$A,P_COMP)>0,"OK — há dados","SEM DADOS no período"))'),
        (23, "Referência dentro da janela?", '=IF(P_REF="","—",IF(AND(P_REF_IDX>=P_START_IDX,P_REF_IDX<=P_REF_IDX),"OK — dentro da janela","FORA DA JANELA"))'),
        (24, "Curso em foco existe?", '=IF(P_COURSE="(todos)","OK — todos os cursos",IF(COUNTIF(LISTAS!$B:$B,P_COURSE)>0,"OK — curso encontrado","CURSO NÃO ENCONTRADO"))'),
        (25, "Disciplina em foco existe?", '=IF(P_DISC="(todas)","OK — todas as disciplinas",IF(COUNTIF(LISTAS!$C:$C,P_DISC)>0,"OK — disciplina encontrada","DISCIPLINA NÃO ENCONTRADA"))'),
        (26, "Janela", '=IF(P_WINDOW="Todo histórico","Todo histórico · gráfico mostra até 12 semestres",P_WINDOW&" semestres")'),
    ]
    for row, label, formula in diagnostics:
        ws[f"A{row}"] = label
        ws[f"B{row}"] = formula
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
        ws[f"B{row}"].fill = PatternFill("solid", fgColor=GRAY_100)
        ws[f"B{row}"].font = Font(name="Arial", size=9, bold=True)
        ws[f"B{row}"].alignment = Alignment(wrap_text=True)

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 3
    ws.column_dimensions["D"].width = 54
    ws.column_dimensions["E"].width = 16
    ws.column_dimensions["F"].width = 3
    ws.protection.sheet = True
    ws.protection.selectUnlockedCells = True
    ws.protection.selectLockedCells = True

    _named(wb, "P_REF", "'PARAMETROS'!$B$6")
    _named(wb, "P_COMP", "'PARAMETROS'!$B$7")
    _named(wb, "P_WINDOW", "'PARAMETROS'!$B$8")
    _named(wb, "P_COURSE", "'PARAMETROS'!$B$9")
    _named(wb, "P_DISC", "'PARAMETROS'!$B$10")
    _named(wb, "P_MATRIX", "'PARAMETROS'!$B$11")
    _named(wb, "P_REF_IDX", "'PARAMETROS'!$B$14")
    _named(wb, "P_START_IDX", "'PARAMETROS'!$B$15")


def _range(sheet: str, col: str, end: int) -> str:
    return f"'{sheet}'!${col}$5:${col}${max(5, end)}"


def _criteria_course(range_ref: str) -> str:
    return f'{range_ref},IF(P_COURSE="(todos)","*",P_COURSE)'


def _criteria_disc(range_ref: str) -> str:
    return f'{range_ref},IF(P_DISC="(todas)","*",P_DISC)'


def _nps_formula(sheet: str, ends: dict[str, int], key: str, period_expr: str, *, course: bool) -> str:
    end = ends[key]
    per = _range(sheet, "A", end)
    resp = _range(sheet, "C" if course else "B", end)
    prom = _range(sheet, "D" if course else "C", end)
    det = _range(sheet, "F" if course else "E", end)
    args = f"{per},{period_expr}"
    if course:
        args += f",{_criteria_course(_range(sheet, 'B', end))}"
    return f'=IFERROR((SUMIFS({prom},{args})-SUMIFS({det},{args}))/SUMIFS({resp},{args})*100,NA())'


def _teacher_formula(ends: dict[str, int], period_expr: str) -> str:
    end = ends["02"]
    per = _range("DADOS_02", "A", end)
    course = _range("DADOS_02", "B", end)
    disc = _range("DADOS_02", "C", end)
    resp = _range("DADOS_02", "D", end)
    weighted = _range("DADOS_02", "F", end)
    args = f"{per},{period_expr},{_criteria_course(course)},{_criteria_disc(disc)}"
    return f'=IFERROR(SUMIFS({weighted},{args})/SUMIFS({resp},{args}),NA())'


def _approval_formula(ends: dict[str, int], period_expr: str) -> str:
    end = ends["03"]
    per = _range("DADOS_03", "A", end)
    course = _range("DADOS_03", "B", end)
    disc = _range("DADOS_03", "C", end)
    fin = _range("DADOS_03", "E", end)
    approved = _range("DADOS_03", "F", end)
    args = f"{per},{period_expr},{_criteria_course(course)},{_criteria_disc(disc)}"
    return f'=IFERROR(SUMIFS({approved},{args})/SUMIFS({fin},{args})*100,NA())'


def _meta_formula(ends: dict[str, int], kpi_code_expr: str, period_expr: str, *, course_expr: str = "P_COURSE", disc_expr: str = '"(todas)"', field_col: str = "E") -> str:
    end = ends["META"]
    per = _range("META_EFETIVA", "A", end)
    kpi = _range("META_EFETIVA", "B", end)
    course = _range("META_EFETIVA", "C", end)
    disc = _range("META_EFETIVA", "D", end)
    val = _range("META_EFETIVA", field_col, end)
    return f'=IF(COUNTIFS({per},{period_expr},{kpi},{kpi_code_expr},{course},{course_expr},{disc},{disc_expr})=0,"",SUMIFS({val},{per},{period_expr},{kpi},{kpi_code_expr},{course},{course_expr},{disc},{disc_expr}))'


def _write_calc(wb: Workbook, payload: dict[str, Any], ends: dict[str, int]) -> None:
    ws = wb.create_sheet("CALC")
    _setup(ws)
    # 12 chart slots. They follow P_REF and P_WINDOW, capped at 12 for readability.
    headers = ["Slot", "Semestre", "01A Diretoria", "01A UNIVC", "01B Curso", "01C Docentes", "02 Avaliação", "03 Aprovação"]
    for col, text in enumerate(headers, 1):
        _header(ws.cell(4, col, text))
    for slot in range(1, 13):
        row = 4 + slot
        ws.cell(row, 1, slot)
        ws.cell(row, 2, f'=IF({slot}<=P_REF_IDX-P_START_IDX+1,INDEX(LISTAS!$A$2:$A$200,P_START_IDX+{slot}-1),"")')
        ws.cell(row, 3, f'=IF(B{row}="",NA(),{_nps_formula("DADOS_01A_DIR", ends, "01A_DIR", f"B{row}", course=True)[1:]})')
        # 01A UNIVC has no course dimension.
        end_u = ends["01A_UNIVC"]
        per_u, resp_u, prom_u, det_u = (_range("DADOS_01A_UNIVC", c, end_u) for c in ["A", "B", "C", "E"])
        ws.cell(row, 4, f'=IF(B{row}="",NA(),IFERROR((SUMIFS({prom_u},{per_u},B{row})-SUMIFS({det_u},{per_u},B{row}))/SUMIFS({resp_u},{per_u},B{row})*100,NA()))')
        ws.cell(row, 5, f'=IF(B{row}="",NA(),{_nps_formula("DADOS_01B", ends, "01B", f"B{row}", course=True)[1:]})')
        # 01C institutional faculty NPS.
        end_c = ends["01C"]
        per_c, resp_c, prom_c, det_c = (_range("DADOS_01C", c, end_c) for c in ["A", "B", "C", "E"])
        ws.cell(row, 6, f'=IF(B{row}="",NA(),IFERROR((SUMIFS({prom_c},{per_c},B{row})-SUMIFS({det_c},{per_c},B{row}))/SUMIFS({resp_c},{per_c},B{row})*100,NA()))')
        ws.cell(row, 7, f'=IF(B{row}="",NA(),{_teacher_formula(ends, f"B{row}")[1:]})')
        ws.cell(row, 8, f'=IF(B{row}="",NA(),{_approval_formula(ends, f"B{row}")[1:]})')
    for row in range(5, 17):
        for col in [3, 4, 5, 6]: ws.cell(row, col).number_format = NPS_FMT
        ws.cell(row, 7).number_format = SCORE_FMT
        ws.cell(row, 8).number_format = PCT_POINT_FMT

    # Reference/comparison KPI summary.
    _section(ws, 19, "RESUMO KPI", end_col=10)
    headers2 = ["KPI", "Código", "Referência", "Comparação", "Variação", "Meta", "Atenção", "Status"]
    for c, h in enumerate(headers2, 1): _header(ws.cell(20, c, h))
    codes = payload["codes"]
    kpis = [
        ("NPS Instituição · Alunos", codes["nps_institution"], lambda p: _nps_formula("DADOS_01A_DIR", ends, "01A_DIR", p, course=True), NPS_FMT, '"(todas)"'),
        ("NPS do Curso", codes["nps_course"], lambda p: _nps_formula("DADOS_01B", ends, "01B", p, course=True), NPS_FMT, '"(todas)"'),
        ("NPS Instituição · Docentes", codes["nps_faculty"], None, NPS_FMT, '"(todas)"'),
        ("Avaliação Docente", codes["teacher"], lambda p: _teacher_formula(ends, p), SCORE_FMT, "P_DISC"),
        ("Aprovação", codes["approval"], lambda p: _approval_formula(ends, p), PCT_POINT_FMT, "P_DISC"),
    ]
    for idx, (name, code, fn, numfmt, disc_expr) in enumerate(kpis, 21):
        ws.cell(idx, 1, name); ws.cell(idx, 2, code)
        if code == codes["nps_faculty"]:
            end_c = ends["01C"]
            per_c, resp_c, prom_c, det_c = (_range("DADOS_01C", c, end_c) for c in ["A", "B", "C", "E"])
            ref_formula = f'=IFERROR((SUMIFS({prom_c},{per_c},P_REF)-SUMIFS({det_c},{per_c},P_REF))/SUMIFS({resp_c},{per_c},P_REF)*100,NA())'
            comp_formula = f'=IF(P_COMP="",NA(),IFERROR((SUMIFS({prom_c},{per_c},P_COMP)-SUMIFS({det_c},{per_c},P_COMP))/SUMIFS({resp_c},{per_c},P_COMP)*100,NA()))'
            course_expr = '"(todos)"'
        else:
            ref_formula = fn("P_REF")
            comp_formula = f'=IF(P_COMP="",NA(),{fn("P_COMP")[1:]})'
            course_expr = '"(todos)"' if code == codes["nps_faculty"] else "P_COURSE"
        ws.cell(idx, 3, ref_formula)
        ws.cell(idx, 4, comp_formula)
        ws.cell(idx, 5, f'=IF(OR(ISNA(C{idx}),ISNA(D{idx})),"",C{idx}-D{idx})')
        ws.cell(idx, 6, _meta_formula(ends, f'B{idx}', "P_REF", course_expr=course_expr, disc_expr=disc_expr, field_col="E"))
        ws.cell(idx, 7, _meta_formula(ends, f'B{idx}', "P_REF", course_expr=course_expr, disc_expr=disc_expr, field_col="F"))
        ws.cell(idx, 8, f'=IF(ISNA(C{idx}),"Sem dados",IF(F{idx}="","Sem meta",IF(C{idx}>=F{idx},"Dentro da meta",IF(AND(G{idx}<>"",C{idx}>=G{idx}),"Atenção","Fora da meta"))))')
        for col in [3,4,5,6,7]: ws.cell(idx, col).number_format = numfmt if col != 5 else VAR_FMT

    # 01A course snapshot at reference, one row per active course.
    _section(ws, 29, "01A · CURSOS NO SEMESTRE DE REFERÊNCIA", end_col=9)
    snap_headers = ["Curso", "NPS", "Meta", "Atenção", "Status", "Dentro", "AtençãoSerie", "Fora", "SemMeta"]
    for c,h in enumerate(snap_headers,1): _header(ws.cell(30,c,h))
    courses = [r.get("curso") for r in payload.get("courses_catalog", []) if r.get("curso")]
    for offset, course_name in enumerate(courses, 31):
        ws.cell(offset,1,course_name)
        end = ends["01A_DIR"]
        per = _range("DADOS_01A_DIR","A",end); cr = _range("DADOS_01A_DIR","B",end)
        resp = _range("DADOS_01A_DIR","C",end); prom = _range("DADOS_01A_DIR","D",end); det = _range("DADOS_01A_DIR","F",end)
        ws.cell(offset,2,f'=IFERROR((SUMIFS({prom},{per},P_REF,{cr},A{offset})-SUMIFS({det},{per},P_REF,{cr},A{offset}))/SUMIFS({resp},{per},P_REF,{cr},A{offset})*100,NA())')
        ws.cell(offset,3,_meta_formula(ends,f'"{codes["nps_institution"]}"',"P_REF",course_expr=f'A{offset}',field_col="E"))
        ws.cell(offset,4,_meta_formula(ends,f'"{codes["nps_institution"]}"',"P_REF",course_expr=f'A{offset}',field_col="F"))
        ws.cell(offset,5,f'=IF(ISNA(B{offset}),"Sem dados",IF(C{offset}="","Sem meta",IF(B{offset}>=C{offset},"Dentro da meta",IF(AND(D{offset}<>"",B{offset}>=D{offset}),"Atenção","Fora da meta"))))')
        ws.cell(offset,6,f'=IF(E{offset}="Dentro da meta",B{offset},NA())')
        ws.cell(offset,7,f'=IF(E{offset}="Atenção",B{offset},NA())')
        ws.cell(offset,8,f'=IF(E{offset}="Fora da meta",B{offset},NA())')
        ws.cell(offset,9,f'=IF(E{offset}="Sem meta",B{offset},NA())')
        for c in range(2,10): ws.cell(offset,c).number_format=NPS_FMT

    # Matrix helper: up to 12 semester columns and one row per course.
    start = 31 + max(1, len(courses)) + 3
    _section(ws, start, "MATRIZ · CURSO × SEMESTRE", end_col=18)
    ws.cell(start+1,1,"Curso")
    for c in range(1,13):
        ws.cell(start+1,1+c,f'=IF({c}<=P_REF_IDX-P_START_IDX+1,INDEX(LISTAS!$A$2:$A$200,P_START_IDX+{c}-1),"")')
    for c,h in enumerate(["Referência","Comparação","Variação","Meta","Status"],14): ws.cell(start+1,c,h)
    for c in range(1,19): _header(ws.cell(start+1,c))
    matrix_first = start+2
    for r_idx, course_name in enumerate(courses, matrix_first):
        ws.cell(r_idx,1,course_name)
        for pos in range(1,13):
            col = 1+pos
            period_ref = f'{get_column_letter(col)}${start+1}'
            # Four supported matrix KPIs.
            f01a = _nps_course_specific_formula("DADOS_01A_DIR", ends["01A_DIR"], period_ref, f'$A{r_idx}')
            f01b = _nps_course_specific_formula("DADOS_01B", ends["01B"], period_ref, f'$A{r_idx}')
            f02 = _teacher_course_specific_formula(ends["02"], period_ref, f'$A{r_idx}')
            f03 = _approval_course_specific_formula(ends["03"], period_ref, f'$A{r_idx}')
            ws.cell(r_idx,col,f'=IF({period_ref}="","",IF(LEFT(P_MATRIX,3)="01A",{f01a},IF(LEFT(P_MATRIX,3)="01B",{f01b},IF(LEFT(P_MATRIX,2)="02",{f02},{f03}))))')
        # Reference/comparison columns use the same formulas against P_REF/P_COMP.
        f01a_ref = _nps_course_specific_formula("DADOS_01A_DIR", ends["01A_DIR"], "P_REF", f'$A{r_idx}')
        f01b_ref = _nps_course_specific_formula("DADOS_01B", ends["01B"], "P_REF", f'$A{r_idx}')
        f02_ref = _teacher_course_specific_formula(ends["02"], "P_REF", f'$A{r_idx}')
        f03_ref = _approval_course_specific_formula(ends["03"], "P_REF", f'$A{r_idx}')
        f01a_cmp = _nps_course_specific_formula("DADOS_01A_DIR", ends["01A_DIR"], "P_COMP", f'$A{r_idx}')
        f01b_cmp = _nps_course_specific_formula("DADOS_01B", ends["01B"], "P_COMP", f'$A{r_idx}')
        f02_cmp = _teacher_course_specific_formula(ends["02"], "P_COMP", f'$A{r_idx}')
        f03_cmp = _approval_course_specific_formula(ends["03"], "P_COMP", f'$A{r_idx}')
        ws.cell(r_idx,14,f'=IF(LEFT(P_MATRIX,3)="01A",{f01a_ref},IF(LEFT(P_MATRIX,3)="01B",{f01b_ref},IF(LEFT(P_MATRIX,2)="02",{f02_ref},{f03_ref})))')
        ws.cell(r_idx,15,f'=IF(P_COMP="","",IF(LEFT(P_MATRIX,3)="01A",{f01a_cmp},IF(LEFT(P_MATRIX,3)="01B",{f01b_cmp},IF(LEFT(P_MATRIX,2)="02",{f02_cmp},{f03_cmp}))))')
        ws.cell(r_idx,16,f'=IF(OR(N{r_idx}="",O{r_idx}=""),"",N{r_idx}-O{r_idx})')
        code_expr = f'IF(LEFT(P_MATRIX,3)="01A","{codes["nps_institution"]}",IF(LEFT(P_MATRIX,3)="01B","{codes["nps_course"]}",IF(LEFT(P_MATRIX,2)="02","{codes["teacher"]}","{codes["approval"]}")))'
        ws.cell(r_idx,17,_meta_formula(ends,code_expr,"P_REF",course_expr=f'A{r_idx}',disc_expr='"(todas)"',field_col="E"))
        ws.cell(r_idx,18,f'=IF(N{r_idx}="","Sem dados",IF(Q{r_idx}="","Sem meta",IF(N{r_idx}>=Q{r_idx},"Dentro da meta","Fora da meta")))')
    ws.sheet_state = "hidden"


def _nps_course_specific_formula(sheet: str, end: int, period_expr: str, course_expr: str) -> str:
    per = _range(sheet,"A",end); cr=_range(sheet,"B",end); resp=_range(sheet,"C",end); prom=_range(sheet,"D",end); det=_range(sheet,"F",end)
    return f'IFERROR((SUMIFS({prom},{per},{period_expr},{cr},{course_expr})-SUMIFS({det},{per},{period_expr},{cr},{course_expr}))/SUMIFS({resp},{per},{period_expr},{cr},{course_expr})*100,"")'


def _teacher_course_specific_formula(end: int, period_expr: str, course_expr: str) -> str:
    per=_range("DADOS_02","A",end); cr=_range("DADOS_02","B",end); resp=_range("DADOS_02","D",end); weighted=_range("DADOS_02","F",end)
    return f'IFERROR(SUMIFS({weighted},{per},{period_expr},{cr},{course_expr})/SUMIFS({resp},{per},{period_expr},{cr},{course_expr}),"")'


def _approval_course_specific_formula(end: int, period_expr: str, course_expr: str) -> str:
    per=_range("DADOS_03","A",end); cr=_range("DADOS_03","B",end); fin=_range("DADOS_03","E",end); app=_range("DADOS_03","F",end)
    return f'IFERROR(SUMIFS({app},{per},{period_expr},{cr},{course_expr})/SUMIFS({fin},{per},{period_expr},{cr},{course_expr})*100,"")'


def _chart_line(ws, title: str, data_sheet, value_col: int, anchor: str, *, number_format: str = "0.0", y_min=None, y_max=None) -> None:
    chart = LineChart()
    chart.title = title
    chart.style = 13
    chart.height = 7.2
    chart.width = 12.8
    chart.legend = None
    data = Reference(data_sheet, min_col=value_col, min_row=4, max_row=16)
    cats = Reference(data_sheet, min_col=2, min_row=5, max_row=16)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.y_axis.numFmt = number_format
    if y_min is not None: chart.y_axis.scaling.min = y_min
    if y_max is not None: chart.y_axis.scaling.max = y_max
    chart.x_axis.title = "Semestre"
    chart.y_axis.majorGridlines = None
    ws.add_chart(chart, anchor)


def _write_readme(wb: Workbook, payload: dict[str, Any]) -> None:
    code = payload["directorate"]
    codes = payload["codes"]
    ws = wb.create_sheet("LEIA-ME", 0)
    _setup(ws, zoom=100)
    _title(ws, f"EXCEL INTERATIVO · {code}", "Uma segunda interface analítica do Data UNIVC. Os dados são exportados do banco; os recortes são controlados dentro do Excel.", end_col=8)
    blocks = [
        (4, "COMO USAR", [
            "1. Vá à aba PARAMETROS e altere somente as células amarelas.",
            "2. Escolha referência, comparação, janela, curso e disciplina.",
            "3. PAINEL e as abas de indicadores respondem automaticamente.",
            "4. MATRIZ permite comparar cursos ao longo da janela e trocar o KPI analisado.",
        ]),
        (11, "FONTE E SEGURANÇA", [
            "O PostgreSQL/Data UNIVC continua sendo a fonte oficial. Este arquivo não escreve dados de volta no sistema.",
            "Bases técnicas ficam ocultas por padrão e contêm apenas agregações gerenciais; linhas individuais de alunos não são exportadas.",
            "DTNH e DCS compartilham o mesmo motor V3. O Excel V2 atual continua disponível em paralelo.",
        ]),
        (17, "INDICADORES", [
            f"{codes['nps_institution']} · NPS da Instituição · Alunos — inclui tendência da diretoria, benchmark geral UNIVC e comparação por curso.",
            f"{codes['nps_course']} · NPS do Curso.",
            f"{codes['nps_faculty']} · NPS da Instituição · Docentes.",
            f"{codes['teacher']} · Avaliação Docente pelo Aluno.",
            f"{codes['approval']} · Aprovação e Resultados.",
        ]),
    ]
    for start, title, lines in blocks:
        _section(ws, start, title, end_col=8)
        for i,line in enumerate(lines,start+1):
            ws.merge_cells(start_row=i,start_column=1,end_row=i,end_column=8)
            ws.cell(i,1,line)
            ws.cell(i,1).font=Font(name="Arial",size=10,color=INK)
            ws.cell(i,1).alignment=Alignment(wrap_text=True,vertical="center")
            ws.row_dimensions[i].height=26
    ws["A26"] = f"Gerado em {payload.get('generated_at','')} · Data UNIVC {payload.get('app_version','')}"
    ws["A26"].font = Font(name="Arial", size=8, color=GRAY_600)
    for col in range(1,9): ws.column_dimensions[get_column_letter(col)].width=16


def _write_panel(wb: Workbook, payload: dict[str, Any], ends: dict[str, int]) -> None:
    code = payload["directorate"]
    ws = wb.create_sheet("PAINEL")
    _setup(ws, freeze="A4", zoom=85)
    _title(ws, f"PAINEL DE GESTÃO · {code}", "A leitura muda pela aba PARAMETROS. O benchmark geral da UNIVC não é afetado pelo curso em foco.", end_col=12)
    # Current controls strip.
    for c,h in enumerate(["Referência","Comparação","Janela","Curso em foco","Disciplina"],1):
        ws.cell(4,(c-1)*2+1,h); _header(ws.cell(4,(c-1)*2+1)); ws.merge_cells(start_row=4,start_column=(c-1)*2+1,end_row=4,end_column=(c-1)*2+2)
    refs=["=P_REF", '=IF(P_COMP="","(sem)",P_COMP)', "='PARAMETROS'!$B$16&\" a \"&'PARAMETROS'!$B$17", "=P_COURSE", "=P_DISC"]
    for c,f in enumerate(refs,1):
        col=(c-1)*2+1; ws.merge_cells(start_row=5,start_column=col,end_row=5,end_column=col+1); ws.cell(5,col,f); ws.cell(5,col).font=Font(name="Arial",size=10,bold=True); ws.cell(5,col).alignment=Alignment(horizontal="center")
    _section(ws,7,"RESUMO EXECUTIVO DOS KPIs",end_col=12)
    headers=["KPI","Referência","Comparação","Variação","Meta","Atenção","Status"]
    for c,h in enumerate(headers,1): _header(ws.cell(8,c,h))
    for r in range(21,26):
        dest=9+(r-21)
        ws.cell(dest,1,f'=CALC!A{r}')
        for c,src in enumerate(range(3,9),2): ws.cell(dest,c,f'=CALC!{get_column_letter(src)}{r}')
        for c in range(1,8): _body(ws.cell(dest,c),formula=True)
        _status_fill(ws.cell(dest,7),formula=True)
    for r in [9,10,11]:
        for c in range(2,7): ws.cell(r,c).number_format=NPS_FMT
    for c in range(2,7): ws.cell(12,c).number_format=SCORE_FMT
    for c in range(2,7): ws.cell(13,c).number_format=PCT_POINT_FMT
    _add_status_cf(ws, "G9:G13")

    _section(ws,16,"NPS INSTITUCIONAL · 3 LEITURAS",end_col=12)
    calc=wb["CALC"]
    _chart_line(ws,f"NPS institucional · {code} / curso em foco",calc,3,"A18",number_format="0.0",y_min=-100,y_max=100)
    _chart_line(ws,"NPS geral · UNIVC",calc,4,"G18",number_format="0.0",y_min=-100,y_max=100)

    # Course status bar chart.
    courses=[r.get("curso") for r in payload.get("courses_catalog",[]) if r.get("curso")]
    if courses:
        first=31; last=30+len(courses)
        chart=BarChart(); chart.type="bar"; chart.grouping="stacked"; chart.overlap=100; chart.style=10
        chart.title="NPS institucional por curso · referência"; chart.height=8.2; chart.width=20.0
        data=Reference(calc,min_col=6,max_col=9,min_row=30,max_row=last)
        cats=Reference(calc,min_col=1,min_row=first,max_row=last)
        chart.add_data(data,titles_from_data=True); chart.set_categories(cats)
        chart.legend.position="b"; chart.x_axis.scaling.min=-100; chart.x_axis.scaling.max=100
        chart.dataLabels=DataLabelList(); chart.dataLabels.showVal=True
        colors=[GREEN,ORANGE,RED,GRAY_600]
        for ser,color in zip(chart.series,colors): ser.graphicalProperties.solidFill=color
        ws.add_chart(chart,"A34")
    ws.column_dimensions["A"].width=30
    for c in range(2,13): ws.column_dimensions[get_column_letter(c)].width=13


def _write_indicator_sheet(wb: Workbook, name: str, title: str, subtitle: str, calc_col: int, *, summary_row: int, y_min=None, y_max=None, fmt="0.0", table_sheet: str | None=None) -> None:
    ws=wb.create_sheet(name); _setup(ws,zoom=88)
    _title(ws,title,subtitle,end_col=12)
    _section(ws,4,"EVOLUÇÃO NA JANELA",end_col=12)
    _chart_line(ws,title,wb["CALC"],calc_col,"A6",number_format=fmt,y_min=y_min,y_max=y_max)
    _section(ws,23,"REFERÊNCIA × COMPARAÇÃO",end_col=12)
    for c,h in enumerate(["Referência","Comparação","Variação","Meta","Atenção","Status"],1): _header(ws.cell(24,c,h))
    for c,src in enumerate(range(3,9),1):
        ws.cell(25,c,f'=CALC!{get_column_letter(src)}{summary_row}')
        _body(ws.cell(25,c),formula=True)
    for c in range(1,6): ws.cell(25,c).number_format = fmt if c != 3 else VAR_FMT
    _status_fill(ws.cell(25,6),formula=True); _add_status_cf(ws,"F25")
    ws["A28"]="Curso em foco"; ws["B28"]="=P_COURSE"; ws["D28"]="Disciplina"; ws["E28"]="=P_DISC"; ws["G28"]="Janela"; ws["H28"]="='PARAMETROS'!$B$16&\" a \"&'PARAMETROS'!$B$17"
    for c in ["A28","D28","G28"]: ws[c].font=Font(name="Arial",size=8,bold=True,color=GRAY_600)
    for c in ["B28","E28","H28"]: ws[c].font=Font(name="Arial",size=9,bold=True,color=INK)
    if table_sheet:
        ws["A31"]="Base técnica correspondente fica oculta por padrão: "+table_sheet
        ws["A31"].font=Font(name="Arial",size=8,color=GRAY_600)


def _write_nps_institution(wb: Workbook, payload: dict[str, Any]) -> None:
    code = payload["directorate"]
    ws=wb.create_sheet("NPS INSTITUICAO"); _setup(ws,zoom=85)
    _title(ws,"NPS DA INSTITUIÇÃO · ALUNOS (01A)",f"As três leituras são complementares: diretoria/curso em foco, benchmark geral UNIVC e comparação entre cursos da {code}.",end_col=12)
    _section(ws,4,"TENDÊNCIAS",end_col=12)
    _chart_line(ws,f"{code} / curso em foco",wb["CALC"],3,"A6",number_format="0.0",y_min=-100,y_max=100)
    _chart_line(ws,"UNIVC geral",wb["CALC"],4,"G6",number_format="0.0",y_min=-100,y_max=100)
    courses=[r.get("curso") for r in payload.get("courses_catalog",[]) if r.get("curso")]
    if courses:
        _section(ws,23,"CURSOS NO SEMESTRE DE REFERÊNCIA",end_col=12)
        calc=wb["CALC"]; first=31; last=30+len(courses)
        chart=BarChart(); chart.type="bar"; chart.grouping="stacked"; chart.overlap=100; chart.title="NPS institucional por curso"; chart.height=9; chart.width=20
        data=Reference(calc,min_col=6,max_col=9,min_row=30,max_row=last); cats=Reference(calc,min_col=1,min_row=first,max_row=last)
        chart.add_data(data,titles_from_data=True); chart.set_categories(cats); chart.legend.position="b"; chart.x_axis.scaling.min=-100; chart.x_axis.scaling.max=100
        colors=[GREEN,ORANGE,RED,GRAY_600]
        for ser,color in zip(chart.series,colors): ser.graphicalProperties.solidFill=color
        chart.dataLabels=DataLabelList(); chart.dataLabels.showVal=True
        ws.add_chart(chart,"A25")


def _write_matrix(wb: Workbook, payload: dict[str, Any]) -> None:
    ws=wb.create_sheet("MATRIZ"); _setup(ws,freeze="B6",zoom=80)
    _title(ws,"MATRIZ · CURSO × SEMESTRE","Troque o KPI em PARAMETROS. A matriz mostra até 12 semestres terminando na referência, além da comparação, variação e meta.",end_col=18)
    ws["A4"]="KPI selecionado"; ws["B4"]="=P_MATRIX"; ws["B4"].fill=PatternFill("solid",fgColor=YELLOW_INPUT); ws["B4"].font=Font(name="Arial",size=10,bold=True)
    calc=wb["CALC"]
    courses=[r.get("curso") for r in payload.get("courses_catalog",[]) if r.get("curso")]
    start=31+max(1,len(courses))+3
    for c in range(1,19): ws.cell(6,c,f'=CALC!{get_column_letter(c)}{start+1}')
    for r_idx,_ in enumerate(courses,7):
        src=start+2+(r_idx-7)
        for c in range(1,19): ws.cell(r_idx,c,f'=CALC!{get_column_letter(c)}{src}')
    for c in range(1,19): _header(ws.cell(6,c))
    for r in range(7,7+len(courses)):
        for c in range(1,19): _body(ws.cell(r,c),formula=True)
        _status_fill(ws.cell(r,18),formula=True)
    if courses:
        _add_status_cf(ws, f"R7:R{6+len(courses)}")
    ws.column_dimensions["A"].width=30
    for c in range(2,19): ws.column_dimensions[get_column_letter(c)].width=12


def _write_goals_actions(wb: Workbook, payload: dict[str, Any]) -> None:
    ws=wb.create_sheet("METAS E PLANOS"); _setup(ws,freeze="A5",zoom=85)
    _title(ws,"METAS E PLANOS","Metas e planos são exportados do banco oficial. Alterações devem ser feitas no Data UNIVC, não nesta planilha.",end_col=10)
    _section(ws,4,"METAS CADASTRADAS",end_col=10)
    headers=["KPI","Recorte","Vigência","Meta","Atenção","Limite superior","Justificativa"]
    for c,h in enumerate(headers,1): _header(ws.cell(5,c,h))
    row=6
    for item in payload.get("goals",[]):
        vals=[item.get("indicador"),item.get("recorte"),item.get("vigencia"),item.get("meta"),item.get("atencao"),item.get("limite_superior"),item.get("justificativa")]
        for c,v in enumerate(vals,1): ws.cell(row,c,v); _body(ws.cell(row,c),imported=True)
        row+=1
    row+=2; _section(ws,row,"PLANOS DE AÇÃO",end_col=10); row+=1
    headers2=["KPI","Recorte","Ação","Responsável","Prazo","Status","Resultado esperado"]
    for c,h in enumerate(headers2,1): _header(ws.cell(row,c,h))
    row+=1
    for item in payload.get("actions",[]):
        vals=[item.get("indicador"),item.get("recorte") or item.get("escopo"),item.get("acao") or item.get("plano") or item.get("descricao"),item.get("responsavel"),item.get("prazo"),item.get("status"),item.get("resultado_esperado")]
        for c,v in enumerate(vals,1): ws.cell(row,c,v); _body(ws.cell(row,c),imported=True)
        row+=1
    _auto_width(ws)


def _write_quality(wb: Workbook, payload: dict[str, Any]) -> None:
    ws=wb.create_sheet("QUALIDADE E GOVERNANCA"); _setup(ws,zoom=95)
    _title(ws,"QUALIDADE E GOVERNANÇA","Controles para interpretar o arquivo com segurança antes de tomar decisão.",end_col=8)
    _section(ws,4,"VALIDAÇÕES AUTOMÁTICAS",end_col=8)
    checks=[
        ("NPS curso · consistência",payload["quality"]["nps_inconsistent_rows"],"Respondentes devem ser Promotores + Neutros + Detratores."),
        ("Resultados · consistência",payload["quality"]["result_inconsistent_rows"],"Finalizados devem ser a soma de aprovados e reprovações."),
        ("Semestres exportados",payload["quality"]["periods"],"Quantidade de períodos disponíveis para análise."),
        ("Cursos ativos",payload["quality"]["courses"],"Catálogo ativo no momento da exportação."),
        ("Disciplinas ativas",payload["quality"]["disciplines"],"Catálogo ativo no momento da exportação."),
    ]
    for c,h in enumerate(["Controle","Resultado","Leitura"],1): _header(ws.cell(5,c,h))
    for r,(label,value,note) in enumerate(checks,6):
        ws.cell(r,1,label); ws.cell(r,2,value); ws.cell(r,3,note)
        for c in range(1,4): _body(ws.cell(r,c), imported=(c==2))
        if r in {6,7}: ws.cell(r,2).fill=PatternFill("solid",fgColor=GREEN_PALE if value==0 else RED_LIGHT)
    _section(ws,13,"DIAGNÓSTICO DO RECORTE ATUAL",end_col=8)
    for r,label,formula in [
        (14,"Referência",'=PARAMETROS!B21'),(15,"Comparação",'=PARAMETROS!B22'),(16,"Curso",'=PARAMETROS!B24'),(17,"Disciplina",'=PARAMETROS!B25'),(18,"Janela",'=PARAMETROS!B26')
    ]:
        ws.cell(r,1,label); ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=5); ws.cell(r,2,formula); ws.cell(r,2).fill=PatternFill("solid",fgColor=GRAY_100); ws.cell(r,2).font=Font(name="Arial",size=9,bold=True)
    _section(ws,21,"REGRAS DE USO",end_col=8)
    rules=[
        "O Excel é uma interface de consulta; a fonte oficial continua sendo o Data UNIVC.",
        "O benchmark NPS geral da UNIVC ignora o filtro de curso por definição.",
        "A matriz é comparativa por curso; por isso o filtro de curso não remove os demais cursos nela.",
        "Metas e planos devem ser alterados no sistema e reexportados.",
    ]
    for i,text in enumerate(rules,22): ws.merge_cells(start_row=i,start_column=1,end_row=i,end_column=8); ws.cell(i,1,text); ws.cell(i,1).alignment=Alignment(wrap_text=True); ws.cell(i,1).font=Font(name="Arial",size=9,color=INK)
    ws.column_dimensions["A"].width=32; ws.column_dimensions["B"].width=18; ws.column_dimensions["C"].width=70


def build_academic_interactive_workbook(payload: dict[str, Any]) -> BytesIO:
    if payload.get("directorate") not in {"DTNH", "DCS"}:
        raise ValueError("O Excel Interativo V3 acadêmico está disponível apenas para DTNH/DCS.")
    wb=Workbook(); wb.remove(wb.active)
    ends=_write_technical_data(wb,payload)
    _write_lists(wb,payload)
    _write_parameters(wb,payload)
    _write_calc(wb,payload,ends)
    _write_readme(wb,payload)
    _write_panel(wb,payload,ends)
    _write_nps_institution(wb,payload)
    _write_indicator_sheet(wb,"NPS CURSO","NPS DO CURSO (01B)",f"Evolução do curso em foco. Com '(todos)', agrega as respostas dos cursos da {payload['directorate']}.",5,summary_row=22,y_min=-100,y_max=100,fmt="0.0",table_sheet="DADOS_01B")
    _write_indicator_sheet(wb,"NPS DOCENTES","NPS DA INSTITUIÇÃO · DOCENTES (01C)","População institucional anônima; o filtro de curso não altera este indicador.",6,summary_row=23,y_min=-100,y_max=100,fmt="0.0",table_sheet="DADOS_01C")
    _write_indicator_sheet(wb,"AVALIACAO DOCENTE","AVALIAÇÃO DOCENTE PELO ALUNO (02)","Média ponderada por respondentes, com recorte por curso e disciplina.",7,summary_row=24,y_min=0,y_max=10,fmt="0.00",table_sheet="DADOS_02")
    _write_indicator_sheet(wb,"APROVACAO RESULTADOS","APROVAÇÃO E RESULTADOS (03)","Taxa de aprovação calculada sobre registros finalizados no curso/disciplina selecionados.",8,summary_row=25,y_min=0,y_max=100,fmt='0.0"%"',table_sheet="DADOS_03")
    _write_matrix(wb,payload)
    _write_goals_actions(wb,payload)
    _write_quality(wb,payload)

    # Reorder visible sheets first, technical sheets after them.
    order = VISIBLE_SHEETS + [s for s in wb.sheetnames if s not in VISIBLE_SHEETS]
    wb._sheets = [wb[name] for name in order]
    for ws in wb.worksheets:
        if ws.title not in VISIBLE_SHEETS:
            ws.sheet_state="hidden"

    # Workbook calculation and metadata.
    try:
        wb.calculation.fullCalcOnLoad=True
        wb.calculation.forceFullCalc=True
        wb.calculation.calcMode="auto"
    except Exception:
        pass
    wb.properties.title=f"Data UNIVC · Excel Interativo {payload['directorate']}"
    wb.properties.subject="Painel acadêmico interativo exportado do Data UNIVC"
    wb.properties.creator="Data UNIVC"
    wb.properties.description=f"Excel Interativo V3 beta · {payload['directorate']}"
    buffer=BytesIO(); wb.save(buffer); wb.close(); buffer.seek(0); return buffer


def build_academic_interactive_workbook_bytes(repo, **kwargs) -> BytesIO:
    return build_academic_interactive_workbook(build_academic_interactive_payload(repo, **kwargs))
