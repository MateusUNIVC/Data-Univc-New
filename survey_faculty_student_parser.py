from __future__ import annotations

import io
import re
import warnings
from collections import Counter, OrderedDict
from pathlib import Path
from typing import BinaryIO

from openpyxl import load_workbook

from survey_faculty_models import ParsedFacultyContext
from survey_models import OptionAggregate, ParsedQuestion
from survey_parser import (
    clean_text,
    detect_metric_type,
    infer_modality,
    normalize_key,
    parse_count,
    parse_date,
    parse_percentage,
)


PERIOD_RE = re.compile(
    r"Per[ií]odo\s+de\s+(\d{1,2}/\d{1,2}/\d{2,4})\s+at[eé]\s+(\d{1,2}/\d{1,2}/\d{2,4})",
    flags=re.IGNORECASE,
)
GENERATED_RE = re.compile(r"\d{1,2}/\d{1,2}/\d{4}\s+\d{1,2}[.:]\d{2}[.:]\d{2}")
REPORT_ID_RE = re.compile(r"/(\d+)(?:_[^/]*)?\.xlsx$", flags=re.IGNORECASE)


EXPLICIT_SEMESTER_RE = re.compile(r"\b(20\d{2})\s*[.\-_/]\s*([12])\b")


def _explicit_semester(title: str | None) -> str | None:
    match = EXPLICIT_SEMESTER_RE.search(str(title or ""))
    return f"{match.group(1)}.{match.group(2)}" if match else None


def _load(source: bytes | BinaryIO | str | Path):
    warnings.filterwarnings(
        "ignore",
        message="Workbook contains no default style",
        category=UserWarning,
    )
    if isinstance(source, bytes):
        source = io.BytesIO(source)
    return load_workbook(source, data_only=True, read_only=False)


def _merge_option(question: ParsedQuestion, option: OptionAggregate) -> None:
    key = normalize_key(option.label)
    existing = next((item for item in question.options if normalize_key(item.label) == key), None)
    if existing is None:
        question.options.append(option)
        return

    # O SEI repete cabeçalhos/perguntas em quebras de página. Uma alternativa
    # literalmente repetida com os mesmos valores é repetição visual, não uma
    # segunda resposta.
    same_count = existing.count == option.count
    same_pct = (
        existing.source_percentage == option.source_percentage
        or (
            existing.source_percentage is not None
            and option.source_percentage is not None
            and abs(existing.source_percentage - option.source_percentage) < 0.0001
        )
    )
    if same_count and same_pct:
        return

    existing.count += option.count
    if option.source_percentage is not None:
        if existing.source_percentage is None:
            existing.source_percentage = option.source_percentage
        else:
            existing.source_percentage += option.source_percentage


def _source_report_id(source_path: str) -> str | None:
    normalized = str(source_path or "").replace("\\\\", "\\")
    # ZIP usa / como separador; a unidade do SEI pode conter uma barra invertida
    # literal em "SÃO MATEUS\\ES", portanto nunca dividimos em "\\".
    match = REPORT_ID_RE.search(normalized)
    return match.group(1) if match else None


def _metadata_from_path(source_path: str) -> dict[str, str | None]:
    parts = [part for part in str(source_path or "").split("/") if part]
    if len(parts) < 5:
        return {
            "unit_name": None,
            "course_name": None,
            "teacher_name": None,
            "discipline_name": None,
            "report_id": _source_report_id(source_path),
        }
    return {
        "unit_name": parts[-5].replace("_", " ").strip() or None,
        "course_name": parts[-4].replace("_", " ").strip() or None,
        "teacher_name": parts[-3].replace("_", " ").strip() or None,
        "discipline_name": parts[-2].replace("_", " ").strip() or None,
        "report_id": _source_report_id(source_path),
    }


def parse_faculty_student_workbook(
    source: bytes | BinaryIO | str | Path,
    *,
    source_path: str = "",
) -> ParsedFacultyContext:
    """Interpreta o relatório real ``Disciplina/Professor`` do SEI.

    Layout validado no HAR de 24/09/2026:
    - B: título e período;
    - E/H: Unidade de Ensino e Questionário;
    - L/M: Curso e Professor;
    - E/G: Disciplina;
    - D: pergunta;
    - J: alternativa;
    - P: percentual;
    - S: quantidade;
    - O: data/hora de geração.

    Os metadados são buscados por rótulo; as posições das respostas são fixadas
    pelo relatório real, com fallback semântico de caminho apenas se o XLSX
    omitir algum cabeçalho.
    """

    wb = _load(source)
    ws = wb.active

    title: str | None = None
    period_start: str | None = None
    period_end: str | None = None
    generated_at: str | None = None
    unit_name: str | None = None
    course_name: str | None = None
    questionnaire_name: str | None = None
    discipline_name: str | None = None
    teacher_name: str | None = None

    questions: "OrderedDict[str, ParsedQuestion]" = OrderedDict()
    current_key: str | None = None

    for row_index in range(1, ws.max_row + 1):
        col_b = clean_text(ws.cell(row_index, 2).value)
        col_d = clean_text(ws.cell(row_index, 4).value)
        col_e = normalize_key(clean_text(ws.cell(row_index, 5).value))
        col_g = clean_text(ws.cell(row_index, 7).value)
        col_h = clean_text(ws.cell(row_index, 8).value)
        col_j = clean_text(ws.cell(row_index, 10).value)
        col_l = normalize_key(clean_text(ws.cell(row_index, 12).value))
        col_m = clean_text(ws.cell(row_index, 13).value)
        col_o = clean_text(ws.cell(row_index, 15).value)
        pct_raw = ws.cell(row_index, 16).value
        count_raw = ws.cell(row_index, 19).value

        if not generated_at and col_o and GENERATED_RE.fullmatch(col_o):
            generated_at = col_o

        if col_b:
            period_match = PERIOD_RE.search(col_b)
            if period_match:
                period_start = period_start or parse_date(period_match.group(1))
                period_end = period_end or parse_date(period_match.group(2))
            elif not title and normalize_key(col_b) != "avaliacao institucional" and not normalize_key(col_b).startswith("pag"):
                title = col_b

        if col_e == "unidade ensino" and col_h:
            unit_name = unit_name or col_h
        elif col_e == "questionario" and col_h:
            questionnaire_name = questionnaire_name or col_h
        elif col_e == "disciplina" and col_g:
            discipline_name = discipline_name or col_g

        if col_l == "curso" and col_m:
            course_name = course_name or col_m
        elif col_l == "professor" and col_m:
            teacher_name = teacher_name or col_m

        if col_d:
            current_key = normalize_key(col_d)
            if current_key not in questions:
                questions[current_key] = ParsedQuestion(
                    text=col_d,
                    normalized_text=current_key,
                    position=len(questions) + 1,
                )

        if current_key and col_j:
            count = parse_count(count_raw)
            pct = parse_percentage(pct_raw)
            if count is None and pct is None:
                questions[current_key].raw_responses.append(col_j)
                continue
            _merge_option(
                questions[current_key],
                OptionAggregate(
                    label=col_j,
                    count=max(0, int(count or 0)),
                    source_percentage=pct,
                ),
            )

    path_meta = _metadata_from_path(source_path)
    unit_name = unit_name or path_meta.get("unit_name")
    course_name = course_name or path_meta.get("course_name")
    teacher_name = teacher_name or path_meta.get("teacher_name")
    discipline_name = discipline_name or path_meta.get("discipline_name")

    missing = [
        label
        for label, value in (
            ("Unidade Ensino", unit_name),
            ("Curso", course_name),
            ("Disciplina", discipline_name),
            ("Professor", teacher_name),
        )
        if not clean_text(value)
    ]
    if missing:
        raise ValueError(
            "Relatório docente sem identificação obrigatória: " + ", ".join(missing)
        )

    parsed_questions = list(questions.values())
    if not parsed_questions:
        raise ValueError(
            f"Não foi possível identificar perguntas no relatório docente {source_path or 'XLSX'}."
        )

    for position, question in enumerate(parsed_questions, start=1):
        question.position = position
        metric_type, nps_candidate = detect_metric_type(question.options)
        if not question.options and question.raw_responses:
            metric_type, nps_candidate = "text", False
        question.metric_type = metric_type
        # Uma avaliação docente pelo discente nunca é promovida implicitamente
        # a NPS, mesmo que uma pergunta futura use 0-10.
        question.nps_candidate = False

    totals = [sum(option.count for option in q.options) for q in parsed_questions if q.options]
    if totals:
        respondent_count = Counter(totals).most_common(1)[0][0]
    else:
        raw_totals = [len(q.raw_responses) for q in parsed_questions if q.raw_responses]
        respondent_count = Counter(raw_totals).most_common(1)[0][0] if raw_totals else 0

    report_id = path_meta.get("report_id")
    stable_parts = [
        normalize_key(str(unit_name or "")),
        normalize_key(str(course_name or "")),
        normalize_key(str(teacher_name or "")),
        normalize_key(str(discipline_name or "")),
        str(report_id or ""),
    ]
    source_key = "|".join(stable_parts)
    if len(source_key) > 500:
        source_key = source_key[:500]

    # A data de aplicação pode ocorrer após o encerramento letivo (ex.: julho).
    # Só inferimos o semestre quando ele estiver explicitamente escrito no título.
    semester = _explicit_semester(title)
    return ParsedFacultyContext(
        source_key=source_key,
        source_path=source_path or "relatorio_docente.xlsx",
        unit_name=unit_name,
        course_name=clean_text(course_name),
        modality=infer_modality(str(course_name or "")),
        teacher_name=clean_text(teacher_name),
        discipline_name=clean_text(discipline_name),
        class_code="",
        survey_title=title,
        questionnaire_name=questionnaire_name,
        period_start=period_start,
        period_end=period_end,
        semester_suggested=semester,
        generated_at=generated_at,
        respondent_count=respondent_count,
        questions=parsed_questions,
    )


def faculty_source_path_metadata(source_path: str) -> dict[str, str | None]:
    """Metadados rápidos do caminho interno do ZIP do SEI."""
    return _metadata_from_path(source_path)
