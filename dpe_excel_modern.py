from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

# Identidade visual Data UNIVC / DPE.
GREEN_DARK = "045233"
GREEN = "0B7A54"
GREEN_LIGHT = "E2F0D9"
TEAL = "398C78"
TEAL_LIGHT = "DDEFE9"
GOLD = "D9B44A"
ORANGE_LIGHT = "FFF4D8"
RED = "C74B50"
RED_LIGHT = "FDE7E7"
BLUE = "4679A6"
GRAY = "F5F6F5"
GRAY_2 = "EEF1EF"
GRAY_TEXT = "5A5A5A"
BORDER = "D9E1DD"
WHITE = "FFFFFF"
BLACK = "000000"
LINKED_GREEN = "008000"
INPUT_BLUE = "0000FF"
PURPLE = "7030A0"

CURRENCY_FMT = 'R$ #,##0.00;[Red](R$ #,##0.00);-'
PERCENT_FMT = '0.0%;[Red](0.0%);-'
PERCENT_POINTS_FMT = '0.0"%";[Red](0.0"%");-'
NUMBER_FMT = '#,##0.00;[Red](#,##0.00);-'
INTEGER_FMT = '#,##0;[Red](#,##0);-'
DATE_FMT = 'dd/mm/yyyy'
DATETIME_FMT = 'dd/mm/yyyy hh:mm'

SHEET_ORDER = [
    "PAINEL",
    "RESULTADO",
    "CURSOS",
    "RECEITAS",
    "DESPESAS",
    "DOCENCIA",
    "RATEIOS",
    "METAS_PLANOS",
    "QUALIDADE",
    "BASE_CURSOS",
    "BASE_RECEITAS",
    "BASE_DESPESAS",
    "BASE_DOCENCIA",
    "BASE_RATEIOS",
    "DICIONARIO",
    "PARAMETROS",
]

STATUS_LABELS = {
    "DRAFT": "Preparação",
    "REVIEW": "Conferência",
    "CALCULATED": "Calculada",
    "CLOSED": "Fechada",
    "PASS": "OK",
    "WARNING": "Atenção",
    "BLOCKER": "Bloqueio",
    "GOOD": "Dentro da meta",
    "ATTENTION": "Atenção",
    "BAD": "Fora da meta",
    "UNAVAILABLE": "Indisponível",
    "INACTIVE": "Fora da vigência",
    "HISTORICAL": "Meta histórica",
}


def _as_float(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _iso_date(value: Any) -> Any:
    if not value:
        return None
    if isinstance(value, (datetime,)):
        return value.replace(tzinfo=None)
    text = str(value)
    try:
        if "T" in text:
            return datetime.fromisoformat(text.replace("Z", "+00:00")).replace(tzinfo=None)
        return datetime.fromisoformat(text).date()
    except ValueError:
        return text


def _safe_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value)


class DPEModernWorkbookBuilder:
    """Workbook operacional da DPE v0.13.

    Bases visíveis representam os fatos exportados. Painéis e consolidações usam
    fórmulas simples sobre essas bases, evitando manter uma segunda regra de
    negócio dentro do Excel.
    """

    def __init__(self, payload: dict[str, Any]):
        self.data = payload
        self.period = payload.get("selected_period") or {}
        self.period_id = self.period.get("id")
        self.period_label = self.period.get("period") or "Sem competência"
        self.period_status = self.period.get("status") or ""
        self.generated_at = datetime.now(timezone.utc).replace(tzinfo=None)
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.ws: dict[str, Any] = {}
        self.ranges: dict[str, dict[str, Any]] = {}
        self.thin = Side(style="thin", color=BORDER)
        self.medium = Side(style="medium", color=GREEN)
        self.title_font = Font(name="Arial", size=16, bold=True, color=GREEN_DARK)
        self.subtitle_font = Font(name="Arial", size=9, color=GRAY_TEXT)
        self.section_font = Font(name="Arial", size=10, bold=True, color=WHITE)
        self.header_font = Font(name="Arial", size=9, bold=True, color=WHITE)
        self.body_font = Font(name="Arial", size=9, color=BLACK)
        self.linked_font = Font(name="Arial", size=9, color=LINKED_GREEN)
        self.formula_font = Font(name="Arial", size=9, color=BLACK)
        self.static_font = Font(name="Arial", size=9, color=GRAY_TEXT)
        self.control_font = Font(name="Arial", size=9, bold=True, color=PURPLE)
        self._table_counter = 0

    # ------------------------------------------------------------------
    # Generic helpers
    # ------------------------------------------------------------------
    def _new_sheet(self, title: str, *, tab_color: str = GREEN):
        ws = self.wb.create_sheet(title)
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 90
        ws.sheet_properties.tabColor = tab_color
        ws.freeze_panes = "A5"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_margins.left = 0.3
        ws.page_margins.right = 0.3
        ws.page_margins.top = 0.45
        ws.page_margins.bottom = 0.45
        self.ws[title] = ws
        return ws

    def _title(self, ws, title: str, subtitle: str, *, end_col: int = 12):
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
        c = ws.cell(1, 1, title)
        c.font = self.title_font
        c.alignment = Alignment(vertical="center")
        c.border = Border(bottom=self.medium)
        ws.row_dimensions[1].height = 28
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_col)
        c2 = ws.cell(2, 1, subtitle)
        c2.font = self.subtitle_font
        c2.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[2].height = 28

    def _section(self, ws, row: int, text: str, *, start_col: int = 1, end_col: int = 12):
        ws.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=end_col)
        c = ws.cell(row, start_col, text)
        c.fill = PatternFill("solid", fgColor=GREEN_DARK)
        c.font = self.section_font
        c.alignment = Alignment(vertical="center")
        ws.row_dimensions[row].height = 22

    def _table(self, ws, start_row: int, headers: list[str], rows: list[list[Any]], *, name: str, widths: list[float] | None = None, formats: dict[int, str] | None = None, linked: bool = True):
        for col, header in enumerate(headers, 1):
            c = ws.cell(start_row, col, header)
            c.fill = PatternFill("solid", fgColor=GREEN)
            c.font = self.header_font
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for r_idx, row in enumerate(rows, start_row + 1):
            for c_idx, value in enumerate(row, 1):
                c = ws.cell(r_idx, c_idx, value)
                c.font = self.linked_font if linked else self.body_font
                c.alignment = Alignment(vertical="top", wrap_text=True)
                if formats and c_idx in formats and isinstance(value, (int, float)):
                    c.number_format = formats[c_idx]
                elif hasattr(value, "year") and hasattr(value, "month"):
                    c.number_format = DATETIME_FMT if isinstance(value, datetime) else DATE_FMT
        end_row = max(start_row + 1, start_row + len(rows))
        if not rows:
            # Excel tables require at least one data row; use an empty row.
            for c_idx in range(1, len(headers) + 1):
                ws.cell(start_row + 1, c_idx, None)
        ref = f"A{start_row}:{self._col(len(headers))}{end_row}"
        self._table_counter += 1
        table = Table(displayName=f"{name}{self._table_counter}", ref=ref)
        style = TableStyleInfo(name="TableStyleMedium4", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        table.tableStyleInfo = style
        ws.add_table(table)
        if widths:
            for i, width in enumerate(widths, 1):
                ws.column_dimensions[self._col(i)].width = width
        else:
            for i, header in enumerate(headers, 1):
                ws.column_dimensions[self._col(i)].width = min(28, max(12, len(header) + 2))
        self.ranges[name] = {"header_row": start_row, "start_row": start_row + 1, "end_row": end_row, "headers": headers, "sheet": ws.title}
        return self.ranges[name]

    @staticmethod
    def _col(index: int) -> str:
        result = ""
        while index:
            index, rem = divmod(index - 1, 26)
            result = chr(65 + rem) + result
        return result

    def _formula(self, ws, row: int, col: int, formula: str, fmt: str | None = None):
        c = ws.cell(row, col, formula)
        c.font = self.formula_font
        if fmt:
            c.number_format = fmt
        return c

    def _kpi(self, ws, start_col: int, row: int, label: str, formula: str, *, fmt: str = CURRENCY_FMT, note: str = ""):
        end_col = start_col + 2
        ws.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=end_col)
        lab = ws.cell(row, start_col, label)
        lab.fill = PatternFill("solid", fgColor=GRAY_2)
        lab.font = Font(name="Arial", size=9, bold=True, color=GREEN_DARK)
        lab.alignment = Alignment(horizontal="left", vertical="center")
        ws.merge_cells(start_row=row + 1, start_column=start_col, end_row=row + 2, end_column=end_col)
        val = ws.cell(row + 1, start_col, formula)
        val.font = Font(name="Arial", size=16, bold=True, color=BLACK)
        val.number_format = fmt
        val.alignment = Alignment(horizontal="left", vertical="center")
        ws.merge_cells(start_row=row + 3, start_column=start_col, end_row=row + 3, end_column=end_col)
        n = ws.cell(row + 3, start_col, note)
        n.font = Font(name="Arial", size=8, color=GRAY_TEXT)
        n.alignment = Alignment(wrap_text=True, vertical="top")
        for rr in range(row, row + 4):
            for cc in range(start_col, end_col + 1):
                ws.cell(rr, cc).border = Border(bottom=Side(style="thin", color=BORDER)) if rr == row + 3 else Border()
        return val

    # ------------------------------------------------------------------
    # Raw bases
    # ------------------------------------------------------------------
    def _base_courses(self):
        ws = self._new_sheet("BASE_CURSOS", tab_color=TEAL)
        self._title(ws, "BASE_CURSOS", f"Fatos de curso/contexto · competência {self.period_label}", end_col=13)
        rows = []
        for item in self.data.get("overview", {}).get("offerings", []) or []:
            rows.append([
                item.get("period_offering_id"), item.get("product"), item.get("label"), item.get("modality"), item.get("shift"), item.get("location"),
                item.get("active_students"), item.get("students_ready"), item.get("revenue_ready"), item.get("cost_available"),
                item.get("revenue"), item.get("allocated_cost"), item.get("economic_result"),
            ])
        self._table(ws, 4,
            ["Contexto ID", "Curso", "Curso / contexto", "Modalidade", "Turno", "Local", "Alunos ativos", "Alunos confirmados", "Receita confirmada", "Custo disponível", "Receita atual", "Custo atual", "Resultado atual"],
            rows, name="tbBaseCursos", widths=[12,28,42,16,14,20,14,16,16,16,16,16,16], formats={7:INTEGER_FMT,11:CURRENCY_FMT,12:CURRENCY_FMT,13:CURRENCY_FMT})

    def _base_revenues(self):
        ws = self._new_sheet("BASE_RECEITAS", tab_color=TEAL)
        self._title(ws, "BASE_RECEITAS", f"Ledger de receitas · competência {self.period_label}", end_col=14)
        rows = []
        for item in self.data.get("revenue_entries", []) or []:
            rows.append([
                item.get("id"), self.period_label, item.get("category_code"), item.get("category_name"), item.get("category_scope"), item.get("description"),
                item.get("amount"), item.get("period_offering_id"), item.get("offering_label"), item.get("course_name"), item.get("source_type"), item.get("source_reference"), item.get("notes"), _iso_date(item.get("created_at")),
            ])
        self._table(ws, 4,
            ["ID", "Competência", "Categoria código", "Categoria", "Escopo", "Descrição", "Valor", "Contexto ID", "Curso / contexto", "Curso", "Origem", "Referência", "Observações", "Criado em"],
            rows, name="tbBaseReceitas", widths=[9,13,20,24,16,34,16,12,42,28,14,22,36,18], formats={7:CURRENCY_FMT})

    def _base_expenses(self):
        ws = self._new_sheet("BASE_DESPESAS", tab_color=TEAL)
        self._title(ws, "BASE_DESPESAS", f"Ledger oficial de despesas · competência {self.period_label}", end_col=19)
        rows = []
        for item in self.data.get("expenses", {}).get("expenses", []) or []:
            rows.append([
                item.get("id"), item.get("period"), _iso_date(item.get("expense_date")), item.get("description"), item.get("amount"), item.get("expense_kind"), item.get("expense_scope"),
                item.get("direct_period_offering_id"), item.get("direct_destination_label"), item.get("cost_center_code"), item.get("cost_center_name"), item.get("category_code"), item.get("category_name"),
                item.get("allocation_rule_code"), item.get("allocation_rule_name"), item.get("allocation_driver"), item.get("counterparty_name"), item.get("document_number"), item.get("source_type"),
            ])
        self._table(ws, 4,
            ["ID", "Competência", "Data", "Descrição", "Valor", "Natureza", "Tratamento", "Contexto direto ID", "Destino direto", "Centro código", "Centro de custo", "Categoria código", "Categoria", "Regra código", "Forma de distribuição", "Base", "Favorecido", "Documento", "Origem"],
            rows, name="tbBaseDespesas", widths=[9,13,12,38,16,14,16,14,38,16,24,18,24,16,28,18,28,18,14], formats={5:CURRENCY_FMT})

    def _base_teaching(self):
        ws = self._new_sheet("BASE_DOCENCIA", tab_color=TEAL)
        self._title(ws, "BASE_DOCENCIA", f"Atividades docentes e carga por curso/contexto · competência {self.period_label}", end_col=16)
        relationship = {int(row.get("teacher_id")): row.get("relationship_type") for row in self.data.get("teaching", {}).get("period_teachers", []) or [] if row.get("teacher_id") is not None}
        rows = []
        for activity in self.data.get("teaching", {}).get("activities", []) or []:
            allocations = activity.get("allocations") or []
            if not allocations:
                allocations = [{}]
            for alloc in allocations:
                rows.append([
                    activity.get("id"), activity.get("teacher_id"), activity.get("teacher_name"), relationship.get(int(activity.get("teacher_id") or 0)),
                    activity.get("subject_id"), activity.get("subject_name"), activity.get("class_group"), _iso_date(activity.get("effective_start_date")), _iso_date(activity.get("effective_end_date")),
                    activity.get("workload_hours"), alloc.get("period_offering_id"), alloc.get("label"), alloc.get("allocated_hours"), activity.get("source_type"), activity.get("status"), activity.get("notes"),
                ])
        self._table(ws, 4,
            ["Atividade ID", "Docente ID", "Docente", "Vínculo", "Disciplina ID", "Disciplina", "Turma", "Início", "Fim", "Carga total da atividade", "Contexto ID", "Curso / contexto", "Horas atribuídas", "Origem", "Status", "Observações"],
            rows, name="tbBaseDocencia", widths=[12,11,28,18,12,34,18,12,12,18,12,42,16,14,12,36], formats={10:NUMBER_FMT,13:NUMBER_FMT})

    def _base_allocations(self):
        ws = self._new_sheet("BASE_RATEIOS", tab_color=TEAL)
        self._title(ws, "BASE_RATEIOS", f"Memória do cálculo oficial · competência {self.period_label}", end_col=14)
        selected = self.data.get("allocation", {}).get("selected_run") or {}
        rows = []
        for item in selected.get("results", []) or []:
            rows.append([
                selected.get("id"), selected.get("run_number"), selected.get("status"), item.get("expense_id"), item.get("expense_description"), item.get("period_offering_id"), item.get("offering_label"),
                item.get("driver_type"), item.get("allocated_amount"), (_as_float(item.get("percentage")) or 0) / 100.0, item.get("numerator"), item.get("denominator"), (item.get("basis") or {}).get("basis_type"),
                selected.get("input_fingerprint"),
            ])
        self._table(ws, 4,
            ["Run ID", "Versão", "Status", "Despesa ID", "Despesa", "Contexto ID", "Curso / contexto", "Base de distribuição", "Valor atribuído", "Percentual", "Numerador", "Denominador", "Tipo da base", "Fingerprint"],
            rows, name="tbBaseRateios", widths=[10,9,12,11,38,12,42,22,16,14,14,14,20,28], formats={9:CURRENCY_FMT,10:PERCENT_FMT,11:NUMBER_FMT,12:NUMBER_FMT})

    # ------------------------------------------------------------------
    # Derived result / management sheets
    # ------------------------------------------------------------------
    def _resultado(self):
        ws = self._new_sheet("RESULTADO", tab_color=GREEN)
        self._title(ws, "RESULTADO POR CURSO / CONTEXTO", f"Consolidação por fórmulas sobre as bases exportadas · {self.period_label}", end_col=11)
        base = self.ranges["tbBaseCursos"]
        rev = self.ranges["tbBaseReceitas"]
        rat = self.ranges["tbBaseRateios"]
        rows = self.data.get("overview", {}).get("offerings", []) or []
        headers = ["Contexto ID", "Curso", "Curso / contexto", "Alunos", "Receita", "Custo atribuído", "Resultado", "Margem", "Receita/aluno", "Custo/aluno", "Situação"]
        start = 4
        for col, h in enumerate(headers,1):
            c=ws.cell(start,col,h);c.fill=PatternFill("solid",fgColor=GREEN);c.font=self.header_font;c.alignment=Alignment(wrap_text=True)
        for idx, item in enumerate(rows, start+1):
            pid=item.get("period_offering_id")
            ws.cell(idx,1,pid).font=self.linked_font
            ws.cell(idx,2,item.get("product")).font=self.linked_font
            ws.cell(idx,3,item.get("label")).font=self.linked_font
            # Alunos linked from base courses, not recomputed.
            ws.cell(idx,4,item.get("active_students")).font=self.linked_font; ws.cell(idx,4).number_format=INTEGER_FMT
            rev_formula = f'=SUMIFS(\'BASE_RECEITAS\'!$G${rev["start_row"]}:$G${rev["end_row"]},\'BASE_RECEITAS\'!$H${rev["start_row"]}:$H${rev["end_row"]},A{idx})'
            cost_formula = f'=SUMIFS(\'BASE_RATEIOS\'!$I${rat["start_row"]}:$I${rat["end_row"]},\'BASE_RATEIOS\'!$F${rat["start_row"]}:$F${rat["end_row"]},A{idx})'
            self._formula(ws,idx,5,rev_formula,CURRENCY_FMT)
            self._formula(ws,idx,6,cost_formula,CURRENCY_FMT)
            self._formula(ws,idx,7,f'=E{idx}-F{idx}',CURRENCY_FMT)
            self._formula(ws,idx,8,f'=IFERROR(G{idx}/E{idx},0)',PERCENT_FMT)
            self._formula(ws,idx,9,f'=IFERROR(E{idx}/D{idx},0)',CURRENCY_FMT)
            self._formula(ws,idx,10,f'=IFERROR(F{idx}/D{idx},0)',CURRENCY_FMT)
            self._formula(ws,idx,11,f'=IF(E{idx}=0,"SEM RECEITA",IF(G{idx}<0,"NEGATIVO",IF(H{idx}<0.1,"ATENÇÃO","POSITIVO")))')
        end=max(start+1,start+len(rows))
        if not rows:
            for col in range(1,len(headers)+1):ws.cell(start+1,col,None)
        tab=Table(displayName="tbResultadoContexto",ref=f"A{start}:K{end}");tab.tableStyleInfo=TableStyleInfo(name="TableStyleMedium4",showRowStripes=True);ws.add_table(tab)
        for col,width in enumerate([12,28,42,12,16,16,16,14,16,16,14],1):ws.column_dimensions[self._col(col)].width=width
        ws.freeze_panes="A5"
        self.ranges["tbResultadoContexto"]={"sheet":"RESULTADO","header_row":start,"start_row":start+1,"end_row":end,"headers":headers}

    def _cursos(self):
        ws = self._new_sheet("CURSOS", tab_color=GREEN)
        self._title(ws, "RESULTADO CONSOLIDADO POR CURSO", f"Agrupa todos os contextos do mesmo curso · {self.period_label}", end_col=10)
        courses = self.data.get("analytics", {}).get("courses", []) or []
        result = self.ranges["tbResultadoContexto"]
        headers=["Curso","Contextos","Alunos","Receita","Custo","Resultado","Margem","Receita/aluno","Custo/aluno","Leitura"]
        start=4
        for col,h in enumerate(headers,1):c=ws.cell(start,col,h);c.fill=PatternFill("solid",fgColor=GREEN);c.font=self.header_font
        for r,item in enumerate(courses,start+1):
            course=item.get("course")
            ws.cell(r,1,course).font=self.linked_font
            ws.cell(r,2,item.get("offering_count")).font=self.linked_font;ws.cell(r,2).number_format=INTEGER_FMT
            self._formula(ws,r,3,f'=SUMIF(RESULTADO!$B${result["start_row"]}:$B${result["end_row"]},A{r},RESULTADO!$D${result["start_row"]}:$D${result["end_row"]})',INTEGER_FMT)
            self._formula(ws,r,4,f'=SUMIF(RESULTADO!$B${result["start_row"]}:$B${result["end_row"]},A{r},RESULTADO!$E${result["start_row"]}:$E${result["end_row"]})',CURRENCY_FMT)
            self._formula(ws,r,5,f'=SUMIF(RESULTADO!$B${result["start_row"]}:$B${result["end_row"]},A{r},RESULTADO!$F${result["start_row"]}:$F${result["end_row"]})',CURRENCY_FMT)
            self._formula(ws,r,6,f'=D{r}-E{r}',CURRENCY_FMT)
            self._formula(ws,r,7,f'=IFERROR(F{r}/D{r},0)',PERCENT_FMT)
            self._formula(ws,r,8,f'=IFERROR(D{r}/C{r},0)',CURRENCY_FMT)
            self._formula(ws,r,9,f'=IFERROR(E{r}/C{r},0)',CURRENCY_FMT)
            self._formula(ws,r,10,f'=IF(F{r}<0,"Resultado negativo",IF(G{r}<0.1,"Margem em atenção","Resultado positivo"))')
        end=max(start+1,start+len(courses))
        if not courses:
            for col in range(1,len(headers)+1):ws.cell(start+1,col,None)
        tab=Table(displayName="tbCursos",ref=f"A{start}:J{end}");tab.tableStyleInfo=TableStyleInfo(name="TableStyleMedium4",showRowStripes=True);ws.add_table(tab)
        for col,width in enumerate([30,12,12,16,16,16,14,16,16,22],1):ws.column_dimensions[self._col(col)].width=width
        self.ranges["tbCursos"]={"sheet":"CURSOS","header_row":start,"start_row":start+1,"end_row":end,"headers":headers}

    def _receitas(self):
        ws=self._new_sheet("RECEITAS",tab_color=GREEN)
        self._title(ws,"ANÁLISE DE RECEITAS",f"Receitas de curso e institucionais · {self.period_label}",end_col=8)
        base=self.ranges["tbBaseReceitas"]
        self._section(ws,4,"Resumo",end_col=8)
        labels=[("Receita total",f'=SUM(\'BASE_RECEITAS\'!$G${base["start_row"]}:$G${base["end_row"]})'),("Receita atribuída a cursos",f'=SUMIFS(\'BASE_RECEITAS\'!$G${base["start_row"]}:$G${base["end_row"]},\'BASE_RECEITAS\'!$H${base["start_row"]}:$H${base["end_row"]},">0")'),("Receita institucional",f'=SUMIFS(\'BASE_RECEITAS\'!$G${base["start_row"]}:$G${base["end_row"]},\'BASE_RECEITAS\'!$H${base["start_row"]}:$H${base["end_row"]},"")')]
        for i,(label,formula) in enumerate(labels,6):ws.cell(i,1,label).font=self.static_font;self._formula(ws,i,2,formula,CURRENCY_FMT)
        categories=self.data.get("revenues",{}).get("categories",[]) or []
        self._section(ws,11,"Por categoria",end_col=8)
        headers=["Categoria código","Categoria","Escopo","Valor"]
        for c,h in enumerate(headers,1):ws.cell(12,c,h).fill=PatternFill("solid",fgColor=GREEN);ws.cell(12,c).font=self.header_font
        for r,cat in enumerate(categories,13):
            ws.cell(r,1,cat.get("code")).font=self.linked_font;ws.cell(r,2,cat.get("name")).font=self.linked_font;ws.cell(r,3,cat.get("scope")).font=self.linked_font
            self._formula(ws,r,4,f'=SUMIF(\'BASE_RECEITAS\'!$C${base["start_row"]}:$C${base["end_row"]},A{r},\'BASE_RECEITAS\'!$G${base["start_row"]}:$G${base["end_row"]})',CURRENCY_FMT)
        end=max(13,12+len(categories))
        tab=Table(displayName="tbReceitaCategorias",ref=f"A12:D{end}");tab.tableStyleInfo=TableStyleInfo(name="TableStyleMedium4",showRowStripes=True);ws.add_table(tab)
        for col,width in enumerate([22,28,18,18,2,18,18,18],1):ws.column_dimensions[self._col(col)].width=width

    def _despesas(self):
        ws=self._new_sheet("DESPESAS",tab_color=GREEN)
        self._title(ws,"ANÁLISE DE DESPESAS",f"Despesas oficiais por tratamento e categoria · {self.period_label}",end_col=9)
        base=self.ranges["tbBaseDespesas"]
        self._section(ws,4,"Tratamento econômico",end_col=9)
        headers=["Tratamento","Valor","Participação"]
        for c,h in enumerate(headers,1):ws.cell(5,c,h).fill=PatternFill("solid",fgColor=GREEN);ws.cell(5,c).font=self.header_font
        scopes=[("DIRECT","Direta"),("SHARED","Compartilhada"),("INSTITUTIONAL","Institucional")]
        for r,(code,label) in enumerate(scopes,6):
            ws.cell(r,1,label).font=self.static_font
            self._formula(ws,r,2,f'=SUMIF(\'BASE_DESPESAS\'!$G${base["start_row"]}:$G${base["end_row"]},"{code}",\'BASE_DESPESAS\'!$E${base["start_row"]}:$E${base["end_row"]})',CURRENCY_FMT)
            self._formula(ws,r,3,f'=IFERROR(B{r}/SUM($B$6:$B$8),0)',PERCENT_FMT)
        tab=Table(displayName="tbDespesaTratamento",ref="A5:C8");tab.tableStyleInfo=TableStyleInfo(name="TableStyleMedium4",showRowStripes=True);ws.add_table(tab)
        cats=self.data.get("expenses",{}).get("categories",[]) or []
        self._section(ws,11,"Por categoria",end_col=9)
        for c,h in enumerate(["Categoria código","Categoria","Valor"],1):ws.cell(12,c,h).fill=PatternFill("solid",fgColor=GREEN);ws.cell(12,c).font=self.header_font
        for r,cat in enumerate(cats,13):
            ws.cell(r,1,cat.get("code")).font=self.linked_font;ws.cell(r,2,cat.get("name")).font=self.linked_font
            self._formula(ws,r,3,f'=SUMIF(\'BASE_DESPESAS\'!$L${base["start_row"]}:$L${base["end_row"]},A{r},\'BASE_DESPESAS\'!$E${base["start_row"]}:$E${base["end_row"]})',CURRENCY_FMT)
        end=max(13,12+len(cats));tab2=Table(displayName="tbDespesaCategorias",ref=f"A12:C{end}");tab2.tableStyleInfo=TableStyleInfo(name="TableStyleMedium4",showRowStripes=True);ws.add_table(tab2)
        for col,width in enumerate([28,30,18,2,18,18,18,18,18],1):ws.column_dimensions[self._col(col)].width=width

    def _docencia(self):
        ws=self._new_sheet("DOCENCIA",tab_color=GREEN)
        self._title(ws,"DOCÊNCIA",f"Vínculos, carga e folha conciliada · {self.period_label}",end_col=9)
        rows=[]
        for item in self.data.get("teaching",{}).get("period_teachers",[]) or []:
            rows.append([item.get("teacher_id"),(item.get("teacher") or {}).get("display_name"),item.get("relationship_type"),item.get("workload_hours"),item.get("payroll_total"),item.get("payroll_count"),None])
        rng=self._table(ws,4,["Docente ID","Docente","Vínculo","Carga no mês","Custo conciliado","Lançamentos de folha","Custo / hora"],rows,name="tbDocenciaResumo",widths=[11,30,18,16,18,18,16],formats={4:NUMBER_FMT,5:CURRENCY_FMT,6:INTEGER_FMT})
        for r in range(rng["start_row"],rng["end_row"]+1):
            ws.cell(r,7,f'=IFERROR(E{r}/D{r},0)').font=self.formula_font;ws.cell(r,7).number_format=CURRENCY_FMT

    def _rateios(self):
        ws=self._new_sheet("RATEIOS",tab_color=GREEN)
        self._title(ws,"DISTRIBUIÇÃO DE CUSTOS",f"Resumo do cálculo oficial · {self.period_label}",end_col=9)
        alloc=self.data.get("allocation",{}); selected=alloc.get("selected_run") or {}; summary=selected.get("summary") or {}
        self._section(ws,4,"Execução oficial",end_col=9)
        items=[("Versão",selected.get("run_number")),("Status",selected.get("status")),("Despesas distribuíveis",selected.get("expense_total")),("Valor atribuído",selected.get("allocated_total")),("Valor pendente",selected.get("unallocated_total")),("Resultados gerados",summary.get("result_count"))]
        for r,(label,value) in enumerate(items,6):
            ws.cell(r,1,label).font=self.static_font;ws.cell(r,2,value).font=self.linked_font
            if r in (8,9,10):ws.cell(r,2).number_format=CURRENCY_FMT
        self._section(ws,14,"Custo atribuído por curso/contexto",end_col=9)
        base=self.ranges["tbBaseRateios"]; contexts=self.data.get("overview",{}).get("offerings",[]) or []
        for c,h in enumerate(["Contexto ID","Curso / contexto","Custo atribuído"],1):ws.cell(15,c,h).fill=PatternFill("solid",fgColor=GREEN);ws.cell(15,c).font=self.header_font
        for r,item in enumerate(contexts,16):
            ws.cell(r,1,item.get("period_offering_id")).font=self.linked_font;ws.cell(r,2,item.get("label")).font=self.linked_font
            self._formula(ws,r,3,f'=SUMIF(\'BASE_RATEIOS\'!$F${base["start_row"]}:$F${base["end_row"]},A{r},\'BASE_RATEIOS\'!$I${base["start_row"]}:$I${base["end_row"]})',CURRENCY_FMT)
        end=max(16,15+len(contexts));tab=Table(displayName="tbRateioCursos",ref=f"A15:C{end}");tab.tableStyleInfo=TableStyleInfo(name="TableStyleMedium4",showRowStripes=True);ws.add_table(tab)
        ws.column_dimensions["A"].width=15;ws.column_dimensions["B"].width=44;ws.column_dimensions["C"].width=18

    def _metas_planos(self):
        ws=self._new_sheet("METAS_PLANOS",tab_color=GOLD)
        self._title(ws,"METAS E PLANOS DE AÇÃO",f"Gestão conectada aos fatos reais da DPE · {self.period_label}",end_col=12)
        mg=self.data.get("management",{})
        self._section(ws,4,"Metas",end_col=12)
        target_rows=[]
        for item in mg.get("targets",[]) or []:
            target=item.get("target")
            if target is None and item.get("target_min") is not None:
                target=f'{item.get("target_min")} a {item.get("target_max")}'
            target_rows.append([item.get("indicator_label"),item.get("metric_label"),item.get("dimension_label") or "Institucional",item.get("current_value"),target,(item.get("status") or {}).get("label"),item.get("valid_from"),item.get("valid_to"),item.get("justification")])
        rng=self._table(ws,5,["Indicador","Métrica","Escopo","Atual","Meta","Situação","Vigência inicial","Vigência final","Justificativa"],target_rows,name="tbMetas",widths=[24,30,28,16,18,18,14,14,38])
        row2=rng["end_row"]+3;self._section(ws,row2,"Planos de ação",end_col=12)
        action_rows=[]
        for item in mg.get("actions",[]) or []:
            action_rows.append([item.get("indicator_label"),item.get("metric_label"),item.get("period"),item.get("problem"),item.get("cause"),item.get("corrective_action"),item.get("responsible"),_iso_date(item.get("due_date")),item.get("effective_status"),item.get("evidence")])
        self._table(ws,row2+1,["Indicador","Métrica","Competência","Problema","Causa","Ação corretiva","Responsável","Prazo","Status","Evidência"],action_rows,name="tbPlanos",widths=[22,28,14,36,32,40,24,14,18,30])

    def _qualidade(self):
        ws=self._new_sheet("QUALIDADE",tab_color=GOLD)
        self._title(ws,"QUALIDADE E FECHAMENTO",f"Checklist, alertas e limitações · {self.period_label}",end_col=10)
        closure=self.data.get("closure",{}); checklist=(closure.get("checklist") or {})
        rows=[]
        for item in checklist.get("checks",[]) or []:
            rows.append([item.get("code"),item.get("label"),STATUS_LABELS.get(str(item.get("status")),item.get("status")),item.get("detail")])
        rng=self._table(ws,4,["Código","Verificação","Situação","Detalhe"],rows,name="tbQualidade",widths=[24,32,18,72])
        row=rng["end_row"]+3
        self._section(ws,row,"Limitações / observações analíticas",end_col=10)
        limitations=self.data.get("analytics",{}).get("limitations",[]) or []
        if limitations:
            for idx,text in enumerate(limitations,row+1):ws.cell(idx,1,"• "+text).font=self.static_font;ws.merge_cells(start_row=idx,start_column=1,end_row=idx,end_column=10);ws.cell(idx,1).alignment=Alignment(wrap_text=True)
        else:
            ws.cell(row+1,1,"Nenhuma limitação adicional registrada.").font=self.static_font
        ws.column_dimensions["A"].width=24

    def _parameters(self):
        ws=self._new_sheet("PARAMETROS",tab_color=PURPLE)
        self._title(ws,"PARÂMETROS DA EXPORTAÇÃO","Contexto do arquivo e fonte de verdade. Esta aba não contém regra de negócio.",end_col=8)
        rows=[
            ["Competência",self.period_label,"Período selecionado no Data UNIVC"],
            ["Status",STATUS_LABELS.get(self.period_status,self.period_status),"Estado da competência no momento da exportação"],
            ["Período ID",self.period_id,"Identificador técnico"],
            ["Gerado em",self.generated_at,"Horário UTC da geração"],
            ["Fonte de verdade","DPE Cost Engine / revenue ledger / expense ledger / official allocation run","Mesmas fontes utilizadas pela aplicação"],
            ["Run oficial",(self.data.get("allocation",{}).get("selected_run") or {}).get("id"),"Execução usada na BASE_RATEIOS"],
            ["Versão do run",(self.data.get("allocation",{}).get("selected_run") or {}).get("run_number"),"Número sequencial da distribuição"],
            ["Workbook","DPE v0.13 moderno","Dados derivados diretamente do Cost Engine e dos ledgers oficiais"],
        ]
        self._table(ws,4,["Parâmetro","Valor","Descrição"],rows,name="tbParametros",widths=[26,58,56])
        ws["B8"].number_format=DATETIME_FMT

    def _dictionary(self):
        ws=self._new_sheet("DICIONARIO",tab_color=PURPLE)
        self._title(ws,"DICIONÁRIO DO WORKBOOK","O que cada aba representa e como deve ser utilizada.",end_col=8)
        rows=[
            ["PAINEL","Executiva","Indicadores e gráficos derivados das bases do arquivo.","Não editar fórmulas."],
            ["RESULTADO","Análise","Resultado por curso/contexto calculado por fórmulas.","Receita e custo vêm das bases."],
            ["CURSOS","Análise","Consolidação por curso, somando seus contextos.","Útil para comparação executiva."],
            ["RECEITAS","Análise","Resumo das receitas por tipo e categoria.","BASE_RECEITAS é a origem."],
            ["DESPESAS","Análise","Resumo das despesas por tratamento e categoria.","BASE_DESPESAS é a origem."],
            ["DOCENCIA","Análise","Carga e custo conciliado por docente.","Atividades detalhadas em BASE_DOCENCIA."],
            ["RATEIOS","Análise","Execução oficial e custo atribuído por contexto.","Memória detalhada em BASE_RATEIOS."],
            ["METAS_PLANOS","Gestão","Metas vigentes e planos de ação.","Valores atuais vêm do Cost Engine."],
            ["QUALIDADE","Governança","Checklist de fechamento e limitações.","Use antes de divulgar números."],
            ["BASE_CURSOS","Base","Curso/contexto, alunos e disponibilidade dos fatos.","Tabela filtrável."],
            ["BASE_RECEITAS","Base","Lançamentos do ledger de receitas.","Tabela filtrável; fonte importada."],
            ["BASE_DESPESAS","Base","Despesas oficiais da competência.","Tabela filtrável; fonte importada."],
            ["BASE_DOCENCIA","Base","Atividade docente por curso/contexto.","Uma atividade pode gerar várias linhas."],
            ["BASE_RATEIOS","Base","Uma linha por despesa × contexto no cálculo oficial.","Somar Valor atribuído para custos."],
            ["PARAMETROS","Técnica","Contexto da exportação.","Sem cálculos financeiros."],
        ]
        self._table(ws,4,["Aba","Tipo","Conteúdo","Observação"],rows,name="tbDicionario",widths=[24,18,62,52])

    def _painel(self):
        ws=self._new_sheet("PAINEL",tab_color=GREEN_DARK)
        ws.freeze_panes="A4"
        self._title(ws,"DPE · PAINEL EXECUTIVO",f"Competência {self.period_label} · {STATUS_LABELS.get(self.period_status,self.period_status)} · dados ligados às bases deste arquivo",end_col=14)
        rev=self.ranges["tbBaseReceitas"]; exp=self.ranges["tbBaseDespesas"]; rat=self.ranges["tbBaseRateios"]
        total_rev=f'=SUM(\'BASE_RECEITAS\'!$G${rev["start_row"]}:$G${rev["end_row"]})'
        total_exp=f'=SUM(\'BASE_DESPESAS\'!$E${exp["start_row"]}:$E${exp["end_row"]})'
        inst_rev=f'=SUMIFS(\'BASE_RECEITAS\'!$G${rev["start_row"]}:$G${rev["end_row"]},\'BASE_RECEITAS\'!$H${rev["start_row"]}:$H${rev["end_row"]},"")'
        alloc=f'=SUM(\'BASE_RATEIOS\'!$I${rat["start_row"]}:$I${rat["end_row"]})'
        self._kpi(ws,1,4,"Receita total",total_rev,note="Todos os lançamentos do ledger")
        self._kpi(ws,5,4,"Despesas do mês",total_exp,note="Todas as despesas oficiais")
        self._kpi(ws,9,4,"Resultado institucional",'=A5-E5',note="Receita total menos despesas")
        self._kpi(ws,13,4,"Margem institucional",'=IFERROR(I5/A5,0)',fmt=PERCENT_FMT,note="Resultado / receita total")
        self._kpi(ws,1,9,"Receita institucional",inst_rev,note="Receita sem curso/contexto")
        self._kpi(ws,5,9,"Custos distribuídos",alloc,note="Run oficial")
        self._kpi(ws,9,9,"Cobertura da distribuição",f'=IFERROR(E10/SUMIFS(\'BASE_DESPESAS\'!$E${exp["start_row"]}:$E${exp["end_row"]},\'BASE_DESPESAS\'!$G${exp["start_row"]}:$G${exp["end_row"]},"<>INSTITUTIONAL"),0)',fmt=PERCENT_FMT,note="Distribuído / distribuível")
        self._kpi(ws,13,9,"Alunos ativos",f'=SUM(\'BASE_CURSOS\'!$G${self.ranges["tbBaseCursos"]["start_row"]}:$G${self.ranges["tbBaseCursos"]["end_row"]})',fmt=INTEGER_FMT,note="Soma dos contextos")

        self._section(ws,14,"Resultado por curso",end_col=14)
        courses=self.ranges["tbCursos"]
        if courses["end_row"]>=courses["start_row"]:
            chart=BarChart();chart.type="bar";chart.style=10;chart.title="Receita × custo por curso";chart.y_axis.title="Curso";chart.x_axis.title="R$";chart.height=8;chart.width=13
            data=Reference(self.ws["CURSOS"],min_col=4,max_col=5,min_row=courses["header_row"],max_row=courses["end_row"])
            cats=Reference(self.ws["CURSOS"],min_col=1,min_row=courses["start_row"],max_row=courses["end_row"])
            chart.add_data(data,titles_from_data=True);chart.set_categories(cats);chart.legend.position="b";ws.add_chart(chart,"A16")
        # Treatment doughnut.
        chart2=DoughnutChart();chart2.title="Despesas por tratamento";chart2.height=7;chart2.width=9;chart2.holeSize=55
        data2=Reference(self.ws["DESPESAS"],min_col=2,min_row=5,max_row=8);cats2=Reference(self.ws["DESPESAS"],min_col=1,min_row=6,max_row=8)
        chart2.add_data(data2,titles_from_data=True);chart2.set_categories(cats2);chart2.dataLabels=DataLabelList();chart2.dataLabels.showPercent=True;chart2.legend.position="r";ws.add_chart(chart2,"J16")
        for col,width in enumerate([14]*14,1):ws.column_dimensions[self._col(col)].width=width

    def build(self) -> BytesIO:
        # Bases first so every analytical sheet can point to stable ranges.
        self._base_courses(); self._base_revenues(); self._base_expenses(); self._base_teaching(); self._base_allocations()
        self._resultado(); self._cursos(); self._receitas(); self._despesas(); self._docencia(); self._rateios(); self._metas_planos(); self._qualidade(); self._dictionary(); self._parameters(); self._painel()
        # Canonical visual order.
        self.wb._sheets = [self.ws[name] for name in SHEET_ORDER if name in self.ws]
        self.wb.active = 0
        # Ask Excel-compatible engines to recalculate formulas when opened.
        try:
            self.wb.calculation.fullCalcOnLoad = True
            self.wb.calculation.forceFullCalc = True
            self.wb.calculation.calcMode = "auto"
        except Exception:
            pass
        output=BytesIO();self.wb.save(output);output.seek(0);return output


def build_dpe_operational_workbook(payload: dict[str, Any]) -> BytesIO:
    return DPEModernWorkbookBuilder(payload).build()
