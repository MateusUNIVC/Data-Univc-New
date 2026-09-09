from __future__ import annotations

"""Workbook acadêmico-visual específico da DPE.

A DPE deixou de usar a planilha gerencial genérica da v0.7.0. Este módulo
reproduz a linguagem visual dos painéis DTNH/DCS e entrega uma página própria
para cada indicador, além de bases independentes que podem ser editadas e
reimportadas no sistema.
"""

from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break

from excel_errors import ExcelExportLimitError
from management_catalog import indicator_spec, indicator_specs, metric_spec, period_sort_key, validate_period
from management_service import build_management_dashboard, row_dimensions, raw_components

# Mesma linguagem visual do AcademicWorkbookBuilder (DTNH/DCS).
GREEN_DARK = "045233"
GREEN = "0B7A54"
GREEN_LIGHT = "E2F0D9"
TEAL_LIGHT = "DDEFE9"
YELLOW = "FFF6DC"
GRAY = "F5F6F5"
GRAY_TEXT = "5A5A5A"
BORDER = "D9E1DD"
RED_LIGHT = "FDE7E7"
ORANGE_LIGHT = "FFF4D8"
WHITE = "FFFFFF"
BLACK = "000000"
LINKED_GREEN = "008000"
INPUT_BLUE = "0000FF"
PURPLE = "7030A0"
PALE_GREEN = "F3F8F5"
PALE_TEAL = "E8F3F0"
PALE_BLUE = "EAF2FA"
PALE_GOLD = "FFF4D6"
SOFT_GRAY = "EEF1EF"
CHART_META = "7F8C8D"
CHART_GREEN = "0B7A54"
CHART_TEAL = "398C78"
CHART_GOLD = "D9B44A"
CHART_RED = "C74B50"
CHART_BLUE = "4679A6"
EXCEL_MAX_ROW = 1_048_576
BASE_HEADER_ROW = 4
BASE_DATA_START = 5
INPUT_RESERVE = 60

CURRENCY_FMT = 'R$ #,##0.00;[Red](R$ #,##0.00);-'
PERCENT_FMT = '0.0%;[Red](0.0%);-'
PERCENT_POINTS_FMT = '0.0"%";[Red](0.0"%");-'
CHANGE_POINTS_FMT = '0.0"%";-0.0"%";-'
DECIMAL_FMT = '0.00;[Red](0.00);-'
INTEGER_FMT = '#,##0;[Red](#,##0);-'

BASE_SHEETS = {
    "DPE-01": "BASE DPE-01",
    "DPE-02": "BASE DPE-02",
    "DPE-03": "BASE DPE-03",
}
INDICATOR_SHEETS = {
    "DPE-01": "DPE-01 MARGEM",
    "DPE-02": "DPE-02 COBERTURA",
    "DPE-03": "DPE-03 FOLHA",
}


def _number(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    if number != number or abs(number) == float("inf"):
        return None
    return number


def _previous_month(period: str) -> str:
    year, month = int(period[:4]), int(period[5:7])
    month -= 1
    if month == 0:
        year -= 1
        month = 12
    return f"{year:04d}-{month:02d}"


def _same_month_previous_year(period: str) -> str:
    return f"{int(period[:4]) - 1:04d}-{period[5:7]}"


def _target_for_period(
    targets: list[dict[str, Any]],
    spec: dict[str, Any],
    metric: dict[str, Any],
    period: str,
    *,
    dimension_key: str = "TOTAL",
    course_dimension: bool = False,
) -> dict[str, Any]:
    candidates: list[tuple[int, int, dict[str, Any]]] = []
    period_key = period_sort_key(period)
    for row in targets:
        if row.get("indicator_code") != spec["code"] or row.get("metric_key") != metric["key"]:
            continue
        target_dimension = row.get("dimension_key") or "TOTAL"
        if target_dimension not in {dimension_key, "TOTAL"}:
            continue
        # Para DPE-01 por curso, a meta institucional de 15% não substitui
        # a regra específica de 10%. Apenas uma meta cadastrada para a
        # dimensão exata pode sobrepor esse parâmetro.
        if course_dimension and dimension_key != "TOTAL" and target_dimension == "TOTAL":
            continue
        start = period_sort_key(row.get("valid_from"))
        end = period_sort_key(row.get("valid_to")) if row.get("valid_to") else 10**12
        if start <= period_key <= end:
            candidates.append((1 if target_dimension == dimension_key else 0, start, row))
    if candidates:
        return dict(max(candidates, key=lambda item: (item[0], item[1]))[2])
    fallback = {
        "target": metric.get("target"),
        "attention": metric.get("attention"),
        "target_min": metric.get("target_min"),
        "target_max": metric.get("target_max"),
        "attention_min": metric.get("attention_min"),
        "attention_max": metric.get("attention_max"),
    }
    if spec["code"] == "DPE-01" and metric["key"] == "net_margin_pct" and course_dimension:
        fallback["target"] = 10
        fallback["attention"] = 0
    return fallback


@dataclass
class DPEExcelContext:
    reference: str | None = None
    comparison: str | None = None
    only_indicator: str | None = None
    template_mode: bool = False


class DPEWorkbookBuilder:
    def __init__(
        self,
        measurements: Iterable[dict[str, Any]],
        targets: Iterable[dict[str, Any]] = (),
        actions: Iterable[dict[str, Any]] = (),
        *,
        reference: str | None = None,
        comparison: str | None = None,
        only_indicator: str | None = None,
        template_mode: bool = False,
    ) -> None:
        self.measurements = [dict(row) for row in measurements if str(row.get("indicator_code", "")).startswith("DPE-")]
        self.targets = [dict(row) for row in targets if str(row.get("indicator_code", "")).startswith("DPE-")]
        self.actions = [dict(row) for row in actions if str(row.get("indicator_code", "")).startswith("DPE-")]
        self.context = DPEExcelContext(reference, comparison, only_indicator, template_mode)
        if only_indicator:
            self.context.only_indicator = indicator_spec("DPE", only_indicator)["code"]
        periods = sorted(
            {str(row.get("period")) for row in self.measurements if period_sort_key(row.get("period")) >= 0},
            key=period_sort_key,
        )
        if reference:
            self.reference = validate_period("DPE", reference)
        elif periods:
            self.reference = periods[-1]
        else:
            today = date.today()
            self.reference = f"{today.year:04d}-{today.month:02d}"
        self.comparison = validate_period("DPE", comparison) if comparison else _previous_month(self.reference)
        self.periods = periods or [self.reference]
        if self.reference not in self.periods:
            self.periods.append(self.reference)
            self.periods.sort(key=period_sort_key)
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.ws: dict[str, Any] = {}
        self.title_font = Font(name="Arial", size=15, bold=True, color=GREEN_DARK)
        self.section_font = Font(name="Arial", size=11, bold=True, color=WHITE)
        self.header_font = Font(name="Arial", size=9, bold=True, color=WHITE)
        self.body_font = Font(name="Arial", size=9, color=BLACK)
        self.static_font = Font(name="Arial", size=9, color=GRAY_TEXT)
        self.formula_font = Font(name="Arial", size=9, color=BLACK)
        self.linked_font = Font(name="Arial", size=9, color=LINKED_GREEN)
        self.input_font = Font(name="Arial", size=9, color=INPUT_BLUE)
        self.control_font = Font(name="Arial", size=9, bold=True, color=PURPLE)
        self.border_box = Border(
            left=Side(style="thin", color=BORDER),
            right=Side(style="thin", color=BORDER),
            top=Side(style="thin", color=BORDER),
            bottom=Side(style="thin", color=BORDER),
        )
        self.dashboards: dict[str, dict[str, Any]] = {}
        for code in ("DPE-01", "DPE-02", "DPE-03"):
            self.dashboards[code] = build_management_dashboard(
                "DPE",
                self.measurements,
                self.targets,
                self.actions,
                reference=self.reference,
                comparison=self.comparison,
                indicator_code=code,
            )

    # ------------------------------------------------------------------
    # Styling helpers
    # ------------------------------------------------------------------
    def _new_sheet(self, title: str):
        ws = self.wb.create_sheet(title)
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 85
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_margins.left = 0.3
        ws.page_margins.right = 0.3
        ws.page_margins.top = 0.45
        ws.page_margins.bottom = 0.45
        ws.page_margins.header = 0.15
        ws.page_margins.footer = 0.15
        self.ws[title] = ws
        return ws

    def _title(self, ws, text: str, subtitle: str | None = None, *, end_col: int = 12):
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
        ws["A1"] = text
        ws["A1"].font = self.title_font
        ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
        ws["A1"].border = Border(bottom=Side(style="medium", color=GREEN))
        ws.row_dimensions[1].height = 28
        if subtitle:
            ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_col)
            ws["A2"] = subtitle
            ws["A2"].font = self.static_font
            ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[2].height = 32

    def _section(self, ws, row: int, text: str, *, start_col: int = 1, end_col: int = 12):
        ws.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=end_col)
        cell = ws.cell(row, start_col, text)
        cell.fill = PatternFill("solid", fgColor=GREEN_DARK)
        cell.font = self.section_font
        cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[row].height = 24

    def _headers(self, ws, row: int, labels: Iterable[str], *, start_col: int = 1):
        for offset, label in enumerate(labels):
            cell = ws.cell(row, start_col + offset, label)
            cell.fill = PatternFill("solid", fgColor=GREEN)
            cell.font = self.header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=Side(style="thin", color=GREEN_DARK))
        ws.row_dimensions[row].height = 32

    def _paint_range(self, ws, min_row: int, max_row: int, min_col: int, max_col: int, *, fill: str | None = None):
        for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
            for cell in row:
                if fill:
                    cell.fill = PatternFill("solid", fgColor=fill)
                cell.border = self.border_box

    def _kpi_card(
        self,
        ws,
        *,
        start_col: int,
        end_col: int,
        label_row: int,
        value_row: int,
        label: str,
        formula: Any,
        number_format: str = "General",
        fill: str = PALE_GREEN,
        font_size: float = 16,
        value_color: str = BLACK,
    ):
        self._paint_range(ws, label_row, value_row + 1, start_col, end_col, fill=fill)
        ws.merge_cells(start_row=label_row, start_column=start_col, end_row=label_row, end_column=end_col)
        ws.merge_cells(start_row=value_row, start_column=start_col, end_row=value_row + 1, end_column=end_col)
        label_cell = ws.cell(label_row, start_col, label)
        label_cell.font = Font(name="Arial", size=8, bold=True, color=GRAY_TEXT)
        label_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        value_cell = ws.cell(value_row, start_col, formula)
        value_cell.font = Font(name="Arial", size=font_size, bold=True, color=value_color)
        value_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        value_cell.number_format = number_format
        ws.row_dimensions[label_row].height = 22
        ws.row_dimensions[value_row].height = 24
        ws.row_dimensions[value_row + 1].height = 24

    @staticmethod
    def _chart_base(chart, *, title: str, x_title: str, y_title: str, width: float = 25.0, height: float = 9.0):
        chart.title = title
        chart.style = 13
        chart.width = width
        chart.height = height
        chart.x_axis.title = x_title
        chart.y_axis.title = y_title
        chart.legend.position = "b"
        chart.display_blanks = "gap"
        chart.visible_cells_only = False
        chart.y_axis.majorGridlines = None
        try:
            chart.x_axis.tickLblSkip = 1
            chart.x_axis.tickMarkSkip = 1
        except Exception:
            pass

    @staticmethod
    def _line_style(chart, colors: list[str], *, labels_first: bool = True, number_format: str = "0.0"):
        for index, series in enumerate(chart.series):
            color = colors[index] if index < len(colors) else CHART_META
            series.graphicalProperties.line.solidFill = color
            series.graphicalProperties.line.width = 25000 if index == 0 else 19050
            if index == len(chart.series) - 1 and len(chart.series) >= 3:
                try:
                    series.graphicalProperties.line.prstDash = "dash"
                except Exception:
                    pass
                series.marker.symbol = "none"
            else:
                series.marker.symbol = "circle"
                series.marker.size = 6 if index else 7
                series.marker.graphicalProperties.solidFill = color
                series.marker.graphicalProperties.line.solidFill = color
            labels = DataLabelList()
            labels.showVal = bool(labels_first and index == 0)
            labels.showLegendKey = False
            labels.showCatName = False
            labels.showSerName = False
            if labels.showVal:
                labels.dLblPos = "t"
                labels.numFmt = number_format
            series.dLbls = labels

    @staticmethod
    def _bar_style(chart, colors: list[str], *, number_format: str = "0.0"):
        for index, series in enumerate(chart.series):
            color = colors[index] if index < len(colors) else CHART_GREEN
            series.graphicalProperties.solidFill = color
            series.graphicalProperties.line.solidFill = color
        labels = DataLabelList()
        labels.showVal = True
        labels.showLegendKey = False
        labels.showCatName = False
        labels.showSerName = False
        labels.dLblPos = "outEnd"
        labels.numFmt = number_format
        chart.dLbls = labels

    # ------------------------------------------------------------------
    # Build orchestration
    # ------------------------------------------------------------------
    def build(self) -> Workbook:
        names = ["LEIA-ME", "PARAMETROS", "PAINEL"]
        if self.context.only_indicator:
            names.extend([INDICATOR_SHEETS[self.context.only_indicator], BASE_SHEETS[self.context.only_indicator]])
        else:
            names.extend([
                "DPE-01 MARGEM", "BASE DPE-01",
                "DPE-02 COBERTURA", "BASE DPE-02",
                "DPE-03 FOLHA", "BASE DPE-03",
            ])
        names.extend([
            "QUALIDADE E GOVERNANÇA", "MATRIZ", "PLANO_DE_ACAO", "METAS", "INDICADORES",
            "CALC", "LISTAS DE APOIO", "DIM_PERIODO",
        ])
        for name in names:
            self._new_sheet(name)

        self._build_base_sheets()
        self._build_lists()
        self._build_indicators()
        self._build_targets()
        self._build_actions()
        self._build_parameters()
        self._build_calc()
        self._build_indicator_pages()
        self._build_panel()
        self._build_matrix()
        self._build_quality()
        self._build_readme()
        self._finish()
        return self.wb

    # ------------------------------------------------------------------
    # Raw bases
    # ------------------------------------------------------------------
    def _rows_for(self, code: str) -> list[dict[str, Any]]:
        return sorted(
            [row for row in self.measurements if row.get("indicator_code") == code],
            key=lambda row: (period_sort_key(row.get("period")), str(row.get("dimension_label", ""))),
        )

    def _validate_capacity(self, count: int, reserve: int = INPUT_RESERVE):
        if BASE_DATA_START + count + reserve > EXCEL_MAX_ROW:
            raise ExcelExportLimitError(
                "A base da DPE ultrapassa o limite físico de 1.048.576 linhas do Excel. "
                "A exportação foi interrompida sem gerar arquivo parcial."
            )

    def _base_common(self, ws, title: str, subtitle: str, headers: list[str], widths: list[float], rows_count: int):
        self._title(ws, title, subtitle, end_col=len(headers))
        self._headers(ws, BASE_HEADER_ROW, headers)
        for index, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(index)].width = width
        ws.freeze_panes = "A5"
        ws.auto_filter.ref = f"A4:{get_column_letter(len(headers))}{BASE_DATA_START + max(rows_count, 1) - 1}"
        ws.print_title_rows = "1:4"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_view.zoomScale = 70

    def _apply_input_style(self, ws, start_row: int, end_row: int, input_cols: set[int], formula_cols: set[int]):
        for r in range(start_row, end_row + 1):
            is_existing = any(ws.cell(r, c).value not in (None, "") for c in input_cols)
            for c in range(1, ws.max_column + 1):
                cell = ws.cell(r, c)
                cell.border = Border(bottom=Side(style="hair", color=BORDER))
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if c in formula_cols:
                    cell.font = self.formula_font
                    cell.fill = PatternFill("solid", fgColor=GRAY)
                elif c in input_cols:
                    cell.font = self.linked_font if is_existing else self.input_font
                    if not is_existing:
                        cell.fill = PatternFill("solid", fgColor=PALE_BLUE)
                else:
                    cell.font = self.body_font
            ws.row_dimensions[r].height = 21

    def _build_base_sheets(self):
        if "BASE DPE-01" in self.ws:
            self._build_base_01()
        if "BASE DPE-02" in self.ws:
            self._build_base_02()
        if "BASE DPE-03" in self.ws:
            self._build_base_03()

    def _build_base_01(self):
        ws = self.ws["BASE DPE-01"]
        rows = self._rows_for("DPE-01")
        reserve = INPUT_RESERVE if self.context.template_mode or not rows else max(25, INPUT_RESERVE // 2)
        self._validate_capacity(len(rows), reserve)
        headers = [
            "ID", "Período", "Curso", "Diretoria acadêmica",
            "Receita líquida", "Custo docente", "Custo de coordenação", "Outros custos diretos", "Rateio indireto",
            "Custo total", "Margem líquida (%)", "Horas-aula semanais", "Nº de docentes", "CH média por docente",
            "Observações", "Validado",
        ]
        widths = [8, 12, 30, 20, 17, 16, 18, 18, 16, 16, 16, 18, 14, 18, 30, 11]
        self._base_common(
            ws,
            "DPE-01 · BASE DE MARGEM E CARGA HORÁRIA",
            "Uma linha por curso e competência. Curso e diretoria vêm do catálogo acadêmico; as colunas cinza são calculadas.",
            headers, widths, len(rows),
        )
        end = BASE_DATA_START + len(rows) + reserve - 1
        for i, row in enumerate(rows, BASE_DATA_START):
            values = raw_components(row); dims = row_dimensions(row)
            raw = [
                row.get("id"), row.get("period"), dims.get("course"), dims.get("academic_directorate"),
                values.get("net_revenue"), values.get("faculty_cost"), values.get("coordination_cost"),
                values.get("other_direct_cost"), values.get("indirect_cost"), None, None,
                values.get("weekly_teaching_hours"), values.get("teacher_count"), None,
                row.get("notes"), "Sim" if row.get("validated") else "Não",
            ]
            for c, value in enumerate(raw, 1): ws.cell(i, c).value = value
        for r in range(BASE_DATA_START, end + 1):
            ws.cell(r, 10).value = f'=IF(COUNTA(E{r}:I{r})=0,"",SUM(F{r}:I{r}))'
            ws.cell(r, 11).value = f'=IFERROR((E{r}-J{r})/E{r}*100,"")'
            ws.cell(r, 14).value = f'=IFERROR(L{r}/M{r},"")'
            for c in (5,6,7,8,9,10): ws.cell(r,c).number_format = CURRENCY_FMT
            ws.cell(r,11).number_format = PERCENT_POINTS_FMT
            for c in (12,14): ws.cell(r,c).number_format = DECIMAL_FMT
            ws.cell(r,13).number_format = INTEGER_FMT
        input_cols = set(range(1,10)) | {12,13,15,16}
        self._apply_input_style(ws, BASE_DATA_START, end, input_cols, {10,11,14})
        dv = DataValidation(type="list", formula1='"Sim,Não"', allow_blank=True); ws.add_data_validation(dv); dv.add(f"P{BASE_DATA_START}:P{end}")

    def _build_base_02(self):
        ws = self.ws["BASE DPE-02"]
        rows = self._rows_for("DPE-02")
        reserve = INPUT_RESERVE if self.context.template_mode or not rows else max(25, INPUT_RESERVE // 2)
        self._validate_capacity(len(rows), reserve)
        headers = [
            "ID", "Período", "Centro de custo", "Natureza da despesa",
            "Receita líquida", "Despesa de pessoal", "Despesa operacional", "Despesa administrativa", "Despesa financeira",
            "Despesa total", "Índice de cobertura", "Margem operacional (%)", "Investimento em imobilizado",
            "Observações", "Validado",
        ]
        widths = [8,12,24,22,17,17,18,19,18,17,16,18,20,30,11]
        self._base_common(ws, "DPE-02 · BASE DE COBERTURA ENTRE RECEITA E DESPESA",
            "Uma linha por competência e recorte. Imobilizado fica separado e não compõe a despesa operacional.", headers, widths, len(rows))
        end = BASE_DATA_START + len(rows) + reserve - 1
        for i,row in enumerate(rows, BASE_DATA_START):
            values=raw_components(row); dims=row_dimensions(row)
            raw=[row.get("id"),row.get("period"),dims.get("cost_center"),dims.get("expense_nature"),
                 values.get("net_revenue"),values.get("personnel_expense"),values.get("operational_expense"),
                 values.get("administrative_expense"),values.get("financial_expense"),None,None,None,values.get("capex"),
                 row.get("notes"),"Sim" if row.get("validated") else "Não"]
            for c,value in enumerate(raw,1): ws.cell(i,c).value=value
        for r in range(BASE_DATA_START,end+1):
            ws.cell(r,10).value=f'=IF(COUNTA(E{r}:I{r})=0,"",SUM(F{r}:I{r}))'
            ws.cell(r,11).value=f'=IFERROR(E{r}/J{r},"")'
            ws.cell(r,12).value=f'=IFERROR((1-1/K{r})*100,"")'
            for c in (5,6,7,8,9,10,13): ws.cell(r,c).number_format=CURRENCY_FMT
            ws.cell(r,11).number_format='0.00"x"'; ws.cell(r,12).number_format=PERCENT_POINTS_FMT
        input_cols=set(range(1,10))|{13,14,15}
        self._apply_input_style(ws,BASE_DATA_START,end,input_cols,{10,11,12})
        dv=DataValidation(type="list",formula1='"Sim,Não"',allow_blank=True); ws.add_data_validation(dv); dv.add(f"O{BASE_DATA_START}:O{end}")

    def _build_base_03(self):
        ws = self.ws["BASE DPE-03"]
        rows = self._rows_for("DPE-03")
        reserve = INPUT_RESERVE if self.context.template_mode or not rows else max(25, INPUT_RESERVE // 2)
        self._validate_capacity(len(rows), reserve)
        headers=["ID","Período","Categoria","Natureza","Folha docente completa","Folha administrativa completa",
                 "Folha total","Salários brutos","Encargos","Provisões","Receita líquida do mês",
                 "Folha total / receita (%)","Folha docente / receita (%)","Folha administrativa / receita (%)",
                 "Observações","Validado"]
        widths=[8,12,20,18,19,22,17,17,15,15,18,19,20,22,30,11]
        self._base_common(ws,"DPE-03 · BASE DE FOLHA DE PAGAMENTO SOBRE A RECEITA",
            "Registre salários, encargos e provisões por competência. A folha total é lida junto da receita do mesmo mês.",headers,widths,len(rows))
        end=BASE_DATA_START+len(rows)+reserve-1
        for i,row in enumerate(rows,BASE_DATA_START):
            values=raw_components(row); dims=row_dimensions(row)
            raw=[row.get("id"),row.get("period"),dims.get("category"),dims.get("nature"),values.get("faculty_payroll"),
                 values.get("administrative_payroll"),None,values.get("gross_salaries"),values.get("charges"),values.get("provisions"),
                 values.get("net_revenue"),None,None,None,row.get("notes"),"Sim" if row.get("validated") else "Não"]
            for c,value in enumerate(raw,1): ws.cell(i,c).value=value
        for r in range(BASE_DATA_START,end+1):
            ws.cell(r,7).value=f'=IF(COUNTA(E{r}:F{r})=0,"",SUM(E{r}:F{r}))'
            ws.cell(r,12).value=f'=IFERROR(G{r}/K{r}*100,"")'
            ws.cell(r,13).value=f'=IFERROR(E{r}/K{r}*100,"")'
            ws.cell(r,14).value=f'=IFERROR(F{r}/K{r}*100,"")'
            for c in (5,6,7,8,9,10,11): ws.cell(r,c).number_format=CURRENCY_FMT
            for c in (12,13,14): ws.cell(r,c).number_format=PERCENT_POINTS_FMT
        input_cols=set(range(1,7))|{8,9,10,11,15,16}
        self._apply_input_style(ws,BASE_DATA_START,end,input_cols,{7,12,13,14})
        dv=DataValidation(type="list",formula1='"Sim,Não"',allow_blank=True); ws.add_data_validation(dv); dv.add(f"P{BASE_DATA_START}:P{end}")

    # ------------------------------------------------------------------
    # Catalog, parameters and helper dimensions
    # ------------------------------------------------------------------
    def _build_lists(self):
        ws = self.ws["LISTAS DE APOIO"]
        self._title(ws, "LISTAS DE APOIO", "Listas usadas nas validações e filtros da DPE.", end_col=8)
        self._headers(ws, 4, ["Indicadores", "Períodos", "Status", "Validado", "Modalidade", "Turno", "Diretorias", "Categorias"])
        for r, value in enumerate(["DPE-01", "DPE-02", "DPE-03"], 5): ws.cell(r, 1).value = value
        for r, value in enumerate(self.periods, 5): ws.cell(r, 2).value = value
        for r, value in enumerate(["Dentro da meta", "Atenção", "Fora da meta", "Não apurado"], 5): ws.cell(r, 3).value = value
        for r, value in enumerate(["Sim", "Não"], 5): ws.cell(r, 4).value = value
        for r, value in enumerate(["Presencial", "Semipresencial", "EAD"], 5): ws.cell(r, 5).value = value
        for r, value in enumerate(["Matutino", "Vespertino", "Noturno", "Integral"], 5): ws.cell(r, 6).value = value
        for r, value in enumerate(["DTNH", "DCS", "DM", "DEAD"], 5): ws.cell(r, 7).value = value
        for r, value in enumerate(["Folha total", "Docente", "Administrativa"], 5): ws.cell(r, 8).value = value
        for col in range(1, 9): ws.column_dimensions[get_column_letter(col)].width = 20
        ws.sheet_state = "hidden"

        dim = self.ws["DIM_PERIODO"]
        self._title(dim, "DIMENSÃO DE PERÍODO", "Calendário mensal efetivamente utilizado pela exportação.", end_col=5)
        self._headers(dim, 4, ["Período", "Chave", "Ano", "Mês", "Rótulo"])
        for r, period in enumerate(self.periods, 5):
            dim.cell(r, 1).value = period
            dim.cell(r, 2).value = period_sort_key(period)
            dim.cell(r, 3).value = int(period[:4])
            dim.cell(r, 4).value = int(period[5:7])
            dim.cell(r, 5).value = f"{int(period[5:7]):02d}/{int(period[:4]) % 100:02d}"
        for col, width in enumerate([14, 12, 10, 10, 16], 1): dim.column_dimensions[get_column_letter(col)].width = width
        dim.sheet_state = "hidden"

    def _build_indicators(self):
        ws = self.ws["INDICADORES"]
        self._title(ws, "CATÁLOGO DE INDICADORES — DPE", "Definição institucional, fonte, periodicidade e métricas de cada indicador.", end_col=10)
        self._headers(ws, 4, ["Código", "Indicador", "Métrica", "Unidade", "Direção", "Meta", "Atenção", "Periodicidade", "Fonte", "Fórmula"])
        r = 5
        for spec in indicator_specs("DPE"):
            for metric in spec.get("metrics") or []:
                ws.cell(r, 1).value = spec["code"]
                ws.cell(r, 2).value = spec["name"]
                ws.cell(r, 3).value = metric["label"]
                ws.cell(r, 4).value = metric.get("unit")
                ws.cell(r, 5).value = metric.get("direction")
                if metric.get("direction") == "range":
                    ws.cell(r, 6).value = f"{metric.get('target_min')} a {metric.get('target_max')}"
                    ws.cell(r, 7).value = f"{metric.get('attention_min')} a {metric.get('attention_max')}"
                else:
                    ws.cell(r, 6).value = metric.get("target")
                    ws.cell(r, 7).value = metric.get("attention")
                ws.cell(r, 8).value = "Mensal"
                ws.cell(r, 9).value = spec.get("source")
                ws.cell(r, 10).value = spec.get("formula_text")
                r += 1
        for col, width in enumerate([12, 32, 30, 12, 12, 14, 14, 14, 40, 55], 1):
            ws.column_dimensions[get_column_letter(col)].width = width
        for row in ws.iter_rows(min_row=5, max_row=r - 1, min_col=1, max_col=10):
            for cell in row:
                cell.font = self.static_font
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = Border(bottom=Side(style="hair", color=BORDER))
        ws.freeze_panes = "A5"
        ws.auto_filter.ref = f"A4:J{max(5, r-1)}"

    def _build_targets(self):
        ws = self.ws["METAS"]
        self._title(ws, "METAS DOS INDICADORES — DPE", "Metas versionadas por vigência e dimensão. As metas do catálogo funcionam como fallback.", end_col=13)
        headers = ["ID", "Indicador", "Métrica", "Dimensão", "Vigência inicial", "Vigência final", "Meta", "Atenção", "Meta mínima", "Meta máxima", "Atenção mínima", "Atenção máxima", "Justificativa"]
        self._headers(ws, 4, headers)
        rows = self.targets
        end = 5 + max(len(rows), 40) - 1
        for r, item in enumerate(rows, 5):
            vals = [item.get("id"), item.get("indicator_code"), item.get("metric_key"), item.get("dimension_label"), item.get("valid_from"), item.get("valid_to"), item.get("target"), item.get("attention"), item.get("target_min"), item.get("target_max"), item.get("attention_min"), item.get("attention_max"), item.get("justification")]
            for c, value in enumerate(vals, 1): ws.cell(r, c).value = value
        for r in range(5, end + 1):
            for c in range(1, 14):
                ws.cell(r, c).font = self.linked_font if r < 5 + len(rows) else self.input_font
                ws.cell(r, c).alignment = Alignment(vertical="center", wrap_text=True)
                ws.cell(r, c).border = Border(bottom=Side(style="hair", color=BORDER))
        for col, width in enumerate([8, 12, 30, 24, 15, 15, 12, 12, 12, 12, 14, 14, 38], 1): ws.column_dimensions[get_column_letter(col)].width = width
        ws.freeze_panes = "A5"
        ws.auto_filter.ref = f"A4:M{end}"

    def _build_actions(self):
        ws = self.ws["PLANO_DE_ACAO"]
        self._title(ws, "PLANOS DE AÇÃO — DPE", "Todo indicador fora da meta deve possuir causa, ação, responsável, prazo e evidência.", end_col=14)
        headers = ["ID", "Indicador", "Métrica", "Período", "Dimensão", "Problema", "Causa provável", "Ação corretiva", "Responsável", "Prazo", "Status", "Evidência", "Criado por", "Atualizado em"]
        self._headers(ws, 4, headers)
        rows = self.actions
        end = 5 + max(len(rows), 40) - 1
        for r, item in enumerate(rows, 5):
            vals = [item.get("id"), item.get("indicator_code"), item.get("metric_key"), item.get("period"), item.get("dimension_label"), item.get("problem"), item.get("probable_cause"), item.get("corrective_action"), item.get("responsible"), item.get("due_date"), item.get("status"), item.get("evidence"), item.get("created_by"), item.get("updated_at")]
            for c, value in enumerate(vals, 1): ws.cell(r, c).value = value
        for r in range(5, end + 1):
            for c in range(1, 15):
                ws.cell(r, c).font = self.linked_font if r < 5 + len(rows) else self.input_font
                ws.cell(r, c).alignment = Alignment(vertical="center", wrap_text=True)
                ws.cell(r, c).border = Border(bottom=Side(style="hair", color=BORDER))
        for col, width in enumerate([8, 12, 28, 13, 24, 36, 32, 40, 24, 14, 15, 34, 24, 20], 1): ws.column_dimensions[get_column_letter(col)].width = width
        ws.freeze_panes = "A5"
        ws.auto_filter.ref = f"A4:N{end}"

    def _build_parameters(self):
        ws = self.ws["PARAMETROS"]
        self._title(ws, "PARÂMETROS DO PAINEL — DPE", "Contexto do arquivo exportado. Para atualizar dados, edite uma base e reimporte no Data UNIVC.", end_col=8)
        self._section(ws, 4, "RECORTE DA EXPORTAÇÃO", end_col=8)
        rows = [
            ("Diretoria", "DPE"),
            ("Periodicidade", "Mensal"),
            ("Mês de referência", self.reference),
            ("Mês de comparação", self.comparison),
            ("Comparação anual", _same_month_previous_year(self.reference)),
            ("Janela gráfica", "12 meses"),
            ("Indicador exportado", self.context.only_indicator or "Todos"),
            ("Gerado em", datetime.now()),
        ]
        for r, (label, value) in enumerate(rows, 5):
            ws.cell(r, 1).value = label
            ws.cell(r, 1).font = Font(name="Arial", size=9, bold=True, color=GREEN_DARK)
            ws.cell(r, 2).value = value
            ws.cell(r, 2).font = self.control_font
            ws.cell(r, 1).fill = PatternFill("solid", fgColor=PALE_GREEN)
            ws.cell(r, 2).fill = PatternFill("solid", fgColor=PALE_BLUE)
            ws.cell(r, 1).border = ws.cell(r, 2).border = self.border_box
        ws.column_dimensions["A"].width = 26
        ws.column_dimensions["B"].width = 28
        for col in range(3, 9): ws.column_dimensions[get_column_letter(col)].width = 14
        self._section(ws, 15, "REGRAS IMPORTANTES", end_col=8)
        notes = [
            "DPE-01: margem institucional ≥ 15%; por curso ≥ 10%; atenção em 0%.",
            "DPE-02: índice de cobertura ≥ 1,11; atenção em 1,05; no máximo dois meses seguidos abaixo de 1,00.",
            "DPE-03: folha total ≤ 55% da receita; atenção em 58%; docente ≤ 38%; administrativa ≤ 17%.",
            "Os acumulados de 12 meses são recalculados pelos numeradores e denominadores, nunca pela média simples dos percentuais.",
            "As bases não possuem corte silencioso: todos os registros exportados permanecem no arquivo.",
        ]
        for r, note in enumerate(notes, 16):
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
            ws.cell(r, 1).value = f"• {note}"
            ws.cell(r, 1).font = self.static_font
            ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r].height = 27

    # ------------------------------------------------------------------
    # CALC formulas and chart sources
    # ------------------------------------------------------------------
    def _build_calc(self):
        ws = self.ws["CALC"]
        self._title(ws, "CÁLCULOS DO PAINEL — DPE", "Faixas auxiliares geradas do zero. Não editar manualmente.", end_col=44)
        self._build_calc_01(ws, start_col=1)
        self._build_calc_02(ws, start_col=15)
        self._build_calc_03(ws, start_col=30)
        self._build_course_calc(ws, start_col=45)
        for col in range(1, 55): ws.column_dimensions[get_column_letter(col)].width = 15
        ws.sheet_state = "hidden"

    def _build_calc_01(self, ws, *, start_col: int):
        c = start_col
        # DPE-01 foi simplificado no fluxo operacional: não depende mais de alunos ativos,
        # modalidade ou turno. As posições abaixo espelham exatamente BASE DPE-01.
        headers = ["Período", "Receita", "Custo", "Margem", "Margem 12m", "CH", "Docentes", "CH/docente", "Meta", "Atenção", "Status"]
        self._headers(ws, 4, headers, start_col=c)
        start = 5
        for index, period in enumerate(self.periods):
            r = start + index
            ws.cell(r, c).value = period
            # BASE DPE-01: E Receita, J Custo total, L Horas-aula, M Nº docentes.
            ws.cell(r, c + 1).value = f'=SUMIFS(\'BASE DPE-01\'!$E:$E,\'BASE DPE-01\'!$B:$B,{get_column_letter(c)}{r})'
            ws.cell(r, c + 2).value = f'=SUMIFS(\'BASE DPE-01\'!$J:$J,\'BASE DPE-01\'!$B:$B,{get_column_letter(c)}{r})'
            ws.cell(r, c + 3).value = f'=IFERROR(ROUND(({get_column_letter(c+1)}{r}-{get_column_letter(c+2)}{r})/{get_column_letter(c+1)}{r}*100,1),"")'
            window_start = max(start, r - 11)
            ws.cell(r, c + 4).value = f'=IFERROR(ROUND((SUM({get_column_letter(c+1)}{window_start}:{get_column_letter(c+1)}{r})-SUM({get_column_letter(c+2)}{window_start}:{get_column_letter(c+2)}{r}))/SUM({get_column_letter(c+1)}{window_start}:{get_column_letter(c+1)}{r})*100,1),"")'
            ws.cell(r, c + 5).value = f'=SUMIFS(\'BASE DPE-01\'!$L:$L,\'BASE DPE-01\'!$B:$B,{get_column_letter(c)}{r})'
            ws.cell(r, c + 6).value = f'=SUMIFS(\'BASE DPE-01\'!$M:$M,\'BASE DPE-01\'!$B:$B,{get_column_letter(c)}{r})'
            ws.cell(r, c + 7).value = f'=IFERROR(ROUND({get_column_letter(c+5)}{r}/{get_column_letter(c+6)}{r},2),"")'
            target = _target_for_period(self.targets, indicator_spec("DPE", "DPE-01"), metric_spec("DPE", "DPE-01", "net_margin_pct"), period)
            ws.cell(r, c + 8).value = _number(target.get("target"))
            ws.cell(r, c + 9).value = _number(target.get("attention"))
            ws.cell(r, c + 10).value = f'=IF({get_column_letter(c+3)}{r}="","Não apurado",IF({get_column_letter(c+3)}{r}>={get_column_letter(c+8)}{r},"Dentro da meta",IF({get_column_letter(c+3)}{r}>={get_column_letter(c+9)}{r},"Atenção","Fora da meta")))'
        self.calc01_start = start
        self.calc01_end = start + len(self.periods) - 1

    def _build_calc_02(self, ws, *, start_col: int):
        c = start_col
        headers = ["Período", "Receita", "Pessoal", "Operacional", "Administrativa", "Financeira", "Despesa total", "Cobertura", "Cobertura 12m", "Margem op.", "Margem op. 12m", "CAPEX", "Meta", "Atenção", "Status"]
        self._headers(ws, 4, headers, start_col=c)
        start = 5
        for index, period in enumerate(self.periods):
            r = start + index
            ws.cell(r, c).value = period
            # BASE DPE-02: E Receita, F Pessoal, G Operacional, H Administrativa, I Financeira.
            base_cols = ["E", "F", "G", "H", "I"]
            for offset, source_col in enumerate(base_cols, 1):
                ws.cell(r, c + offset).value = f'=SUMIFS(\'BASE DPE-02\'!${source_col}:${source_col},\'BASE DPE-02\'!$B:$B,{get_column_letter(c)}{r})'
            ws.cell(r, c + 6).value = f'=SUM({get_column_letter(c+2)}{r}:{get_column_letter(c+5)}{r})'
            ws.cell(r, c + 7).value = f'=IFERROR(ROUND({get_column_letter(c+1)}{r}/{get_column_letter(c+6)}{r},2),"")'
            window_start = max(start, r - 11)
            ws.cell(r, c + 8).value = f'=IFERROR(ROUND(SUM({get_column_letter(c+1)}{window_start}:{get_column_letter(c+1)}{r})/SUM({get_column_letter(c+6)}{window_start}:{get_column_letter(c+6)}{r}),2),"")'
            ws.cell(r, c + 9).value = f'=IFERROR(ROUND((1-1/{get_column_letter(c+7)}{r})*100,1),"")'
            ws.cell(r, c + 10).value = f'=IFERROR(ROUND((1-1/{get_column_letter(c+8)}{r})*100,1),"")'
            ws.cell(r, c + 11).value = f'=SUMIFS(\'BASE DPE-02\'!$M:$M,\'BASE DPE-02\'!$B:$B,{get_column_letter(c)}{r})'
            target = _target_for_period(self.targets, indicator_spec("DPE", "DPE-02"), metric_spec("DPE", "DPE-02", "coverage_index"), period)
            ws.cell(r, c + 12).value = _number(target.get("target"))
            ws.cell(r, c + 13).value = _number(target.get("attention"))
            ws.cell(r, c + 14).value = f'=IF({get_column_letter(c+7)}{r}="","Não apurado",IF({get_column_letter(c+7)}{r}>={get_column_letter(c+12)}{r},"Dentro da meta",IF({get_column_letter(c+7)}{r}>={get_column_letter(c+13)}{r},"Atenção","Fora da meta")))'
        self.calc02_start = start
        self.calc02_end = start + len(self.periods) - 1

    def _build_calc_03(self, ws, *, start_col: int):
        c = start_col
        headers = ["Período", "Folha docente", "Folha adm.", "Folha total", "Receita", "Folha %", "Docente %", "Adm. %", "Folha 3m %", "Variação %", "Meta", "Atenção", "Status"]
        self._headers(ws, 4, headers, start_col=c)
        start = 5
        for index, period in enumerate(self.periods):
            r = start + index
            ws.cell(r, c).value = period
            ws.cell(r, c + 1).value = f'=SUMIFS(\'BASE DPE-03\'!$E:$E,\'BASE DPE-03\'!$B:$B,{get_column_letter(c)}{r})'
            ws.cell(r, c + 2).value = f'=SUMIFS(\'BASE DPE-03\'!$F:$F,\'BASE DPE-03\'!$B:$B,{get_column_letter(c)}{r})'
            ws.cell(r, c + 3).value = f'=SUM({get_column_letter(c+1)}{r}:{get_column_letter(c+2)}{r})'
            ws.cell(r, c + 4).value = f'=SUMIFS(\'BASE DPE-03\'!$K:$K,\'BASE DPE-03\'!$B:$B,{get_column_letter(c)}{r})'
            ws.cell(r, c + 5).value = f'=IFERROR(ROUND({get_column_letter(c+3)}{r}/{get_column_letter(c+4)}{r}*100,1),"")'
            ws.cell(r, c + 6).value = f'=IFERROR(ROUND({get_column_letter(c+1)}{r}/{get_column_letter(c+4)}{r}*100,1),"")'
            ws.cell(r, c + 7).value = f'=IFERROR(ROUND({get_column_letter(c+2)}{r}/{get_column_letter(c+4)}{r}*100,1),"")'
            window_start = max(start, r - 2)
            ws.cell(r, c + 8).value = f'=IFERROR(ROUND({get_column_letter(c+3)}{r}/AVERAGE({get_column_letter(c+4)}{window_start}:{get_column_letter(c+4)}{r})*100,1),"")'
            if r == start:
                ws.cell(r, c + 9).value = '=""'
            else:
                ws.cell(r, c + 9).value = f'=IFERROR(ROUND(({get_column_letter(c+3)}{r}-{get_column_letter(c+3)}{r-1})/{get_column_letter(c+3)}{r-1}*100,1),"")'
            target = _target_for_period(self.targets, indicator_spec("DPE", "DPE-03"), metric_spec("DPE", "DPE-03", "payroll_on_revenue_pct"), period)
            ws.cell(r, c + 10).value = _number(target.get("target"))
            ws.cell(r, c + 11).value = _number(target.get("attention"))
            ws.cell(r, c + 12).value = f'=IF({get_column_letter(c+5)}{r}="","Não apurado",IF({get_column_letter(c+5)}{r}<={get_column_letter(c+10)}{r},"Dentro da meta",IF({get_column_letter(c+5)}{r}<={get_column_letter(c+11)}{r},"Atenção","Fora da meta")))'
        self.calc03_start = start
        self.calc03_end = start + len(self.periods) - 1

    def _build_course_calc(self, ws, *, start_col: int):
        c = start_col
        courses = sorted({row_dimensions(row).get("course") for row in self._rows_for("DPE-01") if row_dimensions(row).get("course")})
        headers = ["Curso", "Receita", "Custo", "Margem", "Meta", "CH/docente", "Status"]
        self._headers(ws, 4, headers, start_col=c)
        start = 5
        for i, course in enumerate(courses, start):
            ws.cell(i, c).value = course
            ref = 'PARAMETROS!$B$7'
            ws.cell(i, c + 1).value = f'=SUMIFS(\'BASE DPE-01\'!$E:$E,\'BASE DPE-01\'!$B:$B,{ref},\'BASE DPE-01\'!$C:$C,{get_column_letter(c)}{i})'
            ws.cell(i, c + 2).value = f'=SUMIFS(\'BASE DPE-01\'!$J:$J,\'BASE DPE-01\'!$B:$B,{ref},\'BASE DPE-01\'!$C:$C,{get_column_letter(c)}{i})'
            ws.cell(i, c + 3).value = f'=IFERROR(ROUND(({get_column_letter(c+1)}{i}-{get_column_letter(c+2)}{i})/{get_column_letter(c+1)}{i}*100,1),"")'
            target = _target_for_period(self.targets, indicator_spec("DPE", "DPE-01"), metric_spec("DPE", "DPE-01", "net_margin_pct"), self.reference, dimension_key=f"course={course}", course_dimension=True)
            ws.cell(i, c + 4).value = _number(target.get("target")) or 10
            ws.cell(i, c + 5).value = f'=IFERROR(ROUND(SUMIFS(\'BASE DPE-01\'!$L:$L,\'BASE DPE-01\'!$B:$B,{ref},\'BASE DPE-01\'!$C:$C,{get_column_letter(c)}{i})/SUMIFS(\'BASE DPE-01\'!$M:$M,\'BASE DPE-01\'!$B:$B,{ref},\'BASE DPE-01\'!$C:$C,{get_column_letter(c)}{i}),2),"")'
            ws.cell(i, c + 6).value = f'=IF({get_column_letter(c+3)}{i}="","Não apurado",IF({get_column_letter(c+3)}{i}>={get_column_letter(c+4)}{i},"Dentro da meta",IF({get_column_letter(c+3)}{i}>=0,"Atenção","Fora da meta")))'
        self.course_calc_start = start
        self.course_calc_end = start + len(courses) - 1

    def _period_row(self, start: int, period: str) -> int:
        try:
            return start + self.periods.index(period)
        except ValueError:
            return start + len(self.periods) - 1

    def _chart_window(self, start: int, end: int, points: int = 12) -> tuple[int, int]:
        reference_row = self._period_row(start, self.reference)
        return max(start, reference_row - points + 1), min(end, reference_row)

    # ------------------------------------------------------------------
    # Visible dashboards
    # ------------------------------------------------------------------
    def _setup_panel_sheet(self, ws, *, title: str, subtitle: str):
        self._title(ws, title, subtitle, end_col=12)
        widths = [14, 14, 14, 14, 14, 14, 14, 14, 14, 14, 14, 14]
        for c, width in enumerate(widths, 1): ws.column_dimensions[get_column_letter(c)].width = width
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.print_title_rows = "1:5"
        ws.freeze_panes = "A6"

    def _context_band(self, ws, row: int):
        labels = [(1, 3, "Mês de referência", "=PARAMETROS!$B$7"), (4, 6, "Comparação", "=PARAMETROS!$B$8"), (7, 9, "Mesmo mês do ano anterior", "=PARAMETROS!$B$9"), (10, 12, "Periodicidade", "Mensal")]
        for start, end, label, value in labels:
            self._paint_range(ws, row, row + 1, start, end, fill=PALE_GREEN)
            ws.merge_cells(start_row=row, start_column=start, end_row=row, end_column=end)
            ws.merge_cells(start_row=row + 1, start_column=start, end_row=row + 1, end_column=end)
            ws.cell(row, start).value = label
            ws.cell(row, start).font = Font(name="Arial", size=8, bold=True, color=GRAY_TEXT)
            ws.cell(row, start).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row + 1, start).value = value
            ws.cell(row + 1, start).font = Font(name="Arial", size=11, bold=True, color=GREEN_DARK)
            ws.cell(row + 1, start).alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[row].height = 19
        ws.row_dimensions[row + 1].height = 24

    def _metric_formula(self, calc_col: int, start_row: int) -> str:
        col = get_column_letter(calc_col)
        period_col = get_column_letter(calc_col - 3 if calc_col >= 30 else (15 if calc_col >= 15 else 1))
        return f'=IFERROR(INDEX(CALC!${col}:${col},MATCH(PARAMETROS!$B$7,CALC!${period_col}:${period_col},0)),"")'

    def _build_panel(self):
        ws = self.ws["PAINEL"]
        self._setup_panel_sheet(ws, title="PAINEL DE GESTÃO — DPE", subtitle="Diretoria de Planejamento Econômico e Oferta · mesmo padrão visual dos painéis DTNH e DCS")
        self._context_band(ws, 4)
        row = 7
        visible_codes = [self.context.only_indicator] if self.context.only_indicator else ["DPE-01", "DPE-02", "DPE-03"]
        for index, code in enumerate(visible_codes):
            if code == "DPE-01": row = self._panel_section_01(ws, row)
            elif code == "DPE-02": row = self._panel_section_02(ws, row)
            elif code == "DPE-03": row = self._panel_section_03(ws, row)
            if index < len(visible_codes) - 1:
                ws.row_breaks.append(Break(id=row - 1))
        ws.print_area = f"A1:L{row-1}"

    def _panel_section_01(self, ws, row: int) -> int:
        self._section(ws, row, "DPE-01 · Margem Líquida por Curso e Carga Horária Alocada")
        r = row + 1
        current = self._period_row(self.calc01_start, self.reference)
        col = lambda n: get_column_letter(n)
        self._kpi_card(ws, start_col=1, end_col=2, label_row=r, value_row=r+1, label="Margem no mês", formula=f"=CALC!D{current}", number_format=PERCENT_POINTS_FMT, fill=PALE_GREEN)
        self._kpi_card(ws, start_col=3, end_col=4, label_row=r, value_row=r+1, label="Margem acumulada 12m", formula=f"=CALC!E{current}", number_format=PERCENT_POINTS_FMT, fill=PALE_TEAL)
        self._kpi_card(ws, start_col=5, end_col=6, label_row=r, value_row=r+1, label="Receita líquida", formula=f"=CALC!B{current}", number_format=CURRENCY_FMT, fill=PALE_BLUE, font_size=13)
        self._kpi_card(ws, start_col=7, end_col=8, label_row=r, value_row=r+1, label="Custo total", formula=f"=CALC!C{current}", number_format=CURRENCY_FMT, fill=PALE_GOLD, font_size=13)
        self._kpi_card(ws, start_col=9, end_col=10, label_row=r, value_row=r+1, label="CH média / docente", formula=f"=CALC!H{current}", number_format=DECIMAL_FMT, fill=PALE_TEAL)
        self._kpi_card(ws, start_col=11, end_col=12, label_row=r, value_row=r+1, label="Meta institucional", formula=f"=CALC!I{current}", number_format=PERCENT_POINTS_FMT, fill=PALE_BLUE)
        chart_start, chart_end = self._chart_window(self.calc01_start, self.calc01_end)
        chart = LineChart()
        chart.add_data(Reference(self.ws["CALC"], min_col=4, max_col=5, min_row=4, max_row=chart_end), titles_from_data=True)
        chart.add_data(Reference(self.ws["CALC"], min_col=9, min_row=4, max_row=chart_end), titles_from_data=True)
        chart.set_categories(Reference(self.ws["DIM_PERIODO"], min_col=5, min_row=chart_start, max_row=chart_end))
        # Rebuild data references to the compact window while retaining titles.
        for idx, series in enumerate(chart.series):
            source_col = [4, 5, 9][idx]
            series.val.numRef.f = f"'CALC'!${get_column_letter(source_col)}${chart_start}:${get_column_letter(source_col)}${chart_end}"
        self._chart_base(chart, title="Margem líquida — evolução e meta", x_title="Competência", y_title="Margem (%)", width=33.5, height=9.0)
        self._line_style(chart, [CHART_GREEN, CHART_TEAL, CHART_META], number_format='0.0"%"')
        ws.add_chart(chart, f"A{r+4}")
        return r + 24

    def _panel_section_02(self, ws, row: int) -> int:
        self._section(ws, row, "DPE-02 · Índice de Cobertura entre Receita e Despesa")
        r = row + 1
        current = self._period_row(self.calc02_start, self.reference)
        # CALC DPE-02 starts in O (15): coverage V(22), 12m W(23), margin X(24), margin12 Y(25), expense U(21), capex Z(26)
        self._kpi_card(ws, start_col=1, end_col=2, label_row=r, value_row=r+1, label="Cobertura no mês", formula=f"=CALC!V{current}", number_format='0.00"x"', fill=PALE_GREEN)
        self._kpi_card(ws, start_col=3, end_col=4, label_row=r, value_row=r+1, label="Cobertura acumulada 12m", formula=f"=CALC!W{current}", number_format='0.00"x"', fill=PALE_TEAL)
        self._kpi_card(ws, start_col=5, end_col=6, label_row=r, value_row=r+1, label="Margem operacional", formula=f"=CALC!X{current}", number_format=PERCENT_POINTS_FMT, fill=PALE_BLUE)
        self._kpi_card(ws, start_col=7, end_col=8, label_row=r, value_row=r+1, label="Margem operacional 12m", formula=f"=CALC!Y{current}", number_format=PERCENT_POINTS_FMT, fill=PALE_GREEN)
        self._kpi_card(ws, start_col=9, end_col=10, label_row=r, value_row=r+1, label="Despesa total", formula=f"=CALC!U{current}", number_format=CURRENCY_FMT, fill=PALE_GOLD, font_size=13)
        self._kpi_card(ws, start_col=11, end_col=12, label_row=r, value_row=r+1, label="Imobilizado (fora da despesa)", formula=f"=CALC!Z{current}", number_format=CURRENCY_FMT, fill=PALE_BLUE, font_size=12)
        chart_start, chart_end = self._chart_window(self.calc02_start, self.calc02_end)
        chart = LineChart()
        for source_col in (22, 23, 27):
            chart.add_data(Reference(self.ws["CALC"], min_col=source_col, min_row=4, max_row=chart_end), titles_from_data=True)
        chart.set_categories(Reference(self.ws["DIM_PERIODO"], min_col=5, min_row=chart_start, max_row=chart_end))
        for idx, series in enumerate(chart.series):
            source_col = [22, 23, 27][idx]
            series.val.numRef.f = f"'CALC'!${get_column_letter(source_col)}${chart_start}:${get_column_letter(source_col)}${chart_end}"
        self._chart_base(chart, title="Cobertura — mês, acumulado de 12 meses e meta", x_title="Competência", y_title="Índice", width=33.5, height=9.0)
        self._line_style(chart, [CHART_GREEN, CHART_TEAL, CHART_META], number_format='0.00"x"')
        ws.add_chart(chart, f"A{r+4}")
        return r + 24

    def _panel_section_03(self, ws, row: int) -> int:
        self._section(ws, row, "DPE-03 · Folha de Pagamento sobre a Receita")
        r = row + 1
        current = self._period_row(self.calc03_start, self.reference)
        # CALC DPE-03 começa em AD (30): folha total AG(33), receita AH(34),
        # total % AI(35), docente % AJ(36), administrativa % AK(37),
        # leitura 3m AL(38), variação AM(39), meta AN(40) e status AP(42).
        self._kpi_card(ws, start_col=1, end_col=2, label_row=r, value_row=r+1, label="Folha total / receita", formula=f"=CALC!AI{current}", number_format=PERCENT_POINTS_FMT, fill=PALE_GREEN)
        self._kpi_card(ws, start_col=3, end_col=4, label_row=r, value_row=r+1, label="Folha / receita média 3m", formula=f"=CALC!AL{current}", number_format=PERCENT_POINTS_FMT, fill=PALE_TEAL)
        self._kpi_card(ws, start_col=5, end_col=6, label_row=r, value_row=r+1, label="Folha docente / receita", formula=f"=CALC!AJ{current}", number_format=PERCENT_POINTS_FMT, fill=PALE_BLUE)
        self._kpi_card(ws, start_col=7, end_col=8, label_row=r, value_row=r+1, label="Folha administrativa / receita", formula=f"=CALC!AK{current}", number_format=PERCENT_POINTS_FMT, fill=PALE_GOLD)
        self._kpi_card(ws, start_col=9, end_col=10, label_row=r, value_row=r+1, label="Variação mensal da folha", formula=f"=CALC!AM{current}", number_format=CHANGE_POINTS_FMT, fill=PALE_TEAL)
        self._kpi_card(ws, start_col=11, end_col=12, label_row=r, value_row=r+1, label="Folha total", formula=f"=CALC!AG{current}", number_format=CURRENCY_FMT, fill=PALE_BLUE, font_size=13)
        chart_start, chart_end = self._chart_window(self.calc03_start, self.calc03_end)
        chart = LineChart()
        for source_col in (35, 38, 40):  # AI, AL, AN
            chart.add_data(Reference(self.ws["CALC"], min_col=source_col, min_row=4, max_row=chart_end), titles_from_data=True)
        chart.set_categories(Reference(self.ws["DIM_PERIODO"], min_col=5, min_row=chart_start, max_row=chart_end))
        for idx, series in enumerate(chart.series):
            source_col = [35, 38, 40][idx]
            series.val.numRef.f = f"'CALC'!${get_column_letter(source_col)}${chart_start}:${get_column_letter(source_col)}${chart_end}"
        self._chart_base(chart, title="Folha sobre a receita — evolução e meta", x_title="Competência", y_title="Folha / receita (%)", width=33.5, height=9.0)
        self._line_style(chart, [CHART_GREEN, CHART_TEAL, CHART_META], number_format='0.0"%"')
        ws.add_chart(chart, f"A{r+4}")
        return r + 24

    def _build_indicator_pages(self):
        if "DPE-01 MARGEM" in self.ws: self._build_page_01()
        if "DPE-02 COBERTURA" in self.ws: self._build_page_02()
        if "DPE-03 FOLHA" in self.ws: self._build_page_03()

    def _build_page_01(self):
        ws = self.ws["DPE-01 MARGEM"]
        self._setup_panel_sheet(ws, title="DPE-01 · MARGEM LÍQUIDA E CARGA HORÁRIA", subtitle="Página própria do indicador · leitura institucional e detalhamento por curso")
        self._context_band(ws, 4)
        self._section(ws, 7, "SÍNTESE DO MÊS E DO ACUMULADO")
        r = 8
        current = self._period_row(self.calc01_start, self.reference)
        formulas = [("Margem no mês", f"=CALC!D{current}", PERCENT_POINTS_FMT, PALE_GREEN), ("Margem 12m", f"=CALC!E{current}", PERCENT_POINTS_FMT, PALE_TEAL), ("Receita líquida", f"=CALC!B{current}", CURRENCY_FMT, PALE_BLUE), ("Custo total", f"=CALC!C{current}", CURRENCY_FMT, PALE_GOLD), ("CH média / docente", f"=CALC!H{current}", DECIMAL_FMT, PALE_TEAL), ("Meta institucional", f"=CALC!I{current}", PERCENT_POINTS_FMT, PALE_BLUE)]
        for idx, (label, formula, fmt, fill) in enumerate(formulas):
            start = idx * 2 + 1
            self._kpi_card(ws, start_col=start, end_col=start+1, label_row=r, value_row=r+1, label=label, formula=formula, number_format=fmt, fill=fill, font_size=13 if "R$" in fmt else 16)
        chart_start, chart_end = self._chart_window(self.calc01_start, self.calc01_end)
        chart = LineChart()
        for source_col in (4, 5, 9): chart.add_data(Reference(self.ws["CALC"], min_col=source_col, min_row=4, max_row=chart_end), titles_from_data=True)
        chart.set_categories(Reference(self.ws["DIM_PERIODO"], min_col=5, min_row=chart_start, max_row=chart_end))
        for idx, series in enumerate(chart.series):
            source_col = [4, 5, 9][idx]
            series.val.numRef.f = f"'CALC'!${get_column_letter(source_col)}${chart_start}:${get_column_letter(source_col)}${chart_end}"
        self._chart_base(chart, title="Margem líquida: mês, acumulado móvel e meta", x_title="Competência", y_title="Margem (%)", width=15.5, height=8.5)
        self._line_style(chart, [CHART_GREEN, CHART_TEAL, CHART_META], number_format='0.0"%"')
        ws.add_chart(chart, "A13")

        bar = BarChart()
        if self.course_calc_end >= self.course_calc_start:
            max_courses = min(self.course_calc_end, self.course_calc_start + 11)
            bar.add_data(Reference(self.ws["CALC"], min_col=48, min_row=4, max_row=max_courses), titles_from_data=True)
            bar.set_categories(Reference(self.ws["CALC"], min_col=45, min_row=self.course_calc_start, max_row=max_courses))
            bar.series[0].val.numRef.f = f"'CALC'!$AV${self.course_calc_start}:$AV${max_courses}"
        self._chart_base(bar, title="Margem por curso no mês de referência", x_title="Curso", y_title="Margem (%)", width=10.0, height=8.5)
        bar.type = "bar"
        bar.grouping = "clustered"
        self._bar_style(bar, [CHART_GREEN], number_format='0.0"%"')
        ws.add_chart(bar, "H13")

        table_row = 31
        self._section(ws, table_row, "DETALHAMENTO POR CURSO")
        headers = ["Curso", "Receita", "Custo", "Margem", "Meta", "CH/docente", "Status"]
        self._headers(ws, table_row + 1, headers)
        out = table_row + 2
        for source_row in range(self.course_calc_start, self.course_calc_end + 1):
            for offset, source_col in enumerate(range(45, 52), 1):
                ws.cell(out, offset).value = f"=CALC!{get_column_letter(source_col)}{source_row}"
            ws.cell(out, 2).number_format = ws.cell(out, 3).number_format = CURRENCY_FMT
            ws.cell(out, 4).number_format = ws.cell(out, 5).number_format = PERCENT_POINTS_FMT
            ws.cell(out, 6).number_format = DECIMAL_FMT
            out += 1
        if out == table_row + 2:
            ws.cell(out, 1).value = "Nenhum curso apurado no período."
            out += 1
        for row in ws.iter_rows(min_row=table_row + 2, max_row=out - 1, min_col=1, max_col=7):
            for cell in row:
                cell.font = self.body_font
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.border = Border(bottom=Side(style="hair", color=BORDER))
        ws.conditional_formatting.add(f"G{table_row+2}:G{out-1}", FormulaRule(formula=[f'ISNUMBER(SEARCH("Fora",G{table_row+2}))'], fill=PatternFill("solid", fgColor=RED_LIGHT)))
        self._section(ws, out + 2, "LEITURA E GOVERNANÇA")
        notes = [
            "A margem institucional é comparada a 15%; a margem de cada curso, a 10%.",
            "A carga horária média deve permanecer entre 12 e 20 horas semanais por docente.",
            "O critério de rateio indireto precisa ser único, escrito e versionado por vigência.",
        ]
        for i, note in enumerate(notes, out + 3):
            ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=12)
            ws.cell(i, 1).value = f"• {note}"
            ws.cell(i, 1).font = self.static_font
            ws.cell(i, 1).alignment = Alignment(wrap_text=True)
        ws.print_area = f"A1:L{out+7}"
        ws.page_setup.fitToHeight = 1

    def _build_page_02(self):
        ws = self.ws["DPE-02 COBERTURA"]
        self._setup_panel_sheet(ws, title="DPE-02 · COBERTURA ENTRE RECEITA E DESPESA", subtitle="Página própria do indicador · mês, acumulado de 12 meses e comparação anual")
        self._context_band(ws, 4)
        self._section(ws, 7, "SÍNTESE DO PERÍODO")
        r = 8
        current = self._period_row(self.calc02_start, self.reference)
        formulas = [("Cobertura no mês", f"=CALC!V{current}", '0.00"x"', PALE_GREEN), ("Cobertura 12m", f"=CALC!W{current}", '0.00"x"', PALE_TEAL), ("Margem operacional", f"=CALC!X{current}", PERCENT_POINTS_FMT, PALE_BLUE), ("Margem operacional 12m", f"=CALC!Y{current}", PERCENT_POINTS_FMT, PALE_GREEN), ("Despesa total", f"=CALC!U{current}", CURRENCY_FMT, PALE_GOLD), ("Imobilizado", f"=CALC!Z{current}", CURRENCY_FMT, PALE_BLUE)]
        for idx, (label, formula, fmt, fill) in enumerate(formulas):
            start = idx * 2 + 1
            self._kpi_card(ws, start_col=start, end_col=start+1, label_row=r, value_row=r+1, label=label, formula=formula, number_format=fmt, fill=fill, font_size=13 if "R$" in fmt else 16)
        chart_start, chart_end = self._chart_window(self.calc02_start, self.calc02_end)
        chart = LineChart()
        for source_col in (22, 23, 27): chart.add_data(Reference(self.ws["CALC"], min_col=source_col, min_row=4, max_row=chart_end), titles_from_data=True)
        chart.set_categories(Reference(self.ws["DIM_PERIODO"], min_col=5, min_row=chart_start, max_row=chart_end))
        for idx, series in enumerate(chart.series):
            source_col = [22, 23, 27][idx]
            series.val.numRef.f = f"'CALC'!${get_column_letter(source_col)}${chart_start}:${get_column_letter(source_col)}${chart_end}"
        self._chart_base(chart, title="Índice de cobertura — evolução e meta", x_title="Competência", y_title="Índice", width=15.5, height=8.5)
        self._line_style(chart, [CHART_GREEN, CHART_TEAL, CHART_META], number_format='0.00"x"')
        ws.add_chart(chart, "A13")

        expense = BarChart()
        expense.type = "col"
        expense.grouping = "stacked"
        for col in range(17, 21):
            expense.add_data(Reference(self.ws["CALC"], min_col=col, min_row=4, max_row=current), titles_from_data=True)
        expense.set_categories(Reference(self.ws["DIM_PERIODO"], min_col=5, min_row=current, max_row=current))
        for idx, series in enumerate(expense.series):
            col = 17 + idx
            series.val.numRef.f = f"'CALC'!${get_column_letter(col)}${current}:${get_column_letter(col)}${current}"
        self._chart_base(expense, title="Composição da despesa no mês", x_title="Competência", y_title="R$", width=10.0, height=8.5)
        self._bar_style(expense, [CHART_GREEN, CHART_TEAL, CHART_GOLD, CHART_BLUE], number_format='R$ #,##0')
        ws.add_chart(expense, "H13")

        table_row = 31
        self._section(ws, table_row, "HISTÓRICO DOS 12 MESES")
        headers = ["Período", "Receita", "Despesa", "Cobertura", "Cobertura 12m", "Margem op.", "CAPEX", "Status"]
        self._headers(ws, table_row + 1, headers)
        out = table_row + 2
        for source_row in range(chart_start, chart_end + 1):
            sources = [15, 16, 21, 22, 23, 24, 26, 29]
            for offset, source_col in enumerate(sources, 1): ws.cell(out, offset).value = f"=CALC!{get_column_letter(source_col)}{source_row}"
            ws.cell(out, 2).number_format = ws.cell(out, 3).number_format = ws.cell(out, 7).number_format = CURRENCY_FMT
            ws.cell(out, 4).number_format = ws.cell(out, 5).number_format = '0.00"x"'
            ws.cell(out, 6).number_format = PERCENT_POINTS_FMT
            out += 1
        for row in ws.iter_rows(min_row=table_row + 2, max_row=out - 1, min_col=1, max_col=8):
            for cell in row:
                cell.font = self.body_font
                cell.alignment = Alignment(vertical="center")
                cell.border = Border(bottom=Side(style="hair", color=BORDER))
        self._section(ws, out + 2, "REGRA DE ALERTA")
        ws.merge_cells(start_row=out + 3, start_column=1, end_row=out + 4, end_column=12)
        ws.cell(out + 3, 1).value = "O índice deve permanecer em pelo menos 1,11 no mês e no acumulado de 12 meses. A atenção começa em 1,05. Mais de dois meses consecutivos abaixo de 1,00 exige plano de ação."
        ws.cell(out + 3, 1).font = Font(name="Arial", size=10, color=GRAY_TEXT)
        ws.cell(out + 3, 1).alignment = Alignment(wrap_text=True, vertical="center")
        self._paint_range(ws, out + 3, out + 4, 1, 12, fill=PALE_GOLD)
        ws.print_area = f"A1:L{out+5}"
        ws.page_setup.fitToHeight = 1

    def _build_page_03(self):
        ws = self.ws["DPE-03 FOLHA"]
        self._setup_panel_sheet(ws, title="DPE-03 · FOLHA DE PAGAMENTO SOBRE A RECEITA", subtitle="Página própria do indicador · folha total, docente, administrativa e leitura sobre a receita média trimestral")
        self._context_band(ws, 4)
        self._section(ws, 7, "SÍNTESE DO PERÍODO")
        r = 8
        current = self._period_row(self.calc03_start, self.reference)
        formulas = [("Folha total / receita", f"=CALC!AI{current}", PERCENT_POINTS_FMT, PALE_GREEN), ("Folha / receita média 3m", f"=CALC!AL{current}", PERCENT_POINTS_FMT, PALE_TEAL), ("Folha docente / receita", f"=CALC!AJ{current}", PERCENT_POINTS_FMT, PALE_BLUE), ("Folha administrativa / receita", f"=CALC!AK{current}", PERCENT_POINTS_FMT, PALE_GOLD), ("Variação mensal da folha", f"=CALC!AM{current}", CHANGE_POINTS_FMT, PALE_TEAL), ("Folha total", f"=CALC!AG{current}", CURRENCY_FMT, PALE_BLUE)]
        for idx, (label, formula, fmt, fill) in enumerate(formulas):
            start = idx * 2 + 1
            self._kpi_card(ws, start_col=start, end_col=start+1, label_row=r, value_row=r+1, label=label, formula=formula, number_format=fmt, fill=fill, font_size=13 if "R$" in fmt else 16)
        chart_start, chart_end = self._chart_window(self.calc03_start, self.calc03_end)
        chart = LineChart()
        for source_col in (35, 38, 40): chart.add_data(Reference(self.ws["CALC"], min_col=source_col, min_row=4, max_row=chart_end), titles_from_data=True)
        chart.set_categories(Reference(self.ws["DIM_PERIODO"], min_col=5, min_row=chart_start, max_row=chart_end))
        for idx, series in enumerate(chart.series):
            source_col = [35, 38, 40][idx]
            series.val.numRef.f = f"'CALC'!${get_column_letter(source_col)}${chart_start}:${get_column_letter(source_col)}${chart_end}"
        self._chart_base(chart, title="Folha total sobre a receita — evolução e meta", x_title="Competência", y_title="Percentual", width=15.5, height=8.5)
        self._line_style(chart, [CHART_GREEN, CHART_TEAL, CHART_META], number_format='0.0"%"')
        ws.add_chart(chart, "A13")

        composition = BarChart()
        composition.type = "col"
        for source_col in (36, 37):
            composition.add_data(Reference(self.ws["CALC"], min_col=source_col, min_row=4, max_row=current), titles_from_data=True)
        composition.set_categories(Reference(self.ws["DIM_PERIODO"], min_col=5, min_row=current, max_row=current))
        for idx, series in enumerate(composition.series):
            source_col = [36, 37][idx]
            series.val.numRef.f = f"'CALC'!${get_column_letter(source_col)}${current}:${get_column_letter(source_col)}${current}"
        self._chart_base(composition, title="Composição da folha sobre a receita", x_title="Competência", y_title="Percentual", width=10.0, height=8.5)
        self._bar_style(composition, [CHART_GREEN, CHART_GOLD], number_format='0.0"%"')
        ws.add_chart(composition, "H13")

        table_row = 31
        self._section(ws, table_row, "HISTÓRICO DOS 12 MESES")
        headers = ["Período", "Folha total", "Receita", "Total %", "Docente %", "Adm. %", "Variação %", "Status"]
        self._headers(ws, table_row + 1, headers)
        out = table_row + 2
        for source_row in range(chart_start, chart_end + 1):
            sources = [30, 33, 34, 35, 36, 37, 39, 42]
            for offset, source_col in enumerate(sources, 1): ws.cell(out, offset).value = f"=CALC!{get_column_letter(source_col)}{source_row}"
            ws.cell(out, 2).number_format = ws.cell(out, 3).number_format = CURRENCY_FMT
            for c in (4, 5, 6): ws.cell(out, c).number_format = PERCENT_POINTS_FMT
            ws.cell(out, 7).number_format = CHANGE_POINTS_FMT
            out += 1
        for row in ws.iter_rows(min_row=table_row + 2, max_row=out - 1, min_col=1, max_col=8):
            for cell in row:
                cell.font = self.body_font
                cell.alignment = Alignment(vertical="center")
                cell.border = Border(bottom=Side(style="hair", color=BORDER))
        self._section(ws, out + 2, "METAS DE REFERÊNCIA")
        notes = ["Folha total ≤ 55% da receita; atenção em 58%.", "Folha docente ≤ 38% da receita.", "Folha administrativa ≤ 17% da receita.", "Variação da folha sobre o mês anterior ≤ 2%.", "Provisões de 13º e férias são apropriadas mensalmente por competência."]
        for i, note in enumerate(notes, out + 3):
            ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=12)
            ws.cell(i, 1).value = f"• {note}"
            ws.cell(i, 1).font = self.static_font
            ws.cell(i, 1).alignment = Alignment(wrap_text=True)
        ws.print_area = f"A1:L{out+8}"
        ws.page_setup.fitToHeight = 1

    # ------------------------------------------------------------------
    # Matrix, quality and instructions
    # ------------------------------------------------------------------
    def _build_matrix(self):
        ws = self.ws["MATRIZ"]
        self._title(ws, "MATRIZ EXECUTIVA — DPE", "Leitura conjunta dos três indicadores no mês de referência.", end_col=12)
        self._context_band(ws, 4)
        self._section(ws, 7, "INDICADORES E MÉTRICAS")
        headers = ["Código", "Indicador", "Métrica", "Valor", "Unidade", "Meta", "Atenção", "Status", "Fonte", "Responsável"]
        self._headers(ws, 8, headers)
        r = 9
        for code, dash in self.dashboards.items():
            if self.context.only_indicator and code != self.context.only_indicator:
                continue
            spec = dash["selected_indicator"]
            for item in dash["selected_metrics"]:
                ws.cell(r, 1).value = code
                ws.cell(r, 2).value = spec["name"]
                ws.cell(r, 3).value = item["label"]
                ws.cell(r, 4).value = item.get("value")
                ws.cell(r, 5).value = item.get("unit")
                ws.cell(r, 6).value = item.get("target_label")
                ws.cell(r, 7).value = item.get("target", {}).get("attention")
                ws.cell(r, 8).value = item.get("status")
                ws.cell(r, 9).value = spec.get("source")
                ws.cell(r, 10).value = spec.get("responsible")
                if item.get("unit") == "R$": ws.cell(r, 4).number_format = CURRENCY_FMT
                elif item.get("unit") == "%": ws.cell(r, 4).number_format = PERCENT_POINTS_FMT
                elif item.get("unit") in {"x", "índice"}: ws.cell(r, 4).number_format = '0.00"x"'
                else: ws.cell(r, 4).number_format = DECIMAL_FMT
                r += 1
        for row in ws.iter_rows(min_row=9, max_row=r-1, min_col=1, max_col=10):
            for cell in row:
                cell.font = self.body_font
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = Border(bottom=Side(style="hair", color=BORDER))
        widths = [12, 34, 31, 15, 12, 16, 14, 18, 45, 28]
        for c, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(c)].width = w
        ws.freeze_panes = "A9"
        ws.auto_filter.ref = f"A8:J{max(9,r-1)}"

    def _build_quality(self):
        ws = self.ws["QUALIDADE E GOVERNANÇA"]
        self._title(ws, "QUALIDADE E GOVERNANÇA — DPE", "Cobertura, validação, fontes e pendências antes da reunião de gestão.", end_col=12)
        self._section(ws, 4, "SÍNTESE DA BASE")
        total = len(self.measurements)
        validated = sum(1 for row in self.measurements if row.get("validated"))
        periods = len({row.get("period") for row in self.measurements})
        open_actions = sum(1 for row in self.actions if row.get("status") not in {"Concluído", "Concluido", "Cancelado"})
        cards = [("Lançamentos", total, INTEGER_FMT, PALE_BLUE), ("Validados", validated, INTEGER_FMT, PALE_GREEN), ("Cobertura de validação", (validated / total * 100 if total else None), PERCENT_POINTS_FMT, PALE_TEAL), ("Meses com dados", periods, INTEGER_FMT, PALE_GOLD), ("Planos em aberto", open_actions, INTEGER_FMT, RED_LIGHT), ("Indicadores ativos", 3, INTEGER_FMT, PALE_GREEN)]
        for idx, (label, value, fmt, fill) in enumerate(cards):
            start = idx * 2 + 1
            self._kpi_card(ws, start_col=start, end_col=start+1, label_row=5, value_row=6, label=label, formula=value, number_format=fmt, fill=fill)
        self._section(ws, 10, "CHECKLIST DO FECHAMENTO MENSAL")
        checklist = [
            "Receita líquida do DPE-01, DPE-02 e DPE-03 conciliada para a mesma competência.",
            "Custos indiretos rateados pelo critério vigente e documentado.",
            "Provisões de férias e 13º apropriadas mensalmente.",
            "Imobilizado separado das despesas que compõem o índice de cobertura.",
            "Registros validados pelo diretor da DPE até o 7º dia útil.",
            "Indicadores fora da meta vinculados a plano de ação com responsável e prazo.",
        ]
        for r, text in enumerate(checklist, 11):
            ws.cell(r, 1).value = "☐"
            ws.cell(r, 2).value = text
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=12)
            ws.cell(r, 1).font = Font(name="Arial", size=12, color=GREEN)
            ws.cell(r, 2).font = self.body_font
            ws.cell(r, 2).alignment = Alignment(wrap_text=True)
            ws.row_dimensions[r].height = 26
        self._section(ws, 19, "ROTINA DE GESTÃO")
        routine = [
            "1º ao 5º dia útil: alimentação dos dados do mês anterior.",
            "Até o 7º dia útil: conferência e validação pelo diretor.",
            "Até o 9º dia útil: consolidação técnica e aplicação das metas.",
            "Até o 12º dia útil: leitura crítica e registro de ações.",
            "Até o 15º dia útil: reunião decisória de gestão.",
        ]
        for r, text in enumerate(routine, 20):
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=12)
            ws.cell(r, 1).value = text
            ws.cell(r, 1).font = self.static_font
            ws.cell(r, 1).alignment = Alignment(wrap_text=True)

    def _build_readme(self):
        ws = self.ws["LEIA-ME"]
        self._title(ws, "DATA UNIVC · PAINEL DPE", "Diretoria de Planejamento Econômico e Oferta · versão padronizada com DTNH/DCS", end_col=10)
        self._section(ws, 4, "COMO O ARQUIVO ESTÁ ORGANIZADO", end_col=10)
        sections = [
            ("PAINEL", "Visão executiva dos três indicadores, com cartões, séries históricas e metas."),
            ("DPE-01 MARGEM", "Página própria de margem, carga horária e comparação por curso."),
            ("DPE-02 COBERTURA", "Página própria de cobertura, acumulado de 12 meses e composição da despesa."),
            ("DPE-03 FOLHA", "Página própria da folha total, docente, administrativa e leitura sobre receita média trimestral."),
            ("BASE DPE-01 / 02 / 03", "Bases independentes para edição, exportação e posterior importação no sistema."),
            ("METAS / PLANO_DE_ACAO", "Vigências e tratamento gerencial dos desvios."),
            ("MATRIZ", "Leitura conjunta das métricas do mês de referência."),
        ]
        self._headers(ws, 5, ["Aba", "Finalidade"])
        for r, (name, purpose) in enumerate(sections, 6):
            ws.cell(r, 1).value = name
            ws.cell(r, 2).value = purpose
            ws.cell(r, 1).font = Font(name="Arial", size=9, bold=True, color=GREEN_DARK)
            ws.cell(r, 2).font = self.body_font
            ws.cell(r, 1).border = ws.cell(r, 2).border = Border(bottom=Side(style="hair", color=BORDER))
            ws.cell(r, 2).alignment = Alignment(wrap_text=True)
            ws.row_dimensions[r].height = 27
        ws.column_dimensions["A"].width = 26
        ws.column_dimensions["B"].width = 88
        self._section(ws, 15, "IMPORTAÇÃO", end_col=10)
        text = (
            "Para importar, mantenha os nomes das abas e dos cabeçalhos. Preencha as colunas de entrada das abas BASE DPE-01, "
            "BASE DPE-02 e BASE DPE-03. As colunas calculadas podem ser mantidas com fórmula ou deixadas vazias: o backend "
            "recalcula os indicadores pelos componentes oficiais. Linhas sem período são ignoradas. Erros são devolvidos por aba e linha; "
            "nenhuma linha é gravada se o arquivo possuir inconsistências."
        )
        ws.merge_cells("A16:J20")
        ws["A16"] = text
        ws["A16"].font = self.body_font
        ws["A16"].alignment = Alignment(wrap_text=True, vertical="top")
        self._paint_range(ws, 16, 20, 1, 10, fill=PALE_BLUE)
        self._section(ws, 22, "REGRAS DOS INDICADORES", end_col=10)
        rules = [
            "DPE-01 é mensal e combina margem com carga horária; a leitura por curso usa meta mínima de 10%.",
            "DPE-02 nunca deve ser interpretado apenas pelo mês: o acumulado móvel de 12 meses aparece ao lado.",
            "DPE-03 compara a folha do mês à receita do mês e também à receita média dos últimos três meses.",
            "Nenhum registro é cortado silenciosamente. Se o limite físico do Excel for atingido, a exportação falha explicitamente.",
        ]
        for r, item in enumerate(rules, 23):
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=10)
            ws.cell(r, 1).value = f"• {item}"
            ws.cell(r, 1).font = self.static_font
            ws.cell(r, 1).alignment = Alignment(wrap_text=True)
            ws.row_dimensions[r].height = 28

    def _finish(self):
        self.wb.active = self.wb.sheetnames.index("PAINEL")
        self.wb.calculation.fullCalcOnLoad = True
        self.wb.calculation.forceFullCalc = True
        self.wb.calculation.calcMode = "auto"
        tab_colors = {
            "PAINEL": GREEN_DARK,
            "DPE-01 MARGEM": GREEN,
            "DPE-02 COBERTURA": GREEN,
            "DPE-03 FOLHA": GREEN,
            "BASE DPE-01": LINKED_GREEN,
            "BASE DPE-02": LINKED_GREEN,
            "BASE DPE-03": LINKED_GREEN,
            "METAS": CHART_GOLD,
            "PLANO_DE_ACAO": CHART_RED,
            "QUALIDADE E GOVERNANÇA": CHART_BLUE,
        }
        for name, color in tab_colors.items():
            if name in self.ws:
                self.ws[name].sheet_properties.tabColor = color
        for ws in self.wb.worksheets:
            ws.sheet_view.showGridLines = False
            ws.sheet_properties.pageSetUpPr.fitToPage = True



def build_dpe_import_template(indicator_code: str) -> BytesIO:
    """Modelo de importação simples: uma única aba DADOS, sem painel auxiliar."""
    code = indicator_spec("DPE", indicator_code)["code"]
    layouts = {
        "DPE-01": [
            "Período", "Curso", "Receita líquida", "Custo docente",
            "Custo de coordenação", "Outros custos diretos", "Rateio indireto",
            "Horas-aula semanais", "Nº de docentes", "Observações", "Validado",
        ],
        "DPE-02": [
            "Período", "Centro de custo", "Natureza da despesa", "Receita líquida",
            "Despesa de pessoal", "Despesa operacional", "Despesa administrativa", "Despesa financeira",
            "Investimento em imobilizado", "Observações", "Validado",
        ],
        "DPE-03": [
            "Período", "Categoria", "Natureza", "Folha docente completa", "Folha administrativa completa",
            "Salários brutos", "Encargos", "Provisões", "Receita líquida do mês", "Observações", "Validado",
        ],
    }
    wb = Workbook(); ws = wb.active; ws.title = "DADOS"
    headers = layouts[code]
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    c = ws.cell(1,1); c.value = f"{code} · MODELO DE IMPORTAÇÃO"; c.font = Font(name="Arial", size=16, bold=True, color=GREEN_DARK); c.alignment=Alignment(vertical="center")
    ws.row_dimensions[1].height = 28
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    ws.cell(2,1).value = "Preencha uma linha por registro. Não altere os cabeçalhos. Período no formato AAAA-MM."
    ws.cell(2,1).font = Font(name="Arial", size=9, color=GRAY_TEXT); ws.cell(2,1).alignment=Alignment(wrap_text=True)
    for col, header in enumerate(headers,1):
        cell=ws.cell(4,col); cell.value=header; cell.font=Font(name="Arial",size=9,bold=True,color=WHITE); cell.fill=PatternFill("solid",fgColor=GREEN_DARK); cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width=max(13,min(28,len(header)+4))
    for row in range(5,65):
        for col in range(1,len(headers)+1):
            cell=ws.cell(row,col); cell.font=Font(name="Arial",size=9,color=INPUT_BLUE); cell.alignment=Alignment(vertical="center",wrap_text=True); cell.border=Border(bottom=Side(style="hair",color=BORDER))
        ws.row_dimensions[row].height=20
    valid_col=headers.index("Validado")+1
    dv=DataValidation(type="list",formula1='"Sim,Não"',allow_blank=True); ws.add_data_validation(dv); dv.add(f"{get_column_letter(valid_col)}5:{get_column_letter(valid_col)}64")
    ws.freeze_panes="A5"; ws.auto_filter.ref=f"A4:{get_column_letter(len(headers))}64"; ws.sheet_view.showGridLines=False; ws.sheet_view.zoomScale=90
    ws.page_setup.orientation="landscape"; ws.page_setup.fitToWidth=1; ws.sheet_properties.pageSetUpPr.fitToPage=True
    out=BytesIO(); wb.save(out); out.seek(0); return out

def build_dpe_workbook(
    measurements: Iterable[dict[str, Any]],
    targets: Iterable[dict[str, Any]] = (),
    actions: Iterable[dict[str, Any]] = (),
    *,
    reference: str | None = None,
    comparison: str | None = None,
    only_indicator: str | None = None,
    template_mode: bool = False,
) -> BytesIO:
    builder = DPEWorkbookBuilder(
        measurements,
        targets,
        actions,
        reference=reference,
        comparison=comparison,
        only_indicator=only_indicator,
        template_mode=template_mode,
    )
    workbook = builder.build()
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
