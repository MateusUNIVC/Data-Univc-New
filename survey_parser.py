from __future__ import annotations

import io
import re
import unicodedata
import warnings
from collections import Counter, OrderedDict
from datetime import datetime
from pathlib import Path
from typing import BinaryIO

from openpyxl import load_workbook

from survey_models import OptionAggregate, ParsedQuestion, ParsedWorkbook


SPACES_RE = re.compile(r"\s+")
SEMESTER_RE = re.compile(r"\b(20\d{2})[./-]?([12])\b")
PERIOD_RE = re.compile(
    r"Per[ií]odo\s+de\s+(\d{1,2}/\d{1,2}/\d{2,4})\s+at[eé]\s+(\d{1,2}/\d{1,2}/\d{2,4})",
    flags=re.IGNORECASE,
)


def clean_text(value: object | None) -> str:
    if value is None:
        return ""
    return SPACES_RE.sub(" ", str(value).replace("\xa0", " ")).strip()


def normalize_key(value: str) -> str:
    value = clean_text(value).casefold()
    value = unicodedata.normalize("NFKD", value)
    value = "".join(ch for ch in value if not unicodedata.combining(ch))
    value = re.sub(r"[^a-z0-9]+", " ", value)
    return SPACES_RE.sub(" ", value).strip()


def parse_percentage(value: object | None) -> float | None:
    if value is None or clean_text(value) == "":
        return None
    if isinstance(value, (int, float)):
        number = float(value)
        return number * 100.0 if 0 <= number <= 1 else number
    text = clean_text(value).replace("%", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def parse_count(value: object | None) -> int | None:
    if value is None or clean_text(value) == "":
        return None
    try:
        return int(round(float(str(value).replace(",", "."))))
    except (TypeError, ValueError):
        return None


def parse_numeric_label(label: str) -> float | None:
    text = clean_text(label).replace(",", ".")
    if not re.fullmatch(r"[-+]?\d+(?:\.\d+)?", text):
        return None
    try:
        return float(text)
    except ValueError:
        return None


def detect_metric_type(options: list[OptionAggregate]) -> tuple[str, bool]:
    if not options:
        return "categorical", False

    numeric_values = [parse_numeric_label(option.label) for option in options]
    if all(value is not None for value in numeric_values):
        values = sorted({float(value) for value in numeric_values if value is not None})
        nps_candidate = values == [float(i) for i in range(11)]
        for option, value in zip(options, numeric_values):
            option.numeric_value = value
        return "numeric", nps_candidate

    return "categorical", False


def infer_modality(course_name: str) -> str:
    key = normalize_key(course_name)
    if "semipresencial" in key:
        return "SEMIPRESENCIAL"
    if "ead" in key or "a distancia" in key:
        return "EAD"
    return "PRESENCIAL"


def parse_date(value: str) -> str | None:
    for fmt in ("%d/%m/%y", "%d/%m/%Y"):
        try:
            return datetime.strptime(value, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def infer_semester(title: str | None, start_date: str | None) -> str | None:
    if title:
        match = SEMESTER_RE.search(title)
        if match:
            return f"{match.group(1)}.{match.group(2)}"
    if start_date:
        date = datetime.fromisoformat(start_date).date()
        return f"{date.year}.{'1' if date.month <= 6 else '2'}"
    return None


def canonical_course_name(course_name: str) -> str:
    return clean_text(course_name)


def _load_workbook(source: bytes | BinaryIO | str | Path):
    warnings.filterwarnings(
        "ignore",
        message="Workbook contains no default style",
        category=UserWarning,
    )
    if isinstance(source, bytes):
        source = io.BytesIO(source)

    # O SEI pode gravar <dimension ref="A1"> em relatórios com milhares de
    # linhas. O modo normal não confia nessa dimensão para materializar as
    # células e, por isso, é deliberadamente utilizado aqui.
    return load_workbook(source, data_only=True, read_only=False)


def _global_metadata(ws, first_course_row: int) -> dict[str, str | None]:
    survey_title: str | None = None
    period_start: str | None = None
    period_end: str | None = None
    generated_at: str | None = None

    for row_index in range(1, min(ws.max_row, max(first_course_row, 60)) + 1):
        col2 = clean_text(ws.cell(row_index, 2).value)
        col14 = clean_text(ws.cell(row_index, 14).value)

        if not generated_at and col14 and re.fullmatch(
            r"\d{1,2}/\d{1,2}/\d{4}\s+\d{1,2}[.:]\d{2}[.:]\d{2}", col14
        ):
            generated_at = col14

        if col2 and "periodo de" in normalize_key(col2):
            match = PERIOD_RE.search(col2)
            if match:
                period_start = parse_date(match.group(1))
                period_end = parse_date(match.group(2))
            continue

        # O título da aplicação costuma ficar na coluna B imediatamente antes
        # da linha "Período de ...". Evita o cabeçalho genérico "Avaliação
        # Institucional", que normalmente aparece no centro da página.
        if col2 and normalize_key(col2) != "avaliacao institucional":
            survey_title = survey_title or col2

    return {
        "survey_title": survey_title,
        "period_start": period_start,
        "period_end": period_end,
        "generated_at": generated_at,
    }


def _course_starts(ws) -> list[int]:
    starts: list[int] = []
    for row_index in range(1, ws.max_row + 1):
        label = normalize_key(clean_text(ws.cell(row_index, 11).value))
        course = clean_text(ws.cell(row_index, 12).value)
        if label == "curso" and course:
            starts.append(row_index)
    return starts


def _merge_option(question: ParsedQuestion, option: OptionAggregate) -> None:
    option_key = normalize_key(option.label)
    existing = next(
        (item for item in question.options if normalize_key(item.label) == option_key),
        None,
    )
    if existing is None:
        question.options.append(option)
        return

    # Em quebras de página o SEI pode repetir a pergunta. Se uma alternativa
    # também vier repetida com exatamente os mesmos valores, tratamos como uma
    # repetição visual, não como novas respostas.
    same_count = existing.count == option.count
    same_pct = (
        existing.source_percentage == option.source_percentage
        or (
            existing.source_percentage is not None
            and option.source_percentage is not None
            and abs(existing.source_percentage - option.source_percentage) < 0.0001
        )
    )
    if same_count and same_pct:
        return

    existing.count += option.count
    if option.source_percentage is not None:
        if existing.source_percentage is None:
            existing.source_percentage = option.source_percentage
        else:
            existing.source_percentage += option.source_percentage


def _block_source_path(base: str, index: int, total: int) -> str:
    base = base or "relatorio.xlsx"
    if total == 1:
        return base
    return f"{base}::report:{index + 1:03d}"


def _parse_course_block(
    ws,
    *,
    start_row: int,
    end_row: int,
    index: int,
    total_blocks: int,
    source_path: str,
    global_meta: dict[str, str | None],
) -> ParsedWorkbook:
    unit_name: str | None = None
    course_name: str | None = None
    questionnaire_name: str | None = None

    questions: "OrderedDict[str, ParsedQuestion]" = OrderedDict()
    current_key: str | None = None

    for row_index in range(start_row, end_row):
        col4 = clean_text(ws.cell(row_index, 4).value)
        col5 = clean_text(ws.cell(row_index, 5).value)
        col7 = clean_text(ws.cell(row_index, 7).value)
        col9 = clean_text(ws.cell(row_index, 9).value)
        col11 = clean_text(ws.cell(row_index, 11).value)
        col12 = clean_text(ws.cell(row_index, 12).value)
        pct_raw = ws.cell(row_index, 15).value
        count_raw = ws.cell(row_index, 18).value

        if normalize_key(col5) == "unidade ensino":
            unit_name = col7 or unit_name

        if normalize_key(col11) == "curso" and col12:
            course_name = col12

        if normalize_key(col5) == "questionario":
            questionnaire_name = col7 or questionnaire_name

        if col4:
            current_key = normalize_key(col4)
            if current_key not in questions:
                questions[current_key] = ParsedQuestion(
                    text=col4,
                    normalized_text=current_key,
                    position=len(questions) + 1,
                )

        if current_key and col9:
            count = parse_count(count_raw)
            pct = parse_percentage(pct_raw)
            if count is None and pct is None:
                # Questões abertas aparecem no relatório como uma resposta por
                # linha na coluna de alternativas, sem percentual/quantidade.
                # Preserve cada ocorrência: respostas idênticas podem ter sido
                # dadas por pessoas diferentes e não devem ser deduplicadas.
                questions[current_key].raw_responses.append(col9)
                continue

            _merge_option(
                questions[current_key],
                OptionAggregate(
                    label=col9,
                    count=count or 0,
                    source_percentage=pct,
                ),
            )

    if not course_name:
        raise ValueError(
            f"Não foi possível identificar o curso no bloco {index + 1} de "
            f"{source_path or 'XLSX'}"
        )

    # O HAR informa a quantidade de perguntas selecionadas. Alguns questionários
    # possuem perguntas abertas/sem agregados numéricos no XLSX. Elas ainda são
    # perguntas válidas e precisam permanecer no modelo, mesmo sem opções para
    # gráfico.
    parsed_questions = list(questions.values())
    for position, question in enumerate(parsed_questions, start=1):
        question.position = position
        metric_type, nps_candidate = detect_metric_type(question.options)
        if not question.options and question.raw_responses:
            metric_type = "text"
            nps_candidate = False
        question.metric_type = metric_type
        question.nps_candidate = nps_candidate

    # O relatório não traz sempre um campo explícito de respondentes. Algumas
    # perguntas podem ter não-resposta e outras podem ser de múltipla escolha.
    # O total modal (o valor que se repete na maior quantidade de perguntas)
    # é mais estável que usar o máximo ou uma pergunta específica.
    totals = [sum(option.count for option in question.options) for question in parsed_questions if question.options]
    if totals:
        respondent_count = Counter(totals).most_common(1)[0][0]
    else:
        raw_totals = [len(question.raw_responses) for question in parsed_questions if question.raw_responses]
        respondent_count = Counter(raw_totals).most_common(1)[0][0] if raw_totals else 0

    survey_title = global_meta.get("survey_title")
    period_start = global_meta.get("period_start")

    return ParsedWorkbook(
        source_path=_block_source_path(source_path, index, total_blocks),
        unit_name=unit_name,
        course_name=canonical_course_name(course_name),
        modality=infer_modality(course_name),
        survey_title=survey_title,
        questionnaire_name=questionnaire_name,
        period_start=period_start,
        period_end=global_meta.get("period_end"),
        semester_suggested=infer_semester(survey_title, period_start),
        generated_at=global_meta.get("generated_at"),
        respondent_count=respondent_count,
        questions=parsed_questions,
    )


def parse_workbooks(
    source: bytes | BinaryIO | str | Path,
    *,
    source_path: str = "",
) -> list[ParsedWorkbook]:
    """Interpreta um XLSX do SEI em um ou mais relatórios de curso.

    O SEI pode devolver:
    - um XLSX por curso (comum em ZIPs grandes), ou
    - um único XLSX contendo vários cursos concatenados verticalmente.

    A saída é sempre uma lista de blocos normalizados, um por curso/contexto.
    """

    wb = _load_workbook(source)
    ws = wb.active
    starts = _course_starts(ws)
    if not starts:
        raise ValueError(f"Não foi possível identificar nenhum bloco 'Curso:' em {source_path or 'XLSX'}")

    global_meta = _global_metadata(ws, starts[0])
    ends = starts[1:] + [ws.max_row + 1]
    total = len(starts)

    return [
        _parse_course_block(
            ws,
            start_row=start,
            end_row=end,
            index=index,
            total_blocks=total,
            source_path=source_path,
            global_meta=global_meta,
        )
        for index, (start, end) in enumerate(zip(starts, ends))
    ]


def parse_workbook(
    source: bytes | BinaryIO | str | Path,
    *,
    source_path: str = "",
) -> ParsedWorkbook:
    """Compatibilidade com a V0.1 para XLSX que contém somente um bloco."""
    parsed = parse_workbooks(source, source_path=source_path)
    if len(parsed) != 1:
        raise ValueError(
            f"{source_path or 'XLSX'} contém {len(parsed)} relatórios. "
            "Use parse_workbooks() para arquivos consolidados."
        )
    return parsed[0]
