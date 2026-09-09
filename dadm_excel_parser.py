from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from management_catalog import indicator_spec, validate_period

BASE_HEADER_ROW = 4
BASE_DATA_START = 5


class DADMExcelImportError(ValueError):
    def __init__(self, message: str, errors: list[dict[str, Any]] | None = None):
        super().__init__(message)
        self.errors = errors or []


@dataclass(frozen=True)
class SheetLayout:
    indicator_code: str
    sheet_name: str
    dimensions: dict[str, str]
    fields: dict[str, str]


LAYOUTS: dict[str, SheetLayout] = {
    "DADM-01": SheetLayout(
        "DADM-01",
        "BASE DADM-01",
        {
            "channel": "Canal",
            "request_type": "Tipo de solicitação",
            "week": "Semana",
        },
        {
            "requests_received": "Solicitações recebidas",
            "requests_within_sla": "Solicitações no prazo",
            "concluded": "Solicitações concluídas",
            "reopened": "Solicitações reabertas",
            "open_balance": "Saldo em aberto",
            "avg_first_response_hours": "Tempo médio da 1ª resposta (h úteis)",
            "avg_resolution_hours": "Tempo médio da resolução (h úteis)",
        },
    ),
    "DADM-02": SheetLayout(
        "DADM-02",
        "BASE DADM-02",
        {
            "channel": "Canal",
            "request_type": "Tipo de solicitação",
            "resolution_band": "Faixa de resolução",
        },
        {
            "eligible_interactions": "Atendimentos elegíveis",
            "respondents": "Respondentes",
            "satisfied": "Respostas 4 e 5",
            "dissatisfied": "Respostas 1 e 2",
        },
    ),
}


def _norm(value: Any) -> str:
    raw = str(value or "").replace("º", "").replace("ª", "").replace("°", "")
    text = unicodedata.normalize("NFKD", raw)
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = text.casefold().strip()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def _bool(value: Any) -> bool:
    normalized = _norm(value)
    if normalized in {"", "sim", "s", "1", "true", "verdadeiro", "validado"}:
        return True
    if normalized in {"nao", "n", "0", "false", "falso", "pendente"}:
        return False
    raise ValueError("Use Sim ou Não na coluna Validado.")


def _number(value: Any, *, required: bool, integer: bool, label: str) -> float | int | None:
    if value in (None, ""):
        if required:
            raise ValueError(f"Informe {label}.")
        return None
    if isinstance(value, str) and value.startswith("="):
        raise ValueError(f"{label} precisa ser um valor, não uma fórmula.")
    if isinstance(value, str):
        text = value.strip().replace(" ", "")
        if "," in text and "." in text:
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", ".")
        value = text
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{label} possui valor numérico inválido.") from exc
    if number != number or abs(number) == float("inf"):
        raise ValueError(f"{label} possui valor numérico inválido.")
    if integer:
        if number != int(number):
            raise ValueError(f"{label} precisa ser um número inteiro.")
        return int(number)
    return number


def _headers(ws) -> dict[str, int]:
    result: dict[str, int] = {}
    for cell in ws[BASE_HEADER_ROW]:
        normalized = _norm(cell.value)
        if normalized:
            result[normalized] = cell.column
    return result


def _required_column(headers: dict[str, int], name: str, *, sheet: str) -> int:
    normalized = _norm(name)
    if normalized not in headers:
        raise DADMExcelImportError(
            f'A aba "{sheet}" não possui a coluna obrigatória "{name}".',
            [{"sheet": sheet, "row": BASE_HEADER_ROW, "field": name, "error": "Cabeçalho ausente."}],
        )
    return headers[normalized]


def parse_dadm_workbook(path: str | Path, *, indicator_code: str | None = None) -> dict[str, Any]:
    """Parse DADM import sheets without trusting calculated KPI columns.

    Only raw components are imported. Derived metrics are always recalculated by
    the backend, preserving the same auditability contract used by the other
    redesigned directorates.
    """
    workbook = load_workbook(path, data_only=False, read_only=False)
    codes: Iterable[str]
    if indicator_code:
        code = indicator_spec("DADM", indicator_code)["code"]
        codes = [code]
    else:
        codes = LAYOUTS.keys()

    payloads: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    sheets_read: list[str] = []

    for code in codes:
        layout = LAYOUTS[code]
        sheet_name = "DADOS" if indicator_code and "DADOS" in workbook.sheetnames else layout.sheet_name
        if sheet_name not in workbook.sheetnames:
            if indicator_code:
                errors.append({"sheet": sheet_name, "row": None, "field": None, "error": "Aba não encontrada."})
            continue

        ws = workbook[sheet_name]
        sheets_read.append(sheet_name)
        headers = _headers(ws)
        try:
            period_col = _required_column(headers, "Período", sheet=sheet_name)
            # Canal é a dimensão operacional mínima dos dois indicadores. As
            # demais dimensões são opcionais e refinam a análise.
            channel_col = _required_column(headers, "Canal", sheet=sheet_name)
            dimension_cols = {
                key: (channel_col if key == "channel" else headers.get(_norm(label)))
                for key, label in layout.dimensions.items()
            }
            field_cols = {
                key: _required_column(headers, label, sheet=sheet_name)
                for key, label in layout.fields.items()
            }
            notes_col = headers.get(_norm("Observações"))
            source_col = headers.get(_norm("Fonte / referência"))
            validated_col = headers.get(_norm("Validado"))
        except DADMExcelImportError as exc:
            errors.extend(exc.errors)
            continue

        spec = indicator_spec("DADM", code)
        required = {field["key"]: bool(field.get("required")) for field in spec.get("fields") or []}
        integer = {field["key"]: field.get("type") == "integer" for field in spec.get("fields") or []}
        labels = {field["key"]: field["label"] for field in spec.get("fields") or []}

        for row_index in range(BASE_DATA_START, ws.max_row + 1):
            raw_period = ws.cell(row_index, period_col).value
            has_any_component = any(ws.cell(row_index, col).value not in (None, "") for col in field_cols.values())
            has_any_dimension = any(col is not None and ws.cell(row_index, col).value not in (None, "") for col in dimension_cols.values())
            if raw_period in (None, "") and not has_any_component and not has_any_dimension:
                continue

            row_errors: list[dict[str, Any]] = []
            try:
                period = validate_period("DADM", raw_period)
            except Exception as exc:
                row_errors.append({"sheet": sheet_name, "row": row_index, "field": "Período", "error": str(exc)})
                period = None

            dimensions = {
                key: str(ws.cell(row_index, col).value or "").strip()
                for key, col in dimension_cols.items()
                if col is not None and str(ws.cell(row_index, col).value or "").strip()
            }
            if not dimensions.get("channel"):
                row_errors.append({"sheet": sheet_name, "row": row_index, "field": "Canal", "error": "O canal é obrigatório."})

            values: dict[str, float | int | None] = {}
            for key, col in field_cols.items():
                try:
                    values[key] = _number(
                        ws.cell(row_index, col).value,
                        required=required.get(key, False),
                        integer=integer.get(key, False),
                        label=labels.get(key, key),
                    )
                except ValueError as exc:
                    row_errors.append({"sheet": sheet_name, "row": row_index, "field": labels.get(key, key), "error": str(exc)})

            try:
                validated = _bool(ws.cell(row_index, validated_col).value) if validated_col else True
            except ValueError as exc:
                row_errors.append({"sheet": sheet_name, "row": row_index, "field": "Validado", "error": str(exc)})
                validated = True

            if row_errors:
                errors.extend(row_errors)
                continue

            payloads.append({
                "indicator_code": code,
                "period": period,
                "dimensions": dimensions,
                "values": values,
                "notes": str(ws.cell(row_index, notes_col).value or "").strip() if notes_col else None,
                "source_reference": str(ws.cell(row_index, source_col).value or "").strip() if source_col else None,
                "validated": validated,
                "_source": {"sheet": sheet_name, "row": row_index},
            })

    if not sheets_read:
        expected = ", ".join(LAYOUTS[code].sheet_name for code in codes)
        raise DADMExcelImportError(
            f"Nenhuma base DADM foi encontrada. Abas esperadas: {expected}.",
            [{"sheet": None, "row": None, "field": None, "error": "Estrutura DADM não encontrada."}],
        )
    if errors:
        raise DADMExcelImportError(
            f"A importação possui {len(errors)} inconsistência(s). Nenhuma linha foi gravada.",
            errors,
        )
    if not payloads:
        raise DADMExcelImportError(
            "A planilha não contém linhas preenchidas para importação.",
            [{"sheet": ", ".join(sheets_read), "row": None, "field": None, "error": "Nenhuma linha preenchida."}],
        )
    return {"payloads": payloads, "sheets": sheets_read, "total": len(payloads)}
