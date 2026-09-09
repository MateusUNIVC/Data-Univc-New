from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

from management_catalog import indicator_spec
from management_excel_builder import build_management_workbook

GREEN_DARK = "006B45"
GREEN = "0E8058"
WHITE = "FFFFFF"
GRAY_TEXT = "53645D"
INPUT_BLUE = "185A78"
BORDER = "D9E4DF"


TEMPLATE_HEADERS = {
    "DADM-01": [
        "Período", "Canal", "Tipo de solicitação", "Semana",
        "Solicitações recebidas", "Solicitações no prazo", "Solicitações concluídas",
        "Solicitações reabertas", "Saldo em aberto", "Tempo médio da 1ª resposta (h úteis)",
        "Tempo médio da resolução (h úteis)", "Fonte / referência", "Observações", "Validado",
    ],
    "DADM-02": [
        "Período", "Canal", "Tipo de solicitação", "Faixa de resolução",
        "Atendimentos elegíveis", "Respondentes", "Respostas 4 e 5", "Respostas 1 e 2",
        "Fonte / referência", "Observações", "Validado",
    ],
}


def build_dadm_import_template(indicator_code: str) -> BytesIO:
    code = indicator_spec("DADM", indicator_code)["code"]
    headers = TEMPLATE_HEADERS[code]
    spec = indicator_spec("DADM", code)

    wb = Workbook()
    ws = wb.active
    ws.title = "DADOS"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    cell = ws.cell(1, 1)
    cell.value = f"{code} · {spec['name'].upper()}"
    cell.font = Font(name="Arial", size=16, bold=True, color=GREEN_DARK)
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 30

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    ws.cell(2, 1).value = (
        "Preencha uma linha por canal/recorte. Não altere os cabeçalhos. "
        "Período no formato AAAA-MM. O sistema recalcula todos os indicadores no servidor."
    )
    ws.cell(2, 1).font = Font(name="Arial", size=9, color=GRAY_TEXT)
    ws.cell(2, 1).alignment = Alignment(wrap_text=True)

    for col, header in enumerate(headers, 1):
        c = ws.cell(4, col)
        c.value = header
        c.font = Font(name="Arial", size=9, bold=True, color=WHITE)
        c.fill = PatternFill("solid", fgColor=GREEN_DARK)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = max(13, min(34, len(header) + 4))

    for row in range(5, 105):
        for col in range(1, len(headers) + 1):
            c = ws.cell(row, col)
            c.font = Font(name="Arial", size=9, color=INPUT_BLUE)
            c.alignment = Alignment(vertical="center", wrap_text=True)
            c.border = Border(bottom=Side(style="hair", color=BORDER))
        ws.row_dimensions[row].height = 20

    valid_col = headers.index("Validado") + 1
    dv = DataValidation(type="list", formula1='"Sim,Não"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{get_column_letter(valid_col)}5:{get_column_letter(valid_col)}104")

    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(headers))}104"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def build_dadm_workbook(
    measurements,
    targets=(),
    actions=(),
    *,
    reference: str | None = None,
    comparison: str | None = None,
    only_indicator: str | None = None,
) -> BytesIO:
    """Dedicated DADM entry point over the shared audited workbook engine."""
    return build_management_workbook(
        "DADM",
        measurements,
        targets,
        actions,
        reference=reference,
        comparison=comparison,
        only_indicator=only_indicator,
    )
